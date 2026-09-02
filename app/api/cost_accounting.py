from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from urllib.parse import quote

from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    Request,
    UploadFile,
    status,
)
from fastapi.encoders import jsonable_encoder
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, field_validator, model_validator
from sqlalchemy import func, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.api.deps import (
    PermissionChecker,
    get_db,
    has_unrestricted_customer_access,
)
from app.core.time_contract import beijing_now_naive, beijing_today
from app.models.customer import Customer
from app.models.finance import FinanceIdempotencyRecord, Statement, StatementItem
from app.models.finance_cost import FinanceCostCenter, FinanceCostPoolEntry
from app.models.finance_payable import FinancePayable
from app.models.user import User
from app.services.audit_log import append_audit_event
from app.services.material_cost_lineage import material_cost_coverage_report


router = APIRouter()
MONEY = Decimal("0.00")
MAX_MONEY = Decimal("999999999999.99")
MONTH_PATTERN = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 500

COST_CATEGORIES: dict[str, dict[str, str]] = {
    "outsourcing": {
        "label": "外协加工",
        "accounting_class": "manufacturing",
        "default_center": "PROD",
    },
    "inbound_freight": {
        "label": "来料运费",
        "accounting_class": "manufacturing",
        "default_center": "PROD",
    },
    "production_wages": {
        "label": "生产工资/社保",
        "accounting_class": "manufacturing",
        "default_center": "PROD",
    },
    "factory_utilities": {
        "label": "车间水电",
        "accounting_class": "manufacturing",
        "default_center": "PROD",
    },
    "factory_rent": {
        "label": "车间租金",
        "accounting_class": "manufacturing",
        "default_center": "PROD",
    },
    "maintenance": {
        "label": "设备维修",
        "accounting_class": "manufacturing",
        "default_center": "PROD",
    },
    "delivery_freight": {
        "label": "送货运费",
        "accounting_class": "selling",
        "default_center": "WH_DELIVERY",
    },
    "driver_wages": {
        "label": "司机工资",
        "accounting_class": "selling",
        "default_center": "WH_DELIVERY",
    },
    "sales_expense": {
        "label": "销售费用",
        "accounting_class": "selling",
        "default_center": "SALES",
    },
    "administrative_wages": {
        "label": "管理工资",
        "accounting_class": "administrative",
        "default_center": "ADMIN",
    },
    "administrative_expense": {
        "label": "管理费用",
        "accounting_class": "administrative",
        "default_center": "ADMIN",
    },
    "finance_expense": {
        "label": "财务费用",
        "accounting_class": "finance",
        "default_center": "FINANCE",
    },
    "finance_wages": {
        "label": "财务工资",
        "accounting_class": "finance",
        "default_center": "FINANCE",
    },
    "tax_fee": {
        "label": "税费",
        "accounting_class": "administrative",
        "default_center": "ADMIN",
    },
    "other": {
        "label": "其他支出",
        "accounting_class": "excluded",
        "default_center": "UNALLOCATED",
    },
}

CENTER_TYPE_LABELS = {
    "production": "生产",
    "warehouse_delivery": "仓储配送",
    "sales": "销售",
    "administration": "管理",
    "finance": "财务",
    "unallocated": "待分类",
}

ACCOUNTING_CLASS_LABELS = {
    "manufacturing": "制造成本",
    "selling": "销售费用",
    "administrative": "管理费用",
    "finance": "财务费用",
    "excluded": "待分类/不参与结转",
}

ALLOCATION_BASIS_LABELS = {
    "unallocated": "暂不分摊",
    "direct": "直接归集",
    "completion_area": "按完工面积",
    "production_quantity": "按完工数量",
    "machine_hours": "按机器工时",
    "delivery_quantity": "按送货数量",
    "manual": "人工确认",
}

STATUS_LABELS = {"draft": "草稿", "confirmed": "已确认", "voided": "已作废"}

PAYABLE_CATEGORY_LABELS = {
    "material": "纸板材料",
    "outsourcing": "外协加工",
    "freight": "运费",
    "utilities": "水电",
    "rent": "租金",
    "wages": "工资",
    "maintenance": "维修",
    "tax_fee": "税费",
    "other": "其他",
}

PAYABLE_STATUS_LABELS = {
    "draft": "草稿",
    "confirmed": "待支付",
    "paid": "已支付",
    "voided": "已作废",
}

SIMPLE_COST_INPUT_CATEGORIES = (
    "production_wages",
    "inbound_freight",
    "delivery_freight",
    "factory_utilities",
    "factory_rent",
    "finance_expense",
    "administrative_expense",
    "maintenance",
    "sales_expense",
    "other",
)

can_finance_view = PermissionChecker("finance.view")
can_cost_view = PermissionChecker("cost.view")
can_cost_manage = PermissionChecker("finance.cost.manage")
can_cost_confirm = PermissionChecker("finance.cost.confirm")
can_cost_export = PermissionChecker("finance.cost.export")


def _require_company_scope(user: User, db: Session) -> None:
    if not has_unrestricted_customer_access(user, db):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="公司级成本费用仅允许全客户范围的财务或管理员查看",
        )


def require_cost_read(
    finance_user: User = Depends(can_finance_view),
    _cost_user: User = Depends(can_cost_view),
    db: Session = Depends(get_db),
) -> User:
    _require_company_scope(finance_user, db)
    return finance_user


def require_cost_manage(
    manager: User = Depends(can_cost_manage),
    _finance_user: User = Depends(can_finance_view),
    _cost_user: User = Depends(can_cost_view),
    db: Session = Depends(get_db),
) -> User:
    _require_company_scope(manager, db)
    return manager


def require_cost_confirm(
    confirmer: User = Depends(can_cost_confirm),
    _finance_user: User = Depends(can_finance_view),
    _cost_user: User = Depends(can_cost_view),
    db: Session = Depends(get_db),
) -> User:
    _require_company_scope(confirmer, db)
    return confirmer


def require_cost_export(
    exporter: User = Depends(can_cost_export),
    _finance_user: User = Depends(can_finance_view),
    _cost_user: User = Depends(can_cost_view),
    db: Session = Depends(get_db),
) -> User:
    _require_company_scope(exporter, db)
    return exporter


def _month(value: str) -> str:
    normalized = str(value or "").strip()
    if not MONTH_PATTERN.fullmatch(normalized):
        raise ValueError("成本月份格式必须为 YYYY-MM")
    return normalized


def _money(value: Decimal | int | float | str) -> Decimal:
    try:
        normalized = Decimal(str(value)).quantize(MONEY, rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError, TypeError) as error:
        raise ValueError("金额格式无效") from error
    if not normalized.is_finite():
        raise ValueError("金额格式无效")
    if abs(normalized) > MAX_MONEY:
        raise ValueError("金额不能超过 999999999999.99")
    return normalized


def _optional_text(value: object | None, max_length: int) -> str | None:
    normalized = str(value or "").strip()
    if not normalized:
        return None
    if len(normalized) > max_length:
        raise ValueError(f"文字长度不能超过 {max_length} 个字符")
    return normalized


def _request_hash(action: str, payload: dict) -> str:
    normalized = json.dumps(
        {"action": action, "payload": payload},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _idempotency_replay(
    db: Session,
    *,
    idempotency_key: str,
    request_hash: str,
    action: str,
    actor: User,
) -> dict | None:
    row = db.scalar(
        select(FinanceIdempotencyRecord).where(
            FinanceIdempotencyRecord.idempotency_key == idempotency_key
        )
    )
    if row is None:
        return None
    if (
        row.actor_user_id != actor.id
        or row.request_hash != request_hash
        or row.action != action
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "finance_idempotency_conflict",
                "message": "该幂等键已用于不同操作者或不同内容，请刷新后重试",
            },
        )
    return json.loads(row.response_json)


def _record_idempotency(
    db: Session,
    *,
    idempotency_key: str,
    request_hash: str,
    action: str,
    actor: User,
    resource_type: str,
    resource_id: int,
    response: dict,
) -> None:
    db.add(
        FinanceIdempotencyRecord(
            idempotency_key=idempotency_key,
            request_hash=request_hash,
            action=action,
            actor_user_id=actor.id,
            resource_type=resource_type,
            resource_id=resource_id,
            response_json=json.dumps(
                jsonable_encoder(response), ensure_ascii=False, sort_keys=True
            ),
        )
    )


def _audit(
    db: Session,
    *,
    request: Request,
    user: User,
    action: str,
    resource: str,
    entity_id: int | None,
    description: str,
    details: dict,
) -> None:
    append_audit_event(
        db,
        request=request,
        actor=user,
        event_category="business",
        result="success",
        source="web",
        module_code="finance.cost",
        action_code=action,
        legacy_action=action[:30],
        resource=resource,
        entity_type=resource.lower(),
        entity_id=entity_id,
        description=description,
        details=details,
    )


class CostCenterCreate(BaseModel):
    code: str = Field(min_length=2, max_length=30)
    name: str = Field(min_length=1, max_length=100)
    center_type: str
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("code")
    @classmethod
    def validate_code(cls, value: str) -> str:
        normalized = value.strip().upper()
        if not re.fullmatch(r"[A-Z0-9_]{2,30}", normalized):
            raise ValueError("成本中心编码只能使用大写字母、数字和下划线")
        return normalized

    @field_validator("name")
    @classmethod
    def normalize_name(cls, value: str) -> str:
        normalized = value.strip()
        if not normalized:
            raise ValueError("成本中心名称不能为空")
        return normalized

    @field_validator("center_type")
    @classmethod
    def validate_type(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in CENTER_TYPE_LABELS:
            raise ValueError("成本中心类型无效")
        return normalized

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        normalized = value.strip()
        if len(normalized) < 8:
            raise ValueError("幂等键长度不能少于 8 个字符")
        return normalized


class CostCenterUpdate(BaseModel):
    name: str = Field(min_length=1, max_length=100)
    is_active: bool
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("name")
    @classmethod
    def normalize_name(cls, value: str) -> str:
        normalized = value.strip()
        if not normalized:
            raise ValueError("成本中心名称不能为空")
        return normalized

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        normalized = value.strip()
        if len(normalized) < 8:
            raise ValueError("幂等键长度不能少于 8 个字符")
        return normalized


class CostEntryFields(BaseModel):
    cost_month: str
    document_date: date
    cost_center_id: int = Field(gt=0)
    cost_category: str
    allocation_basis: str = "unallocated"
    description: str = Field(min_length=1, max_length=300)
    counterparty_name: str | None = Field(default=None, max_length=200)
    document_number: str | None = Field(default=None, max_length=100)
    source_reference: str = Field(min_length=1, max_length=200)
    amount: Decimal = Field(gt=0, le=MAX_MONEY)
    tax_amount: Decimal = Field(default=Decimal("0.00"), ge=0, le=MAX_MONEY)
    note: str | None = Field(default=None, max_length=1000)

    @field_validator("cost_month")
    @classmethod
    def validate_month(cls, value: str) -> str:
        return _month(value)

    @field_validator("cost_category")
    @classmethod
    def validate_category(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in COST_CATEGORIES:
            raise ValueError("成本费用类别无效")
        return normalized

    @field_validator("allocation_basis")
    @classmethod
    def validate_basis(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in ALLOCATION_BASIS_LABELS:
            raise ValueError("分摊依据无效")
        return normalized

    @field_validator("description", "source_reference")
    @classmethod
    def normalize_description(cls, value: str) -> str:
        normalized = value.strip()
        if not normalized:
            raise ValueError("费用事项和来源依据不能为空")
        return normalized

    @field_validator("counterparty_name", "document_number", "note")
    @classmethod
    def normalize_optional_fields(cls, value: str | None) -> str | None:
        return (value or "").strip() or None

    @model_validator(mode="after")
    def validate_amounts(self):
        self.amount = _money(self.amount)
        self.tax_amount = _money(self.tax_amount)
        if self.amount <= 0:
            raise ValueError("计入成本费用金额必须大于 0")
        if self.tax_amount < 0 or self.tax_amount > self.amount:
            raise ValueError("税额必须在 0 到计入成本费用金额之间")
        return self


class CostEntryCreate(CostEntryFields):
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        normalized = value.strip()
        if len(normalized) < 8:
            raise ValueError("幂等键长度不能少于 8 个字符")
        return normalized


class CostEntryUpdate(CostEntryFields):
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        normalized = value.strip()
        if len(normalized) < 8:
            raise ValueError("幂等键长度不能少于 8 个字符")
        return normalized


class CostEntryTransition(BaseModel):
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)
    reason: str | None = Field(default=None, max_length=500)

    @field_validator("reason")
    @classmethod
    def normalize_reason(cls, value: str | None) -> str | None:
        return (value or "").strip() or None

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        normalized = value.strip()
        if len(normalized) < 8:
            raise ValueError("幂等键长度不能少于 8 个字符")
        return normalized


def _center_response(row: FinanceCostCenter) -> dict:
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "center_type": row.center_type,
        "center_type_label": CENTER_TYPE_LABELS.get(row.center_type, row.center_type),
        "is_active": row.is_active,
        "version": row.version,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
    }


def _entry_response(row: FinanceCostPoolEntry) -> dict:
    category = COST_CATEGORIES[row.cost_category]
    return {
        "id": row.id,
        "cost_month": row.cost_month,
        "document_date": row.document_date,
        "cost_center_id": row.cost_center_id,
        "cost_center_code": row.cost_center_code_snapshot,
        "cost_center_name": row.cost_center_name_snapshot,
        "cost_center_type": row.cost_center_type_snapshot,
        "cost_category": row.cost_category,
        "cost_category_label": category["label"],
        "accounting_class": row.accounting_class,
        "accounting_class_label": ACCOUNTING_CLASS_LABELS.get(
            row.accounting_class, row.accounting_class
        ),
        "allocation_basis": row.allocation_basis,
        "allocation_basis_label": ALLOCATION_BASIS_LABELS.get(
            row.allocation_basis, row.allocation_basis
        ),
        "description": row.description,
        "counterparty_name": row.counterparty_name,
        "document_number": row.document_number,
        "amount": row.amount,
        "tax_amount": row.tax_amount,
        "source_type": row.source_type,
        "source_reference": row.source_reference,
        "source_fingerprint": row.source_fingerprint,
        "note": row.note,
        "status": row.status,
        "status_label": STATUS_LABELS.get(row.status, row.status),
        "version": row.version,
        "confirmed_at": row.confirmed_at,
        "voided_at": row.voided_at,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
    }


def _validated_center(
    db: Session,
    center_id: int,
    *,
    accounting_class: str | None = None,
) -> FinanceCostCenter:
    center = db.get(FinanceCostCenter, center_id)
    if center is None:
        raise HTTPException(status_code=409, detail="成本中心不存在")
    if not center.is_active:
        raise HTTPException(status_code=409, detail="成本中心已停用")
    if accounting_class is not None:
        allowed_types = {
            "manufacturing": {"production", "warehouse_delivery", "unallocated"},
            "selling": {"warehouse_delivery", "sales", "unallocated"},
            "administrative": {"administration", "unallocated"},
            "finance": {"finance", "unallocated"},
            "excluded": {"unallocated"},
        }[accounting_class]
        if center.center_type not in allowed_types:
            raise HTTPException(
                status_code=409,
                detail="所选成本中心与费用类别不匹配，请选择对应中心或待分类",
            )
    return center


@router.get("/cost-pool/metadata")
def cost_pool_metadata(
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_read),
) -> dict:
    centers = db.scalars(
        select(FinanceCostCenter).order_by(
            FinanceCostCenter.is_active.desc(), FinanceCostCenter.id
        )
    ).all()
    return {
        "categories": [
            {"value": code, **metadata} for code, metadata in COST_CATEGORIES.items()
        ],
        "input_categories": [
            {
                "value": code,
                **COST_CATEGORIES[code],
                "label": (
                    "人员工资"
                    if code == "production_wages"
                    else "固定月供 / 车贷"
                    if code == "finance_expense"
                    else COST_CATEGORIES[code]["label"]
                ),
            }
            for code in SIMPLE_COST_INPUT_CATEGORIES
        ],
        "accounting_classes": [
            {"value": code, "label": label}
            for code, label in ACCOUNTING_CLASS_LABELS.items()
        ],
        "allocation_bases": [
            {"value": code, "label": label}
            for code, label in ALLOCATION_BASIS_LABELS.items()
        ],
        "centers": [_center_response(row) for row in centers],
        "warning": "新记账只保留人员工资、运费和固定/经营费用；旧细分类仅作历史显示。纸板与外购包材由供应商月结自动归集。",
    }


@router.get("/cost-centers")
def list_cost_centers(
    include_inactive: bool = False,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_read),
) -> dict:
    query = select(FinanceCostCenter)
    if not include_inactive:
        query = query.where(FinanceCostCenter.is_active.is_(True))
    rows = db.scalars(query.order_by(FinanceCostCenter.id)).all()
    return {"items": [_center_response(row) for row in rows]}


@router.post("/cost-centers", status_code=status.HTTP_201_CREATED)
def create_cost_center(
    payload: CostCenterCreate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict:
    action = "finance.cost_center.create"
    request_hash = _request_hash(action, payload.model_dump())
    replay = _idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action=action,
        actor=user,
    )
    if replay is not None:
        return replay
    try:
        if db.scalar(
            select(FinanceCostCenter.id).where(FinanceCostCenter.code == payload.code)
        ) is not None:
            raise HTTPException(status_code=409, detail="成本中心编码已存在")
        row = FinanceCostCenter(
            code=payload.code,
            name=payload.name,
            center_type=payload.center_type,
            is_active=True,
            created_by=user.id,
        )
        db.add(row)
        db.flush()
        response = _center_response(row)
        _audit(
            db,
            request=request,
            user=user,
            action=action,
            resource="FinanceCostCenter",
            entity_id=row.id,
            description="新增成本中心",
            details=response,
        )
        _record_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action=action,
            actor=user,
            resource_type="finance_cost_center",
            resource_id=row.id,
            response=response,
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise
    except IntegrityError as error:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise HTTPException(status_code=409, detail="成本中心或幂等键发生冲突") from error
    except Exception:
        db.rollback()
        raise


@router.put("/cost-centers/{center_id}")
def update_cost_center(
    center_id: int,
    payload: CostCenterUpdate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict:
    action = "finance.cost_center.update"
    request_hash = _request_hash(action, {"center_id": center_id, **payload.model_dump()})
    replay = _idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action=action,
        actor=user,
    )
    if replay is not None:
        return replay
    try:
        row = db.get(FinanceCostCenter, center_id)
        if row is None:
            raise HTTPException(status_code=404, detail="成本中心不存在")
        if row.version != payload.expected_version:
            raise HTTPException(
                status_code=409,
                detail={"message": "成本中心版本已变化", "current_version": row.version},
            )
        before = _center_response(row)
        result = db.execute(
            update(FinanceCostCenter)
            .where(
                FinanceCostCenter.id == center_id,
                FinanceCostCenter.version == payload.expected_version,
            )
            .values(
                name=payload.name,
                is_active=payload.is_active,
                version=payload.expected_version + 1,
                updated_at=beijing_now_naive(),
            )
            .execution_options(synchronize_session=False)
        )
        if result.rowcount != 1:
            db.expire_all()
            current = db.get(FinanceCostCenter, center_id)
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "成本中心版本已变化",
                    "current_version": current.version if current else None,
                },
            )
        db.expire(row)
        db.refresh(row)
        response = _center_response(row)
        _audit(
            db,
            request=request,
            user=user,
            action=action,
            resource="FinanceCostCenter",
            entity_id=row.id,
            description="更新成本中心",
            details={"before": before, "after": response},
        )
        _record_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action=action,
            actor=user,
            resource_type="finance_cost_center",
            resource_id=row.id,
            response=response,
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise
    except IntegrityError as error:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise HTTPException(status_code=409, detail="成本中心更新冲突") from error
    except Exception:
        db.rollback()
        raise


def _validated_entry_values(
    db: Session,
    payload: CostEntryFields,
) -> tuple[dict, FinanceCostCenter]:
    metadata = COST_CATEGORIES[payload.cost_category]
    accounting_class = metadata["accounting_class"]
    center = _validated_center(
        db, payload.cost_center_id, accounting_class=accounting_class
    )
    values = payload.model_dump(
        exclude={"idempotency_key", "expected_version"}
    )
    values["accounting_class"] = accounting_class
    values["cost_center_code_snapshot"] = center.code
    values["cost_center_name_snapshot"] = center.name
    values["cost_center_type_snapshot"] = center.center_type
    return values, center


@router.get("/cost-pool")
def list_cost_pool_entries(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    status_filter: str | None = Query(default=None, alias="status"),
    category: str | None = None,
    center_id: int | None = Query(default=None, gt=0),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_read),
) -> dict:
    try:
        normalized_month = _month(month)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    query = (
        select(FinanceCostPoolEntry, FinanceCostCenter)
        .join(FinanceCostCenter, FinanceCostCenter.id == FinanceCostPoolEntry.cost_center_id)
        .where(FinanceCostPoolEntry.cost_month == normalized_month)
    )
    if status_filter:
        if status_filter not in STATUS_LABELS:
            raise HTTPException(status_code=400, detail="成本费用状态无效")
        query = query.where(FinanceCostPoolEntry.status == status_filter)
    if category:
        if category not in COST_CATEGORIES:
            raise HTTPException(status_code=400, detail="成本费用类别无效")
        query = query.where(FinanceCostPoolEntry.cost_category == category)
    if center_id is not None:
        query = query.where(FinanceCostPoolEntry.cost_center_id == center_id)
    total = int(
        db.scalar(select(func.count()).select_from(query.subquery())) or 0
    )
    rows = db.execute(
        query.order_by(
            FinanceCostPoolEntry.document_date.desc(),
            FinanceCostPoolEntry.id.desc(),
        )
        .offset((page - 1) * page_size)
        .limit(page_size)
    ).all()
    return {
        "month": normalized_month,
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": [_entry_response(entry) for entry, _center in rows],
    }


@router.post("/cost-pool", status_code=status.HTTP_201_CREATED)
def create_cost_pool_entry(
    payload: CostEntryCreate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict:
    action = "finance.cost_pool.create"
    request_hash = _request_hash(action, payload.model_dump())
    replay = _idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action=action,
        actor=user,
    )
    if replay is not None:
        return replay
    try:
        values, center = _validated_entry_values(db, payload)
        row = FinanceCostPoolEntry(
            **values,
            source_type="manual",
            status="draft",
            created_by=user.id,
        )
        db.add(row)
        db.flush()
        response = _entry_response(row)
        _audit(
            db,
            request=request,
            user=user,
            action=action,
            resource="FinanceCostPoolEntry",
            entity_id=row.id,
            description="新增成本费用草稿",
            details={
                "cost_month": row.cost_month,
                "category": row.cost_category,
                "amount": row.amount,
                "center": center.code,
            },
        )
        _record_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action=action,
            actor=user,
            resource_type="finance_cost_entry",
            resource_id=row.id,
            response=response,
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise
    except IntegrityError as error:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise HTTPException(status_code=409, detail="成本费用或幂等键发生冲突") from error
    except Exception:
        db.rollback()
        raise


@router.put("/cost-pool/{entry_id}")
def update_cost_pool_entry(
    entry_id: int,
    payload: CostEntryUpdate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict:
    action = "finance.cost_pool.update"
    request_hash = _request_hash(action, {"entry_id": entry_id, **payload.model_dump()})
    replay = _idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action=action,
        actor=user,
    )
    if replay is not None:
        return replay
    try:
        row = db.get(FinanceCostPoolEntry, entry_id)
        if row is None:
            raise HTTPException(status_code=404, detail="成本费用记录不存在")
        if row.status != "draft":
            raise HTTPException(status_code=409, detail="只有草稿可以修改")
        if row.version != payload.expected_version:
            raise HTTPException(
                status_code=409,
                detail={"message": "成本费用版本已变化", "current_version": row.version},
            )
        before = _entry_response(row)
        values, center = _validated_entry_values(db, payload)
        if row.source_type == "excel_import":
            values["source_reference"] = row.source_reference
        result = db.execute(
            update(FinanceCostPoolEntry)
            .where(
                FinanceCostPoolEntry.id == entry_id,
                FinanceCostPoolEntry.status == "draft",
                FinanceCostPoolEntry.version == payload.expected_version,
            )
            .values(
                **values,
                version=payload.expected_version + 1,
                updated_at=beijing_now_naive(),
            )
            .execution_options(synchronize_session=False)
        )
        if result.rowcount != 1:
            db.expire_all()
            current = db.get(FinanceCostPoolEntry, entry_id)
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "成本费用状态或版本已变化",
                    "current_version": current.version if current else None,
                    "current_status": current.status if current else None,
                },
            )
        db.expire(row)
        db.refresh(row)
        response = _entry_response(row)
        _audit(
            db,
            request=request,
            user=user,
            action=action,
            resource="FinanceCostPoolEntry",
            entity_id=row.id,
            description="修改成本费用草稿",
            details={"before": before, "after": response},
        )
        _record_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action=action,
            actor=user,
            resource_type="finance_cost_entry",
            resource_id=row.id,
            response=response,
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise
    except IntegrityError as error:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise HTTPException(status_code=409, detail="成本费用更新冲突") from error
    except Exception:
        db.rollback()
        raise


def _transition_cost_entry(
    *,
    entry_id: int,
    payload: CostEntryTransition,
    request: Request,
    db: Session,
    user: User,
    transition: str,
) -> dict:
    action = f"finance.cost_pool.{transition}"
    request_hash = _request_hash(
        action, {"entry_id": entry_id, **payload.model_dump()}
    )
    replay = _idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action=action,
        actor=user,
    )
    if replay is not None:
        return replay
    try:
        row = db.get(FinanceCostPoolEntry, entry_id)
        if row is None:
            raise HTTPException(status_code=404, detail="成本费用记录不存在")
        if row.version != payload.expected_version:
            raise HTTPException(
                status_code=409,
                detail={"message": "成本费用版本已变化", "current_version": row.version},
            )
        center = db.get(FinanceCostCenter, row.cost_center_id)
        assert center is not None
        before_status = row.status
        now = beijing_now_naive()
        if transition == "confirm":
            if row.status != "draft":
                raise HTTPException(status_code=409, detail="只有草稿可以确认")
            _validated_center(
                db, row.cost_center_id, accounting_class=row.accounting_class
            )
            expected_statuses = ("draft",)
            transition_values = {
                "status": "confirmed",
                "confirmed_by": user.id,
                "confirmed_at": now,
                "updated_at": now,
            }
            description = "确认成本费用"
        elif transition == "void":
            if row.status not in {"draft", "confirmed"}:
                raise HTTPException(status_code=409, detail="该成本费用不能作废")
            if not payload.reason:
                raise HTTPException(status_code=422, detail="作废必须填写原因")
            expected_statuses = ("draft", "confirmed")
            next_note = "；".join(
                value
                for value in (row.note, f"作废原因：{payload.reason}")
                if value
            )
            transition_values = {
                "status": "voided",
                "voided_by": user.id,
                "voided_at": now,
                "note": next_note,
                "updated_at": now,
            }
            description = "作废成本费用"
        else:
            raise RuntimeError("unsupported cost transition")
        result = db.execute(
            update(FinanceCostPoolEntry)
            .where(
                FinanceCostPoolEntry.id == entry_id,
                FinanceCostPoolEntry.version == payload.expected_version,
                FinanceCostPoolEntry.status.in_(expected_statuses),
            )
            .values(
                **transition_values,
                version=payload.expected_version + 1,
            )
            .execution_options(synchronize_session=False)
        )
        if result.rowcount != 1:
            db.expire_all()
            current = db.get(FinanceCostPoolEntry, entry_id)
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "成本费用状态或版本已变化",
                    "current_version": current.version if current else None,
                    "current_status": current.status if current else None,
                },
            )
        db.expire(row)
        db.refresh(row)
        response = _entry_response(row)
        _audit(
            db,
            request=request,
            user=user,
            action=action,
            resource="FinanceCostPoolEntry",
            entity_id=row.id,
            description=description,
            details={
                "before_status": before_status,
                "after_status": row.status,
                "reason": payload.reason,
                "amount": row.amount,
            },
        )
        _record_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action=action,
            actor=user,
            resource_type="finance_cost_entry",
            resource_id=row.id,
            response=response,
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise
    except IntegrityError as error:
        db.rollback()
        replay = _idempotency_replay(
            db, idempotency_key=payload.idempotency_key,
            request_hash=request_hash, action=action, actor=user,
        )
        if replay is not None:
            return replay
        raise HTTPException(status_code=409, detail="成本费用状态更新冲突") from error
    except Exception:
        db.rollback()
        raise


@router.post("/cost-pool/{entry_id}/confirm")
def confirm_cost_pool_entry(
    entry_id: int,
    payload: CostEntryTransition,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_confirm),
) -> dict:
    return _transition_cost_entry(
        entry_id=entry_id,
        payload=payload,
        request=request,
        db=db,
        user=user,
        transition="confirm",
    )


@router.post("/cost-pool/{entry_id}/void")
def void_cost_pool_entry(
    entry_id: int,
    payload: CostEntryTransition,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_confirm),
) -> dict:
    return _transition_cost_entry(
        entry_id=entry_id,
        payload=payload,
        request=request,
        db=db,
        user=user,
        transition="void",
    )


def _summary_for_month(db: Session, month: str) -> dict:
    rows = db.execute(
        select(FinanceCostPoolEntry, FinanceCostCenter)
        .join(FinanceCostCenter, FinanceCostCenter.id == FinanceCostPoolEntry.cost_center_id)
        .where(FinanceCostPoolEntry.cost_month == month)
        .order_by(FinanceCostPoolEntry.id)
    ).all()
    status_totals = {status: Decimal("0.00") for status in STATUS_LABELS}
    status_counts = {status: 0 for status in STATUS_LABELS}
    class_totals: dict[str, Decimal] = {}
    category_totals: dict[str, Decimal] = {}
    center_totals: dict[int, dict] = {}
    unallocated_manufacturing = Decimal("0.00")
    for entry, center in rows:
        amount = _money(entry.amount)
        status_totals[entry.status] += amount
        status_counts[entry.status] += 1
        if entry.status != "confirmed":
            continue
        class_totals[entry.accounting_class] = (
            class_totals.get(entry.accounting_class, Decimal("0.00")) + amount
        )
        category_totals[entry.cost_category] = (
            category_totals.get(entry.cost_category, Decimal("0.00")) + amount
        )
        current = center_totals.setdefault(
            center.id,
            {
                "center_id": center.id,
                "center_code": entry.cost_center_code_snapshot,
                "center_name": entry.cost_center_name_snapshot,
                "amount": Decimal("0.00"),
            },
        )
        current["amount"] += amount
        if (
            entry.accounting_class == "manufacturing"
            and entry.allocation_basis == "unallocated"
        ):
            unallocated_manufacturing += amount

    material_cost = material_cost_coverage_report(db, month=month)
    blockers = []
    if not material_cost["lineage_ready"]:
        blockers.append(
            {
                "code": "actual_material_cost_lineage_incomplete",
                "message": (
                    "实际材料成本仍有 "
                    f"{material_cost['uncovered_delivery_lines']} 条送货明细未完整冻结，"
                    "禁止把缺失成本按零结转。"
                ),
                "count": material_cost["uncovered_delivery_lines"],
            }
        )
    if unallocated_manufacturing > 0:
        blockers.append(
            {
                "code": "manufacturing_cost_unallocated",
                "message": "仍有已确认制造费用未选择可信分摊依据。",
                "amount": unallocated_manufacturing.quantize(MONEY),
            }
        )
    blockers.append(
        {
            "code": "month_close_workflow_pending",
            "message": "本轮先完成材料成本来源链；试算、复核、锁月和反结转将在下一闭环启用。",
        }
    )
    return {
        "month": month,
        "status_totals": {
            key: {
                "count": status_counts[key],
                "amount": value.quantize(MONEY),
            }
            for key, value in status_totals.items()
        },
        "confirmed_total": status_totals["confirmed"].quantize(MONEY),
        "by_accounting_class": [
            {
                "accounting_class": key,
                "label": ACCOUNTING_CLASS_LABELS.get(key, key),
                "amount": value.quantize(MONEY),
            }
            for key, value in sorted(
                class_totals.items(), key=lambda item: item[1], reverse=True
            )
        ],
        "by_category": [
            {
                "cost_category": key,
                "label": COST_CATEGORIES[key]["label"],
                "amount": value.quantize(MONEY),
            }
            for key, value in sorted(
                category_totals.items(), key=lambda item: item[1], reverse=True
            )
        ],
        "by_center": sorted(
            (
                {**value, "amount": value["amount"].quantize(MONEY)}
                for value in center_totals.values()
            ),
            key=lambda item: item["amount"],
            reverse=True,
        ),
        "material_cost": material_cost,
        "close_ready": False,
        "blockers": blockers,
        "note": "当前为管理成本费用池汇总，不是正式利润结转结果。",
    }


@router.get("/cost-pool/summary")
def cost_pool_summary(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_read),
) -> dict:
    try:
        normalized_month = _month(month)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return _summary_for_month(db, normalized_month)


@router.get("/material-cost/coverage")
def material_cost_coverage(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_read),
) -> dict:
    try:
        normalized_month = _month(month)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return material_cost_coverage_report(db, month=normalized_month)


def _safe_excel_text(value: object | None) -> str:
    text_value = str(value or "")
    if text_value.startswith(("=", "+", "-", "@")):
        return "'" + text_value
    return text_value


def _workbook_bytes(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _workbook_response(content: bytes, filename: str) -> StreamingResponse:
    return StreamingResponse(
        iter((content,)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename)
        },
    )


def _month_bounds(month: str) -> tuple[date, date]:
    year, month_number = (int(part) for part in month.split("-", maxsplit=1))
    start = date(year, month_number, 1)
    if month_number == 12:
        return start, date(year + 1, 1, 1)
    return start, date(year, month_number + 1, 1)


def _style_header(sheet, row_number: int = 1) -> None:
    fill = PatternFill("solid", fgColor="167D74")
    for cell in sheet[row_number]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


@router.get("/cost-pool/template")
def download_cost_pool_template(
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_export),
) -> StreamingResponse:
    centers = db.scalars(
        select(FinanceCostCenter)
        .where(FinanceCostCenter.is_active.is_(True))
        .order_by(FinanceCostCenter.id)
    ).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成本费用导入"
    headers = [
        "成本月份*",
        "单据日期*",
        "类别编码*",
        "成本中心编码*",
        "说明*",
        "计入成本费用金额*",
        "其中税额",
        "对方单位",
        "单据号",
        "分摊依据",
        "备注",
    ]
    sheet.append(headers)
    _style_header(sheet)
    widths = [12, 13, 30, 20, 32, 20, 14, 22, 18, 18, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"

    example = workbook.create_sheet("填写示例")
    example.append(headers)
    example.append(
        [
            beijing_today().strftime("%Y-%m"),
            beijing_today(),
            "production_wages",
            "PROD",
            "本月生产员工工资",
            Decimal("10000.00"),
            Decimal("0.00"),
            "员工工资",
            "",
            "unallocated",
            "确认后再进入月度试算",
        ]
    )
    _style_header(example)
    example.freeze_panes = "A2"

    dictionary = workbook.create_sheet("编码说明")
    dictionary.append(["类别编码", "类别", "归属", "默认中心"])
    for code, metadata in COST_CATEGORIES.items():
        dictionary.append(
            [
                code,
                metadata["label"],
                ACCOUNTING_CLASS_LABELS[metadata["accounting_class"]],
                metadata["default_center"],
            ]
        )
    dictionary.append([])
    dictionary.append(["成本中心编码", "成本中心", "类型", "启用"])
    for center in centers:
        dictionary.append(
            [
                center.code,
                _safe_excel_text(center.name),
                CENTER_TYPE_LABELS[center.center_type],
                "是",
            ]
        )
    dictionary.append([])
    dictionary.append(["分摊依据编码", "说明"])
    for code, label in ALLOCATION_BASIS_LABELS.items():
        dictionary.append([code, label])
    _style_header(dictionary)
    dictionary.column_dimensions["A"].width = 32
    dictionary.column_dimensions["B"].width = 28
    dictionary.column_dimensions["C"].width = 24
    dictionary.column_dimensions["D"].width = 20
    return _workbook_response(_workbook_bytes(workbook), "成本费用导入模板.xlsx")


def _cost_pool_export_rows(
    db: Session,
    *,
    month: str,
    status_filter: str | None,
) -> list[tuple[FinanceCostPoolEntry, FinanceCostCenter]]:
    query = (
        select(FinanceCostPoolEntry, FinanceCostCenter)
        .join(FinanceCostCenter, FinanceCostCenter.id == FinanceCostPoolEntry.cost_center_id)
        .where(FinanceCostPoolEntry.cost_month == month)
    )
    if status_filter:
        query = query.where(FinanceCostPoolEntry.status == status_filter)
    return list(
        db.execute(
            query.order_by(FinanceCostPoolEntry.document_date, FinanceCostPoolEntry.id)
        ).all()
    )


@router.get("/cost-pool/export")
def export_cost_pool_workpaper(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    status_filter: str | None = Query(default="confirmed", alias="status"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_export),
) -> StreamingResponse:
    try:
        normalized_month = _month(month)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    if status_filter not in {None, "draft", "confirmed", "voided"}:
        raise HTTPException(status_code=400, detail="成本费用状态无效")
    rows = _cost_pool_export_rows(
        db, month=normalized_month, status_filter=status_filter
    )
    workbook = Workbook()
    details = workbook.active
    details.title = "成本费用明细"
    details.append(
        [
            "月份",
            "日期",
            "状态",
            "类别",
            "归属",
            "成本中心",
            "说明",
            "对方单位",
            "单据号",
            "金额",
            "其中税额",
            "分摊依据",
            "来源",
            "来源引用",
            "导入行防重指纹",
            "备注",
            "版本",
        ]
    )
    _style_header(details)
    for entry, center in rows:
        details.append(
            [
                entry.cost_month,
                entry.document_date,
                STATUS_LABELS[entry.status],
                COST_CATEGORIES[entry.cost_category]["label"],
                ACCOUNTING_CLASS_LABELS[entry.accounting_class],
                _safe_excel_text(entry.cost_center_name_snapshot),
                _safe_excel_text(entry.description),
                _safe_excel_text(entry.counterparty_name),
                _safe_excel_text(entry.document_number),
                entry.amount,
                entry.tax_amount,
                ALLOCATION_BASIS_LABELS[entry.allocation_basis],
                entry.source_type,
                _safe_excel_text(entry.source_reference),
                _safe_excel_text(entry.source_fingerprint),
                _safe_excel_text(entry.note),
                entry.version,
            ]
        )
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    for column in ("J", "K"):
        for cell in details[column][1:]:
            cell.number_format = "#,##0.00"
    widths = [10, 12, 10, 18, 18, 18, 30, 22, 18, 14, 14, 18, 14, 24, 42, 30, 10]
    for index, width in enumerate(widths, start=1):
        details.column_dimensions[chr(64 + index)].width = width

    summary_data = _summary_for_month(db, normalized_month)
    summary = workbook.create_sheet("月度汇总")
    summary.append([f"{normalized_month} 成本费用汇总", "金额"])
    _style_header(summary)
    summary.append(["已确认成本费用", summary_data["confirmed_total"]])
    summary.append(
        [
            "已冻结实际材料成本（人民币含税采购口径）",
            summary_data["material_cost"]["actual_material_cost"],
        ]
    )
    summary.append(
        [
            "实际材料成本送货明细覆盖率",
            summary_data["material_cost"]["line_coverage_rate"],
        ]
    )
    summary.append([])
    summary.append(["按归属", "金额"])
    for item in summary_data["by_accounting_class"]:
        summary.append([item["label"], item["amount"]])
    summary.append([])
    summary.append(["结转状态", "未就绪"])
    for blocker in summary_data["blockers"]:
        summary.append(["阻断原因", blocker["message"]])
    summary.append(["说明", summary_data["note"]])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 70
    for cell in summary["B"]:
        if isinstance(cell.value, (int, float, Decimal)):
            cell.number_format = "#,##0.00"

    content = _workbook_bytes(workbook)
    return _workbook_response(content, f"{normalized_month}_成本费用底稿.xlsx")


@router.get("/management-report/export")
def export_management_report(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_export),
) -> StreamingResponse:
    """Export one small-company management workbook without pretending to be a GL.

    Receivables, payables and the management cost pool stay separate in both the
    calculations and the workbook.  This prevents an expense that was entered
    in both AP and the cost pool from being silently counted twice.
    """

    try:
        normalized_month = _month(month)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    month_start, next_month_start = _month_bounds(normalized_month)

    statement_rows = list(
        db.execute(
            select(Statement, Customer.name)
            .join(Customer, Customer.id == Statement.customer_id)
            .where(
                Statement.statement_month == normalized_month,
                Statement.confirmation_status == "confirmed",
            )
            .order_by(Customer.name, Statement.id)
        ).all()
    )
    customer_rows: dict[str, dict[str, Decimal | int | str]] = {}
    confirmed_receivable = Decimal("0.00")
    invoiced_against_statements = Decimal("0.00")
    for statement, customer_name in statement_rows:
        receivable = _money(statement.total_receivable)
        invoiced = min(max(_money(statement.invoiced_amount), Decimal("0.00")), receivable)
        pending_invoice = max(receivable - invoiced, Decimal("0.00"))
        billing_name = str(statement.settlement_name_snapshot or customer_name)
        row = customer_rows.setdefault(
            billing_name,
            {
                "customer_name": billing_name,
                "statement_count": 0,
                "confirmed_receivable": Decimal("0.00"),
                "invoiced_amount": Decimal("0.00"),
                "pending_invoice_amount": Decimal("0.00"),
            },
        )
        row["statement_count"] = int(row["statement_count"]) + 1
        row["confirmed_receivable"] = _money(row["confirmed_receivable"] + receivable)
        row["invoiced_amount"] = _money(row["invoiced_amount"] + invoiced)
        row["pending_invoice_amount"] = _money(
            row["pending_invoice_amount"] + pending_invoice
        )
        confirmed_receivable += receivable
        invoiced_against_statements += invoiced

    payable_rows = list(
        db.scalars(
            select(FinancePayable)
            .where(
                FinancePayable.document_date >= month_start,
                FinancePayable.document_date < next_month_start,
                FinancePayable.status != "voided",
            )
            .order_by(FinancePayable.document_date, FinancePayable.id)
        ).all()
    )
    confirmed_payable = sum(
        (_money(row.amount) for row in payable_rows if row.status in {"confirmed", "paid"}),
        Decimal("0.00"),
    )
    open_payable = sum(
        (_money(row.amount) for row in payable_rows if row.status == "confirmed"),
        Decimal("0.00"),
    )
    draft_payable = sum(
        (_money(row.amount) for row in payable_rows if row.status == "draft"),
        Decimal("0.00"),
    )

    cost_summary = _summary_for_month(db, normalized_month)
    cost_rows = _cost_pool_export_rows(
        db, month=normalized_month, status_filter=None
    )
    coverage_rows = list(
        db.execute(
            select(
                StatementItem.receivable_amount,
                StatementItem.gross_profit_amount,
                StatementItem.unit_cost_snapshot,
            )
            .join(Statement, Statement.id == StatementItem.statement_id)
            .where(
                Statement.statement_month == normalized_month,
                Statement.confirmation_status == "confirmed",
                StatementItem.return_receipt_item_id.is_not(None),
            )
        ).all()
    )
    catalog_covered_rows = [
        row
        for row in coverage_rows
        if Decimal(str(row.unit_cost_snapshot or 0)) > Decimal("0.00")
    ]
    catalog_coverage_rate = (
        Decimal(len(catalog_covered_rows)) / Decimal(len(coverage_rows))
        if coverage_rows
        else Decimal("0.00")
    )
    catalog_material_margin_reference = sum(
        (_money(row.gross_profit_amount) for row in catalog_covered_rows),
        Decimal("0.00"),
    )
    actual_material_cost = cost_summary["material_cost"]

    workbook = Workbook()
    summary = workbook.active
    summary.title = "老板月报"
    summary.append([f"{normalized_month} 经营月报", "金额 / 结果", "口径说明"])
    _style_header(summary)
    summary_rows = [
        ("已确认对账收入", confirmed_receivable.quantize(MONEY), "按已确认对账单月份"),
        ("已登记开票", invoiced_against_statements.quantize(MONEY), "对应本月对账单的开票事实"),
        (
            "待开票",
            max(confirmed_receivable - invoiced_against_statements, Decimal("0.00")).quantize(MONEY),
            "不代表客户欠款或银行未到账",
        ),
        ("已确认应付 / 支出", confirmed_payable.quantize(MONEY), "独立应付台账，不与成本池相加"),
        ("其中待支付", open_payable.quantize(MONEY), "仅状态为待支付的记录"),
        ("应付草稿待复核", draft_payable.quantize(MONEY), "尚未进入已确认应付"),
        ("已确认成本费用", cost_summary["confirmed_total"], "独立管理成本池"),
        (
            "成本费用草稿待复核",
            cost_summary["status_totals"]["draft"]["amount"],
            "尚未进入已确认成本费用",
        ),
        (
            "实际材料成本覆盖率",
            actual_material_cost["line_coverage_rate"],
            "按送货日期；仅统计已冻结采购收料来源的送货明细",
        ),
        (
            "已冻结实际材料成本（人民币含税采购口径）",
            actual_material_cost["actual_material_cost"],
            "外币缺汇率时不并入；不含工资、能耗、外协和期间费用",
        ),
        (
            "产品资料成本覆盖率",
            catalog_coverage_rate.quantize(Decimal("0.0001")),
            "旧产品成本资料参考，不作为正式结转成本",
        ),
        (
            "产品资料毛利参考",
            catalog_material_margin_reference.quantize(MONEY),
            "仅供经营比较，不参与正式材料成本结转",
        ),
        ("月末结转", "暂不可结转", "先补齐材料来源，再进入试算、复核与锁月"),
    ]
    for label, value, note in summary_rows:
        summary.append([label, value, note])
    summary.append([])
    summary.append(["当前阻断", "", ""])
    for blocker in cost_summary["blockers"]:
        summary.append([blocker["message"], blocker.get("amount", ""), "先补资料，再由财务确认"])
    summary.append([])
    summary.append(["重要说明", "客户收款不在 ERP 内核销", "已开票不等于已收款；收款仍按银行和承兑记录核对"])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 22
    summary.column_dimensions["C"].width = 58
    for row in summary.iter_rows(min_row=2, min_col=2, max_col=2):
        if isinstance(row[0].value, (int, float, Decimal)):
            row[0].number_format = "#,##0.00"

    receivables = workbook.create_sheet("客户应收")
    receivables.append(["结算客户 / 购方", "对账单数", "已确认对账", "已登记开票", "待开票"])
    _style_header(receivables)
    for row in sorted(
        customer_rows.values(),
        key=lambda item: (-Decimal(str(item["pending_invoice_amount"])), str(item["customer_name"])),
    ):
        receivables.append(
            [
                _safe_excel_text(row["customer_name"]),
                row["statement_count"],
                row["confirmed_receivable"],
                row["invoiced_amount"],
                row["pending_invoice_amount"],
            ]
        )
    receivables.freeze_panes = "A2"
    receivables.auto_filter.ref = receivables.dimensions
    for column in ("C", "D", "E"):
        for cell in receivables[column][1:]:
            cell.number_format = "#,##0.00"
    for column, width in {"A": 30, "B": 12, "C": 18, "D": 18, "E": 18}.items():
        receivables.column_dimensions[column].width = width

    payables = workbook.create_sheet("应付支出")
    payables.append(["日期", "收款单位", "类别", "单据号", "金额", "到期日", "状态", "备注"])
    _style_header(payables)
    for row in payable_rows:
        payables.append(
            [
                row.document_date,
                _safe_excel_text(row.counterparty_name),
                PAYABLE_CATEGORY_LABELS.get(row.category, row.category),
                _safe_excel_text(row.document_number),
                row.amount,
                row.due_date,
                PAYABLE_STATUS_LABELS.get(row.status, row.status),
                _safe_excel_text(row.note),
            ]
        )
    payables.freeze_panes = "A2"
    payables.auto_filter.ref = payables.dimensions
    for cell in payables["E"][1:]:
        cell.number_format = "#,##0.00"
    for column, width in {"A": 12, "B": 28, "C": 16, "D": 18, "E": 16, "F": 12, "G": 12, "H": 36}.items():
        payables.column_dimensions[column].width = width

    costs = workbook.create_sheet("成本费用")
    costs.append(["日期", "状态", "归属", "类别", "成本中心", "事项", "金额", "来源依据", "版本"])
    _style_header(costs)
    for entry, _center in cost_rows:
        costs.append(
            [
                entry.document_date,
                STATUS_LABELS[entry.status],
                ACCOUNTING_CLASS_LABELS[entry.accounting_class],
                COST_CATEGORIES[entry.cost_category]["label"],
                _safe_excel_text(entry.cost_center_name_snapshot),
                _safe_excel_text(entry.description),
                entry.amount,
                _safe_excel_text(entry.source_reference),
                entry.version,
            ]
        )
    costs.freeze_panes = "A2"
    costs.auto_filter.ref = costs.dimensions
    for cell in costs["G"][1:]:
        cell.number_format = "#,##0.00"
    for column, width in {"A": 12, "B": 12, "C": 18, "D": 18, "E": 18, "F": 32, "G": 16, "H": 28, "I": 10}.items():
        costs.column_dimensions[column].width = width

    content = _workbook_bytes(workbook)
    return _workbook_response(content, f"{normalized_month}_老板经营月报.xlsx")


IMPORT_HEADERS = [
    "成本月份*",
    "单据日期*",
    "类别编码*",
    "成本中心编码*",
    "说明*",
    "计入成本费用金额*",
    "其中税额",
    "对方单位",
    "单据号",
    "分摊依据",
    "备注",
]


def _parse_import_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    normalized = str(value or "").strip()
    try:
        return date.fromisoformat(normalized)
    except ValueError as error:
        raise ValueError("单据日期必须为 YYYY-MM-DD") from error


def _parse_import_workbook(
    content: bytes,
) -> tuple[list[dict], list[dict]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as error:
        raise HTTPException(status_code=422, detail="无法读取 Excel 文件") from error
    if "成本费用导入" not in workbook.sheetnames:
        raise HTTPException(status_code=422, detail="缺少“成本费用导入”工作表")
    sheet = workbook["成本费用导入"]
    header_values = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(max_row=1))]
    if header_values[: len(IMPORT_HEADERS)] != IMPORT_HEADERS:
        raise HTTPException(status_code=422, detail="导入表头不匹配，请使用系统模板")
    rows: list[dict] = []
    errors: list[dict] = []
    for excel_row, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
        import_cells = cells[: len(IMPORT_HEADERS)]
        values = [cell.value for cell in import_cells]
        if not any(value not in (None, "") for value in values):
            continue
        if len(rows) + len(errors) >= MAX_IMPORT_ROWS:
            errors.append(
                {"row": excel_row, "message": f"单次最多导入 {MAX_IMPORT_ROWS} 行"}
            )
            break
        formula_fields = [
            IMPORT_HEADERS[index]
            for index, cell in enumerate(import_cells)
            if cell.data_type == "f"
        ]
        if formula_fields:
            errors.append(
                {
                    "row": excel_row,
                    "message": "导入明细不支持公式，请粘贴为数值或文字："
                    + "、".join(formula_fields),
                }
            )
            continue
        try:
            cost_month = _month(str(values[0] or ""))
            document_date = _parse_import_date(values[1])
            category = str(values[2] or "").strip()
            if category not in COST_CATEGORIES:
                raise ValueError("类别编码无效")
            center_code = str(values[3] or "").strip().upper()
            if not center_code:
                raise ValueError("成本中心编码不能为空")
            description = str(values[4] or "").strip()
            if not description:
                raise ValueError("说明不能为空")
            amount = _money(values[5])
            tax_amount = _money(values[6] or 0)
            if amount <= 0:
                raise ValueError("计入成本费用金额必须大于 0")
            if tax_amount < 0 or tax_amount > amount:
                raise ValueError("税额必须在 0 到金额之间")
            basis = str(values[9] or "unallocated").strip()
            if basis not in ALLOCATION_BASIS_LABELS:
                raise ValueError("分摊依据编码无效")
            accounting_class = COST_CATEGORIES[category]["accounting_class"]
            rows.append(
                {
                    "excel_row": excel_row,
                    "cost_month": cost_month,
                    "document_date": document_date,
                    "cost_center_code": center_code,
                    "cost_category": category,
                    "cost_category_label": COST_CATEGORIES[category]["label"],
                    "accounting_class": accounting_class,
                    "allocation_basis": basis,
                    "description": description[:300],
                    "amount": amount,
                    "tax_amount": tax_amount,
                    "counterparty_name": _optional_text(values[7], 200),
                    "document_number": _optional_text(values[8], 100),
                    "note": _optional_text(values[10], 1000),
                }
            )
        except ValueError as error:
            errors.append({"row": excel_row, "message": str(error)})
    return rows, errors


def _resolve_import_centers(
    db: Session,
    rows: list[dict],
) -> tuple[list[dict], list[dict]]:
    centers_by_code = {
        row.code: row for row in db.scalars(select(FinanceCostCenter)).all()
    }
    resolved: list[dict] = []
    errors: list[dict] = []
    allowed_by_class = {
        "manufacturing": {"production", "warehouse_delivery", "unallocated"},
        "selling": {"warehouse_delivery", "sales", "unallocated"},
        "administrative": {"administration", "unallocated"},
        "finance": {"finance", "unallocated"},
        "excluded": {"unallocated"},
    }
    for row in rows:
        center = centers_by_code.get(row["cost_center_code"])
        if center is None or not center.is_active:
            errors.append(
                {"row": row["excel_row"], "message": "成本中心编码不存在或已停用"}
            )
            continue
        if center.center_type not in allowed_by_class[row["accounting_class"]]:
            errors.append(
                {"row": row["excel_row"], "message": "成本中心与费用类别不匹配"}
            )
            continue
        resolved.append(
            {
                **row,
                "cost_center_id": center.id,
                "cost_center_code": center.code,
                "cost_center_name": center.name,
                "cost_center_type": center.center_type,
            }
        )
    return resolved, errors


def _canonical_import_identity(rows: list[dict]) -> tuple[str, list[str]]:
    canonical_rows: list[str] = []
    for row in rows:
        payload = {
            "cost_month": row["cost_month"],
            "document_date": row["document_date"].isoformat(),
            "cost_center_code": row["cost_center_code"],
            "cost_category": row["cost_category"],
            "allocation_basis": row["allocation_basis"],
            "description": row["description"],
            "amount": format(row["amount"], ".2f"),
            "tax_amount": format(row["tax_amount"], ".2f"),
            "counterparty_name": row["counterparty_name"],
            "document_number": row["document_number"],
            "note": row["note"],
        }
        canonical_rows.append(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        )
    sorted_rows = sorted(canonical_rows)
    batch_fingerprint = hashlib.sha256(
        json.dumps(sorted_rows, ensure_ascii=False, separators=(",", ":")).encode(
            "utf-8"
        )
    ).hexdigest()
    occurrences: dict[str, int] = {}
    row_fingerprints: list[str] = []
    for canonical_row in canonical_rows:
        occurrences[canonical_row] = occurrences.get(canonical_row, 0) + 1
        row_hash = hashlib.sha256(canonical_row.encode("utf-8")).hexdigest()
        row_fingerprints.append(
            f"{row_hash}:{occurrences[canonical_row]}"
        )
    return batch_fingerprint, row_fingerprints


@router.post("/cost-pool/import")
async def import_cost_pool_workbook(
    request: Request,
    file: UploadFile = File(...),
    apply: bool = Form(False),
    idempotency_key: str | None = Form(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict:
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="Excel 文件不能超过 5 MB")
    if not content:
        raise HTTPException(status_code=422, detail="Excel 文件为空")
    file_hash = hashlib.sha256(content).hexdigest()
    normalized_key = str(idempotency_key or "").strip()
    action = "finance.cost_pool.import"
    request_hash = _request_hash(action, {"file_hash": file_hash})
    if apply:
        if len(normalized_key) < 8 or len(normalized_key) > 120:
            raise HTTPException(status_code=422, detail="确认导入时需要有效幂等键")
        replay = _idempotency_replay(
            db,
            idempotency_key=normalized_key,
            request_hash=request_hash,
            action=action,
            actor=user,
        )
        if replay is not None:
            return replay
    parsed_rows, errors = _parse_import_workbook(content)
    batch_fingerprint, row_fingerprints = _canonical_import_identity(parsed_rows)
    for parsed_row, row_fingerprint in zip(
        parsed_rows, row_fingerprints, strict=True
    ):
        parsed_row["_source_fingerprint"] = row_fingerprint
    if apply and not errors and row_fingerprints:
        already_imported_count = int(
            db.scalar(
                select(func.count()).where(
                    FinanceCostPoolEntry.source_fingerprint.in_(row_fingerprints)
                )
            )
            or 0
        )
        if already_imported_count:
            raise HTTPException(
                status_code=409,
                detail=(
                    f"Excel 中有 {already_imported_count} 行已导入；请在 ERP 修改未确认草稿，"
                    "或只用仅含新增行的文件导入，不能整表重导"
                ),
            )
    rows, center_errors = _resolve_import_centers(db, parsed_rows)
    errors.extend(center_errors)
    preview = {
        "filename": file.filename,
        "file_hash": file_hash,
        "batch_fingerprint": batch_fingerprint,
        "row_count": len(rows),
        "error_count": len(errors),
        "total_amount": sum(
            (row["amount"] for row in rows), Decimal("0.00")
        ).quantize(MONEY),
        "errors": errors,
        "items": [
            {key: value for key, value in row.items() if key != "_source_fingerprint"}
            for row in rows
        ],
        "applied": False,
    }
    if not apply:
        return preview
    if errors:
        raise HTTPException(
            status_code=422,
            detail={"message": "导入文件仍有错误，未写入任何数据", "errors": errors},
        )
    if not rows:
        raise HTTPException(status_code=422, detail="没有可导入的成本费用行")
    try:
        created_ids = []
        source_name = _optional_text(file.filename, 120) or "cost-import"
        for values in rows:
            excel_row = int(values["excel_row"])
            row_fingerprint = str(values["_source_fingerprint"])
            row = FinanceCostPoolEntry(
                **{
                    key: value
                    for key, value in values.items()
                    if key
                    not in {
                        "excel_row",
                        "_source_fingerprint",
                        "cost_center_code",
                        "cost_center_name",
                        "cost_center_type",
                        "cost_category_label",
                    }
                },
                cost_center_code_snapshot=values["cost_center_code"],
                cost_center_name_snapshot=values["cost_center_name"],
                cost_center_type_snapshot=values["cost_center_type"],
                source_type="excel_import",
                source_reference=f"{source_name}:{file_hash[:16]}#row:{excel_row}",
                source_fingerprint=row_fingerprint,
                status="draft",
                created_by=user.id,
            )
            db.add(row)
            db.flush()
            created_ids.append(row.id)
        response = {
            **{key: value for key, value in preview.items() if key != "items"},
            "created_ids": created_ids,
            "applied": True,
        }
        _audit(
            db,
            request=request,
            user=user,
            action=action,
            resource="FinanceCostPoolImport",
            entity_id=created_ids[0],
            description="导入成本费用草稿",
            details={
                "file_hash": file_hash,
                "batch_fingerprint": batch_fingerprint,
                "filename": file.filename,
                "row_count": len(created_ids),
                "total_amount": preview["total_amount"],
            },
        )
        _record_idempotency(
            db,
            idempotency_key=normalized_key,
            request_hash=request_hash,
            action=action,
            actor=user,
            resource_type="finance_cost_import",
            resource_id=created_ids[0],
            response=response,
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        raise
    except IntegrityError as error:
        db.rollback()
        replay = _idempotency_replay(
            db,
            idempotency_key=normalized_key,
            request_hash=request_hash,
            action=action,
            actor=user,
        )
        if replay is not None:
            return replay
        if row_fingerprints and int(
            db.scalar(
                select(func.count()).where(
                    FinanceCostPoolEntry.source_fingerprint.in_(row_fingerprints)
                )
            )
            or 0
        ):
            raise HTTPException(
                status_code=409,
                detail="Excel 中包含已导入明细，请勿整表重复导入",
            ) from error
        raise HTTPException(status_code=409, detail="成本费用导入或幂等键发生冲突") from error
    except Exception:
        db.rollback()
        raise
