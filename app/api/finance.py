from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.encoders import jsonable_encoder
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, field_validator, model_validator
from sqlalchemy import and_, case, delete, func, or_, select, text, update
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session, aliased

from app.api.deps import (
    PermissionChecker,
    customer_scope_ids,
    get_db,
    has_permission,
    has_unrestricted_customer_access,
    require_customer_access,
)
from app.core.time_contract import beijing_today
from app.models.audit import OperationLog
from app.models.company_config import CompanyConfig
from app.models.customer import Customer
from app.models.customer_charge import CustomerCharge
from app.models.delivery import Delivery, DeliveryItem
from app.models.finance import (
    FinanceIdempotencyRecord,
    FinanceManualMutation,
    Invoice,
    ReturnReceipt,
    ReturnReceiptItem,
    SettlementRecord,
    Statement,
    StatementAdjustment,
    StatementItem,
)
from app.models.finance_payable import FinancePayable
from app.models.invoice_task import (
    CustomerInvoiceProfile,
    FinanceInvoiceTask,
    FinanceSettlementEntity,
)
from app.models.supplier import Supplier
from app.models.fulfillment_reminder import FulfillmentReminder
from app.models.order import Order, OrderItem
from app.models.mold_tool import MoldToolCustomer
from app.models.printing_plate import PrintingPlate
from app.models.product import Product
from app.models.user import User
from app.models.warehouse_inventory import (
    DeliveryInventoryAllocation,
    InventoryReservation,
    UnorderedFinishedDeliveryReversal,
)
from app.services.audit_log import append_audit_event
from app.services.fulfillment_reminders import (
    MAX_INITIAL_REMINDERS,
    FulfillmentReminderError,
    create_reminder,
    create_reminder_with_replay,
    fulfillment_reminder_write_guard,
    list_receipt_reminders,
    record_mutation,
    reminder_product_options,
    replay_mutation,
    request_hash as fulfillment_reminder_request_hash,
    sync_receipt_source,
    transition_reminder_with_replay,
    update_reminder_with_replay,
)
from app.services.order_number_display import build_display_registry, display_order_number
from app.services.customer_price_tax import (
    VALID_PRICE_TAX_MODES,
    resolve_customer_price_tax_terms,
)
from app.services.unordered_finished_delivery import (
    reconsume_unordered_finished_receipt_returns,
    restore_unordered_finished_receipt_shortage,
)
from app.services.ordered_finished_receipt_return import (
    active_ordered_return_location_ids,
    reconsume_ordered_finished_receipt_returns,
    restore_ordered_finished_receipt_shortage,
)
from app.services.warehouse_inventory import WarehouseInventoryError
from app.services.product_specification import resolved_product_specification
from app.services.statement_pdf import render_customer_statement_pdf


router = APIRouter()
can_read = PermissionChecker("finance.view")
can_operate = PermissionChecker("finance.execute")
can_adjust_reconciliation_period = PermissionChecker(
    "finance.return_receipt.period.adjust"
)
can_manage_customer_charge = PermissionChecker("finance.customer_charge.manage")
can_confirm_customer_charge = PermissionChecker("finance.customer_charge.confirm")
MONEY = Decimal("0.00")
STATEMENT_COST_FIELDS = frozenset(
    {"total_gross_profit", "unit_cost_snapshot", "gross_profit_amount"}
)

_CITY_PREFIXES = ("苏州", "昆山", "常熟", "太仓", "上海", "无锡", "南京", "杭州", "深圳", "广州")
_COMPANY_SUFFIXES = ("股份有限公司", "有限责任公司", "科技有限公司", "有限公司")
_ILLEGAL_CHARS = r'/\:*?"<>|'


def _visible_customer_ids(user: User, db: Session) -> set[int] | None:
    if has_unrestricted_customer_access(user, db):
        return None
    return customer_scope_ids(user, db)


def require_company_finance_read(
    user: User = Depends(can_read),
    db: Session = Depends(get_db),
) -> User:
    """Protect company-wide payables and expense facts from scoped accounts."""

    if not has_unrestricted_customer_access(user, db):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="经营概览仅限可查看全公司数据的账号",
        )
    return user


def _settlement_context_for_customer(
    db: Session, customer_id: int
) -> tuple[FinanceSettlementEntity | None, set[int]]:
    profile = db.scalar(
        select(CustomerInvoiceProfile).where(
            CustomerInvoiceProfile.customer_id == customer_id
        )
    )
    if profile is None or profile.settlement_entity_id is None:
        return None, {customer_id}
    entity = db.get(FinanceSettlementEntity, profile.settlement_entity_id)
    if (
        entity is None
        or not entity.is_enabled
        or entity.confirmation_status != "confirmed"
    ):
        raise HTTPException(status_code=409, detail="客户关联的结算对象未确认或已停用")
    customer_ids = set(
        db.scalars(
            select(CustomerInvoiceProfile.customer_id).where(
                CustomerInvoiceProfile.settlement_entity_id == entity.id,
                CustomerInvoiceProfile.is_enabled.is_(True),
                CustomerInvoiceProfile.confirmation_status == "confirmed",
            )
        ).all()
    )
    if customer_id not in customer_ids:
        customer_ids.add(customer_id)
    return entity, customer_ids


def _statement_scope_customer_ids(db: Session, statement: Statement) -> set[int]:
    """Return the customer scope frozen when the statement was created."""

    if statement.settlement_customer_ids_snapshot_json:
        try:
            values = json.loads(statement.settlement_customer_ids_snapshot_json)
            customer_ids = {
                int(value)
                for value in values
                if isinstance(value, int) or str(value).isdigit()
            }
            if customer_ids:
                return customer_ids
        except (TypeError, ValueError, json.JSONDecodeError):
            pass
    source_ids = set(
        db.scalars(
            select(StatementItem.source_customer_id)
            .where(
                StatementItem.statement_id == statement.id,
                StatementItem.source_customer_id.is_not(None),
            )
            .distinct()
        ).all()
    )
    source_ids.add(statement.customer_id)
    return {int(customer_id) for customer_id in source_ids}


def _statement_cycle_day(db: Session, statement: Statement) -> int:
    if statement.statement_cycle_start_day_snapshot is not None:
        return int(statement.statement_cycle_start_day_snapshot)
    customer = db.get(Customer, statement.customer_id)
    return int(customer.statement_cycle_start_day if customer else 1)


def _redact_statement_costs(payload: dict, user: User) -> dict:
    if has_permission(user, "cost.view"):
        return payload
    redacted = {
        key: value
        for key, value in payload.items()
        if key not in STATEMENT_COST_FIELDS
    }
    if "items" in redacted:
        redacted["items"] = [
            {
                key: value
                for key, value in item.items()
                if key not in STATEMENT_COST_FIELDS
            }
            for item in redacted["items"]
        ]
    return redacted


def _statement_for_user(
    db: Session,
    statement_id: int,
    user: User,
) -> Statement:
    statement = db.get(Statement, statement_id)
    if statement is None:
        raise HTTPException(status_code=404, detail="对账单不存在")
    for source_customer_id in _statement_scope_customer_ids(db, statement):
        require_customer_access(source_customer_id, user, db)
    return statement


def _delivery_for_user(
    db: Session,
    delivery_id: int,
    user: User,
) -> Delivery:
    delivery = db.get(Delivery, delivery_id)
    if delivery is None:
        raise HTTPException(status_code=404, detail="送货单不存在")
    require_customer_access(delivery.customer_id, user, db)
    return delivery


def _return_receipt_for_user(
    db: Session,
    receipt_id: int,
    user: User,
) -> ReturnReceipt:
    receipt = db.get(ReturnReceipt, receipt_id)
    if receipt is None:
        raise HTTPException(status_code=404, detail="回单不存在")
    _delivery_for_user(db, receipt.delivery_id, user)
    return receipt


def _fulfillment_reminder_for_user(
    db: Session,
    reminder_id: int,
    user: User,
) -> FulfillmentReminder:
    visible_customer_ids = _visible_customer_ids(user, db)
    query = select(FulfillmentReminder).where(
        FulfillmentReminder.id == reminder_id
    )
    if visible_customer_ids is not None:
        if not visible_customer_ids:
            raise HTTPException(status_code=404, detail="备忘不存在")
        query = query.where(
            FulfillmentReminder.customer_id.in_(visible_customer_ids)
        )
    reminder = db.scalar(query)
    if reminder is None:
        raise HTTPException(status_code=404, detail="备忘不存在")
    return reminder


def _delivery_for_reminder_user(
    db: Session,
    delivery_id: int,
    user: User,
) -> Delivery:
    visible_customer_ids = _visible_customer_ids(user, db)
    query = select(Delivery).where(Delivery.id == delivery_id)
    if visible_customer_ids is not None:
        if not visible_customer_ids:
            raise HTTPException(status_code=404, detail="送货单不存在")
        query = query.where(Delivery.customer_id.in_(visible_customer_ids))
    delivery = db.scalar(query)
    if delivery is None:
        raise HTTPException(status_code=404, detail="送货单不存在")
    return delivery


def _return_receipt_for_reminder_user(
    db: Session,
    receipt_id: int,
    user: User,
) -> ReturnReceipt:
    visible_customer_ids = _visible_customer_ids(user, db)
    query = (
        select(ReturnReceipt)
        .join(Delivery, Delivery.id == ReturnReceipt.delivery_id)
        .where(ReturnReceipt.id == receipt_id)
    )
    if visible_customer_ids is not None:
        if not visible_customer_ids:
            raise HTTPException(status_code=404, detail="回单不存在")
        query = query.where(Delivery.customer_id.in_(visible_customer_ids))
    receipt = db.scalar(query)
    if receipt is None:
        raise HTTPException(status_code=404, detail="回单不存在")
    return receipt


def _raise_fulfillment_reminder_error(error: FulfillmentReminderError) -> None:
    raise HTTPException(
        status_code=error.status_code,
        detail=error.detail,
    ) from error


def _customer_abbr(name: str) -> str:
    for prefix in _CITY_PREFIXES:
        if name.startswith(prefix):
            name = name[len(prefix):]
            break
    for suffix in _COMPANY_SUFFIXES:
        if name.endswith(suffix):
            name = name[: -len(suffix)]
            break
    chinese = [c for c in name if "一" <= c <= "鿿"]
    abbr = "".join(chinese[:2])
    return abbr if abbr else "客户"


def _safe_filename(name: str) -> str:
    for ch in _ILLEGAL_CHARS:
        name = name.replace(ch, "")
    return name.strip()


class ReturnReceiptLineCreate(BaseModel):
    delivery_item_id: int
    actual_received_quantity: int
    resolution_action: str | None = None
    difference_reason: str | None = None
    return_location_id: int | None = None
    expected_return_layout_version: int | None = Field(default=None, gt=0)


class FulfillmentReminderDraft(BaseModel):
    scope_type: str
    reminder_type: str
    content: str
    suggested_quantity: Decimal | None = None
    cadence: str
    remind_on: date | None = None
    product_id: int | None = None


class FulfillmentReminderCreate(FulfillmentReminderDraft):
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        return value.strip()


class FulfillmentReminderUpdate(FulfillmentReminderDraft):
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        return value.strip()


class FulfillmentReminderTransition(BaseModel):
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        return value.strip()


class ReturnReceiptCreate(BaseModel):
    delivery_id: int
    actual_received_date: date
    reconciliation_month: str | None = None
    idempotency_key: str | None = Field(default=None, min_length=8, max_length=120)
    signed_by: str | None = None
    items: list[ReturnReceiptLineCreate]
    reminders: list[FulfillmentReminderDraft] = Field(default_factory=list)
    reminder_bundle_idempotency_key: str | None = Field(
        default=None,
        max_length=120,
    )

    @field_validator("items")
    @classmethod
    def validate_items(cls, value: list[ReturnReceiptLineCreate]):
        if not value:
            raise ValueError("回单至少需要一条明细")
        ids = [item.delivery_item_id for item in value]
        if len(ids) != len(set(ids)):
            raise ValueError("回单明细不能重复")
        return value

    @field_validator("reconciliation_month")
    @classmethod
    def validate_reconciliation_month(cls, value: str | None) -> str | None:
        return None if value is None else _validated_month(value, "对账归属月份")

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str | None) -> str | None:
        return value.strip() if value else None

    @model_validator(mode="after")
    def validate_reminder_bundle(self):
        if len(self.reminders) > MAX_INITIAL_REMINDERS:
            raise ValueError(f"一张回单最多新增{MAX_INITIAL_REMINDERS}条备忘")
        key = (self.reminder_bundle_idempotency_key or "").strip()
        if self.reminders and len(key) < 8:
            raise ValueError("回单含备忘时必须提供幂等键")
        if not self.reminders and key:
            raise ValueError("没有备忘时不要提交备忘幂等键")
        self.reminder_bundle_idempotency_key = key or None
        return self


class ReturnReceiptUpdate(BaseModel):
    actual_received_date: date
    reconciliation_month: str | None = None
    expected_version: int | None = Field(default=None, gt=0)
    idempotency_key: str | None = Field(default=None, min_length=8, max_length=120)
    signed_by: str | None = None
    items: list[ReturnReceiptLineCreate]

    @field_validator("items")
    @classmethod
    def validate_items(cls, value: list[ReturnReceiptLineCreate]):
        if not value:
            raise ValueError("回单至少需要一条明细")
        ids = [item.delivery_item_id for item in value]
        if len(ids) != len(set(ids)):
            raise ValueError("回单明细不能重复")
        return value

    @field_validator("reconciliation_month")
    @classmethod
    def validate_reconciliation_month(cls, value: str | None) -> str | None:
        return None if value is None else _validated_month(value, "对账归属月份")

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str | None) -> str | None:
        return value.strip() if value else None


class ReturnReceiptReconciliationMonthUpdate(BaseModel):
    reconciliation_month: str
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("reconciliation_month")
    @classmethod
    def validate_reconciliation_month(cls, value: str) -> str:
        return _validated_month(value, "对账归属月份")

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        return value.strip()


class StatementCreate(BaseModel):
    customer_id: int
    statement_month: str
    idempotency_key: str | None = Field(default=None, min_length=8, max_length=120)
    delivery_ids: list[int] = Field(default_factory=list)
    # 兼容旧客户端；后端仍会校验这些明细是否覆盖完整送货单。
    return_receipt_item_ids: list[int] = Field(default_factory=list)
    customer_charge_ids: list[int] = Field(default_factory=list)

    @field_validator("statement_month")
    @classmethod
    def validate_month(cls, value: str) -> str:
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", value):
            raise ValueError("对账月份格式必须为 YYYY-MM")
        return value

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str | None) -> str | None:
        return value.strip() if value else None

    @field_validator("delivery_ids", "return_receipt_item_ids", "customer_charge_ids")
    @classmethod
    def validate_ids(cls, value: list[int]) -> list[int]:
        if len(value) != len(set(value)):
            raise ValueError("待对账送货单或明细不能重复")
        if any(item_id <= 0 for item_id in value):
            raise ValueError("待对账送货单或明细ID必须为正整数")
        return value

    @model_validator(mode="after")
    def validate_selection(self):
        if (
            not self.delivery_ids
            and not self.return_receipt_item_ids
            and not self.customer_charge_ids
        ):
            raise ValueError("至少选择一张待对账送货单或一条已确认附加收费")
        if self.delivery_ids and self.return_receipt_item_ids:
            raise ValueError("送货单与旧版明细选择不能同时提交")
        return self


class CustomerChargeBase(BaseModel):
    customer_id: int = Field(gt=0)
    order_id: int = Field(gt=0)
    order_item_id: int | None = Field(default=None, gt=0)
    mold_tool_id: int | None = Field(default=None, gt=0)
    printing_plate_id: int | None = Field(default=None, gt=0)
    charge_type: str
    display_name: str = Field(min_length=1, max_length=200)
    quantity: Decimal = Field(default=Decimal("1"), gt=0)
    unit: str = Field(default="项", min_length=1, max_length=40)
    unit_price: Decimal = Field(ge=0)
    amount: Decimal = Field(gt=0)
    price_tax_mode: str | None = None
    tax_rate: Decimal | None = None
    tax_project_name: str | None = Field(default=None, max_length=200)
    tax_classification_code: str | None = Field(default=None, max_length=80)
    note: str | None = None

    @field_validator("charge_type")
    @classmethod
    def validate_charge_type(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in {"mold", "printing_plate", "sample", "setup", "freight", "other"}:
            raise ValueError("收费类型无效")
        return normalized

    @field_validator("display_name", "unit")
    @classmethod
    def normalize_required_text(cls, value: str) -> str:
        normalized = value.strip()
        if not normalized:
            raise ValueError("收费名称和单位不能为空")
        return normalized

    @field_validator("tax_project_name", "tax_classification_code", "note")
    @classmethod
    def normalize_optional_text(cls, value: str | None) -> str | None:
        normalized = (value or "").strip()
        return normalized or None

    @model_validator(mode="after")
    def validate_tax_fields(self):
        if self.price_tax_mode not in {None, *VALID_PRICE_TAX_MODES}:
            raise ValueError("税价口径无效")
        if self.tax_rate is not None and not Decimal("0") <= self.tax_rate <= Decimal("1"):
            raise ValueError("税率必须在 0 至 1 之间")
        self.quantity = self.quantity.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
        self.unit_price = self.unit_price.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)
        self.amount = self.amount.quantize(MONEY, rounding=ROUND_HALF_UP)
        return self


class CustomerChargeCreate(CustomerChargeBase):
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        return value.strip()


class CustomerChargeUpdate(CustomerChargeBase):
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        return value.strip()


class CustomerChargeTransition(BaseModel):
    expected_version: int = Field(gt=0)
    idempotency_key: str = Field(min_length=8, max_length=120)
    reconciliation_month: str | None = None

    @field_validator("idempotency_key")
    @classmethod
    def normalize_idempotency_key(cls, value: str) -> str:
        return value.strip()

    @field_validator("reconciliation_month")
    @classmethod
    def validate_month(cls, value: str | None) -> str | None:
        return None if value is None else _validated_month(value, "对账归属月份")


class StatementUpdate(BaseModel):
    statement_month: str

    @field_validator("statement_month")
    @classmethod
    def validate_month(cls, value: str) -> str:
        if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", value):
            raise ValueError("对账月份格式必须为 YYYY-MM")
        return value


class StatementReopen(BaseModel):
    expected_version: int = Field(gt=0)
    reason: str = Field(min_length=2, max_length=500)

    @field_validator("reason")
    @classmethod
    def normalize_reason(cls, value: str) -> str:
        return value.strip()


class StatementLineRemoval(BaseModel):
    statement_item_id: int = Field(gt=0)
    target_month: str

    @field_validator("target_month")
    @classmethod
    def validate_target_month(cls, value: str) -> str:
        return _validated_month(value, "移入月份")


class StatementDisputeAdjustment(BaseModel):
    expected_version: int = Field(gt=0)
    reason: str = Field(min_length=2, max_length=500)
    remove_lines: list[StatementLineRemoval] = Field(default_factory=list)
    add_return_receipt_item_ids: list[int] = Field(default_factory=list)

    @field_validator("reason")
    @classmethod
    def normalize_reason(cls, value: str) -> str:
        return value.strip()

    @field_validator("add_return_receipt_item_ids")
    @classmethod
    def validate_add_ids(cls, value: list[int]) -> list[int]:
        if len(value) != len(set(value)) or any(item_id <= 0 for item_id in value):
            raise ValueError("补入明细不能重复且必须为正整数")
        return value

    @model_validator(mode="after")
    def validate_changes(self):
        removal_ids = [item.statement_item_id for item in self.remove_lines]
        if len(removal_ids) != len(set(removal_ids)):
            raise ValueError("移出明细不能重复")
        if not removal_ids and not self.add_return_receipt_item_ids:
            raise ValueError("至少选择一条移出或补入明细")
        return self


class PayableCreate(BaseModel):
    supplier_id: int | None = Field(default=None, gt=0)
    counterparty_name: str = Field(min_length=1, max_length=200)
    category: str
    document_number: str | None = Field(default=None, max_length=100)
    document_date: date
    due_date: date | None = None
    amount: Decimal = Field(gt=0)
    note: str | None = Field(default=None, max_length=1000)
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("category")
    @classmethod
    def validate_category(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in {
            "material", "outsourcing", "freight", "utilities", "rent",
            "wages", "maintenance", "tax_fee", "other",
        }:
            raise ValueError("应付/支出类别无效")
        return normalized

    @field_validator("counterparty_name", "idempotency_key")
    @classmethod
    def normalize_required_payable_text(cls, value: str) -> str:
        return value.strip()

    @field_validator("document_number", "note")
    @classmethod
    def normalize_optional_payable_text(cls, value: str | None) -> str | None:
        normalized = (value or "").strip()
        return normalized or None

    @model_validator(mode="after")
    def validate_payable(self):
        self.amount = self.amount.quantize(MONEY, rounding=ROUND_HALF_UP)
        if self.due_date is not None and self.due_date < self.document_date:
            raise ValueError("到期日不能早于单据日期")
        return self


class PayableTransition(BaseModel):
    expected_version: int = Field(gt=0)
    reason: str | None = Field(default=None, max_length=500)


def _statement_period(
    statement_month: str,
    cycle_start_day: int,
) -> tuple[date, date]:
    """Return the inclusive delivery-date range assigned to a statement month."""
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", statement_month):
        raise ValueError("对账月份格式必须为 YYYY-MM")
    if not 1 <= cycle_start_day <= 28:
        raise ValueError("客户对账结转日必须在 1 至 28 日之间")
    year, month = (int(part) for part in statement_month.split("-"))
    month_start = date(year, month, 1)
    if cycle_start_day == 1:
        next_month = (
            date(year + 1, 1, 1)
            if month == 12
            else date(year, month + 1, 1)
        )
        return month_start, next_month - timedelta(days=1)
    previous_month_end = month_start - timedelta(days=1)
    return (
        date(previous_month_end.year, previous_month_end.month, cycle_start_day),
        date(year, month, cycle_start_day) - timedelta(days=1),
    )


def _validated_month(value: str, label: str) -> str:
    normalized = str(value or "").strip()
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", normalized):
        raise ValueError(f"{label}格式必须为 YYYY-MM")
    return normalized


def _current_reconciliation_month() -> str:
    return beijing_today().strftime("%Y-%m")


def _shift_month(month: str, offset: int) -> str:
    year, month_number = (int(part) for part in month.split("-"))
    absolute = year * 12 + month_number - 1 + offset
    shifted_year, shifted_month = divmod(absolute, 12)
    return f"{shifted_year:04d}-{shifted_month + 1:02d}"


def _historical_reconciliation_month(
    delivery_date: date,
    cycle_start_day: int,
) -> str:
    cycle_day = int(cycle_start_day or 1)
    if cycle_day == 1 or delivery_date.day < cycle_day:
        return delivery_date.strftime("%Y-%m")
    return _shift_month(delivery_date.strftime("%Y-%m"), 1)


def _effective_reconciliation_month(
    receipt: ReturnReceipt,
    *,
    delivery_date: date,
    cycle_start_day: int,
) -> tuple[str, str]:
    if receipt.reconciliation_month:
        return receipt.reconciliation_month, "explicit"
    return (
        _historical_reconciliation_month(delivery_date, cycle_start_day),
        "historical_rule",
    )


def _finance_request_hash(action: str, value: dict) -> str:
    normalized = json.dumps(
        {"action": action, "payload": value},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _finance_idempotency_replay(
    db: Session,
    *,
    idempotency_key: str | None,
    request_hash: str,
    action: str,
    actor: User,
) -> tuple[dict | None, FinanceIdempotencyRecord | None]:
    if not idempotency_key:
        return None, None
    record = db.scalar(
        select(FinanceIdempotencyRecord).where(
            FinanceIdempotencyRecord.idempotency_key == idempotency_key
        )
    )
    if record is None:
        return None, None
    if (
        record.actor_user_id != actor.id
        or record.action != action
        or record.request_hash != request_hash
    ):
        raise HTTPException(
            status_code=409,
            detail={
                "code": "finance_idempotency_conflict",
                "message": "该幂等键已用于不同操作者或不同内容，请刷新后重试",
            },
        )
    return json.loads(record.response_json), record


def _record_finance_idempotency(
    db: Session,
    *,
    idempotency_key: str | None,
    request_hash: str,
    action: str,
    actor: User,
    resource_type: str,
    resource_id: int,
    response: dict,
) -> None:
    if not idempotency_key:
        return
    db.add(
        FinanceIdempotencyRecord(
            idempotency_key=idempotency_key,
            request_hash=request_hash,
            action=action,
            actor_user_id=actor.id,
            resource_type=resource_type,
            resource_id=resource_id,
            response_json=json.dumps(
                jsonable_encoder(response),
                ensure_ascii=False,
                sort_keys=True,
            ),
        )
    )


class LedgerMutationVersionPayload(BaseModel):
    expected_version: int = Field(gt=0)
    expected_ledger_version: int = Field(gt=0)


class InvoiceCreate(LedgerMutationVersionPayload):
    statement_id: int
    invoice_number: str
    invoice_date: date
    invoice_amount: Decimal
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("invoice_number")
    @classmethod
    def validate_invoice_number(cls, value: str) -> str:
        normalized = value.strip()
        if not normalized:
            raise ValueError("发票号码不能为空")
        return normalized

    @field_validator("invoice_amount")
    @classmethod
    def validate_invoice_amount(cls, value: Decimal) -> Decimal:
        amount = value.quantize(MONEY, rounding=ROUND_HALF_UP)
        if amount <= 0:
            raise ValueError("开票金额必须大于 0")
        return amount

    @field_validator("idempotency_key")
    @classmethod
    def validate_idempotency_key(cls, value: str) -> str:
        normalized = value.strip()
        if len(normalized) < 8:
            raise ValueError("开票幂等键至少需要 8 个字符")
        if normalized.casefold().startswith("system:"):
            raise ValueError("开票幂等键不能使用系统保留前缀")
        return normalized


class SettlementCreate(LedgerMutationVersionPayload):
    amount: Decimal
    settlement_date: date
    account: str | None = None
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("amount")
    @classmethod
    def validate_amount(cls, value: Decimal) -> Decimal:
        amount = value.quantize(MONEY, rounding=ROUND_HALF_UP)
        if amount <= 0:
            raise ValueError("收款金额必须大于 0")
        return amount

    @field_validator("account")
    @classmethod
    def validate_account(cls, value: str | None) -> str | None:
        if value is None:
            return None
        normalized = value.strip()
        return normalized or None

    @field_validator("idempotency_key")
    @classmethod
    def validate_idempotency_key(cls, value: str) -> str:
        normalized = value.strip()
        if len(normalized) < 8:
            raise ValueError("收款幂等键至少需要 8 个字符")
        if normalized.casefold().startswith("system:"):
            raise ValueError("收款幂等键不能使用系统保留前缀")
        return normalized


def _customer_charge_for_user(
    db: Session,
    charge_id: int,
    user: User,
) -> CustomerCharge:
    charge = db.get(CustomerCharge, charge_id)
    if charge is None:
        raise HTTPException(status_code=404, detail="客户附加收费不存在")
    require_customer_access(charge.customer_id, user, db)
    return charge


def _customer_charge_statement_id(db: Session, charge_id: int) -> int | None:
    return db.scalar(
        select(StatementItem.statement_id)
        .where(StatementItem.customer_charge_id == charge_id)
        .limit(1)
    )


def _customer_charge_response(db: Session, charge: CustomerCharge) -> dict:
    order = db.get(Order, charge.order_id)
    statement_id = _customer_charge_statement_id(db, charge.id)
    similar_filters = [
        CustomerCharge.id != charge.id,
        CustomerCharge.customer_id == charge.customer_id,
        CustomerCharge.status != "cancelled",
    ]
    if charge.mold_tool_id is not None:
        similar_filters.append(CustomerCharge.mold_tool_id == charge.mold_tool_id)
    elif charge.printing_plate_id is not None:
        similar_filters.append(
            CustomerCharge.printing_plate_id == charge.printing_plate_id
        )
    else:
        similar_filters.append(CustomerCharge.charge_type == charge.charge_type)
        similar_filters.append(CustomerCharge.display_name == charge.display_name)
    historical_similar_charge_count = db.scalar(
        select(func.count(CustomerCharge.id)).where(*similar_filters)
    ) or 0
    return {
        "id": charge.id,
        "customer_id": charge.customer_id,
        "order_id": charge.order_id,
        "order_item_id": charge.order_item_id,
        "customer_po": order.customer_po if order else None,
        "mold_tool_id": charge.mold_tool_id,
        "printing_plate_id": charge.printing_plate_id,
        "charge_type": charge.charge_type,
        "display_name": charge.display_name,
        "quantity": charge.quantity,
        "unit": charge.unit,
        "unit_price": charge.unit_price,
        "amount": charge.amount,
        "price_tax_mode": charge.price_tax_mode,
        "tax_rate": charge.tax_rate,
        "tax_project_name": charge.tax_project_name,
        "tax_classification_code": charge.tax_classification_code,
        "reconciliation_month": charge.reconciliation_month,
        "note": charge.note,
        "status": charge.status,
        "version": charge.version,
        "statement_id": statement_id,
        "locked": statement_id is not None,
        "historical_similar_charge_count": historical_similar_charge_count,
        "created_at": charge.created_at,
        "updated_at": charge.updated_at,
        "confirmed_at": charge.confirmed_at,
        "cancelled_at": charge.cancelled_at,
    }


def _validate_customer_charge_links(
    db: Session,
    payload: CustomerChargeBase,
    user: User,
) -> None:
    require_customer_access(payload.customer_id, user, db)
    customer = db.get(Customer, payload.customer_id)
    if customer is None:
        raise HTTPException(status_code=400, detail="客户不存在")
    order = db.get(Order, payload.order_id)
    if order is None or order.customer_id != payload.customer_id:
        raise HTTPException(status_code=400, detail="订单与客户不匹配")
    if payload.order_item_id is not None:
        order_item = db.get(OrderItem, payload.order_item_id)
        if order_item is None or order_item.order_id != payload.order_id:
            raise HTTPException(status_code=400, detail="订单明细与订单不匹配")
    if payload.mold_tool_id is not None:
        linked = db.scalar(
            select(MoldToolCustomer.id)
            .where(
                MoldToolCustomer.mold_tool_id == payload.mold_tool_id,
                MoldToolCustomer.customer_id == payload.customer_id,
            )
            .limit(1)
        )
        if linked is None:
            raise HTTPException(status_code=400, detail="模具未关联当前客户")
    if payload.printing_plate_id is not None:
        plate = db.get(PrintingPlate, payload.printing_plate_id)
        if plate is None or plate.customer_id != payload.customer_id:
            raise HTTPException(status_code=400, detail="印版未关联当前客户")


def _assign_customer_charge(
    charge: CustomerCharge,
    payload: CustomerChargeBase,
) -> None:
    for field in (
        "customer_id",
        "order_id",
        "order_item_id",
        "mold_tool_id",
        "printing_plate_id",
        "charge_type",
        "display_name",
        "quantity",
        "unit",
        "unit_price",
        "amount",
        "price_tax_mode",
        "tax_rate",
        "tax_project_name",
        "tax_classification_code",
        "note",
    ):
        setattr(charge, field, getattr(payload, field))


@router.get("/customer-charges")
def list_customer_charges(
    customer_id: int | None = Query(default=None, gt=0),
    order_id: int | None = Query(default=None, gt=0),
    charge_status: str | None = Query(default=None, alias="status"),
    reconciliation_month: str | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    if charge_status not in {None, "draft", "confirmed", "cancelled"}:
        raise HTTPException(status_code=400, detail="附加收费状态无效")
    query = select(CustomerCharge).order_by(
        CustomerCharge.created_at.desc(), CustomerCharge.id.desc()
    )
    if customer_id is not None:
        require_customer_access(customer_id, user, db)
        query = query.where(CustomerCharge.customer_id == customer_id)
    else:
        visible_customer_ids = _visible_customer_ids(user, db)
        if visible_customer_ids is not None:
            query = query.where(CustomerCharge.customer_id.in_(visible_customer_ids))
    if order_id is not None:
        query = query.where(CustomerCharge.order_id == order_id)
    if charge_status is not None:
        query = query.where(CustomerCharge.status == charge_status)
    if reconciliation_month is not None:
        try:
            month = _validated_month(reconciliation_month, "对账归属月份")
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        query = query.where(CustomerCharge.reconciliation_month == month)
    rows = db.scalars(query).all()
    return {"total": len(rows), "items": [_customer_charge_response(db, row) for row in rows]}


@router.get("/customer-charges/{charge_id}")
def get_customer_charge(
    charge_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    return _customer_charge_response(
        db, _customer_charge_for_user(db, charge_id, user)
    )


@router.post("/customer-charges", status_code=status.HTTP_201_CREATED)
def create_customer_charge(
    payload: CustomerChargeCreate,
    db: Session = Depends(get_db),
    user: User = Depends(can_manage_customer_charge),
) -> dict:
    request_hash = _finance_request_hash(
        "customer_charge_create", payload.model_dump(exclude={"idempotency_key"})
    )
    replay, _record = _finance_idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="customer_charge_create",
        actor=user,
    )
    if replay is not None:
        return replay
    _validate_customer_charge_links(db, payload, user)
    charge = CustomerCharge(created_by=user.id, updated_by=user.id)
    _assign_customer_charge(charge, payload)
    db.add(charge)
    db.flush()
    response = _customer_charge_response(db, charge)
    _audit(
        db,
        user=user,
        action="CREATE_CUSTOMER_CHARGE",
        resource="CustomerCharge",
        entity_id=charge.id,
        details={"after": response},
        description="新增客户附加收费草稿",
    )
    _record_finance_idempotency(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="customer_charge_create",
        actor=user,
        resource_type="customer_charge",
        resource_id=charge.id,
        response=response,
    )
    db.commit()
    return response


@router.put("/customer-charges/{charge_id}")
def update_customer_charge(
    charge_id: int,
    payload: CustomerChargeUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(can_manage_customer_charge),
) -> dict:
    request_hash = _finance_request_hash(
        "customer_charge_update",
        {"charge_id": charge_id, **payload.model_dump(exclude={"idempotency_key"})},
    )
    replay, _record = _finance_idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="customer_charge_update",
        actor=user,
    )
    if replay is not None:
        return replay
    charge = _customer_charge_for_user(db, charge_id, user)
    if charge.status != "draft" or _customer_charge_statement_id(db, charge.id):
        raise HTTPException(status_code=409, detail="只有未进入对账单的草稿收费可以修改")
    if charge.version != payload.expected_version:
        raise HTTPException(
            status_code=409,
            detail={"message": "收费版本已变化，请刷新后重试", "current_version": charge.version},
        )
    _validate_customer_charge_links(db, payload, user)
    before = _customer_charge_response(db, charge)
    _assign_customer_charge(charge, payload)
    charge.version += 1
    charge.updated_by = user.id
    db.flush()
    response = _customer_charge_response(db, charge)
    _audit(
        db,
        user=user,
        action="UPDATE_CUSTOMER_CHARGE",
        resource="CustomerCharge",
        entity_id=charge.id,
        details={"before": before, "after": response},
        description="修改客户附加收费草稿",
    )
    _record_finance_idempotency(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="customer_charge_update",
        actor=user,
        resource_type="customer_charge",
        resource_id=charge.id,
        response=response,
    )
    db.commit()
    return response


def _customer_charge_transition(
    *,
    db: Session,
    user: User,
    charge_id: int,
    payload: CustomerChargeTransition,
    action: str,
) -> dict:
    request_hash = _finance_request_hash(
        action,
        {"charge_id": charge_id, **payload.model_dump(exclude={"idempotency_key"})},
    )
    replay, _record = _finance_idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action=action,
        actor=user,
    )
    if replay is not None:
        return replay
    charge = _customer_charge_for_user(db, charge_id, user)
    if charge.version != payload.expected_version:
        raise HTTPException(
            status_code=409,
            detail={"message": "收费版本已变化，请刷新后重试", "current_version": charge.version},
        )
    if _customer_charge_statement_id(db, charge.id) is not None:
        raise HTTPException(status_code=409, detail="收费已进入对账单，不能变更状态")
    before = _customer_charge_response(db, charge)
    now = datetime.now()
    if action == "customer_charge_confirm":
        if charge.status != "draft":
            raise HTTPException(status_code=409, detail="只有草稿收费可以确认")
        if payload.reconciliation_month is None:
            raise HTTPException(status_code=400, detail="确认收费必须选择对账归属月份")
        charge.status = "confirmed"
        charge.reconciliation_month = payload.reconciliation_month
        charge.confirmed_by = user.id
        charge.confirmed_at = now
        description = "确认客户附加收费及对账归属月份"
        audit_action = "CONFIRM_CUSTOMER_CHARGE"
    elif action == "customer_charge_cancel_confirmation":
        if charge.status != "confirmed":
            raise HTTPException(status_code=409, detail="只有已确认收费可以取消确认")
        charge.status = "draft"
        charge.reconciliation_month = None
        charge.confirmed_by = None
        charge.confirmed_at = None
        description = "取消客户附加收费确认"
        audit_action = "CANCEL_CUSTOMER_CHARGE_CONFIRMATION"
    else:
        if charge.status == "cancelled":
            return _customer_charge_response(db, charge)
        charge.status = "cancelled"
        charge.cancelled_by = user.id
        charge.cancelled_at = now
        description = "作废客户附加收费"
        audit_action = "VOID_CUSTOMER_CHARGE"
    charge.version += 1
    charge.updated_by = user.id
    db.flush()
    response = _customer_charge_response(db, charge)
    _audit(
        db,
        user=user,
        action=audit_action,
        resource="CustomerCharge",
        entity_id=charge.id,
        details={"before": before, "after": response},
        description=description,
    )
    _record_finance_idempotency(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action=action,
        actor=user,
        resource_type="customer_charge",
        resource_id=charge.id,
        response=response,
    )
    db.commit()
    return response


@router.post("/customer-charges/{charge_id}/confirm")
def confirm_customer_charge(
    charge_id: int,
    payload: CustomerChargeTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_confirm_customer_charge),
) -> dict:
    return _customer_charge_transition(
        db=db, user=user, charge_id=charge_id, payload=payload, action="customer_charge_confirm"
    )


@router.post("/customer-charges/{charge_id}/cancel-confirmation")
def cancel_customer_charge_confirmation(
    charge_id: int,
    payload: CustomerChargeTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_confirm_customer_charge),
) -> dict:
    return _customer_charge_transition(
        db=db,
        user=user,
        charge_id=charge_id,
        payload=payload,
        action="customer_charge_cancel_confirmation",
    )


@router.post("/customer-charges/{charge_id}/void")
def void_customer_charge(
    charge_id: int,
    payload: CustomerChargeTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_manage_customer_charge),
) -> dict:
    return _customer_charge_transition(
        db=db, user=user, charge_id=charge_id, payload=payload, action="customer_charge_void"
    )


@router.get("/statements")
def list_statements(
    customer_id: int | None = None,
    statement_month: str | None = None,
    balance_type: str | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    if balance_type not in {None, "pending_invoice", "pending_payment"}:
        raise HTTPException(status_code=400, detail="未知的财务待办筛选")
    query = (
        select(
            Statement,
            func.coalesce(
                Statement.settlement_name_snapshot, Customer.name
            ).label("customer_name"),
        )
        .join(Customer, Customer.id == Statement.customer_id)
        .order_by(Statement.statement_month.desc(), Statement.id.desc())
    )
    if customer_id is not None:
        require_customer_access(customer_id, user, db)
        query = query.where(Statement.customer_id == customer_id)
    else:
        visible_customer_ids = _visible_customer_ids(user, db)
        if visible_customer_ids is not None:
            query = query.where(Statement.customer_id.in_(visible_customer_ids))
            unauthorized_source = (
                select(StatementItem.id)
                .where(
                    StatementItem.statement_id == Statement.id,
                    StatementItem.source_customer_id.is_not(None),
                    StatementItem.source_customer_id.not_in(visible_customer_ids),
                )
                .exists()
            )
            query = query.where(~unauthorized_source)
    if statement_month:
        query = query.where(Statement.statement_month == statement_month)
    if balance_type == "pending_invoice":
        query = query.where(Statement.total_receivable > Statement.invoiced_amount)
    elif balance_type == "pending_payment":
        query = query.where(Statement.total_receivable > Statement.settled_amount)
    filtered = query.order_by(None).subquery()
    total = db.scalar(select(func.count()).select_from(filtered)) or 0
    customer_count = db.scalar(
        select(func.count(func.distinct(filtered.c.customer_id)))
    ) or 0
    rows = db.execute(
        query.offset((page - 1) * page_size).limit(page_size)
    ).all()
    return {
        "total": total,
        "customer_count": customer_count,
        "page": page,
        "page_size": page_size,
        "items": [
            _redact_statement_costs(
                {
                    "id": statement.id,
                    "statement_number": statement.statement_number,
                    "customer_id": statement.customer_id,
                    "customer_name": customer_name,
                    "statement_month": statement.statement_month,
                    "total_receivable": statement.total_receivable,
                    "total_gross_profit": statement.total_gross_profit,
                    "invoiced_amount": statement.invoiced_amount,
                    "settled_amount": statement.settled_amount,
                    "status": statement.status,
                    "confirmation_status": statement.confirmation_status,
                    "version": statement.version,
                    "ledger_version": statement.ledger_version,
                    "created_at": statement.created_at,
                },
                user,
            )
            for statement, customer_name in rows
        ],
    }


def _aggregate_settlement_customer_summaries(
    db: Session,
    summaries: list[dict],
    *,
    statement_month: str,
    visible_customer_ids: set[int] | None,
) -> list[dict]:
    if not summaries:
        return []
    candidate_ids = {int(row["customer_id"]) for row in summaries}
    entity_rows = db.execute(
        select(
            CustomerInvoiceProfile.customer_id,
            FinanceSettlementEntity.id,
            FinanceSettlementEntity.entity_name,
            FinanceSettlementEntity.statement_cycle_start_day,
        )
        .join(
            FinanceSettlementEntity,
            FinanceSettlementEntity.id == CustomerInvoiceProfile.settlement_entity_id,
        )
        .where(
            CustomerInvoiceProfile.customer_id.in_(candidate_ids),
            CustomerInvoiceProfile.is_enabled.is_(True),
            CustomerInvoiceProfile.confirmation_status == "confirmed",
            FinanceSettlementEntity.is_enabled.is_(True),
            FinanceSettlementEntity.confirmation_status == "confirmed",
        )
    ).all()
    entity_by_customer = {
        int(customer_id): (int(entity_id), entity_name, int(cycle_day))
        for customer_id, entity_id, entity_name, cycle_day in entity_rows
    }
    entity_ids = {value[0] for value in entity_by_customer.values()}
    members_by_entity: dict[int, set[int]] = {}
    if entity_ids:
        for entity_id, customer_id in db.execute(
            select(
                CustomerInvoiceProfile.settlement_entity_id,
                CustomerInvoiceProfile.customer_id,
            ).where(
                CustomerInvoiceProfile.settlement_entity_id.in_(entity_ids),
                CustomerInvoiceProfile.is_enabled.is_(True),
                CustomerInvoiceProfile.confirmation_status == "confirmed",
            )
        ).all():
            members_by_entity.setdefault(int(entity_id), set()).add(int(customer_id))

    grouped: dict[tuple[str, int], dict] = {}
    for source in summaries:
        customer_id = int(source["customer_id"])
        entity_info = entity_by_customer.get(customer_id)
        if entity_info is not None:
            entity_id, entity_name, cycle_day = entity_info
            member_ids = members_by_entity.get(entity_id, {customer_id})
            if (
                visible_customer_ids is not None
                and not member_ids.issubset(visible_customer_ids)
            ):
                # A partial customer scope must not reveal or create a partial
                # consolidated statement for the settlement entity.
                continue
            key = ("entity", entity_id)
            representative_id = min(member_ids)
            period_start, period_end = _statement_period(statement_month, cycle_day)
            target = grouped.setdefault(
                key,
                {
                    "id": representative_id,
                    "name": entity_name,
                    "settlement_entity_id": entity_id,
                    "customer_ids": sorted(member_ids),
                    "pending_count": 0,
                    "blocked_count": 0,
                    "pending_item_count": 0,
                    "amount": Decimal("0.00"),
                    "statement_cycle_start_day": cycle_day,
                    "period_start": period_start,
                    "period_end": period_end,
                    "delivery_ids": [],
                },
            )
        else:
            key = ("customer", customer_id)
            target = grouped.setdefault(
                key,
                {
                    "id": customer_id,
                    "name": source["customer_name"],
                    "settlement_entity_id": None,
                    "customer_ids": [customer_id],
                    "pending_count": 0,
                    "blocked_count": 0,
                    "pending_item_count": 0,
                    "amount": Decimal("0.00"),
                    "statement_cycle_start_day": source[
                        "statement_cycle_start_day"
                    ],
                    "period_start": source["period_start"],
                    "period_end": source["period_end"],
                    "delivery_ids": [],
                },
            )
        target["pending_count"] += int(source.get("pending_count") or 0)
        target["blocked_count"] += int(source.get("blocked_count") or 0)
        target["pending_item_count"] += int(source.get("pending_item_count") or 0)
        target["amount"] = (
            Decimal(str(target["amount"])) + Decimal(str(source.get("amount") or 0))
        ).quantize(MONEY, rounding=ROUND_HALF_UP)
        target["delivery_ids"].extend(source.get("delivery_ids") or [])
    return sorted(grouped.values(), key=lambda row: (str(row["name"]), row["id"]))


@router.get("/statement-customers")
def list_statement_customers(
    statement_month: str | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    visible_customer_ids = _visible_customer_ids(user, db)
    if statement_month:
        try:
            summaries = pending_statement_customer_summaries(
                db,
                statement_month=statement_month,
                visible_customer_ids=visible_customer_ids,
            )
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        return {
            "items": _aggregate_settlement_customer_summaries(
                db,
                summaries,
                statement_month=statement_month,
                visible_customer_ids=visible_customer_ids,
            )
        }
    candidate_query = (
        select(Delivery.customer_id)
        .join(ReturnReceipt, ReturnReceipt.delivery_id == Delivery.id)
        .where(ReturnReceipt.status == "confirmed")
        .distinct()
    )
    if visible_customer_ids is not None:
        candidate_query = candidate_query.where(
            Delivery.customer_id.in_(visible_customer_ids)
        )
    candidate_ids = set(db.scalars(candidate_query).all())
    charge_candidate_query = (
        select(CustomerCharge.customer_id)
        .where(CustomerCharge.status == "confirmed")
        .distinct()
    )
    if visible_customer_ids is not None:
        charge_candidate_query = charge_candidate_query.where(
            CustomerCharge.customer_id.in_(visible_customer_ids)
        )
    candidate_ids.update(db.scalars(charge_candidate_query).all())
    if not candidate_ids:
        return {"items": []}
    query = (
        select(Customer)
        .where(Customer.id.in_(candidate_ids))
        .order_by(Customer.name)
    )
    items = []
    for customer in db.scalars(query).all():
        period_start = period_end = None
        if statement_month:
            try:
                period_start, period_end = _statement_period(
                    statement_month,
                    customer.statement_cycle_start_day,
                )
            except ValueError as error:
                raise HTTPException(status_code=400, detail=str(error)) from error
        grouped = _pending_statement_groups(
            db,
            customer.id,
            statement_month=statement_month,
            cycle_start_day=customer.statement_cycle_start_day,
            period_start=period_start,
            period_end=period_end,
        )
        selectable_count = sum(
            1 for delivery in grouped if not delivery["selection_blocked"]
        )
        blocked_count = sum(
            1 for delivery in grouped if delivery["selection_blocked"]
        )
        if not selectable_count and not blocked_count:
            continue
        items.append(
            {
                "id": customer.id,
                "name": customer.name,
                "pending_count": selectable_count,
                "blocked_count": blocked_count,
                "statement_cycle_start_day": customer.statement_cycle_start_day,
                "period_start": period_start,
                "period_end": period_end,
            }
        )
    return {"items": items}


def _statement_detail_response(
    db: Session,
    statement_id: int,
    user: User,
) -> dict:
    row = db.execute(
        select(Statement, Customer.name.label("customer_name"))
        .join(Customer, Customer.id == Statement.customer_id)
        .where(Statement.id == statement_id)
    ).one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="对账单不存在")
    statement, customer_name = row
    _statement_for_user(db, statement.id, user)
    source_customer = aliased(Customer)
    items = db.execute(
        select(
            StatementItem.id.label("statement_item_id"),
            StatementItem.source_customer_id,
            source_customer.name.label("source_customer_name"),
            StatementItem.return_receipt_item_id,
            ReturnReceipt.actual_received_date,
            Delivery.delivery_date,
            Delivery.delivery_number,
            case(
                (DeliveryItem.source_type == "unordered_finished", "无订单库存"),
                else_=Order.customer_po,
            ).label("customer_po"),
            func.coalesce(
                DeliveryItem.product_code_snapshot,
                Product.product_code,
            ).label("product_code"),
            func.coalesce(
                DeliveryItem.product_name_snapshot,
                OrderItem.snapshot_product_name,
                Product.product_name,
            ).label("product_name"),
            func.coalesce(
                DeliveryItem.specification_snapshot,
                OrderItem.snapshot_spec,
            ).label("specification"),
            Product.length_mm.label("product_length_mm"),
            Product.width_mm.label("product_width_mm"),
            Product.height_mm.label("product_height_mm"),
            OrderItem.snapshot_material.label("material"),
            DeliveryItem.ordered_quantity_snapshot,
            DeliveryItem.delivered_quantity.label("actual_delivery_quantity"),
            DeliveryItem.over_delivery_quantity,
            StatementItem.actual_received_quantity,
            StatementItem.unit_price_snapshot,
            StatementItem.unit_cost_snapshot,
            StatementItem.receivable_amount,
            StatementItem.gross_profit_amount,
            ReturnReceiptItem.difference_reason,
        )
        .select_from(StatementItem)
        .join(
            ReturnReceiptItem,
            ReturnReceiptItem.id == StatementItem.return_receipt_item_id,
        )
        .join(
            ReturnReceipt,
            ReturnReceipt.id == ReturnReceiptItem.return_receipt_id,
        )
        .join(DeliveryItem, DeliveryItem.id == ReturnReceiptItem.delivery_item_id)
        .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
        .outerjoin(OrderItem, OrderItem.id == DeliveryItem.order_item_id)
        .outerjoin(Order, Order.id == OrderItem.order_id)
        .outerjoin(
            Product,
            Product.id
            == func.coalesce(DeliveryItem.product_id, OrderItem.product_id),
        )
        .outerjoin(
            source_customer,
            source_customer.id == StatementItem.source_customer_id,
        )
        .where(StatementItem.statement_id == statement.id)
        .order_by(StatementItem.id)
    ).all()
    invoices = db.execute(
        select(
            Invoice.id,
            Invoice.invoice_number,
            Invoice.invoice_date,
            Invoice.invoice_amount,
        ).where(Invoice.statement_id == statement.id)
    ).all()
    settlements = db.execute(
        select(
            SettlementRecord.id,
            SettlementRecord.settled_amount,
            SettlementRecord.settlement_date,
            SettlementRecord.account,
        ).where(SettlementRecord.statement_id == statement.id)
    ).all()
    statement_items = []
    for item in items:
        data = dict(item._mapping)
        data["source_type"] = "delivery"
        data["customer_charge_id"] = None
        data["specification"] = resolved_product_specification(
            data.get("specification"),
            length_mm=data.pop("product_length_mm", None),
            width_mm=data.pop("product_width_mm", None),
            height_mm=data.pop("product_height_mm", None),
        )
        statement_items.append(data)
    charge_items = db.execute(
        select(
            StatementItem.id.label("statement_item_id"),
            StatementItem.source_customer_id,
            source_customer.name.label("source_customer_name"),
            StatementItem.customer_charge_id,
            CustomerCharge.confirmed_at.label("actual_received_date"),
            Order.customer_po,
            CustomerCharge.display_name.label("product_name"),
            CustomerCharge.charge_type,
            StatementItem.charge_quantity_snapshot,
            StatementItem.unit_snapshot,
            StatementItem.unit_price_snapshot,
            StatementItem.unit_cost_snapshot,
            StatementItem.receivable_amount,
            StatementItem.gross_profit_amount,
            CustomerCharge.note.label("difference_reason"),
        )
        .join(
            CustomerCharge,
            CustomerCharge.id == StatementItem.customer_charge_id,
        )
        .join(Order, Order.id == CustomerCharge.order_id)
        .outerjoin(
            source_customer,
            source_customer.id == StatementItem.source_customer_id,
        )
        .where(StatementItem.statement_id == statement.id)
        .order_by(StatementItem.id)
    ).all()
    for item in charge_items:
        data = dict(item._mapping)
        data.update(
            {
                "source_type": "customer_charge",
                "return_receipt_item_id": None,
                "delivery_number": None,
                "product_code": None,
                "specification": None,
                "material": None,
                "ordered_quantity_snapshot": None,
                "actual_delivery_quantity": None,
                "over_delivery_quantity": None,
                "actual_received_quantity": None,
                "charge_quantity": data.pop("charge_quantity_snapshot"),
            }
        )
        statement_items.append(data)
    statement_items.sort(key=lambda item: item["statement_item_id"])
    return _redact_statement_costs(
        {
            "id": statement.id,
            "statement_number": statement.statement_number,
            "customer_id": statement.customer_id,
            "customer_name": statement.settlement_name_snapshot or customer_name,
            "settlement_entity_id": statement.settlement_entity_id,
            "source_customer_ids": sorted(
                _statement_scope_customer_ids(db, statement)
            ),
            "statement_month": statement.statement_month,
            "total_receivable": statement.total_receivable,
            "total_gross_profit": statement.total_gross_profit,
            "invoiced_amount": statement.invoiced_amount,
            "settled_amount": statement.settled_amount,
            "status": statement.status,
            "confirmation_status": statement.confirmation_status,
            "version": statement.version,
            "ledger_version": statement.ledger_version,
            "status_label": "已结清" if statement.status == "settled" else "未结清",
            "item_count": len(statement_items),
            "invoice_count": len(invoices),
            "settlement_count": len(settlements),
            "items": statement_items,
            "invoices": [dict(row._mapping) for row in invoices],
            "settlements": [dict(row._mapping) for row in settlements],
        },
        user,
    )


@router.get("/statements/{statement_id}/export")
def export_statement_excel(
    statement_id: int,
    request: Request = None,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> StreamingResponse:
    row = db.execute(
        select(Statement, Customer)
        .join(Customer, Customer.id == Statement.customer_id)
        .where(Statement.id == statement_id)
    ).one_or_none()
    if row is None:
        raise HTTPException(status_code=404, detail="对账单不存在")
    statement, customer = row
    require_customer_access(
        statement.customer_id,
        current_user=user,
        db=db,
        request=request,
    )
    lines = db.execute(
        select(
            Delivery.delivery_date,
            Delivery.delivery_number,
            case(
                (DeliveryItem.source_type == "unordered_finished", "无订单库存"),
                else_=Order.customer_po,
            ).label("customer_po"),
            func.coalesce(
                DeliveryItem.product_code_snapshot,
                OrderItem.snapshot_product_code,
            ).label("snapshot_product_code"),
            func.coalesce(
                DeliveryItem.product_name_snapshot,
                OrderItem.snapshot_product_name,
            ).label("snapshot_product_name"),
            func.coalesce(
                DeliveryItem.specification_snapshot,
                OrderItem.snapshot_spec,
            ).label("snapshot_spec"),
            Product.length_mm.label("product_length_mm"),
            Product.width_mm.label("product_width_mm"),
            Product.height_mm.label("product_height_mm"),
            OrderItem.snapshot_material,
            DeliveryItem.ordered_quantity_snapshot,
            DeliveryItem.delivered_quantity.label("actual_delivery_quantity"),
            DeliveryItem.over_delivery_quantity,
            StatementItem.actual_received_quantity,
            StatementItem.unit_price_snapshot,
            StatementItem.receivable_amount,
            ReturnReceiptItem.difference_reason,
        )
        .join(
            ReturnReceiptItem,
            ReturnReceiptItem.id == StatementItem.return_receipt_item_id,
        )
        .join(
            DeliveryItem,
            DeliveryItem.id == ReturnReceiptItem.delivery_item_id,
        )
        .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
        .outerjoin(OrderItem, OrderItem.id == DeliveryItem.order_item_id)
        .outerjoin(Order, Order.id == OrderItem.order_id)
        .outerjoin(
            Product,
            Product.id
            == func.coalesce(DeliveryItem.product_id, OrderItem.product_id),
        )
        .where(StatementItem.statement_id == statement.id)
        .order_by(Delivery.delivery_date, Delivery.delivery_number)
    ).all()
    charge_lines = db.execute(
        select(
            Order.customer_po,
            CustomerCharge.display_name,
            StatementItem.charge_quantity_snapshot,
            StatementItem.unit_snapshot,
            StatementItem.unit_price_snapshot,
            StatementItem.receivable_amount,
            CustomerCharge.note,
        )
        .join(
            CustomerCharge,
            CustomerCharge.id == StatementItem.customer_charge_id,
        )
        .join(Order, Order.id == CustomerCharge.order_id)
        .where(StatementItem.statement_id == statement.id)
        .order_by(StatementItem.id)
    ).all()

    inv = statement.invoiced_amount
    rec = statement.total_receivable
    if inv <= 0:
        invoice_status = "未开票"
    elif inv < rec:
        invoice_status = "部分开票"
    else:
        invoice_status = "已开票"

    sett = statement.settled_amount
    if sett <= 0:
        settlement_status = "未收款"
    elif sett < rec:
        settlement_status = "部分收款"
    else:
        settlement_status = "已结清"

    company = db.scalar(select(CompanyConfig).where(CompanyConfig.id == 1))
    company_name = company.company_name if company and company.company_name else ""

    col_count = 18
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "月结对账单"
    last_col = chr(64 + col_count)
    sheet.merge_cells(f"A1:{last_col}1")
    sheet["A1"] = "月结对账单"
    sheet["A1"].font = Font(size=18, bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells(f"A2:{last_col}2")
    sheet["A2"] = (
        f"客户：{customer.name}    月份：{statement.statement_month}    "
        f"对账单号：{statement.statement_number}"
    )
    sheet.merge_cells(f"A3:{last_col}3")
    sender_parts = [f"供方：{company_name}"] if company_name else []
    if company and company.address:
        sender_parts.append(f"地址：{company.address}")
    if company and company.phone:
        sender_parts.append(f"电话：{company.phone}")
    if company and company.tax_number:
        sender_parts.append(f"税号：{company.tax_number}")
    if company and company.bank_name and company.bank_account:
        sender_parts.append(f"开户行：{company.bank_name}  账号：{company.bank_account}")
    sheet["A3"] = "    ".join(sender_parts)
    sheet["A3"].alignment = Alignment(horizontal="left")
    headers = [
        "客户名称",    # 1
        "客户单号",    # 2
        "存货编码",    # 3
        "送货日期",    # 4
        "送货单号",    # 5
        "产品名称",    # 6
        "规格型号",    # 7
        "材质",        # 8
        "订单数量",    # 9
        "实际送货数量", # 10
        "超订单数量",  # 11
        "实际签收数量", # 12
        "单价",        # 13
        "金额",        # 14
        "备注",        # 15
        "开票状态",    # 16
        "对账状态",    # 17
        "结清状态",    # 18
    ]
    sheet.append([])
    sheet.append(headers)
    for cell in sheet[5]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
        cell.alignment = Alignment(horizontal="center")
    for line in lines:
        sheet.append(
            [
                customer.name,
                line.customer_po,
                line.snapshot_product_code,
                line.delivery_date,
                line.delivery_number,
                line.snapshot_product_name,
                resolved_product_specification(
                    line.snapshot_spec,
                    length_mm=line.product_length_mm,
                    width_mm=line.product_width_mm,
                    height_mm=line.product_height_mm,
                ),
                line.snapshot_material,
                line.ordered_quantity_snapshot,
                line.actual_delivery_quantity,
                line.over_delivery_quantity,
                line.actual_received_quantity,
                float(line.unit_price_snapshot),
                float(line.receivable_amount),
                line.difference_reason,
                invoice_status,
                "已对账",
                settlement_status,
            ]
        )
    for line in charge_lines:
        sheet.append(
            [
                customer.name,
                line.customer_po,
                "",
                "",
                "",
                line.display_name,
                f"{line.charge_quantity_snapshot:g}{line.unit_snapshot}",
                "客户附加收费",
                "",
                "",
                "",
                "",
                float(line.unit_price_snapshot),
                float(line.receivable_amount),
                line.note,
                invoice_status,
                "已对账",
                settlement_status,
            ]
        )
    total_row = sheet.max_row + 2
    sheet.cell(total_row, 13, "合计")
    sheet.cell(total_row, 14, float(statement.total_receivable))
    sheet.cell(total_row, 13).font = Font(bold=True)
    sheet.cell(total_row, 14).font = Font(bold=True)
    widths = [
        22, 18, 16, 13, 20, 28, 20, 14, 12,
        14, 12, 14, 12, 14, 24, 10, 10, 10,
    ]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A6"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    abbr = _customer_abbr(customer.name)
    raw_name = f"{abbr}{statement.statement_month}对账单.xlsx"
    filename = _safe_filename(raw_name)
    encoded = quote(filename, safe="")
    export_bytes = output.getvalue()
    append_audit_event(
        db,
        request=request,
        actor=user,
        event_category="security",
        result="success",
        source="web",
        module_code="finance",
        action_code="finance.statement.export",
        legacy_action="EXPORT_STATEMENT",
        resource="Statement",
        entity_type="statement",
        entity_id=statement.id,
        object_ref=statement.statement_number,
        customer_id=statement.customer_id,
        customer_name=customer.name,
        description="导出月结对账单 Excel",
        details={
            "statement_month": statement.statement_month,
            "line_count": len(lines) + len(charge_lines),
            "filename": filename,
            "file_size": len(export_bytes),
            "file_sha256": hashlib.sha256(export_bytes).hexdigest(),
        },
    )
    db.commit()
    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="statement.xlsx"; filename*=UTF-8\'\'{encoded}'
            ),
        },
    )


def _customer_statement_export_data(
    db: Session,
    *,
    statement: Statement,
    customer: Customer,
    sort_by: str,
) -> dict:
    settlement_entity = (
        db.get(FinanceSettlementEntity, statement.settlement_entity_id)
        if statement.settlement_entity_id is not None
        else None
    )
    display_name = statement.settlement_name_snapshot or customer.name
    short_name = (
        (settlement_entity.short_name or settlement_entity.entity_name).strip()
        if settlement_entity is not None
        else (customer.chinese_short_name or _customer_abbr(customer.name))
    )
    delivery_rows = db.execute(
        select(
            StatementItem.id.label("statement_item_id"),
            Delivery.delivery_date,
            Delivery.delivery_number,
            Order.order_number,
            case(
                (DeliveryItem.source_type == "unordered_finished", "无订单库存"),
                else_=Order.customer_po,
            ).label("customer_po"),
            func.coalesce(
                DeliveryItem.product_code_snapshot,
                OrderItem.snapshot_product_code,
                Product.product_code,
            ).label("product_code"),
            func.coalesce(
                DeliveryItem.product_name_snapshot,
                OrderItem.snapshot_product_name,
                Product.product_name,
            ).label("product_name"),
            StatementItem.actual_received_quantity.label("quantity"),
            StatementItem.unit_price_snapshot.label("unit_price"),
            StatementItem.receivable_amount,
            StatementItem.price_tax_mode_snapshot,
        )
        .join(
            ReturnReceiptItem,
            ReturnReceiptItem.id == StatementItem.return_receipt_item_id,
        )
        .join(DeliveryItem, DeliveryItem.id == ReturnReceiptItem.delivery_item_id)
        .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
        .outerjoin(OrderItem, OrderItem.id == DeliveryItem.order_item_id)
        .outerjoin(Order, Order.id == OrderItem.order_id)
        .outerjoin(
            Product,
            Product.id == func.coalesce(DeliveryItem.product_id, OrderItem.product_id),
        )
        .where(StatementItem.statement_id == statement.id)
    ).all()
    rows: list[dict] = []
    modes: set[str] = set()
    for row in delivery_rows:
        mode = (
            row.price_tax_mode_snapshot
            if row.price_tax_mode_snapshot in VALID_PRICE_TAX_MODES
            else "tax_inclusive"
        )
        modes.add(mode)
        quantity = Decimal(str(row.quantity or 0))
        unit_price = Decimal(str(row.unit_price or 0))
        amount = (
            (quantity * unit_price).quantize(MONEY, rounding=ROUND_HALF_UP)
            if mode == "tax_exclusive"
            else Decimal(str(row.receivable_amount)).quantize(MONEY)
        )
        rows.append(
            {
                "statement_item_id": row.statement_item_id,
                "delivery_date": row.delivery_date,
                "delivery_number": row.delivery_number,
                "order_number": row.order_number or "",
                "customer_po": row.customer_po,
                "product_code": row.product_code,
                "product_name": row.product_name,
                "quantity": quantity,
                "unit_price": unit_price,
                "amount": amount,
                "price_tax_mode": mode,
            }
        )
    charge_rows = db.execute(
        select(
            StatementItem.id.label("statement_item_id"),
            Order.order_number,
            Order.customer_po,
            CustomerCharge.display_name,
            StatementItem.charge_quantity_snapshot,
            StatementItem.unit_price_snapshot,
            StatementItem.receivable_amount,
            StatementItem.price_tax_mode_snapshot,
        )
        .join(CustomerCharge, CustomerCharge.id == StatementItem.customer_charge_id)
        .join(Order, Order.id == CustomerCharge.order_id)
        .where(StatementItem.statement_id == statement.id)
    ).all()
    for row in charge_rows:
        mode = (
            row.price_tax_mode_snapshot
            if row.price_tax_mode_snapshot in VALID_PRICE_TAX_MODES
            else "tax_inclusive"
        )
        modes.add(mode)
        quantity = Decimal(str(row.charge_quantity_snapshot or 0))
        unit_price = Decimal(str(row.unit_price_snapshot or 0))
        amount = Decimal(str(row.receivable_amount)).quantize(MONEY)
        rows.append(
            {
                "statement_item_id": row.statement_item_id,
                "delivery_date": None,
                "delivery_number": "",
                "order_number": row.order_number or "",
                "customer_po": row.customer_po,
                "product_code": "",
                "product_name": row.display_name,
                "quantity": quantity,
                "unit_price": unit_price,
                "amount": amount,
                "price_tax_mode": mode,
            }
        )
    if sort_by == "order_number":
        rows.sort(
            key=lambda item: (
                str(item["order_number"]),
                item["delivery_date"] or date.min,
                str(item["delivery_number"]),
                item["statement_item_id"],
            )
        )
    else:
        rows.sort(
            key=lambda item: (
                item["delivery_date"] or date.max,
                str(item["delivery_number"]),
                item["statement_item_id"],
            )
        )
    only_mode = next(iter(modes)) if len(modes) == 1 else None
    label_prefix = (
        "未税" if only_mode == "tax_exclusive" else "含税"
        if only_mode == "tax_inclusive"
        else "按客户口径"
    )
    return {
        "rows": rows,
        "quantity_total": sum((item["quantity"] for item in rows), Decimal("0")),
        "amount_total": sum((item["amount"] for item in rows), Decimal("0")).quantize(MONEY),
        "price_label": f"{label_prefix}单价",
        "amount_label": f"{label_prefix}金额",
        "customer_name": display_name,
        "customer_short_name": short_name,
    }


def _statement_export_context(
    db: Session,
    *,
    statement_id: int,
    sort_by: str,
    user: User,
) -> tuple[Statement, Customer, dict]:
    if sort_by not in {"business", "order_number"}:
        raise HTTPException(status_code=400, detail="对账单排序方式无效")
    result = db.execute(
        select(Statement, Customer)
        .join(Customer, Customer.id == Statement.customer_id)
        .where(Statement.id == statement_id)
    ).one_or_none()
    if result is None:
        raise HTTPException(status_code=404, detail="对账单不存在")
    statement, customer = result
    _statement_for_user(db, statement.id, user)
    return statement, customer, _customer_statement_export_data(
        db, statement=statement, customer=customer, sort_by=sort_by
    )


@router.get("/statements/{statement_id}/customer-export.xlsx")
def export_customer_statement_excel(
    statement_id: int,
    sort_by: str = Query(default="business"),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> StreamingResponse:
    statement, customer, data = _statement_export_context(
        db, statement_id=statement_id, sort_by=sort_by, user=user
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "对账单"
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "对账单"
    sheet["A1"].font = Font(size=16, bold=True, color="123B3A")
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:H2")
    sheet["A2"] = (
        f"客户：{data['customer_name']}    月份：{statement.statement_month}    "
        f"对账单号：{statement.statement_number}"
    )
    headers = [
        "送货日期", "送货单号", "客户单号", "存货编码", "产品名称", "数量",
        data["price_label"], data["amount_label"],
    ]
    sheet.append([])
    sheet.append(headers)
    for cell in sheet[4]:
        cell.font = Font(bold=True, color="123B3A")
        cell.fill = PatternFill("solid", fgColor="E7F2F0")
        cell.alignment = Alignment(horizontal="center")
    for item in data["rows"]:
        sheet.append(
            [
                item["delivery_date"], item["delivery_number"], item["customer_po"],
                item["product_code"], item["product_name"], float(item["quantity"]),
                float(item["unit_price"]), float(item["amount"]),
            ]
        )
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 5, "合计")
    sheet.cell(total_row, 6, float(data["quantity_total"]))
    sheet.cell(total_row, 8, float(data["amount_total"]))
    for cell in sheet[total_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F5F8F8")
    for column, width in zip("ABCDEFGH", (13, 20, 20, 19, 32, 12, 14, 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:H{max(4, sheet.max_row - 1)}"
    for row in range(5, sheet.max_row + 1):
        sheet.cell(row, 7).number_format = "0.0000"
        sheet.cell(row, 8).number_format = "0.00"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = _safe_filename(
        f"{data['customer_short_name']}{statement.statement_month}对账单.xlsx"
    )
    encoded = quote(filename, safe="")
    _audit(
        db,
        user=user,
        action="EXPORT_CUSTOMER_STATEMENT_XLSX",
        resource="Statement",
        entity_id=statement.id,
        details={"filename": filename, "sort_by": sort_by, "line_count": len(data["rows"])},
        description="导出客户核对版对账单 Excel",
    )
    db.commit()
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=statement.xlsx; filename*=UTF-8''{encoded}"},
    )


@router.get("/statements/{statement_id}/customer-export.pdf")
def export_customer_statement_pdf(
    statement_id: int,
    sort_by: str = Query(default="business"),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> StreamingResponse:
    statement, customer, data = _statement_export_context(
        db, statement_id=statement_id, sort_by=sort_by, user=user
    )
    content = render_customer_statement_pdf(
        title="对账单",
        subtitle=(
            f"客户：{data['customer_name']}    月份：{statement.statement_month}    "
            f"对账单号：{statement.statement_number}"
        ),
        price_label=data["price_label"],
        amount_label=data["amount_label"],
        rows=data["rows"],
        quantity_total=data["quantity_total"],
        amount_total=data["amount_total"],
    )
    filename = _safe_filename(
        f"{data['customer_short_name']}{statement.statement_month}对账单.pdf"
    )
    encoded = quote(filename, safe="")
    _audit(
        db,
        user=user,
        action="EXPORT_CUSTOMER_STATEMENT_PDF",
        resource="Statement",
        entity_id=statement.id,
        details={"filename": filename, "sort_by": sort_by, "line_count": len(data["rows"])},
        description="导出客户核对版对账单 PDF",
    )
    db.commit()
    return StreamingResponse(
        BytesIO(content),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=statement.pdf; filename*=UTF-8''{encoded}"},
    )


@router.get("/invoices")
def list_invoices(
    statement_id: int | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    query = (
        select(
            Invoice,
            Statement.statement_number,
            Customer.name.label("customer_name"),
        )
        .join(Statement, Statement.id == Invoice.statement_id)
        .join(Customer, Customer.id == Statement.customer_id)
        .order_by(Invoice.invoice_date.desc(), Invoice.id.desc())
    )
    if statement_id is not None:
        statement = db.get(Statement, statement_id)
        if statement is not None:
            require_customer_access(statement.customer_id, user, db)
        query = query.where(Invoice.statement_id == statement_id)
    else:
        visible_customer_ids = _visible_customer_ids(user, db)
        if visible_customer_ids is not None:
            query = query.where(Statement.customer_id.in_(visible_customer_ids))
    total = db.scalar(select(func.count()).select_from(query.subquery())) or 0
    rows = db.execute(
        query.offset((page - 1) * page_size).limit(page_size)
    ).all()
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": [
            {
                "id": invoice.id,
                "statement_id": invoice.statement_id,
                "statement_number": statement_number,
                "customer_name": customer_name,
                "invoice_number": invoice.invoice_number,
                "invoice_date": invoice.invoice_date,
                "invoice_amount": invoice.invoice_amount,
                "created_at": invoice.created_at,
            }
            for invoice, statement_number, customer_name in rows
        ],
    }

def _audit(
    db: Session,
    *,
    user: User,
    action: str,
    resource: str,
    entity_id: int,
    details: dict,
    description: str,
) -> None:
    db.add(
        OperationLog(
            user_id=user.id,
            action=action,
            resource=resource,
            details=json.dumps(details, ensure_ascii=False, default=str),
            username=user.username,
            role=user.role,
            entity_type=resource.lower(),
            entity_id=entity_id,
            description=description,
        )
    )


def _finance_manual_request_hash(payload: dict) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest().upper()


def _finance_manual_replay(
    mutation: FinanceManualMutation,
    *,
    mutation_type: str,
    statement_id: int,
    request_hash: str,
    actor_id: int,
) -> dict:
    if (
        mutation.mutation_type != mutation_type
        or mutation.statement_id != statement_id
        or mutation.request_hash != request_hash
        or mutation.actor_id != actor_id
    ):
        raise HTTPException(
            status_code=409,
            detail="该财务幂等键已由不同操作者或不同请求载荷使用",
        )
    try:
        response = json.loads(mutation.response_json)
    except (TypeError, ValueError) as error:
        raise HTTPException(
            status_code=500,
            detail="财务幂等响应事实损坏，请停止重试并联系管理员",
        ) from error
    if not isinstance(response, dict):
        raise HTTPException(
            status_code=500,
            detail="财务幂等响应事实格式错误，请停止重试并联系管理员",
        )
    return response


def _reserve_finance_manual_mutation(
    db: Session,
    *,
    idempotency_key: str,
    mutation_type: str,
    statement_id: int,
    request_hash: str,
    actor_id: int,
) -> tuple[FinanceManualMutation | None, dict | None]:
    existing = db.scalar(
        select(FinanceManualMutation).where(
            FinanceManualMutation.idempotency_key == idempotency_key
        )
    )
    if existing is not None:
        return None, _finance_manual_replay(
            existing,
            mutation_type=mutation_type,
            statement_id=statement_id,
            request_hash=request_hash,
            actor_id=actor_id,
        )

    mutation = FinanceManualMutation(
        idempotency_key=idempotency_key,
        mutation_type=mutation_type,
        statement_id=statement_id,
        request_hash=request_hash,
        actor_id=actor_id,
        response_json="{}",
    )
    db.add(mutation)
    try:
        db.flush()
    except IntegrityError as error:
        db.rollback()
        existing = db.scalar(
            select(FinanceManualMutation).where(
                FinanceManualMutation.idempotency_key == idempotency_key
            )
        )
        if existing is None:
            raise HTTPException(
                status_code=409,
                detail="财务幂等键已被另一请求占用，请刷新后重试",
            ) from error
        return None, _finance_manual_replay(
            existing,
            mutation_type=mutation_type,
            statement_id=statement_id,
            request_hash=request_hash,
            actor_id=actor_id,
        )
    return mutation, None


def _raise_finance_manual_cas_conflict(
    db: Session,
    *,
    statement_id: int,
    expected_version: int,
    expected_ledger_version: int,
    amount_kind: str,
) -> None:
    db.expire_all()
    statement = db.get(Statement, statement_id)
    if statement is None:
        raise HTTPException(status_code=404, detail="对账单不存在")
    if statement.confirmation_status != "confirmed":
        raise HTTPException(
            status_code=409,
            detail="请先核对并确认对账单，再登记人工开票或收款",
        )
    if statement.version != expected_version:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "对账单来源版本已变化，请刷新后重试",
                "current_version": statement.version,
            },
        )
    if statement.ledger_version != expected_ledger_version:
        raise HTTPException(
            status_code=409,
            detail={
                "message": "对账单财务流水版本已变化，请刷新后重试",
                "current_ledger_version": statement.ledger_version,
            },
        )
    message = (
        "累计开票金额不能超过应收总额"
        if amount_kind == "invoice"
        else "累计收款金额不能超过应收总额"
    )
    raise HTTPException(status_code=409, detail=message)


def _record_finance_manual_response(
    mutation: FinanceManualMutation,
    response: dict,
) -> dict:
    encoded = jsonable_encoder(response)
    mutation.response_json = json.dumps(
        encoded,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return encoded


def _next_statement_number(db: Session, month: str) -> str:
    sequence = db.execute(
        text(
            """
            INSERT INTO statement_monthly_sequences (statement_month, last_value)
            VALUES (:month, 1)
            ON CONFLICT(statement_month)
            DO UPDATE SET last_value = last_value + 1
            RETURNING last_value
            """
        ),
        {"month": month},
    ).scalar_one()
    if sequence > 999:
        raise HTTPException(status_code=409, detail="当月对账单流水号已超过999")
    return f"ST-{month.replace('-', '')}-{sequence:03d}"


def _validated_receipt_resolution(
    delivery_item: DeliveryItem,
    line: ReturnReceiptLineCreate,
) -> tuple[str | None, str | None]:
    actual = line.actual_received_quantity
    delivered = int(delivery_item.delivered_quantity or 0)
    action = (line.resolution_action or "").strip() or None
    reason = (line.difference_reason or "").strip() or None
    if actual < 0:
        raise HTTPException(
            status_code=400,
            detail=f"送货明细{delivery_item.id}实收数量不能小于0",
        )
    if actual == delivered:
        return None, reason
    if delivery_item.source_type == "unordered_finished" and actual > delivered:
        raise HTTPException(
            status_code=400,
            detail=f"无订单库存送货明细{delivery_item.id}实收数量不能超过实际发货数量",
        )
    if actual < delivered and action not in {"continue_delivery", "accept_short"}:
        raise HTTPException(
            status_code=400,
            detail=f"送货明细{delivery_item.id}短收后必须选择继续待送或按实收结单",
        )
    if actual > delivered and action != "accept_over":
        raise HTTPException(
            status_code=400,
            detail=f"送货明细{delivery_item.id}超收后必须确认按实际数量入账",
        )
    return action, reason


def _apply_receipt_order_effect(
    db: Session,
    *,
    delivery_item: DeliveryItem,
    actual_received_quantity: int,
    resolution_action: str | None,
    direction: int,
) -> int | None:
    if delivery_item.source_type == "unordered_finished":
        return None
    if resolution_action not in {"continue_delivery", "accept_short", "accept_over"}:
        return None
    order_item = db.get(OrderItem, delivery_item.order_item_id)
    if order_item is None:
        raise HTTPException(status_code=409, detail="回单关联订单明细不存在")
    adjustment = (
        int(actual_received_quantity) - int(delivery_item.delivered_quantity or 0)
    ) * direction
    adjusted_quantity = int(order_item.delivered_quantity or 0) + adjustment
    if adjusted_quantity < 0:
        raise HTTPException(status_code=409, detail="回单数量与订单累计已送数量冲突")
    order_item.delivered_quantity = adjusted_quantity
    if resolution_action == "accept_short":
        order_item.is_force_closed = direction > 0
    return order_item.order_id


def _assert_no_later_dispatched_deliveries(
    db: Session,
    *,
    receipt: ReturnReceipt,
    receipt_items: list[ReturnReceiptItem],
    delivery_items: dict[int, DeliveryItem],
) -> None:
    """Do not reverse a receipt after its released balance has been dispatched."""
    relevant_order_item_ids = {
        delivery_items[item.delivery_item_id].order_item_id
        for item in receipt_items
        if item.resolution_action == "continue_delivery"
        and item.delivery_item_id in delivery_items
        and delivery_items[item.delivery_item_id].order_item_id is not None
    }
    if not relevant_order_item_ids:
        return

    # Serialize edits for databases that support row-level locks. SQLite treats
    # this as a normal read, which is sufficient for the local deployment.
    db.scalars(
        select(OrderItem)
        .where(OrderItem.id.in_(relevant_order_item_ids))
        .with_for_update()
    ).all()

    for receipt_item in receipt_items:
        if receipt_item.resolution_action != "continue_delivery":
            continue
        source_item = delivery_items.get(receipt_item.delivery_item_id)
        if source_item is None:
            continue
        later_delivery_number = db.scalar(
            select(Delivery.delivery_number)
            .join(DeliveryItem, DeliveryItem.delivery_id == Delivery.id)
            .where(
                DeliveryItem.order_item_id == source_item.order_item_id,
                DeliveryItem.delivery_id != source_item.delivery_id,
                DeliveryItem.is_current.is_(True),
                Delivery.status == "dispatched",
                Delivery.dispatched_at.is_not(None),
                Delivery.dispatched_at >= receipt.created_at,
            )
            .order_by(Delivery.dispatched_at, Delivery.id)
            .limit(1)
        )
        if later_delivery_number:
            raise HTTPException(
                status_code=409,
                detail=(
                    "该回单产生的待送数量已用于后续送货单 "
                    f"{later_delivery_number}，请先取消后续发货后再修改或取消旧回单。"
                ),
            )


def _refresh_receipt_order_statuses(db: Session, order_ids: set[int]) -> None:
    if not order_ids:
        return
    from app.api.deliveries import _refresh_order_status

    for order_id in order_ids:
        _refresh_order_status(db, order_id)


def _claim_return_receipt_status(
    db: Session,
    *,
    receipt_id: int,
    expected_status: str,
    next_status: str,
) -> None:
    try:
        claimed = db.execute(
            update(ReturnReceipt)
            .where(
                ReturnReceipt.id == receipt_id,
                ReturnReceipt.status == expected_status,
            )
            .values(status=next_status)
            .execution_options(synchronize_session=False)
        )
    except OperationalError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="回单正在被其他操作修改，请刷新后重试",
        ) from error
    if claimed.rowcount != 1:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="回单状态已变化，请刷新后重试",
        )


def _claim_return_receipt_version(
    db: Session,
    *,
    receipt_id: int,
    expected_status: str,
    expected_version: int,
) -> int:
    new_version = expected_version + 1
    try:
        claimed = db.execute(
            update(ReturnReceipt)
            .where(
                ReturnReceipt.id == receipt_id,
                ReturnReceipt.status == expected_status,
                ReturnReceipt.version == expected_version,
            )
            .values(version=new_version)
            .execution_options(synchronize_session=False)
        )
    except OperationalError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail={
                "code": "return_receipt_write_conflict",
                "message": "回单正在被其他操作修改，请刷新后重试",
            },
        ) from error
    if claimed.rowcount != 1:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail={
                "code": "return_receipt_version_conflict",
                "message": "回单版本已变化，请刷新后重试",
            },
        )
    return new_version


def _receipt_response(db: Session, receipt_id: int) -> dict:
    receipt = db.get(ReturnReceipt, receipt_id)
    if receipt is None:
        raise HTTPException(status_code=404, detail="回单不存在")
    delivery = db.get(Delivery, receipt.delivery_id)
    customer = db.get(Customer, delivery.customer_id) if delivery else None
    if delivery is None or customer is None:
        raise HTTPException(status_code=409, detail="回单关联送货单或客户不存在")
    effective_month, month_source = _effective_reconciliation_month(
        receipt,
        delivery_date=delivery.delivery_date,
        cycle_start_day=customer.statement_cycle_start_day,
    )
    rows = db.execute(
        select(
            ReturnReceiptItem.id,
            ReturnReceiptItem.delivery_item_id,
            DeliveryItem.delivered_quantity,
            ReturnReceiptItem.actual_received_quantity,
            ReturnReceiptItem.resolution_action,
            ReturnReceiptItem.difference_reason,
            DeliveryItem.source_type,
        )
        .join(
            DeliveryItem,
            DeliveryItem.id == ReturnReceiptItem.delivery_item_id,
        )
        .where(ReturnReceiptItem.return_receipt_id == receipt_id)
        .order_by(ReturnReceiptItem.id)
    ).all()
    receipt_item_ids = [int(row.id) for row in rows]
    return_locations = active_ordered_return_location_ids(
        db,
        return_receipt_item_ids=receipt_item_ids,
    )
    inventory_backed_delivery_item_ids = set(
        db.scalars(
            select(DeliveryInventoryAllocation.delivery_item_id)
            .join(
                InventoryReservation,
                InventoryReservation.id == DeliveryInventoryAllocation.reservation_id,
            )
            .where(
                DeliveryInventoryAllocation.delivery_item_id.in_(
                    [int(row.delivery_item_id) for row in rows]
                ),
                InventoryReservation.sales_order_item_bom_component_id.is_(None),
            )
        ).all()
    )
    return {
        "id": receipt.id,
        "delivery_id": receipt.delivery_id,
        "actual_received_date": receipt.actual_received_date,
        "actual_confirmed_at": receipt.created_at,
        "reconciliation_month": receipt.reconciliation_month,
        "effective_reconciliation_month": effective_month,
        "reconciliation_month_source": month_source,
        "version": receipt.version,
        "signed_by": receipt.signed_by,
        "status": receipt.status,
        "items": [
            {
                **dict(row._mapping),
                "return_location_id": return_locations.get(int(row.id)),
                "requires_return_location": (
                    row.source_type != "unordered_finished"
                    and int(row.delivery_item_id)
                    in inventory_backed_delivery_item_ids
                ),
            }
            for row in rows
        ],
    }


@router.get("/deliveries/{delivery_id}/reminder-product-options")
def get_fulfillment_reminder_product_options(
    delivery_id: int,
    q: str = Query(default="", max_length=100),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=100),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    delivery = _delivery_for_reminder_user(db, delivery_id, user)
    return reminder_product_options(
        db,
        delivery_id=delivery.id,
        customer_id=delivery.customer_id,
        query_text=q,
        page=page,
        page_size=page_size,
    )


@router.get("/return_receipts/{receipt_id}/reminders")
def get_return_receipt_reminders(
    receipt_id: int,
    history: bool = False,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    _return_receipt_for_reminder_user(db, receipt_id, user)
    return list_receipt_reminders(
        db,
        receipt_id=receipt_id,
        history=history,
        page=page,
        page_size=page_size,
    )


@router.post(
    "/return_receipts/{receipt_id}/reminders",
    status_code=status.HTTP_201_CREATED,
)
def add_return_receipt_reminder(
    receipt_id: int,
    payload: FulfillmentReminderCreate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _write_guard: None = Depends(fulfillment_reminder_write_guard),
) -> dict:
    del _write_guard
    receipt = _return_receipt_for_reminder_user(db, receipt_id, user)
    try:
        response, _replayed = create_reminder_with_replay(
            db,
            receipt=receipt,
            value=payload.model_dump(exclude={"idempotency_key"}),
            actor=user,
            idempotency_key=payload.idempotency_key,
        )
        db.commit()
        return response
    except FulfillmentReminderError as error:
        db.rollback()
        _raise_fulfillment_reminder_error(error)
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="备忘幂等键或版本已被另一操作使用",
        ) from error
    except Exception:
        db.rollback()
        raise


@router.put("/fulfillment-reminders/{reminder_id}")
def update_fulfillment_reminder(
    reminder_id: int,
    payload: FulfillmentReminderUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _write_guard: None = Depends(fulfillment_reminder_write_guard),
) -> dict:
    del _write_guard
    _fulfillment_reminder_for_user(db, reminder_id, user)
    try:
        response, _replayed = update_reminder_with_replay(
            db,
            reminder_id=reminder_id,
            value=payload.model_dump(
                exclude={"expected_version", "idempotency_key"}
            ),
            expected_version=payload.expected_version,
            actor=user,
            visible_customer_ids=_visible_customer_ids(user, db),
            idempotency_key=payload.idempotency_key,
        )
        db.commit()
        return response
    except FulfillmentReminderError as error:
        db.rollback()
        _raise_fulfillment_reminder_error(error)
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="备忘幂等键或版本已被另一操作使用",
        ) from error
    except Exception:
        db.rollback()
        raise


def _transition_fulfillment_reminder(
    *,
    reminder_id: int,
    payload: FulfillmentReminderTransition,
    action: str,
    db: Session,
    user: User,
) -> dict:
    _fulfillment_reminder_for_user(db, reminder_id, user)
    try:
        response, _replayed = transition_reminder_with_replay(
            db,
            reminder_id=reminder_id,
            expected_version=payload.expected_version,
            action=action,
            actor=user,
            visible_customer_ids=_visible_customer_ids(user, db),
            idempotency_key=payload.idempotency_key,
        )
        db.commit()
        return response
    except FulfillmentReminderError as error:
        db.rollback()
        _raise_fulfillment_reminder_error(error)
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="备忘幂等键或版本已被另一操作使用",
        ) from error
    except Exception:
        db.rollback()
        raise


@router.post("/fulfillment-reminders/{reminder_id}/resolve")
def resolve_fulfillment_reminder(
    reminder_id: int,
    payload: FulfillmentReminderTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _write_guard: None = Depends(fulfillment_reminder_write_guard),
) -> dict:
    del _write_guard
    return _transition_fulfillment_reminder(
        reminder_id=reminder_id,
        payload=payload,
        action="resolve",
        db=db,
        user=user,
    )


@router.post("/fulfillment-reminders/{reminder_id}/cancel")
def cancel_fulfillment_reminder(
    reminder_id: int,
    payload: FulfillmentReminderTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _write_guard: None = Depends(fulfillment_reminder_write_guard),
) -> dict:
    del _write_guard
    return _transition_fulfillment_reminder(
        reminder_id=reminder_id,
        payload=payload,
        action="cancel",
        db=db,
        user=user,
    )


@router.get("/return_receipts/{receipt_id}")
def get_return_receipt(
    receipt_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    _return_receipt_for_user(db, receipt_id, user)
    return _receipt_response(db, receipt_id)


@router.get("/reconciliation-month-options")
def reconciliation_month_options(
    _user: User = Depends(can_read),
) -> dict:
    current = _current_reconciliation_month()
    return {
        "previous": _shift_month(current, -1),
        "current": current,
        "next": _shift_month(current, 1),
    }


@router.post("/return_receipts", status_code=status.HTTP_201_CREATED)
def create_return_receipt(
    payload: ReturnReceiptCreate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _period_permission: User = Depends(can_adjust_reconciliation_period),
    _write_guard: None = Depends(fulfillment_reminder_write_guard),
) -> dict:
    del _period_permission
    del _write_guard
    _delivery_for_user(db, payload.delivery_id, user)
    # Refreshed clients always submit a finance idempotency key.  Calls from an
    # older cached bundle without that contract keep the legacy NULL grouping
    # instead of being silently moved to the server's current month during a
    # rolling deployment.
    reconciliation_month = payload.reconciliation_month or (
        _current_reconciliation_month() if payload.idempotency_key else None
    )
    request_value = payload.model_dump(exclude={"idempotency_key"})
    request_value["reconciliation_month"] = reconciliation_month
    request_hash = _finance_request_hash("return_receipt_create", request_value)
    replay, replay_record = _finance_idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="return_receipt_create",
        actor=user,
    )
    if replay is not None:
        assert replay_record is not None
        _return_receipt_for_user(db, replay_record.resource_id, user)
        return replay
    try:
        reminder_bundle_key = payload.reminder_bundle_idempotency_key
        reminder_bundle_hash = None
        if reminder_bundle_key:
            reminder_bundle_hash = fulfillment_reminder_request_hash(
                payload.model_dump(
                    exclude={"reminder_bundle_idempotency_key"}
                )
            )
            replay = replay_mutation(
                db,
                idempotency_key=reminder_bundle_key,
                request_hash_value=reminder_bundle_hash,
                action="receipt_create_bundle",
                actor=user,
            )
            if replay is not None:
                return replay
        # 与取消发货竞争时，先对同一送货单执行条件写并取得写锁。
        # 第二个事务等待后会重新判断状态，不能同时确认回单和取消发货。
        claimed = db.execute(
            update(Delivery)
            .where(
                Delivery.id == payload.delivery_id,
                Delivery.status == "dispatched",
            )
            .values(status="dispatched")
        )
        if claimed.rowcount != 1:
            exists = db.scalar(
                select(Delivery.id).where(Delivery.id == payload.delivery_id)
            )
            if exists is None:
                raise HTTPException(status_code=404, detail="送货单不存在")
            raise HTTPException(status_code=400, detail="送货单尚未确认发货")
        delivery = db.get(Delivery, payload.delivery_id)
        delivery_items = db.scalars(
            select(DeliveryItem)
            .where(
                DeliveryItem.delivery_id == delivery.id,
                DeliveryItem.is_current.is_(True),
            )
            .order_by(DeliveryItem.id)
        ).all()
        requested = {line.delivery_item_id: line for line in payload.items}
        if set(requested) != {item.id for item in delivery_items}:
            raise HTTPException(
                status_code=400,
                detail="回单必须包含送货单全部明细",
            )
        receipt = db.scalar(
            select(ReturnReceipt).where(ReturnReceipt.delivery_id == delivery.id)
        )
        if receipt is not None and receipt.status == "confirmed":
            raise HTTPException(status_code=409, detail="该送货单已经提交回单")
        if receipt is not None and receipt.status == "cancelled":
            has_unordered_reversal = db.scalar(
                select(UnorderedFinishedDeliveryReversal.id)
                .join(
                    ReturnReceiptItem,
                    ReturnReceiptItem.id
                    == UnorderedFinishedDeliveryReversal.return_receipt_item_id,
                )
                .where(ReturnReceiptItem.return_receipt_id == receipt.id)
                .limit(1)
            )
            if has_unordered_reversal is not None:
                raise HTTPException(
                    status_code=409,
                    detail=(
                        f"该送货单已有已取消回单（ID {receipt.id}）及原批次库存冲回审计，"
                        "请编辑原回单，不能再次创建回单"
                    ),
                )
        if receipt is None:
            receipt = ReturnReceipt(
                delivery_id=delivery.id,
                actual_received_date=payload.actual_received_date,
                reconciliation_month=reconciliation_month,
                version=1,
                signed_by=(payload.signed_by or "").strip() or None,
                status="confirmed",
                created_by=user.id,
            )
            db.add(receipt)
            db.flush()
        else:
            db.execute(
                delete(ReturnReceiptItem).where(
                    ReturnReceiptItem.return_receipt_id == receipt.id
                )
            )
            receipt.actual_received_date = payload.actual_received_date
            receipt.reconciliation_month = reconciliation_month
            receipt.version = int(receipt.version or 1) + 1
            receipt.signed_by = (payload.signed_by or "").strip() or None
            receipt.status = "confirmed"
        affected_order_ids: set[int] = set()
        audit_items: list[dict] = []
        for item in delivery_items:
            line = requested[item.id]
            action, reason = _validated_receipt_resolution(item, line)
            receipt_item = ReturnReceiptItem(
                return_receipt_id=receipt.id,
                delivery_item_id=item.id,
                actual_received_quantity=line.actual_received_quantity,
                resolution_action=action,
                difference_reason=reason,
            )
            db.add(receipt_item)
            db.flush()
            if item.source_type == "unordered_finished":
                restore_unordered_finished_receipt_shortage(
                    db,
                    delivery=delivery,
                    delivery_item=item,
                    return_receipt_item_id=receipt_item.id,
                    actual_received_quantity=line.actual_received_quantity,
                    operator_id=user.id,
                    reason=reason,
                )
            else:
                restore_ordered_finished_receipt_shortage(
                    db,
                    delivery=delivery,
                    delivery_item=item,
                    return_receipt_item_id=receipt_item.id,
                    actual_received_quantity=line.actual_received_quantity,
                    resolution_action=action,
                    return_location_id=line.return_location_id,
                    expected_return_layout_version=(
                        line.expected_return_layout_version
                    ),
                    stock_date=payload.actual_received_date,
                    operator_id=user.id,
                )
            order_id = _apply_receipt_order_effect(
                db,
                delivery_item=item,
                actual_received_quantity=line.actual_received_quantity,
                resolution_action=action,
                direction=1,
            )
            if order_id is not None:
                affected_order_ids.add(order_id)
            audit_items.append(
                {
                    "delivery_item_id": item.id,
                    "delivered_quantity": item.delivered_quantity,
                    "actual_received_quantity": line.actual_received_quantity,
                    "resolution_action": action,
                    "difference_reason": reason,
                    "return_location_id": line.return_location_id,
                    "expected_return_layout_version": (
                        line.expected_return_layout_version
                    ),
                }
            )
        _refresh_receipt_order_statuses(db, affected_order_ids)
        _audit(
            db,
            user=user,
            action="CONFIRM_RETURN_RECEIPT",
            resource="ReturnReceipt",
            entity_id=receipt.id,
            details={
                "delivery_id": delivery.id,
                "actual_received_date": payload.actual_received_date,
                "reconciliation_month": reconciliation_month,
                "item_count": len(delivery_items),
                "items": audit_items,
            },
            description="确认客户送货回单",
        )
        created_reminders: list[dict] = []
        if payload.reminders:
            assert reminder_bundle_key is not None
            assert reminder_bundle_hash is not None
            for reminder_draft in payload.reminders:
                reminder = create_reminder(
                    db,
                    receipt=receipt,
                    value=reminder_draft.model_dump(),
                    actor=user,
                    idempotency_key=reminder_bundle_key,
                )
                created_reminders.append(
                    {
                        "id": reminder.id,
                        "version": reminder.version,
                        "scope_type": reminder.scope_type,
                        "reminder_type": reminder.reminder_type,
                    }
                )
        response = _receipt_response(db, receipt.id)
        response["created_reminders"] = created_reminders
        if reminder_bundle_key:
            record_mutation(
                db,
                reminder_id=None,
                return_receipt_id=receipt.id,
                idempotency_key=reminder_bundle_key,
                request_hash_value=reminder_bundle_hash,
                action="receipt_create_bundle",
                actor=user,
                response=response,
            )
        _record_finance_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action="return_receipt_create",
            actor=user,
            resource_type="return_receipt",
            resource_id=receipt.id,
            response=response,
        )
        db.commit()
        return response
    except FulfillmentReminderError as error:
        db.rollback()
        _raise_fulfillment_reminder_error(error)
    except HTTPException:
        db.rollback()
        raise
    except IntegrityError as error:
        db.rollback()
        replay, replay_record = _finance_idempotency_replay(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action="return_receipt_create",
            actor=user,
        )
        if replay is not None:
            assert replay_record is not None
            _return_receipt_for_user(db, replay_record.resource_id, user)
            return replay
        raise HTTPException(status_code=409, detail="该送货单已经提交回单") from error
    except WarehouseInventoryError as error:
        db.rollback()
        raise HTTPException(
            status_code=error.status_code,
            detail=str(error),
        ) from error
    except Exception:
        db.rollback()
        raise


@router.put("/return_receipts/{receipt_id}")
def update_return_receipt(
    receipt_id: int,
    payload: ReturnReceiptUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _period_permission: User = Depends(can_adjust_reconciliation_period),
    _write_guard: None = Depends(fulfillment_reminder_write_guard),
) -> dict:
    del _period_permission
    del _write_guard
    receipt = _return_receipt_for_user(db, receipt_id, user)
    target_month = (
        payload.reconciliation_month
        if payload.reconciliation_month is not None
        else receipt.reconciliation_month
    )
    request_value = payload.model_dump(exclude={"idempotency_key"})
    request_value["reconciliation_month"] = target_month
    request_hash = _finance_request_hash(
        "return_receipt_update",
        {"receipt_id": receipt_id, **request_value},
    )
    replay, replay_record = _finance_idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="return_receipt_update",
        actor=user,
    )
    if replay is not None:
        assert replay_record is not None
        _return_receipt_for_user(db, replay_record.resource_id, user)
        return replay
    claimed_status = receipt.status
    if claimed_status not in {"confirmed", "cancelled"}:
        raise HTTPException(status_code=409, detail="回单状态不允许修改")
    expected_version = payload.expected_version or int(receipt.version or 1)
    _claim_return_receipt_version(
        db,
        receipt_id=receipt.id,
        expected_status=claimed_status,
        expected_version=expected_version,
    )
    db.expire(receipt)
    receipt_items = db.scalars(
        select(ReturnReceiptItem)
        .where(ReturnReceiptItem.return_receipt_id == receipt.id)
        .order_by(ReturnReceiptItem.id)
    ).all()
    receipt_item_ids = [item.id for item in receipt_items]
    if receipt_item_ids and db.scalar(
        select(StatementItem.id)
        .where(StatementItem.return_receipt_item_id.in_(receipt_item_ids))
        .limit(1)
    ) is not None:
        raise HTTPException(status_code=409, detail="回单已进入对账，禁止修改")
    delivery_items = db.scalars(
        select(DeliveryItem)
        .where(
            DeliveryItem.delivery_id == receipt.delivery_id,
            DeliveryItem.is_current.is_(True),
        )
        .order_by(DeliveryItem.id)
    ).all()
    requested = {line.delivery_item_id: line for line in payload.items}
    if set(requested) != {item.id for item in delivery_items}:
        raise HTTPException(status_code=400, detail="回单必须包含送货单全部明细")
    existing = {item.delivery_item_id: item for item in receipt_items}
    delivery_by_id = {item.id: item for item in delivery_items}
    _assert_no_later_dispatched_deliveries(
        db,
        receipt=receipt,
        receipt_items=receipt_items,
        delivery_items=delivery_by_id,
    )
    before = _receipt_response(db, receipt.id)
    affected_order_ids: set[int] = set()
    if claimed_status == "confirmed":
        try:
            reconsume_unordered_finished_receipt_returns(
                db,
                return_receipt_item_ids=receipt_item_ids,
                operator_id=user.id,
            )
            reconsume_ordered_finished_receipt_returns(
                db,
                return_receipt_item_ids=receipt_item_ids,
                operator_id=user.id,
            )
        except WarehouseInventoryError as error:
            db.rollback()
            raise HTTPException(
                status_code=error.status_code,
                detail=str(error),
            ) from error
        for previous in receipt_items:
            order_id = _apply_receipt_order_effect(
                db,
                delivery_item=delivery_by_id[previous.delivery_item_id],
                actual_received_quantity=previous.actual_received_quantity,
                resolution_action=previous.resolution_action,
                direction=-1,
            )
            if order_id is not None:
                affected_order_ids.add(order_id)
    for delivery_item in delivery_items:
        line = requested[delivery_item.id]
        action, reason = _validated_receipt_resolution(delivery_item, line)
        target = existing[delivery_item.id]
        target.actual_received_quantity = line.actual_received_quantity
        target.resolution_action = action
        target.difference_reason = reason
        if delivery_item.source_type == "unordered_finished":
            delivery = db.get(Delivery, delivery_item.delivery_id)
            if delivery is None:
                raise HTTPException(status_code=409, detail="回单关联送货单不存在")
            try:
                restore_unordered_finished_receipt_shortage(
                    db,
                    delivery=delivery,
                    delivery_item=delivery_item,
                    return_receipt_item_id=target.id,
                    actual_received_quantity=line.actual_received_quantity,
                    operator_id=user.id,
                    reason=reason,
                )
            except WarehouseInventoryError as error:
                db.rollback()
                raise HTTPException(
                    status_code=error.status_code,
                    detail=str(error),
                ) from error
        else:
            try:
                restore_ordered_finished_receipt_shortage(
                    db,
                    delivery=db.get(Delivery, delivery_item.delivery_id),
                    delivery_item=delivery_item,
                    return_receipt_item_id=target.id,
                    actual_received_quantity=line.actual_received_quantity,
                    resolution_action=action,
                    return_location_id=line.return_location_id,
                    expected_return_layout_version=(
                        line.expected_return_layout_version
                    ),
                    stock_date=payload.actual_received_date,
                    operator_id=user.id,
                )
            except WarehouseInventoryError as error:
                db.rollback()
                raise HTTPException(
                    status_code=error.status_code,
                    detail=str(error),
                ) from error
        order_id = _apply_receipt_order_effect(
            db,
            delivery_item=delivery_item,
            actual_received_quantity=line.actual_received_quantity,
            resolution_action=action,
            direction=1,
        )
        if order_id is not None:
            affected_order_ids.add(order_id)
    receipt.actual_received_date = payload.actual_received_date
    receipt.reconciliation_month = target_month
    receipt.signed_by = (payload.signed_by or "").strip() or None
    receipt.status = "confirmed"
    sync_receipt_source(
        db,
        receipt_id=receipt.id,
        source_valid=True,
        received_date=payload.actual_received_date,
        actor=user,
    )
    _refresh_receipt_order_statuses(db, affected_order_ids)
    _audit(
        db,
        user=user,
        action="UPDATE_RETURN_RECEIPT",
        resource="ReturnReceipt",
        entity_id=receipt.id,
        details={
            "before": before,
            "after": {
                **payload.model_dump(exclude={"idempotency_key"}),
                "reconciliation_month": target_month,
                "version": expected_version + 1,
            },
        },
        description="修改客户送货回单",
    )
    response = _receipt_response(db, receipt.id)
    _record_finance_idempotency(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="return_receipt_update",
        actor=user,
        resource_type="return_receipt",
        resource_id=receipt.id,
        response=response,
    )
    db.commit()
    return response


def _reconciliation_month_block_reason(
    db: Session,
    receipt_item_ids: list[int],
) -> str | None:
    if not receipt_item_ids:
        return None
    row = db.execute(
        select(
            Statement.statement_number,
            Statement.confirmation_status,
        )
        .join(StatementItem, StatementItem.statement_id == Statement.id)
        .where(StatementItem.return_receipt_item_id.in_(receipt_item_ids))
        .limit(1)
    ).first()
    if row is None:
        return None
    statement_number, confirmation_status = row
    stage = "已确认对账单" if confirmation_status == "confirmed" else "对账草稿"
    return (
        f"回单已进入{stage} {statement_number}，不能直接调整归属月份；"
        "请先按现有流程取消或撤销下游业务。"
    )


@router.put("/return_receipts/{receipt_id}/reconciliation-month")
def update_return_receipt_reconciliation_month(
    receipt_id: int,
    payload: ReturnReceiptReconciliationMonthUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _period_permission: User = Depends(can_adjust_reconciliation_period),
) -> dict:
    del _period_permission
    receipt = _return_receipt_for_user(db, receipt_id, user)
    request_hash = _finance_request_hash(
        "return_receipt_reconciliation_month_update",
        {
            "receipt_id": receipt_id,
            "reconciliation_month": payload.reconciliation_month,
            "expected_version": payload.expected_version,
        },
    )
    replay, replay_record = _finance_idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="return_receipt_reconciliation_month_update",
        actor=user,
    )
    if replay is not None:
        assert replay_record is not None
        _return_receipt_for_user(db, replay_record.resource_id, user)
        return replay

    receipt_item_ids = list(
        db.scalars(
            select(ReturnReceiptItem.id).where(
                ReturnReceiptItem.return_receipt_id == receipt.id
            )
        ).all()
    )
    block_reason = _reconciliation_month_block_reason(db, receipt_item_ids)
    if block_reason:
        raise HTTPException(
            status_code=409,
            detail={
                "code": "reconciliation_month_locked",
                "message": block_reason,
            },
        )
    before = _receipt_response(db, receipt.id)
    if (
        receipt.reconciliation_month == payload.reconciliation_month
        and int(receipt.version or 1) == payload.expected_version
    ):
        _record_finance_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action="return_receipt_reconciliation_month_update",
            actor=user,
            resource_type="return_receipt",
            resource_id=receipt.id,
            response=before,
        )
        db.commit()
        return before

    new_version = _claim_return_receipt_version(
        db,
        receipt_id=receipt.id,
        expected_status=receipt.status,
        expected_version=payload.expected_version,
    )
    db.expire(receipt)
    receipt.reconciliation_month = payload.reconciliation_month
    _audit(
        db,
        user=user,
        action="UPDATE_RETURN_RECONCILIATION_MONTH",
        resource="ReturnReceipt",
        entity_id=receipt.id,
        details={
            "before": {
                "reconciliation_month": before["reconciliation_month"],
                "effective_reconciliation_month": before[
                    "effective_reconciliation_month"
                ],
                "version": before["version"],
            },
            "after": {
                "reconciliation_month": payload.reconciliation_month,
                "version": new_version,
            },
        },
        description="调整客户回单对账归属月份",
    )
    response = _receipt_response(db, receipt.id)
    _record_finance_idempotency(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="return_receipt_reconciliation_month_update",
        actor=user,
        resource_type="return_receipt",
        resource_id=receipt.id,
        response=response,
    )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        replay, replay_record = _finance_idempotency_replay(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action="return_receipt_reconciliation_month_update",
            actor=user,
        )
        if replay is not None:
            assert replay_record is not None
            _return_receipt_for_user(db, replay_record.resource_id, user)
            return replay
        raise HTTPException(
            status_code=409,
            detail={
                "code": "finance_idempotency_conflict",
                "message": "归属月份已被其他操作修改，请刷新后重试",
            },
        ) from error
    return response


def _pending_statement_query(
    customer_id: int,
    *,
    statement_month: str | None = None,
    period_start: date | None = None,
    period_end: date | None = None,
):
    effective_unit_price = case(
        (
            DeliveryItem.source_type == "unordered_finished",
            DeliveryItem.unit_price_snapshot,
        ),
        else_=OrderItem.unit_price,
    )
    query = (
        select(
            ReturnReceiptItem.id.label("return_receipt_item_id"),
            ReturnReceipt.id.label("return_receipt_id"),
            ReturnReceipt.actual_received_date,
            ReturnReceipt.reconciliation_month,
            ReturnReceiptItem.reconciliation_month_override,
            Delivery.id.label("delivery_id"),
            Delivery.delivery_number,
            Delivery.delivery_date,
            Delivery.customer_id,
            Order.id.label("order_id"),
            case(
                (DeliveryItem.source_type == "unordered_finished", "无订单库存"),
                else_=Order.order_number,
            ).label("order_number"),
            case(
                (DeliveryItem.source_type == "unordered_finished", "无订单库存"),
                else_=Order.customer_po,
            ).label("customer_po"),
            func.coalesce(
                DeliveryItem.product_code_snapshot,
                Product.product_code,
            ).label("product_code"),
            func.coalesce(
                DeliveryItem.product_name_snapshot,
                OrderItem.snapshot_product_name,
                Product.product_name,
            ).label("product_name"),
            func.coalesce(
                DeliveryItem.specification_snapshot,
                OrderItem.snapshot_spec,
            ).label("specification"),
            Product.length_mm.label("product_length_mm"),
            Product.width_mm.label("product_width_mm"),
            Product.height_mm.label("product_height_mm"),
            DeliveryItem.ordered_quantity_snapshot,
            DeliveryItem.delivered_quantity.label("actual_delivery_quantity"),
            DeliveryItem.over_delivery_quantity,
            ReturnReceiptItem.actual_received_quantity,
            effective_unit_price.label("unit_price"),
            (
                ReturnReceiptItem.actual_received_quantity * effective_unit_price
            ).label("receivable_amount"),
            ReturnReceiptItem.difference_reason,
            StatementItem.id.label("statement_item_id"),
            StatementItem.statement_id,
        )
        .join(
            ReturnReceipt,
            ReturnReceipt.id == ReturnReceiptItem.return_receipt_id,
        )
        .join(DeliveryItem, DeliveryItem.id == ReturnReceiptItem.delivery_item_id)
        .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
        .outerjoin(OrderItem, OrderItem.id == DeliveryItem.order_item_id)
        .outerjoin(Order, Order.id == OrderItem.order_id)
        .outerjoin(
            Product,
            Product.id
            == func.coalesce(DeliveryItem.product_id, OrderItem.product_id),
        )
        .outerjoin(
            StatementItem,
            StatementItem.return_receipt_item_id == ReturnReceiptItem.id,
        )
        .where(
            Delivery.customer_id == customer_id,
            ReturnReceipt.status == "confirmed",
        )
        .order_by(
            Delivery.delivery_date,
            Delivery.delivery_number,
            ReturnReceiptItem.id,
        )
    )
    if statement_month is not None:
        if period_start is None or period_end is None:
            raise ValueError("对账月份筛选缺少客户周期边界")
        query = query.where(
            or_(
                ReturnReceiptItem.reconciliation_month_override == statement_month,
                and_(
                    ReturnReceiptItem.reconciliation_month_override.is_(None),
                    or_(
                        ReturnReceipt.reconciliation_month == statement_month,
                        and_(
                            ReturnReceipt.reconciliation_month.is_(None),
                            Delivery.delivery_date >= period_start,
                            Delivery.delivery_date <= period_end,
                        ),
                    ),
                ),
            )
        )
    else:
        if period_start is not None:
            query = query.where(Delivery.delivery_date >= period_start)
        if period_end is not None:
            query = query.where(Delivery.delivery_date <= period_end)
    return query


def _pending_statement_groups(
    db: Session,
    customer_id: int,
    *,
    statement_month: str | None = None,
    cycle_start_day: int = 1,
    period_start: date | None = None,
    period_end: date | None = None,
) -> list[dict]:
    raw_rows = [
        dict(row._mapping)
        for row in db.execute(
            _pending_statement_query(
                customer_id,
                statement_month=statement_month,
                period_start=period_start,
                period_end=period_end,
            )
        )
    ]
    if not raw_rows:
        return []

    registry = build_display_registry(db)
    order_ids = {
        row["order_id"]
        for row in raw_rows
        if row["order_id"] is not None
    }
    orders = {
        order.id: order
        for order in db.scalars(select(Order).where(Order.id.in_(order_ids))).all()
    }
    grouped: dict[int, dict] = {}
    for data in raw_rows:
        explicit_month = data.get("reconciliation_month_override") or data.get(
            "reconciliation_month"
        )
        effective_month = explicit_month or _historical_reconciliation_month(
            data["delivery_date"],
            cycle_start_day,
        )
        data["effective_reconciliation_month"] = effective_month
        data["reconciliation_month_source"] = (
            "line_override"
            if data.get("reconciliation_month_override")
            else ("explicit" if explicit_month else "historical_rule")
        )
        data["specification"] = resolved_product_specification(
            data.get("specification"),
            length_mm=data.pop("product_length_mm", None),
            width_mm=data.pop("product_width_mm", None),
            height_mm=data.pop("product_height_mm", None),
        )
        receivable = Decimal(str(data["receivable_amount"] or 0)).quantize(
            MONEY,
            rounding=ROUND_HALF_UP,
        )
        data["receivable_amount"] = receivable
        data["is_reconciled"] = data["statement_item_id"] is not None
        order = orders.get(data["order_id"])
        display = (
            display_order_number(order, registry)
            if order is not None
            else data["order_number"]
        )
        data["order_number"] = display
        data["display_order_number"] = display
        delivery = grouped.setdefault(
            data["delivery_id"],
            {
                "delivery_id": data["delivery_id"],
                "customer_id": data["customer_id"],
                "delivery_number": data["delivery_number"],
                "delivery_date": data["delivery_date"],
                "actual_received_date": data["actual_received_date"],
                "reconciliation_month": explicit_month,
                "effective_reconciliation_month": effective_month,
                "reconciliation_month_source": data[
                    "reconciliation_month_source"
                ],
                "return_receipt_id": data["return_receipt_id"],
                "items": [],
            },
        )
        delivery["items"].append(data)

    deliveries = []
    for delivery in grouped.values():
        rows = delivery["items"]
        reconciled_count = sum(1 for row in rows if row["is_reconciled"])
        if reconciled_count == len(rows):
            continue
        selection_blocked = reconciled_count > 0
        delivery.update(
            {
                "item_count": len(rows),
                "pending_item_count": len(rows) - reconciled_count,
                "total_received_quantity": sum(
                    int(row["actual_received_quantity"] or 0) for row in rows
                ),
                "total_receivable_amount": sum(
                    (row["receivable_amount"] for row in rows),
                    Decimal("0"),
                ).quantize(MONEY, rounding=ROUND_HALF_UP),
                "selection_blocked": selection_blocked,
                "exception_reason": (
                    "该送货单已有部分明细进入其他对账单，请先处理原对账单。"
                    if selection_blocked
                    else None
                ),
                "return_receipt_item_ids": [
                    row["return_receipt_item_id"]
                    for row in rows
                    if not row["is_reconciled"]
                ],
            }
        )
        deliveries.append(delivery)
    return deliveries


def pending_statement_customer_summaries(
    db: Session,
    *,
    statement_month: str,
    visible_customer_ids: set[int] | None,
) -> list[dict]:
    """Return one pending-reconciliation row per customer for a statement month.

    The finance workbench and dashboard share this query so customer-cycle
    boundaries, customer scope and unreconciled receipt identities cannot drift.
    All candidate deliveries are aggregated in one query; no per-customer query
    loop is used.
    """

    customer_query = select(
        Customer.id,
        Customer.name,
        Customer.statement_cycle_start_day,
    ).order_by(Customer.id)
    if visible_customer_ids is not None:
        if not visible_customer_ids:
            return []
        customer_query = customer_query.where(Customer.id.in_(visible_customer_ids))
    customer_rows = db.execute(customer_query).all()
    if not customer_rows:
        return []

    entity_cycle_days = {
        int(customer_id): int(cycle_day)
        for customer_id, cycle_day in db.execute(
            select(
                CustomerInvoiceProfile.customer_id,
                FinanceSettlementEntity.statement_cycle_start_day,
            )
            .join(
                FinanceSettlementEntity,
                FinanceSettlementEntity.id
                == CustomerInvoiceProfile.settlement_entity_id,
            )
            .where(
                CustomerInvoiceProfile.customer_id.in_(
                    [int(row[0]) for row in customer_rows]
                ),
                CustomerInvoiceProfile.is_enabled.is_(True),
                CustomerInvoiceProfile.confirmation_status == "confirmed",
                FinanceSettlementEntity.is_enabled.is_(True),
                FinanceSettlementEntity.confirmation_status == "confirmed",
            )
        ).all()
    }

    customer_periods: dict[int, tuple[date, date]] = {}
    customer_names: dict[int, str] = {}
    customer_cycle_days: dict[int, int] = {}
    for customer_id, customer_name, cycle_start_day in customer_rows:
        cycle_day = entity_cycle_days.get(
            int(customer_id), int(cycle_start_day or 1)
        )
        customer_periods[int(customer_id)] = _statement_period(
            statement_month,
            cycle_day,
        )
        customer_names[int(customer_id)] = customer_name
        customer_cycle_days[int(customer_id)] = cycle_day

    broad_start = min(period[0] for period in customer_periods.values())
    broad_end = max(period[1] for period in customer_periods.values())
    effective_unit_price = case(
        (
            DeliveryItem.source_type == "unordered_finished",
            DeliveryItem.unit_price_snapshot,
        ),
        else_=OrderItem.unit_price,
    )
    pending_item_count = func.sum(
        case((StatementItem.id.is_(None), 1), else_=0)
    )
    pending_amount = func.sum(
        case(
            (
                StatementItem.id.is_(None),
                ReturnReceiptItem.actual_received_quantity * effective_unit_price,
            ),
            else_=0,
        )
    )
    query = (
        select(
            Delivery.customer_id,
            Delivery.id.label("delivery_id"),
            Delivery.delivery_date,
            ReturnReceipt.reconciliation_month,
            ReturnReceiptItem.reconciliation_month_override,
            func.count(ReturnReceiptItem.id).label("item_count"),
            pending_item_count.label("pending_item_count"),
            func.coalesce(pending_amount, 0).label("pending_amount"),
        )
        .select_from(ReturnReceiptItem)
        .join(ReturnReceipt, ReturnReceipt.id == ReturnReceiptItem.return_receipt_id)
        .join(DeliveryItem, DeliveryItem.id == ReturnReceiptItem.delivery_item_id)
        .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
        .outerjoin(OrderItem, OrderItem.id == DeliveryItem.order_item_id)
        .outerjoin(
            StatementItem,
            StatementItem.return_receipt_item_id == ReturnReceiptItem.id,
        )
        .where(
            ReturnReceipt.status == "confirmed",
            or_(
                ReturnReceiptItem.reconciliation_month_override
                == statement_month,
                and_(
                    ReturnReceiptItem.reconciliation_month_override.is_(None),
                    or_(
                        ReturnReceipt.reconciliation_month == statement_month,
                        and_(
                            ReturnReceipt.reconciliation_month.is_(None),
                            Delivery.delivery_date >= broad_start,
                            Delivery.delivery_date <= broad_end,
                        ),
                    ),
                ),
            ),
            Delivery.customer_id.in_(tuple(customer_periods)),
        )
        .group_by(
            Delivery.customer_id,
            Delivery.id,
            Delivery.delivery_date,
            ReturnReceipt.reconciliation_month,
            ReturnReceiptItem.reconciliation_month_override,
        )
        .order_by(Delivery.customer_id, Delivery.delivery_date, Delivery.id)
    )

    summaries: dict[int, dict] = {}
    for row in db.execute(query).mappings():
        customer_id = int(row["customer_id"])
        period_start, period_end = customer_periods[customer_id]
        delivery_date = row["delivery_date"]
        effective_month = (
            row["reconciliation_month_override"]
            or row["reconciliation_month"]
            or _historical_reconciliation_month(
                delivery_date,
                customer_cycle_days[customer_id],
            )
        )
        if effective_month != statement_month:
            continue
        remaining_count = int(row["pending_item_count"] or 0)
        if remaining_count <= 0:
            continue
        item_count = int(row["item_count"] or 0)
        summary = summaries.setdefault(
            customer_id,
            {
                "customer_id": customer_id,
                "customer_name": customer_names[customer_id],
                "statement_month": statement_month,
                "statement_cycle_start_day": customer_cycle_days[customer_id],
                "period_start": period_start,
                "period_end": period_end,
                "pending_count": 0,
                "blocked_count": 0,
                "pending_item_count": 0,
                "amount": Decimal("0.00"),
                "delivery_ids": [],
            },
        )
        if remaining_count < item_count:
            summary["blocked_count"] += 1
        else:
            summary["pending_count"] += 1
        summary["pending_item_count"] += remaining_count
        summary["amount"] = (
            summary["amount"] + Decimal(str(row["pending_amount"] or 0))
        ).quantize(MONEY, rounding=ROUND_HALF_UP)
        summary["delivery_ids"].append(int(row["delivery_id"]))

    charge_query = (
        select(
            CustomerCharge.customer_id,
            func.count(CustomerCharge.id).label("charge_count"),
            func.coalesce(func.sum(CustomerCharge.amount), 0).label("charge_amount"),
        )
        .outerjoin(
            StatementItem,
            StatementItem.customer_charge_id == CustomerCharge.id,
        )
        .where(
            CustomerCharge.status == "confirmed",
            CustomerCharge.reconciliation_month == statement_month,
            StatementItem.id.is_(None),
            CustomerCharge.customer_id.in_(tuple(customer_periods)),
        )
        .group_by(CustomerCharge.customer_id)
    )
    for row in db.execute(charge_query).mappings():
        customer_id = int(row["customer_id"])
        period_start, period_end = customer_periods[customer_id]
        summary = summaries.setdefault(
            customer_id,
            {
                "customer_id": customer_id,
                "customer_name": customer_names[customer_id],
                "statement_month": statement_month,
                "statement_cycle_start_day": customer_cycle_days[customer_id],
                "period_start": period_start,
                "period_end": period_end,
                "pending_count": 0,
                "blocked_count": 0,
                "pending_item_count": 0,
                "amount": Decimal("0.00"),
                "delivery_ids": [],
                "customer_charge_count": 0,
            },
        )
        charge_count = int(row["charge_count"] or 0)
        summary["pending_count"] += charge_count
        summary["pending_item_count"] += charge_count
        summary["customer_charge_count"] = charge_count
        summary["amount"] = (
            summary["amount"] + Decimal(str(row["charge_amount"] or 0))
        ).quantize(MONEY, rounding=ROUND_HALF_UP)

    return list(summaries.values())


def _money_value(value: object) -> Decimal:
    return Decimal(str(value or 0)).quantize(MONEY, rounding=ROUND_HALF_UP)


def _finance_report_months(year: int) -> list[str]:
    return [f"{year:04d}-{month_number:02d}" for month_number in range(1, 13)]


def _finance_report_selected_month(
    year: int,
    statement_month: str | None,
) -> str:
    if statement_month:
        try:
            _statement_period(statement_month, 1)
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        if int(statement_month[:4]) != year:
            raise HTTPException(status_code=400, detail="月份与年份不一致")
        return statement_month
    today = beijing_today()
    return (
        today.strftime("%Y-%m")
        if today.year == year
        else f"{year:04d}-12"
    )


def _empty_finance_report_month(statement_month: str) -> dict:
    return {
        "statement_month": statement_month,
        "pending_reconciliation_customer_count": 0,
        "pending_reconciliation_item_count": 0,
        "pending_reconciliation_amount": Decimal("0.00"),
        "reconciled_customer_count": 0,
        "reconciled_receivable_amount": Decimal("0.00"),
        "pending_invoice_customer_count": 0,
        "pending_invoice_amount": Decimal("0.00"),
        "invoiced_customer_count": 0,
        "invoiced_amount": Decimal("0.00"),
        "pending_payment_customer_count": 0,
        "pending_payment_amount": Decimal("0.00"),
        "settled_customer_count": 0,
        "settled_amount": Decimal("0.00"),
        "statement_count": 0,
        "status_anomaly_count": 0,
    }


def _month_distance(selected_month: str, source_month: str) -> int:
    selected_year, selected_number = (int(part) for part in selected_month.split("-"))
    source_year, source_number = (int(part) for part in source_month.split("-"))
    return (selected_year - source_year) * 12 + selected_number - source_number


@router.get("/reports/monthly-yearly")
def monthly_yearly_finance_report(
    year: int | None = Query(default=None, ge=2000, le=2100),
    statement_month: str | None = None,
    customer_id: int | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    """Return one read-only report over the existing finance facts.

    The report deliberately keeps unreconciled receipts separate from formal
    statements.  Outstanding and aging values use each statement's real
    remaining balance and classify it by statement month; no second ledger or
    historical balance snapshot is created.
    """

    selected_year = year or beijing_today().year
    selected_month = _finance_report_selected_month(
        selected_year,
        statement_month,
    )
    report_months = _finance_report_months(selected_year)
    visible_customer_ids = _visible_customer_ids(user, db)
    if customer_id is not None:
        require_customer_access(customer_id, user, db)
        visible_customer_ids = {customer_id}

    monthly_by_month = {
        report_month: _empty_finance_report_month(report_month)
        for report_month in report_months
    }
    empty_aging = [
        {
            "key": key,
            "label": label,
            "outstanding_amount": Decimal("0.00"),
            "customer_count": 0,
            "statement_month_count": 0,
        }
        for key, label in (
            ("current", "当月"),
            ("one_month", "1 个月"),
            ("two_months", "2 个月"),
            ("three_plus", "3 个月及以上"),
        )
    ]
    if visible_customer_ids is not None and not visible_customer_ids:
        return {
            "year": selected_year,
            "selected_month": selected_month,
            "as_of": beijing_today(),
            "included_through_statement_month": selected_month,
            "balance_basis": "current_balance_by_statement_month",
            "monthly": list(monthly_by_month.values()),
            "yearly_summary": {
                "customer_count": 0,
                "record_count": 0,
                "pending_reconciliation_amount": Decimal("0.00"),
                "reconciled_receivable_amount": Decimal("0.00"),
                "invoiced_amount": Decimal("0.00"),
                "settled_amount": Decimal("0.00"),
                "pending_payment_amount": Decimal("0.00"),
            },
            "selected_month_summary": monthly_by_month[selected_month],
            "selected_month_customers": [],
            "customer_balances": [],
            "aging": empty_aging,
        }

    pending_invoice_expression = case(
        (
            Statement.total_receivable > Statement.invoiced_amount,
            Statement.total_receivable - Statement.invoiced_amount,
        ),
        else_=0,
    )
    pending_payment_expression = case(
        (
            Statement.total_receivable > Statement.settled_amount,
            Statement.total_receivable - Statement.settled_amount,
        ),
        else_=0,
    )
    status_anomaly_expression = case(
        (
            and_(
                Statement.status == "settled",
                Statement.settled_amount < Statement.total_receivable,
            ),
            1,
        ),
        (
            and_(
                Statement.status != "settled",
                Statement.settled_amount >= Statement.total_receivable,
            ),
            1,
        ),
        else_=0,
    )
    statement_group_query = (
        select(
            Statement.customer_id,
            Customer.name.label("customer_name"),
            Statement.statement_month,
            func.count(Statement.id).label("statement_count"),
            func.sum(Statement.total_receivable).label("total_receivable"),
            func.sum(Statement.invoiced_amount).label("invoiced_amount"),
            func.sum(Statement.settled_amount).label("settled_amount"),
            func.sum(pending_invoice_expression).label("pending_invoice_amount"),
            func.sum(pending_payment_expression).label("pending_payment_amount"),
            func.sum(status_anomaly_expression).label("status_anomaly_count"),
        )
        .join(Customer, Customer.id == Statement.customer_id)
        .where(Statement.statement_month.in_(report_months))
        .group_by(
            Statement.customer_id,
            Customer.name,
            Statement.statement_month,
        )
        .order_by(Statement.statement_month, Customer.name, Statement.customer_id)
    )
    if visible_customer_ids is not None:
        statement_group_query = statement_group_query.where(
            Statement.customer_id.in_(visible_customer_ids)
        )
    statement_groups = [
        {
            "customer_id": int(row["customer_id"]),
            "customer_name": row["customer_name"],
            "statement_month": str(row["statement_month"]),
            "statement_count": int(row["statement_count"] or 0),
            "reconciled_receivable_amount": _money_value(row["total_receivable"]),
            "invoiced_amount": _money_value(row["invoiced_amount"]),
            "settled_amount": _money_value(row["settled_amount"]),
            "pending_invoice_amount": _money_value(row["pending_invoice_amount"]),
            "pending_payment_amount": _money_value(row["pending_payment_amount"]),
            "status_anomaly_count": int(row["status_anomaly_count"] or 0),
        }
        for row in db.execute(statement_group_query).mappings()
    ]

    annual_customer_ids: set[int] = set()
    for group in statement_groups:
        month_row = monthly_by_month[group["statement_month"]]
        group_customer_id = int(group["customer_id"])
        annual_customer_ids.add(group_customer_id)
        month_row["statement_count"] += group["statement_count"]
        month_row["status_anomaly_count"] += group["status_anomaly_count"]
        month_row["reconciled_receivable_amount"] = _money_value(
            month_row["reconciled_receivable_amount"]
            + group["reconciled_receivable_amount"]
        )
        month_row["invoiced_amount"] = _money_value(
            month_row["invoiced_amount"] + group["invoiced_amount"]
        )
        month_row["settled_amount"] = _money_value(
            month_row["settled_amount"] + group["settled_amount"]
        )
        month_row["pending_invoice_amount"] = _money_value(
            month_row["pending_invoice_amount"] + group["pending_invoice_amount"]
        )
        month_row["pending_payment_amount"] = _money_value(
            month_row["pending_payment_amount"] + group["pending_payment_amount"]
        )
        if group["reconciled_receivable_amount"] > 0:
            month_row["reconciled_customer_count"] += 1
        if group["invoiced_amount"] > 0:
            month_row["invoiced_customer_count"] += 1
        if group["settled_amount"] > 0:
            month_row["settled_customer_count"] += 1
        if group["pending_invoice_amount"] > 0:
            month_row["pending_invoice_customer_count"] += 1
        if group["pending_payment_amount"] > 0:
            month_row["pending_payment_customer_count"] += 1

    pending_by_month_customer: dict[tuple[str, int], dict] = {}
    for report_month in report_months:
        pending_rows = pending_statement_customer_summaries(
            db,
            statement_month=report_month,
            visible_customer_ids=visible_customer_ids,
        )
        month_row = monthly_by_month[report_month]
        for pending in pending_rows:
            pending_count = int(pending["pending_count"] or 0)
            blocked_count = int(pending["blocked_count"] or 0)
            if pending_count + blocked_count <= 0:
                continue
            pending_customer_id = int(pending["customer_id"])
            annual_customer_ids.add(pending_customer_id)
            month_row["pending_reconciliation_customer_count"] += 1
            month_row["pending_reconciliation_item_count"] += int(
                pending["pending_item_count"] or 0
            )
            month_row["pending_reconciliation_amount"] = _money_value(
                month_row["pending_reconciliation_amount"]
                + _money_value(pending["amount"])
            )
            pending_by_month_customer[(report_month, pending_customer_id)] = pending

    selected_customer_rows: dict[int, dict] = {}
    for group in statement_groups:
        if group["statement_month"] != selected_month:
            continue
        selected_customer_rows[int(group["customer_id"])] = {
            **group,
            "pending_reconciliation_count": 0,
            "blocked_reconciliation_count": 0,
            "pending_reconciliation_item_count": 0,
            "pending_reconciliation_amount": Decimal("0.00"),
        }
    for (pending_month, pending_customer_id), pending in pending_by_month_customer.items():
        if pending_month != selected_month:
            continue
        row = selected_customer_rows.setdefault(
            pending_customer_id,
            {
                "customer_id": pending_customer_id,
                "customer_name": pending["customer_name"],
                "statement_month": selected_month,
                "statement_count": 0,
                "reconciled_receivable_amount": Decimal("0.00"),
                "invoiced_amount": Decimal("0.00"),
                "settled_amount": Decimal("0.00"),
                "pending_invoice_amount": Decimal("0.00"),
                "pending_payment_amount": Decimal("0.00"),
                "status_anomaly_count": 0,
                "pending_reconciliation_count": 0,
                "blocked_reconciliation_count": 0,
                "pending_reconciliation_item_count": 0,
                "pending_reconciliation_amount": Decimal("0.00"),
            },
        )
        row["pending_reconciliation_count"] = int(pending["pending_count"] or 0)
        row["blocked_reconciliation_count"] = int(pending["blocked_count"] or 0)
        row["pending_reconciliation_item_count"] = int(
            pending["pending_item_count"] or 0
        )
        row["pending_reconciliation_amount"] = _money_value(pending["amount"])
    selected_month_customers = sorted(
        selected_customer_rows.values(),
        key=lambda row: (str(row["customer_name"]), int(row["customer_id"])),
    )

    outstanding_group_query = (
        select(
            Statement.customer_id,
            Customer.name.label("customer_name"),
            Statement.statement_month,
            func.count(Statement.id).label("statement_count"),
            func.sum(Statement.total_receivable).label("total_receivable"),
            func.sum(Statement.invoiced_amount).label("invoiced_amount"),
            func.sum(Statement.settled_amount).label("settled_amount"),
            func.sum(pending_payment_expression).label("outstanding_amount"),
        )
        .join(Customer, Customer.id == Statement.customer_id)
        .where(Statement.statement_month <= selected_month)
        .group_by(
            Statement.customer_id,
            Customer.name,
            Statement.statement_month,
        )
        .having(func.sum(pending_payment_expression) > 0)
        .order_by(Customer.name, Statement.customer_id, Statement.statement_month)
    )
    if visible_customer_ids is not None:
        outstanding_group_query = outstanding_group_query.where(
            Statement.customer_id.in_(visible_customer_ids)
        )
    outstanding_groups = [dict(row) for row in db.execute(outstanding_group_query).mappings()]

    customer_balances_by_id: dict[int, dict] = {}
    aging_by_key = {
        row["key"]: {**row, "_customer_ids": set()}
        for row in empty_aging
    }
    for outstanding in outstanding_groups:
        outstanding_customer_id = int(outstanding["customer_id"])
        source_month = str(outstanding["statement_month"])
        outstanding_amount = _money_value(outstanding["outstanding_amount"])
        balance = customer_balances_by_id.setdefault(
            outstanding_customer_id,
            {
                "customer_id": outstanding_customer_id,
                "customer_name": outstanding["customer_name"],
                "outstanding_amount": Decimal("0.00"),
                "total_receivable": Decimal("0.00"),
                "invoiced_amount": Decimal("0.00"),
                "settled_amount": Decimal("0.00"),
                "statement_count": 0,
                "outstanding_month_count": 0,
                "earliest_statement_month": source_month,
                "latest_statement_month": source_month,
                "months": [],
            },
        )
        balance["outstanding_amount"] = _money_value(
            balance["outstanding_amount"] + outstanding_amount
        )
        balance["total_receivable"] = _money_value(
            balance["total_receivable"] + _money_value(outstanding["total_receivable"])
        )
        balance["invoiced_amount"] = _money_value(
            balance["invoiced_amount"] + _money_value(outstanding["invoiced_amount"])
        )
        balance["settled_amount"] = _money_value(
            balance["settled_amount"] + _money_value(outstanding["settled_amount"])
        )
        balance["statement_count"] += int(outstanding["statement_count"] or 0)
        balance["outstanding_month_count"] += 1
        balance["earliest_statement_month"] = min(
            balance["earliest_statement_month"], source_month
        )
        balance["latest_statement_month"] = max(
            balance["latest_statement_month"], source_month
        )
        balance["months"].append(
            {
                "statement_month": source_month,
                "statement_count": int(outstanding["statement_count"] or 0),
                "outstanding_amount": outstanding_amount,
            }
        )

        distance = _month_distance(selected_month, source_month)
        if distance <= 0:
            aging_key = "current"
        elif distance == 1:
            aging_key = "one_month"
        elif distance == 2:
            aging_key = "two_months"
        else:
            aging_key = "three_plus"
        aging_row = aging_by_key[aging_key]
        aging_row["outstanding_amount"] = _money_value(
            aging_row["outstanding_amount"] + outstanding_amount
        )
        aging_row["statement_month_count"] += 1
        aging_row["_customer_ids"].add(outstanding_customer_id)

    customer_balances = sorted(
        customer_balances_by_id.values(),
        key=lambda row: (
            -row["outstanding_amount"],
            str(row["customer_name"]),
            int(row["customer_id"]),
        ),
    )
    aging = []
    for key in ("current", "one_month", "two_months", "three_plus"):
        aging_row = aging_by_key[key]
        customer_ids = aging_row.pop("_customer_ids")
        aging_row["customer_count"] = len(customer_ids)
        aging.append(aging_row)

    monthly = list(monthly_by_month.values())
    yearly_summary = {
        "customer_count": len(annual_customer_ids),
        "record_count": sum(
            int(row["statement_count"])
            + int(row["pending_reconciliation_item_count"])
            for row in monthly
        ),
        "pending_reconciliation_amount": _money_value(
            sum((row["pending_reconciliation_amount"] for row in monthly), Decimal("0.00"))
        ),
        "reconciled_receivable_amount": _money_value(
            sum((row["reconciled_receivable_amount"] for row in monthly), Decimal("0.00"))
        ),
        "invoiced_amount": _money_value(
            sum((row["invoiced_amount"] for row in monthly), Decimal("0.00"))
        ),
        "settled_amount": _money_value(
            sum((row["settled_amount"] for row in monthly), Decimal("0.00"))
        ),
        "pending_payment_amount": _money_value(
            sum((row["pending_payment_amount"] for row in monthly), Decimal("0.00"))
        ),
    }
    return {
        "year": selected_year,
        "selected_month": selected_month,
        "as_of": beijing_today(),
        "included_through_statement_month": selected_month,
        "balance_basis": "current_balance_by_statement_month",
        "monthly": monthly,
        "yearly_summary": yearly_summary,
        "selected_month_summary": monthly_by_month[selected_month],
        "selected_month_customers": selected_month_customers,
        "customer_balances": customer_balances,
        "aging": aging,
    }


@router.get("/current-customer-months")
def current_customer_months(
    statement_month: str | None = None,
    balance_type: str | None = None,
    customer_id: int | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    """Return the current finance workbench grouped by customer and month.

    A customer can have several statements in the same month.  The grouping is
    therefore performed before pagination and uses the stable customer id, not
    the display name or the possibly stale statement status.
    """

    selected_month = statement_month or beijing_today().strftime("%Y-%m")
    try:
        _statement_period(selected_month, 1)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    allowed_balance_types = {
        None,
        "pending_reconciliation",
        "pending_invoice",
        "pending_payment",
    }
    if balance_type not in allowed_balance_types:
        raise HTTPException(status_code=400, detail="未知的财务待办筛选")

    visible_customer_ids = _visible_customer_ids(user, db)
    if customer_id is not None:
        require_customer_access(customer_id, user, db)
    if visible_customer_ids is not None and not visible_customer_ids:
        return {
            "statement_month": selected_month,
            "as_of": beijing_today(),
            "total": 0,
            "page": page,
            "page_size": page_size,
            "summary": {
                "customer_count": 0,
                "pending_reconciliation_amount": Decimal("0.00"),
                "reconciled_receivable_amount": Decimal("0.00"),
                "pending_invoice_amount": Decimal("0.00"),
                "invoiced_amount": Decimal("0.00"),
                "pending_payment_amount": Decimal("0.00"),
                "settled_amount": Decimal("0.00"),
            },
            "customer_options": [],
            "items": [],
        }

    try:
        pending_rows = pending_statement_customer_summaries(
            db,
            statement_month=selected_month,
            visible_customer_ids=visible_customer_ids,
        )
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error

    statement_query = (
        select(
            Statement.id,
            Statement.statement_number,
            Statement.customer_id,
            Customer.name.label("customer_name"),
            Customer.statement_cycle_start_day,
            Statement.statement_month,
            Statement.total_receivable,
            Statement.invoiced_amount,
            Statement.settled_amount,
            Statement.status,
            Statement.confirmation_status,
            Statement.version,
            Statement.ledger_version,
            Statement.created_at,
        )
        .join(Customer, Customer.id == Statement.customer_id)
        .where(Statement.statement_month == selected_month)
        .order_by(Customer.name, Customer.id, Statement.id.desc())
    )
    if visible_customer_ids is not None:
        statement_query = statement_query.where(
            Statement.customer_id.in_(visible_customer_ids)
        )
    statement_rows = db.execute(statement_query).mappings().all()

    grouped: dict[int, dict] = {}
    for pending in pending_rows:
        pending_customer_id = int(pending["customer_id"])
        grouped[pending_customer_id] = {
            "customer_id": pending_customer_id,
            "customer_name": pending["customer_name"],
            "statement_month": selected_month,
            "statement_cycle_start_day": int(
                pending["statement_cycle_start_day"] or 1
            ),
            "period_start": pending["period_start"],
            "period_end": pending["period_end"],
            "pending_reconciliation_count": int(pending["pending_count"] or 0),
            "blocked_reconciliation_count": int(pending["blocked_count"] or 0),
            "pending_reconciliation_item_count": int(
                pending["pending_item_count"] or 0
            ),
            "pending_reconciliation_amount": _money_value(pending["amount"]),
            "reconciled_receivable_amount": Decimal("0.00"),
            "invoiced_amount": Decimal("0.00"),
            "settled_amount": Decimal("0.00"),
            "pending_invoice_amount": Decimal("0.00"),
            "pending_payment_amount": Decimal("0.00"),
            "statement_count": 0,
            "status_anomaly_count": 0,
            "statements": [],
        }

    for statement in statement_rows:
        statement_customer_id = int(statement["customer_id"])
        cycle_day = int(statement["statement_cycle_start_day"] or 1)
        period_start, period_end = _statement_period(selected_month, cycle_day)
        group = grouped.setdefault(
            statement_customer_id,
            {
                "customer_id": statement_customer_id,
                "customer_name": statement["customer_name"],
                "statement_month": selected_month,
                "statement_cycle_start_day": cycle_day,
                "period_start": period_start,
                "period_end": period_end,
                "pending_reconciliation_count": 0,
                "blocked_reconciliation_count": 0,
                "pending_reconciliation_item_count": 0,
                "pending_reconciliation_amount": Decimal("0.00"),
                "reconciled_receivable_amount": Decimal("0.00"),
                "invoiced_amount": Decimal("0.00"),
                "settled_amount": Decimal("0.00"),
                "pending_invoice_amount": Decimal("0.00"),
                "pending_payment_amount": Decimal("0.00"),
                "statement_count": 0,
                "status_anomaly_count": 0,
                "statements": [],
            },
        )
        is_confirmed = statement["confirmation_status"] == "confirmed"
        receivable = _money_value(statement["total_receivable"])
        invoiced = _money_value(statement["invoiced_amount"]) if is_confirmed else Decimal("0.00")
        settled = _money_value(statement["settled_amount"]) if is_confirmed else Decimal("0.00")
        confirmed_receivable = receivable if is_confirmed else Decimal("0.00")
        invoice_balance = max(receivable - invoiced, Decimal("0.00"))
        payment_balance = max(receivable - settled, Decimal("0.00"))
        if not is_confirmed:
            invoice_balance = Decimal("0.00")
            payment_balance = Decimal("0.00")
        status_is_settled = statement["status"] == "settled"
        balance_is_settled = payment_balance == Decimal("0.00")
        if is_confirmed and status_is_settled != balance_is_settled:
            group["status_anomaly_count"] += 1
        group["statement_count"] += 1
        group["reconciled_receivable_amount"] = _money_value(
            group["reconciled_receivable_amount"] + confirmed_receivable
        )
        group["invoiced_amount"] = _money_value(
            group["invoiced_amount"] + invoiced
        )
        group["settled_amount"] = _money_value(
            group["settled_amount"] + settled
        )
        group["pending_invoice_amount"] = _money_value(
            group["pending_invoice_amount"] + invoice_balance
        )
        group["pending_payment_amount"] = _money_value(
            group["pending_payment_amount"] + payment_balance
        )
        group["statements"].append(
            {
                "id": int(statement["id"]),
                "statement_number": statement["statement_number"],
                "customer_id": statement_customer_id,
                "customer_name": statement["customer_name"],
                "statement_month": statement["statement_month"],
                "total_receivable": receivable,
                "invoiced_amount": invoiced,
                "settled_amount": settled,
                "pending_invoice_amount": invoice_balance,
                "pending_payment_amount": payment_balance,
                "status": statement["status"],
                "confirmation_status": statement["confirmation_status"],
                "version": int(statement["version"]),
                "ledger_version": int(statement["ledger_version"]),
                "invoice_status": (
                    "not_ready"
                    if not is_confirmed
                    else "invoiced"
                    if receivable > Decimal("0.00")
                    and invoice_balance == Decimal("0.00")
                    else "partial"
                    if invoiced > Decimal("0.00")
                    else "pending"
                ),
                "created_at": statement["created_at"],
            }
        )

    eligible_items: list[dict] = []
    for group in grouped.values():
        has_reconciliation = (
            group["pending_reconciliation_count"]
            + group["blocked_reconciliation_count"]
            > 0
        )
        has_invoice = group["pending_invoice_amount"] > 0
        has_payment = group["pending_payment_amount"] > 0
        if not (has_reconciliation or has_invoice or has_payment):
            continue
        if has_reconciliation:
            primary_action = "reconcile"
        elif has_invoice:
            primary_action = "invoice"
        else:
            primary_action = "payment"
        group["primary_action"] = primary_action
        group["primary_action_blocked"] = bool(
            primary_action == "reconcile"
            and group["pending_reconciliation_count"] == 0
            and group["blocked_reconciliation_count"] > 0
        )
        eligible_items.append(group)

    customer_options = [
        {"id": int(row["customer_id"]), "name": str(row["customer_name"])}
        for row in sorted(
            eligible_items,
            key=lambda row: (str(row["customer_name"]), int(row["customer_id"])),
        )
    ]
    items = [
        row
        for row in eligible_items
        if (customer_id is None or int(row["customer_id"]) == customer_id)
        and (
            balance_type is None
            or (
                balance_type == "pending_reconciliation"
                and (
                    row["pending_reconciliation_count"]
                    + row["blocked_reconciliation_count"]
                    > 0
                )
            )
            or (
                balance_type == "pending_invoice"
                and row["pending_invoice_amount"] > 0
            )
            or (
                balance_type == "pending_payment"
                and row["pending_payment_amount"] > 0
            )
        )
    ]

    action_rank = {"reconcile": 0, "invoice": 1, "payment": 2}
    items.sort(
        key=lambda row: (
            action_rank[row["primary_action"]],
            str(row["customer_name"]),
            int(row["customer_id"]),
        )
    )
    total = len(items)
    summary = {
        "customer_count": total,
        "pending_reconciliation_amount": _money_value(
            sum(
                (row["pending_reconciliation_amount"] for row in items),
                Decimal("0.00"),
            )
        ),
        "reconciled_receivable_amount": _money_value(
            sum(
                (row["reconciled_receivable_amount"] for row in items),
                Decimal("0.00"),
            )
        ),
        "pending_invoice_amount": _money_value(
            sum(
                (row["pending_invoice_amount"] for row in items),
                Decimal("0.00"),
            )
        ),
        "invoiced_amount": _money_value(
            sum((row["invoiced_amount"] for row in items), Decimal("0.00"))
        ),
        "pending_payment_amount": _money_value(
            sum(
                (row["pending_payment_amount"] for row in items),
                Decimal("0.00"),
            )
        ),
        "settled_amount": _money_value(
            sum((row["settled_amount"] for row in items), Decimal("0.00"))
        ),
    }
    start = (page - 1) * page_size
    return {
        "statement_month": selected_month,
        "as_of": beijing_today(),
        "total": total,
        "page": page,
        "page_size": page_size,
        "summary": summary,
        "customer_options": customer_options,
        "items": items[start : start + page_size],
    }


@router.get("/settled-customer-months")
def settled_customer_months(
    year: int | None = Query(default=None, ge=2000, le=2100),
    statement_month: str | None = None,
    customer_id: int | None = None,
    statement_number: str | None = Query(default=None, max_length=80),
    invoice_number: str | None = Query(default=None, max_length=120),
    settlement_date_from: date | None = None,
    settlement_date_to: date | None = None,
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=50, ge=1, le=200),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    """Return fully processed finance history by year, month and customer.

    History is a read-only view over the existing statement, invoice and
    settlement facts.  A customer-month enters history only when every
    statement has no invoice or payment balance and no return receipt remains
    pending reconciliation for the same customer-month.
    """

    selected_year = year or beijing_today().year
    if statement_month:
        try:
            _statement_period(statement_month, 1)
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        if int(statement_month[:4]) != selected_year:
            raise HTTPException(status_code=400, detail="月份与年份不一致")
    if (
        settlement_date_from is not None
        and settlement_date_to is not None
        and settlement_date_from > settlement_date_to
    ):
        raise HTTPException(status_code=400, detail="收款开始日期不能晚于结束日期")

    statement_keyword = (statement_number or "").strip()
    invoice_keyword = (invoice_number or "").strip()
    visible_customer_ids = _visible_customer_ids(user, db)
    if customer_id is not None:
        require_customer_access(customer_id, user, db)
        visible_customer_ids = {customer_id}

    empty_summary = {
        "customer_month_count": 0,
        "statement_count": 0,
        "total_receivable": Decimal("0.00"),
        "invoiced_amount": Decimal("0.00"),
        "settled_amount": Decimal("0.00"),
    }
    if visible_customer_ids is not None and not visible_customer_ids:
        return {
            "year": selected_year,
            "statement_month": statement_month,
            "as_of": beijing_today(),
            "total": 0,
            "page": page,
            "page_size": page_size,
            "summary": empty_summary,
            "items": [],
        }

    status_anomaly = case(
        (
            and_(
                Statement.status == "settled",
                Statement.settled_amount < Statement.total_receivable,
            ),
            1,
        ),
        (
            and_(
                Statement.status != "settled",
                Statement.settled_amount >= Statement.total_receivable,
            ),
            1,
        ),
        else_=0,
    )
    selected_year_months = [
        f"{selected_year:04d}-{month_number:02d}"
        for month_number in range(1, 13)
    ]
    group_query = (
        select(
            Statement.customer_id,
            Customer.name.label("customer_name"),
            Statement.statement_month,
            func.count(Statement.id).label("statement_count"),
            func.sum(Statement.total_receivable).label("total_receivable"),
            func.sum(Statement.invoiced_amount).label("invoiced_amount"),
            func.sum(Statement.settled_amount).label("settled_amount"),
            func.sum(status_anomaly).label("status_anomaly_count"),
        )
        .join(Customer, Customer.id == Statement.customer_id)
        .where(Statement.statement_month.in_(selected_year_months))
        .group_by(
            Statement.customer_id,
            Customer.name,
            Statement.statement_month,
        )
        .having(func.sum(Statement.total_receivable) > 0)
        .having(
            func.sum(
                case(
                    (
                        Statement.invoiced_amount < Statement.total_receivable,
                        1,
                    ),
                    else_=0,
                )
            )
            == 0
        )
        .having(
            func.sum(
                case(
                    (
                        Statement.settled_amount < Statement.total_receivable,
                        1,
                    ),
                    else_=0,
                )
            )
            == 0
        )
        .order_by(
            Statement.statement_month.desc(),
            Customer.name,
            Statement.customer_id,
        )
    )
    if statement_month:
        group_query = group_query.where(
            Statement.statement_month == statement_month
        )
    if visible_customer_ids is not None:
        group_query = group_query.where(
            Statement.customer_id.in_(visible_customer_ids)
        )

    if statement_keyword:
        matched_statement = aliased(Statement)
        group_query = group_query.where(
            select(matched_statement.id)
            .where(
                matched_statement.customer_id == Statement.customer_id,
                matched_statement.statement_month == Statement.statement_month,
                matched_statement.statement_number.contains(
                    statement_keyword,
                    autoescape=True,
                ),
            )
            .correlate(Statement)
            .exists()
        )
    if invoice_keyword:
        invoiced_statement = aliased(Statement)
        group_query = group_query.where(
            select(Invoice.id)
            .join(
                invoiced_statement,
                Invoice.statement_id == invoiced_statement.id,
            )
            .where(
                invoiced_statement.customer_id == Statement.customer_id,
                invoiced_statement.statement_month == Statement.statement_month,
                Invoice.invoice_number.contains(
                    invoice_keyword,
                    autoescape=True,
                ),
            )
            .correlate(Statement)
            .exists()
        )
    if settlement_date_from is not None or settlement_date_to is not None:
        settled_statement = aliased(Statement)
        settlement_filters = [
            settled_statement.customer_id == Statement.customer_id,
            settled_statement.statement_month == Statement.statement_month,
        ]
        if settlement_date_from is not None:
            settlement_filters.append(
                SettlementRecord.settlement_date >= settlement_date_from
            )
        if settlement_date_to is not None:
            settlement_filters.append(
                SettlementRecord.settlement_date <= settlement_date_to
            )
        group_query = group_query.where(
            select(SettlementRecord.id)
            .join(
                settled_statement,
                SettlementRecord.statement_id == settled_statement.id,
            )
            .where(*settlement_filters)
            .correlate(Statement)
            .exists()
        )

    candidate_groups = [dict(row) for row in db.execute(group_query).mappings()]

    # Current and historical sections must never contain the same customer-month.
    # The year bound keeps this check to at most twelve authoritative month
    # queries instead of one query per customer or statement.
    pending_keys: set[tuple[str, int]] = set()
    candidate_months = sorted(
        {str(row["statement_month"]) for row in candidate_groups}
    )
    for candidate_month in candidate_months:
        pending_rows = pending_statement_customer_summaries(
            db,
            statement_month=candidate_month,
            visible_customer_ids=visible_customer_ids,
        )
        pending_keys.update(
            (candidate_month, int(row["customer_id"]))
            for row in pending_rows
            if int(row["pending_count"] or 0)
            + int(row["blocked_count"] or 0)
            > 0
        )

    groups: list[dict] = []
    for row in candidate_groups:
        group_key = (str(row["statement_month"]), int(row["customer_id"]))
        if group_key in pending_keys:
            continue
        groups.append(
            {
                "year": int(str(row["statement_month"])[:4]),
                "statement_month": str(row["statement_month"]),
                "customer_id": int(row["customer_id"]),
                "customer_name": row["customer_name"],
                "statement_count": int(row["statement_count"] or 0),
                "total_receivable": _money_value(row["total_receivable"]),
                "invoiced_amount": _money_value(row["invoiced_amount"]),
                "settled_amount": _money_value(row["settled_amount"]),
                "status_anomaly_count": int(
                    row["status_anomaly_count"] or 0
                ),
                "latest_settlement_date": None,
                "statement_numbers": [],
                "invoice_numbers": [],
                "settlement_dates": [],
                "statements": [],
            }
        )

    total = len(groups)
    summary = {
        "customer_month_count": total,
        "statement_count": sum(row["statement_count"] for row in groups),
        "total_receivable": _money_value(
            sum((row["total_receivable"] for row in groups), Decimal("0.00"))
        ),
        "invoiced_amount": _money_value(
            sum((row["invoiced_amount"] for row in groups), Decimal("0.00"))
        ),
        "settled_amount": _money_value(
            sum((row["settled_amount"] for row in groups), Decimal("0.00"))
        ),
    }
    start = (page - 1) * page_size
    page_groups = groups[start : start + page_size]
    if page_groups:
        group_conditions = [
            and_(
                Statement.customer_id == row["customer_id"],
                Statement.statement_month == row["statement_month"],
            )
            for row in page_groups
        ]
        statement_rows = db.execute(
            select(
                Statement.id,
                Statement.statement_number,
                Statement.customer_id,
                Statement.statement_month,
                Statement.total_receivable,
                Statement.invoiced_amount,
                Statement.settled_amount,
                Statement.status,
                Statement.confirmation_status,
                Statement.version,
                Statement.ledger_version,
                Statement.created_at,
            )
            .where(or_(*group_conditions))
            .order_by(Statement.statement_month.desc(), Statement.id.desc())
        ).mappings().all()
        statement_ids = [int(row["id"]) for row in statement_rows]
        invoice_rows = db.execute(
            select(
                Invoice.id,
                Invoice.statement_id,
                Invoice.invoice_number,
                Invoice.invoice_date,
                Invoice.invoice_amount,
            )
            .where(Invoice.statement_id.in_(statement_ids))
            .order_by(Invoice.invoice_date, Invoice.id)
        ).mappings().all()
        settlement_rows = db.execute(
            select(
                SettlementRecord.id,
                SettlementRecord.statement_id,
                SettlementRecord.settled_amount,
                SettlementRecord.settlement_date,
                SettlementRecord.account,
            )
            .where(SettlementRecord.statement_id.in_(statement_ids))
            .order_by(SettlementRecord.settlement_date, SettlementRecord.id)
        ).mappings().all()

        invoices_by_statement: dict[int, list[dict]] = {}
        for invoice in invoice_rows:
            invoices_by_statement.setdefault(
                int(invoice["statement_id"]), []
            ).append(dict(invoice))
        settlements_by_statement: dict[int, list[dict]] = {}
        for settlement in settlement_rows:
            settlements_by_statement.setdefault(
                int(settlement["statement_id"]), []
            ).append(dict(settlement))

        groups_by_key = {
            (row["statement_month"], row["customer_id"]): row
            for row in page_groups
        }
        for statement in statement_rows:
            statement_id = int(statement["id"])
            key = (
                str(statement["statement_month"]),
                int(statement["customer_id"]),
            )
            group = groups_by_key[key]
            invoices = invoices_by_statement.get(statement_id, [])
            settlements = settlements_by_statement.get(statement_id, [])
            group["statement_numbers"].append(statement["statement_number"])
            group["invoice_numbers"].extend(
                invoice["invoice_number"] for invoice in invoices
            )
            group["settlement_dates"].extend(
                settlement["settlement_date"] for settlement in settlements
            )
            group["statements"].append(
                {
                    "id": statement_id,
                    "statement_number": statement["statement_number"],
                    "customer_id": int(statement["customer_id"]),
                    "statement_month": statement["statement_month"],
                    "total_receivable": _money_value(
                        statement["total_receivable"]
                    ),
                    "invoiced_amount": _money_value(
                        statement["invoiced_amount"]
                    ),
                    "settled_amount": _money_value(
                        statement["settled_amount"]
                    ),
                    "status": statement["status"],
                    "confirmation_status": statement["confirmation_status"],
                    "version": int(statement["version"]),
                    "ledger_version": int(statement["ledger_version"]),
                    "created_at": statement["created_at"],
                    "invoices": invoices,
                    "settlements": settlements,
                }
            )
        for group in page_groups:
            group["invoice_numbers"] = list(dict.fromkeys(group["invoice_numbers"]))
            group["settlement_dates"] = sorted(set(group["settlement_dates"]))
            if group["settlement_dates"]:
                group["latest_settlement_date"] = group["settlement_dates"][-1]

    return {
        "year": selected_year,
        "statement_month": statement_month,
        "as_of": beijing_today(),
        "total": total,
        "page": page,
        "page_size": page_size,
        "summary": summary,
        "items": page_groups,
    }


@router.get("/pending_statements")
def pending_statements(
    customer_id: int = Query(gt=0),
    statement_month: str | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    require_customer_access(customer_id, user, db)
    customer = db.get(Customer, customer_id)
    if customer is None:
        raise HTTPException(status_code=404, detail="客户不存在")
    settlement_entity, allowed_customer_ids = _settlement_context_for_customer(
        db, customer_id
    )
    for allowed_customer_id in allowed_customer_ids:
        require_customer_access(allowed_customer_id, user, db)
    cycle_start_day = (
        settlement_entity.statement_cycle_start_day
        if settlement_entity is not None
        else customer.statement_cycle_start_day
    )
    period_start = period_end = None
    if statement_month:
        try:
            period_start, period_end = _statement_period(
                statement_month,
                cycle_start_day,
            )
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
    customer_names = dict(
        db.execute(
            select(Customer.id, Customer.name).where(
                Customer.id.in_(allowed_customer_ids)
            )
        ).all()
    )
    deliveries: list[dict] = []
    for source_customer_id in sorted(allowed_customer_ids):
        source_groups = _pending_statement_groups(
            db,
            source_customer_id,
            statement_month=statement_month,
            cycle_start_day=cycle_start_day,
            period_start=period_start,
            period_end=period_end,
        )
        for group in source_groups:
            group["source_customer_id"] = source_customer_id
            group["source_customer_name"] = customer_names.get(source_customer_id, "")
            for item in group["items"]:
                item["source_customer_id"] = source_customer_id
                item["source_customer_name"] = group["source_customer_name"]
        deliveries.extend(source_groups)
    deliveries.sort(
        key=lambda row: (
            row["delivery_date"],
            str(row["delivery_number"]),
            int(row["delivery_id"]),
        )
    )
    # 保留旧版平铺字段供尚未刷新前端的页面只读使用；只返回可整单选择的明细。
    rows = [
        item
        for delivery in deliveries
        if not delivery["selection_blocked"]
        for item in delivery["items"]
    ]
    charge_query = (
        select(CustomerCharge)
        .outerjoin(
            StatementItem,
            StatementItem.customer_charge_id == CustomerCharge.id,
        )
        .where(
            CustomerCharge.customer_id.in_(allowed_customer_ids),
            CustomerCharge.status == "confirmed",
            StatementItem.id.is_(None),
        )
        .order_by(CustomerCharge.id)
    )
    if statement_month is not None:
        charge_query = charge_query.where(
            CustomerCharge.reconciliation_month == statement_month
        )
    charges = db.scalars(charge_query).all()
    return {
        "items": rows,
        "deliveries": deliveries,
        "customer_charges": [
            _customer_charge_response(db, charge) for charge in charges
        ],
        "statement_cycle_start_day": cycle_start_day,
        "period_start": period_start,
        "period_end": period_end,
        "settlement_entity_id": (
            settlement_entity.id if settlement_entity is not None else None
        ),
        "settlement_name": (
            settlement_entity.entity_name
            if settlement_entity is not None
            else customer.name
        ),
        "customer_ids": sorted(allowed_customer_ids),
    }


@router.post("/statements", status_code=status.HTTP_201_CREATED)
def create_statement(
    payload: StatementCreate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    require_customer_access(payload.customer_id, user, db)
    request_hash = _finance_request_hash(
        "statement_create",
        payload.model_dump(exclude={"idempotency_key"}),
    )
    replay, replay_record = _finance_idempotency_replay(
        db,
        idempotency_key=payload.idempotency_key,
        request_hash=request_hash,
        action="statement_create",
        actor=user,
    )
    if replay is not None:
        assert replay_record is not None
        _statement_for_user(db, replay_record.resource_id, user)
        return replay
    customer = db.get(Customer, payload.customer_id)
    if customer is None:
        raise HTTPException(status_code=400, detail="客户不存在")
    settlement_entity, allowed_customer_ids = _settlement_context_for_customer(
        db, payload.customer_id
    )
    for allowed_customer_id in allowed_customer_ids:
        require_customer_access(allowed_customer_id, user, db)
    representative_customer_id = min(allowed_customer_ids)
    cycle_start_day = (
        settlement_entity.statement_cycle_start_day
        if settlement_entity is not None
        else customer.statement_cycle_start_day
    )
    try:
        period_start, period_end = _statement_period(
            payload.statement_month,
            cycle_start_day,
        )
        compatibility_item_ids = set(payload.return_receipt_item_ids)
        selected_delivery_ids = set(payload.delivery_ids)
        selected_charge_ids = set(payload.customer_charge_ids)
        if compatibility_item_ids:
            mapped_rows = db.execute(
                select(ReturnReceiptItem.id, Delivery.id)
                .join(
                    DeliveryItem,
                    DeliveryItem.id == ReturnReceiptItem.delivery_item_id,
                )
                .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
                .where(ReturnReceiptItem.id.in_(compatibility_item_ids))
            ).all()
            if {row[0] for row in mapped_rows} != compatibility_item_ids:
                raise HTTPException(status_code=400, detail="回单明细不存在")
            selected_delivery_ids = {row[1] for row in mapped_rows}

        selected_rows_query = (
            select(
                ReturnReceiptItem,
                ReturnReceipt,
                Delivery,
                DeliveryItem,
                OrderItem,
                Product,
                StatementItem.id.label("existing_statement_item_id"),
            )
            .join(
                ReturnReceipt,
                ReturnReceipt.id == ReturnReceiptItem.return_receipt_id,
            )
            .join(
                DeliveryItem,
                DeliveryItem.id == ReturnReceiptItem.delivery_item_id,
            )
            .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
            .outerjoin(OrderItem, OrderItem.id == DeliveryItem.order_item_id)
            .outerjoin(
                Product,
                Product.id
                == func.coalesce(DeliveryItem.product_id, OrderItem.product_id),
            )
            .outerjoin(
                StatementItem,
                StatementItem.return_receipt_item_id == ReturnReceiptItem.id,
            )
            .order_by(Delivery.id, ReturnReceiptItem.id)
        )
        if compatibility_item_ids:
            selected_rows_query = selected_rows_query.where(
                ReturnReceiptItem.id.in_(compatibility_item_ids)
            )
        else:
            selected_rows_query = selected_rows_query.where(
                Delivery.id.in_(selected_delivery_ids)
            )
        selected_rows = db.execute(selected_rows_query).all()
        found_delivery_ids = {row[2].id for row in selected_rows}
        if found_delivery_ids != selected_delivery_ids:
            raise HTTPException(
                status_code=400,
                detail="所选送货单不存在已确认的客户回单明细",
            )
        if compatibility_item_ids and not any(
            row[0].reconciliation_month_override for row in selected_rows
        ):
            all_delivery_item_ids = set(
                db.scalars(
                    select(ReturnReceiptItem.id)
                    .join(
                        DeliveryItem,
                        DeliveryItem.id == ReturnReceiptItem.delivery_item_id,
                    )
                    .where(DeliveryItem.delivery_id.in_(selected_delivery_ids))
                ).all()
            )
            if compatibility_item_ids != all_delivery_item_ids:
                raise HTTPException(
                    status_code=400,
                    detail="初次整单对账必须整张送货单选择；客户异议后的跨期明细可单独进入目标月份。",
                )
        claimed_receipt_ids: set[int] = set()
        for row in selected_rows:
            (
                receipt_item,
                receipt,
                delivery,
                _delivery_item,
                _order_item,
                _product,
                existing_id,
            ) = row
            require_customer_access(delivery.customer_id, user, db)
            if delivery.customer_id not in allowed_customer_ids:
                raise HTTPException(status_code=400, detail="送货单客户不匹配")
            if receipt.status != "confirmed":
                raise HTTPException(status_code=409, detail="已取消回单不能生成对账单")
            if existing_id is not None:
                raise HTTPException(
                    status_code=409,
                    detail="该送货单已有明细进入对账单，请先处理原对账单。",
                )
            if receipt_item.reconciliation_month_override:
                effective_month = receipt_item.reconciliation_month_override
            else:
                effective_month, _month_source = _effective_reconciliation_month(
                    receipt,
                    delivery_date=delivery.delivery_date,
                    cycle_start_day=cycle_start_day,
                )
            if effective_month != payload.statement_month:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"送货单 {delivery.delivery_number} 的对账归属月份为"
                        f" {effective_month}，不能进入 {payload.statement_month}。"
                    ),
                )
            if receipt.id not in claimed_receipt_ids:
                _claim_return_receipt_status(
                    db,
                    receipt_id=receipt.id,
                    expected_status="confirmed",
                    next_status="confirmed",
                )
                claimed_receipt_ids.add(receipt.id)

        selected_charge_rows = db.execute(
            select(
                CustomerCharge,
                StatementItem.id.label("existing_statement_item_id"),
            )
            .outerjoin(
                StatementItem,
                StatementItem.customer_charge_id == CustomerCharge.id,
            )
            .where(CustomerCharge.id.in_(selected_charge_ids))
            .order_by(CustomerCharge.id)
        ).all()
        if {row[0].id for row in selected_charge_rows} != selected_charge_ids:
            raise HTTPException(status_code=400, detail="所选客户附加收费不存在")
        for charge, existing_id in selected_charge_rows:
            require_customer_access(charge.customer_id, user, db)
            if charge.customer_id not in allowed_customer_ids:
                raise HTTPException(status_code=400, detail="附加收费客户不匹配")
            if charge.status != "confirmed":
                raise HTTPException(status_code=409, detail="只有已确认收费可以进入对账单")
            if charge.reconciliation_month != payload.statement_month:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"收费 {charge.display_name} 的对账归属月份为 "
                        f"{charge.reconciliation_month}，不能进入 {payload.statement_month}。"
                    ),
                )
            if existing_id is not None:
                raise HTTPException(status_code=409, detail="该附加收费已经进入其他对账单")

        statement = Statement(
            statement_number=_next_statement_number(
                db,
                payload.statement_month,
            ),
            customer_id=representative_customer_id,
            settlement_entity_id=(
                settlement_entity.id if settlement_entity is not None else None
            ),
            settlement_name_snapshot=(
                settlement_entity.entity_name
                if settlement_entity is not None
                else customer.name
            ),
            settlement_customer_ids_snapshot_json=json.dumps(
                sorted(allowed_customer_ids), ensure_ascii=False
            ),
            statement_cycle_start_day_snapshot=cycle_start_day,
            statement_month=payload.statement_month,
            total_receivable=Decimal("0"),
            total_gross_profit=Decimal("0"),
            status="unsettled",
            created_by=user.id,
        )
        db.add(statement)
        db.flush()
        total_receivable = Decimal("0")
        total_profit = Decimal("0")
        price_tax_terms_by_customer = {
            source_customer_id: resolve_customer_price_tax_terms(
                db, source_customer_id
            )
            for source_customer_id in allowed_customer_ids
        }
        statement_price_tax_modes: set[str] = set()
        for row in selected_rows:
            (
                receipt_item,
                _receipt,
                delivery,
                delivery_item,
                order_item,
                product,
                _existing_id,
            ) = row
            if delivery_item.source_type == "unordered_finished":
                if delivery_item.unit_price_snapshot is None:
                    raise HTTPException(
                        status_code=409,
                        detail="无订单库存送货缺少冻结单价，不能生成对账单",
                    )
                unit_price = Decimal(str(delivery_item.unit_price_snapshot))
            else:
                if order_item is None:
                    raise HTTPException(
                        status_code=409,
                        detail="订单送货明细缺少订单关联，不能生成对账单",
                    )
                unit_price = Decimal(str(order_item.unit_price))
            if product is None:
                raise HTTPException(
                    status_code=409,
                    detail="送货明细缺少产品资料，不能生成对账单",
                )
            unit_cost = Decimal(str(product.cost_unit_price or 0))
            quantity = Decimal(receipt_item.actual_received_quantity)
            current_price_tax_terms = price_tax_terms_by_customer[
                delivery.customer_id
            ]
            price_tax_mode = (
                order_item.price_tax_mode_snapshot
                if order_item is not None
                and order_item.price_tax_mode_snapshot in VALID_PRICE_TAX_MODES
                else current_price_tax_terms.price_tax_mode
            )
            tax_rate = Decimal(
                str(
                    order_item.tax_rate_snapshot
                    if order_item is not None
                    and order_item.tax_rate_snapshot is not None
                    else current_price_tax_terms.tax_rate
                )
            )
            line_price_amount = (quantity * unit_price).quantize(
                MONEY,
                rounding=ROUND_HALF_UP,
            )
            if price_tax_mode == "tax_exclusive":
                line_tax_amount = (line_price_amount * tax_rate).quantize(
                    MONEY,
                    rounding=ROUND_HALF_UP,
                )
                receivable = line_price_amount + line_tax_amount
            else:
                receivable = line_price_amount
            profit = (quantity * (unit_price - unit_cost)).quantize(
                MONEY,
                rounding=ROUND_HALF_UP,
            )
            statement_price_tax_modes.add(price_tax_mode)
            db.add(
                StatementItem(
                    statement_id=statement.id,
                    source_customer_id=delivery.customer_id,
                    return_receipt_item_id=receipt_item.id,
                    actual_received_quantity=receipt_item.actual_received_quantity,
                    unit_price_snapshot=unit_price,
                    unit_cost_snapshot=unit_cost,
                    receivable_amount=receivable,
                    gross_profit_amount=profit,
                    price_tax_mode_snapshot=price_tax_mode,
                    tax_rate_snapshot=tax_rate,
                )
            )
            total_receivable += receivable
            total_profit += profit
        for charge, _existing_id in selected_charge_rows:
            db.add(
                StatementItem(
                    statement_id=statement.id,
                    source_customer_id=charge.customer_id,
                    customer_charge_id=charge.id,
                    actual_received_quantity=None,
                    charge_quantity_snapshot=charge.quantity,
                    unit_snapshot=charge.unit,
                    source_label_snapshot=charge.display_name,
                    unit_price_snapshot=charge.unit_price,
                    unit_cost_snapshot=Decimal("0"),
                    receivable_amount=charge.amount,
                    gross_profit_amount=Decimal("0"),
                    price_tax_mode_snapshot=charge.price_tax_mode,
                    tax_rate_snapshot=charge.tax_rate,
                )
            )
            total_receivable += Decimal(str(charge.amount))
            if charge.price_tax_mode in VALID_PRICE_TAX_MODES:
                statement_price_tax_modes.add(charge.price_tax_mode)
        statement.total_receivable = total_receivable.quantize(MONEY)
        statement.total_gross_profit = total_profit.quantize(MONEY)
        _audit(
            db,
            user=user,
            action="CREATE_STATEMENT",
            resource="Statement",
            entity_id=statement.id,
            details={
                "statement_number": statement.statement_number,
                "statement_month": statement.statement_month,
                "delivery_count": len(selected_delivery_ids),
                "customer_charge_count": len(selected_charge_ids),
                "item_count": len(selected_rows),
                "total_receivable": statement.total_receivable,
                "total_gross_profit": statement.total_gross_profit,
                "price_tax_modes": sorted(statement_price_tax_modes),
            },
            description="生成客户月结对账单",
        )
        response = _redact_statement_costs(
            {
                "id": statement.id,
                "statement_number": statement.statement_number,
                "customer_id": statement.customer_id,
                "customer_name": statement.settlement_name_snapshot or customer.name,
                "settlement_entity_id": statement.settlement_entity_id,
                "statement_month": statement.statement_month,
                "total_receivable": statement.total_receivable,
                "total_gross_profit": statement.total_gross_profit,
                "status": statement.status,
                "confirmation_status": statement.confirmation_status,
                "version": statement.version,
                "ledger_version": statement.ledger_version,
            },
            user,
        )
        _record_finance_idempotency(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action="statement_create",
            actor=user,
            resource_type="statement",
            resource_id=statement.id,
            response=response,
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        raise
    except IntegrityError as error:
        db.rollback()
        replay, replay_record = _finance_idempotency_replay(
            db,
            idempotency_key=payload.idempotency_key,
            request_hash=request_hash,
            action="statement_create",
            actor=user,
        )
        if replay is not None:
            assert replay_record is not None
            _statement_for_user(db, replay_record.resource_id, user)
            return replay
        raise HTTPException(status_code=409, detail="对账明细已被其他对账单使用") from error
    except Exception:
        db.rollback()
        raise


@router.post("/invoices", status_code=status.HTTP_201_CREATED)
def create_invoice(
    payload: InvoiceCreate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    _statement_for_user(db, payload.statement_id, user)
    actor_id = int(user.id)
    expected_version = payload.expected_version
    expected_ledger_version = payload.expected_ledger_version
    request_hash = _finance_manual_request_hash(
        {
            "mutation_type": "register_invoice",
            "statement_id": payload.statement_id,
            "expected_version": expected_version,
            "expected_ledger_version": expected_ledger_version,
            "invoice_number": payload.invoice_number,
            "invoice_date": payload.invoice_date.isoformat(),
            "invoice_amount": format(payload.invoice_amount, ".2f"),
        }
    )
    try:
        mutation, replay = _reserve_finance_manual_mutation(
            db,
            idempotency_key=payload.idempotency_key,
            mutation_type="register_invoice",
            statement_id=payload.statement_id,
            request_hash=request_hash,
            actor_id=actor_id,
        )
        if replay is not None:
            return replay
        assert mutation is not None
        updated = db.execute(
            text(
                """
                UPDATE finance_statements
                SET invoiced_amount = invoiced_amount + :amount,
                    ledger_version = ledger_version + 1
                WHERE id = :statement_id
                  AND confirmation_status = 'confirmed'
                  AND version = :expected_version
                  AND ledger_version = :expected_ledger_version
                  AND invoiced_amount + :amount <= total_receivable
                RETURNING id, statement_number, total_receivable,
                          invoiced_amount, settled_amount, status,
                          confirmation_status, version, ledger_version
                """
            ),
            {
                "statement_id": payload.statement_id,
                "amount": str(payload.invoice_amount),
                "expected_version": expected_version,
                "expected_ledger_version": expected_ledger_version,
            },
        ).mappings().one_or_none()
        if updated is None:
            _raise_finance_manual_cas_conflict(
                db,
                statement_id=payload.statement_id,
                expected_version=expected_version,
                expected_ledger_version=expected_ledger_version,
                amount_kind="invoice",
            )

        invoice = Invoice(
            statement_id=payload.statement_id,
            invoice_number=payload.invoice_number,
            invoice_date=payload.invoice_date,
            invoice_amount=payload.invoice_amount,
            created_by=user.id,
        )
        db.add(invoice)
        db.flush()
        _audit(
            db,
            user=user,
            action="REGISTER_INVOICE",
            resource="Invoice",
            entity_id=invoice.id,
            details={
                "statement_id": payload.statement_id,
                "invoice_number": payload.invoice_number,
                "invoice_date": payload.invoice_date,
                "invoice_amount": payload.invoice_amount,
                "invoiced_amount": updated["invoiced_amount"],
                "expected_version": expected_version,
                "expected_ledger_version": expected_ledger_version,
                "version": updated["version"],
                "ledger_version": updated["ledger_version"],
                "idempotency_key": payload.idempotency_key,
            },
            description="登记客户发票",
        )
        response = _record_finance_manual_response(
            mutation,
            {
                "id": invoice.id,
                "statement_id": payload.statement_id,
                "invoice_number": invoice.invoice_number,
                "invoice_date": invoice.invoice_date,
                "invoice_amount": invoice.invoice_amount,
                "total_receivable": updated["total_receivable"],
                "invoiced_amount": updated["invoiced_amount"],
                "settled_amount": updated["settled_amount"],
                "status": updated["status"],
                "confirmation_status": updated["confirmation_status"],
                "version": updated["version"],
                "ledger_version": updated["ledger_version"],
            },
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        raise
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="发票号码已存在") from error
    except Exception:
        db.rollback()
        raise


@router.put("/statements/{statement_id}/settle")
def settle_statement(
    statement_id: int,
    payload: SettlementCreate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    _statement_for_user(db, statement_id, user)
    actor_id = int(user.id)
    expected_version = payload.expected_version
    expected_ledger_version = payload.expected_ledger_version
    request_hash = _finance_manual_request_hash(
        {
            "mutation_type": "settle_statement",
            "statement_id": statement_id,
            "expected_version": expected_version,
            "expected_ledger_version": expected_ledger_version,
            "amount": format(payload.amount, ".2f"),
            "settlement_date": payload.settlement_date.isoformat(),
            "account": payload.account or "",
        }
    )
    try:
        mutation, replay = _reserve_finance_manual_mutation(
            db,
            idempotency_key=payload.idempotency_key,
            mutation_type="settle_statement",
            statement_id=statement_id,
            request_hash=request_hash,
            actor_id=actor_id,
        )
        if replay is not None:
            return replay
        assert mutation is not None
        updated = db.execute(
            text(
                """
                UPDATE finance_statements
                SET settled_amount = settled_amount + :amount,
                    status = CASE
                        WHEN settled_amount + :amount = total_receivable
                        THEN 'settled'
                        ELSE 'unsettled'
                    END,
                    ledger_version = ledger_version + 1
                WHERE id = :statement_id
                  AND confirmation_status = 'confirmed'
                  AND version = :expected_version
                  AND ledger_version = :expected_ledger_version
                  AND settled_amount + :amount <= total_receivable
                RETURNING id, statement_number, total_receivable,
                          invoiced_amount, settled_amount, status,
                          confirmation_status, version, ledger_version
                """
            ),
            {
                "statement_id": statement_id,
                "amount": str(payload.amount),
                "expected_version": expected_version,
                "expected_ledger_version": expected_ledger_version,
            },
        ).mappings().one_or_none()
        if updated is None:
            _raise_finance_manual_cas_conflict(
                db,
                statement_id=statement_id,
                expected_version=expected_version,
                expected_ledger_version=expected_ledger_version,
                amount_kind="settlement",
            )

        settlement = SettlementRecord(
            statement_id=statement_id,
            settled_amount=payload.amount,
            settlement_date=payload.settlement_date,
            account=payload.account or "",
            created_by=user.id,
        )
        db.add(settlement)
        db.flush()
        _audit(
            db,
            user=user,
            action="SETTLE_STATEMENT",
            resource="Statement",
            entity_id=statement_id,
            details={
                "settlement_record_id": settlement.id,
                "amount": payload.amount,
                "settlement_date": payload.settlement_date,
                "account": payload.account or "",
                "settled_amount": updated["settled_amount"],
                "status": updated["status"],
                "expected_version": expected_version,
                "expected_ledger_version": expected_ledger_version,
                "version": updated["version"],
                "ledger_version": updated["ledger_version"],
                "idempotency_key": payload.idempotency_key,
            },
            description="登记客户收款并核销对账单",
        )
        response = _record_finance_manual_response(
            mutation,
            {
                "id": statement_id,
                "statement_number": updated["statement_number"],
                "total_receivable": updated["total_receivable"],
                "invoiced_amount": updated["invoiced_amount"],
                "settled_amount": updated["settled_amount"],
                "status": updated["status"],
                "status_label": (
                    "已结清" if updated["status"] == "settled" else "未结清"
                ),
                "settlement_record_id": settlement.id,
                "confirmation_status": updated["confirmation_status"],
                "version": updated["version"],
                "ledger_version": updated["ledger_version"],
            },
        )
        db.commit()
        return response
    except HTTPException:
        db.rollback()
        raise
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="收款登记冲突，请刷新后重试",
        ) from error
    except Exception:
        db.rollback()
        raise


@router.post("/return_receipts/{receipt_id}/cancel")
def cancel_return_receipt(
    receipt_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
    _write_guard: None = Depends(fulfillment_reminder_write_guard),
) -> dict:
    del _write_guard
    receipt = _return_receipt_for_user(db, receipt_id, user)
    if receipt.status != "confirmed":
        raise HTTPException(status_code=409, detail="回单状态已变化，不能重复取消")
    _claim_return_receipt_status(
        db,
        receipt_id=receipt.id,
        expected_status="confirmed",
        next_status="cancelled",
    )
    db.expire(receipt)
    receipt_item_ids = list(
        db.scalars(
            select(ReturnReceiptItem.id).where(
                ReturnReceiptItem.return_receipt_id == receipt.id
            )
        ).all()
    )
    if receipt_item_ids and db.scalar(
        select(StatementItem.id)
        .where(StatementItem.return_receipt_item_id.in_(receipt_item_ids))
        .limit(1)
    ) is not None:
        raise HTTPException(
            status_code=409,
            detail="该送货单已进入对账/结清流程，请先取消或编辑对应对账单。",
        )
    before = _receipt_response(db, receipt.id)
    before["status"] = "confirmed"
    receipt_items = db.scalars(
        select(ReturnReceiptItem).where(
            ReturnReceiptItem.return_receipt_id == receipt.id
        )
    ).all()
    delivery_items = {
        item.id: item
        for item in db.scalars(
            select(DeliveryItem).where(
                DeliveryItem.id.in_([item.delivery_item_id for item in receipt_items])
            )
        ).all()
    }
    _assert_no_later_dispatched_deliveries(
        db,
        receipt=receipt,
        receipt_items=receipt_items,
        delivery_items=delivery_items,
    )
    try:
        reconsume_unordered_finished_receipt_returns(
            db,
            return_receipt_item_ids=receipt_item_ids,
            operator_id=user.id,
        )
        reconsume_ordered_finished_receipt_returns(
            db,
            return_receipt_item_ids=receipt_item_ids,
            operator_id=user.id,
        )
    except WarehouseInventoryError as error:
        db.rollback()
        raise HTTPException(
            status_code=error.status_code,
            detail=str(error),
        ) from error
    affected_order_ids: set[int] = set()
    for item in receipt_items:
        order_id = _apply_receipt_order_effect(
            db,
            delivery_item=delivery_items[item.delivery_item_id],
            actual_received_quantity=item.actual_received_quantity,
            resolution_action=item.resolution_action,
            direction=-1,
        )
        if order_id is not None:
            affected_order_ids.add(order_id)
    receipt.signed_by = None
    sync_receipt_source(
        db,
        receipt_id=receipt.id,
        source_valid=False,
        received_date=receipt.actual_received_date,
        actor=user,
    )
    _refresh_receipt_order_statuses(db, affected_order_ids)
    _audit(
        db,
        user=user,
        action="CANCEL_RETURN_RECEIPT",
        resource="ReturnReceipt",
        entity_id=receipt.id,
        details={"before": before},
        description="撤销客户送货回单",
    )
    db.commit()
    return _receipt_response(db, receipt.id)


@router.get("/statements/{statement_id}")
def get_statement(
    statement_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    return _statement_detail_response(db, statement_id, user)


def _recalculate_statement_totals(db: Session, statement: Statement) -> None:
    totals = db.execute(
        select(
            func.coalesce(func.sum(StatementItem.receivable_amount), 0),
            func.coalesce(func.sum(StatementItem.gross_profit_amount), 0),
        ).where(StatementItem.statement_id == statement.id)
    ).one()
    statement.total_receivable = Decimal(str(totals[0])).quantize(MONEY)
    statement.total_gross_profit = Decimal(str(totals[1])).quantize(MONEY)


def _statement_adjustment_log(
    db: Session,
    *,
    statement: Statement,
    action: str,
    reason: str,
    before_version: int,
    details: dict,
    user: User,
) -> None:
    db.add(
        StatementAdjustment(
            statement_id=statement.id,
            action=action,
            reason=reason,
            before_version=before_version,
            after_version=statement.version,
            details_json=json.dumps(
                jsonable_encoder(details), ensure_ascii=False, sort_keys=True
            ),
            created_by=user.id,
        )
    )


@router.post("/statements/{statement_id}/reopen")
def reopen_statement_for_dispute(
    statement_id: int,
    payload: StatementReopen,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    try:
        statement = _statement_for_user(db, statement_id, user)
        if statement.version != payload.expected_version:
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "对账单版本已变化，请刷新后重试",
                    "current_version": statement.version,
                },
            )
        issued_invoice = db.scalar(
            select(Invoice.id).where(
                Invoice.statement_id == statement.id,
                Invoice.invoice_status == "issued",
            ).limit(1)
        )
        tasks = db.scalars(
            select(FinanceInvoiceTask).where(
                FinanceInvoiceTask.statement_id == statement.id,
                FinanceInvoiceTask.status != "voided",
            )
        ).all()
        if issued_invoice is not None or any(task.status == "issued" for task in tasks):
            raise HTTPException(
                status_code=409,
                detail="该对账单已登记开票，不能直接撤销；请按红冲/重开流程处理。",
            )
        if statement.confirmation_status != "confirmed":
            return _statement_detail_response(db, statement.id, user)
        before_version = statement.version
        voided_tasks: list[int] = []
        for task in tasks:
            task.status = "voided"
            task.voided_by = user.id
            task.voided_at = datetime.now()
            task.version += 1
            voided_tasks.append(task.id)
        statement.confirmation_status = "draft"
        statement.confirmed_by = None
        statement.confirmed_at = None
        statement.version += 1
        _statement_adjustment_log(
            db,
            statement=statement,
            action="reopen_for_dispute",
            reason=payload.reason,
            before_version=before_version,
            details={"voided_invoice_task_ids": voided_tasks},
            user=user,
        )
        _audit(
            db,
            user=user,
            action="REOPEN_STATEMENT_FOR_DISPUTE",
            resource="Statement",
            entity_id=statement.id,
            details={
                "reason": payload.reason,
                "before_version": before_version,
                "after_version": statement.version,
                "voided_invoice_task_ids": voided_tasks,
            },
            description="客户异议撤销对账确认",
        )
        db.commit()
        return _statement_detail_response(db, statement.id, user)
    except HTTPException:
        db.rollback()
        raise


@router.get("/statements/{statement_id}/adjustment-candidates")
def statement_adjustment_candidates(
    statement_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    statement = _statement_for_user(db, statement_id, user)
    allowed_customer_ids = _statement_scope_customer_ids(db, statement)
    cycle_start_day = _statement_cycle_day(db, statement)
    customer_names = dict(
        db.execute(
            select(Customer.id, Customer.name).where(
                Customer.id.in_(allowed_customer_ids)
            )
        ).all()
    )
    groups: list[dict] = []
    for source_customer_id in sorted(allowed_customer_ids):
        source_groups = _pending_statement_groups(
            db,
            source_customer_id,
            cycle_start_day=cycle_start_day,
        )
        for group in source_groups:
            for item in group["items"]:
                item["source_customer_id"] = source_customer_id
                item["source_customer_name"] = customer_names.get(
                    source_customer_id, ""
                )
        groups.extend(source_groups)
    items = [
        item
        for group in groups
        for item in group["items"]
        if not item["is_reconciled"]
        and item["effective_reconciliation_month"] <= statement.statement_month
    ]
    return {"items": items}


def _new_statement_item_from_receipt(
    db: Session,
    *,
    statement: Statement,
    receipt_item: ReturnReceiptItem,
    source_customer_id: int,
    delivery_item: DeliveryItem,
    order_item: OrderItem | None,
    product: Product,
) -> StatementItem:
    if delivery_item.source_type == "unordered_finished":
        if delivery_item.unit_price_snapshot is None:
            raise HTTPException(status_code=409, detail="无订单库存送货缺少冻结单价")
        unit_price = Decimal(str(delivery_item.unit_price_snapshot))
    else:
        if order_item is None:
            raise HTTPException(status_code=409, detail="订单送货明细缺少订单关联")
        unit_price = Decimal(str(order_item.unit_price))
    terms = resolve_customer_price_tax_terms(db, source_customer_id)
    price_tax_mode = (
        order_item.price_tax_mode_snapshot
        if order_item is not None
        and order_item.price_tax_mode_snapshot in VALID_PRICE_TAX_MODES
        else terms.price_tax_mode
    )
    tax_rate = Decimal(
        str(
            order_item.tax_rate_snapshot
            if order_item is not None and order_item.tax_rate_snapshot is not None
            else terms.tax_rate
        )
    )
    quantity = Decimal(receipt_item.actual_received_quantity)
    line_price = (quantity * unit_price).quantize(MONEY, rounding=ROUND_HALF_UP)
    receivable = (
        line_price
        + (line_price * tax_rate).quantize(MONEY, rounding=ROUND_HALF_UP)
        if price_tax_mode == "tax_exclusive"
        else line_price
    )
    unit_cost = Decimal(str(product.cost_unit_price or 0))
    profit = (quantity * (unit_price - unit_cost)).quantize(
        MONEY, rounding=ROUND_HALF_UP
    )
    return StatementItem(
        statement_id=statement.id,
        source_customer_id=source_customer_id,
        return_receipt_item_id=receipt_item.id,
        actual_received_quantity=receipt_item.actual_received_quantity,
        unit_price_snapshot=unit_price,
        unit_cost_snapshot=unit_cost,
        receivable_amount=receivable,
        gross_profit_amount=profit,
        price_tax_mode_snapshot=price_tax_mode,
        tax_rate_snapshot=tax_rate,
    )


@router.post("/statements/{statement_id}/adjust-dispute")
def adjust_statement_dispute(
    statement_id: int,
    payload: StatementDisputeAdjustment,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    try:
        statement = _statement_for_user(db, statement_id, user)
        if statement.version != payload.expected_version:
            raise HTTPException(
                status_code=409,
                detail={
                    "message": "对账单版本已变化，请刷新后重试",
                    "current_version": statement.version,
                },
            )
        if statement.confirmation_status not in {"draft", "confirmed"}:
            raise HTTPException(status_code=409, detail="当前对账状态不能调整异议明细")
        tasks = db.scalars(
            select(FinanceInvoiceTask).where(
                FinanceInvoiceTask.statement_id == statement.id,
                FinanceInvoiceTask.status != "voided",
            )
        ).all()
        issued_invoice = db.scalar(
            select(Invoice.id).where(
                Invoice.statement_id == statement.id,
                Invoice.invoice_status == "issued",
            ).limit(1)
        )
        if issued_invoice is not None or any(task.status == "issued" for task in tasks):
            raise HTTPException(
                status_code=409,
                detail="该对账单已登记开票，不能直接修改；请按红冲/重开流程处理。",
            )
        if statement.confirmation_status == "draft" and tasks:
            raise HTTPException(status_code=409, detail="仍有有效开票任务，不能调整对账明细")
        voided_tasks: list[int] = []
        reopened_confirmation = statement.confirmation_status == "confirmed"
        if reopened_confirmation:
            for task in tasks:
                task.status = "voided"
                task.voided_by = user.id
                task.voided_at = datetime.now()
                task.version += 1
                voided_tasks.append(task.id)
            statement.confirmation_status = "draft"
            statement.confirmed_by = None
            statement.confirmed_at = None
        allowed_customer_ids = _statement_scope_customer_ids(db, statement)
        cycle_start_day = _statement_cycle_day(db, statement)
        before_version = statement.version
        removed: list[dict] = []
        removal_map = {
            item.statement_item_id: item.target_month for item in payload.remove_lines
        }
        if removal_map:
            remove_rows = db.execute(
                select(StatementItem, ReturnReceiptItem)
                .join(
                    ReturnReceiptItem,
                    ReturnReceiptItem.id == StatementItem.return_receipt_item_id,
                )
                .where(
                    StatementItem.statement_id == statement.id,
                    StatementItem.id.in_(removal_map),
                )
            ).all()
            if {row[0].id for row in remove_rows} != set(removal_map):
                raise HTTPException(status_code=400, detail="所选移出明细不属于当前对账单")
            for statement_item, receipt_item in remove_rows:
                target_month = removal_map[statement_item.id]
                if target_month <= statement.statement_month:
                    raise HTTPException(status_code=400, detail="异议移出月份必须晚于当前对账月份")
                receipt_item.reconciliation_month_override = target_month
                receipt_item.reconciliation_override_reason = payload.reason
                receipt_item.reconciliation_overridden_by = user.id
                receipt_item.reconciliation_overridden_at = datetime.now()
                removed.append(
                    {
                        "statement_item_id": statement_item.id,
                        "return_receipt_item_id": receipt_item.id,
                        "target_month": target_month,
                    }
                )
                db.delete(statement_item)
            db.flush()

        added: list[int] = []
        if payload.add_return_receipt_item_ids:
            add_rows = db.execute(
                select(
                    ReturnReceiptItem,
                    ReturnReceipt,
                    Delivery,
                    DeliveryItem,
                    OrderItem,
                    Product,
                    StatementItem.id.label("existing_statement_item_id"),
                )
                .join(ReturnReceipt, ReturnReceipt.id == ReturnReceiptItem.return_receipt_id)
                .join(DeliveryItem, DeliveryItem.id == ReturnReceiptItem.delivery_item_id)
                .join(Delivery, Delivery.id == DeliveryItem.delivery_id)
                .outerjoin(OrderItem, OrderItem.id == DeliveryItem.order_item_id)
                .outerjoin(
                    Product,
                    Product.id == func.coalesce(DeliveryItem.product_id, OrderItem.product_id),
                )
                .outerjoin(
                    StatementItem,
                    StatementItem.return_receipt_item_id == ReturnReceiptItem.id,
                )
                .where(ReturnReceiptItem.id.in_(payload.add_return_receipt_item_ids))
            ).all()
            if {row[0].id for row in add_rows} != set(payload.add_return_receipt_item_ids):
                raise HTTPException(status_code=400, detail="所选补入明细不存在")
            for receipt_item, receipt, delivery, delivery_item, order_item, product, existing_id in add_rows:
                if delivery.customer_id not in allowed_customer_ids:
                    raise HTTPException(
                        status_code=403,
                        detail="不能补入当前对账归集范围以外的回单明细",
                    )
                if receipt.status != "confirmed" or existing_id is not None:
                    raise HTTPException(status_code=409, detail="只能补入已回单且未被其他对账单使用的明细")
                if product is None:
                    raise HTTPException(status_code=409, detail="补入明细缺少产品资料")
                source_month = receipt_item.reconciliation_month_override
                if not source_month:
                    source_month, _ = _effective_reconciliation_month(
                        receipt,
                        delivery_date=delivery.delivery_date,
                        cycle_start_day=cycle_start_day,
                    )
                if source_month > statement.statement_month:
                    raise HTTPException(status_code=409, detail="不能提前补入未来月份明细")
                receipt_item.reconciliation_month_override = statement.statement_month
                receipt_item.reconciliation_override_reason = payload.reason
                receipt_item.reconciliation_overridden_by = user.id
                receipt_item.reconciliation_overridden_at = datetime.now()
                db.add(
                    _new_statement_item_from_receipt(
                        db,
                        statement=statement,
                        receipt_item=receipt_item,
                        source_customer_id=delivery.customer_id,
                        delivery_item=delivery_item,
                        order_item=order_item,
                        product=product,
                    )
                )
                added.append(receipt_item.id)
            db.flush()
        remaining_count = db.scalar(
            select(func.count(StatementItem.id)).where(
                StatementItem.statement_id == statement.id
            )
        )
        if not remaining_count:
            raise HTTPException(status_code=409, detail="对账单至少保留一条有效明细")
        _recalculate_statement_totals(db, statement)
        statement.version += 1
        _statement_adjustment_log(
            db,
            statement=statement,
            action=(
                "resolve_dispute"
                if reopened_confirmation
                else "adjust_dispute_lines"
            ),
            reason=payload.reason,
            before_version=before_version,
            details={
                "removed": removed,
                "added_return_receipt_item_ids": added,
                "voided_invoice_task_ids": voided_tasks,
            },
            user=user,
        )
        _audit(
            db,
            user=user,
            action="ADJUST_STATEMENT_DISPUTE_LINES",
            resource="Statement",
            entity_id=statement.id,
            details={
                "reason": payload.reason,
                "removed": removed,
                "added_return_receipt_item_ids": added,
                "voided_invoice_task_ids": voided_tasks,
                "before_version": before_version,
                "after_version": statement.version,
            },
            description="客户异议调整对账明细",
        )
        db.commit()
        return _statement_detail_response(db, statement.id, user)
    except HTTPException:
        db.rollback()
        raise


@router.put("/statements/{statement_id}")
def update_statement(
    statement_id: int,
    payload: StatementUpdate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    try:
        statement = _statement_for_user(db, statement_id, user)
        if statement.confirmation_status == "confirmed":
            raise HTTPException(
                status_code=409,
                detail="对账单已确认；如需修改来源，请先作废对应开票任务并重新核对。",
            )
        if db.scalar(
            select(Invoice.id).where(Invoice.statement_id == statement.id).limit(1)
        ) is not None:
            raise HTTPException(
                status_code=409,
                detail="该对账单已有开票记录，请先取消开票后再编辑。",
            )
        if db.scalar(
            select(SettlementRecord.id)
            .where(SettlementRecord.statement_id == statement.id)
            .limit(1)
        ) is not None:
            raise HTTPException(
                status_code=409,
                detail="该对账单已有收款记录，请先撤销收款后再编辑。",
            )
        old_month = statement.statement_month
        if payload.statement_month == old_month:
            return _statement_detail_response(db, statement.id, user)
        customer = db.get(Customer, statement.customer_id)
        if customer is None:
            raise HTTPException(status_code=409, detail="对账单关联客户不存在。")
        period_start, period_end = _statement_period(
            payload.statement_month,
            customer.statement_cycle_start_day,
        )
        delivery_rows = db.execute(
            select(
                Delivery.delivery_number,
                Delivery.delivery_date,
                ReturnReceipt.reconciliation_month,
            )
            .join(DeliveryItem, DeliveryItem.delivery_id == Delivery.id)
            .join(
                ReturnReceiptItem,
                ReturnReceiptItem.delivery_item_id == DeliveryItem.id,
            )
            .join(
                ReturnReceipt,
                ReturnReceipt.id == ReturnReceiptItem.return_receipt_id,
            )
            .join(
                StatementItem,
                StatementItem.return_receipt_item_id == ReturnReceiptItem.id,
            )
            .where(StatementItem.statement_id == statement.id)
            .distinct()
        ).all()
        invalid_deliveries = [
            delivery_number
            for delivery_number, delivery_date, reconciliation_month in delivery_rows
            if (
                reconciliation_month != payload.statement_month
                if reconciliation_month is not None
                else not period_start <= delivery_date <= period_end
            )
        ]
        if invalid_deliveries:
            raise HTTPException(
                status_code=409,
                detail=(
                    "对账月份必须匹配全部回单的归属月份；不匹配的送货单："
                    + "、".join(invalid_deliveries)
                ),
            )
        invalid_charges = db.scalars(
            select(CustomerCharge.display_name)
            .join(
                StatementItem,
                StatementItem.customer_charge_id == CustomerCharge.id,
            )
            .where(
                StatementItem.statement_id == statement.id,
                CustomerCharge.reconciliation_month != payload.statement_month,
            )
        ).all()
        if invalid_charges:
            raise HTTPException(
                status_code=409,
                detail=(
                    "对账月份必须匹配全部附加收费的归属月份；不匹配的收费："
                    + "、".join(invalid_charges)
                ),
            )
        before = _statement_detail_response(db, statement.id, user)
        statement.statement_month = payload.statement_month
        statement.statement_number = _next_statement_number(db, payload.statement_month)
        _audit(
            db,
            user=user,
            action="UPDATE_STATEMENT",
            resource="Statement",
            entity_id=statement.id,
            details={"before": before, "after": {"statement_month": payload.statement_month}},
            description="修改月结对账单基础信息",
        )
        db.commit()
        return _statement_detail_response(db, statement.id, user)
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise


@router.post("/statements/{statement_id}/cancel")
def cancel_statement(
    statement_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    try:
        statement = _statement_for_user(db, statement_id, user)
        if statement.confirmation_status == "confirmed":
            raise HTTPException(
                status_code=409,
                detail="对账单已确认；如需取消，请先作废对应开票任务并重新核对。",
            )
        if db.scalar(
            select(Invoice.id).where(Invoice.statement_id == statement.id).limit(1)
        ) is not None:
            raise HTTPException(
                status_code=409,
                detail="该对账单已有开票记录，请先取消开票后再取消对账单。",
            )
        if db.scalar(
            select(SettlementRecord.id)
            .where(SettlementRecord.statement_id == statement.id)
            .limit(1)
        ) is not None:
            raise HTTPException(
                status_code=409,
                detail="该对账单已有收款记录，请先撤销收款后再取消对账单。",
            )
        before = _statement_detail_response(db, statement.id, user)
        db.execute(delete(StatementItem).where(StatementItem.statement_id == statement.id))
        db.execute(delete(Statement).where(Statement.id == statement.id))
        _audit(
            db,
            user=user,
            action="CANCEL_STATEMENT",
            resource="Statement",
            entity_id=statement.id,
            details={"before": before},
            description="取消月结对账单",
        )
        db.commit()
        return {"status": "cancelled", "statement_id": statement_id}
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise


def _payable_response(row: FinancePayable) -> dict:
    return {
        "id": row.id,
        "supplier_id": row.supplier_id,
        "counterparty_name": row.counterparty_name,
        "category": row.category,
        "document_number": row.document_number,
        "document_date": row.document_date,
        "due_date": row.due_date,
        "amount": row.amount,
        "status": row.status,
        "note": row.note,
        "version": row.version,
        "confirmed_at": row.confirmed_at,
        "paid_at": row.paid_at,
        "created_at": row.created_at,
    }


@router.get("/payables")
def list_payables(
    status_filter: str | None = Query(default=None, alias="status"),
    month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    query = select(FinancePayable).order_by(
        FinancePayable.document_date.desc(), FinancePayable.id.desc()
    )
    if status_filter:
        if status_filter not in {"draft", "confirmed", "paid", "voided"}:
            raise HTTPException(status_code=400, detail="应付状态无效")
        query = query.where(FinancePayable.status == status_filter)
    if month:
        start = date.fromisoformat(f"{month}-01")
        end = date.fromisoformat(f"{_shift_month(month, 1)}-01")
        query = query.where(
            FinancePayable.document_date >= start,
            FinancePayable.document_date < end,
        )
    rows = db.scalars(query.limit(500)).all()
    return {"items": [_payable_response(row) for row in rows]}


@router.get("/payable-suppliers")
def list_payable_suppliers(
    db: Session = Depends(get_db),
    user: User = Depends(can_read),
) -> dict:
    rows = db.scalars(
        select(Supplier)
        .where(Supplier.is_active.is_(True))
        .order_by(Supplier.sort_order, Supplier.id)
    ).all()
    return {
        "items": [
            {
                "id": row.id,
                "name": row.display_name or row.standard_name,
                "standard_name": row.standard_name,
            }
            for row in rows
        ]
    }


@router.post("/payables", status_code=status.HTTP_201_CREATED)
def create_payable(
    payload: PayableCreate,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    existing = db.scalar(
        select(FinancePayable).where(
            FinancePayable.idempotency_key == payload.idempotency_key
        )
    )
    if existing is not None:
        same = (
            existing.supplier_id == payload.supplier_id
            and existing.counterparty_name == payload.counterparty_name
            and existing.category == payload.category
            and existing.document_date == payload.document_date
            and Decimal(str(existing.amount)) == payload.amount
        )
        if not same:
            raise HTTPException(status_code=409, detail="幂等键已用于另一笔应付/支出")
        return _payable_response(existing)
    if payload.supplier_id is not None:
        supplier = db.get(Supplier, payload.supplier_id)
        if supplier is None or not supplier.is_active:
            raise HTTPException(status_code=409, detail="供应商不存在或已停用")
    row = FinancePayable(
        **payload.model_dump(),
        status="draft",
        created_by=user.id,
    )
    db.add(row)
    db.flush()
    _audit(
        db,
        user=user,
        action="CREATE_FINANCE_PAYABLE",
        resource="FinancePayable",
        entity_id=row.id,
        details={"amount": row.amount, "category": row.category},
        description="登记供应商应付或经营支出草稿",
    )
    db.commit()
    return _payable_response(row)


def _transition_payable(
    db: Session,
    *,
    payable_id: int,
    payload: PayableTransition,
    action: str,
    user: User,
) -> dict:
    row = db.get(FinancePayable, payable_id)
    if row is None:
        raise HTTPException(status_code=404, detail="应付/支出记录不存在")
    if row.version != payload.expected_version:
        raise HTTPException(
            status_code=409,
            detail={"message": "记录版本已变化", "current_version": row.version},
        )
    if action == "confirm":
        if row.status == "draft":
            row.status = "confirmed"
            row.confirmed_by = user.id
            row.confirmed_at = datetime.now()
        elif row.status != "confirmed":
            raise HTTPException(status_code=409, detail="只有草稿可以确认")
    elif action == "paid":
        if row.status == "confirmed":
            row.status = "paid"
            row.paid_by = user.id
            row.paid_at = datetime.now()
        elif row.status != "paid":
            raise HTTPException(status_code=409, detail="只有已确认应付可以标记已付")
    elif action == "void":
        if row.status == "paid":
            raise HTTPException(status_code=409, detail="已付记录不能直接作废")
        if row.status != "voided":
            if not (payload.reason or "").strip():
                raise HTTPException(status_code=422, detail="作废必须填写原因")
            row.status = "voided"
            row.voided_by = user.id
            row.voided_at = datetime.now()
            row.note = "；".join(
                part for part in (row.note, f"作废原因：{payload.reason.strip()}") if part
            )
    row.version += 1
    _audit(
        db,
        user=user,
        action=f"{action.upper()}_FINANCE_PAYABLE",
        resource="FinancePayable",
        entity_id=row.id,
        details={"status": row.status, "reason": payload.reason},
        description="更新应付/支出状态",
    )
    db.commit()
    return _payable_response(row)


@router.post("/payables/{payable_id}/confirm")
def confirm_payable(
    payable_id: int,
    payload: PayableTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    return _transition_payable(
        db, payable_id=payable_id, payload=payload, action="confirm", user=user
    )


@router.post("/payables/{payable_id}/paid")
def mark_payable_paid(
    payable_id: int,
    payload: PayableTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    return _transition_payable(
        db, payable_id=payable_id, payload=payload, action="paid", user=user
    )


@router.post("/payables/{payable_id}/void")
def void_payable(
    payable_id: int,
    payload: PayableTransition,
    db: Session = Depends(get_db),
    user: User = Depends(can_operate),
) -> dict:
    return _transition_payable(
        db, payable_id=payable_id, payload=payload, action="void", user=user
    )


@router.get("/overview")
def finance_overview(
    through_month: str | None = Query(default=None, pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    user: User = Depends(require_company_finance_read),
) -> dict:
    through = through_month or beijing_today().strftime("%Y-%m")
    months = [_shift_month(through, offset) for offset in range(-5, 1)]
    visible_customer_ids = _visible_customer_ids(user, db)
    unauthorized_statement_source = None
    if visible_customer_ids is not None:
        unauthorized_statement_source = (
            select(StatementItem.id)
            .where(
                StatementItem.statement_id == Statement.id,
                StatementItem.source_customer_id.is_not(None),
                StatementItem.source_customer_id.not_in(visible_customer_ids),
            )
            .exists()
        )
    trend = {
        month: {
            "month": month,
            "confirmed_statement_amount": Decimal("0"),
            "issued_invoice_amount": Decimal("0"),
            "payable_amount": Decimal("0"),
            "other_expense_amount": Decimal("0"),
        }
        for month in months
    }
    statement_query = select(
        Statement.statement_month, Statement.total_receivable
    ).where(
            Statement.confirmation_status == "confirmed",
            Statement.statement_month.in_(months),
        )
    if visible_customer_ids is not None:
        statement_query = statement_query.where(
            Statement.customer_id.in_(visible_customer_ids),
            ~unauthorized_statement_source,
        )
    statement_rows = db.execute(statement_query).all()
    for row_month, amount in statement_rows:
        trend[row_month]["confirmed_statement_amount"] += Decimal(str(amount))
    invoice_query = (
        select(Statement.statement_month, Invoice.invoice_amount)
        .join(Statement, Statement.id == Invoice.statement_id)
        .where(
            Invoice.invoice_status == "issued",
            Statement.confirmation_status == "confirmed",
            Statement.statement_month.in_(months),
        )
    )
    if visible_customer_ids is not None:
        invoice_query = invoice_query.where(
            Statement.customer_id.in_(visible_customer_ids),
            ~unauthorized_statement_source,
        )
    invoice_rows = db.execute(invoice_query).all()
    for statement_month, amount in invoice_rows:
        trend[statement_month]["issued_invoice_amount"] += Decimal(str(amount))
    payable_rows = db.scalars(
        select(FinancePayable).where(
            FinancePayable.status.in_(("confirmed", "paid")),
            FinancePayable.document_date >= date.fromisoformat(f"{months[0]}-01"),
            FinancePayable.document_date < date.fromisoformat(
                f"{_shift_month(through, 1)}-01"
            ),
        )
    ).all()
    category_totals: dict[str, Decimal] = {}
    for row in payable_rows:
        row_month = row.document_date.strftime("%Y-%m")
        amount = Decimal(str(row.amount))
        trend[row_month]["payable_amount"] += amount
        if row.category not in {"material", "outsourcing"}:
            trend[row_month]["other_expense_amount"] += amount
        category_totals[row.category] = category_totals.get(
            row.category, Decimal("0")
        ) + amount

    today = beijing_today()
    aging = {"not_due": Decimal("0"), "overdue_1_30": Decimal("0"), "overdue_31_60": Decimal("0"), "overdue_61_plus": Decimal("0")}
    open_payables = db.scalars(
        select(FinancePayable).where(FinancePayable.status == "confirmed")
    ).all()
    for row in open_payables:
        amount = Decimal(str(row.amount))
        if row.due_date is None or row.due_date >= today:
            aging["not_due"] += amount
            continue
        days = (today - row.due_date).days
        if days <= 30:
            aging["overdue_1_30"] += amount
        elif days <= 60:
            aging["overdue_31_60"] += amount
        else:
            aging["overdue_61_plus"] += amount

    cost_query = (
        select(
            StatementItem.receivable_amount,
            StatementItem.gross_profit_amount,
            StatementItem.unit_cost_snapshot,
        )
        .join(Statement, Statement.id == StatementItem.statement_id)
        .where(
            Statement.confirmation_status == "confirmed",
            Statement.statement_month.in_(months),
            StatementItem.return_receipt_item_id.is_not(None),
        )
    )
    if visible_customer_ids is not None:
        cost_query = cost_query.where(
            Statement.customer_id.in_(visible_customer_ids),
            ~unauthorized_statement_source,
        )
    cost_rows = db.execute(cost_query).all()
    covered = [row for row in cost_rows if Decimal(str(row.unit_cost_snapshot)) > 0]
    can_view_costs = has_permission(user, "cost.view")
    return {
        "through_month": through,
        "trend": [
            {key: (value.quantize(MONEY) if isinstance(value, Decimal) else value) for key, value in trend[month].items()}
            for month in months
        ],
        "payable_aging": {key: value.quantize(MONEY) for key, value in aging.items()},
        "expense_structure": [
            {"category": key, "amount": value.quantize(MONEY)}
            for key, value in sorted(category_totals.items(), key=lambda item: item[1], reverse=True)
        ],
        "cost_coverage": {
            "covered_lines": len(covered),
            "total_lines": len(cost_rows),
            "coverage_rate": round(len(covered) / len(cost_rows), 4) if cost_rows else 0,
            "covered_revenue": (
                sum((Decimal(str(row.receivable_amount)) for row in covered), Decimal("0")).quantize(MONEY)
                if can_view_costs else None
            ),
            "material_gross_profit_reference": (
                sum((Decimal(str(row.gross_profit_amount)) for row in covered), Decimal("0")).quantize(MONEY)
                if can_view_costs else None
            ),
            "label": (
                "材料毛利参考（不含人工、能耗等）"
                if can_view_costs else "材料毛利参考（无成本查看权限）"
            ),
        },
        "collection_note": "客户收款不在 ERP 内核销；已开票金额仅代表开票事实。",
    }
