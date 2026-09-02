from __future__ import annotations

import calendar
import json
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Any, Callable
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query, Request, status
from fastapi.encoders import jsonable_encoder
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field, field_validator, model_validator
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.api.cost_accounting import (
    _audit,
    _idempotency_replay,
    _money,
    _month,
    _record_idempotency,
    _request_hash,
    _validated_center,
    require_cost_export,
    require_cost_manage,
    require_cost_read,
)
from app.api.deps import PermissionChecker, get_db, has_unrestricted_customer_access
from app.core.time_contract import beijing_now_naive
from app.models.customer import Customer
from app.models.finance import Statement
from app.models.finance_cost import FinanceCostCenter, FinanceCostPoolEntry
from app.models.finance_payable import FinancePayable
from app.models.finance_simplified import (
    FinanceAcceptanceNote,
    FinanceRecurringRule,
    FinanceUtilityReading,
)
from app.models.supplier_settlement import (
    SupplierMonthlyAdjustment,
    SupplierMonthlyInvoice,
    SupplierMonthlyStatement,
    SupplierMonthlyStatementLine,
)
from app.models.user import User
from app.services.supplier_monthly_settlement import (
    SupplierSettlementError,
    add_payment,
)


router = APIRouter()
MONEY = Decimal("0.01")
THREE_DECIMALS = Decimal("0.001")
FOUR_DECIMALS = Decimal("0.0001")
MAX_MONEY = Decimal("999999999999.99")

RULE_LABELS = {
    "employee_wage": "人员工资",
    "annual_rent": "年度房租",
    "fixed_monthly": "固定月供",
}
RULE_CATEGORIES = {
    "employee_wage": "production_wages",
    "annual_rent": "factory_rent",
    "fixed_monthly": "finance_expense",
}
UTILITY_LABELS = {"water": "水费", "electricity": "电费"}
ACCEPTANCE_LABELS = {
    "held": "在手",
    "endorsed": "已背书供应商",
    "matured": "已到期",
    "returned": "已退回",
    "voided": "已作废",
}
WAGE_CATEGORIES = {
    "production_wages",
    "driver_wages",
    "administrative_wages",
    "finance_wages",
}
FREIGHT_CATEGORIES = {"inbound_freight", "delivery_freight"}
CONFIRMED_STATUSES = {
    "confirmed_pending_invoice",
    "invoiced_pending_payment",
    "partial_payment",
    "paid",
}

can_finance_execute = PermissionChecker("finance.execute")


def require_finance_execute(
    user: User = Depends(can_finance_execute),
    db: Session = Depends(get_db),
) -> User:
    if not has_unrestricted_customer_access(user, db):
        raise HTTPException(
            status_code=403,
            detail="承兑与供应商付款仅允许全客户范围的财务或管理员操作",
        )
    return user


def _text(value: str | None, *, required: bool = False) -> str | None:
    normalized = str(value or "").strip()
    if required and not normalized:
        raise ValueError("必填内容不能为空")
    return normalized or None


def _decimal(value: Decimal | int | str, quantum: Decimal = MONEY) -> Decimal:
    normalized = Decimal(str(value or 0)).quantize(quantum, rounding=ROUND_HALF_UP)
    if not normalized.is_finite():
        raise ValueError("数值格式无效")
    return normalized


def _last_day(month: str, preferred_day: int = 28) -> date:
    year, month_number = (int(part) for part in month.split("-"))
    day = min(preferred_day, calendar.monthrange(year, month_number)[1])
    return date(year, month_number, day)


def _month_bounds(month: str) -> tuple[date, date]:
    year, month_number = (int(part) for part in month.split("-"))
    start = date(year, month_number, 1)
    if month_number == 12:
        return start, date(year + 1, 1, 1)
    return start, date(year, month_number + 1, 1)


def _run_mutation(
    db: Session,
    *,
    request: Request,
    user: User,
    action: str,
    key: str,
    payload: dict[str, Any],
    resource: str,
    description: str,
    operation: Callable[[], tuple[dict[str, Any], int]],
) -> dict[str, Any]:
    request_hash = _request_hash(action, payload)
    replay = _idempotency_replay(
        db,
        idempotency_key=key,
        request_hash=request_hash,
        action=action,
        actor=user,
    )
    if replay is not None:
        return replay
    try:
        response, resource_id = operation()
        response = jsonable_encoder(response, custom_encoder={Decimal: str})
        _audit(
            db,
            request=request,
            user=user,
            action=action,
            resource=resource,
            entity_id=resource_id,
            description=description,
            details={"request": payload, "response": response},
        )
        _record_idempotency(
            db,
            idempotency_key=key,
            request_hash=request_hash,
            action=action,
            actor=user,
            resource_type=resource,
            resource_id=resource_id,
            response=response,
        )
        db.commit()
        return response
    except SupplierSettlementError as error:
        db.rollback()
        raise HTTPException(
            status_code=error.status_code,
            detail={"code": error.code, "message": error.message},
        ) from error
    except IntegrityError as error:
        db.rollback()
        replay = _idempotency_replay(
            db,
            idempotency_key=key,
            request_hash=request_hash,
            action=action,
            actor=user,
        )
        if replay is not None:
            return replay
        raise HTTPException(status_code=409, detail="数据已被其他操作修改，请刷新后重试") from error
    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise


class IdempotentPayload(BaseModel):
    idempotency_key: str = Field(min_length=8, max_length=120)

    @field_validator("idempotency_key")
    @classmethod
    def clean_key(cls, value: str) -> str:
        return value.strip()


class RecurringRuleFields(BaseModel):
    rule_type: str
    name: str = Field(min_length=1, max_length=120)
    role_name: str | None = Field(default=None, max_length=80)
    cost_center_id: int = Field(gt=0)
    start_month: str
    end_month: str | None = None
    base_amount: Decimal = Field(default=0, ge=0, le=MAX_MONEY)
    allowance_amount: Decimal = Field(default=0, ge=0, le=MAX_MONEY)
    employer_social_amount: Decimal = Field(default=0, ge=0, le=MAX_MONEY)
    deduction_amount: Decimal = Field(default=0, ge=0, le=MAX_MONEY)
    annual_amount: Decimal = Field(default=0, ge=0, le=MAX_MONEY)
    monthly_amount: Decimal = Field(default=0, ge=0, le=MAX_MONEY)
    due_day: int = Field(default=20, ge=1, le=28)
    note: str | None = Field(default=None, max_length=1000)
    is_active: bool = True

    @field_validator("rule_type")
    @classmethod
    def valid_type(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in RULE_LABELS:
            raise ValueError("固定费用规则类型无效")
        return normalized

    @field_validator("name")
    @classmethod
    def clean_name(cls, value: str) -> str:
        return value.strip()

    @field_validator("role_name", "note")
    @classmethod
    def clean_optional(cls, value: str | None) -> str | None:
        return _text(value)

    @field_validator("start_month")
    @classmethod
    def valid_start_month(cls, value: str) -> str:
        return _month(value)

    @field_validator("end_month")
    @classmethod
    def valid_end_month(cls, value: str | None) -> str | None:
        return _month(value) if value else None

    @field_validator(
        "base_amount",
        "allowance_amount",
        "employer_social_amount",
        "deduction_amount",
        "annual_amount",
        "monthly_amount",
    )
    @classmethod
    def normalize_amount(cls, value: Decimal) -> Decimal:
        return _money(value)

    @model_validator(mode="after")
    def valid_amounts(self):
        if self.end_month and self.end_month < self.start_month:
            raise ValueError("结束月份不能早于开始月份")
        if self.rule_type == "employee_wage":
            if self.base_amount + self.allowance_amount + self.employer_social_amount - self.deduction_amount <= 0:
                raise ValueError("人员工资合计必须大于 0")
        elif self.rule_type == "annual_rent":
            if self.annual_amount <= 0:
                raise ValueError("年度房租必须大于 0")
        elif self.monthly_amount <= 0:
            raise ValueError("固定月供必须大于 0")
        return self


class RecurringRuleCreate(RecurringRuleFields, IdempotentPayload):
    pass


class RecurringRuleUpdate(RecurringRuleFields, IdempotentPayload):
    expected_version: int = Field(gt=0)


class RecurringGenerate(IdempotentPayload):
    month: str

    @field_validator("month")
    @classmethod
    def valid_month(cls, value: str) -> str:
        return _month(value)


class UtilityReadingPayload(IdempotentPayload):
    cost_month: str
    utility_type: str
    cost_center_id: int = Field(gt=0)
    previous_reading: Decimal = Field(ge=0)
    current_reading: Decimal = Field(ge=0)
    unit_price: Decimal = Field(ge=0)
    invoice_number: str | None = Field(default=None, max_length=120)
    invoice_date: date | None = None
    invoice_amount: Decimal | None = Field(default=None, ge=0, le=MAX_MONEY)
    paid_amount: Decimal = Field(default=0, ge=0, le=MAX_MONEY)
    payment_date: date | None = None
    note: str | None = Field(default=None, max_length=1000)
    expected_version: int | None = Field(default=None, gt=0)

    @field_validator("cost_month")
    @classmethod
    def valid_month(cls, value: str) -> str:
        return _month(value)

    @field_validator("utility_type")
    @classmethod
    def valid_type(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in UTILITY_LABELS:
            raise ValueError("水电类型无效")
        return normalized

    @field_validator("previous_reading", "current_reading")
    @classmethod
    def normalize_reading(cls, value: Decimal) -> Decimal:
        return _decimal(value, THREE_DECIMALS)

    @field_validator("unit_price")
    @classmethod
    def normalize_price(cls, value: Decimal) -> Decimal:
        return _decimal(value, FOUR_DECIMALS)

    @field_validator("invoice_amount", "paid_amount")
    @classmethod
    def normalize_money(cls, value: Decimal | None) -> Decimal | None:
        return _money(value) if value is not None else None

    @field_validator("invoice_number", "note")
    @classmethod
    def clean_optional(cls, value: str | None) -> str | None:
        return _text(value)

    @model_validator(mode="after")
    def valid_values(self):
        if self.current_reading < self.previous_reading:
            raise ValueError("本月表数不能小于上月表数")
        calculated = _money((self.current_reading - self.previous_reading) * self.unit_price)
        recognized = self.invoice_amount if self.invoice_amount is not None else calculated
        if recognized <= 0:
            raise ValueError("水电计算金额或发票金额必须大于 0")
        if self.paid_amount > recognized:
            raise ValueError("已付金额不能超过本月确认金额")
        if self.paid_amount > 0 and self.payment_date is None:
            raise ValueError("填写已付金额时必须填写付款日期")
        return self


class AcceptanceCreate(IdempotentPayload):
    bill_number: str = Field(min_length=1, max_length=120)
    customer_id: int = Field(gt=0)
    customer_statement_id: int | None = Field(default=None, gt=0)
    amount: Decimal = Field(gt=0, le=MAX_MONEY)
    received_date: date
    maturity_date: date
    note: str | None = Field(default=None, max_length=1000)

    @field_validator("bill_number")
    @classmethod
    def clean_number(cls, value: str) -> str:
        return value.strip()

    @field_validator("amount")
    @classmethod
    def normalize_amount(cls, value: Decimal) -> Decimal:
        return _money(value)

    @field_validator("note")
    @classmethod
    def clean_note(cls, value: str | None) -> str | None:
        return _text(value)

    @model_validator(mode="after")
    def valid_dates(self):
        if self.maturity_date < self.received_date:
            raise ValueError("承兑到期日不能早于收到日期")
        return self


class AcceptanceEndorse(IdempotentPayload):
    expected_version: int = Field(gt=0)
    supplier_statement_id: int = Field(gt=0)
    endorsement_date: date


class AcceptanceTransition(IdempotentPayload):
    expected_version: int = Field(gt=0)
    action: str

    @field_validator("action")
    @classmethod
    def valid_action(cls, value: str) -> str:
        normalized = value.strip()
        if normalized not in {"matured", "returned", "voided"}:
            raise ValueError("承兑状态操作无效")
        return normalized


def _rule_amount(row: FinanceRecurringRule, month: str) -> Decimal:
    if row.rule_type == "employee_wage":
        return _money(
            row.base_amount
            + row.allowance_amount
            + row.employer_social_amount
            - row.deduction_amount
        )
    if row.rule_type == "annual_rent":
        annual = _money(row.annual_amount)
        monthly = (annual / Decimal("12")).quantize(MONEY, rounding=ROUND_HALF_UP)
        return _money(annual - monthly * 11) if month.endswith("-12") else monthly
    return _money(row.monthly_amount)


def _rule_response(row: FinanceRecurringRule) -> dict[str, Any]:
    return {
        "id": row.id,
        "rule_type": row.rule_type,
        "rule_type_label": RULE_LABELS[row.rule_type],
        "name": row.name,
        "role_name": row.role_name,
        "cost_center_id": row.cost_center_id,
        "cost_category": row.cost_category,
        "start_month": row.start_month,
        "end_month": row.end_month,
        "base_amount": row.base_amount,
        "allowance_amount": row.allowance_amount,
        "employer_social_amount": row.employer_social_amount,
        "deduction_amount": row.deduction_amount,
        "annual_amount": row.annual_amount,
        "monthly_amount": row.monthly_amount,
        "due_day": row.due_day,
        "note": row.note,
        "is_active": row.is_active,
        "version": row.version,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
    }


def _rule_values(payload: RecurringRuleFields) -> dict[str, Any]:
    values = payload.model_dump(exclude={"idempotency_key", "expected_version"})
    values["cost_category"] = RULE_CATEGORIES[payload.rule_type]
    if payload.rule_type != "employee_wage":
        values.update(
            base_amount=Decimal("0"),
            allowance_amount=Decimal("0"),
            employer_social_amount=Decimal("0"),
            deduction_amount=Decimal("0"),
            role_name=None,
        )
    if payload.rule_type != "annual_rent":
        values["annual_amount"] = Decimal("0")
    if payload.rule_type != "fixed_monthly":
        values["monthly_amount"] = Decimal("0")
    return values


def _utility_response(row: FinanceUtilityReading) -> dict[str, Any]:
    return {
        "id": row.id,
        "cost_month": row.cost_month,
        "utility_type": row.utility_type,
        "utility_type_label": UTILITY_LABELS[row.utility_type],
        "cost_center_id": row.cost_center_id,
        "previous_reading": row.previous_reading,
        "current_reading": row.current_reading,
        "usage_quantity": row.usage_quantity,
        "unit_price": row.unit_price,
        "calculated_amount": row.calculated_amount,
        "invoice_number": row.invoice_number,
        "invoice_date": row.invoice_date,
        "invoice_amount": row.invoice_amount,
        "recognized_amount": row.invoice_amount or row.calculated_amount,
        "difference_amount": (
            _money(row.invoice_amount - row.calculated_amount)
            if row.invoice_amount is not None
            else None
        ),
        "paid_amount": row.paid_amount,
        "payment_date": row.payment_date,
        "cost_pool_entry_id": row.cost_pool_entry_id,
        "note": row.note,
        "version": row.version,
    }


def _acceptance_response(row: FinanceAcceptanceNote) -> dict[str, Any]:
    return {
        "id": row.id,
        "bill_number": row.bill_number,
        "customer_id": row.customer_id,
        "customer_name": row.customer_name_snapshot,
        "customer_statement_id": row.customer_statement_id,
        "amount": row.amount,
        "received_date": row.received_date,
        "maturity_date": row.maturity_date,
        "status": row.status,
        "status_label": ACCEPTANCE_LABELS[row.status],
        "supplier_id": row.supplier_id,
        "supplier_name": row.supplier_name_snapshot,
        "supplier_statement_id": row.supplier_statement_id,
        "supplier_payment_id": row.supplier_payment_id,
        "endorsed_date": row.endorsed_date,
        "note": row.note,
        "version": row.version,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
    }


@router.get("/simple-finance/metadata")
def simple_finance_metadata(
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_read),
) -> dict[str, Any]:
    centers = list(
        db.scalars(
            select(FinanceCostCenter)
            .where(FinanceCostCenter.is_active.is_(True))
            .order_by(FinanceCostCenter.id)
        ).all()
    )
    customers = list(
        db.scalars(
            select(Customer)
            .where(Customer.is_active.is_(True))
            .order_by(Customer.name)
        ).all()
    )
    statements = list(
        db.scalars(
            select(SupplierMonthlyStatement)
            .where(
                SupplierMonthlyStatement.active_guard == 1,
                SupplierMonthlyStatement.status.in_(CONFIRMED_STATUSES - {"paid"}),
            )
            .order_by(SupplierMonthlyStatement.period_end.desc())
            .limit(200)
        ).all()
    )
    return {
        "primary_groups": [
            {"value": "material", "label": "材料采购"},
            {"value": "personnel", "label": "人员工资"},
            {"value": "freight", "label": "运费"},
            {"value": "fixed_operating", "label": "固定 / 经营费用"},
        ],
        "rule_types": [
            {"value": value, "label": label} for value, label in RULE_LABELS.items()
        ],
        "centers": [
            {"id": row.id, "code": row.code, "name": row.name} for row in centers
        ],
        "customers": [
            {
                "id": row.id,
                "name": row.name,
                "short_name": row.chinese_short_name or row.name,
            }
            for row in customers
        ],
        "supplier_statements": [
            {
                "id": row.id,
                "statement_number": row.statement_number,
                "supplier_id": row.supplier_id,
                "supplier_name": row.supplier_name_snapshot,
                "settlement_month": row.settlement_month,
                "status": row.status,
                "version": row.version,
                "remaining_amount": _money(
                    max(
                        _money(row.confirmed_amount or 0) - _money(row.paid_amount or 0),
                        Decimal("0"),
                    )
                ),
                "invoice_available_amount": _money(
                    max(
                        _money(row.invoice_allocated_amount or 0)
                        - _money(row.paid_amount or 0),
                        Decimal("0"),
                    )
                ),
            }
            for row in statements
        ],
        "note": "自动延续只生成草稿；承兑背书不是银行现金付款。",
    }


@router.get("/simple-finance/recurring-rules")
def list_recurring_rules(
    include_inactive: bool = True,
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_read),
) -> dict[str, Any]:
    query = select(FinanceRecurringRule)
    if not include_inactive:
        query = query.where(FinanceRecurringRule.is_active.is_(True))
    rows = list(db.scalars(query.order_by(FinanceRecurringRule.rule_type, FinanceRecurringRule.id)).all())
    return {"items": [_rule_response(row) for row in rows]}


@router.post("/simple-finance/recurring-rules", status_code=status.HTTP_201_CREATED)
def create_recurring_rule(
    payload: RecurringRuleCreate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict[str, Any]:
    values = _rule_values(payload)
    _validated_center(
        db,
        payload.cost_center_id,
        accounting_class=(
            "manufacturing" if payload.rule_type != "fixed_monthly" else "finance"
        ),
    )

    def operation():
        row = FinanceRecurringRule(**values, created_by=user.id)
        db.add(row)
        db.flush()
        return _rule_response(row), row.id

    return _run_mutation(
        db,
        request=request,
        user=user,
        action="finance.simple.recurring_rule.create",
        key=payload.idempotency_key,
        payload=values,
        resource="FinanceRecurringRule",
        description="新增固定费用自动延续规则",
        operation=operation,
    )


@router.put("/simple-finance/recurring-rules/{rule_id}")
def update_recurring_rule(
    rule_id: int,
    payload: RecurringRuleUpdate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict[str, Any]:
    values = _rule_values(payload)
    _validated_center(
        db,
        payload.cost_center_id,
        accounting_class=(
            "manufacturing" if payload.rule_type != "fixed_monthly" else "finance"
        ),
    )

    def operation():
        row = db.get(FinanceRecurringRule, rule_id)
        if row is None:
            raise HTTPException(status_code=404, detail="固定费用规则不存在")
        if row.version != payload.expected_version:
            raise HTTPException(status_code=409, detail="规则已被修改，请刷新后重试")
        for key, value in values.items():
            setattr(row, key, value)
        row.version += 1
        row.updated_by = user.id
        row.updated_at = beijing_now_naive()
        db.flush()
        return _rule_response(row), row.id

    return _run_mutation(
        db,
        request=request,
        user=user,
        action="finance.simple.recurring_rule.update",
        key=payload.idempotency_key,
        payload={"rule_id": rule_id, **values, "expected_version": payload.expected_version},
        resource="FinanceRecurringRule",
        description="更新固定费用自动延续规则",
        operation=operation,
    )


@router.post("/simple-finance/recurring-generate")
def generate_recurring_drafts(
    payload: RecurringGenerate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict[str, Any]:
    rules = list(
        db.scalars(
            select(FinanceRecurringRule)
            .where(
                FinanceRecurringRule.is_active.is_(True),
                FinanceRecurringRule.start_month <= payload.month,
                (FinanceRecurringRule.end_month.is_(None))
                | (FinanceRecurringRule.end_month >= payload.month),
            )
            .order_by(FinanceRecurringRule.id)
        ).all()
    )

    def operation():
        created = updated = preserved = 0
        issues: list[dict[str, Any]] = []
        generated_ids: list[int] = []
        for rule in rules:
            center = db.get(FinanceCostCenter, rule.cost_center_id)
            if center is None or not center.is_active:
                issues.append({"rule_id": rule.id, "message": f"{rule.name} 的成本中心已停用"})
                continue
            amount = _rule_amount(rule, payload.month)
            fingerprint = f"recurring:{rule.id}:{payload.month}"
            entry = db.scalar(
                select(FinanceCostPoolEntry).where(
                    FinanceCostPoolEntry.source_fingerprint == fingerprint
                )
            )
            description = f"{RULE_LABELS[rule.rule_type]} · {rule.name}"
            values = {
                "cost_month": payload.month,
                "document_date": _last_day(payload.month, rule.due_day),
                "cost_center_id": center.id,
                "cost_center_code_snapshot": center.code,
                "cost_center_name_snapshot": center.name,
                "cost_center_type_snapshot": center.center_type,
                "cost_category": rule.cost_category,
                "accounting_class": (
                    "manufacturing"
                    if rule.cost_category in {"production_wages", "factory_rent"}
                    else "finance"
                ),
                "allocation_basis": "unallocated",
                "description": description,
                "counterparty_name": rule.name,
                "document_number": None,
                "amount": amount,
                "tax_amount": Decimal("0"),
                "source_type": "recurring_rule",
                "source_reference": f"固定费用规则 #{rule.id}",
                "source_fingerprint": fingerprint,
                "note": rule.note,
            }
            if entry is None:
                entry = FinanceCostPoolEntry(**values, status="draft", created_by=user.id)
                db.add(entry)
                db.flush()
                created += 1
            elif entry.status == "draft":
                for key, value in values.items():
                    setattr(entry, key, value)
                entry.version += 1
                updated += 1
            else:
                preserved += 1
                issues.append(
                    {
                        "rule_id": rule.id,
                        "cost_pool_entry_id": entry.id,
                        "message": f"{rule.name} 已{('确认' if entry.status == 'confirmed' else '作废')}，本次未覆盖",
                    }
                )
            generated_ids.append(entry.id)
        employee_count = sum(rule.rule_type == "employee_wage" for rule in rules)
        response = {
            "month": payload.month,
            "created_count": created,
            "updated_count": updated,
            "preserved_count": preserved,
            "active_employee_count": employee_count,
            "generated_entry_ids": generated_ids,
            "issues": issues,
            "message": "固定费用已生成待确认草稿，不会自动确认或付款。",
        }
        return response, generated_ids[0] if generated_ids else 0

    return _run_mutation(
        db,
        request=request,
        user=user,
        action="finance.simple.recurring.generate",
        key=payload.idempotency_key,
        payload={"month": payload.month},
        resource="FinanceCostPoolEntry",
        description="按月生成固定费用和人员工资草稿",
        operation=operation,
    )


@router.get("/simple-finance/utility-readings")
def list_utility_readings(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_read),
) -> dict[str, Any]:
    normalized_month = _month(month)
    rows = list(
        db.scalars(
            select(FinanceUtilityReading)
            .where(FinanceUtilityReading.cost_month == normalized_month)
            .order_by(FinanceUtilityReading.utility_type)
        ).all()
    )
    previous: dict[str, Decimal] = {}
    for utility_type in UTILITY_LABELS:
        last = db.scalar(
            select(FinanceUtilityReading)
            .where(
                FinanceUtilityReading.utility_type == utility_type,
                FinanceUtilityReading.cost_month < normalized_month,
            )
            .order_by(FinanceUtilityReading.cost_month.desc())
            .limit(1)
        )
        previous[utility_type] = last.current_reading if last is not None else Decimal("0")
    return {
        "month": normalized_month,
        "items": [_utility_response(row) for row in rows],
        "recommended_previous_readings": previous,
    }


@router.post("/simple-finance/utility-readings")
def save_utility_reading(
    payload: UtilityReadingPayload,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_cost_manage),
) -> dict[str, Any]:
    center = _validated_center(db, payload.cost_center_id, accounting_class="manufacturing")
    usage = _decimal(payload.current_reading - payload.previous_reading, THREE_DECIMALS)
    calculated = _money(usage * payload.unit_price)
    recognized = payload.invoice_amount if payload.invoice_amount is not None else calculated
    existing = db.scalar(
        select(FinanceUtilityReading).where(
            FinanceUtilityReading.cost_month == payload.cost_month,
            FinanceUtilityReading.utility_type == payload.utility_type,
        )
    )

    def operation():
        nonlocal existing
        if existing is not None and existing.version != payload.expected_version:
            raise HTTPException(status_code=409, detail="水电记录已被修改，请刷新后重试")
        entry = (
            db.get(FinanceCostPoolEntry, existing.cost_pool_entry_id)
            if existing is not None
            else None
        )
        if entry is not None and entry.status != "draft":
            raise HTTPException(status_code=409, detail="本月水电费用已确认或作废，不能覆盖")
        fingerprint = f"utility:{payload.utility_type}:{payload.cost_month}"
        entry_values = {
            "cost_month": payload.cost_month,
            "document_date": payload.invoice_date or _last_day(payload.cost_month),
            "cost_center_id": center.id,
            "cost_center_code_snapshot": center.code,
            "cost_center_name_snapshot": center.name,
            "cost_center_type_snapshot": center.center_type,
            "cost_category": "factory_utilities",
            "accounting_class": "manufacturing",
            "allocation_basis": "unallocated",
            "description": f"{payload.cost_month} {UTILITY_LABELS[payload.utility_type]}",
            "counterparty_name": None,
            "document_number": payload.invoice_number,
            "amount": recognized,
            "tax_amount": Decimal("0"),
            "source_type": "utility_reading",
            "source_reference": f"{UTILITY_LABELS[payload.utility_type]}表数",
            "source_fingerprint": fingerprint,
            "note": payload.note,
        }
        if entry is None:
            entry = FinanceCostPoolEntry(**entry_values, status="draft", created_by=user.id)
            db.add(entry)
            db.flush()
        else:
            for key, value in entry_values.items():
                setattr(entry, key, value)
            entry.version += 1
        reading_values = {
            "cost_month": payload.cost_month,
            "utility_type": payload.utility_type,
            "cost_center_id": center.id,
            "previous_reading": payload.previous_reading,
            "current_reading": payload.current_reading,
            "usage_quantity": usage,
            "unit_price": payload.unit_price,
            "calculated_amount": calculated,
            "invoice_number": payload.invoice_number,
            "invoice_date": payload.invoice_date,
            "invoice_amount": payload.invoice_amount,
            "paid_amount": payload.paid_amount,
            "payment_date": payload.payment_date,
            "cost_pool_entry_id": entry.id,
            "note": payload.note,
        }
        if existing is None:
            existing = FinanceUtilityReading(**reading_values, created_by=user.id)
            db.add(existing)
        else:
            for key, value in reading_values.items():
                setattr(existing, key, value)
            existing.version += 1
            existing.updated_by = user.id
            existing.updated_at = beijing_now_naive()
        db.flush()
        return _utility_response(existing), existing.id

    return _run_mutation(
        db,
        request=request,
        user=user,
        action="finance.simple.utility.save",
        key=payload.idempotency_key,
        payload=payload.model_dump(exclude={"idempotency_key"}),
        resource="FinanceUtilityReading",
        description="保存月度水电表数和发票付款事实",
        operation=operation,
    )


@router.get("/simple-finance/acceptances")
def list_acceptance_notes(
    status_filter: str | None = Query(default=None, alias="status"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_read),
) -> dict[str, Any]:
    query = select(FinanceAcceptanceNote)
    if status_filter:
        if status_filter not in ACCEPTANCE_LABELS:
            raise HTTPException(status_code=422, detail="承兑状态无效")
        query = query.where(FinanceAcceptanceNote.status == status_filter)
    rows = list(
        db.scalars(
            query.order_by(FinanceAcceptanceNote.maturity_date, FinanceAcceptanceNote.id)
        ).all()
    )
    held = [row for row in rows if row.status == "held"]
    return {
        "items": [_acceptance_response(row) for row in rows],
        "held_count": len(held),
        "held_amount": _money(sum((row.amount for row in held), Decimal("0"))),
    }


@router.post("/simple-finance/acceptances", status_code=status.HTTP_201_CREATED)
def create_acceptance_note(
    payload: AcceptanceCreate,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_finance_execute),
) -> dict[str, Any]:
    customer = db.get(Customer, payload.customer_id)
    if customer is None or not customer.is_active:
        raise HTTPException(status_code=409, detail="客户不存在或已停用")
    if payload.customer_statement_id is not None:
        customer_statement = db.get(Statement, payload.customer_statement_id)
        if customer_statement is None or customer_statement.customer_id != customer.id:
            raise HTTPException(status_code=409, detail="客户对账单与承兑客户不一致")

    def operation():
        row = FinanceAcceptanceNote(
            bill_number=payload.bill_number,
            customer_id=customer.id,
            customer_name_snapshot=customer.chinese_short_name or customer.name,
            customer_statement_id=payload.customer_statement_id,
            amount=payload.amount,
            received_date=payload.received_date,
            maturity_date=payload.maturity_date,
            status="held",
            note=payload.note,
            created_by=user.id,
        )
        db.add(row)
        db.flush()
        return _acceptance_response(row), row.id

    return _run_mutation(
        db,
        request=request,
        user=user,
        action="finance.simple.acceptance.create",
        key=payload.idempotency_key,
        payload=payload.model_dump(exclude={"idempotency_key"}),
        resource="FinanceAcceptanceNote",
        description="登记收到客户承兑",
        operation=operation,
    )


@router.post("/simple-finance/acceptances/{acceptance_id}/endorse")
def endorse_acceptance_note(
    acceptance_id: int,
    payload: AcceptanceEndorse,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_finance_execute),
) -> dict[str, Any]:
    def operation():
        row = db.get(FinanceAcceptanceNote, acceptance_id)
        if row is None:
            raise HTTPException(status_code=404, detail="承兑票据不存在")
        if row.version != payload.expected_version:
            raise HTTPException(status_code=409, detail="承兑票据已被修改，请刷新后重试")
        if row.status != "held":
            raise HTTPException(status_code=409, detail="只有在手承兑可以背书供应商")
        statement = db.get(SupplierMonthlyStatement, payload.supplier_statement_id)
        if statement is None or statement.active_guard != 1:
            raise HTTPException(status_code=404, detail="供应商月结不存在")
        statement, payment = add_payment(
            db,
            statement_id=statement.id,
            expected_version=statement.version,
            payment_date=payload.endorsement_date,
            amount=row.amount,
            reference=f"承兑背书 {row.bill_number}",
            payment_method="acceptance",
            acceptance_note_id=row.id,
            user=user,
        )
        row.status = "endorsed"
        row.supplier_id = statement.supplier_id
        row.supplier_name_snapshot = statement.supplier_name_snapshot
        row.supplier_statement_id = statement.id
        row.supplier_payment_id = payment.id
        row.endorsed_date = payload.endorsement_date
        row.version += 1
        row.updated_by = user.id
        row.updated_at = beijing_now_naive()
        db.flush()
        return {
            "acceptance": _acceptance_response(row),
            "supplier_statement": {
                "id": statement.id,
                "statement_number": statement.statement_number,
                "status": statement.status,
                "paid_amount": statement.paid_amount,
                "version": statement.version,
            },
            "message": "承兑已背书抵付供应商月结；该笔不是银行现金付款。",
        }, row.id

    return _run_mutation(
        db,
        request=request,
        user=user,
        action="finance.simple.acceptance.endorse",
        key=payload.idempotency_key,
        payload={"acceptance_id": acceptance_id, **payload.model_dump(exclude={"idempotency_key"})},
        resource="FinanceAcceptanceNote",
        description="承兑背书抵付供应商月结",
        operation=operation,
    )


@router.post("/simple-finance/acceptances/{acceptance_id}/transition")
def transition_acceptance_note(
    acceptance_id: int,
    payload: AcceptanceTransition,
    request: Request,
    db: Session = Depends(get_db),
    user: User = Depends(require_finance_execute),
) -> dict[str, Any]:
    def operation():
        row = db.get(FinanceAcceptanceNote, acceptance_id)
        if row is None:
            raise HTTPException(status_code=404, detail="承兑票据不存在")
        if row.version != payload.expected_version:
            raise HTTPException(status_code=409, detail="承兑票据已被修改，请刷新后重试")
        if row.status != "held":
            raise HTTPException(status_code=409, detail="只有在手承兑可以到期、退回或作废")
        row.status = payload.action
        row.version += 1
        row.updated_by = user.id
        row.updated_at = beijing_now_naive()
        db.flush()
        return _acceptance_response(row), row.id

    return _run_mutation(
        db,
        request=request,
        user=user,
        action=f"finance.simple.acceptance.{payload.action}",
        key=payload.idempotency_key,
        payload={"acceptance_id": acceptance_id, **payload.model_dump(exclude={"idempotency_key"})},
        resource="FinanceAcceptanceNote",
        description=f"承兑状态更新为{ACCEPTANCE_LABELS[payload.action]}",
        operation=operation,
    )


def _cost_group(category: str) -> str:
    if category in WAGE_CATEGORIES:
        return "personnel"
    if category in FREIGHT_CATEGORIES:
        return "freight"
    return "fixed_operating"


@router.get("/simple-finance/summary")
def simple_finance_summary(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_read),
) -> dict[str, Any]:
    normalized_month = _month(month)
    statements = list(
        db.scalars(
            select(SupplierMonthlyStatement).where(
                SupplierMonthlyStatement.settlement_month == normalized_month,
                SupplierMonthlyStatement.active_guard == 1,
            )
        ).all()
    )
    statement_ids = [row.id for row in statements]
    source_totals = {"paperboard": Decimal("0"), "external_packaging": Decimal("0")}
    if statement_ids:
        for source_type, amount in db.execute(
            select(
                SupplierMonthlyStatementLine.source_type,
                func.coalesce(func.sum(SupplierMonthlyStatementLine.erp_amount), 0),
            )
            .where(
                SupplierMonthlyStatementLine.statement_id.in_(statement_ids),
                SupplierMonthlyStatementLine.active_guard == 1,
            )
            .group_by(SupplierMonthlyStatementLine.source_type)
        ).all():
            source_totals[source_type] = _money(amount)
    linked_payable_ids = {row.finance_payable_id for row in statements if row.finance_payable_id}
    month_start, next_month = _month_bounds(normalized_month)
    standalone_query = select(FinancePayable).where(
        FinancePayable.category == "material",
        FinancePayable.document_date >= month_start,
        FinancePayable.document_date < next_month,
        FinancePayable.status != "voided",
    )
    if linked_payable_ids:
        standalone_query = standalone_query.where(FinancePayable.id.not_in(linked_payable_ids))
    standalone_material = _money(
        sum((row.amount for row in db.scalars(standalone_query).all()), Decimal("0"))
    )
    costs = list(
        db.scalars(
            select(FinanceCostPoolEntry).where(
                FinanceCostPoolEntry.cost_month == normalized_month,
                FinanceCostPoolEntry.status != "voided",
            )
        ).all()
    )
    groups: dict[str, dict[str, Any]] = {
        key: {"draft_amount": Decimal("0"), "confirmed_amount": Decimal("0"), "count": 0}
        for key in ("personnel", "freight", "fixed_operating")
    }
    for row in costs:
        group = groups[_cost_group(row.cost_category)]
        group["count"] += 1
        group[f"{row.status}_amount"] += row.amount
    invoices = list(
        db.scalars(
            select(SupplierMonthlyInvoice).where(
                SupplierMonthlyInvoice.statement_id.in_(statement_ids)
            )
        ).all()
    ) if statement_ids else []
    held = list(
        db.scalars(
            select(FinanceAcceptanceNote).where(FinanceAcceptanceNote.status == "held")
        ).all()
    )
    material_erp = _money(sum((row.adjusted_amount for row in statements), Decimal("0")))
    supplier_actual = _money(
        sum((row.supplier_statement_amount or row.adjusted_amount for row in statements), Decimal("0"))
    )
    invoice_total = _money(sum((row.allocated_amount for row in invoices), Decimal("0")))
    differences: list[dict[str, Any]] = []
    for row in statements:
        supplier_amount = _money(row.supplier_statement_amount or row.adjusted_amount)
        confirmed = _money(row.confirmed_amount or 0)
        row_invoiced = _money(row.invoice_allocated_amount or 0)
        adjustments = list(
            db.scalars(
                select(SupplierMonthlyAdjustment).where(
                    SupplierMonthlyAdjustment.statement_id == row.id
                )
            ).all()
        )
        statement_difference = _money(supplier_amount - row.adjusted_amount)
        invoice_difference = _money(row_invoiced - confirmed)
        if statement_difference or invoice_difference or adjustments:
            differences.append(
                {
                    "statement_id": row.id,
                    "statement_number": row.statement_number,
                    "supplier_name": row.supplier_name_snapshot,
                    "statement_difference_amount": statement_difference,
                    "invoice_difference_amount": invoice_difference,
                    "adjustments": [
                        {
                            "id": item.id,
                            "statement_line_id": item.statement_line_id,
                            "difference_type": item.difference_type,
                            "amount": item.amount,
                            "note": item.note,
                        }
                        for item in adjustments
                    ],
                }
            )
    return {
        "month": normalized_month,
        "primary_groups": {
            "material": {
                "label": "材料采购",
                "erp_amount": material_erp,
                "supplier_statement_amount": supplier_actual,
                "invoice_amount": invoice_total,
                "statement_difference_amount": _money(supplier_actual - material_erp),
                "invoice_difference_amount": _money(invoice_total - supplier_actual),
                "paperboard_amount": source_totals["paperboard"],
                "external_packaging_amount": source_totals["external_packaging"],
                "mold_ink_other_amount": standalone_material,
                "differences": differences,
            },
            "personnel": {"label": "人员工资", **groups["personnel"]},
            "freight": {"label": "运费", **groups["freight"]},
            "fixed_operating": {"label": "固定 / 经营费用", **groups["fixed_operating"]},
        },
        "active_employee_count": int(
            db.scalar(
                select(func.count(FinanceRecurringRule.id)).where(
                    FinanceRecurringRule.rule_type == "employee_wage",
                    FinanceRecurringRule.is_active.is_(True),
                    FinanceRecurringRule.start_month <= normalized_month,
                    (FinanceRecurringRule.end_month.is_(None))
                    | (FinanceRecurringRule.end_month >= normalized_month),
                )
            )
            or 0
        ),
        "held_acceptance_count": len(held),
        "held_acceptance_amount": _money(sum((row.amount for row in held), Decimal("0"))),
        "note": "材料来自供应商20日月结；工资和固定费用仅生成草稿；承兑不计作银行现金。",
    }


def _safe_excel(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


@router.get("/simple-finance/material-export")
def export_monthly_materials(
    month: str = Query(pattern=r"^\d{4}-\d{2}$"),
    db: Session = Depends(get_db),
    _user: User = Depends(require_cost_export),
) -> StreamingResponse:
    normalized_month = _month(month)
    statements = list(
        db.scalars(
            select(SupplierMonthlyStatement)
            .where(
                SupplierMonthlyStatement.settlement_month == normalized_month,
                SupplierMonthlyStatement.active_guard == 1,
            )
            .order_by(SupplierMonthlyStatement.supplier_name_snapshot)
        ).all()
    )
    workbook = Workbook()
    summary = workbook.active
    summary.title = "供应商核票汇总"
    summary.append(
        [
            "供应商",
            "月结单号",
            "ERP金额",
            "调整后金额",
            "供应商账单金额",
            "账单差异",
            "确认应付",
            "已收发票",
            "发票差异",
            "已付",
            "状态",
        ]
    )
    for row in statements:
        supplier_amount = _money(row.supplier_statement_amount or row.adjusted_amount)
        confirmed = _money(row.confirmed_amount or 0)
        summary.append(
            [
                _safe_excel(row.supplier_name_snapshot),
                _safe_excel(row.statement_number),
                float(row.erp_amount),
                float(row.adjusted_amount),
                float(supplier_amount),
                float(supplier_amount - row.adjusted_amount),
                float(confirmed),
                float(row.invoice_allocated_amount),
                float(_money(row.invoice_allocated_amount - confirmed)),
                float(row.paid_amount),
                row.status,
            ]
        )
    detail = workbook.create_sheet("纸板与外购包材明细")
    detail.append(
        [
            "供应商",
            "月结单号",
            "来源",
            "采购单",
            "收料单",
            "收料日期",
            "材料/产品",
            "规格",
            "实收数量",
            "单位",
            "冻结单价",
            "价单位",
            "ERP金额",
            "税额",
        ]
    )
    for statement in statements:
        lines = list(
            db.scalars(
                select(SupplierMonthlyStatementLine)
                .where(
                    SupplierMonthlyStatementLine.statement_id == statement.id,
                    SupplierMonthlyStatementLine.active_guard == 1,
                )
                .order_by(SupplierMonthlyStatementLine.receipt_date, SupplierMonthlyStatementLine.id)
            ).all()
        )
        for row in lines:
            detail.append(
                [
                    _safe_excel(statement.supplier_name_snapshot),
                    _safe_excel(statement.statement_number),
                    "纸板" if row.source_type == "paperboard" else "外购包材",
                    _safe_excel(row.purchase_document_number),
                    _safe_excel(row.receipt_number),
                    row.receipt_date,
                    _safe_excel(row.material_or_product_snapshot),
                    _safe_excel(row.specification_snapshot or ""),
                    float(row.received_quantity),
                    row.quantity_unit,
                    float(row.frozen_unit_price),
                    row.price_unit,
                    float(row.erp_amount),
                    float(row.tax_amount),
                ]
            )
    difference_sheet = workbook.create_sheet("差异明细")
    difference_sheet.append(
        [
            "供应商",
            "月结单号",
            "差异来源",
            "关联收料行",
            "差异金额",
            "说明",
        ]
    )
    difference_type_labels = {
        "price": "价格差异",
        "quantity": "数量差异",
        "tax_rounding": "税额尾差",
        "other": "其他调整",
    }
    for statement in statements:
        supplier_amount = _money(
            statement.supplier_statement_amount or statement.adjusted_amount
        )
        statement_difference = _money(supplier_amount - statement.adjusted_amount)
        confirmed = _money(statement.confirmed_amount or 0)
        invoice_difference = _money(statement.invoice_allocated_amount - confirmed)
        if statement_difference:
            difference_sheet.append(
                [
                    _safe_excel(statement.supplier_name_snapshot),
                    _safe_excel(statement.statement_number),
                    "供应商账单与ERP调整后金额",
                    "",
                    float(statement_difference),
                    _safe_excel(statement.supplier_statement_number or ""),
                ]
            )
        adjustments = list(
            db.scalars(
                select(SupplierMonthlyAdjustment)
                .where(SupplierMonthlyAdjustment.statement_id == statement.id)
                .order_by(SupplierMonthlyAdjustment.id)
            ).all()
        )
        for adjustment in adjustments:
            difference_sheet.append(
                [
                    _safe_excel(statement.supplier_name_snapshot),
                    _safe_excel(statement.statement_number),
                    difference_type_labels.get(
                        adjustment.difference_type, adjustment.difference_type
                    ),
                    adjustment.statement_line_id or "",
                    float(adjustment.amount),
                    _safe_excel(adjustment.note or ""),
                ]
            )
        if invoice_difference:
            difference_sheet.append(
                [
                    _safe_excel(statement.supplier_name_snapshot),
                    _safe_excel(statement.statement_number),
                    "已收发票与确认应付",
                    "",
                    float(invoice_difference),
                    "正数为发票多，负数为发票不足",
                ]
            )
    header_fill = PatternFill("solid", fgColor="167D74")
    for sheet in (summary, detail, difference_sheet):
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 36)
            sheet.column_dimensions[column[0].column_letter].width = width
    stream = BytesIO()
    workbook.save(stream)
    filename = f"供应商材料核票_{normalized_month}.xlsx"
    return StreamingResponse(
        iter((stream.getvalue(),)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''" + quote(filename)},
    )
