from __future__ import annotations

from copy import copy
from datetime import date
from io import BytesIO
from typing import Literal
from urllib.parse import quote

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Response,
    UploadFile,
    status,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session, selectinload

from app.api.deps import get_db
from app.api.stocktake import (
    require_stocktake_review,
    require_stocktake_submit,
    require_stocktake_view,
)
from app.models.inventory_onboarding import InventoryOnboardingBatch
from app.models.user import User
from app.models.warehouse_inventory import (
    FinishedGoodsInventoryDetail,
    InventoryLot,
    InventoryPalletItem,
    WarehouseLocation,
)
from app.services import inventory_onboarding as onboarding_service
from app.services import (
    inventory_onboarding_posting as posting_service,
)
from app.services.inventory_onboarding_uploads import (
    cleanup_inventory_onboarding_upload,
    store_and_parse_inventory_onboarding_upload,
)
from app.services.secure_uploads import UploadValidationError


router = APIRouter()

TEMPLATE_HEADERS = (
    "盘点日期",
    "盘点人",
    "库存类型",
    "归属类型",
    "楼层",
    "区域",
    "库位编码",
    "栈板号",
    "客户编码",
    "客户名称",
    "存货编码",
    "产品名称",
    "材质编码",
    "数量",
    "单位",
    "入库日期",
    "日期可信度",
    "入库日期原文",
    "供应商",
    "层数",
    "楞型",
    "纸板长",
    "纸板宽",
    "片料类型",
    "组件类型",
    "每箱片数",
    "每张产出",
    "压线类型",
    "左压线",
    "中压线",
    "右压线",
    "开料备注",
    "备注",
    "处理决定",
)

FIELD_SHEET_HEADERS = (
    "现场序号",
    "系统位置(仅参考)",
    "大概位置(可空)",
    "客户(大概即可)",
    "产品或存货编码",
    "系统数量(仅参考)",
    "现场数量",
    "备注",
    "库存批次ID",
    "库存版本",
    "可用数量",
    "预占数量",
    "库存类型",
    "归属类型",
    "楼层",
    "区域",
    "栈板号",
    "客户ID",
    "产品ID",
    "单位",
    "入库日期",
    "日期可信度",
    "入库日期原文",
)


class BatchVersionRequest(BaseModel):
    expected_version: int = Field(gt=0)


class SubmitBatchRequest(BatchVersionRequest):
    dry_run_fingerprint: str = Field(min_length=64, max_length=64)
    idempotency_key: str = Field(min_length=1, max_length=120)
    confirmed: bool

    @field_validator("idempotency_key")
    @classmethod
    def strip_idempotency_key(cls, value: str) -> str:
        stripped = value.strip()
        if not stripped:
            raise ValueError("幂等键不能为空")
        return stripped


class UpdateLineRequest(BaseModel):
    expected_version: int = Field(gt=0)
    batch_expected_version: int | None = Field(default=None, gt=0)
    inventory_type: str | None = None
    ownership_type: str | None = None
    location_id: int | None = Field(default=None, gt=0)
    location_code: str | None = Field(default=None, max_length=50)
    pallet_code: str | None = Field(default=None, max_length=100)
    customer_id: int | None = Field(default=None, gt=0)
    customer_code: str | None = Field(default=None, max_length=100)
    customer_name: str | None = Field(default=None, max_length=200)
    product_id: int | None = Field(default=None, gt=0)
    inventory_code: str | None = Field(default=None, max_length=150)
    product_name: str | None = Field(default=None, max_length=250)
    material_id: int | None = Field(default=None, gt=0)
    material_code: str | None = Field(default=None, max_length=100)
    quantity: int | None = Field(default=None, ge=0)
    unit: str | None = None
    stocktake_date: date | None = None
    stocktaker_name: str | None = Field(default=None, max_length=100)
    stock_date: date | None = None
    stock_date_accuracy: str | None = None
    stock_date_original_text: str | None = Field(default=None, max_length=100)
    supplier_name: str | None = Field(default=None, max_length=200)
    layer_count: int | None = None
    flute_type: str | None = Field(default=None, max_length=20)
    board_length_mm: int | None = Field(default=None, gt=0)
    board_width_mm: int | None = Field(default=None, gt=0)
    sheet_type: str | None = None
    component_type: str | None = None
    pieces_per_box: int | None = Field(default=None, gt=0)
    stock_yield_per_sheet: int | None = Field(default=None, gt=0)
    crease_type: str | None = Field(default=None, max_length=20)
    crease_left_mm: int | None = Field(default=None, ge=0)
    crease_middle_mm: int | None = Field(default=None, ge=0)
    crease_right_mm: int | None = Field(default=None, ge=0)
    cutting_note: str | None = Field(default=None, max_length=1000)
    remarks: str | None = Field(default=None, max_length=1000)
    action_decision: str | None = None


def _service_error(error: onboarding_service.InventoryOnboardingError) -> None:
    raise HTTPException(
        status_code=error.status_code,
        detail={"code": error.code, "message": str(error)},
    ) from error


def _concurrency_error(error: IntegrityError) -> None:
    raise HTTPException(
        status_code=409,
        detail={
            "code": "INVENTORY_ONBOARDING_CONCURRENCY_CONFLICT",
            "message": "草稿已被其他操作更新，请刷新后重试",
        },
    ) from error


def _batch_response(
    db: Session,
    batch: InventoryOnboardingBatch,
) -> dict[str, object]:
    lines = onboarding_service._batch_lines(db, batch.id)
    batch_data = onboarding_service.batch_payload(batch)
    batch_data.update(posting_service.posting_state_payload(db, batch))
    return {
        "batch": batch_data,
        "items": [onboarding_service.line_payload(line) for line in lines],
    }


def _csv_template() -> bytes:
    import csv
    from io import StringIO

    stream = StringIO(newline="")
    csv.writer(stream).writerow(TEMPLATE_HEADERS)
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _xlsx_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "库存建账"
    sheet.append(TEMPLATE_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:AH1"
    for cell in sheet[1]:
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = 14
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


def _field_sheet_lots(db: Session) -> list[InventoryLot]:
    return db.scalars(
        select(InventoryLot)
        .join(
            FinishedGoodsInventoryDetail,
            FinishedGoodsInventoryDetail.inventory_lot_id == InventoryLot.id,
        )
        .join(
            WarehouseLocation,
            WarehouseLocation.id == InventoryLot.warehouse_location_id,
        )
        .where(
            InventoryLot.inventory_type == "finished",
            InventoryLot.status.in_(("active", "frozen")),
            WarehouseLocation.is_active.is_(True),
        )
        .options(
            selectinload(InventoryLot.location),
            selectinload(InventoryLot.finished_detail),
            selectinload(InventoryLot.pallet_item).selectinload(
                InventoryPalletItem.pallet
            ),
        )
        .order_by(
            WarehouseLocation.sort_order,
            WarehouseLocation.location_code,
            InventoryLot.id,
        )
    ).all()


def _field_sheet_row(index: int, lot: InventoryLot) -> tuple[object, ...]:
    detail = lot.finished_detail
    pallet_item = lot.pallet_item
    pallet = pallet_item.pallet if pallet_item is not None else None
    system_quantity = int(lot.quantity_available) + int(lot.quantity_reserved)
    return (
        f"P{index:03d}",
        lot.location.location_code,
        None,
        detail.owner_customer_name_snapshot if detail is not None else None,
        (
            " / ".join(
                value
                for value in (
                    detail.inventory_code_snapshot,
                    detail.product_name_snapshot,
                )
                if value
            )
            if detail is not None
            else None
        ),
        system_quantity,
        system_quantity,
        None,
        lot.id,
        lot.version,
        int(lot.quantity_available),
        int(lot.quantity_reserved),
        "成品",
        (
            "通用"
            if detail is not None and detail.is_general
            else "客户专用"
        ),
        lot.location.warehouse_floor,
        lot.location.area_code,
        pallet.pallet_code if pallet is not None else None,
        detail.owner_customer_id if detail is not None else None,
        detail.product_id if detail is not None else None,
        "个",
        lot.stock_date.isoformat(),
        {
            "exact": "精确",
            "estimated": "估算",
            "unknown": "不明",
        }.get(lot.stock_date_accuracy, "不明"),
        lot.stock_date_original_text,
    )


def _inventory_field_sheet(db: Session) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "成品现场盘点"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(FIELD_SHEET_HEADERS)
    lots = _field_sheet_lots(db)
    for index, lot in enumerate(lots, start=1):
        sheet.append(_field_sheet_row(index, lot))

    visible_count = 8
    header_fill = PatternFill("solid", fgColor="1F4E78")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    reference_fill = PatternFill("solid", fgColor="E2F0D9")
    border = Border(
        left=Side(style="thin", color="C9D2DC"),
        right=Side(style="thin", color="C9D2DC"),
        top=Side(style="thin", color="C9D2DC"),
        bottom=Side(style="thin", color="C9D2DC"),
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=12)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    last_formatted_row = max(len(lots) + 31, 31)
    for row_index in range(2, last_formatted_row + 1):
        sheet.row_dimensions[row_index].height = 25
        for column_index in range(1, visible_count + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=column_index in {3, 5, 8},
            )
            cell.font = Font(size=12)
            cell.fill = (
                input_fill
                if column_index in {3, 4, 5, 7, 8}
                else reference_fill
            )

    widths = {
        "A": 10,
        "B": 20,
        "C": 24,
        "D": 24,
        "E": 38,
        "F": 12,
        "G": 18,
        "H": 30,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for column_index in range(visible_count + 1, len(FIELD_SHEET_HEADERS) + 1):
        sheet.column_dimensions[
            sheet.cell(row=1, column=column_index).column_letter
        ].hidden = True

    sheet.auto_filter.ref = f"A1:H{max(len(lots) + 1, 2)}"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f"A1:H{max(len(lots) + 1, 31)}"
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


@router.get("/inventory-onboarding/template")
def download_inventory_onboarding_template(
    format: Literal["csv", "xlsx"] = "xlsx",
    _user: User = Depends(require_stocktake_view),
) -> Response:
    if format == "csv":
        content = _csv_template()
        media_type = "text/csv; charset=utf-8"
        filename = "N081_inventory_onboarding_template.csv"
    else:
        content = _xlsx_template()
        media_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        filename = "N081_inventory_onboarding_template.xlsx"
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/inventory-onboarding/field-sheet")
def download_inventory_field_sheet(
    db: Session = Depends(get_db),
    _user: User = Depends(require_stocktake_view),
) -> Response:
    content = _inventory_field_sheet(db)
    filename = "天明ERP_当前成品库存现场盘点表.xlsx"
    encoded = quote(filename)
    return Response(
        content=content,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; filename="inventory_field_count.xlsx"; '
                f"filename*=UTF-8''{encoded}"
            )
        },
    )


@router.post(
    "/inventory-onboarding/batches/import",
    status_code=status.HTTP_201_CREATED,
)
async def import_inventory_onboarding_batch(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: User = Depends(require_stocktake_submit),
) -> dict[str, object]:
    stored = None
    created: bool | None = None
    committed = False
    try:
        stored = await store_and_parse_inventory_onboarding_upload(file)
        batch, created = onboarding_service.create_onboarding_draft(
            db,
            upload=stored,
            creator=user,
        )
        response = _batch_response(db, batch)
        response["created"] = created
        db.commit()
        committed = True
        if not created:
            cleanup_inventory_onboarding_upload(stored)
        return response
    except UploadValidationError as error:
        db.rollback()
        if stored is not None:
            cleanup_inventory_onboarding_upload(stored)
        raise HTTPException(status_code=400, detail=str(error)) from error
    except onboarding_service.InventoryOnboardingError as error:
        db.rollback()
        if stored is not None:
            cleanup_inventory_onboarding_upload(stored)
        _service_error(error)
    except IntegrityError as error:
        db.rollback()
        if stored is None:
            raise
        existing = db.scalar(
            select(InventoryOnboardingBatch).where(
                InventoryOnboardingBatch.source_file_sha256 == stored.sha256
            )
        )
        cleanup_inventory_onboarding_upload(stored)
        if existing is None:
            raise HTTPException(
                status_code=409,
                detail={
                    "code": "INVENTORY_ONBOARDING_IMPORT_CONFLICT",
                    "message": "盘点文件导入发生冲突，请刷新后重试",
                },
            ) from error
        response = _batch_response(db, existing)
        response["created"] = False
        return response
    except Exception:
        db.rollback()
        if stored is not None and (not committed or created is False):
            cleanup_inventory_onboarding_upload(stored)
        raise


@router.get("/inventory-onboarding/batches")
def get_inventory_onboarding_batches(
    db: Session = Depends(get_db),
    _user: User = Depends(require_stocktake_view),
) -> dict[str, object]:
    return {
        "items": [
            {
                **onboarding_service.batch_payload(batch),
                **posting_service.posting_state_payload(db, batch),
            }
            for batch in onboarding_service.list_onboarding_batches(db)
        ]
    }


@router.get("/inventory-onboarding/batches/{batch_id}")
def get_inventory_onboarding_batch(
    batch_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(require_stocktake_view),
) -> dict[str, object]:
    try:
        batch = onboarding_service.get_onboarding_batch(db, batch_id)
        return _batch_response(db, batch)
    except onboarding_service.InventoryOnboardingError as error:
        _service_error(error)


@router.get("/inventory-onboarding/batches/{batch_id}/lines")
def get_inventory_onboarding_lines(
    batch_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(require_stocktake_view),
) -> dict[str, object]:
    try:
        onboarding_service.get_onboarding_batch(db, batch_id)
        return {
            "items": [
                onboarding_service.line_payload(line)
                for line in onboarding_service._batch_lines(db, batch_id)
            ]
        }
    except onboarding_service.InventoryOnboardingError as error:
        _service_error(error)


@router.patch(
    "/inventory-onboarding/batches/{batch_id}/lines/{line_id}"
)
def patch_inventory_onboarding_line(
    batch_id: int,
    line_id: int,
    payload: UpdateLineRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_stocktake_submit),
) -> dict[str, object]:
    excluded = {"expected_version", "batch_expected_version"}
    values = payload.model_dump(exclude=excluded, exclude_unset=True)
    try:
        line = onboarding_service.update_onboarding_line(
            db,
            batch_id=batch_id,
            line_id=line_id,
            expected_version=payload.expected_version,
            batch_expected_version=payload.batch_expected_version,
            values=values,
            operator=user,
        )
        db.commit()
        batch = onboarding_service.get_onboarding_batch(db, batch_id)
        return {
            "batch": onboarding_service.batch_payload(batch),
            "line": onboarding_service.line_payload(line),
        }
    except onboarding_service.InventoryOnboardingError as error:
        db.rollback()
        _service_error(error)
    except IntegrityError as error:
        db.rollback()
        _concurrency_error(error)


@router.post("/inventory-onboarding/batches/{batch_id}/rematch")
def rematch_inventory_onboarding_batch(
    batch_id: int,
    payload: BatchVersionRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_stocktake_submit),
) -> dict[str, object]:
    try:
        onboarding_service.rematch_onboarding_batch(
            db,
            batch_id=batch_id,
            expected_version=payload.expected_version,
            operator=user,
        )
        db.commit()
        return _batch_response(
            db,
            onboarding_service.get_onboarding_batch(db, batch_id),
        )
    except onboarding_service.InventoryOnboardingError as error:
        db.rollback()
        _service_error(error)
    except IntegrityError as error:
        db.rollback()
        _concurrency_error(error)


@router.post("/inventory-onboarding/batches/{batch_id}/dry-run")
def dry_run_inventory_onboarding_batch(
    batch_id: int,
    payload: BatchVersionRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_stocktake_submit),
) -> dict[str, object]:
    try:
        onboarding_service.dry_run_onboarding_batch(
            db,
            batch_id=batch_id,
            expected_version=payload.expected_version,
            operator=user,
        )
        db.commit()
        return _batch_response(
            db,
            onboarding_service.get_onboarding_batch(db, batch_id),
        )
    except onboarding_service.InventoryOnboardingError as error:
        db.rollback()
        _service_error(error)
    except IntegrityError as error:
        db.rollback()
        _concurrency_error(error)


@router.post("/inventory-onboarding/batches/{batch_id}/submit")
def submit_inventory_onboarding_batch(
    batch_id: int,
    payload: SubmitBatchRequest,
    db: Session = Depends(get_db),
    user: User = Depends(require_stocktake_submit),
) -> dict[str, object]:
    try:
        onboarding_service.submit_onboarding_batch(
            db,
            batch_id=batch_id,
            expected_version=payload.expected_version,
            dry_run_fingerprint=payload.dry_run_fingerprint,
            idempotency_key=payload.idempotency_key,
            confirmed=payload.confirmed,
            operator=user,
        )
        db.commit()
        return _batch_response(
            db,
            onboarding_service.get_onboarding_batch(db, batch_id),
        )
    except onboarding_service.InventoryOnboardingError as error:
        db.rollback()
        _service_error(error)
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail={
                "code": "INVENTORY_ONBOARDING_SUBMIT_CONFLICT",
                "message": "提交发生并发冲突，请刷新后重试",
            },
        ) from error


@router.post("/inventory-onboarding/batches/{batch_id}/post")
def post_inventory_onboarding_batch(
    batch_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(require_stocktake_review),
) -> dict[str, object]:
    try:
        posting_service.post_submitted_batch(
            db,
            batch_id=batch_id,
            operator=user,
        )
        db.commit()
        return _batch_response(
            db,
            onboarding_service.get_onboarding_batch(db, batch_id),
        )
    except posting_service.InventoryOnboardingPostingError as error:
        db.rollback()
        raise HTTPException(
            status_code=error.status_code,
            detail={"code": error.code, "message": str(error)},
        ) from error
    except onboarding_service.InventoryOnboardingError as error:
        db.rollback()
        _service_error(error)
    except (IntegrityError, OperationalError) as error:
        db.rollback()
        replay = posting_service.get_posting_for_batch(db, batch_id)
        if replay is not None:
            return _batch_response(
                db,
                onboarding_service.get_onboarding_batch(db, batch_id),
            )
        raise HTTPException(
            status_code=409,
            detail={
                "code": "INVENTORY_ONBOARDING_POSTING_CONFLICT",
                "message": "入账发生并发冲突，请刷新批次后重试",
            },
        ) from error


@router.get("/inventory-onboarding/batches/{batch_id}/errors.csv")
def download_inventory_onboarding_errors(
    batch_id: int,
    db: Session = Depends(get_db),
    _user: User = Depends(require_stocktake_view),
) -> Response:
    try:
        batch = onboarding_service.get_onboarding_batch(db, batch_id)
        content = onboarding_service.error_csv_bytes(
            onboarding_service._batch_lines(db, batch.id)
        )
        return Response(
            content=content,
            media_type="text/csv; charset=utf-8",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="{batch.batch_number}_errors.csv"'
                )
            },
        )
    except onboarding_service.InventoryOnboardingError as error:
        _service_error(error)
