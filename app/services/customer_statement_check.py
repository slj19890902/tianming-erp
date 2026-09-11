"""Memory-only comparison of customer spreadsheets with a scoped ERP statement."""
from collections import defaultdict, Counter
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json, re, unicodedata, zipfile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill

MAX_BYTES=8*1024*1024
FIELDS={"delivery_number":"送货单号","customer_po":"客户单号","product_code":"存货编码","product_name":"产品名称","quantity":"数量","unit_price":"单价","amount":"金额"}
ALIASES={"delivery_number":("送货单号","送货单编号","送货编号","deliveryno"),"customer_po":("客户单号","客户订单号","采购订单号","订单号","pono"),"product_code":("存货编码","存货编号","物料编码","物料号","客户料号","料号","产品编码","产品编号"),"product_name":("产品名称","品名","物料名称","名称"),"quantity":("数量","收货数量","实收数量","送货数量","对账数量"),"unit_price":("单价","含税单价","未税单价","不含税单价"),"amount":("金额","含税金额","未税金额","不含税金额","总金额","价税合计")}


def normalize(value):
    if value is None:return ""
    if isinstance(value,float) and value.is_integer(): value=int(value)
    return re.sub(r"\s+","",unicodedata.normalize("NFKC",str(value))).strip().casefold()


def number(value):
    if value is None or str(value).strip()=="": raise ValueError("数值为空")
    try:
        result=Decimal(re.sub(r"[,，￥¥元\s]","",str(value)))
    except InvalidOperation as exc: raise ValueError("数值格式无效") from exc
    if not result.is_finite() or abs(result)>Decimal('1000000000000'): raise ValueError("数值超出范围")
    return result


def load_customer(content, *, sheet_name="", header_row=0, columns=None):
    if not content or len(content)>MAX_BYTES: raise ValueError("请上传不超过8MB的xlsx文件")
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist())>2000 or sum(i.file_size for i in archive.infolist())>32*1024*1024: raise ValueError("表格解压内容过大，请只保留本期对账页")
        book=load_workbook(BytesIO(content),read_only=True,data_only=True,keep_links=False)
    except (zipfile.BadZipFile,KeyError,OSError,ValueError,ParseError,InvalidFileException) as exc: raise ValueError("文件不是有效的xlsx工作簿") from exc
    try:
        if sheet_name and sheet_name not in book.sheetnames: raise ValueError("工作表已变化，请重新选择")
        sheet=book[sheet_name] if sheet_name else book.active
        if sheet.max_row>10050 or sheet.max_column>100: raise ValueError("单页最多10000条、100列，请拆分文件")
        cells=list(sheet.iter_rows())
        raw=[[cell.value for cell in row] for row in cells]
        def identify(row):return {key:i for i,value in enumerate(row) for key,names in ALIASES.items() if normalize(value) in names}
        if not header_row:
            header_row=max(range(1,min(len(raw),30)+1),key=lambda n:len(identify(raw[n-1])),default=1)
        if not 1<=header_row<=min(len(raw),100): raise ValueError("表头行不正确")
        mapped=columns if columns else identify(raw[header_row-1])
        if any(k not in FIELDS or not isinstance(v,int) or isinstance(v,bool) or v<0 or v>=len(raw[header_row-1]) for k,v in mapped.items()):raise ValueError("列对应设置无效")
        if len(set(mapped.values()))!=len(mapped):raise ValueError("同一列不能对应多个字段")
        meta={"sheets":book.sheetnames,"sheet_name":sheet.title,"header_row":header_row,"headers":[str(v or f"第{i+1}列") for i,v in enumerate(raw[header_row-1])],"columns":mapped,"fields":FIELDS}
        if not {"quantity","amount"}<=mapped.keys() or not ({"product_code","product_name"}&mapped.keys()):return {**meta,"needs_mapping":True,"rows":[],"errors":[]}
        rows=[];errors=[]
        for n,rawrow in enumerate(raw[header_row:],header_row+1):
            if not any(v is not None and str(v).strip() for v in rawrow):continue
            values={k:rawrow[i] if i<len(rawrow) else None for k,i in mapped.items()}
            if any(normalize(v) in ("合计","总计","本页合计","累计","小计") for v in values.values()):continue
            identity=values.get('product_code') or values.get('product_name')
            if not identity:
                errors.append({"row":n,"message":"缺少存货编码或产品名称"});continue
            try:
                item={k:str(v or '').strip() for k,v in values.items() if k not in ('quantity','amount','unit_price')}
                # Numeric codes with an explicit all-zero Excel display format retain their leading zeros.
                for key in ('product_code','customer_po','delivery_number'):
                    if key in mapped and isinstance(values.get(key),(int,float)):
                        fmt=cells[n-1][mapped[key]].number_format
                        item[key]=str(int(values[key])).zfill(len(fmt)) if re.fullmatch(r'0+',fmt) else str(values[key]).removesuffix('.0')
                item.update(quantity=number(values.get('quantity')),amount=number(values.get('amount')),source_row=n)
                item['unit_price']=number(values['unit_price']) if values.get('unit_price') is not None and str(values['unit_price']).strip() else None
                rows.append(item)
            except ValueError as exc:errors.append({"row":n,"message":str(exc)})
        return {**meta,"needs_mapping":False,"rows":rows,"errors":errors}
    except (ParseError, zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("工作表内容损坏，请重新另存为xlsx后上传") from exc
    finally:book.close()


def compare(parsed, erp_rows):
    identity='product_code' if 'product_code' in parsed['columns'] else 'product_name'
    keys=[field for field in ('delivery_number','customer_po') if field in parsed['columns']]+[identity]
    def grouped(rows):
        groups=defaultdict(list)
        for row in rows:
            groups[tuple(normalize(row.get(k) or ("name:"+str(row.get('product_name') or '')) if k==identity else row.get(k)) for k in keys)].append(row)
        return groups
    left,right=grouped(parsed['rows']),grouped(erp_rows);results=[]
    for key in sorted(set(left)|set(right)):
        customer=left.get(key,[]);erp=right.get(key,[]); sample=(customer or erp)[0]
        cq=sum((number(r['quantity']) for r in customer),Decimal(0));eq=sum((number(r['quantity']) for r in erp),Decimal(0))
        ca=sum((number(r['amount']) for r in customer),Decimal(0));ea=sum((number(r['amount']) for r in erp),Decimal(0))
        status='客户多出' if not erp else '客户缺少' if not customer else '一致' if cq==eq and abs(ca-ea)<=Decimal('.01') else '数量/金额差异'
        cp={number(r['unit_price']) for r in customer if r.get('unit_price') is not None};ep={number(r['unit_price']) for r in erp if r.get('unit_price') is not None}
        if status=='一致' and cp and cp!=ep:status='单价差异'
        results.append({"status":status,**{k:sample.get(k) or '' for k in ('delivery_number','customer_po','product_code','product_name')},'customer_quantity':str(cq),'erp_quantity':str(eq),'quantity_difference':str(cq-eq),'customer_amount':str(ca),'erp_amount':str(ea),'amount_difference':str(ca-ea),'customer_prices':','.join(map(str,sorted(cp))),'erp_prices':','.join(map(str,sorted(ep))),'customer_rows':[r['source_row'] for r in customer],'statement_item_ids':[r['statement_item_id'] for r in erp]})
    return {'matching_fields':[FIELDS[k] for k in keys],'amount_tolerance':'0.01','quantity_tolerance':'0','summary':dict(Counter(r['status'] for r in results)),'items':results,'errors':parsed['errors'],'complete':not parsed['errors'],'read_only':True,'needs_mapping':False,**{k:parsed[k] for k in ('sheets','sheet_name','header_row','headers','columns','fields')}}


def export_report(result):
    book=Workbook();sheet=book.active;sheet.title='对账差异'
    sheet.append(['客户',result['customer_name'],'月份',result['statement_month'],'ERP对账单',result['statement_number']])
    sheet.append(['匹配字段','、'.join(result['matching_fields']),'金额容差','0.01元','数量容差','0'])
    sheet.append(['核对状态','完整' if result['complete'] else '有无法解析的行，结果不完整'])
    headers=['结果','送货单号','客户单号','存货编码','产品名称','客户数量','ERP数量','数量差','客户金额','ERP金额','金额差','客户单价','ERP单价','客户文件行','ERP明细ID'];sheet.append(headers)
    for row in result['items']:
        values=[row[k] for k in ('status','delivery_number','customer_po','product_code','product_name')]+[Decimal(row[k]) for k in ('customer_quantity','erp_quantity','quantity_difference','customer_amount','erp_amount','amount_difference')]+[row['customer_prices'],row['erp_prices'],','.join(map(str,row['customer_rows'])),','.join(map(str,row['statement_item_ids']))]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
            if row['status']!='一致':cell.fill=PatternFill('solid',fgColor='FFF1E6')
    for cell in sheet[4]:cell.font=Font(bold=True)
    for column in 'ABCDEFGHIJKLMNO':sheet.column_dimensions[column].width=19
    sheet.column_dimensions['E'].width=32;sheet.freeze_panes='A5';sheet.auto_filter.ref=f'A4:O{sheet.max_row}'
    errors=book.create_sheet('未识别行');errors.append(['原文件行号','原因'])
    for row in result['errors']:errors.append([row['row'],row['message']])
    for tab in book:
        for row in tab:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
    stream=BytesIO();book.save(stream);stream.seek(0);return stream
