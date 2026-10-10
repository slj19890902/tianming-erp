from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO, Iterable, Mapping, Sequence


DEFAULT_CUSTOMER_SHEET_NAME = "采购入库打印"
MAX_XLSX_SIZE_BYTES = 20 * 1024 * 1024
MAX_WORKSHEET_ROWS = 50_000
MAX_WORKSHEET_COLUMNS = 200
MAX_HEADER_SCAN_ROWS = 50

QUANTITY_TOLERANCE = Decimal("0")
UNIT_PRICE_TOLERANCE = Decimal("0.01")
AMOUNT_TOLERANCE = Decimal("0.10")

MATCHED = "完全匹配"
CUSTOMER_ONLY = "客户有，ERP没有"
ERP_ONLY = "ERP有，客户没有"
QUANTITY_MISMATCH = "数量不一致"
UNIT_PRICE_MISMATCH = "单价不一致"
AMOUNT_MISMATCH = "金额不一致"
PRODUCT_NAME_MISMATCH = "产品名称不一致"
SPECIFICATION_MISMATCH = "规格型号不一致"
UNIT_MISMATCH = "单位不一致"
CUSTOMER_PO_MISMATCH = "客户订单号不一致"
CUSTOMER_DUPLICATE = "客户文件重复行"
ERP_DUPLICATE = "ERP对账单重复行"
MANUAL_REVIEW = "待人工确认"
AGGREGATED_NOTE = "同一客户订单号 + 存货编码存在多行，已合并核对"


class CustomerStatementReconciliationError(ValueError):
    pass


HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "sequence": ("序号", "行号"),
    "receipt_date": ("入库日期", "收货日期", "实收日期"),
    "receipt_number": ("入库单号", "入库单编号", "收货单号"),
    "supplier_code": ("供货商编码", "供应商编码"),
    "supplier_name": ("供货商名称", "供应商名称"),
    "product_code": ("存货编码", "物料编码", "产品编码"),
    "product_name": ("存货名称", "物料名称", "产品名称"),
    "specification": ("规格型号", "规格", "型号"),
    "unit": ("单位", "计量单位"),
    "quantity": ("实收数量", "入库数量", "签收数量", "数量"),
    "customer_po": (
        "采购订单号",
        "采购单号",
        "客户订单号",
        "客户单号",
        "客户PO",
        "PO号",
    ),
    "unit_price": ("原币含税单价", "含税单价", "客户含税单价", "单价"),
    "net_amount": ("原币金额", "未税金额", "不含税金额"),
    "amount": ("原币价税合计", "价税合计", "含税金额", "对账金额", "金额"),
}

HEADER_REQUIRED_FIELDS = frozenset(
    {"customer_po", "product_code", "quantity", "unit_price", "amount"}
)

ERP_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "source_row": ("source_row", "source_row_number", "row_number", "line_number"),
    "reference_id": (
        "reference_id",
        "statement_item_id",
        "erp_statement_line_id",
        "id",
    ),
    "delivery_date": ("delivery_date", "送货日期"),
    "delivery_number": ("delivery_number", "送货单号"),
    "customer_po": ("customer_po", "customer_order_number", "客户单号", "客户订单号"),
    "product_code": ("product_code", "snapshot_product_code", "存货编码"),
    "product_name": ("product_name", "snapshot_product_name", "产品名称"),
    "specification": ("specification", "spec", "snapshot_spec", "规格型号"),
    "unit": ("unit", "snapshot_unit", "单位"),
    "quantity": (
        "quantity",
        "actual_received_quantity",
        "reconciled_quantity",
        "实收数量",
    ),
    "unit_price": ("unit_price", "unit_price_snapshot", "含税单价", "单价"),
    "amount": ("amount", "receivable_amount", "reconciliation_amount", "含税金额"),
}


@dataclass(frozen=True, slots=True)
class ReconciliationTolerance:
    quantity: Decimal = QUANTITY_TOLERANCE
    unit_price: Decimal = UNIT_PRICE_TOLERANCE
    amount: Decimal = AMOUNT_TOLERANCE

    def __post_init__(self) -> None:
        if self.quantity < 0 or self.unit_price < 0 or self.amount < 0:
            raise CustomerStatementReconciliationError("核对容差不能为负数")


@dataclass(frozen=True, slots=True)
class StatementRow:
    source_kind: str
    source_row: int
    reference_id: str | None = None
    sequence: str | None = None
    receipt_date: date | None = None
    receipt_number: str | None = None
    delivery_date: date | None = None
    delivery_number: str | None = None
    supplier_code: str | None = None
    supplier_name: str | None = None
    customer_po: str | None = None
    product_code: str | None = None
    product_name: str | None = None
    specification: str | None = None
    unit: str | None = None
    quantity: Decimal | None = None
    unit_price: Decimal | None = None
    net_amount: Decimal | None = None
    amount: Decimal | None = None
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class ParsedCustomerStatement:
    sheet_name: str
    header_row: int
    field_mapping: Mapping[str, int]
    rows: tuple[StatementRow, ...]
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class StatementGroup:
    source_kind: str
    rows: tuple[StatementRow, ...]
    customer_po: str | None
    product_code: str | None
    product_name: str | None
    specification: str | None
    unit: str | None
    quantity: Decimal | None
    unit_price: Decimal | None
    amount: Decimal | None
    receipt_date: date | None
    delivery_date: date | None
    customer_po_keys: frozenset[str]
    product_name_keys: frozenset[str]
    specification_keys: frozenset[str]
    unit_keys: frozenset[str]

    @property
    def source_rows(self) -> tuple[int, ...]:
        return tuple(row.source_row for row in self.rows)

    @property
    def references(self) -> tuple[str, ...]:
        return tuple(
            row.reference_id or str(row.source_row)
            for row in self.rows
        )

    @property
    def is_aggregated(self) -> bool:
        return len(self.rows) > 1


@dataclass(frozen=True, slots=True)
class ReconciliationResult:
    difference_type: str
    match_status: str
    match_level: str | None
    difference_types: tuple[str, ...]
    customer_group: StatementGroup | None = None
    erp_group: StatementGroup | None = None
    quantity_difference: Decimal | None = None
    unit_price_difference: Decimal | None = None
    amount_difference: Decimal | None = None
    handling_suggestion: str = ""
    notes: tuple[str, ...] = ()
    is_match: bool = False
    requires_manual_review: bool = False


@dataclass(frozen=True, slots=True)
class ReconciliationMetadata:
    customer_name: str
    statement_month: str
    customer_file_name: str
    erp_statement_range: str
    operator: str
    generated_at: datetime = field(default_factory=datetime.now)

    def __post_init__(self) -> None:
        if not re.fullmatch(r"\d{4}-\d{2}", self.statement_month):
            raise CustomerStatementReconciliationError("对账月份格式必须为 YYYY-MM")


@dataclass(frozen=True, slots=True)
class ReconciliationReport:
    customer_statement: ParsedCustomerStatement
    erp_rows: tuple[StatementRow, ...]
    results: tuple[ReconciliationResult, ...]
    metadata: ReconciliationMetadata
    tolerances: ReconciliationTolerance

    @property
    def matched_results(self) -> tuple[ReconciliationResult, ...]:
        return tuple(result for result in self.results if result.is_match)

    @property
    def manual_review_results(self) -> tuple[ReconciliationResult, ...]:
        return tuple(result for result in self.results if result.requires_manual_review)

    @property
    def difference_results(self) -> tuple[ReconciliationResult, ...]:
        return tuple(
            result
            for result in self.results
            if not result.is_match and not result.requires_manual_review
        )


def _normalize_header(value: Any) -> str:
    text = _clean_text(value) or ""
    return re.sub(r"[\s:：()（）/\\_-]+", "", text).upper()


def _clean_text(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        text = format(value, "f")
    elif isinstance(value, int):
        text = str(value)
    elif isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value)
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


HEADER_LOOKUP: dict[str, str] = {
    _normalize_header(alias): field_name
    for field_name, aliases in HEADER_ALIASES.items()
    for alias in aliases
}


def _identifier_key(value: Any) -> str:
    text = (_clean_text(value) or "").upper()
    return re.sub(r"\s+", "", text)


def _descriptive_key(value: Any) -> str:
    text = (_clean_text(value) or "").upper()
    text = text.replace("×", "*").replace("X", "*")
    return re.sub(r"[\s\-_/()（）*.,，。:：]+", "", text)


def _to_decimal(value: Any) -> Decimal | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = _clean_text(value)
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = re.sub(r"[,，￥¥$\s]", "", text)
    try:
        numeric = Decimal(text)
    except InvalidOperation as error:
        raise CustomerStatementReconciliationError(f"无法解析数字：{value}") from error
    return -numeric if negative else numeric


def _to_date(value: Any, *, epoch: datetime | None = None) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)) and epoch is not None:
        numeric = float(value)
        if 20_000 <= numeric <= 90_000:
            from openpyxl.utils.datetime import from_excel

            converted = from_excel(numeric, epoch)
            return converted.date() if isinstance(converted, datetime) else converted
    text = _clean_text(value)
    if not text:
        return None
    for pattern in (
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%Y.%m.%d",
        "%Y%m%d",
        "%Y-%m-%d %H:%M:%S",
    ):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise CustomerStatementReconciliationError(f"无法解析日期：{value}")


def _first_value(row: Mapping[str, Any], aliases: Sequence[str]) -> Any:
    for alias in aliases:
        if alias in row:
            return row[alias]
    return None


def _sheet_source(source: bytes | bytearray | Path | str | BinaryIO) -> BytesIO | Path | str:
    if isinstance(source, (bytes, bytearray)):
        if len(source) > MAX_XLSX_SIZE_BYTES:
            raise CustomerStatementReconciliationError("客户对账文件超过 20MB 限制")
        return BytesIO(bytes(source))
    if isinstance(source, (str, Path)):
        path = Path(source)
        if not path.is_file():
            raise CustomerStatementReconciliationError(f"客户对账文件不存在：{path}")
        if path.stat().st_size > MAX_XLSX_SIZE_BYTES:
            raise CustomerStatementReconciliationError("客户对账文件超过 20MB 限制")
        return path
    payload = source.read(MAX_XLSX_SIZE_BYTES + 1)
    if len(payload) > MAX_XLSX_SIZE_BYTES:
        raise CustomerStatementReconciliationError("客户对账文件超过 20MB 限制")
    return BytesIO(payload)


def _detect_header_mapping(
    buffered_rows: Sequence[Sequence[Any]],
) -> tuple[int, dict[str, int]]:
    best: tuple[int, int, dict[str, int]] | None = None
    for row_number, values in enumerate(buffered_rows, 1):
        mapping: dict[str, int] = {}
        for column_number, value in enumerate(values, 1):
            field_name = HEADER_LOOKUP.get(_normalize_header(value))
            if field_name and field_name not in mapping:
                mapping[field_name] = column_number
        required_count = len(HEADER_REQUIRED_FIELDS.intersection(mapping))
        score = required_count * 10 + len(mapping)
        if "product_code" not in mapping or "quantity" not in mapping:
            continue
        if required_count < 3 or len(mapping) < 5:
            continue
        candidate = (score, -row_number, mapping)
        if best is None or candidate[:2] > best[:2]:
            best = candidate
    if best is None:
        raise CustomerStatementReconciliationError(
            "未识别到客户对账表头；至少需要存货编码、数量及订单/价格/金额字段"
        )
    return -best[1], best[2]


def _mapped_value(values: Sequence[Any], mapping: Mapping[str, int], name: str) -> Any:
    column_number = mapping.get(name)
    if not column_number or column_number > len(values):
        return None
    return values[column_number - 1]


def _parse_decimal_field(
    value: Any,
    *,
    label: str,
    warnings: list[str],
) -> Decimal | None:
    try:
        return _to_decimal(value)
    except CustomerStatementReconciliationError:
        warnings.append(f"{label}无法解析：{_clean_text(value) or ''}")
        return None


def _parse_date_field(
    value: Any,
    *,
    label: str,
    epoch: datetime,
    warnings: list[str],
) -> date | None:
    try:
        return _to_date(value, epoch=epoch)
    except CustomerStatementReconciliationError:
        warnings.append(f"{label}无法解析：{_clean_text(value) or ''}")
        return None


def _looks_like_data_row(values: Sequence[Any], mapping: Mapping[str, int]) -> bool:
    identity = (
        _clean_text(_mapped_value(values, mapping, "customer_po")),
        _clean_text(_mapped_value(values, mapping, "product_code")),
        _clean_text(_mapped_value(values, mapping, "product_name")),
        _clean_text(_mapped_value(values, mapping, "receipt_number")),
    )
    return sum(bool(value) for value in identity) >= 2 and bool(identity[1] or identity[2])


def _parse_customer_row(
    values: Sequence[Any],
    *,
    source_row: int,
    mapping: Mapping[str, int],
    epoch: datetime,
) -> StatementRow | None:
    if not _looks_like_data_row(values, mapping):
        return None
    warnings: list[str] = []
    quantity = _parse_decimal_field(
        _mapped_value(values, mapping, "quantity"),
        label="实收数量",
        warnings=warnings,
    )
    unit_price = _parse_decimal_field(
        _mapped_value(values, mapping, "unit_price"),
        label="含税单价",
        warnings=warnings,
    )
    net_amount = _parse_decimal_field(
        _mapped_value(values, mapping, "net_amount"),
        label="原币金额",
        warnings=warnings,
    )
    amount = _parse_decimal_field(
        _mapped_value(values, mapping, "amount"),
        label="价税合计",
        warnings=warnings,
    )
    receipt_date = _parse_date_field(
        _mapped_value(values, mapping, "receipt_date"),
        label="入库日期",
        epoch=epoch,
        warnings=warnings,
    )
    if quantity is None:
        warnings.append("实收数量为空")
    if unit_price is None:
        warnings.append("含税单价为空")
    if amount is None:
        warnings.append("价税合计为空")
    return StatementRow(
        source_kind="customer",
        source_row=source_row,
        sequence=_clean_text(_mapped_value(values, mapping, "sequence")),
        receipt_date=receipt_date,
        receipt_number=_clean_text(_mapped_value(values, mapping, "receipt_number")),
        supplier_code=_clean_text(_mapped_value(values, mapping, "supplier_code")),
        supplier_name=_clean_text(_mapped_value(values, mapping, "supplier_name")),
        customer_po=_clean_text(_mapped_value(values, mapping, "customer_po")),
        product_code=_clean_text(_mapped_value(values, mapping, "product_code")),
        product_name=_clean_text(_mapped_value(values, mapping, "product_name")),
        specification=_clean_text(_mapped_value(values, mapping, "specification")),
        unit=_clean_text(_mapped_value(values, mapping, "unit")),
        quantity=quantity,
        unit_price=unit_price,
        net_amount=net_amount,
        amount=amount,
        warnings=tuple(dict.fromkeys(warnings)),
    )


def parse_customer_statement_xlsx(
    source: bytes | bytearray | Path | str | BinaryIO,
    *,
    sheet_name: str = DEFAULT_CUSTOMER_SHEET_NAME,
) -> ParsedCustomerStatement:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    workbook_source = _sheet_source(source)
    try:
        workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError) as error:
        raise CustomerStatementReconciliationError(
            "客户对账文件不是可读取的 xlsx 工作簿"
        ) from error
    try:
        normalized_target = _normalize_header(sheet_name)
        matching_names = [
            name
            for name in workbook.sheetnames
            if _normalize_header(name) == normalized_target
        ]
        if not matching_names:
            raise CustomerStatementReconciliationError(
                f"客户对账文件缺少工作表：{sheet_name}"
            )
        worksheet = workbook[matching_names[0]]
        if worksheet.max_row > MAX_WORKSHEET_ROWS:
            raise CustomerStatementReconciliationError(
                f"客户对账工作表超过 {MAX_WORKSHEET_ROWS} 行限制"
            )
        if worksheet.max_column > MAX_WORKSHEET_COLUMNS:
            raise CustomerStatementReconciliationError(
                f"客户对账工作表超过 {MAX_WORKSHEET_COLUMNS} 列限制"
            )

        row_iterator = worksheet.iter_rows(values_only=True)
        buffered_rows: list[tuple[Any, ...]] = []
        for _ in range(MAX_HEADER_SCAN_ROWS):
            try:
                buffered_rows.append(tuple(next(row_iterator)))
            except StopIteration:
                break
        header_row, mapping = _detect_header_mapping(buffered_rows)
        parsed_rows: list[StatementRow] = []
        for row_number, values in enumerate(buffered_rows[header_row:], header_row + 1):
            parsed = _parse_customer_row(
                values,
                source_row=row_number,
                mapping=mapping,
                epoch=workbook.epoch,
            )
            if parsed:
                parsed_rows.append(parsed)
        for row_number, values in enumerate(row_iterator, len(buffered_rows) + 1):
            parsed = _parse_customer_row(
                tuple(values),
                source_row=row_number,
                mapping=mapping,
                epoch=workbook.epoch,
            )
            if parsed:
                parsed_rows.append(parsed)
        if not parsed_rows:
            raise CustomerStatementReconciliationError("客户对账文件没有可核对的明细行")
        warnings = tuple(
            f"第 {row.source_row} 行：{warning}"
            for row in parsed_rows
            for warning in row.warnings
        )
        return ParsedCustomerStatement(
            sheet_name=worksheet.title,
            header_row=header_row,
            field_mapping=dict(mapping),
            rows=tuple(parsed_rows),
            warnings=warnings,
        )
    finally:
        workbook.close()


def normalize_erp_statement_rows(
    rows: Iterable[Mapping[str, Any] | StatementRow],
) -> tuple[StatementRow, ...]:
    normalized: list[StatementRow] = []
    for index, source in enumerate(rows, 1):
        if isinstance(source, StatementRow):
            if source.source_kind != "erp":
                source = StatementRow(
                    **{
                        field_name: getattr(source, field_name)
                        for field_name in StatementRow.__dataclass_fields__
                        if field_name != "source_kind"
                    },
                    source_kind="erp",
                )
            normalized.append(source)
            continue
        warnings: list[str] = []
        quantity = _parse_decimal_field(
            _first_value(source, ERP_FIELD_ALIASES["quantity"]),
            label="ERP实收数量",
            warnings=warnings,
        )
        unit_price = _parse_decimal_field(
            _first_value(source, ERP_FIELD_ALIASES["unit_price"]),
            label="ERP含税单价",
            warnings=warnings,
        )
        amount = _parse_decimal_field(
            _first_value(source, ERP_FIELD_ALIASES["amount"]),
            label="ERP含税金额",
            warnings=warnings,
        )
        delivery_date_value = _first_value(source, ERP_FIELD_ALIASES["delivery_date"])
        try:
            delivery_date = _to_date(delivery_date_value)
        except CustomerStatementReconciliationError:
            delivery_date = None
            warnings.append(f"ERP送货日期无法解析：{_clean_text(delivery_date_value) or ''}")
        source_row_value = _first_value(source, ERP_FIELD_ALIASES["source_row"])
        try:
            source_row = int(source_row_value) if source_row_value not in (None, "") else index
        except (TypeError, ValueError):
            source_row = index
            warnings.append("ERP行号无法解析，已使用输入顺序")
        normalized.append(
            StatementRow(
                source_kind="erp",
                source_row=source_row,
                reference_id=_clean_text(_first_value(source, ERP_FIELD_ALIASES["reference_id"])),
                delivery_date=delivery_date,
                delivery_number=_clean_text(_first_value(source, ERP_FIELD_ALIASES["delivery_number"])),
                customer_po=_clean_text(_first_value(source, ERP_FIELD_ALIASES["customer_po"])),
                product_code=_clean_text(_first_value(source, ERP_FIELD_ALIASES["product_code"])),
                product_name=_clean_text(_first_value(source, ERP_FIELD_ALIASES["product_name"])),
                specification=_clean_text(_first_value(source, ERP_FIELD_ALIASES["specification"])),
                unit=_clean_text(_first_value(source, ERP_FIELD_ALIASES["unit"])),
                quantity=quantity,
                unit_price=unit_price,
                amount=amount,
                warnings=tuple(dict.fromkeys(warnings)),
            )
        )
    return tuple(normalized)


def _row_fingerprint(row: StatementRow) -> tuple[Any, ...]:
    return (
        _identifier_key(row.customer_po),
        _identifier_key(row.product_code),
        _descriptive_key(row.product_name),
        _descriptive_key(row.specification),
        _descriptive_key(row.unit),
        row.quantity,
        row.unit_price,
        row.amount,
        row.receipt_date or row.delivery_date,
        _identifier_key(row.receipt_number or row.delivery_number),
    )


def _deduplicate_rows(
    rows: Sequence[StatementRow],
) -> tuple[tuple[StatementRow, ...], tuple[tuple[StatementRow, ...], ...]]:
    groups: dict[tuple[Any, ...], list[StatementRow]] = defaultdict(list)
    for row in rows:
        groups[_row_fingerprint(row)].append(row)
    unique_rows: list[StatementRow] = []
    duplicates: list[tuple[StatementRow, ...]] = []
    for grouped_rows in groups.values():
        unique_rows.append(grouped_rows[0])
        if len(grouped_rows) > 1:
            duplicates.append(tuple(grouped_rows))
    unique_rows.sort(key=lambda row: row.source_row)
    duplicates.sort(key=lambda group: group[0].source_row)
    return tuple(unique_rows), tuple(duplicates)


def _distinct_text(rows: Sequence[StatementRow], field_name: str) -> str | None:
    values = [
        _clean_text(getattr(row, field_name))
        for row in rows
        if _clean_text(getattr(row, field_name))
    ]
    distinct = tuple(dict.fromkeys(value for value in values if value))
    return " / ".join(distinct) if distinct else None


def _sum_if_complete(rows: Sequence[StatementRow], field_name: str) -> Decimal | None:
    values = [getattr(row, field_name) for row in rows]
    if any(value is None for value in values):
        return None
    return sum(values, Decimal("0"))


def _group_unit_price(
    rows: Sequence[StatementRow],
    *,
    quantity: Decimal | None,
    amount: Decimal | None,
) -> Decimal | None:
    prices = [row.unit_price for row in rows]
    if any(price is None for price in prices):
        return None
    if len(set(prices)) == 1:
        return prices[0]
    if quantity not in (None, Decimal("0")) and amount is not None:
        return amount / quantity
    return None


def _build_group(rows: Sequence[StatementRow]) -> StatementGroup:
    quantity = _sum_if_complete(rows, "quantity")
    amount = _sum_if_complete(rows, "amount")
    return StatementGroup(
        source_kind=rows[0].source_kind,
        rows=tuple(rows),
        customer_po=_distinct_text(rows, "customer_po"),
        product_code=_distinct_text(rows, "product_code"),
        product_name=_distinct_text(rows, "product_name"),
        specification=_distinct_text(rows, "specification"),
        unit=_distinct_text(rows, "unit"),
        quantity=quantity,
        unit_price=_group_unit_price(rows, quantity=quantity, amount=amount),
        amount=amount,
        receipt_date=min(
            (row.receipt_date for row in rows if row.receipt_date),
            default=None,
        ),
        delivery_date=min(
            (row.delivery_date for row in rows if row.delivery_date),
            default=None,
        ),
        customer_po_keys=frozenset(
            _identifier_key(row.customer_po) for row in rows if _identifier_key(row.customer_po)
        ),
        product_name_keys=frozenset(
            _descriptive_key(row.product_name)
            for row in rows
            if _descriptive_key(row.product_name)
        ),
        specification_keys=frozenset(
            _descriptive_key(row.specification)
            for row in rows
            if _descriptive_key(row.specification)
        ),
        unit_keys=frozenset(
            _descriptive_key(row.unit) for row in rows if _descriptive_key(row.unit)
        ),
    )


def _comparison_groups(rows: Sequence[StatementRow]) -> tuple[StatementGroup, ...]:
    grouped: dict[tuple[str, ...], list[StatementRow]] = defaultdict(list)
    for row in rows:
        customer_po = _identifier_key(row.customer_po)
        product_code = _identifier_key(row.product_code)
        if customer_po and product_code:
            key = ("business", customer_po, product_code)
        else:
            key = ("row", str(row.source_row))
        grouped[key].append(row)
    return tuple(
        _build_group(group)
        for _, group in sorted(grouped.items(), key=lambda item: item[1][0].source_row)
    )


def _business_key(group: StatementGroup) -> tuple[str, str] | None:
    if len(group.customer_po_keys) != 1:
        return None
    product_code = _identifier_key(group.product_code)
    if not product_code:
        return None
    return next(iter(group.customer_po_keys)), product_code


def _difference(
    customer_value: Decimal | None,
    erp_value: Decimal | None,
) -> Decimal | None:
    if customer_value is None or erp_value is None:
        return None
    return customer_value - erp_value


def _sets_mismatch(left: frozenset[str], right: frozenset[str]) -> bool:
    return bool(left and right and left.isdisjoint(right))


def _suggestion(differences: Sequence[str]) -> str:
    if not differences:
        return "无需处理"
    if MANUAL_REVIEW in differences:
        return "请人工选择唯一对应行后重新核对"
    if CUSTOMER_ONLY in differences:
        return "核对客户文件是否属于当前月份，或检查 ERP 回单是否遗漏"
    if ERP_ONLY in differences:
        return "核对 ERP 明细是否属于当前客户文件范围"
    if CUSTOMER_DUPLICATE in differences or ERP_DUPLICATE in differences:
        return "请确认重复行是否为真实分批入库，确认前不要自动覆盖正式数据"
    return "请按差异字段核对客户原单与 ERP 回单，不自动修改正式数据"


def _compare_groups(
    customer: StatementGroup,
    erp: StatementGroup,
    *,
    match_level: str,
    tolerances: ReconciliationTolerance,
    compare_customer_po: bool,
) -> ReconciliationResult:
    quantity_difference = _difference(customer.quantity, erp.quantity)
    unit_price_difference = _difference(customer.unit_price, erp.unit_price)
    amount_difference = _difference(customer.amount, erp.amount)
    differences: list[str] = []
    missing_required = any(
        value is None
        for value in (
            customer.quantity,
            erp.quantity,
            customer.unit_price,
            erp.unit_price,
            customer.amount,
            erp.amount,
        )
    )
    if missing_required:
        differences.append(MANUAL_REVIEW)
    else:
        if abs(quantity_difference or Decimal("0")) > tolerances.quantity:
            differences.append(QUANTITY_MISMATCH)
        if abs(unit_price_difference or Decimal("0")) > tolerances.unit_price:
            differences.append(UNIT_PRICE_MISMATCH)
        if abs(amount_difference or Decimal("0")) > tolerances.amount:
            differences.append(AMOUNT_MISMATCH)
    if _sets_mismatch(customer.product_name_keys, erp.product_name_keys):
        differences.append(PRODUCT_NAME_MISMATCH)
    if _sets_mismatch(customer.specification_keys, erp.specification_keys):
        differences.append(SPECIFICATION_MISMATCH)
    if _sets_mismatch(customer.unit_keys, erp.unit_keys):
        differences.append(UNIT_MISMATCH)
    if compare_customer_po and customer.customer_po_keys != erp.customer_po_keys:
        differences.append(CUSTOMER_PO_MISMATCH)
    differences = list(dict.fromkeys(differences))
    notes: list[str] = []
    if customer.is_aggregated or erp.is_aggregated:
        notes.append(AGGREGATED_NOTE)
    notes.extend(
        f"客户第 {row.source_row} 行：{warning}"
        for row in customer.rows
        for warning in row.warnings
    )
    notes.extend(
        f"ERP第 {row.source_row} 行：{warning}"
        for row in erp.rows
        for warning in row.warnings
    )
    is_match = not differences
    requires_manual_review = MANUAL_REVIEW in differences
    return ReconciliationResult(
        difference_type=MATCHED if is_match else "；".join(differences),
        match_status=MATCHED if is_match else (MANUAL_REVIEW if requires_manual_review else "存在差异"),
        match_level=match_level,
        difference_types=tuple(differences) if differences else (MATCHED,),
        customer_group=customer,
        erp_group=erp,
        quantity_difference=quantity_difference,
        unit_price_difference=unit_price_difference,
        amount_difference=amount_difference,
        handling_suggestion=_suggestion(differences),
        notes=tuple(dict.fromkeys(notes)),
        is_match=is_match,
        requires_manual_review=requires_manual_review,
    )


def _duplicate_result(
    rows: Sequence[StatementRow],
    *,
    difference_type: str,
) -> ReconciliationResult:
    group = _build_group(rows)
    return ReconciliationResult(
        difference_type=difference_type,
        match_status=MANUAL_REVIEW,
        match_level=None,
        difference_types=(difference_type,),
        customer_group=group if group.source_kind == "customer" else None,
        erp_group=group if group.source_kind == "erp" else None,
        handling_suggestion=_suggestion((difference_type,)),
        notes=(f"重复来源行：{', '.join(str(row.source_row) for row in rows)}",),
        requires_manual_review=True,
    )


def _unmatched_result(
    group: StatementGroup,
    *,
    difference_type: str,
) -> ReconciliationResult:
    return ReconciliationResult(
        difference_type=difference_type,
        match_status="存在差异",
        match_level=None,
        difference_types=(difference_type,),
        customer_group=group if group.source_kind == "customer" else None,
        erp_group=group if group.source_kind == "erp" else None,
        handling_suggestion=_suggestion((difference_type,)),
    )


def _auxiliary_candidate(customer: StatementGroup, erp: StatementGroup) -> bool:
    if _identifier_key(customer.product_code) != _identifier_key(erp.product_code):
        return False
    if customer.quantity is None or erp.quantity is None or customer.quantity != erp.quantity:
        return False
    name_matches = bool(customer.product_name_keys.intersection(erp.product_name_keys))
    spec_matches = bool(customer.specification_keys.intersection(erp.specification_keys))
    return name_matches or spec_matches


def reconcile_statement_rows(
    customer_rows: Sequence[StatementRow],
    erp_rows: Iterable[Mapping[str, Any] | StatementRow],
    *,
    tolerances: ReconciliationTolerance | None = None,
) -> tuple[tuple[StatementRow, ...], tuple[ReconciliationResult, ...]]:
    effective_tolerances = tolerances or ReconciliationTolerance()
    normalized_erp = normalize_erp_statement_rows(erp_rows)
    unique_customer, customer_duplicates = _deduplicate_rows(tuple(customer_rows))
    unique_erp, erp_duplicates = _deduplicate_rows(normalized_erp)
    customer_groups = _comparison_groups(unique_customer)
    erp_groups = _comparison_groups(unique_erp)
    results: list[ReconciliationResult] = []
    results.extend(
        _duplicate_result(group, difference_type=CUSTOMER_DUPLICATE)
        for group in customer_duplicates
    )
    results.extend(
        _duplicate_result(group, difference_type=ERP_DUPLICATE)
        for group in erp_duplicates
    )

    erp_by_key = {
        key: group
        for group in erp_groups
        if (key := _business_key(group)) is not None
    }
    matched_erp_ids: set[int] = set()
    unmatched_customers: list[StatementGroup] = []
    for customer in customer_groups:
        key = _business_key(customer)
        erp = erp_by_key.get(key) if key else None
        if erp is None or id(erp) in matched_erp_ids:
            unmatched_customers.append(customer)
            continue
        quantity_matches = (
            customer.quantity is not None
            and erp.quantity is not None
            and customer.quantity == erp.quantity
        )
        price_matches = (
            customer.unit_price is not None
            and erp.unit_price is not None
            and abs(customer.unit_price - erp.unit_price) <= effective_tolerances.unit_price
        )
        results.append(
            _compare_groups(
                customer,
                erp,
                match_level="强匹配" if quantity_matches and price_matches else "次强匹配",
                tolerances=effective_tolerances,
                compare_customer_po=False,
            )
        )
        matched_erp_ids.add(id(erp))

    available_erp = [group for group in erp_groups if id(group) not in matched_erp_ids]
    ambiguous_erp_ids: set[int] = set()
    for customer in unmatched_customers:
        candidates = [
            group
            for group in available_erp
            if id(group) not in matched_erp_ids and _auxiliary_candidate(customer, group)
        ]
        if len(candidates) == 1:
            erp = candidates[0]
            results.append(
                _compare_groups(
                    customer,
                    erp,
                    match_level="辅助匹配",
                    tolerances=effective_tolerances,
                    compare_customer_po=True,
                )
            )
            matched_erp_ids.add(id(erp))
        elif len(candidates) > 1:
            ambiguous_erp_ids.update(id(candidate) for candidate in candidates)
            candidate_refs = "、".join(
                ",".join(candidate.references) for candidate in candidates
            )
            results.append(
                ReconciliationResult(
                    difference_type=MANUAL_REVIEW,
                    match_status=MANUAL_REVIEW,
                    match_level="辅助匹配",
                    difference_types=(MANUAL_REVIEW,),
                    customer_group=customer,
                    handling_suggestion=_suggestion((MANUAL_REVIEW,)),
                    notes=(f"存在多个 ERP 候选行：{candidate_refs}",),
                    requires_manual_review=True,
                )
            )
        else:
            results.append(_unmatched_result(customer, difference_type=CUSTOMER_ONLY))

    for erp in erp_groups:
        if id(erp) in matched_erp_ids or id(erp) in ambiguous_erp_ids:
            continue
        results.append(_unmatched_result(erp, difference_type=ERP_ONLY))

    results.sort(
        key=lambda result: (
            min(
                result.customer_group.source_rows
                if result.customer_group
                else (10**9,)
            ),
            min(result.erp_group.source_rows if result.erp_group else (10**9,)),
            result.difference_type,
        )
    )
    return normalized_erp, tuple(results)


def reconcile_customer_statement_xlsx(
    source: bytes | bytearray | Path | str | BinaryIO,
    erp_rows: Iterable[Mapping[str, Any] | StatementRow],
    *,
    metadata: ReconciliationMetadata,
    sheet_name: str = DEFAULT_CUSTOMER_SHEET_NAME,
    tolerances: ReconciliationTolerance | None = None,
) -> ReconciliationReport:
    parsed = parse_customer_statement_xlsx(source, sheet_name=sheet_name)
    effective_tolerances = tolerances or ReconciliationTolerance()
    normalized_erp, results = reconcile_statement_rows(
        parsed.rows,
        erp_rows,
        tolerances=effective_tolerances,
    )
    return ReconciliationReport(
        customer_statement=parsed,
        erp_rows=normalized_erp,
        results=results,
        metadata=metadata,
        tolerances=effective_tolerances,
    )


def _excel_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _excel_number(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _join(values: Iterable[Any]) -> str:
    return ", ".join(str(value) for value in values)


def _group_value(group: StatementGroup | None, field_name: str) -> Any:
    return getattr(group, field_name) if group is not None else None


def _style_tabular_sheet(sheet: Any, *, widths: Sequence[int]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 42)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _duplicate_labels(rows: Sequence[StatementRow]) -> dict[tuple[str, int], str]:
    labels: dict[tuple[str, int], str] = {}
    _, groups = _deduplicate_rows(rows)
    for index, group in enumerate(groups, 1):
        label = f"DUP-{index:03d}"
        for row in group:
            labels[(row.source_kind, row.source_row)] = label
    return labels


def generate_reconciliation_xlsx(report: ReconciliationReport) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "核对汇总"
    detail_sheet = workbook.create_sheet("差异明细")
    customer_sheet = workbook.create_sheet("客户原始数据标准化结果")
    erp_sheet = workbook.create_sheet("ERP对账数据标准化结果")

    customer_rows = report.customer_statement.rows
    erp_rows = report.erp_rows
    customer_quantity = sum(
        (row.quantity or Decimal("0") for row in customer_rows),
        Decimal("0"),
    )
    erp_quantity = sum(
        (row.quantity or Decimal("0") for row in erp_rows),
        Decimal("0"),
    )
    customer_amount = sum(
        (row.amount or Decimal("0") for row in customer_rows),
        Decimal("0"),
    )
    erp_amount = sum(
        (row.amount or Decimal("0") for row in erp_rows),
        Decimal("0"),
    )
    summary_rows = [
        ("客户", report.metadata.customer_name),
        ("对账月份", report.metadata.statement_month),
        ("客户文件名", report.metadata.customer_file_name),
        ("ERP对账单范围", report.metadata.erp_statement_range),
        ("客户文件总行数", len(customer_rows)),
        ("ERP对账单总行数", len(erp_rows)),
        ("客户文件总数量", _excel_number(customer_quantity)),
        ("ERP总数量", _excel_number(erp_quantity)),
        ("客户文件总金额", _excel_number(customer_amount)),
        ("ERP总金额", _excel_number(erp_amount)),
        ("匹配成功行数", len(report.matched_results)),
        ("差异行数", len(report.difference_results)),
        ("待人工确认行数", len(report.manual_review_results)),
        ("数量容差", _excel_number(report.tolerances.quantity)),
        ("单价容差", _excel_number(report.tolerances.unit_price)),
        ("金额容差", _excel_number(report.tolerances.amount)),
        ("生成时间", report.metadata.generated_at),
        ("操作人", report.metadata.operator),
    ]
    summary_sheet.append(["核对项目", "结果"])
    for label, value in summary_rows:
        summary_sheet.append([label, _excel_safe(value)])
    summary_sheet.sheet_view.showGridLines = False
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 52
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center")
    for row_number in range(2, summary_sheet.max_row + 1):
        summary_sheet.cell(row_number, 1).font = Font(bold=True)
    for row_number in (8, 9, 10, 11, 15, 16, 17):
        summary_sheet.cell(row_number, 2).number_format = "#,##0.00####"
    summary_sheet.cell(18, 2).number_format = "yyyy-mm-dd hh:mm:ss"

    detail_headers = [
        "差异类型",
        "匹配状态",
        "匹配层级",
        "客户文件行号",
        "ERP对账单行号",
        "客户订单号",
        "ERP客户订单号",
        "存货编码",
        "客户文件产品名称",
        "ERP产品名称",
        "客户文件规格",
        "ERP规格",
        "客户文件单位",
        "ERP单位",
        "客户文件数量",
        "ERP数量",
        "数量差异",
        "客户文件单价",
        "ERP单价",
        "单价差异",
        "客户文件金额",
        "ERP金额",
        "金额差异",
        "客户入库日期",
        "ERP送货日期",
        "处理建议",
        "备注",
    ]
    detail_sheet.append(detail_headers)
    for result in report.results:
        customer = result.customer_group
        erp = result.erp_group
        detail_sheet.append(
            [
                _excel_safe(result.difference_type),
                result.match_status,
                result.match_level,
                _join(customer.source_rows) if customer else None,
                _join(erp.references) if erp else None,
                _excel_safe(_group_value(customer, "customer_po")),
                _excel_safe(_group_value(erp, "customer_po")),
                _excel_safe(
                    _group_value(customer, "product_code")
                    or _group_value(erp, "product_code")
                ),
                _excel_safe(_group_value(customer, "product_name")),
                _excel_safe(_group_value(erp, "product_name")),
                _excel_safe(_group_value(customer, "specification")),
                _excel_safe(_group_value(erp, "specification")),
                _excel_safe(_group_value(customer, "unit")),
                _excel_safe(_group_value(erp, "unit")),
                _excel_number(_group_value(customer, "quantity")),
                _excel_number(_group_value(erp, "quantity")),
                _excel_number(result.quantity_difference),
                _excel_number(_group_value(customer, "unit_price")),
                _excel_number(_group_value(erp, "unit_price")),
                _excel_number(result.unit_price_difference),
                _excel_number(_group_value(customer, "amount")),
                _excel_number(_group_value(erp, "amount")),
                _excel_number(result.amount_difference),
                _group_value(customer, "receipt_date"),
                _group_value(erp, "delivery_date"),
                _excel_safe(result.handling_suggestion),
                _excel_safe("；".join(result.notes)),
            ]
        )
    _style_tabular_sheet(
        detail_sheet,
        widths=(
            24, 14, 12, 15, 16, 18, 18, 16, 24, 24, 28, 28, 12, 12,
            14, 14, 14, 14, 14, 14, 15, 15, 15, 14, 14, 34, 42,
        ),
    )
    for row in detail_sheet.iter_rows(min_row=2, min_col=15, max_col=23):
        for cell in row:
            cell.number_format = "#,##0.00####"
    for row in detail_sheet.iter_rows(min_row=2, min_col=24, max_col=25):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"

    customer_headers = [
        "客户文件行号", "序号", "入库日期", "入库单号", "供货商编码",
        "供货商名称", "存货编码", "存货名称", "规格型号", "单位",
        "实收数量", "采购订单号", "原币含税单价", "原币金额",
        "原币价税合计", "重复组", "解析警告",
    ]
    customer_sheet.append(customer_headers)
    customer_duplicate_labels = _duplicate_labels(customer_rows)
    for row in customer_rows:
        customer_sheet.append(
            [
                row.source_row,
                _excel_safe(row.sequence),
                row.receipt_date,
                _excel_safe(row.receipt_number),
                _excel_safe(row.supplier_code),
                _excel_safe(row.supplier_name),
                _excel_safe(row.product_code),
                _excel_safe(row.product_name),
                _excel_safe(row.specification),
                _excel_safe(row.unit),
                _excel_number(row.quantity),
                _excel_safe(row.customer_po),
                _excel_number(row.unit_price),
                _excel_number(row.net_amount),
                _excel_number(row.amount),
                customer_duplicate_labels.get((row.source_kind, row.source_row)),
                _excel_safe("；".join(row.warnings)),
            ]
        )
    _style_tabular_sheet(
        customer_sheet,
        widths=(12, 10, 13, 18, 14, 24, 16, 26, 36, 10, 14, 18, 14, 14, 15, 12, 34),
    )
    for cell in customer_sheet["C"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for row in customer_sheet.iter_rows(min_row=2, min_col=11, max_col=15):
        for cell in row:
            cell.number_format = "#,##0.00####"

    erp_headers = [
        "ERP行号", "ERP明细ID", "送货日期", "送货单号", "客户订单号",
        "存货编码", "产品名称", "规格型号", "单位", "实收数量",
        "含税单价", "含税金额", "重复组", "解析警告",
    ]
    erp_sheet.append(erp_headers)
    erp_duplicate_labels = _duplicate_labels(erp_rows)
    for row in erp_rows:
        erp_sheet.append(
            [
                row.source_row,
                _excel_safe(row.reference_id),
                row.delivery_date,
                _excel_safe(row.delivery_number),
                _excel_safe(row.customer_po),
                _excel_safe(row.product_code),
                _excel_safe(row.product_name),
                _excel_safe(row.specification),
                _excel_safe(row.unit),
                _excel_number(row.quantity),
                _excel_number(row.unit_price),
                _excel_number(row.amount),
                erp_duplicate_labels.get((row.source_kind, row.source_row)),
                _excel_safe("；".join(row.warnings)),
            ]
        )
    _style_tabular_sheet(
        erp_sheet,
        widths=(12, 14, 13, 18, 18, 16, 26, 36, 10, 14, 14, 15, 12, 34),
    )
    for cell in erp_sheet["C"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for row in erp_sheet.iter_rows(min_row=2, min_col=10, max_col=12):
        for cell in row:
            cell.number_format = "#,##0.00####"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


generate_reconciliation_report_xlsx = generate_reconciliation_xlsx
