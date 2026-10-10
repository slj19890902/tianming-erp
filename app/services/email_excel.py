"""Bounded spreadsheet interpretation; returns drafts, never order facts."""
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from xml.etree.ElementTree import ParseError

FIELDS = {'product_code':'存货编码','product_name':'产品名称','spec':'规格','quantity':'数量',
          'unit_price':'单价','unit':'单位','customer_po':'客户单号','delivery_date':'交期'}
ALIASES = {'product_code':('存货编码','物料编码','产品编码','料号','客户料号','产品编号'),
    'product_name':('产品名称','物料名称','品名','名称'), 'spec':('规格','规格型号','尺寸'),
    'quantity':('数量','订单数量','订购数量','采购数量'), 'unit_price':('单价','含税单价','未税单价'),
    'unit':('单位',), 'customer_po':('客户单号','客户订单号','订单号','采购订单号'),
    'delivery_date':('交期','交货日期','交货时间','交货期')}


def text(value, number_format=''):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        if re.fullmatch(r'0{2,30}', number_format or ''):
            return str(int(value)).zfill(len(number_format))
        return str(int(value))
    return str(value).strip()[:2000]


def table(content, sheet_name='', header_row=1):
    if not content or len(content) > 8*1024*1024:
        raise ValueError('Excel附件需为不超过8MB的xlsx')
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            files = archive.infolist()
            if len(files)>2000 or sum(f.file_size for f in files)>32*1024*1024:
                raise ValueError('工作簿解压内容过大')
            if any('vbaproject' in f.filename.casefold() for f in files):
                raise ValueError('请另存为不含宏的xlsx')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError, InvalidFileException, ParseError) as error:
        raise ValueError('无法读取工作簿，请另存为xlsx') from error
    try:
        if not 1 <= header_row <= 100:
            raise ValueError('标题行需为1至100')
        if len(book.sheetnames)>50:
            raise ValueError('工作表超过50张，请保留订单页')
        name = sheet_name or book.sheetnames[0]
        if name not in book.sheetnames:
            raise ValueError('工作表不存在，请重新选择')
        sheet = book[name]
        if (sheet.max_column or 0)>100 or (sheet.max_row or 0)>5000:
            raise ValueError('工作表超过5000行或100列，请拆分后再处理')
        rows, formula_rows = [], []
        for row_no,cells in enumerate(sheet.iter_rows(max_row=min(sheet.max_row or 5001,5001), max_col=min(sheet.max_column or 1,100)),1):
            if row_no>5000:
                if any(c.value is not None for c in cells):
                    raise ValueError('工作表超过5000行')
                break
            values=[]
            for cell in cells:
                if cell.data_type=='f':
                    formula_rows.append(row_no)
                    values.append('')
                else:
                    values.append(text(cell.value,cell.number_format))
            rows.append(values)
        headers = rows[header_row-1] if len(rows)>=header_row else []
        recommended={}
        for key, aliases in ALIASES.items():
            matches=[i for i,h in enumerate(headers) if re.sub(r'\s+','',h) in aliases]
            if len(matches)==1:
                recommended[key]=matches[0]
        return {'sheets':book.sheetnames,'sheet_name':name,'header_row':header_row,'headers':headers,
                'columns':recommended,'fields':FIELDS,'column_labels':[get_column_letter(i+1) for i in range(len(headers))],
                'sample':rows[header_row:header_row+5],'rows':rows[header_row:], 'formula_rows':sorted(set(formula_rows))}
    finally:
        book.close()


def parse_number(value, quantity=False):
    try:
        number=Decimal(re.sub(r'[,，￥¥\s]','',value))
    except InvalidOperation:
        return None
    if not number.is_finite() or abs(number)>Decimal('1000000000000'):
        return None
    if quantity:
        return int(number) if number>0 and number==number.to_integral_value() else None
    return str(number) if number>=0 else None


def drafts(parsed, columns, *, customer_po='', order_date=None, delivery_date=None):
    if not isinstance(columns,dict) or set(columns)-set(FIELDS):
        raise ValueError('列对应字段无效')
    if 'quantity' not in columns or not {'product_code','product_name'} & set(columns):
        raise ValueError('请对应数量列，以及存货编码或产品名称列')
    if any(type(v) is not int or not 0<=v<len(parsed['headers']) for v in columns.values()):
        raise ValueError('列对应超出工作表范围')
    if len(set(columns.values()))!=len(columns):
        raise ValueError('同一列不能对应两个字段')
    groups=defaultdict(list); warnings=defaultdict(list); skipped=[]
    for line_no,row in enumerate(parsed['rows'],parsed['header_row']+1):
        fields={key:row[index] if index<len(row) else '' for key,index in columns.items()}
        if not any(fields.values()) and line_no not in parsed['formula_rows']:
            continue
        identity=[fields.get('product_code',''),fields.get('product_name','')]
        if any(value.casefold() in ('合计','总计','小计','total','subtotal') for value in identity):
            skipped.append(line_no); continue
        po=fields.get('customer_po') or customer_po
        quantity=parse_number(fields.get('quantity',''),True)
        price=parse_number(fields.get('unit_price',''))
        item_warnings=[]
        if quantity is None:item_warnings.append(f'原表第{line_no}行数量缺失或不是正整数，请人工填写')
        if 'unit_price' in columns and price is None:item_warnings.append(f'原表第{line_no}行单价无效，请人工核对')
        if line_no in parsed['formula_rows']:item_warnings.append(f'原表第{line_no}行含公式，未计算公式，请核对空白字段')
        raw_date=fields.get('delivery_date') or delivery_date
        parsed_date=None
        if raw_date:
            try:parsed_date=date.fromisoformat(str(raw_date).replace('/','-').replace('.','-')).isoformat()
            except ValueError:item_warnings.append(f'原表第{line_no}行交期格式需核对：{raw_date}')
        groups[po].append({'line_no':line_no,'product_code':identity[0],'raw_product_code':identity[0],
            'product_name':identity[1],'raw_product_name':identity[1], 'spec':fields.get('spec',''),
            'raw_spec_model':fields.get('spec',''),'quantity':quantity,'unit_price':price,
            'unit':fields.get('unit',''),'delivery_date':parsed_date,'warnings':item_warnings})
        warnings[po].extend(item_warnings)
    if not groups:
        raise ValueError('没有找到订单明细，请检查工作表、标题行及列对应')
    if sum(map(len,groups.values()))>500:
        raise ValueError('单次超过500条明细，请拆分后处理')
    result=[]
    for po,items in groups.items():
        notes=warnings[po]+['请核对原表单价的含税口径；识别不会改写现有订单。']
        if skipped:notes.append('已跳过明确标注合计/小计的原表行：'+','.join(map(str,skipped)))
        result.append({'customer_po':po,'order_date':order_date,'delivery_date':delivery_date,'items':items,
            'item_count':len(items),'recognition_status':'needs_confirmation','warnings':notes,
            'integrity_check':{'integrity_status':'needs_confirmation','integrity_errors':warnings[po]}})
    return result
