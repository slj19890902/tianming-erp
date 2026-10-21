from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import hashlib
import json
import re
from pathlib import PurePath
from typing import Any, Literal
from zipfile import BadZipFile, ZipFile


MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_SHEETS = 30
MAX_ROWS = 20_000
MAX_COLUMNS = 200
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
PARSER_VERSION = "excel-order-predelivery-v4"
PICK_LIST_TEMPLATE = "yg_gy_location_order_v1"
SHORTAGE_LIST_TEMPLATE = "yg_gy_shortage_order_v1"
ORDER_LIST_TEMPLATES = frozenset({PICK_LIST_TEMPLATE, SHORTAGE_LIST_TEMPLATE})


class ExcelImportError(ValueError):
    pass


@dataclass(frozen=True)
class SourceCell:
    row: int
    column: int
    coordinate: str
    raw_value: Any
    display_value: str


@dataclass
class ImportedRow:
    sheet: str
    source_row: int
    source_no: str | None
    customer_po: str | None
    stock_code: str
    product_name: str | None
    requested_quantity: int
    order_quantity: int | None = None
    unit: str | None = None
    unit_price: Decimal | None = None
    amount: Decimal | None = None
    drawing_number: str | None = None
    category: str | None = None
    model: str | None = None
    source_location: str | None = None
    source_stock_reference: dict[str, str | None] | None = None
    cells: dict[str, SourceCell] = field(default_factory=dict)
    issues: list[str] = field(default_factory=list)

    def normalized_identity(self) -> dict[str, Any]:
        return {
            "sheet": self.sheet,
            "source_row": self.source_row,
            "stock_code": self.stock_code,
            "requested_quantity": self.requested_quantity,
            "order_quantity": self.order_quantity,
            "customer_po": self.customer_po,
        }


@dataclass
class ImportedDocument:
    document_type: Literal["order", "pre_delivery"]
    source_format: Literal["xls", "xlsx"]
    template: str
    filename: str
    source_hash: str
    parser_version: str
    customer_code: str | None
    customer_name_evidence: str | None
    source_date: date | None
    delivery_date: date | None
    customer_po: str | None
    rows: list[ImportedRow]
    skipped_sheets: list[str]
    warnings: list[str]

    @property
    def business_fingerprint(self) -> str:
        payload = {
            "document_type": self.document_type,
            "customer_code": self.customer_code,
            "source_date": self.source_date.isoformat() if self.source_date else None,
            "customer_po": self.customer_po,
            "rows": [row.normalized_identity() for row in self.rows],
        }
        return hashlib.sha256(
            json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest()


def _safe_filename(filename: str) -> str:
    value = PurePath(str(filename or "Excel来源").replace("\\", "/")).name
    value = "".join(character for character in value if character >= " " and character != "\x7f")
    return value[:255] or "Excel来源"


def _decimal(value: Any, *, label: str, required: bool = False) -> Decimal | None:
    if value in (None, ""):
        if required:
            raise ExcelImportError(f"{label}为空")
        return None
    text = str(value).strip().replace(",", "")
    try:
        result = Decimal(text)
    except (InvalidOperation, ValueError) as error:
        raise ExcelImportError(f"{label}不是有效数字：{text}") from error
    if not result.is_finite():
        raise ExcelImportError(f"{label}不是有限数字")
    return result


def _positive_integer(value: Any, *, label: str) -> int:
    result = _decimal(value, label=label, required=True)
    assert result is not None
    if result <= 0 or result != result.to_integral_value():
        raise ExcelImportError(f"{label}必须是正整数")
    return int(result)


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in (r"(20\d{2})[-/.年](\d{1,2})[-/.月](\d{1,2})", r"(20\d{2})(\d{2})(\d{2})"):
        match = re.search(pattern, text)
        if match:
            try:
                return date(*(int(part) for part in match.groups()))
            except ValueError:
                return None
    return None


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\x00", "").strip()
    return text or None


def _code_text(raw: Any, display: str) -> str:
    text = display.strip()
    if not text:
        return ""
    if isinstance(raw, float) and raw.is_integer() and text.endswith(".0"):
        text = text[:-2]
    return text


class _Sheet:
    def __init__(self, name: str, rows: list[list[SourceCell]], *, hidden: bool = False):
        self.name = name
        self.rows = rows
        self.hidden = hidden

    def cell(self, row: int, column: int) -> SourceCell:
        if row < 1 or column < 1 or row > len(self.rows) or column > len(self.rows[row - 1]):
            return SourceCell(row, column, _coordinate(row, column), None, "")
        return self.rows[row - 1][column - 1]

    def text(self, row: int, column: int) -> str:
        return self.cell(row, column).display_value.strip()

    @property
    def max_row(self) -> int:
        return len(self.rows)


def _coordinate(row: int, column: int) -> str:
    letters = ""
    value = column
    while value:
        value, remainder = divmod(value - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"


def _formatted_numeric(value: Any, number_format: str | None) -> str:
    fmt = str(number_format or "").split(";")[0]
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        # Only implement the unambiguous date tokens used by order templates;
        # this preserves Excel's yyyymmd (e.g. 2026-01-03 -> 2026013) rather
        # than serialising a datetime or its Excel serial number.
        token_format = fmt.casefold()
        if "y" in token_format and "m" in token_format and "d" in token_format:
            rendered: list[str] = []
            index = 0
            while index < len(fmt):
                char = fmt[index]
                if char == '"':
                    closing = fmt.find('"', index + 1)
                    if closing < 0:
                        return value.isoformat()
                    rendered.append(fmt[index + 1:closing])
                    index = closing + 1
                    continue
                if char == "\\" and index + 1 < len(fmt):
                    rendered.append(fmt[index + 1])
                    index += 2
                    continue
                token = char.casefold()
                if token in {"y", "m", "d"}:
                    end = index + 1
                    while end < len(fmt) and fmt[end].casefold() == token:
                        end += 1
                    width = end - index
                    if token in {"m", "d"} and width > 2:
                        return value.isoformat()
                    raw = value.year % 100 if token == "y" and width == 2 else {
                        "y": value.year, "m": value.month, "d": value.day,
                    }[token]
                    rendered.append(str(raw).zfill(width) if width > 1 else str(raw))
                    index = end
                    continue
                rendered.append(char)
                index += 1
            return "".join(rendered)
        return value.isoformat()
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return _clean_text(value) or ""
    integer = int(value) if float(value).is_integer() else None
    zero_match = re.fullmatch(r"0+", fmt)
    if integer is not None and zero_match:
        return f"{integer:0{len(fmt)}d}"
    return str(integer if integer is not None else value)


def _read_xlsx(content: bytes) -> list[_Sheet]:
    try:
        with ZipFile(BytesIO(content)) as archive:
            names = archive.namelist()
            total = sum(info.file_size for info in archive.infolist())
            compressed = sum(max(info.compress_size, 1) for info in archive.infolist())
            if total > MAX_UNCOMPRESSED_BYTES or total / max(compressed, 1) > MAX_COMPRESSION_RATIO:
                raise ExcelImportError("Excel 解压体积或压缩比超过安全限制")
            lowered = {name.casefold() for name in names}
            if any("vbaproject.bin" in name for name in lowered):
                raise ExcelImportError("不接受含宏的 Excel 文件")
            if any(name.startswith("xl/externallinks/") for name in lowered):
                raise ExcelImportError("不接受含外部链接的 Excel 文件")
    except BadZipFile as error:
        raise ExcelImportError("xlsx 文件结构损坏") from error
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(BytesIO(content), read_only=False, data_only=True, keep_links=False)
    except Exception as error:
        raise ExcelImportError("xlsx 无法读取，可能已损坏或加密") from error
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise ExcelImportError("Excel 工作表数量超过限制")
        result: list[_Sheet] = []
        for worksheet in workbook.worksheets:
            if worksheet.max_row > MAX_ROWS or worksheet.max_column > MAX_COLUMNS:
                raise ExcelImportError(f"工作表“{worksheet.title}”行列数超过限制")
            rows: list[list[SourceCell]] = []
            for row_index, cells in enumerate(worksheet.iter_rows(), start=1):
                rows.append([
                    SourceCell(row_index, cell.column, cell.coordinate, cell.value,
                               _formatted_numeric(cell.value, cell.number_format))
                    for cell in cells
                ])
            result.append(_Sheet(worksheet.title, rows, hidden=worksheet.sheet_state != "visible"))
        return result
    finally:
        workbook.close()


def _read_xls(content: bytes) -> list[_Sheet]:
    try:
        import xlrd
    except ImportError as error:
        raise ExcelImportError("服务器缺少 xls 读取组件，请安装 xlrd 2.0.2") from error
    try:
        workbook = xlrd.open_workbook(file_contents=content, on_demand=True, formatting_info=True)
    except Exception as error:
        raise ExcelImportError("xls 无法读取，可能已损坏、加密或不是有效工作簿") from error
    try:
        if workbook.nsheets > MAX_SHEETS:
            raise ExcelImportError("Excel 工作表数量超过限制")
        result: list[_Sheet] = []
        for index in range(workbook.nsheets):
            worksheet = workbook.sheet_by_index(index)
            if worksheet.nrows > MAX_ROWS or worksheet.ncols > MAX_COLUMNS:
                raise ExcelImportError(f"工作表“{worksheet.name}”行列数超过限制")
            rows: list[list[SourceCell]] = []
            for row_index in range(worksheet.nrows):
                values: list[SourceCell] = []
                for column_index in range(worksheet.ncols):
                    cell = worksheet.cell(row_index, column_index)
                    raw: Any = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            raw = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode)
                        except Exception:
                            raw = cell.value
                    number_format = None
                    if cell.xf_index is not None and cell.xf_index < len(workbook.xf_list):
                        format_key = workbook.xf_list[cell.xf_index].format_key
                        format_info = workbook.format_map.get(format_key)
                        number_format = format_info.format_str if format_info else None
                    values.append(SourceCell(row_index + 1, column_index + 1,
                        _coordinate(row_index + 1, column_index + 1), raw,
                        _formatted_numeric(raw, number_format)))
                rows.append(values)
            visibility = getattr(workbook, "sheet_visibility", [0] * workbook.nsheets)
            result.append(_Sheet(worksheet.name, rows, hidden=bool(visibility[index])))
        return result
    finally:
        workbook.release_resources()


def _read_sheets(content: bytes) -> tuple[Literal["xls", "xlsx"], list[_Sheet]]:
    if content.startswith(OLE_SIGNATURE):
        return "xls", _read_xls(content)
    if content.startswith(ZIP_SIGNATURES):
        return "xlsx", _read_xlsx(content)
    raise ExcelImportError("文件签名不是有效的 xls 或 xlsx，不能只修改扩展名导入")


def _row(sheet: _Sheet, row_number: int, mapping: dict[str, int], *, document_type: str) -> ImportedRow | None:
    cells = {name: sheet.cell(row_number, column) for name, column in mapping.items()}
    stock = _code_text(cells["stock_code"].raw_value, cells["stock_code"].display_value)
    quantity_cell = cells["order_quantity"] if document_type == "order" else cells["requested_quantity"]
    if not stock and not quantity_cell.display_value.strip():
        return None
    if not stock:
        return None
    quantity = _positive_integer(quantity_cell.raw_value, label=f"{sheet.name}!{quantity_cell.coordinate} 数量")
    requested = quantity
    order_quantity = quantity if document_type == "order" else None
    issues: list[str] = []
    if document_type == "order" and "requested_quantity" in cells:
        requested_value = _positive_integer(cells["requested_quantity"].raw_value,
            label=f"{sheet.name}!{cells['requested_quantity'].coordinate} 要求数")
        requested = requested_value
        if requested_value != quantity:
            issues.append(f"要求数 {requested_value} 与发注数 {quantity} 不一致，请确认订单数量")
    return ImportedRow(
        sheet=sheet.name,
        source_row=row_number,
        source_no=_clean_text(cells.get("source_no").display_value if cells.get("source_no") else None),
        customer_po=_clean_text(cells.get("customer_po").display_value if cells.get("customer_po") else None),
        stock_code=stock,
        product_name=_clean_text(cells.get("product_name").display_value if cells.get("product_name") else None),
        requested_quantity=requested,
        order_quantity=order_quantity,
        unit=_clean_text(cells.get("unit").display_value if cells.get("unit") else None),
        unit_price=_decimal(cells.get("unit_price").raw_value, label="单价") if cells.get("unit_price") else None,
        amount=_decimal(cells.get("amount").raw_value, label="金额") if cells.get("amount") else None,
        drawing_number=_clean_text(cells.get("drawing_number").display_value if cells.get("drawing_number") else None),
        category=_clean_text(cells.get("category").display_value if cells.get("category") else None),
        model=_clean_text(cells.get("model").display_value if cells.get("model") else None),
        cells=cells,
        issues=issues,
    )


def _sheet_text(sheet: _Sheet) -> str:
    return " ".join(cell.display_value for row in sheet.rows for cell in row if cell.display_value)


def _management_no(sheet: _Sheet) -> str | None:
    values: set[str] = set()
    for row in sheet.rows:
        for index, cell in enumerate(row):
            label = re.sub(r"\s+", "", cell.display_value).upper()
            if not re.fullmatch(r"管理NO\.?[:：]?", label):
                continue
            if index + 1 >= len(row):
                raise ExcelImportError(f"{sheet.name}!{cell.coordinate} 管理NO缺少右侧订单号")
            value = _clean_text(row[index + 1].display_value)
            if not value:
                raise ExcelImportError(f"{sheet.name}!{cell.coordinate} 管理NO缺少右侧订单号")
            values.add(value)
    if len(values) > 1:
        raise ExcelImportError(f"{sheet.name} 存在多个管理NO，不能静默合并订单号")
    return next(iter(values), None)


def _parse_yanguang_order(sheets: list[_Sheet]) -> tuple[list[ImportedRow], date | None, str | None, list[str]]:
    rows: list[ImportedRow] = []
    warnings: list[str] = []
    source_date: date | None = None
    customer_po: str | None = None
    mapping = {"source_no": 1, "customer_po": 2, "stock_code": 3, "product_name": 4,
               "drawing_number": 5, "requested_quantity": 6, "unit": 7,
               "order_quantity": 10, "unit_price": 12, "amount": 13}
    for sheet in sheets:
        text = _sheet_text(sheet)
        if "品目" not in text or ("發注" not in text and "发注" not in text):
            continue
        if sheet.hidden:
            warnings.append(f"隐藏工作表“{sheet.name}”未自动导入")
            continue
        management_no = _management_no(sheet)
        if management_no:
            if customer_po and customer_po != management_no:
                raise ExcelImportError("工作簿存在多个管理NO，不能静默合并订单号")
            customer_po = management_no
        source_date = source_date or _date_value(sheet.cell(4, 12).raw_value) or _date_value(sheet.cell(2, 16).raw_value)
        for number in range(11, sheet.max_row + 1):
            try:
                parsed = _row(sheet, number, mapping, document_type="order")
            except ExcelImportError as error:
                if sheet.text(number, 3):
                    raise
                continue
            if parsed:
                if management_no:
                    if parsed.customer_po and parsed.customer_po != management_no:
                        raise ExcelImportError("工作簿存在多个订单号，不能静默合并")
                    parsed.customer_po = management_no
                if parsed.customer_po:
                    if customer_po and customer_po != parsed.customer_po:
                        raise ExcelImportError("工作簿存在多个订单号，不能静默合并")
                    customer_po = parsed.customer_po
                rows.append(parsed)
    if not rows:
        raise ExcelImportError("未识别到购买要求书有效明细")
    return rows, source_date, customer_po, warnings


def _parse_yanguang_predelivery(sheets: list[_Sheet]) -> tuple[list[ImportedRow], date | None, list[str]]:
    mapping = {"source_no": 1, "product_name": 2, "stock_code": 3,
               "drawing_number": 4, "requested_quantity": 5, "unit_price": 6, "amount": 7}
    rows: list[ImportedRow] = []
    warnings: list[str] = []
    delivery_date: date | None = None
    for sheet in sheets:
        text = _sheet_text(sheet)
        if "品目" not in text or "数量" not in text:
            continue
        if sheet.hidden:
            warnings.append(f"隐藏工作表“{sheet.name}”未自动导入")
            continue
        delivery_date = delivery_date or _date_value(sheet.text(2, 1)) or _date_value(text[:300])
        for number in range(4, sheet.max_row + 1):
            parsed = _row(sheet, number, mapping, document_type="pre_delivery")
            if parsed:
                rows.append(parsed)
    if not rows:
        raise ExcelImportError("未识别到研光预送货有效明细")
    return rows, delivery_date, warnings


def _parse_guangyang_predelivery(sheets: list[_Sheet]) -> tuple[list[ImportedRow], date | None, list[str]]:
    mapping = {"source_no": 1, "model": 2, "drawing_number": 3, "stock_code": 4,
               "category": 5, "requested_quantity": 6, "unit_price": 7, "amount": 8}
    rows: list[ImportedRow] = []
    warnings: list[str] = []
    delivery_date: date | None = None
    for sheet in sheets:
        text = _sheet_text(sheet)
        sample_codes = sum(1 for number in range(1, min(sheet.max_row, 40) + 1)
                           if re.fullmatch(r"C?\d{6,}", sheet.text(number, 4), re.IGNORECASE))
        if sample_codes < 3:
            continue
        if sheet.hidden:
            warnings.append(f"隐藏工作表“{sheet.name}”未自动导入")
            continue
        delivery_date = delivery_date or _date_value(text[:500])
        for number in range(1, sheet.max_row + 1):
            parsed = _row(sheet, number, mapping, document_type="pre_delivery")
            if parsed:
                rows.append(parsed)
    if not rows:
        raise ExcelImportError("未识别到光洋预送货有效明细")
    return rows, delivery_date, warnings


def _location_price_columns(sheet: _Sheet) -> tuple[int, int] | None:
    layouts = set()
    for number in range(1, sheet.max_row + 1):
        if (re.fullmatch(r"C?\d{6,}", sheet.text(number, 4), re.IGNORECASE)
                and sheet.text(number, 5)
                and not sheet.text(number, 6) and not sheet.text(number, 7)):
            if sheet.text(number, 12):
                layouts.add((12, 13))
            elif not sheet.text(number, 13) and sheet.text(number, 14):
                layouts.add((14, 15))
    if len(layouts) > 1:
        raise ExcelImportError(f"{sheet.name} 单价列不统一，请核对L/M列或N/O列")
    return next(iter(layouts), None)


def _is_location_order(sheet: _Sheet) -> bool:
    return _location_price_columns(sheet) is not None


def _shortage_header(sheet: _Sheet):
    aliases = {
        "stock_code": {"品目号", "品目號", "存货编码"},
        "order_quantity": {"欠品数量", "欠品數量"},
        "product_name": {"品名", "产品名称"},
        "model": {"使用型号", "使用型號"},
        "drawing_number": {"规格/图号", "規格/圖號"},
    }
    for number in range(1, min(sheet.max_row, 20) + 1):
        mapping = {}
        for col in range(1, len(sheet.rows[number - 1]) + 1):
            text = re.sub(r"\s+", "", sheet.text(number, col)).replace("／", "/")
            for field, labels in aliases.items():
                if text in labels:
                    if field in mapping:
                        raise ExcelImportError(f"{sheet.name} 第{number}行{field}表头重复")
                    mapping[field] = col
        if {"stock_code", "order_quantity"}.issubset(mapping):
            mapping["requested_quantity"] = mapping["order_quantity"]
            return number, mapping
    return None


def _parse_shortage_order(sheets: list[_Sheet], document_type: str):
    rows, warnings = [], []
    for sheet in sheets:
        header = _shortage_header(sheet)
        if sheet.hidden or header is None:
            continue
        number, mapping = header
        for row_number in range(number + 1, sheet.max_row + 1):
            code_cell = sheet.cell(row_number, mapping["stock_code"])
            qty_cell = sheet.cell(row_number, mapping["order_quantity"])
            if not code_cell.display_value and not qty_cell.display_value:
                continue
            code = _code_text(code_cell.raw_value, code_cell.display_value)
            if not re.fullmatch(r"C?\d{6,}", code, re.IGNORECASE):
                raise ExcelImportError(f"{sheet.name}!{code_cell.coordinate} 存货编码无效，请核对原表")
            row = _row(sheet, row_number, mapping, document_type=document_type)
            assert row is not None
            rows.append(row)
    if not rows:
        raise ExcelImportError("未识别到欠品清单有效明细")
    warnings.append("订单数量采用欠品数量，不乘配套数；原表未提供单价时请核对常用箱报价，并确认交货日期。")
    return rows, warnings


def _parse_location_order(sheets: list[_Sheet], document_type: str) -> tuple[list[ImportedRow], list[str]]:
    mapping = dict(product_name=1, drawing_number=2, source_location=3,
                   stock_code=4, requested_quantity=5, order_quantity=5,
                   category=9, source_on_hand=10, source_remaining=11,
                   unit_price=12, amount=13)
    rows, warnings = [], []
    for sheet in sheets:
        if sheet.hidden:
            warnings.append(f"隐藏工作表“{sheet.name}”未自动导入")
            continue
        price_col, amount_col = _location_price_columns(sheet)
        mapping.update(unit_price=price_col, amount=amount_col)
        for number in range(1, sheet.max_row + 1):
            code = sheet.text(number, 4)
            if not re.fullmatch(r"C?\d{6,}", code, re.IGNORECASE):
                # Summary/footer prose has no product identity. A populated
                # detail with a malformed code must not silently disappear.
                footer = (not sheet.text(number, 1) and not sheet.text(number, 2)
                          and re.match(r"^(研光|光洋)", code)
                          and not re.search(r"^\d+(?:\.\d+)?$", sheet.text(number, 5)))
                if footer:
                    warnings.append(f"{sheet.name} 第{number}行备注：" + " / ".join(
                        sheet.text(number, col) for col in (3, 4, 5, 8) if sheet.text(number, col)))
                    continue
                if sheet.text(number, 8) or (sheet.text(number, 1) and sheet.text(number, 5)):
                    raise ExcelImportError(f"{sheet.name}!D{number} 存货编码无效，请核对原表")
                if any(sheet.text(number, col) for col in (3, 4, 5)):
                    warnings.append(f"{sheet.name} 第{number}行备注：" +
                                    " / ".join(sheet.text(number, col) for col in (3, 4, 5) if sheet.text(number, col)))
                continue
            duplicate_code = sheet.text(number, 8)
            if duplicate_code and duplicate_code != code:
                raise ExcelImportError(f"{sheet.name}!D{number} 与 H{number} 编码不一致")
            row = _row(sheet, number, mapping, document_type=document_type)
            assert row is not None
            if row.unit_price is None or row.unit_price < 0:
                raise ExcelImportError(f"{sheet.name}!{sheet.cell(number, price_col).coordinate} 单价必须是非负数字")
            calculated = Decimal(row.requested_quantity) * row.unit_price
            if row.amount is not None and abs(row.amount - calculated) > Decimal("0.000001"):
                raise ExcelImportError(f"{sheet.name}!{sheet.cell(number, amount_col).coordinate} 金额与 E列数量 × 单价不一致，请核对")
            row.amount = calculated if row.amount is None else row.amount
            row.source_location = _clean_text(sheet.text(number, 3))
            on_hand = _decimal(sheet.cell(number, 10).raw_value, label=f"J{number} 原表库存")
            remaining = _decimal(sheet.cell(number, 11).raw_value, label=f"K{number} 原表剩余")
            if remaining is None and on_hand is not None:
                remaining = on_hand - row.requested_quantity
            row.source_stock_reference = {
                "on_hand": str(on_hand) if on_hand is not None else None,
                "remaining": str(remaining) if remaining is not None else None,
            }
            rows.append(row)
    if not rows:
        raise ExcelImportError("未识别到位置清单有效明细")
    warnings.append("C列为指定拿货位置，J/K列仅作原表库存参考；请确认交货日期。未生产和缺货行也可打印提前备货单。")
    return rows, warnings


def parse_excel_document(
    content: bytes,
    filename: str,
    *,
    document_type: Literal["order", "pre_delivery"],
    expected_customer_code: str | None = None,
) -> ImportedDocument:
    if not content:
        raise ExcelImportError("Excel 文件为空")
    if len(content) > MAX_FILE_BYTES:
        raise ExcelImportError("Excel 文件超过 20MB 限制")
    safe_name = _safe_filename(filename)
    source_format, sheets = _read_sheets(bytes(content))
    nonempty = [sheet for sheet in sheets if any(cell.display_value for row in sheet.rows for cell in row)]
    if not nonempty:
        raise ExcelImportError("Excel 没有可读取内容")
    all_text = " ".join(_sheet_text(sheet) for sheet in nonempty)
    source_hash = hashlib.sha256(content).hexdigest()
    skipped = [sheet.name for sheet in sheets if sheet not in nonempty]
    warnings: list[str] = []
    customer_name_evidence: str | None = None
    customer_po: str | None = None
    source_date: date | None = None
    delivery_date: date | None = None

    purchase_sheets = [sheet for sheet in nonempty
                       if "品目" in _sheet_text(sheet)
                       and any(label in _sheet_text(sheet) for label in ("發注", "发注"))]
    # Layout and requested operation are separate: customers use the same
    # purchase workbook for ordering and for the pre-delivery request.
    if purchase_sheets:
        purchase_text = " ".join(_sheet_text(sheet) for sheet in purchase_sheets if not sheet.hidden)
        yanguang = "研光" in purchase_text or "YKE" in purchase_text.upper()
        guangyang = "光洋" in purchase_text or "KEW" in purchase_text.upper()
        if yanguang and guangyang:
            raise ExcelImportError("购买要求书客户识别不唯一，请分别上传研光或光洋文件")
        if not (yanguang or guangyang):
            raise ExcelImportError("购买要求书未识别到研光/YKE或光洋/KEW客户，不能猜测客户")
        customer_code = "YG" if yanguang else "GY"
        template = "yanguang_purchase_order_v1" if yanguang else "guangyang_purchase_order_v1"
        customer_name_evidence = "工作簿内容包含研光/YKE" if yanguang else "工作簿内容包含光洋/KEW"
        rows, source_date, customer_po, extra = _parse_yanguang_order(purchase_sheets)
        warnings.extend(extra)
        if document_type == "pre_delivery":
            for row in rows:
                row.issues = [issue.replace("请确认订单数量", "预送数量采用要求数，请核对本次预送数量") for issue in row.issues]
    elif any(_is_location_order(sheet) for sheet in nonempty if not sheet.hidden):
        template = PICK_LIST_TEMPLATE
        visible_text = " ".join(_sheet_text(sheet) for sheet in nonempty if not sheet.hidden)
        yg, gy = "研光" in visible_text, "光洋" in visible_text
        if yg and gy:
            raise ExcelImportError("位置清单客户不唯一，请分别上传研光或光洋文件")
        customer_code = "YG" if yg else "GY" if gy else (expected_customer_code or "").strip().upper()
        if customer_code not in {"YG", "GY"}:
            raise ExcelImportError("位置清单请先选择研光或光洋客户，再识别")
        customer_name_evidence = "清单客户文字或人工选择的研光/光洋客户"
        matching = [sheet for sheet in nonempty if _is_location_order(sheet)]
        unrecognized = [sheet.name for sheet in nonempty if sheet not in matching]
        skipped.extend(unrecognized)
        if unrecognized:
            warnings.append("未按位置清单导入的工作表：" + "、".join(unrecognized))
        rows, extra = _parse_location_order(matching, document_type)
        warnings.extend(extra)
    elif any(_shortage_header(sheet) for sheet in nonempty if not sheet.hidden):
        template = SHORTAGE_LIST_TEMPLATE
        visible_text = " ".join(_sheet_text(sheet) for sheet in nonempty if not sheet.hidden)
        yg, gy = "研光" in visible_text, "光洋" in visible_text
        if yg and gy:
            raise ExcelImportError("欠品清单客户不唯一，请分别上传")
        customer_code = "YG" if yg else "GY" if gy else (expected_customer_code or "").strip().upper()
        if customer_code not in {"YG", "GY"}:
            raise ExcelImportError("欠品清单请先选择研光或光洋客户，再识别")
        customer_name_evidence = "欠品清单客户文字或人工选择的研光/光洋客户"
        matching = [sheet for sheet in nonempty if not sheet.hidden and _shortage_header(sheet)]
        skipped.extend(sheet.name for sheet in nonempty if sheet not in matching)
        rows, extra = _parse_shortage_order(matching, document_type)
        warnings.extend(extra)
    elif document_type == "order":
        raise ExcelImportError("订单 Excel 模板未识别，请先进行列映射")
    elif "光洋" in all_text or any(
        sum(1 for number in range(1, min(sheet.max_row, 40) + 1)
            if re.fullmatch(r"C\d{6,}", sheet.text(number, 4), re.IGNORECASE)) >= 3
        for sheet in nonempty
    ):
        template = "guangyang_pre_delivery_v1"
        customer_code = "GY"
        customer_name_evidence = "工作簿内容或C前缀品目号符合光洋模板"
        rows, delivery_date, extra = _parse_guangyang_predelivery(nonempty)
        warnings.extend(extra)
    elif "研光" in all_text or "捷太格特" in all_text:
        template = "yanguang_pre_delivery_v1"
        customer_code = "YG"
        customer_name_evidence = "工作簿内容包含研光/捷太格特"
        rows, delivery_date, extra = _parse_yanguang_predelivery(nonempty)
        warnings.extend(extra)
    else:
        raise ExcelImportError("预送货 Excel 模板未识别，请选择客户并进行列映射")

    if expected_customer_code and expected_customer_code.strip().upper() != customer_code:
        raise ExcelImportError(
            f"所选客户 {expected_customer_code.strip().upper()} 与文件识别客户 {customer_code} 不一致"
        )
    return ImportedDocument(
        document_type=document_type,
        source_format=source_format,
        template=template,
        filename=safe_name,
        source_hash=source_hash,
        parser_version=PARSER_VERSION,
        customer_code=customer_code,
        customer_name_evidence=customer_name_evidence,
        source_date=source_date,
        delivery_date=delivery_date,
        customer_po=customer_po,
        rows=rows,
        skipped_sheets=skipped,
        warnings=warnings,
    )
