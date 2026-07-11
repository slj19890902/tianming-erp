from __future__ import annotations

import os
import re
import tempfile
import zipfile
from dataclasses import asdict, dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import PROJECT_ROOT, settings
from app.models.historical_purchase import HistoricalPurchaseEntry


DEFAULT_WORKBOOK_NAMES = ("2025年采购单.xlsx", "2025采购单.xlsx")
DEFAULT_SHEET_NAME = "2020.1-2026"
PAIR_RE = re.compile(
    r"^\s*(\d+(?:\.\d+)?)\s*[*xX×]\s*(\d+(?:\.\d+)?)\s*$"
)
CREASE_RE = re.compile(
    r"^\s*(\d+(?:\.\d+)?)\s*[*xX×]\s*"
    r"(\d+(?:\.\d+)?)\s*[*xX×]\s*(\d+(?:\.\d+)?)\s*$"
)


class HistoricalPurchaseLookupError(RuntimeError):
    pass


@dataclass(frozen=True, slots=True)
class HistoricalPurchaseRecord:
    source_row: int
    supplier_name: str | None
    record_date: date | None
    product_reference: str
    search_text: str
    material_code: str
    historical_quantity: int | None
    board_length_cm: Decimal
    board_width_cm: Decimal
    report_length_mm: int
    report_width_mm: int
    crease_text: str | None
    crease_type: str | None
    crease_left_mm: int | None
    crease_middle_mm: int | None
    crease_right_mm: int | None
    notes: str | None

    def as_response(self, *, match_score: float) -> dict[str, Any]:
        payload = asdict(self)
        payload["record_date"] = self.record_date.isoformat() if self.record_date else None
        payload["board_length_cm"] = str(self.board_length_cm)
        payload["board_width_cm"] = str(self.board_width_cm)
        payload["source_ref"] = f"{DEFAULT_SHEET_NAME}!{self.source_row}"
        payload["match_score"] = round(match_score, 4)
        return payload


def normalize_lookup_text(value: Any) -> str:
    text = str(value or "").strip().upper()
    text = text.replace("×", "*").replace("X", "*")
    return re.sub(r"[^0-9A-Z\u4e00-\u9fff]+", "", text)


def _text(value: Any) -> str:
    return str(value).strip() if value not in (None, "") else ""


def _decimal(value: str) -> Decimal:
    try:
        return Decimal(value)
    except (InvalidOperation, TypeError) as error:
        raise HistoricalPurchaseLookupError(f"无法解析历史采购尺寸：{value}") from error


def _cm_to_mm(value: Decimal) -> int:
    return int((value * Decimal("10")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def _quantity(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    try:
        numeric = Decimal(str(value))
    except InvalidOperation:
        return None
    if numeric < 0 or numeric != numeric.to_integral_value():
        return None
    return int(numeric)


def _excel_date(value: Any, *, epoch: datetime, forced_year: int | None = None) -> date | None:
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    elif isinstance(value, (int, float)) and 30_000 <= float(value) <= 80_000:
        from openpyxl.utils.datetime import from_excel

        converted = from_excel(value, epoch)
        parsed = converted.date() if isinstance(converted, datetime) else converted
    else:
        text = _text(value)
        parsed = None
        for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                parsed = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                continue
        if parsed is None:
            return None
    if forced_year is not None:
        try:
            parsed = parsed.replace(year=forced_year)
        except ValueError:
            return None
    return parsed


def _crease(value: Any) -> tuple[str | None, str | None, int | None, int | None, int | None]:
    text = _text(value).replace("×", "*").replace("X", "*")
    if not text:
        return None, None, None, None, None
    match = CREASE_RE.fullmatch(text)
    if match:
        parts = [_cm_to_mm(_decimal(part)) for part in match.groups()]
        return text, "压线", parts[0], parts[1], parts[2]
    if "毛" in text:
        return text, "毛片", None, None, None
    if "净" in text:
        return text, "净料", None, None, None
    return text, "其他", None, None, None


def _record_date(values: list[Any], *, epoch: datetime, inherited: date | None) -> date | None:
    # Newer rows store the purchase date in column G. Older repeated blocks
    # inherit month/day from the batch header in columns D/E.
    own_date = _excel_date(values[6] if len(values) > 6 else None, epoch=epoch)
    if own_date:
        return own_date
    candidate = values[4] if len(values) > 4 else None
    if isinstance(candidate, (date, datetime)):
        return _excel_date(candidate, epoch=epoch)
    return inherited


def parse_historical_purchase_rows(
    rows: Iterable[Iterable[Any]],
    *,
    epoch: datetime,
) -> tuple[HistoricalPurchaseRecord, ...]:
    records: list[HistoricalPurchaseRecord] = []
    inherited_supplier: str | None = None
    inherited_date: date | None = None

    for row_number, source in enumerate(rows, 1):
        values = list(source)
        values.extend([None] * max(0, 24 - len(values)))

        year_value = values[3]
        batch_date = values[4]
        if (
            isinstance(year_value, (int, float))
            and 2018 <= int(year_value) <= 2035
            and _excel_date(batch_date, epoch=epoch) is not None
        ):
            inherited_supplier = _text(values[0]) or inherited_supplier
            inherited_date = _excel_date(
                batch_date,
                epoch=epoch,
                forced_year=int(year_value),
            )
            continue

        board_match = PAIR_RE.fullmatch(_text(values[0]).replace("×", "*").replace("X", "*"))
        if not board_match:
            continue
        board_length_cm = _decimal(board_match.group(1))
        board_width_cm = _decimal(board_match.group(2))
        material_code = _text(values[3]).upper()
        product_reference = _text(values[12])
        if not material_code or not product_reference:
            continue

        search_parts = [_text(value) for value in values[12:24] if _text(value)]
        search_text = " | ".join(dict.fromkeys(search_parts))
        if not normalize_lookup_text(search_text):
            continue

        crease_text, crease_type, left, middle, right = _crease(values[1])
        notes = " | ".join(
            dict.fromkeys(
                _text(values[index])
                for index in (4, 5, 7, 11, 13, 14, 15)
                if _text(values[index])
                and _text(values[index]) not in search_parts
            )
        )
        records.append(
            HistoricalPurchaseRecord(
                source_row=row_number,
                supplier_name=inherited_supplier,
                record_date=_record_date(values, epoch=epoch, inherited=inherited_date),
                product_reference=product_reference,
                search_text=search_text,
                material_code=material_code,
                historical_quantity=_quantity(values[2]),
                board_length_cm=board_length_cm,
                board_width_cm=board_width_cm,
                report_length_mm=_cm_to_mm(board_length_cm),
                report_width_mm=_cm_to_mm(board_width_cm),
                crease_text=crease_text,
                crease_type=crease_type,
                crease_left_mm=left,
                crease_middle_mm=middle,
                crease_right_mm=right,
                notes=notes or None,
            )
        )
    return tuple(records)


def _repair_invalid_workbook_styles(source: Path, temp_dir: Path) -> Path:
    """Clamp invalid Excel font-family values without modifying the source file."""
    repaired = temp_dir / f"{source.stem}_style_repaired.xlsx"
    with zipfile.ZipFile(source, "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    styles_name = "xl/styles.xml"
    if styles_name in members:
        styles = members[styles_name].decode("utf-8")

        def clamp_family(match: re.Match[str]) -> str:
            return f"{match.group(1)}{min(int(match.group(2)), 14)}{match.group(3)}"

        styles = re.sub(
            r'(<family\b[^>]*\bval=["\'])(\d+)(["\'][^>]*/>)',
            clamp_family,
            styles,
        )
        members[styles_name] = styles.encode("utf-8")
    with zipfile.ZipFile(repaired, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    return repaired


def _load_historical_purchase_workbook(
    workbook_path: Path,
    *,
    sheet_name: str,
) -> tuple[HistoricalPurchaseRecord, ...]:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise HistoricalPurchaseLookupError(
                f"历史采购工作簿缺少工作表：{sheet_name}"
            )
        sheet = workbook[sheet_name]
        return parse_historical_purchase_rows(
            sheet.iter_rows(values_only=True),
            epoch=workbook.epoch,
        )
    finally:
        workbook.close()


def load_historical_purchase_records(
    workbook_path: Path,
    *,
    sheet_name: str = DEFAULT_SHEET_NAME,
) -> tuple[HistoricalPurchaseRecord, ...]:
    try:
        return _load_historical_purchase_workbook(
            workbook_path,
            sheet_name=sheet_name,
        )
    except ValueError as error:
        if "stylesheet" not in str(error).lower() and "max value is 14" not in str(error).lower():
            raise
        with tempfile.TemporaryDirectory(prefix="tm-erp-history-xlsx-") as temp_name:
            repaired = _repair_invalid_workbook_styles(
                workbook_path,
                Path(temp_name),
            )
            return _load_historical_purchase_workbook(
                repaired,
                sheet_name=sheet_name,
            )


@lru_cache(maxsize=4)
def _cached_records(
    path_text: str,
    modified_ns: int,
    size: int,
    sheet_name: str,
) -> tuple[HistoricalPurchaseRecord, ...]:
    del modified_ns, size
    return load_historical_purchase_records(Path(path_text), sheet_name=sheet_name)


def resolve_historical_workbook_path(explicit: str | Path | None = None) -> Path:
    configured = explicit or os.getenv("ERP_HISTORICAL_PURCHASE_WORKBOOK", "").strip()
    candidates: list[Path] = []
    if configured:
        candidates.append(Path(configured).expanduser())
    for root in (PROJECT_ROOT, settings.database_path.parent.parent):
        candidates.extend(root / name for name in DEFAULT_WORKBOOK_NAMES)
    for candidate in candidates:
        try:
            resolved = candidate.resolve(strict=False)
        except OSError:
            resolved = candidate
        if resolved.is_file():
            return resolved
    checked = "、".join(str(path) for path in candidates)
    raise FileNotFoundError(f"未找到历史采购工作簿；已检查：{checked}")


def historical_purchase_records(
    workbook_path: str | Path | None = None,
    *,
    sheet_name: str = DEFAULT_SHEET_NAME,
) -> tuple[HistoricalPurchaseRecord, ...]:
    path = resolve_historical_workbook_path(workbook_path)
    stat = path.stat()
    return _cached_records(str(path), stat.st_mtime_ns, stat.st_size, sheet_name)


def _match_score(query: str, record: HistoricalPurchaseRecord) -> float | None:
    normalized_query = normalize_lookup_text(query)
    normalized_text = normalize_lookup_text(record.search_text)
    if not normalized_query or not normalized_text:
        return None
    if normalized_query == normalized_text:
        return 2_000.0
    if normalized_query in normalized_text:
        return 1_200.0 + min(len(normalized_query), 100) / 100

    tokens = [normalize_lookup_text(token) for token in re.split(r"\s+", query.strip())]
    tokens = [token for token in tokens if token]
    if tokens and all(token in normalized_text for token in tokens):
        return 1_000.0 + sum(len(token) for token in tokens) / 100

    compact_query = re.sub(r"\s+", "", query.strip().upper())
    if (
        len(normalized_query) >= 5
        and any(character.isdigit() for character in normalized_query)
        and not re.search(r"[\u4e00-\u9fff]", compact_query)
    ):
        # Codes and dimensions are intentionally substring-searchable, but a
        # low similarity score must not flood an exact code search with
        # unrelated historical rows.
        return None

    ratio = SequenceMatcher(None, normalized_query, normalized_text).ratio()
    minimum = 0.72 if len(normalized_query) <= 4 else 0.62
    return ratio * 800 if ratio >= minimum else None


def search_historical_purchases(
    query: str,
    *,
    limit: int = 20,
    workbook_path: str | Path | None = None,
    sheet_name: str = DEFAULT_SHEET_NAME,
) -> dict[str, Any]:
    normalized_query = normalize_lookup_text(query)
    if not normalized_query:
        return {
            "query": query,
            "items": [],
            "total_matches": 0,
            "indexed_records": 0,
            "source_workbook": None,
            "source_sheet": sheet_name,
        }
    path = resolve_historical_workbook_path(workbook_path)
    records = historical_purchase_records(path, sheet_name=sheet_name)
    scored = [
        (score, record)
        for record in records
        if (score := _match_score(query, record)) is not None
    ]
    scored.sort(
        key=lambda row: (
            row[0],
            row[1].record_date or date.min,
            row[1].source_row,
        ),
        reverse=True,
    )
    bounded_limit = max(1, min(int(limit), 100))
    return {
        "query": query,
        "items": [
            record.as_response(match_score=score)
            for score, record in scored[:bounded_limit]
        ],
        "total_matches": len(scored),
        "indexed_records": len(records),
        "source_workbook": path.name,
        "source_sheet": sheet_name,
    }


def historical_purchase_entry_response(
    row: HistoricalPurchaseEntry,
) -> dict[str, Any]:
    return {
        "id": row.id,
        "source_row": row.source_row,
        "supplier_name": row.supplier_name,
        "record_date": row.record_date.isoformat() if row.record_date else None,
        "product_reference": row.product_reference,
        "search_text": row.search_text,
        "material_code": row.material_code,
        "historical_quantity": row.historical_quantity,
        "board_length_cm": str(Decimal(row.report_length_mm) / Decimal("10")),
        "board_width_cm": str(Decimal(row.report_width_mm) / Decimal("10")),
        "report_length_mm": row.report_length_mm,
        "report_width_mm": row.report_width_mm,
        "crease_text": row.crease_text,
        "crease_type": row.crease_type,
        "crease_left_mm": row.crease_left_mm,
        "crease_middle_mm": row.crease_middle_mm,
        "crease_right_mm": row.crease_right_mm,
        "notes": row.notes,
        "product_id": row.product_id,
        "customer_id": row.customer_id,
        "source_ref": f"{row.source_sheet}!{row.source_row}",
        "source_workbook": row.source_workbook,
        "source_sheet": row.source_sheet,
    }


def _historical_purchase_display_key(row: HistoricalPurchaseEntry) -> tuple[Any, ...]:
    """Collapse repeated purchases without deleting their source-row history."""
    return (
        row.product_id,
        row.customer_id,
        normalize_lookup_text(row.product_reference),
        normalize_lookup_text(row.material_code),
        row.report_length_mm,
        row.report_width_mm,
        row.crease_type,
        row.crease_left_mm,
        row.crease_middle_mm,
        row.crease_right_mm,
        normalize_lookup_text(row.crease_text),
        normalize_lookup_text(row.notes),
    )


def _historical_purchase_group_response(
    rows: list[HistoricalPurchaseEntry],
) -> dict[str, Any]:
    latest = rows[0]
    response = historical_purchase_entry_response(latest)
    response["history_count"] = len(rows)
    response["source_rows"] = [row.source_row for row in rows]
    response["quantity_history"] = [
        {
            "record_date": row.record_date.isoformat() if row.record_date else None,
            "quantity": row.historical_quantity,
        }
        for row in rows[:10]
    ]
    return response


def search_historical_purchase_database(
    db: Session,
    query: str,
    *,
    limit: int = 20,
) -> dict[str, Any]:
    """Search imported ERP history; runtime never depends on the source xlsx."""
    normalized_query = normalize_lookup_text(query)
    indexed_records = int(
        db.scalar(select(func.count(HistoricalPurchaseEntry.id))) or 0
    )
    if not normalized_query:
        return {
            "query": query,
            "items": [],
            "total_matches": 0,
            "indexed_records": indexed_records,
            "source_workbook": None,
            "source_sheet": DEFAULT_SHEET_NAME,
        }

    condition = HistoricalPurchaseEntry.normalized_search_text.contains(
        normalized_query
    )
    bounded_limit = max(1, min(int(limit), 100))
    matching_rows = list(
        db.scalars(
            select(HistoricalPurchaseEntry)
            .where(condition)
            .order_by(
                HistoricalPurchaseEntry.record_date.desc(),
                HistoricalPurchaseEntry.source_row.desc(),
            )
        )
    )
    grouped_rows: dict[tuple[Any, ...], list[HistoricalPurchaseEntry]] = {}
    for row in matching_rows:
        grouped_rows.setdefault(_historical_purchase_display_key(row), []).append(row)
    groups = list(grouped_rows.values())
    displayed_groups = groups[:bounded_limit]
    source_workbook = matching_rows[0].source_workbook if matching_rows else None
    source_sheet = matching_rows[0].source_sheet if matching_rows else DEFAULT_SHEET_NAME
    return {
        "query": query,
        "items": [
            _historical_purchase_group_response(group) for group in displayed_groups
        ],
        "total_matches": len(groups),
        "source_record_matches": len(matching_rows),
        "indexed_records": indexed_records,
        "source_workbook": source_workbook,
        "source_sheet": source_sheet,
    }
