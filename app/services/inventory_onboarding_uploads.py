from __future__ import annotations

import asyncio
import csv
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from hashlib import sha256
from io import BytesIO, StringIO
import json
import math
from pathlib import Path, PurePath, PurePosixPath
import re
from typing import Final, Literal
import unicodedata
import zipfile

from fastapi import UploadFile
from openpyxl import load_workbook

from app.services.secure_uploads import (
    CHUNK_SIZE,
    EXCEL_POLICY,
    UPLOAD_READ_TIMEOUT_SECONDS,
    UploadValidationError,
    ValidatedUpload,
    read_validated_upload,
    remove_stored_reference,
    store_private_upload,
)


InventoryOnboardingFormat = Literal["csv", "xlsx"]

PRIVATE_CATEGORY: Final = "inventory_onboarding"
PRIVATE_REFERENCE_PREFIX: Final = f"private:{PRIVATE_CATEGORY}/"
CSV_MIME_TYPES: Final = frozenset(
    {
        "text/csv",
        "application/csv",
        "text/plain",
        "application/vnd.ms-excel",
    }
)
MAX_UPLOAD_BYTES: Final = EXCEL_POLICY.max_bytes
MAX_ROWS: Final = 10_000
MAX_COLUMNS: Final = 200
MAX_CELL_CHARACTERS: Final = 20_000
MAX_RAW_ROW_CHARACTERS: Final = 100_000
MAX_XLSX_ENTRIES: Final = 1_024
MAX_XLSX_MEMBER_BYTES: Final = 24 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES: Final = 80 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO: Final = 150

_FORMULA_PREFIXES: Final = ("=", "+", "-", "@")
_FORMULA_XML_PATTERN: Final = re.compile(
    br"<(?:[A-Za-z_][A-Za-z0-9_.-]*:)?f(?:\s|/?>)",
    re.IGNORECASE,
)
_EXTERNAL_REL_PATTERN: Final = re.compile(
    br"TargetMode\s*=\s*[\"']External[\"']",
    re.IGNORECASE,
)
_UNSAFE_XLSX_PREFIXES: Final = (
    "xl/activex/",
    "xl/charts/",
    "xl/ctrlprops/",
    "xl/drawings/",
    "xl/embeddings/",
    "xl/externallinks/",
    "xl/media/",
    "xl/model/",
    "xl/oleobjects/",
    "xl/printersettings/",
    "customxml/",
)


@dataclass(frozen=True)
class InventoryOnboardingRawRow:
    sheet_name: str
    row_number: int
    raw_text: str
    original_values: tuple[str | int | float | bool | None, ...]


@dataclass(frozen=True)
class InventoryOnboardingRawSheet:
    name: str
    rows: tuple[InventoryOnboardingRawRow, ...]


@dataclass(frozen=True)
class StoredInventoryOnboardingUpload:
    private_reference: str
    private_path: Path
    filename: str
    content_type: str
    size: int
    sha256: str
    format: InventoryOnboardingFormat
    encoding: str | None
    sheets: tuple[InventoryOnboardingRawSheet, ...]

    @property
    def reference(self) -> str:
        return self.private_reference

    @property
    def path(self) -> Path:
        return self.private_path

    @property
    def rows(self) -> tuple[InventoryOnboardingRawRow, ...]:
        return tuple(row for sheet in self.sheets for row in sheet.rows)

    def cleanup(self) -> list[str]:
        """Remove the private payload and metadata after an upper-layer rollback."""

        return cleanup_inventory_onboarding_upload(self)


def _validated_filename(filename: str | None) -> tuple[str, str]:
    supplied = filename or ""
    if not supplied or supplied != supplied.strip():
        raise UploadValidationError("盘点文件名不能为空或包含首尾空格")
    if any(unicodedata.category(character) in {"Cc", "Cf"} for character in supplied):
        raise UploadValidationError("盘点文件名包含控制字符")
    normalized = supplied.replace("\\", "/")
    if "/" in normalized or ":" in normalized:
        raise UploadValidationError("盘点文件名不得包含路径")
    pure = PurePath(normalized)
    if pure.name != normalized or any(part in {"", ".", ".."} for part in pure.parts):
        raise UploadValidationError("盘点文件名不得包含路径")
    safe_name = pure.name
    if len(safe_name) > 255:
        raise UploadValidationError("盘点文件名过长")
    extension = Path(safe_name).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise UploadValidationError("盘点文件仅支持 .csv 或 .xlsx")
    return safe_name, extension


def _normalized_content_type(content_type: str | None) -> str:
    return (content_type or "").split(";", 1)[0].strip().lower()


async def _read_bounded_content(upload: UploadFile) -> bytes:
    chunks: list[bytes] = []
    total = 0
    while True:
        try:
            chunk = await asyncio.wait_for(
                upload.read(CHUNK_SIZE),
                timeout=UPLOAD_READ_TIMEOUT_SECONDS,
            )
        except TimeoutError as error:
            raise UploadValidationError("盘点文件上传读取超时") from error
        if not chunk:
            break
        total += len(chunk)
        if total > MAX_UPLOAD_BYTES:
            raise UploadValidationError(
                f"盘点文件不能超过 {MAX_UPLOAD_BYTES // (1024 * 1024)}MB"
            )
        chunks.append(chunk)
    content = b"".join(chunks)
    if not content:
        raise UploadValidationError("盘点文件不能为空")
    return content


def _decode_csv(content: bytes) -> tuple[str, str]:
    signature_probe = content[:512].lstrip(b"\xef\xbb\xbf\x00\t\r\n ")
    lowered_probe = signature_probe.lower()
    if (
        signature_probe.startswith(
            (
                b"PK\x03\x04",
                b"%PDF-",
                b"\x89PNG\r\n\x1a\n",
                b"\xff\xd8\xff",
                b"MZ",
                b"Rar!",
                b"\x1f\x8b",
            )
        )
        or lowered_probe.startswith(
            (
                b"<!doctype",
                b"<html",
                b"<script",
                b"<?xml",
                b"<svg",
            )
        )
    ):
        raise UploadValidationError("CSV 文件签名与扩展名不一致")
    attempts: tuple[tuple[str, str], ...]
    if content.startswith(b"\xef\xbb\xbf"):
        attempts = (("utf-8-sig", "utf-8-sig"),)
    else:
        attempts = (
            ("utf-8", "utf-8"),
            ("gb18030", "gb18030"),
        )
    for codec, reported_encoding in attempts:
        try:
            decoded = content.decode(codec, errors="strict")
        except UnicodeDecodeError:
            continue
        return decoded, reported_encoding
    raise UploadValidationError("CSV 必须是 UTF-8 BOM、UTF-8 或 GB18030 编码")


def _has_disallowed_control(value: str, *, allow_line_endings: bool) -> bool:
    for character in value:
        if allow_line_endings and character in {"\r", "\n"}:
            continue
        if unicodedata.category(character) in {"Cc", "Cf"}:
            return True
    return False


def _formula_like(value: str) -> bool:
    candidate = value.lstrip()
    return bool(candidate) and candidate.startswith(_FORMULA_PREFIXES)


def _validated_string(value: str, *, location: str) -> str:
    if len(value) > MAX_CELL_CHARACTERS:
        raise UploadValidationError(f"{location} 单元格内容过长")
    if _has_disallowed_control(value, allow_line_endings=False):
        raise UploadValidationError(f"{location} 包含 NUL 或控制字符")
    if _formula_like(value):
        raise UploadValidationError(f"{location} 包含公式或公式注入内容")
    return value


def _cell_value(
    value: object,
    *,
    location: str,
) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise UploadValidationError(f"{location} 包含非有限数字")
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, str):
        return _validated_string(value, location=location)
    raise UploadValidationError(f"{location} 包含不支持的单元格类型")


def _row_is_nonempty(values: tuple[object, ...]) -> bool:
    return any(
        value is not None and (not isinstance(value, str) or bool(value.strip()))
        for value in values
    )


def _parse_csv(
    content: bytes,
) -> tuple[str, tuple[InventoryOnboardingRawSheet, ...]]:
    text, encoding = _decode_csv(content)
    if _has_disallowed_control(text, allow_line_endings=True):
        raise UploadValidationError("CSV 包含 NUL 或控制字符")
    physical_lines = text.splitlines(keepends=True)
    reader = csv.reader(StringIO(text, newline=""), strict=True)
    rows: list[InventoryOnboardingRawRow] = []
    previous_end_line = 0
    try:
        for parsed in reader:
            end_line = reader.line_num
            start_line = previous_end_line + 1
            raw_text = "".join(physical_lines[previous_end_line:end_line]).rstrip(
                "\r\n"
            )
            previous_end_line = end_line
            if len(parsed) > MAX_COLUMNS:
                raise UploadValidationError(
                    f"CSV 第 {start_line} 行超过 {MAX_COLUMNS} 列"
                )
            values = tuple(
                _cell_value(
                    value,
                    location=f"CSV 第 {start_line} 行第 {index} 列",
                )
                for index, value in enumerate(parsed, start=1)
            )
            if not _row_is_nonempty(values):
                continue
            if len(raw_text) > MAX_RAW_ROW_CHARACTERS:
                raise UploadValidationError(f"CSV 第 {start_line} 行原文过长")
            rows.append(
                InventoryOnboardingRawRow(
                    sheet_name="CSV",
                    row_number=start_line,
                    raw_text=raw_text,
                    original_values=values,
                )
            )
            if len(rows) > MAX_ROWS:
                raise UploadValidationError(f"CSV 不能超过 {MAX_ROWS} 个非空行")
    except csv.Error as error:
        raise UploadValidationError(f"CSV 格式无效：{error}") from error
    if not rows:
        raise UploadValidationError("CSV 没有非空数据")
    if len(rows[0].original_values) < 2:
        raise UploadValidationError("CSV 至少需要两列，不能用纯文本伪装")
    return encoding, (InventoryOnboardingRawSheet(name="CSV", rows=tuple(rows)),)


def _safe_zip_member_name(name: str) -> str:
    if (
        not name
        or "\x00" in name
        or "\\" in name
        or name.startswith("/")
        or ":" in name
    ):
        raise UploadValidationError("XLSX 包含不安全的 ZIP 路径")
    path = PurePosixPath(name)
    if any(part in {"", ".", ".."} for part in path.parts):
        raise UploadValidationError("XLSX 包含路径穿越成员")
    return path.as_posix()


def _zip_ratio_too_high(*, compressed: int, uncompressed: int) -> bool:
    if uncompressed <= 1024 * 1024:
        return False
    if compressed <= 0:
        return uncompressed > 0
    return uncompressed / compressed > MAX_XLSX_COMPRESSION_RATIO


def _inspect_xlsx_xml(name: str, payload: bytes) -> None:
    lowered = payload.lower()
    if b"<!doctype" in lowered or b"<!entity" in lowered:
        raise UploadValidationError("XLSX XML 包含 DTD 或实体声明")
    if name.endswith(".rels") and _EXTERNAL_REL_PATTERN.search(payload):
        raise UploadValidationError("XLSX 包含外部链接")
    if (
        name.startswith("xl/worksheets/")
        and name.endswith(".xml")
        and _FORMULA_XML_PATTERN.search(payload)
    ):
        raise UploadValidationError("XLSX 包含公式")
    if name == "[content_types].xml" and (
        b"macroenabled" in lowered or b"vbaproject" in lowered
    ):
        raise UploadValidationError("XLSX 包含宏")


def _preflight_xlsx_container(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ENTRIES:
                raise UploadValidationError("XLSX ZIP 成员过多")
            seen_names: set[str] = set()
            total_compressed = 0
            total_uncompressed = 0
            normalized_infos: list[tuple[str, zipfile.ZipInfo]] = []
            for info in infos:
                name = _safe_zip_member_name(info.filename)
                folded = name.casefold()
                if folded in seen_names:
                    raise UploadValidationError("XLSX ZIP 包含重复成员")
                seen_names.add(folded)
                if info.flag_bits & 0x1:
                    raise UploadValidationError("XLSX 不允许加密 ZIP 成员")
                if info.compress_type not in {
                    zipfile.ZIP_STORED,
                    zipfile.ZIP_DEFLATED,
                }:
                    raise UploadValidationError("XLSX 使用了不允许的 ZIP 压缩算法")
                if info.file_size > MAX_XLSX_MEMBER_BYTES:
                    raise UploadValidationError("XLSX 单个 ZIP 成员解压后过大")
                if _zip_ratio_too_high(
                    compressed=info.compress_size,
                    uncompressed=info.file_size,
                ):
                    raise UploadValidationError("XLSX ZIP 压缩比异常，疑似压缩炸弹")
                total_compressed += info.compress_size
                total_uncompressed += info.file_size
                normalized_infos.append((folded, info))
            if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise UploadValidationError("XLSX 解压后总体积过大")
            if total_uncompressed > 5 * 1024 * 1024 and _zip_ratio_too_high(
                compressed=total_compressed,
                uncompressed=total_uncompressed,
            ):
                raise UploadValidationError("XLSX 总压缩比异常，疑似压缩炸弹")
            required = {"[content_types].xml", "xl/workbook.xml"}
            if not required.issubset(seen_names):
                raise UploadValidationError("XLSX 缺少 Office 工作簿结构")
            for name, info in normalized_infos:
                if (
                    "vbaproject" in name
                    or name.endswith(".bin")
                    or name.startswith(_UNSAFE_XLSX_PREFIXES)
                ):
                    raise UploadValidationError(
                        "XLSX 不允许宏、对象、图片、图表或外链资源"
                    )
                if name.endswith((".xml", ".rels")):
                    _inspect_xlsx_xml(name, archive.read(info))
    except UploadValidationError:
        raise
    except (OSError, RuntimeError, zipfile.BadZipFile, zipfile.LargeZipFile) as error:
        raise UploadValidationError("XLSX ZIP 结构无效") from error


def _parse_xlsx(
    content: bytes,
) -> tuple[InventoryOnboardingRawSheet, ...]:
    _preflight_xlsx_container(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise UploadValidationError("XLSX 工作簿无法安全读取") from error
    sheets: list[InventoryOnboardingRawSheet] = []
    nonempty_sheet_count = 0
    total_rows = 0
    try:
        for worksheet in workbook.worksheets:
            max_row = int(worksheet.max_row or 0)
            max_column = int(worksheet.max_column or 0)
            if max_row > MAX_ROWS:
                raise UploadValidationError(
                    f"XLSX 工作表“{worksheet.title}”超过 {MAX_ROWS} 行"
                )
            if max_column > MAX_COLUMNS:
                raise UploadValidationError(
                    f"XLSX 工作表“{worksheet.title}”超过 {MAX_COLUMNS} 列"
                )
            parsed_rows: list[InventoryOnboardingRawRow] = []
            for row_number, cells in enumerate(worksheet.iter_rows(), start=1):
                converted: list[str | int | float | bool | None] = []
                for column_number, cell in enumerate(cells, start=1):
                    if cell.data_type == "f":
                        raise UploadValidationError(
                            f"XLSX 工作表“{worksheet.title}”"
                            f"第 {row_number} 行第 {column_number} 列包含公式"
                        )
                    if cell.data_type == "e":
                        raise UploadValidationError(
                            f"XLSX 工作表“{worksheet.title}”"
                            f"第 {row_number} 行第 {column_number} 列包含错误值"
                        )
                    converted.append(
                        _cell_value(
                            cell.value,
                            location=(
                                f"XLSX 工作表“{worksheet.title}”"
                                f"第 {row_number} 行第 {column_number} 列"
                            ),
                        )
                    )
                values = tuple(converted)
                if not _row_is_nonempty(values):
                    continue
                raw_text = json.dumps(
                    values,
                    ensure_ascii=False,
                    separators=(",", ":"),
                )
                if len(raw_text) > MAX_RAW_ROW_CHARACTERS:
                    raise UploadValidationError(
                        f"XLSX 工作表“{worksheet.title}”第 {row_number} 行原文过长"
                    )
                parsed_rows.append(
                    InventoryOnboardingRawRow(
                        sheet_name=worksheet.title,
                        row_number=row_number,
                        raw_text=raw_text,
                        original_values=values,
                    )
                )
                total_rows += 1
                if total_rows > MAX_ROWS:
                    raise UploadValidationError(
                        f"XLSX 合计不能超过 {MAX_ROWS} 个非空行"
                    )
            if parsed_rows:
                nonempty_sheet_count += 1
            sheets.append(
                InventoryOnboardingRawSheet(
                    name=worksheet.title,
                    rows=tuple(parsed_rows),
                )
            )
    finally:
        workbook.close()
    if nonempty_sheet_count == 0:
        raise UploadValidationError("XLSX 没有非空工作表")
    if nonempty_sheet_count > 1:
        raise UploadValidationError("XLSX 只允许一个非空工作表")
    return tuple(sheets)


def cleanup_inventory_onboarding_upload(
    upload_or_reference: StoredInventoryOnboardingUpload | str,
) -> list[str]:
    reference = (
        upload_or_reference.private_reference
        if isinstance(upload_or_reference, StoredInventoryOnboardingUpload)
        else upload_or_reference
    )
    if not reference.startswith(PRIVATE_REFERENCE_PREFIX):
        return ["不是 N081 盘点私有文件引用，拒绝清理"]
    return remove_stored_reference(reference)


async def store_and_parse_inventory_onboarding_upload(
    upload: UploadFile,
) -> StoredInventoryOnboardingUpload:
    """Validate, parse and privately store one N081 onboarding source file.

    Parsing happens before storage, so rejected content leaves no private file.
    The returned object exposes ``cleanup()`` for an upper database transaction
    that fails after the file has been accepted.
    """

    filename, extension = _validated_filename(upload.filename)
    supplied_type = _normalized_content_type(upload.content_type)
    if extension == ".xlsx":
        validated = await read_validated_upload(upload, EXCEL_POLICY)
        sheets = _parse_xlsx(validated.content)
        upload_format: InventoryOnboardingFormat = "xlsx"
        encoding = None
    else:
        if supplied_type not in CSV_MIME_TYPES:
            raise UploadValidationError("CSV MIME 类型不允许")
        content = await _read_bounded_content(upload)
        encoding, sheets = _parse_csv(content)
        validated = ValidatedUpload(
            content=content,
            original_filename=filename,
            extension=".csv",
            content_type="text/csv",
            size=len(content),
            sha256=sha256(content).hexdigest(),
        )
        upload_format = "csv"

    stored = store_private_upload(validated, category=PRIVATE_CATEGORY)
    try:
        return StoredInventoryOnboardingUpload(
            private_reference=stored.reference,
            private_path=stored.path,
            filename=stored.original_filename,
            content_type=stored.content_type,
            size=stored.size,
            sha256=stored.sha256,
            format=upload_format,
            encoding=encoding,
            sheets=sheets,
        )
    except Exception:
        remove_stored_reference(stored.reference)
        raise
