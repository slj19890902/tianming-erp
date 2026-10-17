"""FIN-001 tax-bureau invoice item-template export helpers.

The service deliberately accepts an explicit, hash-verified empty template.  It
never mutates that source file: every export starts with a byte-for-byte copy so
the tax-bureau helper worksheets and validations remain intact.
"""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
import hashlib
import re
import json
import os
from pathlib import Path
import shutil
from typing import Any, Mapping, Sequence
from xml.etree import ElementTree
import zipfile

from openpyxl import load_workbook


TAX_TEMPLATE_SHA256 = "7FA373E248E48FCAA00B9732606F8B653E53EEC86B6D2A49A0A69D070FB78C73"
TAX_TEMPLATE_SHEETS = ("1-明细模板", "excelVersion", "xzqhdm", "2-特定业务信息")
DETAIL_SHEET = TAX_TEMPLATE_SHEETS[0]
FIRST_DETAIL_ROW = 4
WARNING_LINE_COUNT = 200
BLOCKING_LINE_COUNT = 2000


class InvoiceTaxTemplateError(ValueError):
    """Raised when an export cannot meet the tax-template gate."""


@dataclass(frozen=True)
class InvoiceTaxTemplateExport:
    path: Path
    sha256: str
    template_sha256: str
    line_count: int
    total_amount: Decimal
    warnings: tuple[str, ...]


def sha256_file(path: str | Path) -> str:
    """Return an upper-case SHA-256 without loading a whole file into memory."""

    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def verify_invoice_tax_template(
    template_path: str | Path,
    *,
    expected_sha256: str = TAX_TEMPLATE_SHA256,
) -> str:
    """Verify the supplied empty tax-bureau template before it is copied."""

    template = Path(template_path)
    if not template.is_file():
        raise InvoiceTaxTemplateError("税局空模板不存在")
    actual_sha256 = sha256_file(template)
    if actual_sha256 != expected_sha256.upper():
        raise InvoiceTaxTemplateError("税局空模板 SHA-256 不匹配，已拒绝导出")
    _verify_workbook_layout(template)
    return actual_sha256


def _verify_workbook_layout(path: Path) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if tuple(workbook.sheetnames) != TAX_TEMPLATE_SHEETS:
                raise InvoiceTaxTemplateError("税局空模板工作表结构不符合已验收版本")
        finally:
            workbook.close()
    except InvoiceTaxTemplateError:
        raise
    except Exception as error:  # pragma: no cover - openpyxl error text is version-specific
        raise InvoiceTaxTemplateError("税局空模板无法读取") from error


def _sheet_fingerprint(path: Path, sheet_names: Sequence[str]) -> str:
    """Fingerprint helper worksheets so export validation can prove preservation."""

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        data: list[tuple[str, int, int, Any]] = []
        for name in sheet_names:
            sheet = workbook[name]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        data.append((name, cell.row, cell.column, cell.value))
    finally:
        workbook.close()
    encoded = json.dumps(data, ensure_ascii=False, default=str, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest().upper()


def _required_text(line: Mapping[str, Any], key: str, label: str) -> str:
    value = line.get(key)
    if value is None or not str(value).strip():
        raise InvoiceTaxTemplateError(f"第 {line.get('_line_number', '?')} 行缺少{label}")
    return str(value).strip()


def _decimal(line: Mapping[str, Any], key: str, label: str) -> Decimal:
    value = line.get(key)
    if value is None or str(value).strip() == "":
        raise InvoiceTaxTemplateError(f"第 {line.get('_line_number', '?')} 行缺少{label}")
    try:
        decimal_value = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise InvoiceTaxTemplateError(f"第 {line.get('_line_number', '?')} 行{label}格式无效") from error
    if not decimal_value.is_finite() or decimal_value < 0:
        raise InvoiceTaxTemplateError(f"第 {line.get('_line_number', '?')} 行{label}必须为非负数")
    return decimal_value


def _normalise_lines(lines: Sequence[Mapping[str, Any]]) -> tuple[list[dict[str, Any]], Decimal]:
    if not lines:
        raise InvoiceTaxTemplateError("开票明细不能为空")
    if len(lines) >= BLOCKING_LINE_COUNT:
        raise InvoiceTaxTemplateError("明细行数达到 2000 行，必须拆分后再导出")

    normalised: list[dict[str, Any]] = []
    total_amount = Decimal("0")
    for position, raw_line in enumerate(lines, start=1):
        line = dict(raw_line)
        line["_line_number"] = position
        amount = _decimal(line, "amount", "金额")
        tax_rate = _decimal(line, "tax_rate", "税率")
        if amount != amount.quantize(Decimal("0.01")):
            raise InvoiceTaxTemplateError("税局导入金额必须实际保留两位，不能只设置显示格式")
        normalised.append(
            {
                "project_name": _required_text(line, "project_name", "项目名称"),
                "tax_classification_code": _required_text(
                    line, "tax_classification_code", "税收分类编码"
                ),
                "specification": line.get("specification"),
                "unit": line.get("unit"),
                "quantity": line.get("quantity"),
                "unit_price": line.get("unit_price"),
                "amount": amount,
                "tax_rate": tax_rate,
                "discount_amount": line.get("discount_amount"),
                "preferential_policy_type": line.get("preferential_policy_type"),
                "coal_type": line.get("coal_type"),
                "professional_service_invoice": line.get("professional_service_invoice"),
                "professional_service_agreement_no": line.get(
                    "professional_service_agreement_no"
                ),
            }
        )
        total_amount += amount
    return normalised, total_amount


_SPREADSHEET_NAMESPACE = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"


def _qualified(name: str) -> str:
    return f"{{{_SPREADSHEET_NAMESPACE}}}{name}"


def _detail_sheet_archive_name(archive: zipfile.ZipFile) -> str:
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    relationship_id = next(
        (
            sheet.attrib[f"{{{_RELATIONSHIP_NAMESPACE}}}id"]
            for sheet in workbook.findall(f".//{_qualified('sheet')}")
            if sheet.attrib.get("name") == DETAIL_SHEET
        ),
        None,
    )
    if relationship_id is None:
        raise InvoiceTaxTemplateError("导出副本缺少明细工作表")
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    target = next(
        (
            relation.attrib.get("Target")
            for relation in relationships.findall(
                f"{{{_PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationship"
            )
            if relation.attrib.get("Id") == relationship_id
        ),
        None,
    )
    if not target:
        raise InvoiceTaxTemplateError("导出副本无法定位明细工作表")
    return f"xl/{target.lstrip('/')}".replace("xl/xl/", "xl/")


def _write_detail_rows(path: Path, lines: Sequence[Mapping[str, Any]]) -> None:
    """Patch native sheetData and shared strings; never reserialize the workbook.

    Tax importers may reject a spreadsheet that Excel can recover. Preserve the
    vendor's namespaces, style XML, relationships and helper parts byte for byte.
    Text uses the template's shared-string representation, not inlineStr.
    """
    from xml.sax.saxutils import escape
    fields = {"A": "project_name", "B": "tax_classification_code", "C": "specification",
              "D": "unit", "E": "quantity", "F": "unit_price", "G": "amount",
              "H": "tax_rate", "I": "discount_amount", "J": "preferential_policy_type",
              "K": "coal_type", "U": "professional_service_invoice", "V": "professional_service_agreement_no"}
    temporary = path.with_name(f".{path.name}.writing")
    try:
        with zipfile.ZipFile(path) as archive:
            name = _detail_sheet_archive_name(archive)
            sheet = archive.read(name).decode("utf-8")
            shared = archive.read("xl/sharedStrings.xml").decode("utf-8")
            if "<worksheet " not in sheet or "<sst " not in shared:
                raise InvoiceTaxTemplateError("模板不是已验证的原生结构，请更新税务模板")
            block = re.search(r"(<sheetData\b[^>]*>)(.*?)(</sheetData>)", sheet, re.S)
            if block is None:
                raise InvoiceTaxTemplateError("税务模板缺少明细区")
            header_rows = []
            for row in re.findall(r"<row\b[^>]*(?:/>|>.*?</row>)", block[2], re.S):
                number = re.search(r'\br="(\d+)"', row)
                if number and int(number[1]) < FIRST_DETAIL_ROW:
                    header_rows.append(row)
            root = ElementTree.fromstring(shared)
            strings = ["".join(node.itertext()) for node in root]
            lookup = {}
            for i, value in enumerate(strings):
                lookup.setdefault(value, i)
            additions = []
            rows = []
            for i, line in enumerate(lines, FIRST_DETAIL_ROW):
                cells = []
                for column, key in fields.items():
                    value = line.get(key)
                    if value is None or value == "":
                        continue
                    ref = f"{column}{i}"
                    if column in {"E", "F", "G", "H", "I"}:
                        number = Decimal(str(value))
                        if not number.is_finite():
                            raise InvoiceTaxTemplateError("导出数字无效")
                        cells.append(f'<c r="{ref}"><v>{format(number, "f")}</v></c>')
                    else:
                        value = str(value)
                        if value not in lookup:
                            lookup[value] = len(strings) + len(additions)
                            additions.append('<si><t xml:space="preserve">' + escape(value) + '</t></si>')
                        cells.append(f'<c r="{ref}" t="s"><v>{lookup[value]}</v></c>')
                rows.append(f'<row r="{i}">' + ''.join(cells) + '</row>')
            sheet = sheet[:block.start(2)] + ''.join(header_rows + rows) + sheet[block.end(2):]
            sheet = re.sub(r'(<dimension\b[^>]*\bref=")[^"]*',
                           lambda m: m[1] + f'A1:V{FIRST_DETAIL_ROW + len(lines) - 1}', sheet, count=1)
            shared = shared.replace('</sst>', ''.join(additions) + '</sst>')
            count = len(re.findall(r'<c\b[^>]*\bt="s"', sheet))
            for member in archive.namelist():
                if member.startswith('xl/worksheets/') and member.endswith('.xml') and member != name:
                    count += len(re.findall(rb'<(?:\w+:)?c\b[^>]*\bt="s"', archive.read(member)))
            shared = re.sub(r'(<sst\b[^>]*\bcount=")[^"]*', lambda m: m[1] + str(count), shared, count=1)
            shared = re.sub(r'(<sst\b[^>]*\buniqueCount=")[^"]*', lambda m: m[1] + str(len(strings) + len(additions)), shared, count=1)
            with zipfile.ZipFile(temporary, 'w') as out:
                for info in archive.infolist():
                    data = sheet.encode('utf-8') if info.filename == name else shared.encode('utf-8') if info.filename == 'xl/sharedStrings.xml' else archive.read(info.filename)
                    out.writestr(info, data)
        temporary.replace(path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _verify_export(path: Path, lines: Sequence[Mapping[str, Any]], helper_fingerprint: str) -> None:
    _verify_workbook_layout(path)
    helper_sheets = TAX_TEMPLATE_SHEETS[1:]
    if _sheet_fingerprint(path, helper_sheets) != helper_fingerprint:
        raise InvoiceTaxTemplateError("导出文件的税局辅助数据发生变化")
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[DETAIL_SHEET]
        for offset, line in enumerate(lines):
            row = FIRST_DETAIL_ROW + offset
            for column, key in (("A", "project_name"), ("B", "tax_classification_code")):
                if sheet[f"{column}{row}"].value != line[key]:
                    raise InvoiceTaxTemplateError("导出文件明细回读校验失败")
            for column, key in (("C", "specification"), ("D", "unit"), ("E", "quantity")):
                actual = sheet[f"{column}{row}"].value
                expected = line.get(key)
                if key == "quantity" and expected is not None:
                    if actual is None or Decimal(str(actual)) != Decimal(str(expected)):
                        raise InvoiceTaxTemplateError("导出文件数量回读校验失败")
                elif (actual or None) != (expected or None):
                    raise InvoiceTaxTemplateError("导出文件规格或单位回读校验失败")
            for column, key in (("G", "amount"), ("H", "tax_rate")):
                actual = sheet[f"{column}{row}"].value
                if actual is None or Decimal(str(actual)) != line[key]:
                    raise InvoiceTaxTemplateError("导出文件金额或税率回读校验失败")
            if line.get("unit_price") is not None:
                actual_price = sheet[f"F{row}"].value
                if (
                    actual_price is None
                    or Decimal(str(actual_price)) != Decimal(str(line["unit_price"]))
                ):
                    raise InvoiceTaxTemplateError("导出文件未税单价回读校验失败")
    finally:
        workbook.close()


def generate_invoice_tax_template(
    *,
    template_path: str | Path,
    output_path: str | Path,
    lines: Sequence[Mapping[str, Any]],
    expected_template_sha256: str = TAX_TEMPLATE_SHA256,
) -> InvoiceTaxTemplateExport:
    """Copy the approved template, write detail rows, then reopen and validate it."""

    normalised_lines, total_amount = _normalise_lines(lines)
    source = Path(template_path)
    source_sha256 = verify_invoice_tax_template(
        source, expected_sha256=expected_template_sha256
    )
    helper_fingerprint = _sheet_fingerprint(source, TAX_TEMPLATE_SHEETS[1:])
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    # ``Path.resolve`` can fail for a reachable NAS mapped drive on Windows
    # (WinError 1005).  This comparison only guards against self-overwrite, so
    # lexical absolute paths are both sufficient and safer here.
    if os.path.normcase(os.path.abspath(destination)) == os.path.normcase(
        os.path.abspath(source)
    ):
        raise InvoiceTaxTemplateError("导出文件不能覆盖税局空模板")

    shutil.copyfile(source, destination)
    try:
        _write_detail_rows(destination, normalised_lines)
        _verify_export(destination, normalised_lines, helper_fingerprint)
        if sha256_file(source) != source_sha256:
            raise InvoiceTaxTemplateError("税局空模板在导出期间发生变化")
    except Exception:
        destination.unlink(missing_ok=True)
        raise

    warnings: tuple[str, ...] = ()
    if len(normalised_lines) > WARNING_LINE_COUNT:
        warnings = ("明细超过 200 行，请在导入税局前人工复核拆分需求。",)
    return InvoiceTaxTemplateExport(
        path=destination,
        sha256=sha256_file(destination),
        template_sha256=source_sha256,
        line_count=len(normalised_lines),
        total_amount=total_amount,
        warnings=warnings,
    )
