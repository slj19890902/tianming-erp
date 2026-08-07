from __future__ import annotations

from collections import OrderedDict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePath, PurePosixPath
import re
import secrets
from threading import RLock
import time
from typing import Any, Final, Iterable
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, UnidentifiedImageError
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models.customer import Customer
from app.models.product import Product
from app.models.user import User
from app.services.secure_uploads import UploadValidationError, discard_temporary_token


TEMPLATE_VERSION = "P1-22-v2"
INFO_SHEET = "导入信息"
GUIDE_SHEET = "使用说明"
PRODUCT_SHEET = "样品录入"
MOLD_SHEET = "模具档案"
EXPECTED_SHEETS = {
    INFO_SHEET,
    GUIDE_SHEET,
    PRODUCT_SHEET,
    MOLD_SHEET,
}
PRODUCT_HEADERS = (
    "样品号",
    "手写型号*",
    "尺寸(mm)",
    "楞型",
    "成型方式*",
    "印刷*",
    "印刷颜色",
    "结合方式*",
    "二次粘合*",
    "图纸文件名(多个用分号)",
    "模具编号",
    "现场备注",
    "操作*",
    "系统ID",
    "当前版本",
    "客户料号*",
    "产品名称*",
    "材质代码",
    "单位",
    "默认含税单价",
    "报料长(mm)",
    "报料宽(mm)",
    "启用",
)
PRODUCT_HEADER_ALIASES = {
    1: {"手写型号/ERP存货编码*", "手写型号*"},
}
MOLD_HEADERS = ("操作*", "模具编号*", "模具名称*", "固定位置*", "备注")
MAX_PRODUCT_ROWS = 2_000
MAX_MOLD_ROWS = 1_000
PREVIEW_TTL_SECONDS = 10 * 60
MAX_PREVIEWS = 32
DRAWING_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".pdf"}
SINGLE_PHOTO_CONFIRMED_REMARK = "图片要求=1张（已确认）"
NAIL_PROCESS_TOKENS = {"打钉", "钉箱", "打钉箱"}
GLUE_PROCESS_TOKENS = {"粘合", "粘贴", "粘箱", "糊箱"}
NO_PRINT_CONTENT = {"无印刷", "无", "否", "不印刷"}
MAX_XLSX_ENTRIES: Final = 1_024
MAX_XLSX_MEMBER_BYTES: Final = 24 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES: Final = 80 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO: Final = 150
MAX_EMBEDDED_DRAWINGS: Final = MAX_PRODUCT_ROWS * 2
MAX_EMBEDDED_IMAGE_DIMENSION: Final = 12_000
MAX_EMBEDDED_IMAGE_PIXELS: Final = 40_000_000
_FORMULA_XML_PATTERN: Final = re.compile(
    br"<(?:[A-Za-z_][A-Za-z0-9_.-]*:)?f(?:\s|/?>)",
    re.IGNORECASE,
)
_EXTERNAL_REL_PATTERN: Final = re.compile(
    br"TargetMode\s*=\s*[\"']External[\"']",
    re.IGNORECASE,
)
_DRAWING_MEMBER_PATTERN: Final = re.compile(
    r"^xl/drawings/drawing[0-9]+\.xml$",
    re.IGNORECASE,
)
_DRAWING_RELS_MEMBER_PATTERN: Final = re.compile(
    r"^xl/drawings/_rels/drawing[0-9]+\.xml\.rels$",
    re.IGNORECASE,
)
_MEDIA_MEMBER_PATTERN: Final = re.compile(
    r"^xl/media/[^/]+\.(?:jpe?g|png)$",
    re.IGNORECASE,
)
_UNSAFE_PRODUCT_XLSX_PREFIXES: Final = (
    "xl/activex/",
    "xl/charts/",
    "xl/comments",
    "xl/connections",
    "xl/ctrlprops/",
    "xl/embeddings/",
    "xl/externallinks/",
    "xl/model/",
    "xl/objects/",
    "xl/oleobjects/",
    "xl/persons/",
    "xl/pivot",
    "xl/printersettings/",
    "xl/querytables/",
    "xl/richdata/",
    "xl/slicer",
    "xl/threadedcomments/",
    "xl/webextensions/",
    "customxml/",
)


def single_photo_requirement_confirmed(value: object) -> bool:
    """Return true only for the explicit, standalone one-photo marker."""

    return any(
        part.strip() == SINGLE_PHOTO_CONFIRMED_REMARK
        for part in re.split(r"[;；]", str(value or ""))
    )


class ProductImportWorkbookError(ValueError):
    def __init__(self, code: str, message: str, *, status_code: int = 400) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


@dataclass(frozen=True)
class EmbeddedProductDrawing:
    row_number: int
    column_number: int
    column_offset: int
    content: bytes


@dataclass(frozen=True)
class ProductImportPreview:
    actor: str
    owner_id: int
    expires_at: float
    customer_id: int
    source_sha256: str
    source_files: tuple[dict[str, Any], ...]
    product_items: tuple[dict[str, Any], ...]
    mold_items: tuple[dict[str, Any], ...]
    drawing_tokens: dict[str, str]
    summary: dict[str, int]


def _safe_product_xlsx_member_name(name: str) -> str:
    if (
        not name
        or "\x00" in name
        or "\\" in name
        or name.startswith("/")
        or ":" in name
    ):
        raise UploadValidationError("XLSX 包含不安全的 ZIP 路径")
    path = PurePosixPath(name)
    if any(part in {"", ".", ".."} for part in path.parts):
        raise UploadValidationError("XLSX 包含路径穿越成员")
    return path.as_posix()


def _product_xlsx_ratio_too_high(*, compressed: int, uncompressed: int) -> bool:
    if uncompressed <= 1024 * 1024:
        return False
    if compressed <= 0:
        return uncompressed > 0
    return uncompressed / compressed > MAX_XLSX_COMPRESSION_RATIO


def _inspect_product_xlsx_xml(name: str, payload: bytes) -> None:
    lowered = payload.lower()
    if b"<!doctype" in lowered or b"<!entity" in lowered:
        raise UploadValidationError("XLSX XML 包含 DTD 或实体声明")
    if name.endswith(".rels") and _EXTERNAL_REL_PATTERN.search(payload):
        raise UploadValidationError("XLSX 包含外部链接")
    if name.endswith(".rels") and re.search(
        br"Type\s*=\s*[\"'][^\"']*/(?:activeX|chart|control|externalLink|"
        br"oleObject|package|queryTable|pivotTable|webextension)[\"']",
        payload,
        re.IGNORECASE,
    ):
        raise UploadValidationError("XLSX 包含不允许的对象关系")
    if (
        name.startswith("xl/worksheets/")
        and name.endswith(".xml")
        and _FORMULA_XML_PATTERN.search(payload)
    ):
        raise UploadValidationError("XLSX 包含公式")
    if name == "[content_types].xml" and (
        b"macroenabled" in lowered or b"vbaproject" in lowered
    ):
        raise UploadValidationError("XLSX 包含宏")
    if _DRAWING_MEMBER_PATTERN.fullmatch(name):
        try:
            root = ElementTree.fromstring(payload)
        except ElementTree.ParseError as error:
            raise UploadValidationError("XLSX 图形层 XML 无效") from error
        local_names = {
            element.tag.rsplit("}", 1)[-1].casefold()
            for element in root.iter()
            if isinstance(element.tag, str)
        }
        if local_names.intersection(
            {"graphicframe", "sp", "cxnsp", "grpsp", "contentpart", "oleobj"}
        ):
            raise UploadValidationError("XLSX 图形层只允许普通图片")
        if "pic" not in local_names:
            raise UploadValidationError("XLSX 图形层缺少可识别图片")
    if _DRAWING_RELS_MEMBER_PATTERN.fullmatch(name):
        try:
            root = ElementTree.fromstring(payload)
        except ElementTree.ParseError as error:
            raise UploadValidationError("XLSX 图形层关系 XML 无效") from error
        relationships = [
            (
                str(element.attrib.get("Type", "")),
                str(element.attrib.get("Target", "")),
            )
            for element in root
            if element.tag.rsplit("}", 1)[-1].casefold() == "relationship"
        ]
        if not relationships:
            raise UploadValidationError("XLSX 图形层关系无效")
        for relation_type, target in relationships:
            normalized_target = target.casefold()
            if (
                not relation_type.casefold().endswith("/image")
                or not re.fullmatch(
                    r"(?:\.\./media/|/xl/media/)[^/]+\.(?:jpe?g|png)",
                    normalized_target,
                )
            ):
                raise UploadValidationError("XLSX 图形层只允许引用内嵌 JPG/PNG 图片")


def _preflight_product_xlsx_container(content: bytes) -> None:
    """Validate the product workbook without weakening the shared XLSX gate.

    This workflow deliberately permits only worksheet picture drawings backed
    by embedded JPG/JPEG/PNG media.  Active content, arbitrary drawing objects,
    external relationships and high-risk Office package members remain blocked.
    """

    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ENTRIES:
                raise UploadValidationError("XLSX ZIP 成员过多")
            seen_names: set[str] = set()
            total_compressed = 0
            total_uncompressed = 0
            normalized_infos: list[tuple[str, zipfile.ZipInfo]] = []
            for info in infos:
                is_directory = info.is_dir()
                name = _safe_product_xlsx_member_name(info.filename)
                folded = (
                    f"{name.casefold().rstrip('/')}/"
                    if is_directory
                    else name.casefold()
                )
                if folded in seen_names:
                    raise UploadValidationError("XLSX ZIP 包含重复成员")
                seen_names.add(folded)
                if info.flag_bits & 0x1:
                    raise UploadValidationError("XLSX 不允许加密 ZIP 成员")
                if info.compress_type not in {
                    zipfile.ZIP_STORED,
                    zipfile.ZIP_DEFLATED,
                }:
                    raise UploadValidationError("XLSX 使用了不允许的 ZIP 压缩算法")
                if info.file_size > MAX_XLSX_MEMBER_BYTES:
                    raise UploadValidationError("XLSX 单个 ZIP 成员解压后过大")
                if _product_xlsx_ratio_too_high(
                    compressed=info.compress_size,
                    uncompressed=info.file_size,
                ):
                    raise UploadValidationError("XLSX ZIP 压缩比异常，疑似压缩炸弹")
                total_compressed += info.compress_size
                total_uncompressed += info.file_size
                if is_directory:
                    if info.file_size != 0:
                        raise UploadValidationError("XLSX ZIP 目录成员包含异常数据")
                    continue
                normalized_infos.append((folded, info))
            if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise UploadValidationError("XLSX 解压后总体积过大")
            if total_uncompressed > 5 * 1024 * 1024 and _product_xlsx_ratio_too_high(
                compressed=total_compressed,
                uncompressed=total_uncompressed,
            ):
                raise UploadValidationError("XLSX 总压缩比异常，疑似压缩炸弹")
            required = {"[content_types].xml", "xl/workbook.xml"}
            if not required.issubset(seen_names):
                raise UploadValidationError("XLSX 缺少 Office 工作簿结构")
            embedded_media_count = sum(
                bool(_MEDIA_MEMBER_PATTERN.fullmatch(name))
                for name, _info in normalized_infos
            )
            if embedded_media_count > MAX_EMBEDDED_DRAWINGS:
                raise UploadValidationError(
                    f"Excel 内嵌图最多允许 {MAX_EMBEDDED_DRAWINGS} 张"
                )
            for name, info in normalized_infos:
                is_drawing = bool(_DRAWING_MEMBER_PATTERN.fullmatch(name))
                is_drawing_rels = bool(
                    _DRAWING_RELS_MEMBER_PATTERN.fullmatch(name)
                )
                is_media = name.startswith("xl/media/")
                if is_media and not _MEDIA_MEMBER_PATTERN.fullmatch(name):
                    raise UploadValidationError("XLSX 内嵌图只允许 JPG/JPEG/PNG")
                if name.startswith("xl/drawings/") and not (
                    is_drawing or is_drawing_rels
                ):
                    raise UploadValidationError("XLSX 包含不允许的图形层成员")
                if (
                    "vbaproject" in name
                    or name.endswith(".bin")
                    or name.startswith(_UNSAFE_PRODUCT_XLSX_PREFIXES)
                ):
                    raise UploadValidationError(
                        "XLSX 不允许宏、ActiveX、图表、OLE、对象或外链资源"
                    )
                if name.endswith((".xml", ".rels")):
                    _inspect_product_xlsx_xml(name, archive.read(info))
    except UploadValidationError:
        raise
    except (OSError, RuntimeError, zipfile.BadZipFile, zipfile.LargeZipFile) as error:
        raise UploadValidationError("XLSX ZIP 结构无效") from error


_PREVIEWS: OrderedDict[str, ProductImportPreview] = OrderedDict()
_PREVIEW_LOCK = RLock()


def _actor(user: User) -> str:
    return f"id:{user.id}" if user.id is not None else f"username:{user.username}"


def _discard_preview_files(preview: ProductImportPreview) -> None:
    for token in preview.drawing_tokens.values():
        try:
            discard_temporary_token(token, owner_id=preview.owner_id)
        except Exception:
            continue


def _purge_previews(now: float) -> None:
    expired = [token for token, item in _PREVIEWS.items() if item.expires_at <= now]
    for token in expired:
        item = _PREVIEWS.pop(token, None)
        if item is not None:
            _discard_preview_files(item)
    while len(_PREVIEWS) >= MAX_PREVIEWS:
        _token, item = _PREVIEWS.popitem(last=False)
        _discard_preview_files(item)


def store_preview(preview: ProductImportPreview) -> str:
    token = secrets.token_urlsafe(32)
    with _PREVIEW_LOCK:
        _purge_previews(time.monotonic())
        _PREVIEWS[token] = preview
    return token


def get_preview(token: str, user: User) -> ProductImportPreview:
    with _PREVIEW_LOCK:
        _purge_previews(time.monotonic())
        preview = _PREVIEWS.get(token)
        if preview is None or preview.actor != _actor(user):
            raise ProductImportWorkbookError(
                "PRODUCT_IMPORT_PREVIEW_STALE",
                "常用箱批量预检已过期或不属于当前操作人，请重新上传预检",
                status_code=409,
            )
        _PREVIEWS.move_to_end(token)
        return preview


def consume_preview(token: str, *, keep_files: bool = True) -> None:
    with _PREVIEW_LOCK:
        preview = _PREVIEWS.pop(token, None)
    if preview is not None and not keep_files:
        _discard_preview_files(preview)


def discard_preview(token: str) -> None:
    consume_preview(token, keep_files=False)


def _style_header(worksheet, row: int, *, fill: str = "2F75B5") -> None:
    for cell in worksheet[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
    worksheet.row_dimensions[row].height = 34


def _set_widths(worksheet, widths: Iterable[int]) -> None:
    for index, width in enumerate(widths, start=1):
        column = get_column_letter(index)
        worksheet.column_dimensions[column].width = width


def _as_export_number(value: object) -> object:
    if isinstance(value, Decimal):
        return float(value)
    return value


def _process_tokens(value: object) -> set[str]:
    return {
        item.strip()
        for item in re.split(r"[,，、;；]+", _cell_text(value))
        if item.strip()
    }


def product_printing_enabled(
    *,
    print_content: object,
    printing_colors: object,
    production_process: object,
) -> bool:
    process = _process_tokens(production_process)
    if "印刷" in process:
        return True
    content = _cell_text(print_content)
    if content:
        return content not in NO_PRINT_CONTENT
    return bool(_cell_text(printing_colors))


def project_workbook_managed_process_tokens(
    *,
    box_category: object,
    print_content: object,
    printing_colors: object,
    production_process: object,
) -> tuple[str, ...]:
    """Project rich/legacy process data into the v2 sheet's four selectors."""
    process = _process_tokens(production_process)
    items: list[str] = []
    if _cell_text(box_category) == "die_cut" or "模切" in process:
        items.append("模切")
    elif "开槽" in process:
        items.append("开槽")
    if product_printing_enabled(
        print_content=print_content,
        printing_colors=printing_colors,
        production_process=production_process,
    ):
        items.append("印刷")
    secondary_gluing = "二次粘合" in process
    if secondary_gluing and process.intersection(GLUE_PROCESS_TOKENS):
        items.append("粘合")
    elif process.intersection(NAIL_PROCESS_TOKENS):
        items.append("打钉")
    elif process.intersection(GLUE_PROCESS_TOKENS):
        items.append("粘合")
    if secondary_gluing:
        items.append("二次粘合")
    return tuple(items)


def _export_dimensions(product: Product) -> str:
    values: list[str] = []
    for value in (product.length_mm, product.width_mm, product.height_mm):
        if value is None:
            continue
        number = Decimal(str(value))
        if number == number.to_integral_value():
            values.append(str(int(number)))
        else:
            values.append(format(number.normalize(), "f"))
    return "×".join(values)


def _product_export_row(product: Product, *, sample_id: str) -> list[object]:
    material_code = ""
    if product.material is not None:
        supplier_name = (product.material.supplier_name or "").strip()
        material_code = (
            f"{supplier_name}|{product.material.code}"
            if supplier_name
            else product.material.code
        )
    mold_code = product.mold_tool.mold_code if product.mold_tool is not None else ""
    managed_process = project_workbook_managed_process_tokens(
        box_category=product.box_category,
        print_content=product.print_content,
        printing_colors=product.printing_colors,
        production_process=product.production_process,
    )
    if "模切" in managed_process:
        forming_method = "模切"
    elif "开槽" in managed_process:
        forming_method = "开槽"
    else:
        forming_method = "无需"
    printed = "印刷" in managed_process
    if "打钉" in managed_process:
        joining_method = "打钉"
    elif "粘合" in managed_process:
        joining_method = "粘合"
    else:
        joining_method = "无需结合"
    return [
        sample_id,
        product.product_code,
        _export_dimensions(product),
        product.flute_type or "",
        forming_method,
        "是" if printed else "否",
        (product.printing_colors or "黑色") if printed else "",
        joining_method,
        "是" if "二次粘合" in managed_process else "否",
        "",
        mold_code,
        product.remark or "",
        "自动",
        product.id,
        product.version,
        product.customer_material_code,
        product.product_name,
        material_code,
        product.unit,
        _as_export_number(product.sale_unit_price),
        product.report_length_mm,
        product.report_width_mm,
        "是" if product.is_active else "否",
    ]


def build_product_import_workbook(
    db: Session,
    *,
    customer: Customer,
    blank_rows: int = 40,
) -> bytes:
    products = list(
        db.scalars(
            select(Product)
            .where(
                Product.customer_id == customer.id,
                Product.deleted_at.is_(None),
            )
            .options(
                selectinload(Product.material),
                selectinload(Product.mold_tool),
                selectinload(Product.drawings),
            )
            .order_by(Product.product_code, Product.id)
        ).all()
    )
    workbook = Workbook()
    guide = workbook.active
    guide.title = GUIDE_SHEET
    guide.sheet_view.showGridLines = False
    guide.merge_cells("A1:G1")
    guide["A1"] = "天明 ERP｜样品常用箱批量导入导出模板"
    guide["A1"].font = Font(bold=True, color="FFFFFF", size=16)
    guide["A1"].fill = PatternFill("solid", fgColor="17365D")
    guide["A1"].alignment = Alignment(vertical="center")
    guide.row_dimensions[1].height = 30
    guide.merge_cells("A2:G2")
    guide["A2"] = f"客户：{customer.name}（系统ID {customer.id}）。本模板只能回导到该客户。"
    guide["A2"].fill = PatternFill("solid", fgColor="DDEBF7")
    guide["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    guide.append([])
    guide.append(["现场只填左侧", "在“样品录入”A:L 填现场事实；A4 横向默认只打印 A:I，J:L 的照片文件名、模具号和备注在电脑补，需纸面全列时使用 A3。"])
    guide.append(["型号规则", "手写清单型号直接成为 ERP 存货编码；客户料号仍由办公室明确填写，不能自动与手写型号相同。"])
    guide.append(["照片改名", "通常每件样品两张：YP001_实物.jpg、YP001_展开.jpg；确认只需一张时，在配对CSV的照片2写“无需”。照片里同时拍到样品号纸卡，插入前先压缩。"])
    guide.append(["办公室只补右侧", "M列操作可以留空由 ERP 自动识别：存货编码已存在时更新/补图，不存在时按新增处理。新产品仍必须明确客户料号和产品名称。"])
    guide.append(["新增最低必填", "手写型号、客户料号、产品名称，并选择成型方式、印刷、结合方式和二次粘合；不要让系统猜客户或产品名称。"])
    guide.append(["材质同码", "材质代码唯一时直接填代码；多个供应商有同码时填写“供应商|材质代码”，例如“胜源|G9G”。"])
    guide.append(["图纸与分卷", "把压缩后的 JPG/PNG 插入同一行 J 列并保留文件名。可一次选择多卷，ERP 会合并重复产品行和跨卷图片；通常要求2张，现场备注明确单图时允许1张。"])
    guide.append(["模具", "模切必须填模具编号；新模具还要在“模具档案”填写编号、名称、固定位置。"])
    guide.append(["预检", "任何阻断错误时整批不写。预检通过后由管理员一次确认。"])
    guide.append(["禁止", "不要用相近尺寸、模糊照片或最低报价自动决定型号、材质、供应商或模具。"])
    for row in range(4, 14):
        guide.cell(row=row, column=1).font = Font(bold=True, color="17365D")
        guide.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    _set_widths(guide, (18, 90, 4, 4, 4, 4, 4))

    product_sheet = workbook.create_sheet(PRODUCT_SHEET)
    product_sheet.sheet_view.showGridLines = False
    product_sheet.append(PRODUCT_HEADERS)
    _style_header(product_sheet, 1)
    for cell in product_sheet[1][12:]:
        cell.fill = PatternFill("solid", fgColor="548235")
    product_sheet.freeze_panes = "M2"
    for index, product in enumerate(products, start=1):
        product_sheet.append(
            _product_export_row(product, sample_id=f"YP{index:03d}")
        )
    first_blank_index = len(products) + 1
    for index in range(blank_rows):
        row = [""] * len(PRODUCT_HEADERS)
        sample_id = f"YP{first_blank_index + index:03d}"
        row[0] = sample_id
        row[6] = "黑色"
        row[8] = "否"
        row[9] = f"{sample_id}_实物.jpg；{sample_id}_展开.jpg"
        row[18] = "只"
        row[22] = "是"
        product_sheet.append(row)
    product_sheet.auto_filter.ref = f"A1:W{product_sheet.max_row}"
    action_validation = DataValidation(
        type="list",
        formula1='"自动,新增,更新,不变"',
        allow_blank=True,
    )
    flute_validation = DataValidation(
        type="list",
        formula1='"A,B,E,AB,BE,AAA,ABC,其他,不确定"',
        allow_blank=True,
    )
    forming_validation = DataValidation(
        type="list",
        formula1='"模切,开槽,无需,不确定"',
        allow_blank=True,
    )
    printed_validation = DataValidation(
        type="list",
        formula1='"是,否,不确定"',
        allow_blank=True,
    )
    color_validation = DataValidation(
        type="list",
        formula1='"黑色,红色,其他,无"',
        allow_blank=True,
    )
    joining_validation = DataValidation(
        type="list",
        formula1='"打钉,粘合,无需结合,不确定"',
        allow_blank=True,
    )
    yes_no_validation = DataValidation(
        type="list",
        formula1='"是,否"',
        allow_blank=True,
    )
    for validation in (
        action_validation,
        flute_validation,
        forming_validation,
        printed_validation,
        color_validation,
        joining_validation,
        yes_no_validation,
    ):
        product_sheet.add_data_validation(validation)
    flute_validation.add(f"D2:D{product_sheet.max_row}")
    forming_validation.add(f"E2:E{product_sheet.max_row}")
    printed_validation.add(f"F2:F{product_sheet.max_row}")
    color_validation.add(f"G2:G{product_sheet.max_row}")
    joining_validation.add(f"H2:H{product_sheet.max_row}")
    yes_no_validation.add(f"I2:I{product_sheet.max_row}")
    yes_no_validation.add(f"W2:W{product_sheet.max_row}")
    action_validation.add(f"M2:M{product_sheet.max_row}")
    _set_widths(
        product_sheet,
        (
            11, 22, 22, 10, 12, 10, 12, 13, 12, 48, 18, 30,
            10, 10, 10, 20, 24, 18, 9, 15, 13, 13, 10,
        ),
    )
    for row in product_sheet.iter_rows(min_row=2):
        product_sheet.row_dimensions[row[0].row].height = 88
        row[0].fill = PatternFill("solid", fgColor="E7E6E6")
        for cell in row[1:12]:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        row[9].alignment = Alignment(
            wrap_text=True,
            vertical="bottom",
            horizontal="center",
        )
        for cell in row[12:]:
            cell.fill = PatternFill("solid", fgColor="E2F0D9")
        for cell in (row[13], row[14]):
            cell.fill = PatternFill("solid", fgColor="E7E6E6")
    print_last_row = min(
        product_sheet.max_row,
        max(len(products) + 1, 23),
    )
    product_sheet.print_title_rows = "1:1"
    product_sheet.print_area = f"A1:I{print_last_row}"
    product_sheet.page_setup.orientation = "landscape"
    product_sheet.page_setup.fitToWidth = 1
    product_sheet.page_setup.fitToHeight = 0
    product_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    product_sheet.sheet_view.zoomScale = 85

    mold_sheet = workbook.create_sheet(MOLD_SHEET)
    mold_sheet.sheet_view.showGridLines = False
    mold_sheet.append(MOLD_HEADERS)
    _style_header(mold_sheet, 1)
    mold_sheet.freeze_panes = "A2"
    linked_molds: dict[int, Any] = {}
    for product in products:
        if product.mold_tool is not None:
            linked_molds[product.mold_tool.id] = product.mold_tool
    for mold in sorted(linked_molds.values(), key=lambda item: item.mold_code):
        mold_sheet.append(
            ["已有", mold.mold_code, mold.mold_name, mold.rack_location, mold.remarks or ""]
        )
    for _ in range(100):
        mold_sheet.append(["", "", "", "", ""])
    mold_action_validation = DataValidation(
        type="list",
        formula1='"新增,已有"',
        allow_blank=True,
    )
    mold_sheet.add_data_validation(mold_action_validation)
    mold_action_validation.add(f"A2:A{mold_sheet.max_row}")
    mold_sheet.auto_filter.ref = f"A1:E{mold_sheet.max_row}"
    _set_widths(mold_sheet, (12, 22, 32, 32, 42))

    info = workbook.create_sheet(INFO_SHEET)
    info.append(["template_version", TEMPLATE_VERSION])
    info.append(["customer_id", customer.id])
    info.append(["customer_name", customer.name])
    info.append(["generated_at", datetime.now().isoformat(timespec="seconds")])
    info.sheet_state = "hidden"

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _cell_text(value: object) -> str:
    return str(value or "").strip()


def _optional_int(
    value: object,
    label: str,
    *,
    positive: bool = False,
    allow_zero: bool = True,
) -> int | None:
    if value is None or _cell_text(value) == "":
        return None
    if isinstance(value, bool):
        raise ValueError(f"{label}必须是整数")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as error:
        raise ValueError(f"{label}必须是整数") from error
    if not number.is_finite() or number != number.to_integral_value():
        raise ValueError(f"{label}必须是整数")
    result = int(number)
    if positive and result <= 0:
        raise ValueError(f"{label}必须大于0")
    if not allow_zero and result == 0:
        raise ValueError(f"{label}不能为0")
    if result < 0:
        raise ValueError(f"{label}不能小于0")
    return result


def _optional_decimal(
    value: object,
    label: str,
    *,
    positive: bool = False,
) -> Decimal | None:
    if value is None or _cell_text(value) == "":
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError) as error:
        raise ValueError(f"{label}必须是数字") from error
    if not result.is_finite():
        raise ValueError(f"{label}必须是有限数字")
    if positive and result <= 0:
        raise ValueError(f"{label}必须大于0")
    if not positive and result < 0:
        raise ValueError(f"{label}不能小于0")
    return result


def _enabled(value: object) -> bool:
    if value is None or _cell_text(value) == "":
        return True
    text = _cell_text(value).lower()
    if text in {"是", "启用", "1", "true", "yes", "y"}:
        return True
    if text in {"否", "停用", "0", "false", "no", "n"}:
        return False
    raise ValueError("启用列只能填写是/否")


def _box_category(value: object) -> str:
    text = _cell_text(value).lower()
    if not text:
        return "normal"
    if text in {"normal", "普通", "普通箱"}:
        return "normal"
    if text in {"die_cut", "模切", "模切箱"}:
        return "die_cut"
    raise ValueError("箱类别只能填写 normal/普通 或 die_cut/模切")


def _crease_type(value: object) -> str | None:
    text = _cell_text(value)
    if not text:
        return None
    mapping = {"净": "净料", "毛": "毛片"}
    text = mapping.get(text, text)
    if text not in {"压线", "净料", "毛片", "其他"}:
        raise ValueError("压线类型只能填写压线、净料、毛片或其他")
    return text


_FLUTE_LAYERS = {
    "A": 3,
    "B": 3,
    "E": 3,
    "AB": 5,
    "BE": 5,
    "AAA": 7,
    "ABC": 7,
}


def _dimensions(value: object) -> tuple[int | None, int | None, int | None]:
    text = _cell_text(value)
    if not text:
        return None, None, None
    normalized = re.sub(r"(毫米|mm)", "", text, flags=re.IGNORECASE)
    normalized = re.sub(r"\s+", "", normalized)
    parts = re.split(r"[xX×*＊]", normalized)
    if len(parts) not in {2, 3} or any(not part.isdigit() for part in parts):
        raise ValueError("尺寸只能填写“长×宽”或“长×宽×高”，单位为 mm")
    numbers = [int(part) for part in parts]
    if any(number <= 0 for number in numbers):
        raise ValueError("尺寸中的数值必须大于0")
    if len(numbers) == 2:
        return numbers[0], numbers[1], None
    return numbers[0], numbers[1], numbers[2]


def _flute_and_layer(value: object) -> tuple[str | None, int | None]:
    text = _cell_text(value).upper()
    if not text or text in {"其他", "不确定"}:
        return None, None
    layer_count = _FLUTE_LAYERS.get(text)
    if layer_count is None:
        raise ValueError("楞型只能填写 A/B/E/AB/BE/AAA/ABC/其他/不确定")
    return text, layer_count


def _required_choice(
    value: object,
    label: str,
    allowed: set[str],
) -> str:
    text = _cell_text(value)
    if not text or text == "不确定":
        raise ValueError(f"{label}必须在正式导入前确认")
    if text not in allowed:
        raise ValueError(f"{label}只能填写{'/'.join(sorted(allowed))}")
    return text


def _yes_no(value: object, label: str) -> bool:
    text = _required_choice(value, label, {"是", "否"})
    return text == "是"


def _structured_process(
    *,
    forming_method: str,
    printed: bool,
    joining_method: str,
    secondary_gluing: bool,
) -> tuple[str, str, str | None]:
    if secondary_gluing and (
        forming_method != "模切" or joining_method != "粘合"
    ):
        raise ValueError("二次粘合=是时，成型方式必须为模切且结合方式必须为粘合")
    items: list[str] = []
    if forming_method in {"模切", "开槽"}:
        items.append(forming_method)
    if printed:
        items.append("印刷")
    if joining_method == "打钉":
        items.append("打钉")
    elif joining_method == "粘合":
        items.append("粘合")
    if secondary_gluing:
        items.append("二次粘合")
    box_category = "die_cut" if forming_method == "模切" else "normal"
    box_style = "模切内盒" if secondary_gluing else None
    return "、".join(items), box_category, box_style


def split_drawing_filenames(value: object) -> tuple[str, ...]:
    names: list[str] = []
    seen: set[str] = set()
    for raw in re.split(r"[;；\r\n]+", _cell_text(value)):
        name = raw.strip()
        if not name:
            continue
        pure = PurePath(name.replace("\\", "/"))
        if pure.name != name or "/" in name or "\\" in name or ":" in name:
            raise ValueError(f"图纸文件名不得包含路径：{name}")
        if pure.suffix.lower() not in DRAWING_EXTENSIONS:
            raise ValueError(f"图纸文件类型不允许：{name}")
        key = name.casefold()
        if key in seen:
            raise ValueError(f"图纸文件名在本单元格重复：{name}")
        seen.add(key)
        names.append(name)
    return tuple(names)


def _validate_headers(
    worksheet,
    expected: tuple[str, ...],
    *,
    aliases: dict[int, set[str]] | None = None,
) -> None:
    actual = tuple(_cell_text(cell.value) for cell in worksheet[1])
    invalid = len(actual) != len(expected) or any(
        value != expected[index]
        and value not in (aliases or {}).get(index, set())
        for index, value in enumerate(actual)
    )
    if invalid:
        raise ProductImportWorkbookError(
            "PRODUCT_IMPORT_HEADERS_INVALID",
            f"工作表“{worksheet.title}”表头不匹配，请重新下载系统模板",
        )


def _append_error(
    errors: list[dict[str, Any]],
    *,
    sheet: str,
    row_number: int,
    message: str,
) -> None:
    if len(errors) < 300:
        errors.append(
            {"sheet": sheet, "row_number": row_number, "message": message}
        )


def _parse_product_rows(worksheet, errors: list[dict[str, Any]]) -> list[dict]:
    rows: list[dict] = []
    for row_number, cells in enumerate(
        worksheet.iter_rows(min_row=2, max_col=len(PRODUCT_HEADERS)),
        start=2,
    ):
        values = [cell.value for cell in cells]
        action = _cell_text(values[12])
        meaningful_without_action = any(
            _cell_text(value)
            for index, value in enumerate(values)
            if index not in {0, 6, 8, 9, 12, 18, 22}
        )
        if not action and not meaningful_without_action:
            continue
        if not action:
            action = "自动"
        if action not in {"自动", "新增", "更新", "不变"}:
            _append_error(
                errors,
                sheet=PRODUCT_SHEET,
                row_number=row_number,
                message="操作只能填写自动、新增、更新或不变",
            )
            continue
        if action == "不变":
            continue
        try:
            system_id = _optional_int(values[13], "系统ID", positive=True)
            version = _optional_int(values[14], "当前版本", positive=True)
            product_code = _cell_text(values[1])
            customer_material_code = _cell_text(values[15])
            product_name = _cell_text(values[16])
            if not product_code:
                raise ValueError("手写型号不能为空；该字段直接作为 ERP 存货编码")
            if action != "自动" and not customer_material_code:
                raise ValueError("客户料号不能为空")
            if action != "自动" and not product_name:
                raise ValueError("产品名称不能为空")
            if action == "新增" and (system_id is not None or version is not None):
                raise ValueError("新增行的系统ID和当前版本必须留空")
            if action == "更新" and (system_id is None or version is None):
                raise ValueError("更新行必须保留系统ID和当前版本")
            if action == "自动" and ((system_id is None) != (version is None)):
                raise ValueError("自动识别行的系统ID和当前版本必须同时填写或同时留空")
            length_mm, width_mm, height_mm = _dimensions(values[2])
            flute_source = _cell_text(values[3])
            flute_type, layer_count = _flute_and_layer(flute_source)
            forming_method = _required_choice(
                values[4],
                "成型方式",
                {"模切", "开槽", "无需"},
            )
            printed = _yes_no(values[5], "印刷")
            printing_colors = _cell_text(values[6]) or "黑色"
            if printed and printing_colors == "无":
                raise ValueError("印刷=是时，印刷颜色不能填写无")
            if not printed:
                printing_colors = None
            joining_method = _required_choice(
                values[7],
                "结合方式",
                {"打钉", "粘合", "无需结合"},
            )
            secondary_gluing = _yes_no(values[8], "二次粘合")
            production_process, box_category, box_style = _structured_process(
                forming_method=forming_method,
                printed=printed,
                joining_method=joining_method,
                secondary_gluing=secondary_gluing,
            )
            rows.append(
                {
                    "row_number": row_number,
                    "action": action,
                    "system_id": system_id,
                    "version": version,
                    "sample_id": _cell_text(values[0]) or None,
                    "product_code": product_code,
                    "customer_material_code": customer_material_code,
                    "product_name": product_name,
                    "box_category": box_category,
                    "box_style": box_style,
                    "length_mm": length_mm,
                    "width_mm": width_mm,
                    "height_mm": height_mm,
                    "material_code": _cell_text(values[17]).upper() or None,
                    "layer_count": layer_count,
                    "flute_type": flute_type,
                    "flute_was_blank": not flute_source,
                    "unit": _cell_text(values[18]) or "只",
                    "sale_unit_price": _optional_decimal(
                        values[19],
                        "默认含税单价",
                    ),
                    "report_length_mm": _optional_int(
                        values[20],
                        "报料长(mm)",
                        positive=True,
                    ),
                    "report_width_mm": _optional_int(
                        values[21],
                        "报料宽(mm)",
                        positive=True,
                    ),
                    "crease_type": None,
                    "crease_left_mm": None,
                    "crease_middle_mm": None,
                    "crease_right_mm": None,
                    "production_process": production_process or None,
                    "managed_process_tokens": tuple(
                        item
                        for item in re.split(r"[,，、;；]+", production_process)
                        if item
                    ),
                    "is_printed": printed,
                    "print_content": "单色印刷" if printed else "无印刷",
                    "printing_colors": printing_colors,
                    "mold_code": _cell_text(values[10]).upper() or None,
                    "drawing_filenames": split_drawing_filenames(values[9]),
                    "report_notes": None,
                    "remark": _cell_text(values[11]) or None,
                    "is_active": _enabled(values[22]),
                }
            )
        except ValueError as error:
            _append_error(
                errors,
                sheet=PRODUCT_SHEET,
                row_number=row_number,
                message=str(error),
            )
    return rows


def _parse_mold_rows(worksheet, errors: list[dict[str, Any]]) -> list[dict]:
    rows: list[dict] = []
    seen_codes: set[str] = set()
    for row_number, cells in enumerate(
        worksheet.iter_rows(min_row=2, max_col=len(MOLD_HEADERS)),
        start=2,
    ):
        values = [cell.value for cell in cells]
        action = _cell_text(values[0])
        if not action and not any(_cell_text(value) for value in values[1:]):
            continue
        if action == "已有":
            continue
        try:
            if action != "新增":
                raise ValueError("模具操作只能填写新增或已有")
            code = _cell_text(values[1]).upper()
            name = _cell_text(values[2])
            location = _cell_text(values[3])
            if not code:
                raise ValueError("模具编号不能为空")
            if not name:
                raise ValueError("模具名称不能为空")
            if not location:
                raise ValueError("固定位置不能为空")
            if code in seen_codes:
                raise ValueError(f"模具编号 {code} 在表内重复")
            seen_codes.add(code)
            rows.append(
                {
                    "row_number": row_number,
                    "mold_code": code,
                    "mold_name": name,
                    "rack_location": location,
                    "remarks": _cell_text(values[4]) or None,
                }
            )
        except ValueError as error:
            _append_error(
                errors,
                sheet=MOLD_SHEET,
                row_number=row_number,
                message=str(error),
            )
    return rows


def _extract_embedded_product_drawings(
    workbook,
    errors: list[dict[str, Any]],
) -> dict[int, tuple[EmbeddedProductDrawing, ...]]:
    total_images = sum(
        len(tuple(getattr(worksheet, "_images", ())))
        for worksheet in workbook.worksheets
    )
    if total_images > MAX_EMBEDDED_DRAWINGS:
        raise ProductImportWorkbookError(
            "PRODUCT_IMPORT_DRAWINGS_EXCEEDED",
            f"Excel 内嵌图最多允许 {MAX_EMBEDDED_DRAWINGS} 张",
        )
    by_row: dict[int, list[EmbeddedProductDrawing]] = {}
    for worksheet in workbook.worksheets:
        images = tuple(getattr(worksheet, "_images", ()))
        for image in images:
            anchor = getattr(image, "anchor", None)
            marker = getattr(anchor, "_from", None)
            row_number = int(getattr(marker, "row", -1)) + 1
            column_number = int(getattr(marker, "col", -1)) + 1
            if worksheet.title != PRODUCT_SHEET:
                _append_error(
                    errors,
                    sheet=worksheet.title,
                    row_number=max(row_number, 1),
                    message="内嵌图只能放在“样品录入”工作表 J 列",
                )
                continue
            if marker is None or row_number < 2 or row_number > MAX_PRODUCT_ROWS + 1:
                _append_error(
                    errors,
                    sheet=PRODUCT_SHEET,
                    row_number=max(row_number, 1),
                    message="内嵌图起始锚点必须位于有效数据行",
                )
                continue
            if column_number != 10:
                _append_error(
                    errors,
                    sheet=PRODUCT_SHEET,
                    row_number=row_number,
                    message="内嵌图起始锚点必须放在本行 J 列",
                )
                continue
            try:
                payload = image._data()
                with Image.open(BytesIO(payload)) as source:
                    width, height = source.size
                    if (
                        width <= 0
                        or height <= 0
                        or width > MAX_EMBEDDED_IMAGE_DIMENSION
                        or height > MAX_EMBEDDED_IMAGE_DIMENSION
                        or width * height > MAX_EMBEDDED_IMAGE_PIXELS
                    ):
                        raise ValueError(
                            "内嵌图像素尺寸过大，请先压缩后重新插入"
                        )
                    if (source.format or "").upper() not in {"JPEG", "PNG"}:
                        raise ValueError("内嵌图只允许 JPG/JPEG/PNG")
                    source.load()
            except ValueError as error:
                _append_error(
                    errors,
                    sheet=PRODUCT_SHEET,
                    row_number=row_number,
                    message=str(error),
                )
                continue
            except (
                Image.DecompressionBombError,
                OSError,
                UnidentifiedImageError,
            ):
                _append_error(
                    errors,
                    sheet=PRODUCT_SHEET,
                    row_number=row_number,
                    message="内嵌图无法安全读取，请重新插入压缩后的 JPG/PNG",
                )
                continue
            by_row.setdefault(row_number, []).append(
                EmbeddedProductDrawing(
                    row_number=row_number,
                    column_number=column_number,
                    column_offset=int(getattr(marker, "colOff", 0) or 0),
                    content=payload,
                )
            )
    return {
        row_number: tuple(
            sorted(
                images,
                key=lambda image: (
                    image.column_number,
                    image.column_offset,
                ),
            )
        )
        for row_number, images in by_row.items()
    }


def read_product_import_workbook(
    content: bytes,
) -> tuple[
    int,
    list[dict],
    list[dict],
    list[dict],
    dict[int, tuple[EmbeddedProductDrawing, ...]],
]:
    _preflight_product_xlsx_container(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise ProductImportWorkbookError(
            "PRODUCT_IMPORT_WORKBOOK_INVALID",
            "Excel 工作簿无法安全读取",
        ) from error
    errors: list[dict[str, Any]] = []
    try:
        if set(workbook.sheetnames) != EXPECTED_SHEETS:
            raise ProductImportWorkbookError(
                "PRODUCT_IMPORT_SHEETS_INVALID",
                "Excel 工作表结构不正确，请重新下载系统模板",
            )
        info = workbook[INFO_SHEET]
        metadata = {
            _cell_text(info.cell(row=row, column=1).value): info.cell(
                row=row,
                column=2,
            ).value
            for row in range(1, info.max_row + 1)
        }
        if _cell_text(metadata.get("template_version")) != TEMPLATE_VERSION:
            raise ProductImportWorkbookError(
                "PRODUCT_IMPORT_TEMPLATE_VERSION_INVALID",
                "Excel 模板版本不受支持，请重新下载系统模板",
            )
        try:
            customer_id = int(metadata.get("customer_id"))
        except (TypeError, ValueError) as error:
            raise ProductImportWorkbookError(
                "PRODUCT_IMPORT_CUSTOMER_INVALID",
                "Excel 模板缺少有效客户身份",
            ) from error
        product_sheet = workbook[PRODUCT_SHEET]
        mold_sheet = workbook[MOLD_SHEET]
        if product_sheet.max_row > MAX_PRODUCT_ROWS + 1:
            raise ProductImportWorkbookError(
                "PRODUCT_IMPORT_ROWS_EXCEEDED",
                f"样品录入最多允许 {MAX_PRODUCT_ROWS} 行数据",
            )
        if mold_sheet.max_row > MAX_MOLD_ROWS + 1:
            raise ProductImportWorkbookError(
                "PRODUCT_IMPORT_MOLD_ROWS_EXCEEDED",
                f"模具档案最多允许 {MAX_MOLD_ROWS} 行数据",
            )
        _validate_headers(
            product_sheet,
            PRODUCT_HEADERS,
            aliases=PRODUCT_HEADER_ALIASES,
        )
        _validate_headers(mold_sheet, MOLD_HEADERS)
        product_rows = _parse_product_rows(product_sheet, errors)
        mold_rows = _parse_mold_rows(mold_sheet, errors)
        embedded_drawings = _extract_embedded_product_drawings(workbook, errors)
        actionable_rows = {row["row_number"] for row in product_rows}
        for row_number in embedded_drawings:
            if row_number not in actionable_rows:
                _append_error(
                    errors,
                    sheet=PRODUCT_SHEET,
                    row_number=row_number,
                    message="本行已有内嵌图，但没有可识别的产品资料",
                )
        if not product_rows and not errors:
            errors.append(
                {
                    "sheet": PRODUCT_SHEET,
                    "row_number": 2,
                    "message": "工作簿没有可识别的常用箱资料",
                }
            )
        return customer_id, product_rows, mold_rows, errors, embedded_drawings
    finally:
        workbook.close()
