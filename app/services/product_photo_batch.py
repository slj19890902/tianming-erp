from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime
import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata
from typing import Final, Iterable
import warnings
import zipfile

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import PatternFill
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image, ImageCms, ImageOps, UnidentifiedImageError

from app.services.product_import_workbook import (
    MOLD_SHEET,
    PRODUCT_SHEET,
    SINGLE_PHOTO_CONFIRMED_REMARK,
    _preflight_product_xlsx_container,
    single_photo_requirement_confirmed,
)
from app.services.secure_uploads import UploadValidationError

try:
    from pillow_heif import register_heif_opener
except ImportError:  # pragma: no cover - exercised by the CLI error path
    register_heif_opener = None
else:
    register_heif_opener()


PHOTO_EXTENSIONS: Final = {".heic", ".heif", ".jpg", ".jpeg", ".png"}
MAPPING_HEADERS: Final = ("样品号", "照片1", "照片2")
PRODUCT_HEADER_ROW: Final = 1
SAMPLE_ID_COLUMN: Final = 1
DRAWING_COLUMN: Final = 10
MOLD_CODE_COLUMN: Final = 11
ACTION_COLUMN: Final = 13
SYSTEM_ID_COLUMN: Final = 14
REMARK_COLUMN: Final = 12
MAX_SOURCE_PHOTO_BYTES: Final = 30 * 1024 * 1024
MAX_MAPPING_CSV_BYTES: Final = 2 * 1024 * 1024
MAX_SOURCE_DIMENSION: Final = 12_000
MAX_SOURCE_PIXELS: Final = 40_000_000
ERP_ABSOLUTE_XLSX_BYTES: Final = 20 * 1024 * 1024
DEFAULT_TARGET_XLSX_BYTES: Final = 16 * 1024 * 1024
DEFAULT_MAX_SAMPLES_PER_VOLUME: Final = 200
DEFAULT_MAX_EDGE: Final = 1_600
DEFAULT_JPEG_QUALITY: Final = 76
DISPLAY_SLOT_WIDTH_PX: Final = 148
DISPLAY_SLOT_HEIGHT_PX: Final = 100
DISPLAY_SLOT_GAP_PX: Final = 10
DISPLAY_CELL_LEFT_PX: Final = 5
DISPLAY_CELL_TOP_PX: Final = 5
DISPLAY_ROW_HEIGHT_POINTS: Final = 102
DISPLAY_COLUMN_WIDTH: Final = 48
SUPPORTED_DETECTED_FORMATS: Final = {
    ".heic": {"HEIF", "HEIC"},
    ".heif": {"HEIF", "HEIC"},
    ".jpg": {"JPEG"},
    ".jpeg": {"JPEG"},
    ".png": {"PNG"},
}
ROLE_SUFFIXES: Final = {
    0: ("图1", "照片1", "实物", "1"),
    1: ("图2", "照片2", "展开", "2"),
}
INVALID_LOGICAL_FILENAME: Final = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
NO_SECOND_PHOTO_TOKEN: Final = "无需"


@dataclass(frozen=True)
class PhotoBatchIssue:
    code: str
    message: str
    sample_id: str = ""
    filename: str = ""
    row_number: int | None = None


class ProductPhotoBatchError(ValueError):
    def __init__(
        self,
        message: str,
        *,
        issues: Iterable[PhotoBatchIssue] = (),
        report_path: Path | None = None,
    ) -> None:
        super().__init__(message)
        self.issues = tuple(issues)
        self.report_path = report_path


@dataclass(frozen=True)
class PhotoFile:
    path: Path
    filename: str
    normalized_filename: str
    size: int
    sha256: str


@dataclass(frozen=True)
class PreparedPhoto:
    source_name: str
    source_sha256: str
    source_bytes: int
    compressed_sha256: str
    compressed_bytes: int
    width: int
    height: int
    content: bytes
    from_workbook: bool = False


@dataclass(frozen=True)
class SampleRow:
    row_number: int
    sample_id: str
    normalized_sample_id: str
    action: str
    mold_code: str
    drawing_names: tuple[str, ...]
    remark: str


@dataclass(frozen=True)
class PreparedSample:
    row_number: int
    sample_id: str
    logical_names: tuple[str, ...]
    photos: tuple[PreparedPhoto, ...]
    mold_code: str

    @property
    def compressed_bytes(self) -> int:
        return sum(photo.compressed_bytes for photo in self.photos)


@dataclass(frozen=True)
class GeneratedVolume:
    volume_number: int
    volume_count: int
    path: Path
    size: int
    sha256: str
    sample_ids: tuple[str, ...]


@dataclass(frozen=True)
class ProductPhotoBatchResult:
    batch_id: str
    source_workbook: Path
    source_workbook_sha256: str
    photo_directory: Path
    volumes: tuple[GeneratedVolume, ...]
    report_csv: Path
    report_json: Path
    source_photo_bytes: int
    compressed_photo_bytes: int
    sample_count: int
    photo_count: int
    formal_import_status: str


def _cell_text(value: object) -> str:
    return str(value or "").strip()


def _remark_with_single_photo_requirement(value: object, *, enabled: bool) -> str:
    parts = [
        part.strip()
        for part in re.split(r"[;；]", _cell_text(value))
        if part.strip() and part.strip() != SINGLE_PHOTO_CONFIRMED_REMARK
    ]
    if enabled:
        parts.append(SINGLE_PHOTO_CONFIRMED_REMARK)
    return "；".join(parts)


def _normalized_key(value: str) -> str:
    return unicodedata.normalize("NFKC", value).strip().casefold()


def _normalized_filename(value: str) -> str:
    return _normalized_key(Path(value).name)


def _split_names(value: object) -> tuple[str, ...]:
    return tuple(
        item.strip()
        for item in re.split(r"[;；]+", _cell_text(value))
        if item.strip()
    )


def _safe_logical_stem(sample_id: str) -> str:
    value = INVALID_LOGICAL_FILENAME.sub("_", sample_id).strip(" ._")
    value = re.sub(r"\s+", "_", value)
    if not value:
        raise ProductPhotoBatchError("样品号不能生成安全的图片文件名")
    return value[:80]


def _safe_output_stem(value: str) -> str:
    stem = INVALID_LOGICAL_FILENAME.sub("_", value).strip(" ._")
    return stem[:120] or "常用箱图片批次"


def _sha256_file(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def _write_issue_report(
    output_dir: Path,
    *,
    batch_id: str,
    issues: Iterable[PhotoBatchIssue],
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    for sequence in range(1, 1_000):
        suffix = "" if sequence == 1 else f"_{sequence:03d}"
        report_path = (
            output_dir / f"批量图片处理_{batch_id}_异常报告{suffix}.csv"
        )
        try:
            with report_path.open(
                "x",
                encoding="utf-8-sig",
                newline="",
            ) as handle:
                writer = csv.writer(handle)
                writer.writerow(
                    ("错误代码", "样品号", "Excel行", "原文件名", "错误原因")
                )
                for issue in issues:
                    writer.writerow(
                        (
                            issue.code,
                            issue.sample_id,
                            issue.row_number or "",
                            issue.filename,
                            issue.message,
                        )
                    )
            return report_path
        except FileExistsError:
            continue
    raise ProductPhotoBatchError("同一批次异常报告过多，请更换输出目录或批次号")


def _raise_with_report(
    output_dir: Path,
    *,
    batch_id: str,
    issues: Iterable[PhotoBatchIssue],
) -> None:
    issue_list = tuple(issues)
    report_path = _write_issue_report(
        output_dir,
        batch_id=batch_id,
        issues=issue_list,
    )
    raise ProductPhotoBatchError(
        f"批量图片处理发现 {len(issue_list)} 个阻断问题，未生成可导入文件",
        issues=issue_list,
        report_path=report_path,
    )


def _load_template(
    template_path: Path,
) -> tuple[bytes, list[SampleRow], dict[int, tuple[bytes, ...]], dict[str, int]]:
    if not template_path.is_file() or template_path.suffix.lower() != ".xlsx":
        raise ProductPhotoBatchError("请选择有效的 .xlsx 常用箱登记表")
    content = template_path.read_bytes()
    try:
        _preflight_product_xlsx_container(content)
    except UploadValidationError as error:
        raise ProductPhotoBatchError(str(error)) from error
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    try:
        if PRODUCT_SHEET not in workbook.sheetnames:
            raise ProductPhotoBatchError(f"Excel 缺少“{PRODUCT_SHEET}”工作表")
        if MOLD_SHEET not in workbook.sheetnames:
            raise ProductPhotoBatchError(f"Excel 缺少“{MOLD_SHEET}”工作表")
        sheet = workbook[PRODUCT_SHEET]
        expected = {
            SAMPLE_ID_COLUMN: "样品号",
            DRAWING_COLUMN: "图纸文件名",
            ACTION_COLUMN: "操作",
        }
        for column, label in expected.items():
            actual = _cell_text(sheet.cell(PRODUCT_HEADER_ROW, column).value)
            if label not in actual:
                raise ProductPhotoBatchError(
                    f"“{PRODUCT_SHEET}”第 {column} 列表头应包含“{label}”"
                )

        rows: list[SampleRow] = []
        seen_sample_ids: dict[str, int] = {}
        for row_number in range(2, sheet.max_row + 1):
            sample_id = _cell_text(sheet.cell(row_number, SAMPLE_ID_COLUMN).value)
            action = _cell_text(sheet.cell(row_number, ACTION_COLUMN).value)
            drawing_names = _split_names(
                sheet.cell(row_number, DRAWING_COLUMN).value
            )
            if not sample_id:
                if action in {"新增", "更新"} or drawing_names:
                    raise ProductPhotoBatchError(
                        f"“{PRODUCT_SHEET}”第 {row_number} 行缺少样品号"
                    )
                continue
            normalized = _normalized_key(sample_id)
            if normalized in seen_sample_ids:
                raise ProductPhotoBatchError(
                    f"样品号重复：{sample_id}（第 {seen_sample_ids[normalized]}、"
                    f"{row_number} 行）"
                )
            seen_sample_ids[normalized] = row_number
            rows.append(
                SampleRow(
                    row_number=row_number,
                    sample_id=sample_id,
                    normalized_sample_id=normalized,
                    action=action,
                    mold_code=_cell_text(
                        sheet.cell(row_number, MOLD_CODE_COLUMN).value
                    ),
                    drawing_names=drawing_names,
                    remark=_cell_text(
                        sheet.cell(row_number, REMARK_COLUMN).value
                    ),
                )
            )

        existing_images: dict[int, list[tuple[int, bytes]]] = {}
        for drawing in sheet._images:
            anchor = drawing.anchor
            marker = getattr(anchor, "_from", None)
            if marker is None:
                raise ProductPhotoBatchError("Excel 存在无法识别锚点的内嵌图片")
            row_number = int(marker.row) + 1
            column_number = int(marker.col) + 1
            if column_number != DRAWING_COLUMN or row_number < 2:
                raise ProductPhotoBatchError(
                    f"Excel 内嵌图片必须放在“{PRODUCT_SHEET}”对应数据行的 J 列"
                )
            try:
                payload = drawing._data()
            except Exception as error:  # pragma: no cover - corrupt workbook path
                raise ProductPhotoBatchError(
                    f"第 {row_number} 行内嵌图片无法读取"
                ) from error
            offset = int(getattr(marker, "colOff", 0))
            existing_images.setdefault(row_number, []).append((offset, payload))

        ordered_images: dict[int, tuple[bytes, ...]] = {}
        for row_number, items in existing_images.items():
            items.sort(key=lambda item: item[0])
            if len(items) > 2:
                raise ProductPhotoBatchError(
                    f"第 {row_number} 行已有 {len(items)} 张内嵌图片；每款只能有2张"
                )
            offsets = [item[0] for item in items]
            if len(offsets) == 2 and offsets[0] == offsets[1]:
                raise ProductPhotoBatchError(
                    f"第 {row_number} 行两张内嵌图片锚点重叠，无法确认图1/图2顺序"
                )
            ordered_images[row_number] = tuple(payload for _offset, payload in items)

        mold_rows: dict[str, int] = {}
        mold_sheet = workbook[MOLD_SHEET]
        for row_number in range(2, mold_sheet.max_row + 1):
            action = _cell_text(mold_sheet.cell(row_number, 1).value)
            code = _cell_text(mold_sheet.cell(row_number, 2).value)
            if action == "新增" and code:
                key = _normalized_key(code)
                if key in mold_rows:
                    raise ProductPhotoBatchError(f"新增模具编号重复：{code}")
                mold_rows[key] = row_number
        return content, rows, ordered_images, mold_rows
    finally:
        workbook.close()


def _load_mapping_csv(path: Path | None) -> dict[str, tuple[str, str]]:
    if path is None:
        return {}
    if not path.is_file():
        raise ProductPhotoBatchError("照片配对 CSV 不存在")
    raw = path.read_bytes()
    if not raw or len(raw) > MAX_MAPPING_CSV_BYTES:
        raise ProductPhotoBatchError("照片配对 CSV 必须大于0且不超过2MiB")
    text: str | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None or "\x00" in text:
        raise ProductPhotoBatchError("照片配对 CSV 编码无法识别")
    with StringIO(text, newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = tuple(reader.fieldnames or ())
        missing = [header for header in MAPPING_HEADERS if header not in fieldnames]
        if missing:
            raise ProductPhotoBatchError(
                f"照片配对 CSV 缺少列：{'、'.join(missing)}"
            )
        result: dict[str, tuple[str, str]] = {}
        for csv_row_number, row in enumerate(reader, start=2):
            sample_id = _cell_text(row.get("样品号"))
            photo1 = _cell_text(row.get("照片1"))
            photo2 = _cell_text(row.get("照片2"))
            if not sample_id and not photo1 and not photo2:
                continue
            if not sample_id:
                raise ProductPhotoBatchError(
                    f"照片配对 CSV 第 {csv_row_number} 行缺少样品号"
                )
            key = _normalized_key(sample_id)
            if key in result:
                raise ProductPhotoBatchError(
                    f"照片配对 CSV 样品号重复：{sample_id}"
                )
            result[key] = (photo1, photo2)
        return result


def _inventory_photos(photo_dir: Path) -> tuple[dict[str, PhotoFile], list[PhotoBatchIssue]]:
    if not photo_dir.is_dir():
        raise ProductPhotoBatchError("请选择有效的照片文件夹")
    result: dict[str, PhotoFile] = {}
    issues: list[PhotoBatchIssue] = []
    for path in sorted(
        photo_dir.iterdir(),
        key=lambda item: _normalized_key(item.name),
    ):
        if path.is_symlink() or not path.is_file():
            continue
        if path.suffix.lower() not in PHOTO_EXTENSIONS:
            continue
        normalized = _normalized_filename(path.name)
        if normalized in result:
            issues.append(
                PhotoBatchIssue(
                    "PHOTO_NAME_COLLISION",
                    "照片文件名在忽略大小写/全半角后重复",
                    filename=path.name,
                )
            )
            continue
        size = path.stat().st_size
        if size <= 0 or size > MAX_SOURCE_PHOTO_BYTES:
            issues.append(
                PhotoBatchIssue(
                    "PHOTO_SIZE_INVALID",
                    f"原图必须大于0且不超过 {MAX_SOURCE_PHOTO_BYTES // (1024 * 1024)}MB",
                    filename=path.name,
                )
            )
            continue
        result[normalized] = PhotoFile(
            path=path,
            filename=path.name,
            normalized_filename=normalized,
            size=size,
            sha256=_sha256_file(path),
        )
    return result, issues


def _auto_role_for_filename(
    filename: str,
    sample_rows: dict[str, SampleRow],
) -> tuple[SampleRow, int] | None:
    stem = unicodedata.normalize("NFKC", Path(filename).stem).strip()
    for role, suffixes in ROLE_SUFFIXES.items():
        for suffix in suffixes:
            for separator in ("_", "-", " "):
                ending = f"{separator}{suffix}"
                if not stem.casefold().endswith(ending.casefold()):
                    continue
                candidate = stem[: -len(ending)]
                row = sample_rows.get(_normalized_key(candidate))
                if row is not None:
                    return row, role
    return None


def _drawing_names_are_default_placeholders(row: SampleRow) -> bool:
    if len(row.drawing_names) != 2:
        return False
    logical_stem = _safe_logical_stem(row.sample_id)
    normalized = tuple(_normalized_filename(name) for name in row.drawing_names)
    accepted = {
        (
            _normalized_filename(f"{logical_stem}_图1.jpg"),
            _normalized_filename(f"{logical_stem}_图2.jpg"),
        ),
        (
            _normalized_filename(f"{logical_stem}_实物.jpg"),
            _normalized_filename(f"{logical_stem}_展开.jpg"),
        ),
    }
    return normalized in accepted


def _open_image(
    payload: bytes,
    *,
    source_name: str,
    source_extension: str | None,
) -> Image.Image:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            image = Image.open(BytesIO(payload))
            detected_format = (image.format or "").upper()
            if source_extension is not None:
                allowed = SUPPORTED_DETECTED_FORMATS.get(source_extension.lower())
                if allowed is None or detected_format not in allowed:
                    raise ProductPhotoBatchError(
                        f"{source_name} 的扩展名与真实图片格式不一致"
                    )
            width, height = image.size
            if (
                width <= 0
                or height <= 0
                or width > MAX_SOURCE_DIMENSION
                or height > MAX_SOURCE_DIMENSION
                or width * height > MAX_SOURCE_PIXELS
            ):
                raise ProductPhotoBatchError(
                    f"{source_name} 像素尺寸过大或无效，最大允许 "
                    f"{MAX_SOURCE_DIMENSION}px/边、{MAX_SOURCE_PIXELS:,} 像素"
                )
            image.load()
            return image.copy()
    except ProductPhotoBatchError:
        raise
    except Image.DecompressionBombError as error:
        raise ProductPhotoBatchError(f"{source_name} 像素尺寸异常") from error
    except Image.DecompressionBombWarning as error:
        raise ProductPhotoBatchError(f"{source_name} 像素尺寸异常") from error
    except (OSError, ValueError, UnidentifiedImageError) as error:
        if source_extension in {".heic", ".heif"} and register_heif_opener is None:
            raise ProductPhotoBatchError(
                "当前环境缺少 pillow-heif，无法读取 HEIC/HEIF 原图"
            ) from error
        raise ProductPhotoBatchError(f"{source_name} 不是可完整解码的图片") from error


def _convert_to_srgb(image: Image.Image) -> Image.Image:
    icc_profile = image.info.get("icc_profile")
    if icc_profile:
        try:
            source_profile = ImageCms.ImageCmsProfile(BytesIO(icc_profile))
            target_profile = ImageCms.createProfile("sRGB")
            return ImageCms.profileToProfile(
                image,
                source_profile,
                target_profile,
                outputMode="RGB",
            )
        except (OSError, ValueError):
            pass
    if image.mode in {"RGBA", "LA"}:
        rgba = image.convert("RGBA")
        background = Image.new("RGB", rgba.size, "white")
        background.paste(rgba, mask=rgba.getchannel("A"))
        return background
    return image.convert("RGB")


def _prepare_photo_payload(
    payload: bytes,
    *,
    source_name: str,
    source_extension: str | None,
    cache_dir: Path,
    max_edge: int,
    jpeg_quality: int,
    from_workbook: bool,
) -> PreparedPhoto:
    source_digest = sha256(payload).hexdigest()
    if from_workbook:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(BytesIO(payload)) as existing:
                    width, height = existing.size
                    existing.load()
                    if (
                        (existing.format or "").upper() == "JPEG"
                        and width <= max_edge
                        and height <= max_edge
                        and width * height <= MAX_SOURCE_PIXELS
                        and len(existing.getexif()) == 0
                    ):
                        return PreparedPhoto(
                            source_name=source_name,
                            source_sha256=source_digest,
                            source_bytes=len(payload),
                            compressed_sha256=source_digest,
                            compressed_bytes=len(payload),
                            width=width,
                            height=height,
                            content=payload,
                            from_workbook=True,
                        )
        except (
            Image.DecompressionBombError,
            Image.DecompressionBombWarning,
            OSError,
            ValueError,
            UnidentifiedImageError,
        ):
            pass
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_dir / f"{source_digest}_{max_edge}_{jpeg_quality}.jpg"
    if cache_path.is_file():
        cached = cache_path.read_bytes()
        try:
            image = _open_image(
                cached,
                source_name=cache_path.name,
                source_extension=".jpg",
            )
            width, height = image.size
            image.close()
            return PreparedPhoto(
                source_name=source_name,
                source_sha256=source_digest,
                source_bytes=len(payload),
                compressed_sha256=sha256(cached).hexdigest(),
                compressed_bytes=len(cached),
                width=width,
                height=height,
                content=cached,
                from_workbook=from_workbook,
            )
        except ProductPhotoBatchError:
            cache_path.unlink(missing_ok=True)

    image = _open_image(
        payload,
        source_name=source_name,
        source_extension=source_extension,
    )
    try:
        image = ImageOps.exif_transpose(image)
        converted = _convert_to_srgb(image)
        converted.thumbnail((max_edge, max_edge), Image.Resampling.LANCZOS)
        output = BytesIO()
        converted.save(
            output,
            format="JPEG",
            quality=jpeg_quality,
            optimize=True,
            progressive=True,
            subsampling="4:2:0",
        )
        content = output.getvalue()
        width, height = converted.size
    finally:
        image.close()
        if "converted" in locals() and converted is not image:
            converted.close()
    temporary = cache_path.with_name(f".{cache_path.name}.tmp")
    temporary.write_bytes(content)
    temporary.replace(cache_path)
    return PreparedPhoto(
        source_name=source_name,
        source_sha256=source_digest,
        source_bytes=len(payload),
        compressed_sha256=sha256(content).hexdigest(),
        compressed_bytes=len(content),
        width=width,
        height=height,
        content=content,
        from_workbook=from_workbook,
    )


def _prepare_samples(
    *,
    rows: list[SampleRow],
    existing_images: dict[int, tuple[bytes, ...]],
    photo_files: dict[str, PhotoFile],
    mapping: dict[str, tuple[str, str]],
    cache_dir: Path,
    max_edge: int,
    jpeg_quality: int,
    strict_unused: bool,
) -> tuple[list[PreparedSample], list[PhotoBatchIssue]]:
    issues: list[PhotoBatchIssue] = []
    rows_by_id = {row.normalized_sample_id: row for row in rows}
    for sample_key in mapping:
        if sample_key not in rows_by_id:
            issues.append(
                PhotoBatchIssue(
                    "MAPPING_SAMPLE_NOT_FOUND",
                    "照片配对 CSV 的样品号不在 Excel 中",
                    sample_id=sample_key,
                )
            )

    inferred: dict[tuple[int, int], list[PhotoFile]] = {}
    for photo in photo_files.values():
        role_match = _auto_role_for_filename(photo.filename, rows_by_id)
        if role_match is None:
            continue
        row, role = role_match
        inferred.setdefault((row.row_number, role), []).append(photo)

    used_photos: set[str] = set()
    prepared_samples: list[PreparedSample] = []
    raw_hash_owners: dict[str, tuple[str, str]] = {}
    logical_name_owners: dict[str, str] = {}
    mapping_is_authoritative = bool(mapping)

    for row in rows:
        csv_names = mapping.get(row.normalized_sample_id)
        single_photo_requested = (
            csv_names is not None
            and _normalized_key(csv_names[1])
            == _normalized_key(NO_SECOND_PHOTO_TOKEN)
        ) or (
            csv_names is None
            and single_photo_requirement_confirmed(row.remark)
        )
        required_roles = (0,) if single_photo_requested else (0, 1)
        existing = (
            ()
            if csv_names is not None
            else existing_images.get(row.row_number, ())
        )
        j_names = row.drawing_names
        inferred_for_row = {
            role: inferred.get((row.row_number, role), [])
            for role in (0, 1)
        }
        explicit_j_mapping = bool(j_names) and not (
            _drawing_names_are_default_placeholders(row)
        )
        if mapping_is_authoritative:
            should_process = bool(existing or csv_names)
        else:
            should_process = bool(
                existing
                or csv_names
                or explicit_j_mapping
                or row.action in {"新增", "更新"}
                or inferred_for_row[0]
                or inferred_for_row[1]
            )
        if not should_process:
            continue
        if len(existing) > len(required_roles):
            issues.append(
                PhotoBatchIssue(
                    "TOO_MANY_EMBEDDED",
                    (
                        f"本行要求{len(required_roles)}张图片，"
                        f"但已有{len(existing)}张内嵌图片"
                    ),
                    sample_id=row.sample_id,
                    row_number=row.row_number,
                )
            )
            continue
        if csv_names is not None:
            requested_names = csv_names
        elif j_names:
            if len(j_names) != 2:
                issues.append(
                    PhotoBatchIssue(
                        "DRAWING_NAME_COUNT",
                        f"J列必须恰好写2个照片文件名，当前为{len(j_names)}个",
                        sample_id=row.sample_id,
                        row_number=row.row_number,
                    )
                )
                continue
            requested_names = (j_names[0], j_names[1])
        else:
            requested_names = ("", "")

        role_photos: list[PreparedPhoto | None] = [None, None]
        for role, payload in enumerate(existing):
            try:
                role_photos[role] = _prepare_photo_payload(
                    payload,
                    source_name=f"Excel第{row.row_number}行内嵌图{role + 1}",
                    source_extension=None,
                    cache_dir=cache_dir,
                    max_edge=max_edge,
                    jpeg_quality=jpeg_quality,
                    from_workbook=True,
                )
            except ProductPhotoBatchError as error:
                issues.append(
                    PhotoBatchIssue(
                        "EMBEDDED_DECODE_FAILED",
                        str(error),
                        sample_id=row.sample_id,
                        row_number=row.row_number,
                    )
                )

        for role in required_roles:
            if role_photos[role] is not None:
                continue
            requested_name = requested_names[role].strip()
            selected: PhotoFile | None = None
            if requested_name:
                if Path(requested_name).name != requested_name:
                    issues.append(
                        PhotoBatchIssue(
                            "MAPPING_PATH_NOT_ALLOWED",
                            "照片配对只能填写原文件名，不能包含目录",
                            sample_id=row.sample_id,
                            filename=requested_name,
                            row_number=row.row_number,
                        )
                    )
                    continue
                selected = photo_files.get(_normalized_filename(requested_name))
                if selected is None and csv_names is not None:
                    issues.append(
                        PhotoBatchIssue(
                            "PHOTO_NOT_FOUND",
                            f"配对文件不存在：{requested_name}",
                            sample_id=row.sample_id,
                            filename=requested_name,
                            row_number=row.row_number,
                        )
                    )
                    continue
            if selected is None:
                candidates = inferred_for_row[role]
                if len(candidates) > 1:
                    issues.append(
                        PhotoBatchIssue(
                            "PHOTO_ROLE_DUPLICATE",
                            f"图{role + 1}匹配到多张照片，不能自动选择",
                            sample_id=row.sample_id,
                            row_number=row.row_number,
                        )
                    )
                    continue
                if len(candidates) == 1:
                    selected = candidates[0]
            if selected is None:
                issues.append(
                    PhotoBatchIssue(
                        "PHOTO_MISSING",
                        f"缺少图{role + 1}；请在配对CSV/J列写原文件名，"
                        f"或把照片命名为“{row.sample_id}_图{role + 1}”",
                        sample_id=row.sample_id,
                        filename=requested_name,
                        row_number=row.row_number,
                    )
                )
                continue
            if selected.normalized_filename in used_photos:
                issues.append(
                    PhotoBatchIssue(
                        "PHOTO_REUSED",
                        "同一原图被分配给多个位置",
                        sample_id=row.sample_id,
                        filename=selected.filename,
                        row_number=row.row_number,
                    )
                )
                continue
            used_photos.add(selected.normalized_filename)
            try:
                source_content = selected.path.read_bytes()
                if (
                    len(source_content) != selected.size
                    or sha256(source_content).hexdigest() != selected.sha256
                ):
                    raise ProductPhotoBatchError(
                        f"{selected.filename} 在清点后发生变化，请重新运行"
                    )
                role_photos[role] = _prepare_photo_payload(
                    source_content,
                    source_name=selected.filename,
                    source_extension=selected.path.suffix.lower(),
                    cache_dir=cache_dir,
                    max_edge=max_edge,
                    jpeg_quality=jpeg_quality,
                    from_workbook=False,
                )
            except ProductPhotoBatchError as error:
                issues.append(
                    PhotoBatchIssue(
                        "PHOTO_DECODE_FAILED",
                        str(error),
                        sample_id=row.sample_id,
                        filename=selected.filename,
                        row_number=row.row_number,
                    )
                )

        if any(role_photos[role] is None for role in required_roles):
            continue
        complete_photos = tuple(
            role_photos[role] for role in required_roles
        )
        assert all(photo is not None for photo in complete_photos)
        typed_photos = tuple(
            photo for photo in complete_photos if photo is not None
        )
        if (
            len(typed_photos) == 2
            and typed_photos[0].source_sha256
            == typed_photos[1].source_sha256
        ):
            issues.append(
                PhotoBatchIssue(
                    "PHOTO_CONTENT_DUPLICATE",
                    "图1和图2内容完全相同，不能用同一张照片凑两图",
                    sample_id=row.sample_id,
                    row_number=row.row_number,
                )
            )
            continue
        for role, photo in enumerate(typed_photos, start=1):
            previous = raw_hash_owners.get(photo.source_sha256)
            owner = (row.sample_id, f"图{role}")
            if previous is not None and previous != owner:
                issues.append(
                    PhotoBatchIssue(
                        "PHOTO_HASH_REUSED",
                        f"图片内容已分配给 {previous[0]} {previous[1]}",
                        sample_id=row.sample_id,
                        filename=photo.source_name,
                        row_number=row.row_number,
                    )
                )
            else:
                raw_hash_owners[photo.source_sha256] = owner

        logical_stem = _safe_logical_stem(row.sample_id)
        logical_names = tuple(
            f"{logical_stem}_图{role + 1}.jpg"
            for role in required_roles
        )
        logical_collision = False
        for logical_name in logical_names:
            key = _normalized_filename(logical_name)
            previous_sample = logical_name_owners.get(key)
            if previous_sample is not None and previous_sample != row.sample_id:
                issues.append(
                    PhotoBatchIssue(
                        "LOGICAL_NAME_COLLISION",
                        f"压缩后逻辑文件名与样品 {previous_sample} 冲突",
                        sample_id=row.sample_id,
                        filename=logical_name,
                        row_number=row.row_number,
                    )
                )
                logical_collision = True
            else:
                logical_name_owners[key] = row.sample_id
        if logical_collision:
            continue
        prepared_samples.append(
            PreparedSample(
                row_number=row.row_number,
                sample_id=row.sample_id,
                logical_names=logical_names,
                photos=typed_photos,
                mold_code=row.mold_code,
            )
        )

    # A mapping CSV is an explicit selection list: photos not named there are
    # future/unrecorded work and must not block the current mapped batch.
    if strict_unused and not mapping_is_authoritative:
        for key, photo in photo_files.items():
            if key not in used_photos:
                issues.append(
                    PhotoBatchIssue(
                        "PHOTO_UNUSED",
                        "照片文件夹中的图片未归属任何样品",
                        filename=photo.filename,
                    )
                )
    return prepared_samples, issues


def _partition_samples(
    samples: list[PreparedSample],
    *,
    source_workbook_bytes: int,
    target_xlsx_bytes: int,
    max_samples_per_volume: int,
) -> list[list[PreparedSample]]:
    if source_workbook_bytes >= target_xlsx_bytes:
        raise ProductPhotoBatchError(
            "原始 Excel 已达到分卷目标大小，请先使用未嵌图模板或提高目标上限"
        )
    image_budget = max(1, int((target_xlsx_bytes - source_workbook_bytes) * 0.86))
    partitions: list[list[PreparedSample]] = []
    current: list[PreparedSample] = []
    current_bytes = 0
    for sample in sorted(samples, key=lambda item: item.row_number):
        would_exceed = (
            current
            and (
                len(current) >= max_samples_per_volume
                or current_bytes + sample.compressed_bytes > image_budget
            )
        )
        if would_exceed:
            partitions.append(current)
            current = []
            current_bytes = 0
        current.append(sample)
        current_bytes += sample.compressed_bytes
    if current:
        partitions.append(current)
    return partitions


def _estimated_base_workbook_bytes(template_content: bytes) -> int:
    """Estimate the workbook size after existing embedded photos are removed."""

    with zipfile.ZipFile(BytesIO(template_content)) as archive:
        embedded_compressed = sum(
            info.compress_size
            for info in archive.infolist()
            if info.filename.casefold().startswith("xl/media/")
        )
    return max(64 * 1024, len(template_content) - embedded_compressed)


def _add_prepared_image(
    sheet,
    *,
    row_number: int,
    slot: int,
    photo: PreparedPhoto,
    streams: list[BytesIO],
) -> None:
    stream = BytesIO(photo.content)
    streams.append(stream)
    image = WorksheetImage(stream)
    aspect = photo.width / photo.height
    width_px = DISPLAY_SLOT_WIDTH_PX
    height_px = round(width_px / aspect)
    if height_px > DISPLAY_SLOT_HEIGHT_PX:
        height_px = DISPLAY_SLOT_HEIGHT_PX
        width_px = round(height_px * aspect)
    slot_left = (
        DISPLAY_CELL_LEFT_PX
        + slot * (DISPLAY_SLOT_WIDTH_PX + DISPLAY_SLOT_GAP_PX)
    )
    centered_left = slot_left + max(
        0,
        round((DISPLAY_SLOT_WIDTH_PX - width_px) / 2),
    )
    centered_top = DISPLAY_CELL_TOP_PX + max(
        0,
        round((DISPLAY_SLOT_HEIGHT_PX - height_px) / 2),
    )
    marker = AnchorMarker(
        col=DRAWING_COLUMN - 1,
        colOff=pixels_to_EMU(centered_left),
        row=row_number - 1,
        rowOff=pixels_to_EMU(centered_top),
    )
    image.anchor = OneCellAnchor(
        _from=marker,
        ext=XDRPositiveSize2D(
            cx=pixels_to_EMU(width_px),
            cy=pixels_to_EMU(height_px),
        ),
    )
    sheet.add_image(image)


def _write_volume(
    template_content: bytes,
    *,
    all_samples: list[PreparedSample],
    volume_samples: list[PreparedSample],
    mold_rows: dict[str, int],
    first_mold_volume: dict[str, int],
    volume_index: int,
    target_path: Path,
) -> None:
    workbook = load_workbook(
        BytesIO(template_content),
        data_only=False,
        read_only=False,
    )
    streams: list[BytesIO] = []
    try:
        sheet = workbook[PRODUCT_SHEET]
        sheet._images = []
        selected_rows = {sample.row_number for sample in volume_samples}
        for sample in all_samples:
            if sample.row_number not in selected_rows:
                sheet.cell(sample.row_number, ACTION_COLUMN).value = None
                sheet.cell(sample.row_number, DRAWING_COLUMN).value = None
        for sample in volume_samples:
            sheet.cell(sample.row_number, DRAWING_COLUMN).value = ";".join(
                sample.logical_names
            )
            prior_remark = _cell_text(
                sheet.cell(sample.row_number, REMARK_COLUMN).value
            )
            if prior_remark.startswith("仅1张：") and "ERP必须阻止导入" in prior_remark:
                prior_remark = ""
                sheet.cell(sample.row_number, DRAWING_COLUMN).fill = PatternFill(
                    "solid",
                    fgColor="FFF2CC",
                )
            updated_remark = _remark_with_single_photo_requirement(
                prior_remark,
                enabled=len(sample.photos) == 1,
            )
            sheet.cell(sample.row_number, REMARK_COLUMN).value = (
                updated_remark or None
            )
            sheet.row_dimensions[sample.row_number].height = (
                DISPLAY_ROW_HEIGHT_POINTS
            )
            for slot, photo in enumerate(sample.photos):
                _add_prepared_image(
                    sheet,
                    row_number=sample.row_number,
                    slot=slot,
                    photo=photo,
                    streams=streams,
                )
        sheet.column_dimensions["J"].width = DISPLAY_COLUMN_WIDTH

        mold_sheet = workbook[MOLD_SHEET]
        selected_molds = {
            _normalized_key(sample.mold_code)
            for sample in volume_samples
            if sample.mold_code
        }
        for mold_key, row_number in mold_rows.items():
            if (
                mold_key not in selected_molds
                or first_mold_volume.get(mold_key) != volume_index
            ):
                mold_sheet.cell(row_number, 1).value = None

        workbook.save(target_path)
    finally:
        workbook.close()
        for stream in streams:
            stream.close()


def _verify_volume(
    path: Path,
    *,
    samples: list[PreparedSample],
    absolute_xlsx_bytes: int,
) -> None:
    size = path.stat().st_size
    if size > absolute_xlsx_bytes:
        raise ProductPhotoBatchError(
            f"{path.name} 为 {size / 1024 / 1024:.2f}MiB，超过 ERP 绝对上限"
        )
    content = path.read_bytes()
    try:
        _preflight_product_xlsx_container(content)
    except UploadValidationError as error:
        raise ProductPhotoBatchError(
            f"{path.name} 未通过 ERP XLSX 安全检查：{error}"
        ) from error
    workbook = load_workbook(BytesIO(content), data_only=False, read_only=False)
    try:
        sheet = workbook[PRODUCT_SHEET]
        expected_rows = {sample.row_number for sample in samples}
        anchors: dict[int, list[int]] = {}
        for drawing in sheet._images:
            marker = getattr(drawing.anchor, "_from", None)
            if marker is None or int(marker.col) + 1 != DRAWING_COLUMN:
                raise ProductPhotoBatchError(
                    f"{path.name} 存在未锚定在 J 列的图片"
                )
            row_number = int(marker.row) + 1
            anchors.setdefault(row_number, []).append(
                int(getattr(marker, "colOff", 0))
            )
        if set(anchors) != expected_rows:
            raise ProductPhotoBatchError(
                f"{path.name} 的图片行与本卷样品行不一致"
            )
        by_row = {sample.row_number: sample for sample in samples}
        for row_number, offsets in anchors.items():
            expected_photo_count = len(by_row[row_number].photos)
            if (
                len(offsets) != expected_photo_count
                or len(set(offsets)) != len(offsets)
            ):
                raise ProductPhotoBatchError(
                    f"{path.name} 第 {row_number} 行图片数量或顺序"
                    f"与本卷要求不一致（应为{expected_photo_count}张）"
                )
            expected_names = ";".join(by_row[row_number].logical_names)
            if _cell_text(sheet.cell(row_number, DRAWING_COLUMN).value) != expected_names:
                raise ProductPhotoBatchError(
                    f"{path.name} 第 {row_number} 行 J 列文件名与图片不一致"
                )
    finally:
        workbook.close()


def _fit_partitions_by_actual_size(
    template_content: bytes,
    *,
    all_samples: list[PreparedSample],
    initial_partitions: list[list[PreparedSample]],
    mold_rows: dict[str, int],
    target_xlsx_bytes: int,
    temporary_dir: Path,
) -> list[list[PreparedSample]]:
    partitions = [list(partition) for partition in initial_partitions]
    while True:
        first_mold_volume: dict[str, int] = {}
        for index, partition in enumerate(partitions):
            for sample in partition:
                if sample.mold_code:
                    first_mold_volume.setdefault(
                        _normalized_key(sample.mold_code),
                        index,
                    )
        oversized_index: int | None = None
        for index, partition in enumerate(partitions):
            probe = temporary_dir / f"probe_{index:04d}.xlsx"
            _write_volume(
                template_content,
                all_samples=all_samples,
                volume_samples=partition,
                mold_rows=mold_rows,
                first_mold_volume=first_mold_volume,
                volume_index=index,
                target_path=probe,
            )
            if probe.stat().st_size > target_xlsx_bytes:
                oversized_index = index
                break
        for probe in temporary_dir.glob("probe_*.xlsx"):
            probe.unlink(missing_ok=True)
        if oversized_index is None:
            return partitions
        oversized = partitions[oversized_index]
        if len(oversized) <= 1:
            raise ProductPhotoBatchError(
                f"样品 {oversized[0].sample_id} 单独成卷仍超过目标大小，"
                "请降低图片尺寸/质量或提高分卷上限"
            )
        midpoint = len(oversized) // 2
        partitions[oversized_index : oversized_index + 1] = [
            oversized[:midpoint],
            oversized[midpoint:],
        ]


def _write_success_reports(
    *,
    output_dir: Path,
    batch_id: str,
    result_data: dict,
    samples: list[PreparedSample],
    volumes: list[GeneratedVolume],
) -> tuple[Path, Path]:
    report_csv = output_dir / f"批量图片处理_{batch_id}_明细.csv"
    volume_by_sample = {
        sample_id: volume.volume_number
        for volume in volumes
        for sample_id in volume.sample_ids
    }
    with report_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            (
                "分卷",
                "样品号",
                "Excel行",
                "图位",
                "原文件名",
                "逻辑文件名",
                "原图字节",
                "压缩字节",
                "原图SHA256",
                "压缩图SHA256",
                "来源",
            )
        )
        for sample in samples:
            for role, (logical_name, photo) in enumerate(
                zip(sample.logical_names, sample.photos, strict=True),
                start=1,
            ):
                writer.writerow(
                    (
                        volume_by_sample[sample.sample_id],
                        sample.sample_id,
                        sample.row_number,
                        f"图{role}",
                        photo.source_name,
                        logical_name,
                        photo.source_bytes,
                        photo.compressed_bytes,
                        photo.source_sha256,
                        photo.compressed_sha256,
                        "Excel已有内嵌图" if photo.from_workbook else "照片文件夹",
                    )
                )

    report_json = output_dir / f"批量图片处理_{batch_id}_汇总.json"
    report_json.write_text(
        json.dumps(result_data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return report_csv, report_json


def _build_product_photo_workbooks(
    *,
    template_path: Path,
    photo_dir: Path,
    output_dir: Path,
    mapping_path: Path | None = None,
    batch_id: str | None = None,
    strict_unused: bool = True,
    target_xlsx_bytes: int = DEFAULT_TARGET_XLSX_BYTES,
    absolute_xlsx_bytes: int = ERP_ABSOLUTE_XLSX_BYTES,
    max_samples_per_volume: int = DEFAULT_MAX_SAMPLES_PER_VOLUME,
    max_edge: int = DEFAULT_MAX_EDGE,
    jpeg_quality: int = DEFAULT_JPEG_QUALITY,
) -> ProductPhotoBatchResult:
    """Create deterministic, ERP-ready photo workbooks without touching ERP data.

    The source workbook and source photos are read-only.  The function never
    fills formal product/customer codes and never changes a row to ``新增``.
    """

    template_path = template_path.resolve()
    photo_dir = photo_dir.resolve()
    output_dir = output_dir.resolve()
    mapping_path = mapping_path.resolve() if mapping_path is not None else None
    batch_id = batch_id or datetime.now().strftime("%Y%m%d_%H%M%S")
    batch_id = _safe_output_stem(batch_id)
    if not (1 * 1024 * 1024 <= target_xlsx_bytes <= absolute_xlsx_bytes):
        raise ProductPhotoBatchError(
            "分卷目标必须介于 1MiB 和 ERP 绝对上限之间"
        )
    if absolute_xlsx_bytes > ERP_ABSOLUTE_XLSX_BYTES:
        raise ProductPhotoBatchError("不得提高 ERP 20MiB 绝对上限")
    if max_samples_per_volume < 1 or max_samples_per_volume > 200:
        raise ProductPhotoBatchError("每卷样品数必须介于1和200之间")
    if max_edge < 800 or max_edge > 2_000:
        raise ProductPhotoBatchError("图片长边必须介于800和2000像素之间")
    if jpeg_quality < 60 or jpeg_quality > 90:
        raise ProductPhotoBatchError("JPEG质量必须介于60和90之间")

    output_dir.mkdir(parents=True, exist_ok=True)
    template_content, rows, existing_images, mold_rows = _load_template(
        template_path
    )
    mapping = _load_mapping_csv(mapping_path)
    photo_files, inventory_issues = _inventory_photos(photo_dir)
    cache_dir = output_dir / ".product_photo_cache"
    samples, preparation_issues = _prepare_samples(
        rows=rows,
        existing_images=existing_images,
        photo_files=photo_files,
        mapping=mapping,
        cache_dir=cache_dir,
        max_edge=max_edge,
        jpeg_quality=jpeg_quality,
        strict_unused=strict_unused,
    )
    issues = [*inventory_issues, *preparation_issues]
    if not samples and not issues:
        issues.append(
            PhotoBatchIssue(
                "NO_SAMPLES",
                "没有找到需要处理的样品行；请填写J列配对、提供配对CSV，"
                "或按“样品号_图1/图2”命名照片",
            )
        )
    if issues:
        _raise_with_report(
            output_dir,
            batch_id=batch_id,
            issues=issues,
        )

    initial_partitions = _partition_samples(
        samples,
        source_workbook_bytes=_estimated_base_workbook_bytes(template_content),
        target_xlsx_bytes=target_xlsx_bytes,
        max_samples_per_volume=max_samples_per_volume,
    )
    temporary_root = Path(
        tempfile.mkdtemp(prefix=".product_photo_batch_", dir=output_dir)
    )
    try:
        partitions = _fit_partitions_by_actual_size(
            template_content,
            all_samples=samples,
            initial_partitions=initial_partitions,
            mold_rows=mold_rows,
            target_xlsx_bytes=target_xlsx_bytes,
            temporary_dir=temporary_root,
        )
        first_mold_volume: dict[str, int] = {}
        for index, partition in enumerate(partitions):
            for sample in partition:
                if sample.mold_code:
                    first_mold_volume.setdefault(
                        _normalized_key(sample.mold_code),
                        index,
                    )

        source_stem = _safe_output_stem(template_path.stem)
        staged: list[tuple[Path, Path, list[PreparedSample]]] = []
        total = len(partitions)
        for index, partition in enumerate(partitions):
            temporary_path = temporary_root / f"volume_{index + 1:04d}.xlsx"
            final_path = output_dir / (
                f"{source_stem}_{batch_id}_第{index + 1:03d}卷_共{total:03d}卷.xlsx"
            )
            if final_path.exists():
                raise ProductPhotoBatchError(
                    f"输出文件已存在，为避免覆盖已停止：{final_path.name}"
                )
            _write_volume(
                template_content,
                all_samples=samples,
                volume_samples=partition,
                mold_rows=mold_rows,
                first_mold_volume=first_mold_volume,
                volume_index=index,
                target_path=temporary_path,
            )
            _verify_volume(
                temporary_path,
                samples=partition,
                absolute_xlsx_bytes=absolute_xlsx_bytes,
            )
            if temporary_path.stat().st_size > target_xlsx_bytes:
                raise ProductPhotoBatchError(
                    f"{temporary_path.name} 超过分卷目标大小"
                )
            staged.append((temporary_path, final_path, partition))

        volumes: list[GeneratedVolume] = []
        finalized_paths: list[Path] = []
        try:
            for index, (temporary_path, final_path, partition) in enumerate(staged):
                temporary_path.replace(final_path)
                finalized_paths.append(final_path)
                volumes.append(
                    GeneratedVolume(
                        volume_number=index + 1,
                        volume_count=len(staged),
                        path=final_path,
                        size=final_path.stat().st_size,
                        sha256=_sha256_file(final_path),
                        sample_ids=tuple(sample.sample_id for sample in partition),
                    )
                )
        except Exception:
            for finalized_path in finalized_paths:
                finalized_path.unlink(missing_ok=True)
            raise
    except Exception:
        for candidate in temporary_root.glob("*.xlsx"):
            candidate.unlink(missing_ok=True)
        raise
    finally:
        shutil.rmtree(temporary_root, ignore_errors=True)

    source_photo_bytes = sum(
        photo.source_bytes for sample in samples for photo in sample.photos
    )
    compressed_photo_bytes = sum(
        photo.compressed_bytes for sample in samples for photo in sample.photos
    )
    formal_status = (
        "待核对：客户正式存货编码映射未确认前，不得把手写型号自动作为ERP存货编码"
    )
    result_data = {
        "batch_id": batch_id,
        "source_workbook": str(template_path),
        "source_workbook_sha256": sha256(template_content).hexdigest(),
        "photo_directory": str(photo_dir),
        "mapping_file": str(mapping_path) if mapping_path else None,
        "sample_count": len(samples),
        "photo_count": sum(len(sample.photos) for sample in samples),
        "source_photo_bytes": source_photo_bytes,
        "compressed_photo_bytes": compressed_photo_bytes,
        "compression_ratio": (
            round(compressed_photo_bytes / source_photo_bytes, 6)
            if source_photo_bytes
            else 0
        ),
        "max_edge": max_edge,
        "jpeg_quality": jpeg_quality,
        "target_xlsx_bytes": target_xlsx_bytes,
        "formal_import_status": formal_status,
        "volume_transaction_boundary": (
            "当前ERP每卷独立预检、独立事务；必须按卷号顺序逐卷导入，"
            "不得宣称多卷整体回滚"
        ),
        "volumes": [
            {
                **asdict(volume),
                "path": str(volume.path),
                "sample_ids": list(volume.sample_ids),
            }
            for volume in volumes
        ],
    }
    try:
        report_csv, report_json = _write_success_reports(
            output_dir=output_dir,
            batch_id=batch_id,
            result_data=result_data,
            samples=samples,
            volumes=volumes,
        )
    except Exception:
        for volume in volumes:
            volume.path.unlink(missing_ok=True)
        raise
    return ProductPhotoBatchResult(
        batch_id=batch_id,
        source_workbook=template_path,
        source_workbook_sha256=sha256(template_content).hexdigest(),
        photo_directory=photo_dir,
        volumes=tuple(volumes),
        report_csv=report_csv,
        report_json=report_json,
        source_photo_bytes=source_photo_bytes,
        compressed_photo_bytes=compressed_photo_bytes,
        sample_count=len(samples),
        photo_count=sum(len(sample.photos) for sample in samples),
        formal_import_status=formal_status,
    )


def build_product_photo_workbooks(
    *,
    template_path: Path,
    photo_dir: Path,
    output_dir: Path,
    mapping_path: Path | None = None,
    batch_id: str | None = None,
    strict_unused: bool = True,
    target_xlsx_bytes: int = DEFAULT_TARGET_XLSX_BYTES,
    absolute_xlsx_bytes: int = ERP_ABSOLUTE_XLSX_BYTES,
    max_samples_per_volume: int = DEFAULT_MAX_SAMPLES_PER_VOLUME,
    max_edge: int = DEFAULT_MAX_EDGE,
    jpeg_quality: int = DEFAULT_JPEG_QUALITY,
) -> ProductPhotoBatchResult:
    """Create photo workbooks and always leave a CSV report on business errors."""

    effective_batch_id = _safe_output_stem(
        batch_id or datetime.now().strftime("%Y%m%d_%H%M%S")
    )
    try:
        return _build_product_photo_workbooks(
            template_path=template_path,
            photo_dir=photo_dir,
            output_dir=output_dir,
            mapping_path=mapping_path,
            batch_id=effective_batch_id,
            strict_unused=strict_unused,
            target_xlsx_bytes=target_xlsx_bytes,
            absolute_xlsx_bytes=absolute_xlsx_bytes,
            max_samples_per_volume=max_samples_per_volume,
            max_edge=max_edge,
            jpeg_quality=jpeg_quality,
        )
    except ProductPhotoBatchError as error:
        if error.report_path is not None:
            raise
        issues = error.issues or (
            PhotoBatchIssue(
                "INPUT_VALIDATION",
                str(error),
            ),
        )
        _raise_with_report(
            output_dir.resolve(),
            batch_id=effective_batch_id,
            issues=issues,
        )
        raise AssertionError("unreachable")
