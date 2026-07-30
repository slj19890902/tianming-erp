from __future__ import annotations

from collections import OrderedDict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import re
import secrets
from threading import RLock
import time
from typing import Any

from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.master_data_common import audit_master_change
from app.models.material import Material
from app.models.material_price_history import MaterialPriceHistory
from app.models.supplier_paper_code import SupplierPaperCode
from app.models.supplier import Supplier
from app.models.user import User
from app.services.inventory_onboarding_uploads import _preflight_xlsx_container
from app.services.master_data_versioning import (
    apply_versioned_update,
    canonical_json,
    record_versioned_create,
)
from app.services.secure_uploads import (
    EXCEL_POLICY,
    UploadValidationError,
    read_validated_upload,
)
from app.services.supplier_master import (
    SupplierLookupError,
    normalize_supplier_identity,
    resolve_supplier,
)


PAPER_SHEET = "基础纸种"
MATERIAL_SHEET = "组合材质"
PAPER_HEADERS = (
    "系统ID",
    "供应商标准名",
    "基础代码",
    "纸种名称",
    "克重(g)",
    "等级/说明",
    "纸张用途",
    "备注",
    "启用",
)
MATERIAL_HEADERS = (
    "系统ID",
    "当前版本",
    "供应商标准名",
    "组合代码",
    "层数",
    "平方单价(元/㎡)",
    "报价日期",
    "价格单位",
    "备注",
    "启用",
)
MAX_ROWS_PER_SHEET = 5_000
PREVIEW_TTL_SECONDS = 5 * 60
MAX_PREVIEWS = 64

class SupplierMaterialWorkbookError(ValueError):
    def __init__(self, code: str, message: str, *, status_code: int = 400) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


@dataclass(frozen=True)
class _WorkbookPreview:
    actor: str
    expires_at: float
    source_sha256: str
    paper_items: tuple[dict[str, Any], ...]
    material_items: tuple[dict[str, Any], ...]
    supplier_references: tuple[dict[str, Any], ...]
    summary: dict[str, int]


_PREVIEWS: OrderedDict[str, _WorkbookPreview] = OrderedDict()
_PREVIEW_LOCK = RLock()


def normalize_supplier_name(value: object) -> str:
    """Only clean workbook input; canonicalization belongs to SupplierAlias."""
    return str(value or "").strip()


def _actor(user: User) -> str:
    return f"id:{user.id}" if user.id is not None else f"username:{user.username}"


def _purge_previews(now: float) -> None:
    expired = [token for token, item in _PREVIEWS.items() if item.expires_at <= now]
    for token in expired:
        _PREVIEWS.pop(token, None)
    while len(_PREVIEWS) >= MAX_PREVIEWS:
        _PREVIEWS.popitem(last=False)


def _store_preview(item: _WorkbookPreview) -> str:
    token = secrets.token_urlsafe(32)
    with _PREVIEW_LOCK:
        _purge_previews(time.monotonic())
        _PREVIEWS[token] = item
    return token


def _take_preview(token: str, user: User) -> _WorkbookPreview:
    with _PREVIEW_LOCK:
        now = time.monotonic()
        _purge_previews(now)
        item = _PREVIEWS.get(token)
        if item is None or item.actor != _actor(user):
            raise SupplierMaterialWorkbookError(
                "SUPPLIER_MATERIAL_PREVIEW_STALE",
                "批量预览已过期或不属于当前操作人，请重新导入预览",
                status_code=409,
            )
        _PREVIEWS.move_to_end(token)
        return item


def _consume_preview(token: str) -> None:
    with _PREVIEW_LOCK:
        _PREVIEWS.pop(token, None)


def _paper_snapshot(row: SupplierPaperCode) -> dict[str, Any]:
    return {
        "id": row.id,
        "supplier_name": row.supplier_name,
        "code_char": row.code_char,
        "paper_name": row.paper_name,
        "gram_weight": row.gram_weight,
        "paper_grade": row.paper_grade,
        "paper_role": row.paper_role,
        "remark": row.remark,
        "is_active": row.is_active,
    }


def _fingerprint(value: dict[str, Any]) -> str:
    return sha256(canonical_json(value).encode("utf-8")).hexdigest()


def _style_sheet(worksheet, widths: tuple[int, ...]) -> None:
    header_fill = PatternFill("solid", fgColor="DCEAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width


def build_supplier_material_workbook(db: Session) -> bytes:
    workbook = Workbook()
    paper_sheet = workbook.active
    paper_sheet.title = PAPER_SHEET
    paper_sheet.append(PAPER_HEADERS)
    paper_notes = (
        "新增时留空；修改已有资料时也可留空，系统按供应商+基础代码匹配。",
        "填写已启用的供应商标准名或主档中维护的别名。",
        "单个字母或数字，例如 C、4、6。",
        "该基础代码对应的纸种名称。",
        "只填数字，例如 130。",
        "可空。",
        "可填面纸、芯纸、中纸、里纸或留空。",
        "可空。",
        "填写是或否；留空按是处理。",
    )
    for cell, note in zip(paper_sheet[1], paper_notes):
        cell.comment = Comment(note, "天明ERP")
    _style_sheet(paper_sheet, (10, 22, 12, 24, 12, 18, 14, 28, 10))

    material_sheet = workbook.create_sheet(MATERIAL_SHEET)
    material_sheet.append(MATERIAL_HEADERS)
    material_notes = (
        "新增时留空；修改已有资料时也可留空，系统按组合代码匹配。",
        "通常留空；使用旧导出表修改时可保留版本号。",
        "填写已启用的供应商标准名或主档中维护的别名。",
        "填写3位、5位或7位组合代码；代码中的基础字符必须已存在或在本文件基础纸种表同时新增。",
        "只填3、5或7。",
        "当前平方报价，单位元/平方米。",
        "可空；填写格式 YYYY-MM-DD。",
        "默认元/㎡，通常无需修改。",
        "可空。",
        "填写是或否；留空按是处理。",
    )
    for cell, note in zip(material_sheet[1], material_notes):
        cell.comment = Comment(note, "天明ERP")
    _style_sheet(material_sheet, (10, 10, 22, 16, 10, 18, 14, 14, 30, 10))

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _cell_text(value: object) -> str:
    return str(value or "").strip()


def _positive_int(value: object, label: str, *, optional: bool = False) -> int | None:
    if value is None or _cell_text(value) == "":
        if optional:
            return None
        raise ValueError(f"{label}不能为空")
    if isinstance(value, bool):
        raise ValueError(f"{label}必须是正整数")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"{label}必须是正整数") from error
    if number != number.to_integral_value() or number <= 0:
        raise ValueError(f"{label}必须是正整数")
    return int(number)


def _price(value: object) -> Decimal:
    if value is None or _cell_text(value) == "":
        raise ValueError("平方单价不能为空")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise ValueError("平方单价必须是数字") from error
    if not result.is_finite() or result < 0:
        raise ValueError("平方单价必须大于等于0")
    return result.quantize(Decimal("0.0001"))


def _date(value: object) -> date | None:
    if value is None or _cell_text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_text(value).replace("/", "-")
    try:
        return date.fromisoformat(text)
    except ValueError as error:
        raise ValueError("报价日期必须为 YYYY-MM-DD") from error


def _enabled(value: object) -> bool:
    if value is None or _cell_text(value) == "":
        return True
    text = _cell_text(value).lower()
    if text in {"是", "启用", "1", "true", "yes", "y"}:
        return True
    if text in {"否", "停用", "0", "false", "no", "n"}:
        return False
    raise ValueError("启用列只能填写是/否")


def _validate_headers(worksheet, expected: tuple[str, ...]) -> None:
    actual = tuple(_cell_text(cell.value) for cell in worksheet[1])
    if actual != expected:
        raise SupplierMaterialWorkbookError(
            "SUPPLIER_MATERIAL_HEADERS_INVALID",
            (
                f"工作表“{worksheet.title}”表头不匹配；"
                "请重新下载系统模板后填写"
            ),
        )


def _append_error(
    errors: list[dict[str, Any]],
    *,
    sheet: str,
    row_number: int,
    message: str,
) -> None:
    if len(errors) < 200:
        errors.append(
            {"sheet": sheet, "row_number": row_number, "message": message}
        )


def _parse_paper_rows(worksheet, errors: list[dict[str, Any]]) -> list[dict]:
    rows: list[dict] = []
    seen_ids: set[int] = set()
    seen_keys: set[tuple[str, str]] = set()
    for row_number, cells in enumerate(
        worksheet.iter_rows(min_row=2, max_col=len(PAPER_HEADERS)),
        start=2,
    ):
        values = [cell.value for cell in cells]
        if not any(value is not None and _cell_text(value) != "" for value in values):
            continue
        try:
            system_id = _positive_int(values[0], "系统ID", optional=True)
            supplier = normalize_supplier_name(values[1])
            if not supplier:
                raise ValueError("供应商标准名不能为空")
            code = _cell_text(values[2]).upper()
            if not re.fullmatch(r"[A-Z0-9]", code):
                raise ValueError("基础代码必须是单个字母或数字")
            paper_name = _cell_text(values[3])
            if not paper_name:
                raise ValueError("纸种名称不能为空")
            gram_weight = _positive_int(values[4], "克重")
            if gram_weight is None or gram_weight > 2000:
                raise ValueError("克重必须为1至2000的整数")
            key = (supplier, code)
            if system_id is not None and system_id in seen_ids:
                raise ValueError(f"系统ID {system_id} 在表内重复")
            if key in seen_keys:
                raise ValueError(f"供应商 {supplier} 的基础代码 {code} 在表内重复")
            seen_ids.add(system_id) if system_id is not None else None
            seen_keys.add(key)
            rows.append(
                {
                    "row_number": row_number,
                    "system_id": system_id,
                    "supplier_name": supplier,
                    "code_char": code,
                    "paper_name": paper_name,
                    "gram_weight": gram_weight,
                    "paper_grade": _cell_text(values[5]) or None,
                    "paper_role": _cell_text(values[6]) or None,
                    "remark": _cell_text(values[7]) or None,
                    "is_active": _enabled(values[8]),
                }
            )
        except ValueError as error:
            _append_error(
                errors,
                sheet=PAPER_SHEET,
                row_number=row_number,
                message=str(error),
            )
    return rows


def _parse_material_rows(worksheet, errors: list[dict[str, Any]]) -> list[dict]:
    rows: list[dict] = []
    seen_ids: set[int] = set()
    seen_codes: set[str] = set()
    for row_number, cells in enumerate(
        worksheet.iter_rows(min_row=2, max_col=len(MATERIAL_HEADERS)),
        start=2,
    ):
        values = [cell.value for cell in cells]
        if not any(value is not None and _cell_text(value) != "" for value in values):
            continue
        try:
            system_id = _positive_int(values[0], "系统ID", optional=True)
            file_version = _positive_int(values[1], "当前版本", optional=True)
            supplier = normalize_supplier_name(values[2])
            if not supplier:
                raise ValueError("供应商标准名不能为空")
            code = _cell_text(values[3]).upper()
            if not code:
                raise ValueError("组合代码不能为空")
            layer_count = _positive_int(values[4], "层数")
            if layer_count not in {3, 5, 7}:
                raise ValueError("层数只能填写3、5或7")
            if system_id is None and not re.fullmatch(
                rf"[A-Z0-9]{{{layer_count}}}", code
            ):
                raise ValueError(
                    f"新增组合代码必须是{layer_count}位字母或数字"
                )
            if system_id is not None and system_id in seen_ids:
                raise ValueError(f"系统ID {system_id} 在表内重复")
            if code in seen_codes:
                raise ValueError(f"组合代码 {code} 在表内重复")
            seen_ids.add(system_id) if system_id is not None else None
            seen_codes.add(code)
            rows.append(
                {
                    "row_number": row_number,
                    "system_id": system_id,
                    "file_version": file_version,
                    "supplier_name": supplier,
                    "code": code,
                    "layer_count": layer_count,
                    "quote_price": _price(values[5]),
                    "quote_date": _date(values[6]),
                    "price_unit": _cell_text(values[7]) or "元/㎡",
                    "remarks": _cell_text(values[8]) or None,
                    "is_active": _enabled(values[9]),
                }
            )
        except ValueError as error:
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=str(error),
            )
    return rows


def _read_workbook(content: bytes) -> tuple[list[dict], list[dict], list[dict]]:
    _preflight_xlsx_container(content)
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as error:
        raise SupplierMaterialWorkbookError(
            "SUPPLIER_MATERIAL_WORKBOOK_INVALID",
            "Excel 工作簿无法安全读取",
        ) from error
    errors: list[dict[str, Any]] = []
    try:
        expected_sheets = {PAPER_SHEET, MATERIAL_SHEET}
        if set(workbook.sheetnames) != expected_sheets:
            raise SupplierMaterialWorkbookError(
                "SUPPLIER_MATERIAL_SHEETS_INVALID",
                "Excel 必须且只能包含“基础纸种”和“组合材质”两张工作表",
            )
        paper_sheet = workbook[PAPER_SHEET]
        material_sheet = workbook[MATERIAL_SHEET]
        if paper_sheet.max_row > MAX_ROWS_PER_SHEET + 1:
            raise SupplierMaterialWorkbookError(
                "SUPPLIER_MATERIAL_ROWS_EXCEEDED",
                f"“基础纸种”最多允许 {MAX_ROWS_PER_SHEET} 行数据",
            )
        if material_sheet.max_row > MAX_ROWS_PER_SHEET + 1:
            raise SupplierMaterialWorkbookError(
                "SUPPLIER_MATERIAL_ROWS_EXCEEDED",
                f"“组合材质”最多允许 {MAX_ROWS_PER_SHEET} 行数据",
            )
        _validate_headers(paper_sheet, PAPER_HEADERS)
        _validate_headers(material_sheet, MATERIAL_HEADERS)
        paper_rows = _parse_paper_rows(paper_sheet, errors)
        material_rows = _parse_material_rows(material_sheet, errors)
        if not paper_rows and not material_rows and not errors:
            errors.append(
                {
                    "sheet": PAPER_SHEET,
                    "row_number": 2,
                    "message": "工作簿没有可导入的数据",
                }
            )
        return paper_rows, material_rows, errors
    finally:
        workbook.close()


def _resolve_plan(
    db: Session,
    *,
    paper_rows: list[dict],
    material_rows: list[dict],
    errors: list[dict],
) -> tuple[list[dict], list[dict], tuple[dict[str, Any], ...], dict[str, int]]:
    current_papers = list(db.scalars(select(SupplierPaperCode)).all())
    papers_by_id = {row.id: row for row in current_papers}
    papers_by_key = {
        (row.supplier_name, row.code_char.upper()): row
        for row in current_papers
    }
    current_materials = list(db.scalars(select(Material)).all())
    materials_by_id = {row.id: row for row in current_materials}
    materials_by_code = {row.code.upper(): row for row in current_materials}

    def existing_material_for(raw: dict[str, Any]) -> Material | None:
        if raw["system_id"] is not None:
            return materials_by_id.get(raw["system_id"])
        return materials_by_code.get(raw["code"])

    def historical_supplier_is_unchanged(
        raw: dict[str, Any],
        existing: Material | None,
    ) -> bool:
        if existing is None or not existing.supplier_name:
            return False
        if normalize_supplier_identity(raw["supplier_name"]) == (
            normalize_supplier_identity(existing.supplier_name)
        ):
            return True
        try:
            requested = resolve_supplier(
                db,
                raw["supplier_name"],
                require_active=False,
            )
            stored = resolve_supplier(
                db,
                existing.supplier_name,
                require_active=False,
            )
        except SupplierLookupError:
            return False
        return requested.id == stored.id

    def historical_paper_supplier_is_unchanged(
        raw: dict[str, Any],
    ) -> bool:
        if raw["system_id"] is None:
            return False
        existing = papers_by_id.get(raw["system_id"])
        if existing is None:
            return False
        return normalize_supplier_identity(raw["supplier_name"]) == (
            normalize_supplier_identity(existing.supplier_name)
        )

    supplier_references: dict[int, dict[str, Any]] = {}
    for sheet_name, rows in (
        (PAPER_SHEET, paper_rows),
        (MATERIAL_SHEET, material_rows),
    ):
        for raw in rows:
            try:
                supplier = resolve_supplier(
                    db,
                    raw["supplier_name"],
                    require_active=True,
                )
            except SupplierLookupError as error:
                if (
                    sheet_name == PAPER_SHEET
                    and historical_paper_supplier_is_unchanged(raw)
                ):
                    existing = papers_by_id[raw["system_id"]]
                    raw["supplier_name"] = existing.supplier_name
                    raw["_historical_supplier_unchanged"] = True
                    continue
                existing = (
                    existing_material_for(raw)
                    if sheet_name == MATERIAL_SHEET
                    else None
                )
                if historical_supplier_is_unchanged(raw, existing):
                    # Historical rows may keep an inactive or not-yet-mastered
                    # supplier while other fields are maintained. This does not
                    # authorize creating a row or changing its supplier.
                    raw["supplier_name"] = existing.supplier_name
                    raw["_historical_supplier_unchanged"] = True
                    continue
                raw["_supplier_invalid"] = True
                _append_error(
                    errors,
                    sheet=sheet_name,
                    row_number=raw["row_number"],
                    message=error.message,
                )
                continue
            raw["supplier_name"] = supplier.standard_name
            raw["supplier_id"] = supplier.id
            raw["supplier_version"] = supplier.version
            supplier_references[supplier.id] = {
                "id": supplier.id,
                "standard_name": supplier.standard_name,
                "version": supplier.version,
            }

    paper_items: list[dict] = []
    resolved_paper_ids: set[int] = set()
    planned_paper_keys: set[tuple[str, str]] = set()

    for raw in paper_rows:
        if raw.get("_supplier_invalid"):
            continue
        row_number = raw["row_number"]
        target_key = (raw["supplier_name"], raw["code_char"])
        if target_key in planned_paper_keys:
            _append_error(
                errors,
                sheet=PAPER_SHEET,
                row_number=row_number,
                message=(
                    f"供应商 {target_key[0]} 的基础代码 {target_key[1]}"
                    " 在表内重复"
                ),
            )
            continue
        planned_paper_keys.add(target_key)
        existing = (
            papers_by_id.get(raw["system_id"])
            if raw["system_id"] is not None
            else papers_by_key.get(target_key)
        )
        if raw["system_id"] is not None and existing is None:
            _append_error(
                errors,
                sheet=PAPER_SHEET,
                row_number=row_number,
                message=f"系统ID {raw['system_id']} 不存在，请重新导出模板",
            )
            continue
        if existing is not None and existing.id in resolved_paper_ids:
            _append_error(
                errors,
                sheet=PAPER_SHEET,
                row_number=row_number,
                message=(
                    f"系统基础纸种 ID {existing.id} 在导入计划中被重复引用"
                ),
            )
            continue
        conflict = papers_by_key.get(target_key)
        if conflict is not None and (existing is None or conflict.id != existing.id):
            _append_error(
                errors,
                sheet=PAPER_SHEET,
                row_number=row_number,
                message=(
                    f"供应商 {target_key[0]} 已存在基础代码 {target_key[1]}"
                ),
            )
            continue
        if existing is not None:
            resolved_paper_ids.add(existing.id)
        updates = {
            key: raw[key]
            for key in (
                "supplier_name",
                "code_char",
                "paper_name",
                "gram_weight",
                "paper_grade",
                "paper_role",
                "remark",
                "is_active",
            )
        }
        before = _paper_snapshot(existing) if existing else None
        changed = existing is None or any(
            getattr(existing, key) != value for key, value in updates.items()
        )
        paper_items.append(
            {
                "row_number": row_number,
                "action": "create" if existing is None else "update",
                "changed": changed,
                "existing_id": existing.id if existing else None,
                "before_fingerprint": _fingerprint(before) if before else None,
                "updates": updates,
            }
        )

    available_papers: dict[tuple[str, str], dict[str, Any]] = {
        (row.supplier_name, row.code_char.upper()): {
            "paper_name": row.paper_name,
            "gram_weight": row.gram_weight,
            "is_active": row.is_active,
        }
        for row in current_papers
    }
    for item in paper_items:
        updates = item["updates"]
        if item["existing_id"] is not None:
            current = papers_by_id[item["existing_id"]]
            available_papers.pop(
                (current.supplier_name, current.code_char.upper()),
                None,
            )
        available_papers[(updates["supplier_name"], updates["code_char"])] = {
            "paper_name": updates["paper_name"],
            "gram_weight": updates["gram_weight"],
            "is_active": updates["is_active"],
        }

    material_items: list[dict] = []
    target_codes: dict[str, int | None] = {}
    resolved_material_ids: set[int] = set()

    for raw in material_rows:
        if raw.get("_supplier_invalid"):
            continue
        row_number = raw["row_number"]
        existing = (
            materials_by_id.get(raw["system_id"])
            if raw["system_id"] is not None
            else materials_by_code.get(raw["code"])
        )
        if raw["system_id"] is not None and existing is None:
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=f"系统ID {raw['system_id']} 不存在，请重新导出模板",
            )
            continue
        if existing is not None and existing.id in resolved_material_ids:
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=f"系统材质 ID {existing.id} 在导入计划中被重复引用",
            )
            continue
        if (
            raw["system_id"] is None
            and existing is not None
            and (existing.supplier_name or "") != raw["supplier_name"]
        ):
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=(
                    f"组合代码 {raw['code']} 已被供应商"
                    f"“{existing.supplier_name or '未设置'}”使用；"
                    "当前系统组合代码全局唯一，不能跨供应商重复"
                ),
            )
            continue
        if existing is not None and raw["file_version"] not in {
            None,
            existing.version,
        }:
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=(
                    f"材质 {existing.code} 当前为 v{existing.version}，"
                    f"模板仍是 v{raw['file_version']}，请重新导出"
                ),
            )
            continue
        conflict = materials_by_code.get(raw["code"])
        if conflict is not None and (existing is None or conflict.id != existing.id):
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=(
                    f"组合代码 {raw['code']} 已被供应商"
                    f"“{conflict.supplier_name or '未设置'}”使用；"
                    "当前系统组合代码全局唯一，不能跨供应商重复"
                ),
            )
            continue
        if existing is not None and existing.code.upper() != raw["code"]:
            if not re.fullmatch(
                rf"[A-Z0-9]{{{raw['layer_count']}}}",
                raw["code"],
            ):
                _append_error(
                    errors,
                    sheet=MATERIAL_SHEET,
                    row_number=row_number,
                    message=(
                        f"修改后的组合代码必须是{raw['layer_count']}位字母或数字"
                    ),
                )
                continue
        if raw["code"] in target_codes and target_codes[raw["code"]] != (
            existing.id if existing else None
        ):
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=f"组合代码 {raw['code']} 在导入计划中冲突",
            )
            continue
        target_codes[raw["code"]] = existing.id if existing else None
        if existing is not None:
            resolved_material_ids.add(existing.id)

        compact_code = re.sub(r"[^A-Z0-9]", "", raw["code"])[
            : raw["layer_count"]
        ]
        if len(compact_code) != raw["layer_count"]:
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=(
                    f"组合代码无法解析为{raw['layer_count']}位基础代码"
                ),
            )
            continue
        roles = (
            ["面纸", "第一楞纸", "第一芯纸", "第二楞纸", "第二芯纸", "第三楞纸", "里纸"]
            if raw["layer_count"] == 7
            else ["面纸", "瓦楞纸", "里纸"]
            if raw["layer_count"] == 3
            else ["面纸", "B楞瓦纸", "芯纸", "A楞瓦纸", "里纸"]
        )
        layers: list[tuple[str, dict[str, Any], str]] = []
        missing: list[str] = []
        for code_char, role in zip(compact_code, roles):
            paper = available_papers.get((raw["supplier_name"], code_char))
            if paper is None or not paper["is_active"]:
                if code_char not in missing:
                    missing.append(code_char)
            else:
                layers.append((code_char, paper, role))
        if missing:
            _append_error(
                errors,
                sheet=MATERIAL_SHEET,
                row_number=row_number,
                message=(
                    f"供应商 {raw['supplier_name']} 缺少启用的基础代码："
                    + "、".join(missing)
                ),
            )
            continue
        composition = " | ".join(
            (
                f"{role}:{code_char}={paper['gram_weight']}g "
                f"{paper['paper_name']}"
            )
            for code_char, paper, role in layers
        )
        weights = "/".join(
            f"{paper['gram_weight']}g" for _, paper, _ in layers
        )
        updates = {
            "code": raw["code"],
            "paper_composition": composition,
            "layer_count": raw["layer_count"],
            "flute_type": None,
            "basis_weight_description": weights,
            "quote_price": raw["quote_price"],
            "price_source": "供应商材质 Excel 批量维护",
            "price_unit": raw["price_unit"],
            "supplier_name": raw["supplier_name"],
            "quote_date": raw["quote_date"],
            "remarks": raw["remarks"],
            "is_active": raw["is_active"],
        }
        changed = existing is None or any(
            getattr(existing, key) != value for key, value in updates.items()
        )
        material_items.append(
            {
                "row_number": row_number,
                "action": "create" if existing is None else "update",
                "changed": changed,
                "existing_id": existing.id if existing else None,
                "expected_version": existing.version if existing else None,
                "updates": updates,
            }
        )

    summary = {
        "paper_create": sum(
            item["action"] == "create" for item in paper_items
        ),
        "paper_update": sum(
            item["action"] == "update" and item["changed"]
            for item in paper_items
        ),
        "paper_unchanged": sum(not item["changed"] for item in paper_items),
        "material_create": sum(
            item["action"] == "create" for item in material_items
        ),
        "material_update": sum(
            item["action"] == "update" and item["changed"]
            for item in material_items
        ),
        "material_unchanged": sum(
            not item["changed"] for item in material_items
        ),
    }
    return (
        paper_items,
        material_items,
        tuple(
            supplier_references[key]
            for key in sorted(supplier_references)
        ),
        summary,
    )


async def preview_supplier_material_workbook(
    db: Session,
    *,
    upload: UploadFile,
    user: User,
) -> dict[str, Any]:
    try:
        validated = await read_validated_upload(upload, EXCEL_POLICY)
        paper_rows, material_rows, errors = _read_workbook(validated.content)
    except UploadValidationError as error:
        raise SupplierMaterialWorkbookError(
            "SUPPLIER_MATERIAL_UPLOAD_INVALID",
            str(error),
        ) from error
    paper_items, material_items, supplier_references, summary = _resolve_plan(
        db,
        paper_rows=paper_rows,
        material_rows=material_rows,
        errors=errors,
    )
    result: dict[str, Any] = {
        "valid": not errors,
        "filename": validated.original_filename,
        "sha256": validated.sha256,
        "summary": summary,
        "errors": errors,
        "paper_codes": [
            {
                "row_number": item["row_number"],
                "action": item["action"],
                "changed": item["changed"],
                **item["updates"],
            }
            for item in paper_items
        ],
        "materials": [
            {
                "row_number": item["row_number"],
                "action": item["action"],
                "changed": item["changed"],
                "expected_version": item["expected_version"],
                **item["updates"],
            }
            for item in material_items
        ],
    }
    if errors:
        return result
    preview = _WorkbookPreview(
        actor=_actor(user),
        expires_at=time.monotonic() + PREVIEW_TTL_SECONDS,
        source_sha256=validated.sha256,
        paper_items=tuple(paper_items),
        material_items=tuple(material_items),
        supplier_references=supplier_references,
        summary=summary,
    )
    result["preview_token"] = _store_preview(preview)
    result["expires_in_seconds"] = PREVIEW_TTL_SECONDS
    return result


def _material_update_with_confirmation(
    db: Session,
    *,
    material: Material,
    updates: dict[str, Any],
    expected_version: int,
    user: User,
) -> None:
    kwargs = {
        "object_type": "material",
        "entity": material,
        "updates": updates,
        "expected_version": expected_version,
        "user": user,
        "reason": "供应商材质 Excel 批量维护",
        "source": "api.materials.supplier_workbook.apply",
        "action": "system_mapping",
    }
    try:
        apply_versioned_update(db, **kwargs)
    except HTTPException as error:
        detail = error.detail if isinstance(error.detail, dict) else {}
        if detail.get("code") != "MASTER_CHANGE_CONFIRMATION_REQUIRED":
            raise
        apply_versioned_update(
            db,
            **kwargs,
            confirmation_token=detail.get("confirmation_token"),
        )


def apply_supplier_material_workbook(
    db: Session,
    *,
    preview_token: str,
    user: User,
) -> dict[str, Any]:
    preview = _take_preview(preview_token, user)

    # Fail closed before the first write. This also guarantees that a stale
    # material version cannot leave earlier paper-code updates behind.
    for reference in preview.supplier_references:
        supplier = db.get(Supplier, reference["id"])
        if (
            supplier is None
            or not supplier.is_active
            or supplier.version != reference["version"]
            or supplier.standard_name != reference["standard_name"]
        ):
            raise SupplierMaterialWorkbookError(
                "SUPPLIER_MATERIAL_SUPPLIER_STALE",
                (
                    f"供应商“{reference['standard_name']}”已停用或资料已变化，"
                    "请重新导入预览"
                ),
                status_code=409,
            )

    for item in preview.paper_items:
        existing_id = item["existing_id"]
        updates = item["updates"]
        if existing_id is None:
            conflict = db.scalar(
                select(SupplierPaperCode.id).where(
                    SupplierPaperCode.supplier_name == updates["supplier_name"],
                    SupplierPaperCode.code_char == updates["code_char"],
                )
            )
            if conflict is not None:
                raise SupplierMaterialWorkbookError(
                    "SUPPLIER_MATERIAL_PREVIEW_STALE",
                    "基础纸种已发生变化，请重新导入预览",
                    status_code=409,
                )
        else:
            current = db.get(SupplierPaperCode, existing_id)
            if current is None or _fingerprint(_paper_snapshot(current)) != item[
                "before_fingerprint"
            ]:
                raise SupplierMaterialWorkbookError(
                    "SUPPLIER_MATERIAL_PREVIEW_STALE",
                    "基础纸种已发生变化，请重新导入预览",
                    status_code=409,
                )

    for item in preview.material_items:
        existing_id = item["existing_id"]
        updates = item["updates"]
        conflict_query = select(Material).where(
            func.upper(Material.code) == updates["code"].upper()
        )
        if existing_id is not None:
            conflict_query = conflict_query.where(Material.id != existing_id)
        conflict = db.scalar(conflict_query)
        if conflict is not None:
            raise SupplierMaterialWorkbookError(
                "SUPPLIER_MATERIAL_GLOBAL_CODE_CONFLICT",
                (
                    f"组合代码 {updates['code']} 已被供应商"
                    f"“{conflict.supplier_name or '未设置'}”使用"
                ),
                status_code=409,
            )
        if existing_id is not None:
            current = db.get(Material, existing_id)
            if current is None or current.version != item["expected_version"]:
                raise SupplierMaterialWorkbookError(
                    "SUPPLIER_MATERIAL_PREVIEW_STALE",
                    "材质版本已发生变化，请重新导入预览",
                    status_code=409,
                )

    paper_created = paper_updated = 0
    material_created = material_updated = 0
    for item in preview.paper_items:
        updates = item["updates"]
        if item["existing_id"] is None:
            row = SupplierPaperCode(**updates)
            db.add(row)
            db.flush()
            audit_master_change(
                db,
                user=user,
                action="CREATE",
                resource="SupplierPaperCode",
                resource_id=row.id,
                details={
                    "source": "supplier_material_workbook",
                    "source_sha256": preview.source_sha256,
                    "after": updates,
                },
            )
            paper_created += 1
        elif item["changed"]:
            row = db.get(SupplierPaperCode, item["existing_id"])
            before = _paper_snapshot(row)
            for key, value in updates.items():
                setattr(row, key, value)
            db.flush()
            audit_master_change(
                db,
                user=user,
                action="UPDATE",
                resource="SupplierPaperCode",
                resource_id=row.id,
                details={
                    "source": "supplier_material_workbook",
                    "source_sha256": preview.source_sha256,
                    "before": before,
                    "after": updates,
                },
            )
            paper_updated += 1

    for item in preview.material_items:
        updates = item["updates"]
        if item["existing_id"] is None:
            material = Material(**updates)
            db.add(material)
            db.flush()
            record_versioned_create(
                db,
                object_type="material",
                entity=material,
                user=user,
                reason="供应商材质 Excel 批量维护",
                source="api.materials.supplier_workbook.apply",
            )
            audit_master_change(
                db,
                user=user,
                action="CREATE",
                resource="Material",
                resource_id=material.id,
                details={
                    "source": "supplier_material_workbook",
                    "source_sha256": preview.source_sha256,
                    "after": updates,
                },
            )
            material_created += 1
        elif item["changed"]:
            material = db.get(Material, item["existing_id"])
            before = {
                key: getattr(material, key) for key in updates
            }
            _material_update_with_confirmation(
                db,
                material=material,
                updates=updates,
                expected_version=item["expected_version"],
                user=user,
            )
            if before["quote_price"] != updates["quote_price"]:
                db.add(
                    MaterialPriceHistory(
                        material_id=material.id,
                        supplier_name=updates["supplier_name"],
                        material_code=updates["code"],
                        old_price=before["quote_price"],
                        new_price=updates["quote_price"],
                        effective_date=updates["quote_date"],
                        adjust_reason="供应商材质 Excel 批量维护",
                        operator=user.username,
                    )
                )
            audit_master_change(
                db,
                user=user,
                action="UPDATE",
                resource="Material",
                resource_id=material.id,
                details={
                    "source": "supplier_material_workbook",
                    "source_sha256": preview.source_sha256,
                    "before": before,
                    "after": updates,
                },
            )
            material_updated += 1

    db.flush()
    _consume_preview(preview_token)
    return {
        "ok": True,
        "source_sha256": preview.source_sha256,
        "paper_created": paper_created,
        "paper_updated": paper_updated,
        "material_created": material_created,
        "material_updated": material_updated,
        "message": (
            f"基础纸种新增 {paper_created}、更新 {paper_updated}；"
            f"组合材质新增 {material_created}、更新 {material_updated}"
        ),
    }
