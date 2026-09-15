"""Read-only internal trace and supplier-only export projection.

Reservations are order-item current facts, not an invented purchase-time snapshot.
Exports deliberately whitelist supplier fields and never serialize internal sources.
"""
from io import BytesIO
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from sqlalchemy import select
from sqlalchemy.orm import joinedload
from app.models.order import Order, OrderItem
from app.models.warehouse_inventory import InventoryLot, InventoryReservation, WarehouseLocation, FinishedGoodsInventoryDetail


def internal_trace(db, purchase):
    sources = [dict(source, report_length_mm=line.get('report_length_mm'),
                    report_width_mm=line.get('report_width_mm'))
               for line in purchase.get('lines', []) for source in line.get('source_items', [])]
    ids = {row['order_item_id'] for row in sources if row.get('order_item_id')}
    customer_pos = dict(db.execute(select(OrderItem.id, Order.customer_po).join(
        Order, Order.id == OrderItem.order_id).where(OrderItem.id.in_(ids))).all()) if ids else {}
    records = {}
    if ids:
        rows = db.execute(select(InventoryReservation, InventoryLot, WarehouseLocation)
            .options(joinedload(InventoryLot.semi_finished_detail), joinedload(InventoryLot.finished_detail).joinedload(FinishedGoodsInventoryDetail.product))
            .join(InventoryLot, InventoryLot.id == InventoryReservation.inventory_lot_id)
            .join(WarehouseLocation, WarehouseLocation.id == InventoryLot.warehouse_location_id)
            .where(InventoryReservation.order_item_id.in_(ids),
                   InventoryReservation.status.in_(['active', 'partial', 'consumed']))
            .order_by(InventoryReservation.id)).all()
        for reservation, lot, location in rows:
            consumed = int(reservation.consumed_stock_quantity or 0)
            pending = max(int(reservation.reserved_stock_quantity or 0) - consumed
                          - int(reservation.released_stock_quantity or 0), 0)
            if not pending and not consumed:
                continue
            semi = lot.semi_finished_detail
            kind = 'finished' if lot.inventory_type == 'finished' else (
                'raw' if semi and semi.sheet_type == 'raw_board' else 'semi')
            detail = lot.finished_detail
            records.setdefault(reservation.order_item_id, []).append(dict(
                id=reservation.id, lot_number=lot.lot_number, kind=kind,
                kind_label={'finished':'成品','semi':'半成品','raw':'原料'}[kind],
                product_code=detail.inventory_code_snapshot if detail else None,
                pending_qty=pending, consumed_qty=consumed,
                unit='张' if lot.unit == 'sheets' else ((detail.product.unit if detail and detail.product else None) or '只'),
                location_id=location.id,
                location=f"{location.warehouse_floor or ''}楼 · {location.location_name}",
                location_active=location.is_active,
                bom_component_id=reservation.sales_order_item_bom_component_id))
    return {'scope':'current_order_item', 'location_basis':'current_lot_location', 'sources':[
        dict(source_id=row.get('id'), order_item_id=row.get('order_item_id'),
             customer_po=customer_pos.get(row.get('order_item_id')),
             product_code=row.get('product_code'),
             report_length_mm=row.get('report_length_mm'), report_width_mm=row.get('report_width_mm'),
             requisition_qty=row.get('requisition_qty'),
             order_purpose_sheet_qty=row.get('order_purpose_sheet_qty'),
             stock_purpose_sheet_qty=row.get('stock_purpose_sheet_qty'),
             inventory_records=records.get(row.get('order_item_id'), [])) for row in sources]}


def supplier_rows(purchase):
    return [[index + 1, f"{line.get('report_length_mm') or ''}×{line.get('report_width_mm') or ''}",
             '' if line.get('crease_display') == '-' else line.get('crease_display') or '',
             line.get('material_display') or line.get('material_code') or '',
             int(line.get('requisition_qty') or 0), line.get('remark') or '']
            for index, line in enumerate(purchase.get('lines', []))]


HEADERS = ['序号', '纸板长宽（mm）', '压线（mm）', '材质/楞型', '数量（张）', '报料备注']


def purchase_date(purchase):
    from app.core.time_contract import BEIJING_TIMEZONE
    value = str(purchase.get('created_at') or '')
    if not value or len(value) == 10:
        return value
    parsed = datetime.fromisoformat(value.replace('Z', '+00:00'))
    return parsed.replace(tzinfo=parsed.tzinfo or timezone.utc).astimezone(BEIJING_TIMEZONE).date().isoformat()


def supplier_pdf(purchase):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from app.services.contract_pdf import _registered_font
    font, _ = _registered_font()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=12*mm, rightMargin=12*mm,
                            topMargin=12*mm, bottomMargin=14*mm)
    style = ParagraphStyle('supplier', fontName=font, fontSize=10, leading=14, wordWrap='CJK')
    title = ParagraphStyle('supplier_title', parent=style, fontSize=17, leading=22, alignment=1)
    def p(value):
        return Paragraph(escape(str(value if value is not None else '')).replace('\n','<br/>'), style)
    sender = purchase.get('sender') or {}
    story = [Paragraph(escape(sender.get('company_name') or '苏州工业园区天明纸品包装厂'), title),
             Paragraph('采购单', title), Spacer(1, 5*mm),
             p(f"供应商：{purchase.get('supplier_name') or ''}"),
             p(f"单号：{purchase.get('order_number') or ''}    日期：{purchase_date(purchase)}"), Spacer(1, 3*mm)]
    table = Table([[p(v) for v in row] for row in [HEADERS, *supplier_rows(purchase)]],
                  colWidths=[12*mm, 37*mm, 36*mm, 36*mm, 20*mm, 45*mm], repeatRows=1)
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.black),
        ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),
        ('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),6),
        ('BOTTOMPADDING',(0,0),(-1,-1),6)]))
    story.extend([table, Spacer(1, 4*mm), p(f"地址：{sender.get('address') or ''}    电话：{sender.get('phone') or ''}")])
    def footer(canvas, document):
        canvas.setFont(font,9); canvas.drawRightString(198*mm,8*mm,f"第 {document.page} 页")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def supplier_xlsx(purchase):
    # Use the ERP's existing production dependency, not an external desktop tool.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side, Font
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title='采购单'
    sender = purchase.get('sender') or {}
    for row in [[sender.get('company_name') or '苏州工业园区天明纸品包装厂'], ['采购单'],
                [f"供应商：{purchase.get('supplier_name') or ''}"],
                [f"单号：{purchase.get('order_number') or ''}    日期：{purchase_date(purchase)}"], HEADERS, *supplier_rows(purchase)]:
        ws.append(row)
    for row in range(1,5): ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=6)
    widths=[8,24,26,24,14,40]
    border=Border(*( [Side(style='thin',color='999999')]*4 ))
    for index,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(index)].width=width
    for row in ws:
        for cell in row:
            if isinstance(cell.value,str): cell.data_type='s'  # Formula injection protection, including remarks.
            cell.font=Font(name='宋体',size=11,bold=cell.row<=5)
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            if cell.row>=5: cell.border=border
        ws.row_dimensions[row[0].row].height=max(26, 16*max((len(str(c.value or ''))//max(1,int(widths[i]/2))+1 for i,c in enumerate(row))))
    for row in (1,2): ws.cell(row,1).alignment=Alignment(horizontal='center'); ws.cell(row,1).font=Font(name='宋体',size=16,bold=True)
    ws.append([f"地址：{sender.get('address') or ''}    电话：{sender.get('phone') or ''}"])
    ws.merge_cells(start_row=ws.max_row,start_column=1,end_row=ws.max_row,end_column=6)
    ws.freeze_panes='A6'; ws.print_title_rows='1:5'; ws.print_options.horizontalCentered=True
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='portrait'; ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
    ws.print_area=f'A1:F{ws.max_row}'
    output=BytesIO(); wb.save(output); return output.getvalue()
