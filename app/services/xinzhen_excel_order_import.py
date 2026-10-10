"""New Zhen carton-marking Excel parsing and immutable source audit support.

Parsing produces editable dictionaries only.  The preview API may persist an
immutable source hash/row ledger, but it never creates products or formal
orders; formal order creation remains the explicit responsibility of
``POST /api/orders`` after a signed confirmation round trip.
"""

from __future__ import annotations

from collections.abc import Iterable
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO
import hashlib
import json
import re
import zipfile

from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.customer import Customer
from app.models.excel_order_import import ExcelOrderImportBatch, ExcelOrderImportRow
from app.models.product import Product


_DIMENSIONS_RE = re.compile(
    r"(?P<length>\d+(?:\.\d+)?)\s*[x×*]\s*"
    r"(?P<width>\d+(?:\.\d+)?)\s*[x×*]\s*"
    r"(?P<height>\d+(?:\.\d+)?)\s*(?P<unit>cm|mm)?",
    re.IGNORECASE,
)
_NUMBER_RE = re.compile(r"^\s*(\d+(?:\.\d+)?)\s*$")
_TOTAL_MARKERS = ("total", "grand total", "合计", "总计")
MAX_XLSX_ZIP_ENTRIES = 2_000
MAX_XLSX_ENTRY_BYTES = 20 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_WORKBOOK_SHEETS = 20
MAX_WORKBOOK_ROWS = 5_000
MAX_WORKBOOK_COLUMNS = 256
_GENERIC_REQUIRED_HEADERS = {
    "vendor_style",
    "units_per_carton",
    "carton_count",
    "garment_total",
}
XINZHEN_CUSTOMER_CODES = frozenset({"XINZHEN"})
XINZHEN_EXCEL_PARSER_VERSION = "n042-xinzhen-v2.2"


class XinzhenExcelParseError(ValueError):
    """The uploaded workbook is not a usable New Zhen carton-marking sheet."""


@dataclass(frozen=True)
class _SheetRow:
    number: int
    values: tuple[object, ...]


@dataclass(frozen=True)
class _WorkbookSheet:
    name: str
    rows: tuple[_SheetRow, ...]


class XinzhenWorksheetSelectionRequired(XinzhenExcelParseError):
    """More than one visible worksheet contains order evidence."""

    def __init__(self, sheet_names: list[str]) -> None:
        self.sheet_names = sheet_names
        super().__init__(
            "多个可见工作表包含订单证据，请明确选择一张工作表后重新预览："
            + "、".join(sheet_names)
        )


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value: object) -> str:
    return re.sub(r"[\s_\-()（）:：./]+", "", _text(value).casefold())


def _header_field(value: object) -> str | None:
    key = _header_key(value)
    if not key:
        return None
    aliases = {
        "customer_po": ("po", "pono", "ponumber", "订单号", "采购单号"),
        "vendor_style": ("vendorstyle", "style", "styleno", "款号", "款式"),
        "color": ("colour", "color", "颜色", "色号"),
        "size": ("size", "尺码", "尺寸"),
        "product_code": ("productcode", "itemcode", "materialcode", "产品编码", "物料编码"),
        "carton_dimensions": (
            "cartondimensions",
            "cartondimension",
            "cartonsize",
            "cartonmeasurement",
            "ctnsize",
            "外箱尺寸",
            "纸箱尺寸",
        ),
        "units_per_carton": (
            "pcspercarton",
            "pcscarton",
            "pcsctn",
            "piecespercarton",
            "unitspercarton",
            "qtypercarton",
            "每箱数量",
            "每箱件数",
        ),
        "carton_count": (
            "cartoncount",
            "cartonqty",
            "cartonquantity",
            "numberofcartons",
            "cartons",
            "箱数",
            "箱量",
        ),
        "garment_total": (
            "totalpcs",
            "totalqty",
            "totalquantity",
            "garmenttotal",
            "garmentqty",
            "总件数",
            "总数量",
        ),
    }
    for field, candidates in aliases.items():
        if key in candidates:
            return field
    return None


def _decimal(value: object) -> Decimal | None:
    match = _NUMBER_RE.match(_text(value))
    if not match:
        return None
    try:
        return Decimal(match.group(1))
    except InvalidOperation:
        return None


def _json_number(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _parse_dimensions(value: object) -> tuple[dict[str, int | float | str], str] | None:
    raw = _text(value)
    match = _DIMENSIONS_RE.search(raw)
    if not match:
        return None
    unit = (match.group("unit") or "cm").casefold()
    multiplier = Decimal("10") if unit == "cm" else Decimal("1")
    values = [Decimal(match.group(name)) * multiplier for name in ("length", "width", "height")]
    if any(value != value.to_integral_value() for value in values):
        return None
    return (
        {
            "length_mm": int(values[0]),
            "width_mm": int(values[1]),
            "height_mm": int(values[2]),
            "source_unit": unit,
        },
        raw,
    )


def _load_carton_marking_sheets(
    content: bytes, filename: str
) -> list[_WorkbookSheet]:
    suffix = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if suffix == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:  # pragma: no cover - requirements guard
            raise XinzhenExcelParseError("缺少 openpyxl，无法读取 .xlsx 文件") from error
        if not zipfile.is_zipfile(BytesIO(content)):
            raise XinzhenExcelParseError("上传的 .xlsx 文件结构无效")
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > MAX_XLSX_ZIP_ENTRIES:
                    raise XinzhenExcelParseError("Excel 文件内部条目过多，已拒绝读取")
                if any(member.flag_bits & 0x1 for member in members):
                    raise XinzhenExcelParseError("不支持加密的 Excel 文件")
                if any(member.file_size > MAX_XLSX_ENTRY_BYTES for member in members):
                    raise XinzhenExcelParseError("Excel 文件内部单项内容过大，已拒绝读取")
                if sum(member.file_size for member in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise XinzhenExcelParseError("Excel 文件解压后内容过大，已拒绝读取")
        except XinzhenExcelParseError:
            raise
        except (OSError, zipfile.BadZipFile) as error:
            raise XinzhenExcelParseError("上传的 .xlsx 文件结构无效") from error
        workbook = None
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            if len(workbook.sheetnames) > MAX_WORKBOOK_SHEETS:
                raise XinzhenExcelParseError("Excel 工作表数量过多，已拒绝读取")
            visible_sheets: list[_WorkbookSheet] = []
            for sheet in workbook.worksheets:
                if sheet.sheet_state != "visible":
                    continue
                if (
                    sheet.max_row > MAX_WORKBOOK_ROWS
                    or sheet.max_column > MAX_WORKBOOK_COLUMNS
                ):
                    raise XinzhenExcelParseError(
                        f"Excel 工作表“{sheet.title}”范围过大，已拒绝读取"
                    )
                rows = tuple(
                    _SheetRow(number=row_number, values=tuple(values))
                    for row_number, values in enumerate(
                        sheet.iter_rows(
                            min_row=1,
                            max_row=max(sheet.max_row, 1),
                            min_col=2,
                            max_col=14,
                            values_only=True,
                        ),
                        start=1,
                    )
                )
                visible_sheets.append(_WorkbookSheet(name=sheet.title, rows=rows))
            if not visible_sheets:
                raise XinzhenExcelParseError("Excel 没有可见工作表，无法读取")
            return visible_sheets
        except XinzhenExcelParseError:
            raise
        except Exception as error:
            raise XinzhenExcelParseError("无法读取 New Zhen .xlsx 文件") from error
        finally:
            if workbook is not None:
                workbook.close()
    if suffix == "xls":
        try:
            import xlrd
        except ImportError as error:  # pragma: no cover - requirements guard
            raise XinzhenExcelParseError("缺少 xlrd，无法读取 .xls 文件") from error
        try:
            workbook = xlrd.open_workbook(file_contents=content)
            if workbook.nsheets > MAX_WORKBOOK_SHEETS:
                raise XinzhenExcelParseError("Excel 工作表数量过多，已拒绝读取")
            visible_sheets: list[_WorkbookSheet] = []
            for sheet in workbook.sheets():
                if int(getattr(sheet, "visibility", 0) or 0) != 0:
                    continue
                if sheet.nrows > MAX_WORKBOOK_ROWS or sheet.ncols > MAX_WORKBOOK_COLUMNS:
                    raise XinzhenExcelParseError(
                        f"Excel 工作表“{sheet.name}”范围过大，已拒绝读取"
                    )
                visible_sheets.append(
                    _WorkbookSheet(
                        name=sheet.name,
                        rows=tuple(
                            _SheetRow(
                                number=row_number,
                                values=tuple(
                                    sheet.cell_value(row_number - 1, column)
                                    if column < sheet.ncols
                                    else None
                                    for column in range(1, 14)
                                ),
                            )
                            for row_number in range(1, max(sheet.nrows, 1) + 1)
                        ),
                    )
                )
            if not visible_sheets:
                raise XinzhenExcelParseError("Excel 没有可见工作表，无法读取")
            return visible_sheets
        except XinzhenExcelParseError:
            raise
        except Exception as error:
            raise XinzhenExcelParseError("无法读取 New Zhen .xls 文件") from error
    raise XinzhenExcelParseError("仅支持 .xls 或 .xlsx 文件")


def _fixed_cell(rows: dict[int, _SheetRow], row_number: int, column: str) -> object:
    """Return a B:N cell from the verified New Zhen carton-marking layout."""

    index = ord(column.upper()) - ord("B")
    row = rows.get(row_number)
    if row is None or index < 0 or index >= len(row.values):
        return None
    return row.values[index]


def _prefixed_value(value: object, prefix: str) -> str | None:
    text = _text(value)
    match = re.search(
        rf"{re.escape(prefix)}\s*#?\s*[:：]?\s*(.+)$", text, re.IGNORECASE
    )
    return match.group(1).strip() if match else None


def _count_from_text(value: object, pattern: str) -> Decimal | None:
    match = re.search(pattern, _text(value), re.IGNORECASE)
    return Decimal(match.group(1)) if match else None


def _is_verified_layout(rows: Iterable[_SheetRow]) -> bool:
    indexed = {row.number: row for row in rows}
    header_values = (
        _header_key(_fixed_cell(indexed, 16, "B")),
        _header_key(_fixed_cell(indexed, 16, "D")),
        _header_key(_fixed_cell(indexed, 16, "F")),
        _header_key(_fixed_cell(indexed, 16, "H")),
    )
    return bool(
        "vendorstyle" in header_values[0]
        and header_values[1] in {"color", "colour"}
        and header_values[2] == "size"
        and header_values[3] == "units"
    )


def _row_has_order_evidence(row: _SheetRow) -> bool:
    values = tuple(row.values)
    mapped_fields = {
        field for value in values if (field := _header_field(value)) is not None
    }
    if len(mapped_fields) >= 2:
        return True
    if any(_parse_dimensions(value) is not None for value in values):
        numeric_values = [value for value in values if _decimal(value) is not None]
        return bool(numeric_values)
    text = " ".join(_text(value) for value in values if _text(value))
    if re.search(r"\bPO\s*#?\s*[:：]?\s*\S+", text, re.IGNORECASE):
        return True
    populated_fixed_fields = sum(
        bool(_text(values[index]))
        for index in (0, 2, 4, 6)
        if index < len(values)
    )
    return populated_fixed_fields >= 3


def _sheet_has_order_evidence(sheet: _WorkbookSheet) -> bool:
    if _is_verified_layout(sheet.rows):
        return True
    header_row, fields = _find_header(sheet.rows)
    return header_row is not None and len(fields) >= 2


def _select_order_sheet(
    sheets: list[_WorkbookSheet], requested_sheet_name: str | None
) -> _WorkbookSheet:
    evidence_sheets = [sheet for sheet in sheets if _sheet_has_order_evidence(sheet)]
    if requested_sheet_name:
        selected = next(
            (sheet for sheet in evidence_sheets if sheet.name == requested_sheet_name),
            None,
        )
        if selected is None:
            raise XinzhenExcelParseError(
                "所选工作表不存在或没有订单证据，请重新选择。"
            )
        return selected
    if len(evidence_sheets) > 1:
        raise XinzhenWorksheetSelectionRequired(
            [sheet.name for sheet in evidence_sheets]
        )
    if len(evidence_sheets) == 1:
        return evidence_sheets[0]
    raise XinzhenExcelParseError(
        "所有可见工作表均未识别到新振订单表，已停止解析。"
    )


def _parse_verified_carton_marking(
    rows: list[_SheetRow], *, filename: str, sheet_name: str
) -> dict | None:
    """Parse the actual New Zhen B10:N30 carton-marking form.

    The garment style table, PO, carton count and carton dimensions live in
    different blocks.  Treating each physical row as an ERP order line would
    create false lines, so this layout deliberately produces one carton line.
    """

    indexed = {row.number: row for row in rows}
    if not _is_verified_layout(rows):
        return None
    overflow_rows = [
        row.number for row in rows if row.number > 30 and _row_has_order_evidence(row)
    ]
    if overflow_rows:
        raise XinzhenExcelParseError(
            "固定版式第31行后仍存在订单证据，疑似截断；"
            f"请整理工作表或明确完整范围后重试（行：{', '.join(map(str, overflow_rows[:8]))}）。"
        )

    styles: list[dict] = []
    for row_number in range(17, 21):
        style = _text(_fixed_cell(indexed, row_number, "B"))
        color = _text(_fixed_cell(indexed, row_number, "D"))
        size = _text(_fixed_cell(indexed, row_number, "F"))
        units = _decimal(_fixed_cell(indexed, row_number, "H"))
        if style or color or size or units is not None:
            styles.append(
                {
                    "source_row": row_number,
                    "vendor_style_code": style or None,
                    "color_raw": color or None,
                    "size_raw": size or None,
                    "units_per_carton": _json_number(units),
                }
            )
    if not styles:
        raise XinzhenExcelParseError("未识别到 Vendor Style 明细")

    customer_po = _prefixed_value(_fixed_cell(indexed, 13, "B"), "PO")
    department_code = _prefixed_value(_fixed_cell(indexed, 14, "B"), "Dept")
    vendor_parts = [
        _prefixed_value(_fixed_cell(indexed, 10, "B"), "From"),
        _text(_fixed_cell(indexed, 11, "C")) or None,
    ]
    vendor_name = " ".join(part for part in vendor_parts if part) or None
    customer_name_raw = _prefixed_value(_fixed_cell(indexed, 10, "F"), "TO")
    ship_to_address = ", ".join(
        value
        for row_number in range(11, 14)
        if (value := _text(_fixed_cell(indexed, row_number, "F")))
    ) or None

    dimensions_result = _parse_dimensions(_fixed_cell(indexed, 30, "L"))
    if dimensions_result is None:
        for row in rows:
            dimensions_result = next(
                (
                    parsed
                    for value in row.values
                    if (parsed := _parse_dimensions(value)) is not None
                ),
                None,
            )
            if dimensions_result is not None:
                break
    dimensions_mm, raw_dimensions = (
        dimensions_result if dimensions_result is not None else (None, "")
    )

    style_units = [
        Decimal(str(style["units_per_carton"]))
        for style in styles
        if style["units_per_carton"] is not None
    ]
    style_units_complete = len(style_units) == len(styles)
    style_units_total = sum(style_units, Decimal("0")) if style_units else None
    units = _decimal(_fixed_cell(indexed, 21, "H"))
    garment_total = _decimal(_fixed_cell(indexed, 29, "F"))
    reported_cartons = _decimal(_fixed_cell(indexed, 23, "H"))
    range_cartons = _count_from_text(_fixed_cell(indexed, 30, "J"), r"#?\s*1\s*(?:到|to|[-~])\s*#?\s*(\d+)")
    dimension_cartons = _count_from_text(_fixed_cell(indexed, 30, "L"), r"(\d+)\s*(?:个|箱|ctns?)")
    derived_cartons = (
        garment_total / units
        if garment_total is not None and units is not None and units > 0
        else None
    )
    signals = [
        value
        for value in (reported_cartons, range_cartons, dimension_cartons, derived_cartons)
        if value is not None
    ]
    carton_count = signals[0] if signals else None
    conflicts: list[str] = []
    if not style_units_complete:
        conflicts.append("款式行 Units 存在空值")
    if style_units_total is None:
        conflicts.append("未识别到款式行 Units")
    if units is None:
        conflicts.append("未识别到 H21 Units 汇总")
    if (
        style_units_total is not None
        and units is not None
        and style_units_total != units
    ):
        conflicts.append("款式行 Units 合计与 H21 汇总不一致")
    for label, value in (
        ("H23 箱数", reported_cartons),
        ("箱号范围", range_cartons),
        ("尺寸箱数", dimension_cartons),
        ("总件数", garment_total),
        ("推导箱数", derived_cartons),
    ):
        if value is None:
            conflicts.append(f"未识别到{label}")
    comparable_cartons = [
        value
        for value in (reported_cartons, range_cartons, dimension_cartons, derived_cartons)
        if value is not None
    ]
    if comparable_cartons and any(
        value != comparable_cartons[0] for value in comparable_cartons[1:]
    ):
        conflicts.append("H23箱数、箱号范围、尺寸箱数或推导箱数不一致")
    reconciled = bool(
        carton_count is not None
        and carton_count > 0
        and carton_count == carton_count.to_integral_value()
        and style_units_total is not None
        and style_units_total > 0
        and style_units_total == units
        and derived_cartons is not None
        and reported_cartons is not None
        and range_cartons is not None
        and dimension_cartons is not None
        and garment_total is not None
        and all(value == carton_count for value in comparable_cartons)
        and not conflicts
    )
    quantity_check = {
        "style_units_total": _json_number(style_units_total),
        "units_per_carton": _json_number(units),
        "garment_total": _json_number(garment_total),
        "reported_carton_count": _json_number(reported_cartons),
        "range_carton_count": _json_number(range_cartons),
        "dimension_carton_count": _json_number(dimension_cartons),
        "derived_carton_count": _json_number(derived_cartons),
        "carton_count": _json_number(carton_count),
        "conflicts": conflicts,
        "status": "passed" if reconciled else "needs_review",
        "message": (
            "款式行Units、H21汇总、总件数、H23箱数、箱号范围和尺寸箱数已交叉核对。"
            if reconciled
            else "固定版式数量证据冲突或不完整，不能转入订单。"
        ),
    }
    item_warnings: list[str] = []
    if dimensions_mm is None:
        item_warnings.append("未识别到纸箱长宽高，不能自动匹配常用箱。")
    if not reconciled:
        item_warnings.append(str(quantity_check["message"]))
    warnings: list[str] = []
    if len(styles) > 1:
        warnings.append(
            "该文件包含多个服装款号，但共用一个纸箱尺寸；已按一条纸箱生产明细生成草稿。"
        )
    return {
        "source_name": filename,
        "source_type": "xinzhen_excel_carton_marking",
        "sheet_name": sheet_name,
        "customer_po": customer_po,
        "department_code": department_code,
        "vendor_name": vendor_name,
        "customer_name_raw": customer_name_raw,
        "ship_to_address": ship_to_address,
        "items": [
            {
                "source_row": styles[0]["source_row"],
                "raw_vendor_style": " / ".join(
                    str(style["vendor_style_code"])
                    for style in styles
                    if style["vendor_style_code"]
                ) or None,
                "raw_color": " / ".join(
                    str(style["color_raw"]) for style in styles if style["color_raw"]
                ) or None,
                "raw_size": " / ".join(
                    str(style["size_raw"]) for style in styles if style["size_raw"]
                ) or None,
                "raw_product_code": None,
                "raw_carton_dimensions": raw_dimensions or None,
                "carton_dimensions_mm": dimensions_mm,
                "style_rows": styles,
                "quantity": _json_number(carton_count),
                "quantity_check": quantity_check,
                "unit_price": None,
                "order_date": None,
                "delivery_date": None,
                "matched_product_id": None,
                "match_status": "unmatched",
                "product_candidates": [],
                "product_default_price": None,
                "warnings": item_warnings,
            }
        ],
        "warnings": warnings,
        "quantity_check": {"status": quantity_check["status"], "items": [quantity_check]},
        "requires_manual_confirmation": True,
    }


def _find_header(
    rows: Iterable[_SheetRow],
) -> tuple[_SheetRow | None, dict[str, tuple[int, ...]]]:
    best_row: _SheetRow | None = None
    best_mapping: dict[str, tuple[int, ...]] = {}
    for row in rows:
        collected: dict[str, list[int]] = {}
        for index, value in enumerate(row.values):
            field = _header_field(value)
            if field is not None:
                collected.setdefault(field, []).append(index)
        mapping = {field: tuple(indices) for field, indices in collected.items()}
        if len(mapping) > len(best_mapping):
            best_row, best_mapping = row, mapping
    return (best_row, best_mapping) if len(best_mapping) >= 2 else (None, {})


def _normalized_mapped_value(field: str, value: object) -> object:
    if field in {"units_per_carton", "carton_count", "garment_total"}:
        decimal_value = _decimal(value)
        return ("number", str(decimal_value.normalize())) if decimal_value is not None else (
            "text",
            re.sub(r"\s+", "", _text(value)).casefold(),
        )
    if field == "carton_dimensions":
        parsed = _parse_dimensions(value)
        if parsed is not None:
            dimensions, _raw = parsed
            return (
                "dimensions",
                dimensions["length_mm"],
                dimensions["width_mm"],
                dimensions["height_mm"],
            )
    return ("text", re.sub(r"\s+", " ", _text(value)).strip().casefold())


def _mapped_value(
    values: Iterable[object],
    fields: dict[str, tuple[int, ...]],
    field: str,
    *,
    row_number: int,
) -> object:
    value_list = tuple(values)
    candidates = [
        value_list[index]
        for index in fields.get(field, ())
        if index < len(value_list) and _text(value_list[index])
    ]
    if not candidates:
        return None
    normalized = {_normalized_mapped_value(field, value) for value in candidates}
    if len(normalized) > 1:
        raise XinzhenExcelParseError(
            f"第{row_number}行重复表头“{field}”映射到冲突值，已停止解析。"
        )
    return candidates[0]


def _quantity_check(
    values: Iterable[object],
    fields: dict[str, tuple[int, ...]],
    *,
    row_number: int,
) -> tuple[Decimal | None, dict]:
    units = _decimal(
        _mapped_value(
            values,
            fields,
            "units_per_carton",
            row_number=row_number,
        )
    )
    cartons = _decimal(
        _mapped_value(values, fields, "carton_count", row_number=row_number)
    )
    total = _decimal(
        _mapped_value(values, fields, "garment_total", row_number=row_number)
    )
    if cartons is None and units is not None and total is not None and units > 0:
        cartons = total / units
    if units is None and cartons is not None and total is not None and cartons > 0:
        units = total / cartons
    expected = units * cartons if units is not None and cartons is not None else None
    reconciled = expected is not None and total is not None and expected == total
    status = "passed" if reconciled else "needs_review"
    return cartons, {
        "units_per_carton": _json_number(units),
        "carton_count": _json_number(cartons),
        "garment_total": _json_number(total),
        "derived_garment_total": _json_number(expected),
        "status": status,
        "message": (
            "箱数已由每箱数量和成衣总数核对。"
            if reconciled
            else "无法核对每箱数量、箱数与成衣总数；保存前必须人工确认。"
        ),
    }


def _looks_like_total(values: Iterable[object]) -> bool:
    text = " ".join(_text(value).casefold() for value in values if _text(value))
    return any(marker in text for marker in _TOTAL_MARKERS)


def parse_xinzhen_excel_order(
    content: bytes,
    filename: str,
    *,
    worksheet_name: str | None = None,
) -> dict:
    """Parse B10:N30 into source-faithful, writable-by-the-user draft rows."""

    sheets = _load_carton_marking_sheets(content, filename)
    selected_sheet = _select_order_sheet(sheets, worksheet_name)
    sheet_name = selected_sheet.name
    rows = list(selected_sheet.rows)
    verified = _parse_verified_carton_marking(
        rows, filename=filename, sheet_name=sheet_name
    )
    if verified is not None:
        return verified
    header_row, fields = _find_header(rows)
    if header_row is None:
        raise XinzhenExcelParseError(
            "未识别到新振 CARTON MARKING 固定版式或完整标准表头，已停止解析以避免生成错误明细"
        )
    missing_headers = sorted(_GENERIC_REQUIRED_HEADERS - set(fields))
    if missing_headers or not ({"product_code", "carton_dimensions"} & set(fields)):
        raise XinzhenExcelParseError(
            "标准表头不完整，必须明确包含款式、箱数、每箱件数、总件数，以及产品编码或纸箱尺寸"
        )
    data_rows = [row for row in rows if header_row is None or row.number > header_row.number]
    items: list[dict] = []
    warnings: list[str] = []
    customer_pos: str | None = None
    customer_po_values: set[str] = set()

    for row in data_rows:
        if not any(_text(value) for value in row.values) or _looks_like_total(row.values):
            continue
        dimensions = None
        raw_dimensions = ""
        if "carton_dimensions" in fields:
            dimensions = _parse_dimensions(
                _mapped_value(
                    row.values,
                    fields,
                    "carton_dimensions",
                    row_number=row.number,
                )
            )
        if dimensions is not None:
            dimensions_mm, raw_dimensions = dimensions
        else:
            dimensions_mm = None

        source_fields = {
            "raw_vendor_style": None,
            "raw_color": None,
            "raw_size": None,
            "raw_product_code": None,
        }
        for output_key, header_key in (
            ("raw_vendor_style", "vendor_style"),
            ("raw_color", "color"),
            ("raw_size", "size"),
            ("raw_product_code", "product_code"),
        ):
            if header_key in fields:
                source_fields[output_key] = _text(
                    _mapped_value(
                        row.values,
                        fields,
                        header_key,
                        row_number=row.number,
                    )
                ) or None
        if "customer_po" in fields:
            row_po = _text(
                _mapped_value(
                    row.values,
                    fields,
                    "customer_po",
                    row_number=row.number,
                )
            ) or None
            if row_po:
                customer_po_values.add(row_po)
                if customer_pos is None:
                    customer_pos = row_po

        carton_count, quantity_check = _quantity_check(
            row.values,
            fields,
            row_number=row.number,
        )
        has_source_content = bool(
            dimensions_mm
            or source_fields["raw_vendor_style"]
            or source_fields["raw_product_code"]
            or carton_count is not None
        )
        if not has_source_content:
            continue
        item_warnings: list[str] = []
        if dimensions_mm is None:
            item_warnings.append("未识别到三段纸箱尺寸，不能自动匹配产品。")
        if carton_count is None or carton_count <= 0:
            item_warnings.append("未识别到有效箱数，不能直接保存。")
        if quantity_check["status"] != "passed":
            item_warnings.append(str(quantity_check["message"]))
        items.append(
            {
                "source_row": row.number,
                **source_fields,
                "raw_carton_dimensions": raw_dimensions or None,
                "carton_dimensions_mm": dimensions_mm,
                # Production quantity is carton count only.  Never substitute
                # units/carton or garment-total for this field.
                "quantity": _json_number(carton_count),
                "quantity_check": quantity_check,
                "unit_price": None,
                "order_date": None,
                "delivery_date": None,
                "matched_product_id": None,
                "match_status": "unmatched",
                "product_candidates": [],
                "product_default_price": None,
                "warnings": item_warnings,
            }
        )

    if not items:
        raise XinzhenExcelParseError("B10:N30 未识别到可编辑的 CARTON MARKING 明细")
    if len(customer_po_values) > 1:
        raise XinzhenExcelParseError(
            "Excel 中存在多个客户单号，当前导入只能生成一张订单草稿，已停止解析。"
        )
    return {
        "source_name": filename,
        "source_type": "xinzhen_excel_carton_marking",
        "sheet_name": sheet_name,
        "customer_po": customer_pos,
        "items": items,
        "warnings": warnings,
        "quantity_check": {
            "status": "passed" if all(item["quantity_check"]["status"] == "passed" for item in items) else "needs_review",
            "items": [item["quantity_check"] for item in items],
        },
        "requires_manual_confirmation": True,
    }


def _canonical_json(value: object) -> str:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )


def normalized_xinzhen_source_payload(draft: dict) -> dict:
    """Return the source-only payload protected by the preview signature."""

    source_items: list[dict] = []
    for item in draft.get("items") or []:
        source_items.append(
            {
                "source_row": int(item.get("source_row") or 0),
                "raw_vendor_style": item.get("raw_vendor_style"),
                "raw_color": item.get("raw_color"),
                "raw_size": item.get("raw_size"),
                "raw_product_code": item.get("raw_product_code"),
                "raw_carton_dimensions": item.get("raw_carton_dimensions"),
                "carton_dimensions_mm": item.get("carton_dimensions_mm"),
                "style_rows": item.get("style_rows") or [],
                "source_quantity": item.get("quantity"),
                "quantity_check": item.get("quantity_check") or {},
            }
        )
    return {
        "source_type": draft.get("source_type"),
        "sheet_name": draft.get("sheet_name"),
        "customer_po": draft.get("customer_po"),
        "department_code": draft.get("department_code"),
        "vendor_name": draft.get("vendor_name"),
        "customer_name_raw": draft.get("customer_name_raw"),
        "ship_to_address": draft.get("ship_to_address"),
        "items": source_items,
    }


def register_xinzhen_import_batch(
    db: Session,
    *,
    draft: dict,
    customer: Customer,
    source_filename: str,
    source_sha256: str,
    operator_id: int | None,
) -> tuple[ExcelOrderImportBatch, str]:
    """Insert or reuse the immutable source ledger for one uploaded file."""

    source_payload = normalized_xinzhen_source_payload(draft)
    source_payload_json = _canonical_json(source_payload)
    source_payload_hash = hashlib.sha256(
        source_payload_json.encode("utf-8")
    ).hexdigest()
    lookup = (
        ExcelOrderImportBatch.customer_id == customer.id,
        ExcelOrderImportBatch.source_type == draft.get("source_type"),
        ExcelOrderImportBatch.source_sha256 == source_sha256,
    )
    existing = db.scalar(select(ExcelOrderImportBatch).where(*lookup))
    if existing is not None:
        if (
            existing.parser_version != XINZHEN_EXCEL_PARSER_VERSION
            or existing.normalized_source_hash != source_payload_hash
            or existing.worksheet_name != str(draft.get("sheet_name") or "")
        ):
            raise XinzhenExcelParseError(
                "同一 Excel 文件的解析版本或规范化结果已变化，禁止静默覆盖既有导入台账。"
            )
        return existing, source_payload_hash

    batch = ExcelOrderImportBatch(
        customer_id=customer.id,
        customer_code_snapshot=(customer.customer_code or "").strip() or None,
        customer_name_snapshot=customer.name,
        source_type=str(draft.get("source_type") or "xinzhen_excel_carton_marking"),
        source_filename=source_filename,
        source_sha256=source_sha256,
        parser_version=XINZHEN_EXCEL_PARSER_VERSION,
        worksheet_name=str(draft.get("sheet_name") or ""),
        normalized_source_hash=source_payload_hash,
        normalized_source_json=source_payload_json,
        created_by=operator_id,
    )
    try:
        db.add(batch)
        db.flush()
        for item in source_payload["items"]:
            row_json = _canonical_json(item)
            db.add(
                ExcelOrderImportRow(
                    batch_id=batch.id,
                    source_row_number=int(item["source_row"]),
                    source_row_hash=hashlib.sha256(
                        row_json.encode("utf-8")
                    ).hexdigest(),
                    normalized_row_json=row_json,
                )
            )
        db.flush()
        return batch, source_payload_hash
    except IntegrityError:
        db.rollback()
        existing = db.scalar(select(ExcelOrderImportBatch).where(*lookup))
        if existing is None:
            raise
        if (
            existing.parser_version != XINZHEN_EXCEL_PARSER_VERSION
            or existing.normalized_source_hash != source_payload_hash
            or existing.worksheet_name != str(draft.get("sheet_name") or "")
        ):
            raise XinzhenExcelParseError(
                "同一 Excel 文件的解析结果发生冲突，已停止预览。"
            )
        return existing, source_payload_hash


def find_active_xinzhen_customer(db: Session) -> Customer:
    coded_customers = db.scalars(
        select(Customer)
        .where(
            Customer.customer_code.is_not(None),
            func.upper(func.trim(Customer.customer_code)).in_(XINZHEN_CUSTOMER_CODES),
            Customer.is_active.is_(True),
            Customer.status == "active",
        )
        .order_by(Customer.id)
    ).all()
    if len(coded_customers) == 1:
        return coded_customers[0]
    if len(coded_customers) > 1:
        raise XinzhenExcelParseError(
            "存在多个启用的新振客户编码，无法安全导入。"
        )
    named_customers = db.scalars(
        select(Customer)
        .where(
            Customer.name.contains("新振"),
            Customer.is_active.is_(True),
            Customer.status == "active",
        )
        .order_by(Customer.id)
    ).all()
    if len(named_customers) == 1:
        return named_customers[0]
    if not named_customers:
        raise XinzhenExcelParseError("未找到唯一启用的新振客户，无法导入。")
    raise XinzhenExcelParseError("存在多个启用的新振客户，无法安全导入。")


def _candidate(product: Product, *, include_sale_price: bool) -> dict:
    candidate = {
        "id": product.id,
        "product_code": product.product_code,
        "customer_material_code": product.customer_material_code,
        "product_name": product.product_name,
        "specification": (
            f"{_json_number(product.length_mm)}×{_json_number(product.width_mm)}×{_json_number(product.height_mm)}"
            if all(
                value is not None
                for value in (product.length_mm, product.width_mm, product.height_mm)
            )
            else None
        ),
        "material_id": product.material_id,
        "layer_count": product.layer_count,
        "flute_type": product.flute_type,
        "dimensions_mm": {
            "length_mm": _json_number(product.length_mm),
            "width_mm": _json_number(product.width_mm),
            "height_mm": _json_number(product.height_mm),
        },
    }
    if include_sale_price:
        candidate["sale_unit_price"] = (
            str(product.sale_unit_price)
            if product.sale_unit_price is not None
            else None
        )
    return candidate


def _product_dimensions_match(product: Product, dimensions: dict[str, object] | None) -> bool:
    if not dimensions or any(value is None for value in (product.length_mm, product.width_mm, product.height_mm)):
        return False
    return (
        product.length_mm == Decimal(str(dimensions["length_mm"]))
        and product.width_mm == Decimal(str(dimensions["width_mm"]))
        and product.height_mm == Decimal(str(dimensions["height_mm"]))
    )


def _exact_code_match(product: Product, code: str | None) -> bool:
    normalized = re.sub(r"\s+", "", (code or "")).casefold()
    if not normalized:
        return False
    return normalized in {
        re.sub(r"\s+", "", product.product_code or "").casefold(),
        re.sub(r"\s+", "", product.customer_material_code or "").casefold(),
    }


def match_xinzhen_excel_draft(
    db: Session,
    draft: dict,
    customer: Customer,
    *,
    include_sale_price: bool = True,
) -> dict:
    """Attach candidates only; selection is fail-closed on non-unique matches."""

    products = db.scalars(
        select(Product)
        .where(
            Product.customer_id == customer.id,
            Product.is_active.is_(True),
            Product.deleted_at.is_(None),
        )
        .order_by(Product.id)
    ).all()
    matched_items: list[dict] = []
    warnings = list(draft.get("warnings", []))
    for raw_item in draft["items"]:
        item = {**raw_item, "warnings": list(raw_item.get("warnings", []))}
        source_dimensions = item.get("carton_dimensions_mm")
        dimension_matches = [
            product
            for product in products
            if _product_dimensions_match(product, source_dimensions)
        ]
        code_matches = [
            product for product in products if _exact_code_match(product, item.get("raw_product_code"))
        ]
        has_source_code = bool(re.sub(r"\s+", "", item.get("raw_product_code") or ""))
        unmatched_source_code = has_source_code and not code_matches
        selected: Product | None = None
        match_basis: str | None = None
        dimension_code_conflict = bool(
            source_dimensions
            and has_source_code
            and (
                unmatched_source_code
                or not any(
                    _product_dimensions_match(product, source_dimensions)
                    for product in code_matches
                )
            )
        )
        candidates = dimension_matches or code_matches
        if dimension_code_conflict:
            candidates = products
        elif dimension_matches and code_matches:
            candidates = list({product.id: product for product in [*dimension_matches, *code_matches]}.values())
            code_candidate = code_matches[0] if len(code_matches) == 1 else None
            if code_candidate is not None and any(
                product.id == code_candidate.id for product in dimension_matches
            ):
                selected, match_basis = code_candidate, "exact_dimensions_and_product_code"
        elif not has_source_code and len(dimension_matches) == 1:
            selected, match_basis = dimension_matches[0], "exact_carton_dimensions_cm_to_mm"
        elif not source_dimensions and len(code_matches) == 1:
            selected, match_basis = code_matches[0], "exact_product_code"
        ambiguous_match = selected is None and len(candidates) > 1
        if not candidates:
            # An unmatched source stays fail-closed, but the operator must still
            # be able to choose an existing product belonging to this customer.
            candidates = products
        item["product_candidates"] = [
            _candidate(product, include_sale_price=include_sale_price)
            for product in candidates
        ]
        item["matched_product_id"] = selected.id if selected is not None else None
        item["match_status"] = (
            "conflict"
            if dimension_code_conflict
            else "matched"
            if selected is not None
            else "ambiguous"
            if ambiguous_match
            else "unmatched"
        )
        item["match_basis"] = match_basis
        item["product_default_price"] = (
            str(selected.sale_unit_price)
            if include_sale_price
            and selected is not None
            and selected.sale_unit_price is not None
            else None
        )
        if dimension_code_conflict:
            item["warnings"].append(
                "Excel 原始产品编码未匹配，或其常用箱尺寸与 Excel 尺寸不一致；已取消自动匹配，请人工选择。"
            )
        elif item["match_status"] != "matched":
            item["warnings"].append("产品匹配不唯一或未匹配；请在草稿中人工选择。")
        matched_items.append(item)
    if any(item["match_status"] != "matched" for item in matched_items):
        warnings.append("存在未匹配或歧义产品，保存前必须人工选择。")
    return {
        **draft,
        "matched_customer_id": customer.id,
        "matched_customer_name": customer.name,
        "customer_match_status": "matched",
        "items": matched_items,
        "warnings": warnings,
        "requires_manual_confirmation": True,
    }
