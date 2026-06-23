from __future__ import annotations

from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from phase1_postgres.database import get_session
from phase1_postgres.models import Customer, ReturnReceipt, ReturnReceiptItem, Statement, StatementItem


router = APIRouter(prefix="/api/statements", tags=["statements"])


class StatementGeneratePayload(BaseModel):
    customer_id: int
    start_date: date
    end_date: date
    note: str | None = None


def money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.00"), rounding=ROUND_HALF_UP)


def _next_statement_number(session: Session, end_date: date) -> str:
    prefix = f"ST-{end_date:%Y%m}-"
    max_number = session.scalar(
        select(func.max(Statement.statement_number)).where(Statement.statement_number.like(f"{prefix}%"))
    )
    if not max_number:
        return f"{prefix}001"
    try:
        serial = int(str(max_number).rsplit("-", 1)[1]) + 1
    except (IndexError, ValueError):
        serial = 1
    return f"{prefix}{serial:03d}"


def statement_item_dict(item: StatementItem) -> dict[str, Any]:
    return {
        "id": item.id,
        "receipt_item_id": item.receipt_item_id,
        "delivery_date": item.delivery_date.isoformat(),
        "delivery_number": item.delivery_number,
        "order_number": item.order_number,
        "customer_po": item.customer_po,
        "product_code": item.product_code,
        "product_name": item.product_name,
        "spec": item.spec,
        "actual_signed_qty": item.actual_signed_qty,
        "unit_price": str(item.unit_price_snapshot),
        "amount": str(item.amount),
        "remark": item.remark,
    }


def statement_dict(statement: Statement) -> dict[str, Any]:
    return {
        "id": statement.id,
        "statement_number": statement.statement_number,
        "customer_id": statement.customer_id,
        "customer_name": statement.customer.name if statement.customer else None,
        "statement_month": statement.statement_month,
        "start_date": statement.start_date.isoformat(),
        "end_date": statement.end_date.isoformat(),
        "total_quantity": statement.total_quantity,
        "total_amount": str(statement.total_amount),
        "status": statement.status,
        "items": [statement_item_dict(item) for item in statement.items],
    }


def _eligible_receipt_items(session: Session, payload: StatementGeneratePayload) -> list[ReturnReceiptItem]:
    query = (
        select(ReturnReceiptItem)
        .join(ReturnReceipt)
        .where(ReturnReceipt.status == "OWNER_REVIEWED")
        .where(~ReturnReceiptItem.statement_items.any())
        .where(ReturnReceipt.actual_received_date >= payload.start_date)
        .where(ReturnReceipt.actual_received_date <= payload.end_date)
        .where(ReturnReceipt.delivery.has(customer_id=payload.customer_id))
        .order_by(ReturnReceipt.actual_received_date, ReturnReceiptItem.id)
    )
    return session.scalars(query).unique().all()


@router.post("/generate", status_code=status.HTTP_201_CREATED)
def generate_statement(payload: StatementGeneratePayload, session: Session = Depends(get_session)) -> dict[str, Any]:
    customer = session.get(Customer, payload.customer_id)
    if customer is None:
        raise HTTPException(status_code=400, detail="客户不存在")

    receipt_items = _eligible_receipt_items(session, payload)
    if not receipt_items:
        raise HTTPException(status_code=400, detail="没有老板已核对且未对账的回单明细")

    statement = Statement(
        statement_number=_next_statement_number(session, payload.end_date),
        customer_id=payload.customer_id,
        statement_month=f"{payload.end_date:%Y-%m}",
        start_date=payload.start_date,
        end_date=payload.end_date,
        status="DRAFT",
        note=payload.note,
    )
    total_qty = 0
    total_amount = Decimal("0.00")
    for receipt_item in receipt_items:
        delivery_item = receipt_item.delivery_item
        order_item = delivery_item.order_item
        delivery = delivery_item.delivery
        order = order_item.order
        amount = money(receipt_item.amount)
        statement.items.append(
            StatementItem(
                receipt_item_id=receipt_item.id,
                delivery_date=delivery.delivery_date,
                delivery_number=delivery.delivery_number,
                order_number=order.order_number,
                customer_po=order.customer_po,
                product_code=order_item.snapshot_product_code,
                product_name=order_item.snapshot_product_name,
                spec=order_item.snapshot_spec,
                actual_signed_qty=receipt_item.actual_signed_qty,
                unit_price_snapshot=receipt_item.unit_price_snapshot,
                amount=amount,
                remark=receipt_item.difference_reason,
            )
        )
        total_qty += receipt_item.actual_signed_qty
        total_amount += amount
    statement.total_quantity = total_qty
    statement.total_amount = money(total_amount)
    session.add(statement)
    session.commit()
    session.refresh(statement)
    return statement_dict(statement)


@router.get("/{statement_id}")
def get_statement(statement_id: int, session: Session = Depends(get_session)) -> dict[str, Any]:
    statement = session.get(Statement, statement_id)
    if statement is None:
        raise HTTPException(status_code=404, detail="对账单不存在")
    return statement_dict(statement)


@router.get("/{statement_id}/export")
def export_statement(statement_id: int, session: Session = Depends(get_session)) -> StreamingResponse:
    statement = session.get(Statement, statement_id)
    if statement is None:
        raise HTTPException(status_code=404, detail="对账单不存在")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "月结对账"
    worksheet.merge_cells("A1:J1")
    worksheet["A1"] = "天明包装月结对账清单"
    worksheet["A1"].font = Font(size=18, bold=True)
    worksheet["A1"].alignment = Alignment(horizontal="center")
    worksheet["A2"] = f"客户：{statement.customer.name if statement.customer else ''}"
    worksheet["D2"] = f"周期：{statement.start_date:%Y-%m-%d} 至 {statement.end_date:%Y-%m-%d}"
    worksheet["H2"] = f"单号：{statement.statement_number}"
    headers = ["送货日期", "送货单号", "订单号", "客户单号", "品号", "产品名称", "规格", "实签数量", "单价", "金额"]
    worksheet.append([])
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[4]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for item in statement.items:
        worksheet.append(
            [
                item.delivery_date.isoformat(),
                item.delivery_number,
                item.order_number,
                item.customer_po,
                item.product_code,
                item.product_name,
                item.spec,
                item.actual_signed_qty,
                float(item.unit_price_snapshot),
                float(item.amount),
            ]
        )
    total_row = worksheet.max_row + 1
    worksheet.cell(total_row, 7, "合计")
    worksheet.cell(total_row, 8, statement.total_quantity)
    worksheet.cell(total_row, 10, float(statement.total_amount))
    for column, width in {
        "A": 14,
        "B": 18,
        "C": 18,
        "D": 16,
        "E": 14,
        "F": 22,
        "G": 18,
        "H": 12,
        "I": 12,
        "J": 14,
    }.items():
        worksheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"{statement.statement_number}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
