from decimal import Decimal
from io import BytesIO
import pytest
from openpyxl import Workbook,load_workbook
from fastapi.testclient import TestClient
from sqlalchemy import text
from tests.test_p1_130_statement_invoice_finance import p1_130_app,_login
from app.services.customer_statement_check import load_customer,compare,export_report


def workbook(rows):
    book=Workbook();sheet=book.active
    for row in rows:sheet.append(row)
    stream=BytesIO();book.save(stream);return stream.getvalue()


def test_duplicate_aggregation_differences_mapping_and_invalid_numbers():
    content=workbook([['对账清单'],['客户单号','存货编码','数量','金额'],['PO','0006',2,10],['PO','0006',3,15],['PO','extra',1,3],['PO','broken','abc',6]])
    parsed=load_customer(content)
    assert parsed['header_row']==2 and len(parsed['errors'])==1
    erp=[dict(customer_po='PO',product_code='0006',quantity=5,amount=Decimal('25'),unit_price=5,statement_item_id=7)]
    result=compare(parsed,erp);assert result['summary']=={'一致':1,'客户多出':1} and not result['complete']
    row=next(r for r in result['items'] if r['status']=='一致');assert row['customer_rows']==[3,4]
    assert load_customer(workbook([['A','B','C'],['code',2,3]]))['needs_mapping']
    mapped=load_customer(workbook([['A','B','C'],['code',2,3]]),columns={'product_code':0,'quantity':1,'amount':2});assert not mapped['needs_mapping']
    with pytest.raises(ValueError):load_customer(content,columns={'quantity':1,'amount':1,'product_code':0})


def test_report_does_not_execute_cell_formula_and_preserves_code():
    content=workbook([['存货编码','数量','金额'],['=HYPERLINK("x")',1,2]])
    # cached formula is absent: never treat it as an identified product
    assert load_customer(content)['errors']
    result={'customer_name':'customer','statement_month':'2026-09','statement_number':'ST','matching_fields':['存货编码'],'complete':True,'items':[{'status':'客户多出','delivery_number':'','customer_po':'=1+1','product_code':'0006','product_name':'name','customer_quantity':'1','erp_quantity':'0','quantity_difference':'1','customer_amount':'2','erp_amount':'0','amount_difference':'2','customer_prices':'2','erp_prices':'','customer_rows':[2],'statement_item_ids':[]}],'errors':[]}
    sheet=load_workbook(export_report(result),data_only=False).active
    assert sheet['C5'].data_type=='s' and sheet['C5'].value=='=1+1'
    assert sheet['D5'].value=='0006'


def test_authorized_api_roundtrip_readonly_and_stale_download(p1_130_app):
    app,factory=p1_130_app
    with factory() as db:before=db.execute(text('select * from finance_statements')).all()
    with TestClient(app) as client:
        url='/api/finance/statements/1/customer-check'
        assert client.post(url,files={'file':('x.xlsx',b'x')}).status_code==401
        _login(client)
        original=client.get('/api/finance/statements/1/customer-export.xlsx');assert original.status_code==200
        result=client.post(url,files={'file':('customer.xlsx',original.content)})
        assert result.status_code==200,result.text
        data=result.json();assert data['complete'] and set(data['summary'])=={'一致'}
        report=client.post(url,files={'file':('customer.xlsx',original.content)},data={'output_format':'xlsx','expected_version':data['statement_version']})
        assert report.status_code==200 and load_workbook(BytesIO(report.content)).active.title=='对账差异'
        stale=client.post(url,files={'file':('customer.xlsx',original.content)},data={'output_format':'xlsx','expected_version':999})
        assert stale.status_code==409
        assert client.post('/api/finance/statements/999/customer-check',files={'file':('customer.xlsx',original.content)}).status_code==404
    with factory() as db:assert db.execute(text('select * from finance_statements')).all()==before
