from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.services.customer_statement_reconciliation import (
    AGGREGATED_NOTE,
    AMOUNT_MISMATCH,
    CUSTOMER_DUPLICATE,
    CUSTOMER_ONLY,
    CUSTOMER_PO_MISMATCH,
    ERP_DUPLICATE,
    ERP_ONLY,
    MANUAL_REVIEW,
    MATCHED,
    PRODUCT_NAME_MISMATCH,
    QUANTITY_MISMATCH,
    SPECIFICATION_MISMATCH,
    UNIT_MISMATCH,
    UNIT_PRICE_MISMATCH,
    ReconciliationMetadata,
    generate_reconciliation_xlsx,
    parse_customer_statement_xlsx,
    reconcile_customer_statement_xlsx,
)


SAMPLE_WORKBOOK = Path(r"D:\Edge浏览器下载\124507.xlsx")

CUSTOMER_COLUMNS = {
    "sequence": "A",
    "receipt_date": "B",
    "receipt_number": "D",
    "supplier_code": "G",
    "supplier_name": "H",
    "product_code": "J",
    "product_name": "L",
    "specification": "N",
    "unit": "P",
    "quantity": "Q",
    "customer_po": "U",
    "unit_price": "W",
    "net_amount": "Y",
    "amount": "AC",
}

CUSTOMER_HEADERS = {
    "sequence": "序号",
    "receipt_date": "入库日期",
    "receipt_number": "入库单号",
    "supplier_code": "供货商编码",
    "supplier_name": "供货商名称",
    "product_code": "存货编码",
    "product_name": "存货名称",
    "specification": "规格型号",
    "unit": "单位",
    "quantity": "实收数量",
    "customer_po": "采购订单号",
    "unit_price": "原币含税单价",
    "net_amount": "原币金额",
    "amount": "原币价税合计",
}


def _customer_row(**overrides):
    row = {
        "sequence": "1",
        "receipt_date": date(2026, 7, 1),
        "receipt_number": "WR-001",
        "supplier_code": "1245",
        "supplier_name": "苏州天明包装有限公司",
        "product_code": "21301021",
        "product_name": "中性内箱",
        "specification": "115*67*2.5cm",
        "unit": "个",
        "quantity": "10",
        "customer_po": "PO-001",
        "unit_price": "5.31",
        "net_amount": "46.99",
        "amount": "53.10",
    }
    row.update(overrides)
    return row


def _workbook_bytes(rows: list[dict], *, sheet_name: str = "采购入库打印") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet["K2"] = "采购入库单列表"
    for field_name, column in CUSTOMER_COLUMNS.items():
        sheet[f"{column}5"] = CUSTOMER_HEADERS[field_name]
    placeholder_columns = ("C", "E", "F", "I", "K", "M", "O", "R", "S", "T", "V", "X", "Z", "AA", "AB", "AD")
    for offset, source in enumerate(rows, 6):
        for field_name, column in CUSTOMER_COLUMNS.items():
            sheet[f"{column}{offset}"] = source.get(field_name)
        for column in placeholder_columns:
            sheet[f"{column}{offset}"] = 1
    # 模拟客户导出文件尾部的合计行；没有订单和产品身份时必须忽略。
    total_row = 6 + len(rows)
    sheet[f"Q{total_row}"] = sum(Decimal(str(row["quantity"]).replace(",", "")) for row in rows)
    sheet[f"AC{total_row}"] = sum(Decimal(str(row["amount"]).replace(",", "")) for row in rows)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _metadata() -> ReconciliationMetadata:
    return ReconciliationMetadata(
        customer_name="测试客户",
        statement_month="2026-07",
        customer_file_name="customer.xlsx",
        erp_statement_range="ST-001",
        operator="admin",
        generated_at=datetime(2026, 7, 20, 10, 30, 0),
    )


def _erp_row(**overrides):
    row = {
        "source_row": 1,
        "statement_item_id": 101,
        "delivery_date": date(2026, 7, 1),
        "delivery_number": "DN-001",
        "customer_po": "PO-001",
        "product_code": "21301021",
        "product_name": "中性内箱",
        "specification": "115*67*2.5cm",
        "unit": "个",
        "actual_received_quantity": "10",
        "unit_price_snapshot": "5.31",
        "receivable_amount": "53.10",
    }
    row.update(overrides)
    return row


def test_parser_detects_sparse_header_and_normalizes_values() -> None:
    payload = _workbook_bytes(
        [
            _customer_row(
                sequence=1,
                receipt_date="2026/07/18",
                product_code=" 001-ABC ",
                product_name="  测试   产品 ",
                quantity="1,000.00",
                unit_price="5.3100",
                net_amount="4,699.12",
                amount="5,310.00",
            )
        ]
    )

    parsed = parse_customer_statement_xlsx(payload)

    assert parsed.sheet_name == "采购入库打印"
    assert parsed.header_row == 5
    assert parsed.field_mapping == {
        "sequence": 1,
        "receipt_date": 2,
        "receipt_number": 4,
        "supplier_code": 7,
        "supplier_name": 8,
        "product_code": 10,
        "product_name": 12,
        "specification": 14,
        "unit": 16,
        "quantity": 17,
        "customer_po": 21,
        "unit_price": 23,
        "net_amount": 25,
        "amount": 29,
    }
    assert len(parsed.rows) == 1
    row = parsed.rows[0]
    assert row.source_row == 6
    assert row.receipt_date == date(2026, 7, 18)
    assert row.product_code == "001-ABC"
    assert row.product_name == "测试 产品"
    assert row.quantity == Decimal("1000.00")
    assert row.unit_price == Decimal("5.3100")
    assert row.net_amount == Decimal("4699.12")
    assert row.amount == Decimal("5310.00")
    assert row.warnings == ()


def test_duplicate_rows_are_flagged_and_nonduplicate_rows_aggregate_for_matching() -> None:
    first = _customer_row(quantity="10", amount="20", net_amount="18", unit_price="2")
    duplicate = dict(first, sequence="2")
    second_receipt = _customer_row(
        sequence="3",
        receipt_date=date(2026, 7, 2),
        receipt_number="WR-002",
        quantity="20",
        amount="40",
        net_amount="36",
        unit_price="2",
    )
    report = reconcile_customer_statement_xlsx(
        _workbook_bytes([first, duplicate, second_receipt]),
        [
            _erp_row(
                actual_received_quantity="30",
                unit_price_snapshot="2",
                receivable_amount="60",
            )
        ],
        metadata=_metadata(),
    )

    duplicate_result = next(
        result for result in report.results if result.difference_type == CUSTOMER_DUPLICATE
    )
    matched_result = next(result for result in report.results if result.is_match)

    assert duplicate_result.requires_manual_review is True
    assert duplicate_result.customer_group is not None
    assert duplicate_result.customer_group.source_rows == (6, 7)
    assert matched_result.difference_type == MATCHED
    assert matched_result.customer_group is not None
    assert matched_result.customer_group.quantity == Decimal("30")
    assert matched_result.customer_group.is_aggregated is True
    assert AGGREGATED_NOTE in matched_result.notes


def test_tolerances_are_inclusive_and_larger_differences_are_reported() -> None:
    report = reconcile_customer_statement_xlsx(
        _workbook_bytes(
            [
                _customer_row(customer_po="PO-A", product_code="A", quantity="10", unit_price="5.00", amount="50.00"),
                _customer_row(
                    sequence="2",
                    customer_po="PO-B",
                    product_code="B",
                    product_name="客户名称",
                    specification="客户规格",
                    unit="个",
                    quantity="10",
                    unit_price="5.00",
                    amount="50.00",
                ),
            ]
        ),
        [
            _erp_row(
                source_row=1,
                customer_po="PO-A",
                product_code="A",
                actual_received_quantity="10",
                unit_price_snapshot="4.99",
                receivable_amount="49.90",
            ),
            _erp_row(
                source_row=2,
                statement_item_id=102,
                customer_po="PO-B",
                product_code="B",
                product_name="ERP名称",
                specification="ERP规格",
                unit="箱",
                actual_received_quantity="9",
                unit_price_snapshot="4.98",
                receivable_amount="44.00",
            ),
        ],
        metadata=_metadata(),
    )

    within_tolerance = next(
        result
        for result in report.results
        if result.customer_group and result.customer_group.customer_po == "PO-A"
    )
    outside_tolerance = next(
        result
        for result in report.results
        if result.customer_group and result.customer_group.customer_po == "PO-B"
    )

    assert within_tolerance.is_match is True
    assert within_tolerance.match_level == "强匹配"
    assert set(outside_tolerance.difference_types) == {
        QUANTITY_MISMATCH,
        UNIT_PRICE_MISMATCH,
        AMOUNT_MISMATCH,
        PRODUCT_NAME_MISMATCH,
        SPECIFICATION_MISMATCH,
        UNIT_MISMATCH,
    }
    assert outside_tolerance.quantity_difference == Decimal("1")
    assert outside_tolerance.unit_price_difference == Decimal("0.02")
    assert outside_tolerance.amount_difference == Decimal("6.00")


def test_auxiliary_match_reports_order_mismatch_and_ambiguous_candidates_wait_for_review() -> None:
    report = reconcile_customer_statement_xlsx(
        _workbook_bytes(
            [
                _customer_row(customer_po="CUSTOMER-1", product_code="AUX-1", quantity="10", amount="20", unit_price="2"),
                _customer_row(sequence="2", customer_po="CUSTOMER-2", product_code="AUX-2", quantity="5", amount="10", unit_price="2"),
            ]
        ),
        [
            _erp_row(
                source_row=1,
                customer_po="ERP-1",
                product_code="AUX-1",
                actual_received_quantity="10",
                unit_price_snapshot="2",
                receivable_amount="20",
            ),
            _erp_row(
                source_row=2,
                statement_item_id=102,
                customer_po="ERP-2A",
                product_code="AUX-2",
                actual_received_quantity="5",
                unit_price_snapshot="2",
                receivable_amount="10",
            ),
            _erp_row(
                source_row=3,
                statement_item_id=103,
                customer_po="ERP-2B",
                product_code="AUX-2",
                actual_received_quantity="5",
                unit_price_snapshot="2",
                receivable_amount="10",
            ),
        ],
        metadata=_metadata(),
    )

    order_mismatch = next(
        result
        for result in report.results
        if result.customer_group and result.customer_group.customer_po == "CUSTOMER-1"
    )
    ambiguous = next(
        result
        for result in report.results
        if result.customer_group and result.customer_group.customer_po == "CUSTOMER-2"
    )

    assert order_mismatch.match_level == "辅助匹配"
    assert order_mismatch.difference_types == (CUSTOMER_PO_MISMATCH,)
    assert ambiguous.difference_type == MANUAL_REVIEW
    assert ambiguous.requires_manual_review is True
    assert "多个 ERP 候选" in ambiguous.notes[0]
    assert not any(
        result.difference_type == ERP_ONLY
        and result.erp_group
        and result.erp_group.product_code == "AUX-2"
        for result in report.results
    )


def test_unmatched_and_erp_duplicate_results_are_preserved() -> None:
    erp_duplicate = _erp_row(
        source_row=2,
        statement_item_id=202,
        customer_po="ERP-ONLY",
        product_code="ERP-X",
    )
    report = reconcile_customer_statement_xlsx(
        _workbook_bytes([_customer_row(customer_po="CUSTOMER-ONLY", product_code="CUSTOMER-X")]),
        [
            dict(erp_duplicate, source_row=1, statement_item_id=201),
            erp_duplicate,
        ],
        metadata=_metadata(),
    )

    difference_types = [result.difference_type for result in report.results]
    assert CUSTOMER_ONLY in difference_types
    assert ERP_ONLY in difference_types
    assert ERP_DUPLICATE in difference_types


def test_generated_report_contains_four_sheets_and_safe_standardized_values() -> None:
    report = reconcile_customer_statement_xlsx(
        _workbook_bytes([_customer_row(product_name="@HYPERLINK")]),
        [_erp_row(product_name="@HYPERLINK")],
        metadata=_metadata(),
    )

    output = generate_reconciliation_xlsx(report)
    workbook = load_workbook(BytesIO(output), data_only=False)

    assert workbook.sheetnames == [
        "核对汇总",
        "差异明细",
        "客户原始数据标准化结果",
        "ERP对账数据标准化结果",
    ]
    summary = workbook["核对汇总"]
    assert summary["A1"].value == "核对项目"
    assert summary["B2"].value == "测试客户"
    assert summary["B12"].value == 1
    detail = workbook["差异明细"]
    assert detail["A1"].value == "差异类型"
    assert detail["A2"].value == MATCHED
    customer = workbook["客户原始数据标准化结果"]
    assert customer["H2"].value == "'@HYPERLINK"
    assert customer.freeze_panes == "A2"
    assert customer.auto_filter.ref is not None
    workbook.close()


@pytest.mark.skipif(not SAMPLE_WORKBOOK.exists(), reason="真实 124507.xlsx 样例不存在")
def test_real_124507_sample_read_only_smoke() -> None:
    before = SAMPLE_WORKBOOK.read_bytes()

    parsed = parse_customer_statement_xlsx(SAMPLE_WORKBOOK)

    assert parsed.sheet_name == "采购入库打印"
    assert parsed.header_row == 5
    assert len(parsed.rows) == 457
    assert parsed.rows[0].source_row == 6
    assert parsed.rows[0].customer_po == "PO2026050685"
    assert parsed.rows[0].product_code == "21301021"
    assert sum((row.quantity or Decimal("0") for row in parsed.rows), Decimal("0")) == Decimal("67455")
    assert sum((row.net_amount or Decimal("0") for row in parsed.rows), Decimal("0")) == Decimal("161717.89")
    assert sum((row.amount or Decimal("0") for row in parsed.rows), Decimal("0")) == Decimal("182741.28")
    assert SAMPLE_WORKBOOK.read_bytes() == before
