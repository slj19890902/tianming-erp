from __future__ import annotations

from collections.abc import Generator
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from app.api import finance
from app.api.deps import get_db
from app.core.database import create_sqlite_engine
from app.core.security import hash_password
from app.models import Base
from app.models.audit import OperationLog
from app.models.customer import Customer
from app.models.finance import Statement
from app.models.user import User


def _customer_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "采购入库打印"
    headers = {
        "A5": "序号",
        "B5": "入库日期",
        "D5": "入库单号",
        "G5": "供货商编码",
        "H5": "供货商名称",
        "J5": "存货编码",
        "L5": "存货名称",
        "N5": "规格型号",
        "P5": "单位",
        "Q5": "实收数量",
        "U5": "采购订单号",
        "W5": "原币含税单价",
        "Y5": "原币金额",
        "AC5": "原币价税合计",
    }
    for cell, value in headers.items():
        sheet[cell] = value
    values = {
        "A6": 1,
        "B6": date(2026, 7, 18),
        "D6": "RK-001",
        "G6": "1245",
        "H6": "苏州天明包装有限公司",
        "J6": "21301021",
        "L6": "中性内箱",
        "N6": "115*67*2.5cm",
        "P6": "只",
        "Q6": 10,
        "U6": "PO-001",
        "W6": 5.31,
        "Y6": 46.99,
        "AC6": 53.10,
    }
    for cell, value in values.items():
        sheet[cell] = value
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.fixture()
def reconciliation_api(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    engine = create_sqlite_engine(tmp_path / "customer-reconciliation-api.sqlite3")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        user = User(
            username="customer-reconciliation-admin",
            password_hash=hash_password("TestOnly123!"),
            role="admin",
            real_name="Reconciliation Admin",
            must_change_password=False,
            customer_access_mode="all",
        )
        outsider = User(
            username="customer-reconciliation-outsider",
            password_hash=hash_password("TestOnly123!"),
            role="finance",
            real_name="Out Of Scope Finance",
            must_change_password=False,
            customer_access_mode="selected",
        )
        customer = Customer(name="客户对账测试客户")
        db.add_all([user, outsider, customer])
        db.flush()
        statement = Statement(
            statement_number="ST-202607-001",
            customer_id=customer.id,
            statement_month="2026-07",
            total_receivable=Decimal("53.10"),
            total_gross_profit=Decimal("0"),
            status="unsettled",
            created_by=user.id,
        )
        db.add(statement)
        db.commit()
        user_id = user.id
        outsider_id = outsider.id
        statement_id = statement.id

    erp_rows = [
        {
            "source_row": 1,
            "statement_item_id": 1,
            "delivery_date": date(2026, 7, 18),
            "delivery_number": "DN-001",
            "customer_po": "PO-001",
            "product_code": "21301021",
            "product_name": "中性内箱",
            "specification": "115*67*2.5cm",
            "unit": "只",
            "actual_received_quantity": 10,
            "unit_price_snapshot": Decimal("5.31"),
            "receivable_amount": Decimal("53.10"),
        }
    ]
    monkeypatch.setattr(
        finance,
        "_customer_statement_erp_rows",
        lambda *_args, **_kwargs: erp_rows,
    )

    app = FastAPI()
    app.include_router(finance.router, prefix="/api/finance")

    def override_get_db() -> Generator[Session, None, None]:
        with factory() as db:
            yield db

    current_user_id = {"value": user_id}

    def override_can_operate() -> User:
        with factory() as db:
            return db.get(User, current_user_id["value"])

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[finance.can_operate] = override_can_operate
    yield app, factory, statement_id, current_user_id, outsider_id


def test_preview_reconciles_without_modifying_statement_and_writes_audit(
    reconciliation_api,
) -> None:
    app, factory, statement_id, _current_user_id, _outsider_id = reconciliation_api
    payload = _customer_workbook()
    with TestClient(app) as client:
        response = client.post(
            f"/api/finance/statements/{statement_id}/customer-reconciliation/preview",
            files={
                "file": (
                    "customer-statement.xlsx",
                    payload,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["summary"]["customer_file_total_rows"] == 1
    assert body["summary"]["erp_total_rows"] == 1
    assert body["summary"]["matched_rows"] == 1
    assert body["summary"]["difference_rows"] == 0
    assert body["differences"] == []
    with factory() as db:
        statement = db.get(Statement, statement_id)
        assert statement.total_receivable == Decimal("53.10")
        log = db.scalar(
            select(OperationLog).where(
                OperationLog.action == "RECONCILE_CUSTOMER_XLSX"
            )
        )
        assert log is not None
        assert "file_sha256" in (log.details or "")


def test_export_returns_four_sheet_report_and_rejects_wrong_extension(
    reconciliation_api,
) -> None:
    app, _factory, statement_id, _current_user_id, _outsider_id = reconciliation_api
    payload = _customer_workbook()
    with TestClient(app) as client:
        rejected = client.post(
            f"/api/finance/statements/{statement_id}/customer-reconciliation/preview",
            files={"file": ("customer.csv", payload, "text/csv")},
        )
        exported = client.post(
            f"/api/finance/statements/{statement_id}/customer-reconciliation/export",
            files={
                "file": (
                    "customer.xlsx",
                    payload,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert rejected.status_code == 422
    assert "xlsx" in rejected.json()["detail"]
    assert exported.status_code == 200, exported.text
    assert "spreadsheetml.sheet" in exported.headers["content-type"]
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == [
        "核对汇总",
        "差异明细",
        "客户原始数据标准化结果",
        "ERP对账数据标准化结果",
    ]
    workbook.close()


def test_customer_scope_is_enforced_before_reconciliation(
    reconciliation_api,
) -> None:
    app, _factory, statement_id, current_user_id, outsider_id = reconciliation_api
    current_user_id["value"] = outsider_id
    with TestClient(app) as client:
        response = client.post(
            f"/api/finance/statements/{statement_id}/customer-reconciliation/preview",
            files={
                "file": (
                    "customer.xlsx",
                    _customer_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

    assert response.status_code == 403
