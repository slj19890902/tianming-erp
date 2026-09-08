from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, func
from test_n029_production_integration import n029_delivery_app, _login, _prepare_103_finished_stock
from app.models.delivery import DeliveryItem
from app.models.order import Order, OrderItem
from app.models.customer import Customer
from app.models.warehouse_inventory import InventoryMovement


def test_customer_po_changes_only_delivery_and_statement(n029_delivery_app):
    app, factory, ids = n029_delivery_app
    _prepare_103_finished_stock(factory, ids)
    with factory() as db:
        item = db.get(OrderItem, ids["task_completed"])
        original_order_id = item.order_id
        original_po = db.get(Order, original_order_id).customer_po
        db.get(Customer, ids["customer"]).statement_cycle_start_day = 1
        db.commit()
    with TestClient(app) as client:
        _login(client)
        response = client.post("/api/deliveries", json={"customer_id":ids["customer"],
            "items":[{"order_item_id":ids["task_completed"], "delivered_quantity":100}]})
        assert response.status_code == 201, response.text
        delivery = response.json()
        did, line_id = delivery["id"], delivery["items"][0]["id"]
        dispatched = client.put(f"/api/deliveries/{did}/dispatch")
        assert dispatched.status_code == 200, dispatched.text
        with factory() as db:
            movement_count = db.scalar(select(func.count()).select_from(InventoryMovement))
        payload = {"expected_version":dispatched.json()["version"],
            "idempotency_key":"test-delivery-customer-po-01",
            "items":[{"delivery_item_id":line_id,"customer_po":"NEW-PO-2026"}]}
        changed = client.put(f"/api/deliveries/{did}/customer-po", json=payload)
        assert changed.status_code == 200, changed.text
        assert changed.json()["items"][0]["customer_po"] == "NEW-PO-2026"
        assert changed.json()["status"] == "dispatched"
        assert client.put(f"/api/deliveries/{did}/customer-po", json=payload).json() == changed.json()
        bad = {**payload,"items":[{"delivery_item_id":line_id,"customer_po":"OTHER"}]}
        assert client.put(f"/api/deliveries/{did}/customer-po", json=bad).status_code == 409
        stale = {**payload,"idempotency_key":"test-stale-po-01"}
        assert client.put(f"/api/deliveries/{did}/customer-po", json=stale).status_code == 409
        foreign = {**payload,"expected_version":changed.json()["version"],
            "idempotency_key":"test-foreign-po-01", "items":[{"delivery_item_id":999999,"customer_po":"BAD"}]}
        assert client.put(f"/api/deliveries/{did}/customer-po", json=foreign).status_code == 409
        printed = client.get(f"/api/deliveries/{did}/print")
        assert printed.status_code == 200, printed.text
        assert printed.json()["items"][0]["customer_po"] == "NEW-PO-2026"
        receipt = client.post("/api/finance/return_receipts",json={"delivery_id":did,
            "actual_received_date":date.today().isoformat(),
            "items":[{"delivery_item_id":line_id,"actual_received_quantity":100}]})
        assert receipt.status_code == 201, receipt.text
        statement = client.post("/api/finance/statements",json={"customer_id":ids["customer"],
            "statement_month":date.today().strftime("%Y-%m"),"delivery_ids":[did]})
        assert statement.status_code == 201, statement.text
        sid=statement.json()["id"]
        detail=client.get(f"/api/finance/statements/{sid}")
        assert detail.json()["items"][0]["customer_po"] == "NEW-PO-2026"
        exported=client.get(f"/api/finance/statements/{sid}/export")
        assert exported.status_code == 200, exported.text
        wb=load_workbook(BytesIO(exported.content),read_only=True)
        assert any("NEW-PO-2026" in row for row in wb.active.iter_rows(values_only=True))
    with factory() as db:
        assert db.get(Order, original_order_id).customer_po == original_po
        assert db.get(OrderItem, ids["task_completed"]).order_id == original_order_id
        assert db.get(DeliveryItem,line_id).order_item_id == ids["task_completed"]
        assert db.scalar(select(func.count()).select_from(InventoryMovement)) == movement_count


def test_po_survives_normal_edit_and_audit_failure_rolls_back(n029_delivery_app, monkeypatch):
    import app.api.deliveries as api
    app, factory, ids = n029_delivery_app
    _prepare_103_finished_stock(factory, ids)
    with TestClient(app) as client:
        _login(client)
        payload={"customer_id":ids["customer"],"items":[{
            "order_item_id":ids["task_completed"],"delivered_quantity":100,"customer_po":"DELIVERY-ONLY"}]}
        created=client.post("/api/deliveries",json=payload)
        assert created.status_code == 201, created.text
        did=created.json()["id"]
        updated=client.put(f"/api/deliveries/{did}",json={"items":[{
            "order_item_id":ids["task_completed"],"delivered_quantity":99}]})
        assert updated.status_code == 200, updated.text
        assert updated.json()["items"][0]["customer_po"] == "DELIVERY-ONLY"
        current=updated.json()
        def fail(*args, **kwargs): raise RuntimeError("audit failure")
        monkeypatch.setattr(api,"_write_audit",fail)
        import pytest
        with pytest.raises(RuntimeError,match="audit failure"):
            client.put(f"/api/deliveries/{did}/customer-po",json={
                "expected_version":current["version"],"idempotency_key":"test-audit-po-failure",
                "items":[{"delivery_item_id":current["items"][0]["id"],"customer_po":"SHOULD-ROLLBACK"}]})
        after=client.get(f"/api/deliveries/{did}").json()
        assert after["version"] == current["version"]
        assert after["items"][0]["customer_po"] == "DELIVERY-ONLY"
