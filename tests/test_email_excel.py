from io import BytesIO
from openpyxl import Workbook
import pytest
from app.services.email_excel import table, drafts


def workbook():
    book=Workbook(); sheet=book.active; sheet.title='订单明细'
    sheet.append(['说明页标题'])
    sheet.append(['存货编码','产品名称','数量','单价','客户单号','交货日期'])
    sheet.append([6,'测试箱',20,'2.30','PO-A','2026-09-15'])
    sheet['A3'].number_format='0000'
    sheet.append(['0007','测试衬板',5,'1.20','PO-B','2026-09-16'])
    sheet.append(['合计',None,25,None,None,None])
    book.create_sheet('其他页')
    content=BytesIO();book.save(content);return content.getvalue()


def test_workbook_mapping_preserves_codes_and_splits_orders():
    parsed=table(workbook(),header_row=2)
    assert parsed['sheets']==['订单明细','其他页']
    assert parsed['rows'][0][0]=='0006'
    result=drafts(parsed,parsed['columns'])
    assert [row['customer_po'] for row in result]==['PO-A','PO-B']
    assert result[0]['items'][0]['quantity']==20
    assert result[0]['items'][0]['unit_price']=='2.30'
    assert result[1]['items'][0]['delivery_date']=='2026-09-16'
    assert any('合计' in w for w in result[0]['warnings'])


def test_formula_invalid_quantity_and_unnamed_lines_are_retained():
    book=Workbook();sheet=book.active
    sheet.append(['编码','名称','数量','单价'])
    sheet.append(['A','箱','=SUM(1,2)','2.00'])
    sheet.append(['B','箱','0.5','not-price'])
    sheet.append([None,None,12,'3'])
    output=BytesIO();book.save(output)
    parsed=table(output.getvalue())
    result=drafts(parsed,{'product_code':0,'product_name':1,'quantity':2,'unit_price':3})
    assert len(result[0]['items'])==3
    assert result[0]['items'][0]['quantity'] is None
    assert result[0]['items'][1]['quantity'] is None
    assert result[0]['items'][2]['quantity']==12
    assert parsed['formula_rows']==[2]
    assert any('公式' in w for w in result[0]['warnings'])


def test_duplicate_headers_and_invalid_mapping_do_not_guess():
    book=Workbook();sheet=book.active;sheet.append(['存货编码','数量','数量']);sheet.append(['A',1,2])
    output=BytesIO();book.save(output);parsed=table(output.getvalue())
    assert 'quantity' not in parsed['columns']
    with pytest.raises(ValueError,match='数量列'):drafts(parsed,parsed['columns'])
    with pytest.raises(ValueError,match='两个字段'):drafts(parsed,{'product_code':0,'quantity':0})
    with pytest.raises(ValueError):table(b'not a workbook')
