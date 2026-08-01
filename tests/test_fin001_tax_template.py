from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
import pytest

from app.services.invoice_tax_template import (
    BLOCKING_LINE_COUNT,
    TAX_TEMPLATE_SHA256,
    TAX_TEMPLATE_SHEETS,
    InvoiceTaxTemplateError,
    generate_invoice_tax_template,
    sha256_file,
    verify_invoice_tax_template,
)


NAS_TEMPLATE = Path(
    r"Z:\sata1-18015598002\天明ERP知识库\09_待整理收件箱\2026-07-27_月结对账开票任务导入与收款闭环计划_附件\发票开具项目信息导入模板.xlsx"
)


@pytest.fixture(scope="module")
def verified_template() -> Path:
    if not NAS_TEMPLATE.is_file():
        pytest.skip("FIN-001 税局空模板当前不可访问")
    assert verify_invoice_tax_template(NAS_TEMPLATE) == TAX_TEMPLATE_SHA256
    return NAS_TEMPLATE


def _line(**overrides: object) -> dict[str, object]:
    line: dict[str, object] = {
        "project_name": "纸箱加工服务",
        "tax_classification_code": "3040101000000000000",
        "specification": "520×350×300mm",
        "unit": "个",
        "quantity": 1000,
        "unit_price": Decimal("2.35"),
        "amount": Decimal("2350.00"),
        "tax_rate": Decimal("0.13"),
        "discount_amount": Decimal("0"),
        "preferential_policy_type": "",
        "coal_type": "",
        "professional_service_invoice": "否",
        "professional_service_agreement_no": "HT-TEST-001",
    }
    line.update(overrides)
    return line


def test_export_copies_verified_template_and_preserves_helpers(
    verified_template: Path, tmp_path: Path
) -> None:
    source_hash = sha256_file(verified_template)
    output = tmp_path / "invoice-items.xlsx"

    result = generate_invoice_tax_template(
        template_path=verified_template,
        output_path=output,
        lines=[_line()],
    )

    assert output.is_file()
    assert result.template_sha256 == TAX_TEMPLATE_SHA256
    assert result.sha256 == sha256_file(output)
    assert result.total_amount == Decimal("2350.00")
    assert result.warnings == ()
    assert sha256_file(verified_template) == source_hash

    source_workbook = load_workbook(verified_template, read_only=True, data_only=False)
    exported_workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        assert tuple(exported_workbook.sheetnames) == TAX_TEMPLATE_SHEETS
        for name in TAX_TEMPLATE_SHEETS[1:]:
            assert exported_workbook[name].max_row == source_workbook[name].max_row
            assert exported_workbook[name].max_column == source_workbook[name].max_column
        detail = exported_workbook["1-明细模板"]
        assert detail["A4"].value == "纸箱加工服务"
        assert detail["B4"].value == "3040101000000000000"
        assert Decimal(str(detail["G4"].value)) == Decimal("2350.00")
        assert Decimal(str(detail["H4"].value)) == Decimal("0.13")
        assert detail["V4"].value == "HT-TEST-001"
    finally:
        source_workbook.close()
        exported_workbook.close()


def test_more_than_200_rows_returns_manual_review_warning(
    verified_template: Path, tmp_path: Path
) -> None:
    result = generate_invoice_tax_template(
        template_path=verified_template,
        output_path=tmp_path / "more-than-200.xlsx",
        lines=[_line() for _ in range(201)],
    )

    assert result.line_count == 201
    assert result.warnings == ("明细超过 200 行，请在导入税局前人工复核拆分需求。",)


def test_2000_rows_are_blocked_before_copying_template(
    verified_template: Path, tmp_path: Path
) -> None:
    output = tmp_path / "blocked.xlsx"

    with pytest.raises(InvoiceTaxTemplateError, match="2000"):
        generate_invoice_tax_template(
            template_path=verified_template,
            output_path=output,
            lines=[_line() for _ in range(BLOCKING_LINE_COUNT)],
        )

    assert not output.exists()


@pytest.mark.parametrize("field", ["project_name", "tax_classification_code", "amount", "tax_rate"])
def test_required_import_fields_are_blocking(
    verified_template: Path, tmp_path: Path, field: str
) -> None:
    line = _line(**{field: ""})
    with pytest.raises(InvoiceTaxTemplateError):
        generate_invoice_tax_template(
            template_path=verified_template,
            output_path=tmp_path / f"missing-{field}.xlsx",
            lines=[line],
        )
