from copy import deepcopy
from decimal import Decimal as D
from pathlib import Path
import zipfile
import pytest
from openpyxl import load_workbook
from app.services.invoice_import_rounding import balance_invoice_lines, money, InvoiceRoundingError
from app.services.invoice_tax_template import generate_invoice_tax_template, InvoiceTaxTemplateError, TAX_TEMPLATE_SHA256

TEMPLATE = Path(__file__).parents[1] / "app/resources/invoice_tax_templates/发票开具项目信息导入模板.xlsx"


def line(amount, quantity=1, rate="0.13"):
    return dict(project_name="test",tax_classification_code="1060105010000000000",specification=None,
                unit="PCS",quantity=quantity,unit_price=D(amount)/quantity,amount=D(amount),tax_rate=D(rate))


def plan(rows):
    net=money(sum(x["amount"] for x in rows));tax=money(sum(x["amount"]*x["tax_rate"] for x in rows))
    return balance_invoice_lines(rows,target_net=net,target_tax=tax,target_total=net+tax)


def test_joint_cent_balance_preserves_prices_quantities_source_and_is_deterministic(tmp_path):
    rows=[line("783.12",120),line("481.005",15),line("68.475",5),line("1.04")]
    before=deepcopy(rows);out,report=plan(rows)
    assert report["net"]=="1333.64" and report["tax"]=="173.37"
    assert out[0]["amount"]==D("783.11")
    assert rows==before and plan(rows)==(out,report)
    for a,b in zip(rows,out):
        assert a["unit_price"]==b["unit_price"] and a["quantity"]==b["quantity"]
        assert abs(a["amount"]-b["amount"])<=D(".01")
    export=generate_invoice_tax_template(template_path=TEMPLATE,output_path=tmp_path/'out.xlsx',lines=out)
    w=load_workbook(export.path,data_only=True);s=w.worksheets[0]
    values=[D(str(s.cell(r,7).value)) for r in range(4,8)]
    assert sum(values)==D('1333.64')
    assert sum(money(v*D('.13')) for v in values)==D('173.37')
    assert all(v==money(v) for v in values)
    assert s['C4'].value is None and s.max_row==7
    with zipfile.ZipFile(TEMPLATE) as a,zipfile.ZipFile(export.path) as b:
        for name in a.namelist():
            if name not in ('xl/worksheets/sheet1.xml','xl/sharedStrings.xml'):
                assert a.read(name)==b.read(name)


def test_no_adjustment_when_already_exact():
    rows=[line('100.00')];out,r=plan(rows)
    assert out==rows and r['adjustments']==[]


def test_impossible_tax_target_is_blocked():
    with pytest.raises(InvoiceRoundingError,match='无法同时'):
        plan([line('.02'),line('.02')])


def test_source_mismatch_cannot_be_hidden():
    rows=[line('100')];rows[0]['unit_price']=D('99')
    with pytest.raises(InvoiceRoundingError,match='数量单价'):
        plan(rows)


def test_wrong_frozen_target_is_blocked():
    with pytest.raises(InvoiceRoundingError,match='原始明细'):
        balance_invoice_lines([line('100')],target_net='99',target_tax='13',target_total='112')


def test_mixed_rates_do_not_transfer_amounts_between_rates():
    out,r=plan([line('100',rate='.13'),line('100',rate='.09')])
    assert r['total']=='222.00' and out[0]['amount']==out[1]['amount']==D('100')


def test_raw_three_decimal_amount_cannot_be_exported(tmp_path):
    with pytest.raises(InvoiceTaxTemplateError,match='实际保留两位'):
        generate_invoice_tax_template(template_path=TEMPLATE,output_path=tmp_path/'out.xlsx',lines=[line('1.005')])
    assert not (tmp_path/'out.xlsx').exists()


def test_new_template_is_empty_and_keeps_version():
    w=load_workbook(TEMPLATE,data_only=True)
    assert w.worksheets[0].max_row==3
    assert w['excelVersion']['A4'].value=='mx;20250501'
    assert w['xzqhdm'].max_row==2863


def test_export_preserves_qname_namespace_aliases_required_by_strict_importers(tmp_path):
    from io import BytesIO
    from xml.etree import ElementTree as ET
    output = tmp_path / 'strict.xlsx'
    generate_invoice_tax_template(template_path=TEMPLATE, output_path=output, lines=[line('100')])
    with zipfile.ZipFile(output) as z:
        raw = z.read('xl/styles.xml')
        ns = dict(value for _, value in ET.iterparse(BytesIO(raw), events=('start-ns',)))
        root = ET.fromstring(raw)
        key = '{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable'
        for node in root.iter():
            for prefix in node.get(key, '').split():
                assert prefix in ns, 'dangling mc:Ignorable namespace: ' + prefix
        assert ns['xr9'] == 'http://schemas.microsoft.com/office/spreadsheetml/2016/revision9'
