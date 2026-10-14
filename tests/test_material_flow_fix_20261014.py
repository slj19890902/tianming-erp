from datetime import date
import json

import pytest
from sqlalchemy import select

from tests.test_phase11_requisition import requisition_app
from app.models.requisition import Requisition, RequisitionItem
from app.models.procurement_source import ProcurementSourceLink
from app.models.supplier_requisition_order import SupplierRequisitionOrder, SupplierRequisitionOrderItem
from app.models.user import User
from app.services.sheet_cutting_contract import SheetCuttingContract


def seed(db, *, quantities=(267, 800, 850), typed=True):
    header = SupplierRequisitionOrder(order_number='SRO208-REPRO', supplier_name='欢达',
        crease_type='净料', total_quantity=sum(quantities), requisition_qty=sum(quantities))
    batch = Requisition(requisition_number='BOM-REPRO', requisition_date=date(2026,10,14))
    db.add_all([header, batch]); db.flush()
    rows = []
    for index, qty in enumerate(quantities):
        snapshot = SheetCuttingContract(375,236,2,1,True,3 if index == 0 else 1,
            actual_supplier_length_mm=750 if index else None,
            actual_supplier_width_mm=236 if index else None).to_snapshot()
        row = SupplierRequisitionOrderItem(supplier_order_id=header.id,
            order_item_id=None if typed else 1, quantity=qty, requisition_qty=qty,
            material_code_snapshot='VSNIV', layer_count_snapshot=5, flute_type_snapshot='AB',
            report_length_mm=750, report_width_mm=236, sheet_cutting_snapshot=snapshot,
            cutting_mode='一开六' if index == 0 else '一开二')
        db.add(row); db.flush()
        if typed:
            source = RequisitionItem(requisition_id=batch.id, order_item_id=1,
                requisition_qty=qty, cardboard_len=750, cardboard_width=236,
                product_name_snapshot=f'组件{index}')
            db.add(source); db.flush()
            db.add(ProcurementSourceLink(supplier_item_id=row.id,
                material_requisition_item_id=source.id, source_quantity=qty, created_by=1,
                source_snapshot_json=json.dumps({'crease_type':'净料','direction_note':None})))
        rows.append(row)
    db.commit()
    return header, rows


def test_same_delivered_board_merges_schema_and_internal_mold_yield(requisition_app):
    from app.api.requisition import _supplier_order_purchase_lines
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        lines = _supplier_order_purchase_lines(header, db)
        assert len(lines) == 1
        assert lines[0]['requisition_qty'] == 1917
        assert [s['requisition_qty'] for s in lines[0]['source_items']] == [267,800,850]


def test_admin_withdraw_finds_indirect_and_preserves_history(requisition_app):
    from app.services.admin_order_reversal import _void_order_purchase_lines
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        _void_order_purchase_lines(db, user=db.get(User,1), item_ids=[1],
            operation_key='typed-withdraw', reason='测试')
        db.commit()
        assert all(r.status == 'voided' for r in rows)
        assert all(l.status == 'voided' for l in db.scalars(select(ProcurementSourceLink)))
        assert len(header.items) == 3


@pytest.mark.parametrize('difference', ['direction', 'sheet_type', 'crease', 'supplier_cut', 'remark', 'missing'])
def test_delivery_merge_respects_real_supplier_requirements(requisition_app, difference):
    from app.api.requisition import _supplier_order_purchase_lines
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db, quantities=(267,800))
        link = db.scalar(select(ProcurementSourceLink).where(ProcurementSourceLink.supplier_item_id == rows[1].id))
        snapshot = json.loads(link.source_snapshot_json)
        if difference == 'direction': snapshot['direction_note'] = '瓦楞沿宽'
        if difference == 'sheet_type': snapshot['sheet_type'] = 'raw_board'
        if difference == 'crease': snapshot['crease_type'] = '毛片'
        if difference == 'remark': snapshot['supplier_work'] = '供应商先修边'
        if difference == 'missing': snapshot.pop('crease_type')
        if difference == 'supplier_cut':
            rows[1].sheet_cutting_snapshot = SheetCuttingContract(750,118,1,2,True,1).to_snapshot()
        link.source_snapshot_json = json.dumps(snapshot)
        db.flush()
        lines = _supplier_order_purchase_lines(header, db)
        assert len(lines) == 2
        assert sum(r['requisition_qty'] for r in lines) == 1067
        if difference == 'missing':
            assert 'crease_type' in lines[1]['supplier_delivery_contract_issues']


def other_source(db, header):
    from app.models.order import Order, OrderItem
    order = Order(order_number='OTHER-PURCHASE', customer_id=1, order_date=date(2026,10,14))
    db.add(order); db.flush()
    item = OrderItem(order_id=order.id, product_id=1, quantity=5, unit_price=1,
        subtotal=5, snapshot_product_name='其他来源', snapshot_product_code='OTHER')
    db.add(item); db.flush()
    row = SupplierRequisitionOrderItem(supplier_order_id=header.id, order_item_id=item.id,
        quantity=5, requisition_qty=5)
    db.add(row); db.flush()
    header.total_quantity += 5; header.requisition_qty += 5
    return row


@pytest.mark.parametrize('legacy', [False, True])
def test_shared_mixed_direct_typed_keeps_foreign_source_and_is_idempotent(requisition_app, legacy):
    from app.api.orders import _rollback_supplier_requisition_items
    from app.services.admin_order_reversal import _void_order_purchase_lines
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        rows[0].order_item_id = 1  # Both direct and typed; must only count once.
        foreign = other_source(db, header); db.commit()
        user = db.get(User,1)
        def action():
            if legacy:
                return _rollback_supplier_requisition_items(db, order_item_ids=[1], user=user)
            return _void_order_purchase_lines(db, user=user, item_ids=[1], operation_key='mixed-withdraw', reason='测试')
        action(); db.commit()
        versions = [r.version for r in rows]
        action(); db.commit()
        assert [r.version for r in rows] == versions
        assert foreign.status == 'active' and header.requisition_qty == header.total_quantity == 5
        assert len(header.items) == 4
        assert all(l.status == 'voided' for l in db.scalars(select(ProcurementSourceLink)))


def receipt(db, source_id):
    from datetime import datetime
    from app.models.incoming_receipt import IncomingReceipt, IncomingReceiptItem
    header = IncomingReceipt(receipt_number='PARTIAL', idempotency_key='partial-receipt', received_at=datetime(2026,10,14))
    db.add(header); db.flush()
    fact = IncomingReceiptItem(receipt_id=header.id, order_id=1, order_item_id=1,
        requisition_item_id=source_id, planned_quantity=800, received_quantity=100,
        cumulative_received_quantity=100, variance_quantity=-700, variance_type='short',
        resolution_status='pending', status='posted')
    db.add(fact); db.commit()
    return fact


def test_partial_receipt_blocks_source_void_and_transaction_restores_earlier_lines(requisition_app):
    from fastapi import HTTPException
    from app.services.admin_order_reversal import _void_order_purchase_lines
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        link = db.scalar(select(ProcurementSourceLink).where(ProcurementSourceLink.supplier_item_id == rows[1].id))
        fact = receipt(db, link.material_requisition_item_id)
        with pytest.raises(HTTPException, match='有效收料'):
            _void_order_purchase_lines(db, user=db.get(User,1), item_ids=[1], operation_key='partial-void', reason='测试')
        db.rollback()
        assert all(r.status == 'active' and r.version == 1 for r in rows)
        assert all(l.status == 'active' for l in db.scalars(select(ProcurementSourceLink)))
        assert fact.status == 'posted' and header.requisition_qty == 1917


def test_preview_execute_hash_replay_and_audit_rollback(requisition_app, monkeypatch):
    from app.services import admin_order_reversal as service, audit_log
    from app.models.order import Order
    from fastapi import HTTPException
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        user = db.get(User,1)
        plan = service.preview(db, user=user, order_ids=[1], mode='withdraw')
        assert plan['counts']['purchases'] == 1 and not plan['blockers']
        args = dict(user=user, order_ids=[1],mode='withdraw', reviewed_hash=plan['reviewed_hash'],
            operation_key='atomic-typed-withdraw', trial_confirmed=False, reason='测试')
        original = audit_log.append_audit_event
        def fail(*a, **kw):
            if kw.get('action_code') == 'order.admin_disposition':
                raise HTTPException(409, '故障注入审计')
            return original(*a, **kw)
        with monkeypatch.context() as patch:
            patch.setattr(audit_log, 'append_audit_event', fail)
            with pytest.raises(HTTPException, match='故障注入审计'):
                service.execute(db, **args)
        db.rollback()
        assert all(r.status == 'active' for r in rows)
        result = service.execute(db, **args); db.commit()
        assert service.execute(db, **args) == result
        with pytest.raises(HTTPException, match='其他内容'):
            service.execute(db, **{**args, 'reason':'changed'})
        db.rollback()
        assert db.get(Order,1).status == 'pending_production'
        assert all(r.status == 'voided' for r in rows)


def test_cancelled_history_preview_is_read_only_and_cannot_execute(requisition_app):
    from app.services import admin_order_reversal as service
    from app.models.order import Order
    from fastapi import HTTPException
    from sqlalchemy import event
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db, quantities=(75,150))
        db.get(Order,1).status = 'cancelled'; db.commit()
        before = list(db.connection().connection.driver_connection.iterdump())
        user = db.get(User,1)
        def read_only(conn,cursor,statement,params,ctx,many):
            assert not statement.lstrip().upper().startswith(('INSERT','UPDATE','DELETE','REPLACE'))
        event.listen(factory.kw['bind'], 'before_cursor_execute', read_only)
        try:
            plan = service.preview(db, user=user, order_ids=[1], mode='delete_order')
        finally:
            event.remove(factory.kw['bind'], 'before_cursor_execute', read_only)
        assert [r['supplier_item_id'] for r in plan['procurement_corrections']] == [r.id for r in rows]
        assert all(not r['executable'] and r['requires_supplier_commitment_review'] for r in plan['procurement_corrections'])
        assert plan['blockers']
        assert before == list(db.connection().connection.driver_connection.iterdump())
        with pytest.raises(HTTPException, match='供应商下单承诺'):
            service.execute(db, user=user, order_ids=[1], mode='delete_order', reviewed_hash=plan['reviewed_hash'],
                operation_key='cancelled-history', trial_confirmed=False, reason='不可执行')
        db.rollback()
        assert all(r.status == 'active' for r in rows)


def test_external_exports_share_merged_projection(requisition_app):
    from app.api.requisition import _supplier_order_dict
    from app.services.supplier_purchase_view import supplier_rows, supplier_pdf, supplier_xlsx
    from io import BytesIO
    from openpyxl import load_workbook
    from pypdf import PdfReader
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        purchase = _supplier_order_dict(header, db)
        assert len(purchase['lines']) == len(supplier_rows(purchase)) == 1
        assert load_workbook(BytesIO(supplier_xlsx(purchase))).active['E6'].value == 1917
        text = ''.join(p.extract_text() for p in PdfReader(BytesIO(supplier_pdf(purchase))).pages)
        assert '1917' in text and '750' in text and '236' in text and 'VSNIV' in text
        assert '一开六' not in text and 'mold_count' not in text
        assert '一开六' not in str(list(load_workbook(BytesIO(supplier_xlsx(purchase))).active.values))
        from pathlib import Path
        html = (Path(__file__).resolve().parents[1] / 'static/index.html').read_text()
        rendered = html[html.index('id="supplier-order-print-area"'):html.index('厂内采购用途只读详情')]
        assert 'line.cutting_mode' not in rendered and 'line.sheet_cutting_snapshot' not in rendered


def test_inactive_purchase_header_retains_old_rollback_semantics(requisition_app):
    from app.api.orders import _rollback_supplier_requisition_items
    from app.services.procurement_sources import order_purchase_lines
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        header.status = 'voided'; db.commit()
        assert order_purchase_lines(db,[1],active_only=True) == []
        assert _rollback_supplier_requisition_items(db,order_item_ids=[1],user=db.get(User,1)) == []
        assert all(r.status == 'active' for r in rows)


def test_draft_requisition_order_still_deletes_without_retained_fk(requisition_app):
    from app.api.orders import _delete_orders_in_transaction
    from app.models.order import Order
    _, factory = requisition_app
    with factory() as db:
        batch = Requisition(requisition_number='DRAFT-DELETE',requisition_date=date(2026,10,14),status='draft')
        db.add(batch);db.flush()
        db.add(RequisitionItem(requisition_id=batch.id,order_item_id=1,requisition_qty=10,
            cardboard_len=750,cardboard_width=236,product_name_snapshot='草稿',status='已取消'))
        db.commit()
        _delete_orders_in_transaction(db,orders=[db.get(Order,1)],user=db.get(User,1))
        assert db.get(Order,1) is None
        assert list(db.scalars(select(RequisitionItem))) == []


def test_legacy_workflow_preserves_typed_cancelled_history_and_baseline(requisition_app):
    from app.api.orders import rollback_order_workflow, WorkflowRollbackRequest, _already_at_workflow_rollback_baseline
    from app.models.order import Order
    _, factory = requisition_app
    with factory() as db:
        header, rows = seed(db)
        user = db.get(User,1)
        rollback_order_workflow(order_id=1,payload=WorkflowRollbackRequest(),db=db,user=user)
        assert all(r.status=='voided' for r in rows)
        assert all(r.status=='已取消' for r in db.scalars(select(RequisitionItem)))
        assert _already_at_workflow_rollback_baseline(db,order=db.get(Order,1),item_ids=[1])


def test_reported_reversal_uses_permission_checked_history_scope():
    from pathlib import Path
    source = (Path(__file__).resolve().parents[1] / 'static/index.html').read_text()
    assert 'this.openOrderEditor(reversal.order_group,{scope:"all"})' in source
    assert 'scope: scope || this.filters.orderScope || "active"' in source


def cancelled_source(db):
    from app.models.order import Order
    header, rows = seed(db, quantities=(75,150,10))
    db.get(Order,1).status = 'cancelled'
    for source in db.scalars(select(RequisitionItem)):
        source.status = '已取消'
    db.commit()
    return header, rows


def test_reviewed_cancelled_correction_exact_scope_versions_audit_and_replay(requisition_app,monkeypatch):
    from app.services import procurement_cancelled_correction as service, audit_log
    from app.models.order import Order
    from fastapi import HTTPException
    _, factory = requisition_app
    with factory() as db:
        header, rows = cancelled_source(db)
        targets = [r.id for r in rows[:2]]
        args = dict(user=db.get(User,1),order_ids=[1],supplier_item_ids=targets,supplier_not_sent_confirmed=True)
        plan = service.preview(db, **args)
        assert plan['executable'] and plan['unchanged_shared_item_ids'] == [rows[2].id]
        execute = dict(**args,reviewed_hash=plan['reviewed_hash'],operation_key='review-cancelled',reason='用户确认两条均未向供应商下单')
        original = audit_log.append_audit_event
        with monkeypatch.context() as patch:
            def fail(*a,**kw):
                if kw['action_code'] == 'procurement.cancelled_order_correction':
                    raise HTTPException(409,'更正审计故障')
                return original(*a,**kw)
            patch.setattr(audit_log,'append_audit_event',fail)
            with pytest.raises(HTTPException, match='更正审计故障'):
                service.execute(db,**execute)
        db.rollback()
        assert all(r.status == 'active' and r.version == 1 for r in rows)
        rows[0].version += 1; db.commit()
        with pytest.raises(HTTPException,match='已变化'):
            service.execute(db,**execute)
        db.rollback()
        plan = service.preview(db,**args); execute['reviewed_hash'] = plan['reviewed_hash']
        result = service.execute(db,**execute); db.commit()
        assert service.execute(db,**execute) == result
        with pytest.raises(HTTPException,match='其他内容'):
            service.execute(db,**{**execute,'supplier_item_ids':[rows[2].id]})
        db.rollback()
        assert [r.status for r in rows] == ['voided','voided','active']
        assert header.requisition_qty == 10
        assert db.get(Order,1).status == 'cancelled'
        assert all(s.status == '已取消' for s in db.scalars(select(RequisitionItem)))


def test_cancelled_correction_requires_physical_fact_and_no_receipt(requisition_app):
    from app.services import procurement_cancelled_correction as service
    _, factory = requisition_app
    with factory() as db:
        header, rows = cancelled_source(db)
        args = dict(user=db.get(User,1),order_ids=[1],supplier_item_ids=[rows[0].id])
        assert not service.preview(db,**args)['executable']
        link = db.scalar(select(ProcurementSourceLink).where(ProcurementSourceLink.supplier_item_id==rows[0].id))
        receipt(db,link.material_requisition_item_id)
        plan = service.preview(db,**args,supplier_not_sent_confirmed=True)
        assert not plan['executable'] and any('有效实收' in b for b in plan['blockers'])


def verify_real_snapshot():
    """Foreground read-only probe, only against run_task's owned child copy."""
    import os
    from sqlalchemy import create_engine, event
    from sqlalchemy.orm import Session
    from app.api.requisition import _supplier_order_dict
    from app.services.admin_order_reversal import preview
    database = os.environ['ERP_DATABASE_PATH']
    assert os.environ.get('ERP_UAT_ISOLATION_ID') and '/copies/' in database
    engine = create_engine(f'sqlite:///file:{database}?mode=ro&uri=true')
    @event.listens_for(engine, 'connect')
    def query_only(conn, _): conn.execute('PRAGMA query_only=ON')
    with Session(engine) as db:
        purchase = _supplier_order_dict(db.get(SupplierRequisitionOrder,208), db)
        groups = [dict(quantity=r['requisition_qty'], source_ids=[s['id'] for s in r['source_items']],
            contract=r['supplier_delivery_contract'], issues=r['supplier_delivery_contract_issues']) for r in purchase['lines']]
        assert len(groups) == 1 and groups[0]['quantity'] == 1917 and groups[0]['source_ids'] == [620,621,622], groups
        assert [s['requisition_qty'] for s in purchase['lines'][0]['source_items']] == [267,800,850]
        admin = db.scalar(select(User).where(User.role == 'admin').order_by(User.id))
        plan = preview(db, user=admin, order_ids=[9781], mode='delete_order')
        assert [r['supplier_item_id'] for r in plan['procurement_corrections']] == [605,606], plan
        assert all(not r['executable'] for r in plan['procurement_corrections']) and plan['blockers']
        print(json.dumps(dict(database=database,business_written=False,groups=groups,
            corrections=plan['procurement_corrections'],reviewed_hash=plan['reviewed_hash'],blockers=plan['blockers']),ensure_ascii=False))


def verify_real_correction():
    import os
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from app.models.order import Order
    from app.models.audit import OperationLog
    from app.services import procurement_cancelled_correction as service
    database = os.environ['ERP_DATABASE_PATH']
    assert os.environ.get('ERP_UAT_ISOLATION_ID') and '/copies/' in database
    with Session(create_engine(f'sqlite:///{database}')) as db:
        user = db.scalar(select(User).where(User.role=='admin').order_by(User.id))
        args = dict(user=user,order_ids=[9781],supplier_item_ids=[605,606],supplier_not_sent_confirmed=True)
        plan = service.preview(db,**args)
        assert plan['executable'],plan
        assert [(r['id'],r['version'],r['requisition_qty']) for r in plan['before']] == [(605,1,75),(606,1,150)]
        payload = dict(**args,reviewed_hash=plan['reviewed_hash'],operation_key='material-flow-owned-rehearsal',
            reason='2026-10-14用户明确确认80012500本体75张、补强150张两条均未向供应商下单；仅隔离副本演练')
        result = service.execute(db,**payload); db.commit()
        assert service.execute(db,**payload)==result
        assert db.get(Order,9781).status=='cancelled'
        assert [db.get(SupplierRequisitionOrderItem,i).status for i in [605,606]] == ['voided','voided']
        assert [db.get(RequisitionItem,i).status for i in [204,205]] == ['已取消','已取消']
        logs = list(db.scalars(select(OperationLog).where(OperationLog.action_code=='procurement.cancelled_order_correction')))
        assert len(logs)==1
        print(json.dumps(dict(database=database,owned_copy_written=True,production_written=False,preview=plan,
            result=result,audit_id=logs[0].id,order_status='cancelled',internal_statuses=['已取消','已取消']),ensure_ascii=False,default=str))


if __name__ == '__main__':
    import sys
    verify_real_correction() if '--correct-owned-copy' in sys.argv else verify_real_snapshot()
