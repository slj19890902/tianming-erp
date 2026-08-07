from __future__ import annotations

from io import BytesIO
import json

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import select
from sqlalchemy.orm import sessionmaker


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def mixed_sample_app(tmp_path, monkeypatch):
    from app.api.auth import router as auth_router
    from app.api.deps import get_db
    from app.api.product_import import router as product_import_router
    from app.core.database import create_sqlite_engine
    from app.core.security import hash_password
    from app.models import Base
    from app.models.customer import Customer
    from app.models.user import User

    monkeypatch.setenv("ERP_FILE_STORAGE_DIR", str(tmp_path / "private_uploads"))
    monkeypatch.setenv("ERP_UPLOAD_TEMP_DIR", str(tmp_path / "temporary_uploads"))
    engine = create_sqlite_engine(tmp_path / "mixed-samples.sqlite3")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add_all(
            [
                User(
                    username="mixed-sample-admin",
                    password_hash=hash_password("WorkbookPass123!"),
                    role="admin",
                    real_name="混合样品管理员",
                    must_change_password=False,
                ),
                Customer(customer_number=9101, customer_code="YL", name="仪菱"),
                Customer(customer_number=9102, customer_code="YKE", name="研光"),
                Customer(customer_number=9103, customer_code="KEW", name="光洋"),
            ]
        )
        db.commit()

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(product_import_router, prefix="/api/master/products")

    def override_get_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    app.state.session_factory = factory
    yield app
    engine.dispose()


def _login(client: TestClient) -> None:
    response = client.post(
        "/api/auth/login",
        json={"username": "mixed-sample-admin", "password": "WorkbookPass123!"},
    )
    assert response.status_code == 200


def _xlsx(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _reference_workbook() -> bytes:
    workbook = Workbook()
    yl = workbook.active
    yl.title = "Y L"
    yl.append(
        [
            "物料代号", "料号", "箱数", "制造尺寸", "压线或剖切", "门幅",
            "开数", "总长", "系数", "张数", "材质", "纸价", "总成本",
            "报价", "库存单价", "天明单价",
        ]
    )
    yl.append(
        [
            "Z.001.000001", "YL测试箱", 1, "100*80*60", "70*60*70=200",
            200, 1, 400, 1, 1, "VIK（B）", 1, 1, 1, 1, 2.5,
        ]
    )
    for code, name in (("YKE", "研光共享箱"), ("KEW", "光洋共享箱")):
        sheet = workbook.create_sheet(code)
        row = [None] * 19
        row[1] = name
        row[2] = f"DRAW-{code}"
        row[3] = "SHARED001"
        row[8] = "300*200*100"
        row[9] = "120*100*120=340"
        row[10] = 340
        row[12] = 1000
        row[15] = "VSNIV/AB"
        row[17] = 6.8
        row[18] = "红钉"
        sheet.append(row)
    return _xlsx(workbook)


def _registration_workbook() -> bytes:
    from app.services.mixed_sample_import import build_mixed_sample_workbook

    workbook = load_workbook(BytesIO(build_mixed_sample_workbook(blank_rows=2)))
    sheet = workbook["样品录入"]
    sheet["A2"] = "YP001"
    sheet["B2"] = "SHARED001"
    sheet["C2"] = ""
    sheet["D2"] = "AB"
    sheet["E2"] = "开槽"
    sheet["F2"] = "是"
    sheet["G2"] = "红色"
    sheet["H2"] = "打钉"
    sheet["I2"] = "否"
    sheet["J2"] = ""
    return _xlsx(workbook)


def _jpeg_bytes() -> bytes:
    from PIL import Image

    output = BytesIO()
    Image.new("RGB", (160, 100), "white").save(output, "JPEG", quality=80)
    return output.getvalue()


def test_mixed_template_replaces_csv_and_bat_workflow() -> None:
    from app.services.mixed_sample_import import build_mixed_sample_workbook

    workbook = load_workbook(BytesIO(build_mixed_sample_workbook(blank_rows=5)))
    try:
        assert workbook.sheetnames == ["使用说明", "样品录入", "导入信息"]
        assert workbook["样品录入"]["C1"].value == "客户代码(可空)"
        guide_text = " ".join(
            str(cell.value or "") for row in workbook["使用说明"] for cell in row
        )
        assert "无需 CSV" in guide_text
        assert "无需先压缩" in guide_text
        assert workbook["导入信息"].sheet_state == "veryHidden"
    finally:
        workbook.close()


def test_mixed_preview_requires_explicit_customer_then_applies_atomically(
    mixed_sample_app: FastAPI,
) -> None:
    from app.models.customer import Customer
    from app.models.product import Product
    from app.models.product_drawing import ProductDrawing

    with TestClient(mixed_sample_app) as client:
        _login(client)
        files = [
            ("reference_file", ("三客户基础资料.xlsx", _reference_workbook(), EXCEL_MIME)),
            ("files", ("现场登记.xlsx", _registration_workbook(), EXCEL_MIME)),
            ("drawings", ("YP001_1.jpg", _jpeg_bytes(), "image/jpeg")),
        ]
        first = client.post(
            "/api/master/products/mixed-import/preview",
            data={"overrides": "{}"},
            files=files,
        )
        assert first.status_code == 200, first.text
        first_payload = first.json()
        assert first_payload["valid"] is False
        assert first_payload["summary"]["samples"] == 1, first_payload
        assert len(first_payload["unresolved"]) == 1
        candidates = first_payload["unresolved"][0]["candidates"]
        assert {item["customer_code"] for item in candidates} == {"YKE", "KEW"}
        yke_key = next(item["key"] for item in candidates if item["customer_code"] == "YKE")

        second = client.post(
            "/api/master/products/mixed-import/preview",
            data={"overrides": '{"YP001":"' + yke_key + '"}'},
            files=[
                ("reference_file", ("三客户基础资料.xlsx", _reference_workbook(), EXCEL_MIME)),
                ("files", ("现场登记.xlsx", _registration_workbook(), EXCEL_MIME)),
                ("drawings", ("YP001_1.jpg", _jpeg_bytes(), "image/jpeg")),
            ],
        )
        assert second.status_code == 200, second.text
        payload = second.json()
        assert payload["valid"] is True, payload
        assert payload["summary"]["drawings"] == 1
        assert payload["groups"][0]["customer_code"] == "YKE"

        applied = client.post(
            "/api/master/products/mixed-import/apply",
            json={"preview_token": payload["preview_token"]},
        )
        assert applied.status_code == 200, applied.text

    with mixed_sample_app.state.session_factory() as db:
        products = db.scalars(select(Product).order_by(Product.id)).all()
        assert len(products) == 1
        customer = db.get(Customer, products[0].customer_id)
        assert customer.customer_code == "YKE"
        assert products[0].product_code == "SHARED001"
        assert products[0].legacy_material_text == "VSNIV/AB"
        assert products[0].report_width_mm == 340
        assert products[0].report_length_mm == 1000
        assert products[0].production_process == "开槽、印刷、打钉"
        assert db.scalar(select(ProductDrawing).where(ProductDrawing.product_id == products[0].id)) is not None


def test_mixed_preview_can_apply_one_sample_to_two_customers_with_shared_photo(
    mixed_sample_app: FastAPI,
) -> None:
    from app.models.customer import Customer
    from app.models.product import Product
    from app.models.product_drawing import ProductDrawing

    with TestClient(mixed_sample_app) as client:
        _login(client)
        first = client.post(
            "/api/master/products/mixed-import/preview",
            data={"overrides": "{}"},
            files=[
                ("reference_file", ("三客户基础资料.xlsx", _reference_workbook(), EXCEL_MIME)),
                ("files", ("现场登记.xlsx", _registration_workbook(), EXCEL_MIME)),
                ("drawings", ("YP001_1.jpg", _jpeg_bytes(), "image/jpeg")),
            ],
        )
        candidates = first.json()["unresolved"][0]["candidates"]
        selected = [
            candidate["key"]
            for candidate in candidates
            if candidate["customer_code"] in {"YKE", "KEW"}
        ]
        second = client.post(
            "/api/master/products/mixed-import/preview",
            data={"overrides": json.dumps({"YP001": selected})},
            files=[
                ("reference_file", ("三客户基础资料.xlsx", _reference_workbook(), EXCEL_MIME)),
                ("files", ("现场登记.xlsx", _registration_workbook(), EXCEL_MIME)),
                ("drawings", ("YP001_1.jpg", _jpeg_bytes(), "image/jpeg")),
            ],
        )
        payload = second.json()
        assert second.status_code == 200, second.text
        assert payload["valid"] is True, payload
        assert payload["summary"]["resolved_samples"] == 2
        assert {group["customer_code"] for group in payload["groups"]} == {"YKE", "KEW"}
        applied = client.post(
            "/api/master/products/mixed-import/apply",
            json={"preview_token": payload["preview_token"]},
        )
        assert applied.status_code == 200, applied.text

    with mixed_sample_app.state.session_factory() as db:
        products = db.scalars(select(Product).order_by(Product.id)).all()
        assert len(products) == 2
        assert {
            db.get(Customer, product.customer_id).customer_code for product in products
        } == {"YKE", "KEW"}
        drawings = db.scalars(select(ProductDrawing).order_by(ProductDrawing.id)).all()
        assert len(drawings) == 2
        assert {drawing.product_id for drawing in drawings} == {
            product.id for product in products
        }


def test_legacy_embedded_rows_with_invalid_old_choices_are_still_read() -> None:
    from openpyxl.drawing.image import Image as ExcelImage
    from app.services.mixed_sample_import import read_mixed_sample_workbook

    # Build only the legacy workbook shape; the DB-backed export is covered by
    # the existing product-import tests, so here a minimal compatible package is enough.
    workbook = Workbook()
    guide = workbook.active
    guide.title = "使用说明"
    product = workbook.create_sheet("样品录入")
    product.append(
        [
            "样品号", "手写型号*", "尺寸(mm)", "楞型", "成型方式*", "印刷*",
            "印刷颜色", "结合方式*", "二次粘合*", "图纸文件名(多个用分号)",
            "模具编号", "现场备注", "操作*", "系统ID", "当前版本", "客户料号*",
            "产品名称*", "材质代码", "单位", "默认含税单价", "报料长(mm)",
            "报料宽(mm)", "启用",
        ]
    )
    product.append(["YP174", "80011965", None, "E", "模切", "否", "无", "不确定", "否"])
    image_stream = BytesIO(_jpeg_bytes())
    image = ExcelImage(image_stream)
    image.anchor = "J2"
    product.add_image(image)
    workbook.create_sheet("模具档案").append(["操作*", "模具编号*", "模具名称*", "固定位置*", "备注"])
    info = workbook.create_sheet("导入信息")
    info.append(["template_version", "P1-22-v2"])
    info.append(["customer_id", 0])
    content = _xlsx(workbook)

    rows, errors = read_mixed_sample_workbook(content)
    assert errors == []
    assert len(rows) == 1
    assert rows[0]["sample_id"] == "YP174"
    assert rows[0]["drawing_filenames"] == ("YP174_1.jpg",)
    assert "结合方式" in rows[0]["unresolved_fields"]
