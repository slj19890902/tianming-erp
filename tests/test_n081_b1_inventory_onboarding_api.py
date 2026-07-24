from __future__ import annotations

from collections.abc import Generator
import csv
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker

from app.api import inventory_onboarding as onboarding_api
from app.api.auth import router as auth_router
from app.api.deps import get_db
from app.api.inventory_onboarding import (
    FIELD_SHEET_HEADERS,
    TEMPLATE_HEADERS,
    router as inventory_onboarding_router,
)
from app.core.database import create_sqlite_engine
from app.core.security import hash_password
from app.models import Base
from app.models.access_control import UserCustomerScope, UserPermissionOverride
from app.models.customer import Customer
from app.models.inventory_onboarding import InventoryOnboardingBatch
from app.models.material import Material
from app.models.product import Product
from app.models.user import User
from app.models.warehouse_inventory import (
    InventoryLot,
    InventoryMovement,
    InventoryPallet,
    InventoryPalletItem,
    InventoryReservation,
    WarehouseLocation,
)
from app.services.secure_uploads import resolve_stored_reference
from app.services.warehouse_inventory import manual_finished_in


@pytest.fixture()
def onboarding_api_app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    private_root = tmp_path / "private-uploads"
    monkeypatch.setenv("ERP_FILE_STORAGE_DIR", str(private_root))
    engine = create_sqlite_engine(tmp_path / "n081-b1-api.sqlite3")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)

    with factory() as db:
        admin = User(
            username="n081-api-admin",
            password_hash=hash_password("123456"),
            role="admin",
            real_name="N081 API 管理员",
            must_change_password=False,
            customer_access_mode="all",
        )
        workshop = User(
            username="n081-api-workshop",
            password_hash=hash_password("123456"),
            role="workshop",
            real_name="N081 API 仓库员",
            must_change_password=False,
            customer_access_mode="all",
        )
        view_only = User(
            username="n081-api-view",
            password_hash=hash_password("123456"),
            role="workshop",
            real_name="N081 API 只读员",
            must_change_password=False,
            customer_access_mode="all",
        )
        scoped = User(
            username="n081-api-scoped",
            password_hash=hash_password("123456"),
            role="workshop",
            real_name="N081 API 客户受限员",
            must_change_password=False,
            customer_access_mode="selected",
        )
        sales = User(
            username="n081-api-sales",
            password_hash=hash_password("123456"),
            role="sales",
            real_name="N081 API 销售员",
            must_change_password=False,
            customer_access_mode="all",
        )
        customer = Customer(
            customer_number=8181,
            customer_code="N081-API-C",
            name="N081 API 测试客户",
            payment_term_days=0,
            credit_limit=0,
        )
        material = Material(
            code="K=A",
            layer_count=5,
            flute_type="AB",
            is_active=True,
        )
        db.add_all(
            [
                admin,
                workshop,
                view_only,
                scoped,
                sales,
                customer,
                material,
            ]
        )
        db.flush()
        db.add_all(
            [
                UserPermissionOverride(
                    user_id=view_only.id,
                    permission_code="warehouse.stocktake.submit",
                    is_allowed=False,
                    granted_by=admin.id,
                ),
                UserCustomerScope(
                    user_id=scoped.id,
                    customer_id=customer.id,
                    assigned_by=admin.id,
                ),
            ]
        )
        product = Product(
            customer_id=customer.id,
            product_code="N081-API-P001",
            customer_material_code="N081-API-P001",
            product_name="N081 API 成品纸箱",
            box_category="normal",
            material_id=material.id,
            default_material_code=material.code,
            flute_type="AB",
            layer_count=5,
        )
        location = WarehouseLocation(
            location_code="E1-R01",
            location_name="三楼 E1-R01",
            warehouse_type="finished",
            warehouse_floor=3,
            area_code="E1",
            storage_type="ground",
            placement_status="placed",
            source_version="V11",
        )
        db.add_all([product, location])
        db.commit()
        ids = {
            "admin": admin.id,
            "customer": customer.id,
            "product": product.id,
            "location": location.id,
        }

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(
        inventory_onboarding_router,
        prefix="/api/warehouse",
    )

    def override_get_db() -> Generator[Session, None, None]:
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    try:
        yield app, ids, factory, private_root
    finally:
        engine.dispose()


def _login(client: TestClient, username: str) -> None:
    response = client.post(
        "/api/auth/login",
        json={"username": username, "password": "123456"},
    )
    assert response.status_code == 200, response.text


def _inventory_csv(*, quantity: int = 12) -> bytes:
    values = {
        "盘点日期": "2026-07-23",
        "盘点人": "API盘点员",
        "库存类型": "成品",
        "归属类型": "客户专用",
        "楼层": 3,
        "区域": "E1",
        "库位编码": "E1-R01",
        "栈板号": "STK-API-FG-001",
        "客户编码": "N081-API-C",
        "客户名称": "N081 API 测试客户",
        "存货编码": "N081-API-P001",
        "产品名称": "N081 API 成品纸箱",
        "数量": quantity,
        "单位": "boxes",
        "日期可信度": "unknown",
        "入库日期原文": "历史入库日期不明",
        "备注": "N081-B1 API 定向测试",
    }
    stream = StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=TEMPLATE_HEADERS)
    writer.writeheader()
    writer.writerow(values)
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _formal_inventory_counts(factory: sessionmaker) -> tuple[int, ...]:
    with factory() as db:
        return tuple(
            int(db.scalar(select(func.count(model.id))) or 0)
            for model in (
                InventoryLot,
                InventoryMovement,
                InventoryPallet,
                InventoryPalletItem,
                InventoryReservation,
            )
        )


def _import_batch(client: TestClient) -> dict[str, object]:
    response = client.post(
        "/api/warehouse/inventory-onboarding/batches/import",
        files={"file": ("N081盘点.csv", _inventory_csv(), "text/csv")},
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["created"] is True
    return payload


def test_templates_require_view_permission_and_are_valid(onboarding_api_app) -> None:
    app, _ids, _factory, _private_root = onboarding_api_app
    with TestClient(app) as client:
        unauthenticated = client.get(
            "/api/warehouse/inventory-onboarding/template?format=csv"
        )
        assert unauthenticated.status_code == 401

        _login(client, "n081-api-sales")
        forbidden = client.get(
            "/api/warehouse/inventory-onboarding/template?format=csv"
        )
        assert forbidden.status_code == 403

        _login(client, "n081-api-view")
        csv_response = client.get(
            "/api/warehouse/inventory-onboarding/template?format=csv"
        )
        assert csv_response.status_code == 200
        assert csv_response.content.startswith(b"\xef\xbb\xbf")
        assert "盘点日期" in csv_response.content.decode("utf-8-sig")

        xlsx_response = client.get(
            "/api/warehouse/inventory-onboarding/template?format=xlsx"
        )
        assert xlsx_response.status_code == 200
        workbook = load_workbook(BytesIO(xlsx_response.content), read_only=True)
        sheet = workbook["库存建账"]
        assert tuple(cell.value for cell in next(sheet.iter_rows())) == TEMPLATE_HEADERS
        workbook.close()

        write_forbidden = client.post(
            "/api/warehouse/inventory-onboarding/batches/import",
            files={"file": ("N081盘点.csv", _inventory_csv(), "text/csv")},
        )
        assert write_forbidden.status_code == 403


def test_field_sheet_exports_current_inventory_and_round_trips_count(
    onboarding_api_app,
) -> None:
    app, ids, factory, _private_root = onboarding_api_app
    with factory() as db:
        lot = manual_finished_in(
            db,
            customer_id=ids["customer"],
            product_id=ids["product"],
            location_id=ids["location"],
            quantity=12,
            stock_date=date(2026, 7, 23),
            source_type="stocktake",
            remarks="现场盘点导出测试",
            operator_id=ids["admin"],
            idempotency_key="n081-field-sheet-lot",
            pallet_code="PLT-N081-FIELD-SHEET",
        )
        db.commit()
        lot_id = lot.id
        lot_version = lot.version

    with TestClient(app) as client:
        unauthenticated = client.get(
            "/api/warehouse/inventory-onboarding/field-sheet"
        )
        assert unauthenticated.status_code == 401

        _login(client, "n081-api-view")
        response = client.get(
            "/api/warehouse/inventory-onboarding/field-sheet"
        )
        assert response.status_code == 200, response.text
        assert "filename*=UTF-8''" in response.headers["content-disposition"]

        workbook = load_workbook(BytesIO(response.content))
        assert workbook.sheetnames == ["成品现场盘点"]
        sheet = workbook["成品现场盘点"]
        assert tuple(cell.value for cell in sheet[1]) == FIELD_SHEET_HEADERS
        assert sheet["A2"].value == "P001"
        assert sheet["B2"].value == "E1-R01"
        assert sheet["C2"].value is None
        assert sheet["D2"].value == "N081 API 测试客户"
        assert "N081-API-P001" in sheet["E2"].value
        assert sheet["F2"].value == 12
        assert sheet["G2"].value == 12
        assert sheet["I2"].value == lot_id
        assert sheet["J2"].value == lot_version
        assert all(
            sheet.column_dimensions[
                sheet.cell(row=1, column=column).column_letter
            ].hidden
            is True
            for column in range(9, len(FIELD_SHEET_HEADERS) + 1)
        )
        assert sheet.freeze_panes == "A2"
        sheet["G2"] = 15
        sheet["C3"] = "东墙第3垛"
        sheet["D3"] = "N081 API 测试客户"
        sheet["E3"] = "N081-API-P001"
        sheet["G3"] = 5
        payload = BytesIO()
        workbook.save(payload)
        workbook.close()

        _login(client, "n081-api-workshop")
        imported = client.post(
            "/api/warehouse/inventory-onboarding/batches/import",
            files={
                "file": (
                    "现场盘点.xlsx",
                    payload.getvalue(),
                    (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )
            },
        )
        assert imported.status_code == 201, imported.text
        existing_line, new_line = imported.json()["items"]
        assert existing_line["action_decision"] == "route_n035"
        assert existing_line["match_status"] == "routed"
        assert existing_line["existing_lot_id"] == lot_id
        assert existing_line["quantity"] == 15
        assert new_line["action_decision"] == "create_new"
        assert new_line["match_status"] == "ready"
        assert new_line["inventory_type"] == "finished"
        assert new_line["ownership_type"] == "customer_specific"
        assert new_line["unit"] == "boxes"
        assert new_line["customer_id"] == ids["customer"]
        assert new_line["product_id"] == ids["product"]
        assert new_line["quantity"] == 5
        assert new_line["location_id"] is None
        assert new_line["location_code"].startswith("PD-")
        assert new_line["area_code"] == "待定位"


def test_customer_scoped_account_is_rejected_even_with_role_permission(
    onboarding_api_app,
) -> None:
    app, _ids, _factory, _private_root = onboarding_api_app
    with TestClient(app) as client:
        _login(client, "n081-api-scoped")
        template = client.get(
            "/api/warehouse/inventory-onboarding/template?format=csv"
        )
        assert template.status_code == 403
        assert "可访问全部客户" in template.text

        imported = client.post(
            "/api/warehouse/inventory-onboarding/batches/import",
            files={"file": ("N081盘点.csv", _inventory_csv(), "text/csv")},
        )
        assert imported.status_code == 403
        assert "可访问全部客户" in imported.text


def test_import_list_detail_and_private_reference(onboarding_api_app) -> None:
    app, ids, factory, private_root = onboarding_api_app
    with TestClient(app) as client:
        _login(client, "n081-api-workshop")
        payload = _import_batch(client)
        batch = payload["batch"]
        line = payload["items"][0]

        assert batch["status"] == "draft"
        assert batch["source_file_reference"].startswith(
            "private:inventory_onboarding/"
        )
        assert not batch["source_file_reference"].startswith("/static/")
        stored_path = resolve_stored_reference(batch["source_file_reference"])
        assert stored_path.is_file()
        assert private_root in stored_path.parents
        assert line["match_status"] == "ready"
        assert line["action_decision"] == "create_new"
        assert line["customer_id"] == ids["customer"]
        assert line["product_id"] == ids["product"]
        assert line["location_id"] == ids["location"]

        listed = client.get(
            "/api/warehouse/inventory-onboarding/batches"
        )
        assert listed.status_code == 200
        assert [item["id"] for item in listed.json()["items"]] == [batch["id"]]

        detail = client.get(
            f"/api/warehouse/inventory-onboarding/batches/{batch['id']}"
        )
        assert detail.status_code == 200
        assert detail.json()["batch"]["id"] == batch["id"]
        assert len(detail.json()["items"]) == 1

        lines = client.get(
            f"/api/warehouse/inventory-onboarding/batches/{batch['id']}/lines"
        )
        assert lines.status_code == 200
        assert lines.json()["items"][0]["id"] == line["id"]

        with factory() as db:
            persisted = db.get(InventoryOnboardingBatch, batch["id"])
            assert persisted is not None
            assert persisted.source_file_reference == batch["source_file_reference"]


def test_dry_run_confirmed_submit_freezes_and_replays_idempotently(
    onboarding_api_app,
) -> None:
    app, _ids, factory, _private_root = onboarding_api_app
    before = _formal_inventory_counts(factory)
    assert before == (0, 0, 0, 0, 0)

    with TestClient(app) as client:
        _login(client, "n081-api-admin")
        imported = _import_batch(client)
        batch = imported["batch"]

        dry_run = client.post(
            (
                "/api/warehouse/inventory-onboarding/batches/"
                f"{batch['id']}/dry-run"
            ),
            json={"expected_version": batch["version"]},
        )
        assert dry_run.status_code == 200, dry_run.text
        dry_batch = dry_run.json()["batch"]
        fingerprint = dry_batch["dry_run_fingerprint"]
        assert len(fingerprint) == 64
        assert dry_batch["dry_run_summary"]["error_count"] == 0
        assert _formal_inventory_counts(factory) == before

        unconfirmed = client.post(
            (
                "/api/warehouse/inventory-onboarding/batches/"
                f"{batch['id']}/submit"
            ),
            json={
                "expected_version": dry_batch["version"],
                "dry_run_fingerprint": fingerprint,
                "idempotency_key": "n081-b1-api-submit-1",
                "confirmed": False,
            },
        )
        assert unconfirmed.status_code == 422
        assert (
            unconfirmed.json()["detail"]["code"]
            == "INVENTORY_ONBOARDING_CONFIRMATION_REQUIRED"
        )

        submitted = client.post(
            (
                "/api/warehouse/inventory-onboarding/batches/"
                f"{batch['id']}/submit"
            ),
            json={
                "expected_version": dry_batch["version"],
                "dry_run_fingerprint": fingerprint,
                "idempotency_key": "n081-b1-api-submit-1",
                "confirmed": True,
            },
        )
        assert submitted.status_code == 200, submitted.text
        submitted_batch = submitted.json()["batch"]
        assert submitted_batch["status"] == "submitted"
        assert _formal_inventory_counts(factory) == before

        replay = client.post(
            (
                "/api/warehouse/inventory-onboarding/batches/"
                f"{batch['id']}/submit"
            ),
            json={
                "expected_version": 1,
                "dry_run_fingerprint": fingerprint,
                "idempotency_key": "n081-b1-api-submit-1",
                "confirmed": True,
            },
        )
        assert replay.status_code == 200, replay.text
        assert replay.json()["batch"]["status"] == "submitted"
        assert replay.json()["batch"]["id"] == batch["id"]
        assert _formal_inventory_counts(factory) == before

        frozen_line = submitted.json()["items"][0]
        frozen_patch = client.patch(
            (
                "/api/warehouse/inventory-onboarding/batches/"
                f"{batch['id']}/lines/{frozen_line['id']}"
            ),
            json={
                "expected_version": frozen_line["version"],
                "remarks": "提交后不得修改",
            },
        )
        assert frozen_patch.status_code == 409
        assert (
            frozen_patch.json()["detail"]["code"]
            == "INVENTORY_ONBOARDING_FROZEN"
        )
        assert _formal_inventory_counts(factory) == before


def test_b1_api_exposes_no_formal_apply_endpoint(onboarding_api_app) -> None:
    app, _ids, factory, _private_root = onboarding_api_app
    route_paths = {route.path for route in app.routes}
    assert not any(
        path.endswith("/apply")
        for path in route_paths
        if "inventory-onboarding" in path
    )

    with TestClient(app) as client:
        _login(client, "n081-api-admin")
        response = client.post(
            "/api/warehouse/inventory-onboarding/batches/1/apply",
            json={},
        )
        assert response.status_code == 404
    assert _formal_inventory_counts(factory) == (0, 0, 0, 0, 0)


def test_database_trigger_concurrency_conflict_returns_409(
    onboarding_api_app,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    app, _ids, _factory, _private_root = onboarding_api_app

    def raise_trigger_conflict(*_args, **_kwargs):
        raise IntegrityError(
            "version trigger rejected stale update",
            {},
            RuntimeError("stale"),
        )

    monkeypatch.setattr(
        onboarding_api.onboarding_service,
        "update_onboarding_line",
        raise_trigger_conflict,
    )
    with TestClient(app) as client:
        _login(client, "n081-api-admin")
        response = client.patch(
            (
                "/api/warehouse/inventory-onboarding/"
                "batches/1/lines/1"
            ),
            json={"expected_version": 1, "remarks": "stale"},
        )

    assert response.status_code == 409
    assert (
        response.json()["detail"]["code"]
        == "INVENTORY_ONBOARDING_CONCURRENCY_CONFLICT"
    )
