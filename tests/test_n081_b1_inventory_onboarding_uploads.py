from __future__ import annotations

import asyncio
from io import BytesIO
from pathlib import Path
import zipfile

from openpyxl import Workbook
import pytest
from starlette.datastructures import Headers, UploadFile

from app.services.inventory_onboarding_uploads import (
    StoredInventoryOnboardingUpload,
    cleanup_inventory_onboarding_upload,
    store_and_parse_inventory_onboarding_upload,
)
from app.services.secure_uploads import UploadValidationError


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(filename: str, content: bytes, content_type: str) -> UploadFile:
    return UploadFile(
        filename=filename,
        file=BytesIO(content),
        headers=Headers({"content-type": content_type}),
    )


def _store(filename: str, content: bytes, content_type: str):
    return asyncio.run(
        store_and_parse_inventory_onboarding_upload(
            _upload(filename, content, content_type)
        )
    )


def _xlsx_bytes(
    rows: list[list[object]],
    *,
    title: str = "盘点",
    second_sheet_rows: list[list[object]] | None = None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    if second_sheet_rows is not None:
        second = workbook.create_sheet("第二表")
        for row in second_sheet_rows:
            second.append(row)
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()
    return payload.getvalue()


def _with_zip_member(content: bytes, name: str, payload: bytes) -> bytes:
    source = BytesIO(content)
    output = BytesIO()
    with zipfile.ZipFile(source, "r") as current, zipfile.ZipFile(
        output,
        "w",
        compression=zipfile.ZIP_DEFLATED,
    ) as rewritten:
        for info in current.infolist():
            rewritten.writestr(info, current.read(info.filename))
        rewritten.writestr(name, payload)
    return output.getvalue()


@pytest.fixture()
def private_root(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    root = tmp_path / "private"
    monkeypatch.setenv("ERP_FILE_STORAGE_DIR", str(root))
    return root


def test_utf8_bom_csv_preserves_rows_and_private_metadata(
    private_root: Path,
) -> None:
    content = (
        "\ufeff存货编码,数量,备注\r\n"
        "A-001,10,第一板\r\n"
        "\r\n"
        'A-002,20,"含,逗号"\r\n'
    ).encode("utf-8")

    result = _store("盘点.csv", content, "text/csv; charset=utf-8")

    assert isinstance(result, StoredInventoryOnboardingUpload)
    assert result.format == "csv"
    assert result.encoding == "utf-8-sig"
    assert result.filename == "盘点.csv"
    assert result.content_type == "text/csv"
    assert result.size == len(content)
    assert len(result.sha256) == 64
    assert result.private_reference.startswith("private:inventory_onboarding/")
    assert result.reference == result.private_reference
    assert result.path == result.private_path
    assert result.private_path.is_file()
    assert result.private_path.resolve().is_relative_to(private_root.resolve())
    assert result.private_path.name != result.filename
    assert [row.row_number for row in result.rows] == [1, 2, 4]
    assert {row.sheet_name for row in result.rows} == {"CSV"}
    assert result.rows[-1].raw_text == 'A-002,20,"含,逗号"'
    assert result.rows[-1].original_values == ("A-002", "20", "含,逗号")
    metadata = result.private_path.with_name(
        f"{result.private_path.name}.metadata.json"
    )
    assert metadata.is_file()

    assert result.cleanup() == []
    assert not result.private_path.exists()
    assert not metadata.exists()
    assert result.cleanup() == []


def test_gb18030_csv_is_decoded_without_rewriting_source(
    private_root: Path,
) -> None:
    content = "客户,数量\r\n天华,300\r\n".encode("gb18030")
    result = _store("factory.csv", content, "application/vnd.ms-excel")

    assert result.encoding == "gb18030"
    assert result.rows[1].original_values == ("天华", "300")
    assert result.private_path.read_bytes() == content
    assert result.sha256 == __import__("hashlib").sha256(content).hexdigest()


@pytest.mark.parametrize(
    ("filename", "content", "content_type", "message"),
    [
        ("../盘点.csv", b"a,b\n1,2\n", "text/csv", "路径"),
        ("fake.xlsx", b"a,b\n1,2\n", XLSX_MIME, "签名"),
        ("fake.csv", b"PK\x03\x04binary", "text/csv", "签名"),
        (
            "active.csv",
            b"<html>,value\n<script>,bad\n",
            "text/csv",
            "签名",
        ),
        ("盘点.txt", b"a,b\n1,2\n", "text/plain", "仅支持"),
        ("盘点.csv", b"a,b\n1,2\n", "application/pdf", "MIME"),
        ("one.csv", b"only-one-column\nvalue\n", "text/csv", "至少需要两列"),
    ],
)
def test_filename_mime_extension_and_content_must_match(
    private_root: Path,
    filename: str,
    content: bytes,
    content_type: str,
    message: str,
) -> None:
    with pytest.raises(UploadValidationError, match=message):
        _store(filename, content, content_type)
    assert not list(private_root.rglob("*")) if private_root.exists() else True


@pytest.mark.parametrize(
    "content",
    [
        b"a,b\n1,\x00bad\n",
        "a,b\n1,\tbad\n".encode(),
        "a,b\n1,=SUM(A1:A2)\n".encode(),
        "a,b\n1,+cmd\n".encode(),
        "a,b\n1,-cmd\n".encode(),
        "a,b\n1,@cmd\n".encode(),
    ],
)
def test_csv_rejects_controls_and_formula_injection(
    private_root: Path,
    content: bytes,
) -> None:
    with pytest.raises(UploadValidationError, match="控制字符|公式"):
        _store("unsafe.csv", content, "text/csv")
    assert not list(private_root.rglob("*")) if private_root.exists() else True


def test_csv_rejects_upload_over_p0b_limit(private_root: Path) -> None:
    content = b"a,b\n" + (b"x,y\n" * (5 * 1024 * 1024 + 1))
    with pytest.raises(UploadValidationError, match="不能超过 20MB"):
        _store("too-large.csv", content, "text/csv")
    assert not list(private_root.rglob("*")) if private_root.exists() else True


def test_xlsx_preserves_sheet_name_row_number_values_and_raw_text(
    private_root: Path,
) -> None:
    content = _xlsx_bytes(
        [
            ["存货编码", "数量", "盘点日期"],
            [None, None, None],
            ["FG-001", 300, "2026-07-23"],
        ],
        second_sheet_rows=[],
    )
    result = _store("盘点.xlsx", content, XLSX_MIME)

    assert result.format == "xlsx"
    assert result.encoding is None
    assert result.content_type == XLSX_MIME
    assert [sheet.name for sheet in result.sheets] == ["盘点", "第二表"]
    assert [row.row_number for row in result.sheets[0].rows] == [1, 3]
    assert result.sheets[1].rows == ()
    assert result.rows[-1].original_values == ("FG-001", 300, "2026-07-23")
    assert result.rows[-1].raw_text == '["FG-001",300,"2026-07-23"]'


def test_xlsx_rejects_more_than_one_nonempty_sheet(
    private_root: Path,
) -> None:
    content = _xlsx_bytes(
        [["编码", "数量"], ["FG-1", 1]],
        second_sheet_rows=[["编码", "数量"], ["FG-2", 2]],
    )
    with pytest.raises(UploadValidationError, match="一个非空工作表"):
        _store("multi.xlsx", content, XLSX_MIME)


def test_xlsx_rejects_formula_cells_and_formula_like_text(
    private_root: Path,
) -> None:
    formula = _xlsx_bytes([["编码", "数量"], ["FG-1", "=SUM(1,2)"]])
    with pytest.raises(UploadValidationError, match="公式"):
        _store("formula.xlsx", formula, XLSX_MIME)

    formula_text = _xlsx_bytes([["编码", "备注"], ["FG-1", "'safe"]])
    # A leading apostrophe is inert and retained as text.
    accepted = _store("inert.xlsx", formula_text, XLSX_MIME)
    assert accepted.rows[1].original_values == ("FG-1", "'safe")

    injected_text = _xlsx_bytes([["编码", "备注"], ["FG-1", "@cmd"]])
    with pytest.raises(UploadValidationError, match="公式"):
        _store("formula-text.xlsx", injected_text, XLSX_MIME)


def test_xlsx_rejects_external_links(private_root: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["编码", "说明"])
    sheet.append(["FG-1", "链接"])
    sheet["B2"].hyperlink = "https://example.invalid/inventory"
    payload = BytesIO()
    workbook.save(payload)
    workbook.close()

    with pytest.raises(UploadValidationError, match="外部链接"):
        _store("external.xlsx", payload.getvalue(), XLSX_MIME)


@pytest.mark.parametrize(
    ("member_name", "payload", "message"),
    [
        ("../evil.xml", b"<evil />", "路径穿越"),
        ("xl/vbaProject.bin", b"macro", "宏"),
        ("xl/drawings/drawing1.xml", b"<drawing />", "对象"),
        ("xl/embeddings/object1.bin", b"object", "对象"),
    ],
)
def test_xlsx_rejects_zip_traversal_macros_and_objects(
    private_root: Path,
    member_name: str,
    payload: bytes,
    message: str,
) -> None:
    content = _with_zip_member(
        _xlsx_bytes([["编码", "数量"], ["FG-1", 1]]),
        member_name,
        payload,
    )
    with pytest.raises(UploadValidationError, match=message):
        _store("unsafe.xlsx", content, XLSX_MIME)


def test_xlsx_rejects_zip_bomb_ratio(private_root: Path) -> None:
    content = _with_zip_member(
        _xlsx_bytes([["编码", "数量"], ["FG-1", 1]]),
        "xl/bomb.xml",
        b"A" * (2 * 1024 * 1024),
    )
    with pytest.raises(UploadValidationError, match="压缩比"):
        _store("bomb.xlsx", content, XLSX_MIME)


def test_cleanup_refuses_other_private_categories(private_root: Path) -> None:
    assert cleanup_inventory_onboarding_upload("private:drawings/not-ours.pdf") == [
        "不是 N081 盘点私有文件引用，拒绝清理"
    ]
    assert not private_root.exists()
