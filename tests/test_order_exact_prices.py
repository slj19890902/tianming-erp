from decimal import Decimal

import pytest
from sqlalchemy import select, func

from tests.test_p1_130_statement_invoice_finance import p1_130_app
from tests.test_phase5_orders import order_api_app


def test_order_price_roundtrip_and_total_do_not_round(p1_130_app):
    from app.models.order import Order, OrderItem
    from app.api.orders import _refresh_total
    _, factory = p1_130_app
    with factory() as db:
        item = db.scalar(select(OrderItem).order_by(OrderItem.id))
        order = db.get(Order, item.order_id)
        item.unit_price = Decimal('1.234567')
        item.quantity = 7
        item.subtotal = item.unit_price * item.quantity
        db.flush()
        _refresh_total(db, order)
        db.commit()
        db.expire_all()
        assert item.unit_price == Decimal('1.234567')
        assert item.subtotal == Decimal('8.641969')
        assert order.total_amount == Decimal('8.641969')
        assert db.scalar(select(OrderItem.unit_price).where(OrderItem.id == item.id)) == item.unit_price


@pytest.mark.parametrize('value, expected', [('222.525','222.525'),('1.2345','1.2345'),
    ('222.500','222.5'),('100.000000','100'),('-0.0004','-0.0004'),('0','0')])
def test_display_keeps_all_significant_decimal_places(value, expected):
    from app.services.reconciliation_number_display import format_reconciliation_number
    assert format_reconciliation_number(Decimal(value)) == expected


def test_exact_storage_and_sql_sum_preserve_high_value_tail(p1_130_app):
    from app.models.finance import Statement, StatementItem
    _, factory = p1_130_app
    with factory() as db:
        rows = db.scalars(select(StatementItem).order_by(StatementItem.id)).all()
        rows[0].receivable_amount = Decimal('123456789.1234567891')
        rows[1].receivable_amount = Decimal('0.0000000001')
        db.commit()
        db.expire_all()
        assert rows[0].receivable_amount == Decimal('123456789.1234567891')
        assert db.scalar(select(func.sum(StatementItem.receivable_amount))) == Decimal('123456789.1234567892')


def test_exact_decimal_bounds_reject_instead_of_rounding():
    from app.core.exact_money import decimal_value, money_text, price_value
    assert money_text(Decimal('6.67500')) == '6.675'
    for value in ['NaN', 'Infinity', '1e80', '1e-50']:
        with pytest.raises(ValueError):
            decimal_value(value)
    with pytest.raises(ValueError):
        price_value('1.1234567')


def test_real_order_save_reload_edit_and_replay_keep_original_price(order_api_app):
    from fastapi.testclient import TestClient
    from tests.test_phase5_orders import _login
    from tests.test_order_create_readback_contract import api_payload
    app, factory = order_api_app
    from app.core.config import settings
    app.state.erp_settings = settings
    payload = api_payload(factory)
    payload['items'][0].update(quantity=7, unit_price='1.234567')
    with TestClient(app) as client:
        _login(client)
        saved = client.post('/api/orders', json=payload)
        assert saved.status_code == 201, saved.text
        order = saved.json()
        assert order['total_amount'] == '8.641969'
        assert order['items'][0]['unit_price'] == '1.234567'
        assert order['items'][0]['subtotal'] == '8.641969'
        assert client.get('/api/orders/' + str(order['id'])).json()['total_amount'] == '8.641969'
        replay = client.post('/api/orders', json=payload)
        assert replay.status_code in (200, 201), replay.text
        assert replay.json()['id'] == order['id']
        item = order['items'][0]
        edit = dict(quantity=3, unit_price='2.225', product_code=item['snapshot_product_code'],
                    product_name=item['snapshot_product_name'], material=item['snapshot_material'],
                    specification=item['snapshot_spec'], quantity_adjustment_idempotency_key='exact-money-edit-once')
        updated = client.put('/api/orders/items/' + str(item['id']), json=edit)
        assert updated.status_code == 200, updated.text
        assert updated.json()['subtotal'] == '6.675'
        assert client.get('/api/orders/' + str(order['id'])).json()['total_amount'] == '6.675'
        edit['unit_price'] = '1.1234567'
        assert client.put('/api/orders/items/' + str(item['id']), json=edit).status_code == 422


def test_exact_sql_outstanding_comparisons_and_legacy_read(p1_130_app):
    from sqlalchemy import text
    from app.models.finance import Statement
    from app.models.order import OrderItem
    _, factory = p1_130_app
    with factory() as db:
        statement = db.get(Statement, 1)
        statement.total_receivable = Decimal('123456789.1234567891')
        statement.settled_amount = Decimal('123456789.12')
        db.commit()
        assert db.scalar(select(Statement.total_receivable - Statement.settled_amount).where(Statement.id == 1)) == Decimal('.0034567891')
        assert db.scalar(select(Statement.id).where(Statement.total_receivable > Decimal('123456789.1234567890'))) == 1
        # An untouched pre-upgrade row keeps the original amount; no inference.
        db.execute(text('UPDATE sales_order_items SET subtotal=6.68, subtotal_exact=NULL WHERE id=1'))
        db.commit(); db.expire_all()
        assert db.get(OrderItem, 1).subtotal == Decimal('6.68')


def test_reader_refuses_older_program_after_exact_fact(p1_130_app):
    from desktop_assistant.exact_sales_reader import require_reader
    _, factory = p1_130_app
    with factory() as db:
        connection = db.connection().connection.driver_connection
        with pytest.raises(ValueError, match='精确销售金额'):
            require_reader(connection, {})
        require_reader(connection, {'reader_capabilities': {'exact_sales_decimal_v1': 1}})
        with pytest.raises(ValueError):
            require_reader(connection, {'reader_capabilities': {'exact_sales_decimal_v1': True}})


def test_original_quote_tax_reference_never_drives_amount():
    from types import SimpleNamespace as NS
    from app.services.sales_money import price_terms, order_totals
    inclusive = price_terms(Decimal('3.29'), 'tax_inclusive', Decimal('.13'))
    assert inclusive['tax_reference']['approximate'] is True
    assert inclusive['tax_reference']['value'] == '2.911504'
    exclusive = price_terms(Decimal('3.29'), 'tax_exclusive', Decimal('.13'))
    assert exclusive['tax_reference']['value'] == '3.7177'
    def line(mode):
        return NS(subtotal=Decimal('658'), price_tax_mode_snapshot=mode, tax_rate_snapshot=Decimal('.13'))
    assert order_totals([line('tax_inclusive')])['receivable_amount'] == '658'
    result = order_totals([line('tax_exclusive')])
    assert result['original_amount'] == '658' and result['receivable_amount'] == '743.54'
    mixed = order_totals([line('tax_inclusive'), line('tax_exclusive')])
    assert mixed['original_amount'] is None and len(mixed['groups']) == 2
    assert mixed['receivable_amount'] == '1401.54'


def test_excel_long_decimal_fact_is_explicit_text_not_a_rounded_number(p1_130_app):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from io import BytesIO
    from app.models.finance import StatementItem, Statement
    from tests.test_p1_130_statement_invoice_finance import _login
    app, factory = p1_130_app
    with factory() as db:
        line = db.scalar(select(StatementItem).where(StatementItem.statement_id == 1).order_by(StatementItem.id))
        line.receivable_amount = Decimal('123456789.1234567891')
        db.get(Statement, 1).total_receivable = Decimal('123456889.1234567891')
        db.commit()
    with TestClient(app) as client:
        _login(client)
        response = client.get('/api/finance/statements/1/customer-export.xlsx')
        assert response.status_code == 200, response.text
        body = client.get('/api/finance/statements/1').json()
        assert body['total_receivable'] == '123456889.1234567891'
    cell = load_workbook(BytesIO(response.content))['对账单']['H5']
    assert cell.value == '123456789.1234567891' and cell.data_type == 's'
    assert cell.number_format == '@' and '15位' in cell.comment.text


def test_mixed_tax_export_separates_quote_subtotals_and_gross_receivable(p1_130_app):
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from io import BytesIO
    import fitz
    from app.models.finance import StatementItem, Statement
    from tests.test_p1_130_statement_invoice_finance import _login
    app, factory = p1_130_app
    with factory() as db:
        rows = db.scalars(select(StatementItem).where(StatementItem.statement_id == 1).order_by(StatementItem.id)).all()
        assert len(rows) == 2
        for row, mode, amount in zip(rows, ['tax_inclusive','tax_exclusive'], ['658','743.54']):
            row.actual_received_quantity = 200
            row.unit_price_snapshot = Decimal('3.29')
            row.price_tax_mode_snapshot = mode
            row.tax_rate_snapshot = Decimal('.13')
            row.receivable_amount = Decimal(amount)
        db.get(Statement, 1).total_receivable = Decimal('1401.54')
        db.commit()
    with TestClient(app) as client:
        _login(client)
        response = client.get('/api/finance/statements/1/customer-export.xlsx')
        assert response.status_code == 200, response.text
        sheet = load_workbook(BytesIO(response.content))['对账单']
        assert sheet['H5'].value == sheet['H6'].value == 658
        assert [sheet['I5'].value, sheet['I6'].value] == ['含税','未税']
        assert sheet['E9'].value == '含税应收合计' and sheet['H9'].value == 1401.54
        response = client.get('/api/finance/statements/1/customer-export.pdf')
        assert response.status_code == 200, response.text
    with fitz.open(stream=response.content, filetype='pdf') as pdf:
        text = ''.join(page.get_text() for page in pdf)
    assert '1401.54' in text and '含税应收合计' in text and '未税小计' in text
    import os
    if os.environ.get('ORDER_EXACT_PDF_SAMPLE'):
        from pathlib import Path
        Path(os.environ['ORDER_EXACT_PDF_SAMPLE']).write_bytes(response.content)
