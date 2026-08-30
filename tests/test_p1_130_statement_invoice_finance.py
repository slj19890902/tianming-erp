from __future__ import annotations

from collections.abc import Generator
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, sessionmaker


@pytest.fixture()
def p1_130_app(tmp_path: Path):
    from app.api.auth import router as auth_router
    from app.api.deps import get_db
    from app.api.finance import router as finance_router
    from app.api.invoice_tasks import router as invoice_router
    from app.core.database import create_sqlite_engine
    from app.core.security import hash_password
    from app.models import Base
    from app.models.customer import Customer
    from app.models.delivery import Delivery, DeliveryItem
    from app.models.finance import ReturnReceipt, ReturnReceiptItem, Statement, StatementItem
    from app.models.order import Order, OrderItem
    from app.models.product import Product
    from app.models.user import User

    engine = create_sqlite_engine(tmp_path / "p1-130.sqlite3")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        admin = User(
            username="p1130-admin",
            password_hash=hash_password("RolePass123!"),
            role="admin",
            real_name="Admin",
            must_change_password=False,
        )
        customer = Customer(
            name="苏州思迈尔包装有限公司",
            chinese_short_name="思迈尔",
            customer_code="P1130",
            statement_cycle_start_day=20,
            is_active=True,
        )
        db.add_all([admin, customer])
        db.flush()
        product = Product(
            customer_id=customer.id,
            product_code="BOX-001",
            customer_material_code="BOX-001",
            product_name="三层瓦楞外箱",
            box_category="normal",
            cost_unit_price=Decimal("1.00"),
        )
        db.add(product)
        db.flush()
        receipt_items = []
        for index, (order_no, po, delivery_day, quantity, price, month_value) in enumerate(
            [
                ("SO-002", "PO-B", date(2026, 8, 3), 10, "10.00", "2026-08"),
                ("SO-001", "PO-A", date(2026, 8, 4), 5, "20.00", "2026-08"),
                ("SO-000", "PO-OLD", date(2026, 7, 4), 3, "30.00", "2026-07"),
            ],
            start=1,
        ):
            order = Order(
                order_number=order_no,
                customer_id=customer.id,
                customer_po=po,
                order_date=date(2026, 7, 1),
                delivery_date=delivery_day,
                status="pending_reconciliation",
                payment_status="unpaid",
                total_amount=Decimal(price) * quantity,
            )
            db.add(order)
            db.flush()
            order_item = OrderItem(
                order_id=order.id,
                product_id=product.id,
                quantity=quantity,
                delivered_quantity=quantity,
                unit_price=Decimal(price),
                subtotal=Decimal(price) * quantity,
                material_status="received",
                snapshot_product_code=f"BOX-00{index}",
                snapshot_product_name=f"纸箱 {index}",
                price_tax_mode_snapshot="tax_inclusive",
                tax_rate_snapshot=Decimal("0.13"),
            )
            db.add(order_item)
            db.flush()
            delivery = Delivery(
                delivery_number=f"DN-00{index}",
                customer_id=customer.id,
                delivery_date=delivery_day,
                status="dispatched",
                total_quantity=quantity,
            )
            db.add(delivery)
            db.flush()
            delivery_item = DeliveryItem(
                delivery_id=delivery.id,
                order_item_id=order_item.id,
                delivered_quantity=quantity,
            )
            db.add(delivery_item)
            db.flush()
            receipt = ReturnReceipt(
                delivery_id=delivery.id,
                actual_received_date=delivery_day,
                reconciliation_month=month_value,
                status="confirmed",
            )
            db.add(receipt)
            db.flush()
            receipt_item = ReturnReceiptItem(
                return_receipt_id=receipt.id,
                delivery_item_id=delivery_item.id,
                actual_received_quantity=quantity,
            )
            db.add(receipt_item)
            db.flush()
            receipt_items.append((receipt_item, Decimal(price), quantity))
        statement = Statement(
            statement_number="ST-P1-130",
            customer_id=customer.id,
            statement_month="2026-08",
            total_receivable=Decimal("200.00"),
            total_gross_profit=Decimal("185.00"),
            status="unsettled",
        )
        db.add(statement)
        db.flush()
        for receipt_item, price, quantity in receipt_items[:2]:
            db.add(
                StatementItem(
                    statement_id=statement.id,
                    return_receipt_item_id=receipt_item.id,
                    actual_received_quantity=quantity,
                    unit_price_snapshot=price,
                    unit_cost_snapshot=Decimal("1.00"),
                    receivable_amount=price * quantity,
                    gross_profit_amount=(price - Decimal("1.00")) * quantity,
                    price_tax_mode_snapshot="tax_inclusive",
                    tax_rate_snapshot=Decimal("0.13"),
                )
            )
        db.commit()

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(finance_router, prefix="/api/finance")
    app.include_router(invoice_router, prefix="/api/finance")

    def override_get_db() -> Generator[Session, None, None]:
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    yield app, factory
    engine.dispose()


def _login(client: TestClient) -> None:
    response = client.post(
        "/api/auth/login",
        json={"username": "p1130-admin", "password": "RolePass123!"},
    )
    assert response.status_code == 200, response.text


def test_statement_dispute_reopens_moves_and_adds_without_rewriting_receipt_facts(
    p1_130_app,
) -> None:
    from app.models.finance import ReturnReceiptItem, Statement, StatementAdjustment

    app, factory = p1_130_app
    with TestClient(app) as client:
        _login(client)
        confirmed = client.post(
            "/api/finance/statements/1/confirm", json={"expected_version": 1}
        )
        assert confirmed.status_code == 200, confirmed.text
        candidates = client.get("/api/finance/statements/1/adjustment-candidates")
        assert candidates.status_code == 200, candidates.text
        assert [row["return_receipt_item_id"] for row in candidates.json()["items"]] == [3]

        reopened = client.post(
            "/api/finance/statements/1/reopen",
            json={"expected_version": 2, "reason": "客户要求一条延后并补入遗漏"},
        )
        assert reopened.status_code == 200, reopened.text
        assert reopened.json()["confirmation_status"] == "draft"
        first_statement_item_id = next(
            item["statement_item_id"]
            for item in reopened.json()["items"]
            if item["return_receipt_item_id"] == 1
        )
        adjusted = client.post(
            "/api/finance/statements/1/adjust-dispute",
            json={
                "expected_version": 3,
                "reason": "客户要求一条延后并补入遗漏",
                "remove_lines": [
                    {
                        "statement_item_id": first_statement_item_id,
                        "target_month": "2026-09",
                    }
                ],
                "add_return_receipt_item_ids": [3],
            },
        )
        assert adjusted.status_code == 200, adjusted.text
        assert adjusted.json()["version"] == 4
        assert Decimal(str(adjusted.json()["total_receivable"])) == Decimal("190.00")
        assert {
            item["return_receipt_item_id"] for item in adjusted.json()["items"]
        } == {2, 3}
        with factory() as db:
            first = db.get(ReturnReceiptItem, 1)
            added = db.get(ReturnReceiptItem, 3)
            statement = db.get(Statement, 1)
            assert first.actual_received_quantity == 10
            assert first.reconciliation_month_override == "2026-09"
            assert added.actual_received_quantity == 3
            assert added.reconciliation_month_override == "2026-08"
            assert statement.confirmation_status == "draft"
            assert db.scalar(select(func.count(StatementAdjustment.id))) == 2


def test_statement_dispute_can_reopen_and_adjust_atomically(p1_130_app) -> None:
    from app.models.finance import ReturnReceiptItem, Statement, StatementAdjustment

    app, factory = p1_130_app
    with TestClient(app) as client:
        _login(client)
        confirmed = client.post(
            "/api/finance/statements/1/confirm", json={"expected_version": 1}
        )
        assert confirmed.status_code == 200, confirmed.text
        detail = client.get("/api/finance/statements/1").json()
        first_item_id = next(
            item["statement_item_id"]
            for item in detail["items"]
            if item["return_receipt_item_id"] == 1
        )
        adjusted = client.post(
            "/api/finance/statements/1/adjust-dispute",
            json={
                "expected_version": 2,
                "reason": "客户要求本月暂不核对该送货明细",
                "remove_lines": [
                    {"statement_item_id": first_item_id, "target_month": "2026-09"}
                ],
                "add_return_receipt_item_ids": [],
            },
        )
        assert adjusted.status_code == 200, adjusted.text
        assert adjusted.json()["confirmation_status"] == "draft"
        assert adjusted.json()["version"] == 3
        with factory() as db:
            assert db.get(ReturnReceiptItem, 1).actual_received_quantity == 10
            assert db.get(ReturnReceiptItem, 1).reconciliation_month_override == "2026-09"
            assert db.get(Statement, 1).confirmation_status == "draft"
            assert db.scalar(select(func.count(StatementAdjustment.id))) == 1


def test_issued_statement_cannot_be_reopened_or_adjusted(p1_130_app) -> None:
    from app.models.finance import Invoice

    app, factory = p1_130_app
    with TestClient(app) as client:
        _login(client)
        assert client.post(
            "/api/finance/statements/1/confirm", json={"expected_version": 1}
        ).status_code == 200
        with factory() as db:
            db.add(
                Invoice(
                    statement_id=1,
                    invoice_number="FP-P1-130",
                    invoice_date=date(2026, 8, 31),
                    invoice_amount=Decimal("200.00"),
                    invoice_status="issued",
                    source="legacy_manual",
                )
            )
            db.commit()
        response = client.post(
            "/api/finance/statements/1/adjust-dispute",
            json={
                "expected_version": 2,
                "reason": "客户提出异议但已经开票",
                "remove_lines": [{"statement_item_id": 1, "target_month": "2026-09"}],
                "add_return_receipt_item_ids": [],
            },
        )
        assert response.status_code == 409
        assert "红冲" in response.text


def test_settlement_entity_groups_customers_and_freezes_scope(p1_130_app) -> None:
    from app.models.customer import Customer
    from app.models.delivery import Delivery, DeliveryItem
    from app.models.finance import ReturnReceipt, ReturnReceiptItem, Statement, StatementItem
    from app.models.invoice_task import (
        CustomerInvoiceProfile,
        FinanceSettlementEntity,
        InvoiceSellerEntity,
    )
    from app.models.order import Order, OrderItem
    from app.models.product import Product

    app, factory = p1_130_app
    with factory() as db:
        first_customer = db.get(Customer, 1)
        second_customer = Customer(
            name="昆山合作客户有限公司",
            chinese_short_name="合作客户",
            customer_code="P1130-B",
            statement_cycle_start_day=1,
            is_active=True,
        )
        seller = InvoiceSellerEntity(
            seller_code="TM-PACK",
            seller_name="苏州天明包装有限公司",
            tax_no="913205000000000001",
            address="苏州市",
            phone="0512-00000001",
            bank_name="中国银行",
            bank_account="100000000001",
            confirmation_status="confirmed",
        )
        db.add_all([second_customer, seller])
        db.flush()
        entity = FinanceSettlementEntity(
            entity_code="PARTNER-01",
            entity_name="合作纸箱厂有限公司",
            short_name="合作纸箱厂",
            tax_no="913205000000000002",
            invoice_address="昆山市",
            invoice_phone="0512-00000002",
            bank_name="工商银行",
            bank_account="200000000002",
            default_seller_id=seller.id,
            statement_cycle_start_day=20,
            confirmation_status="confirmed",
        )
        db.add(entity)
        db.flush()
        for customer in (first_customer, second_customer):
            db.add(
                CustomerInvoiceProfile(
                    customer_id=customer.id,
                    settlement_entity_id=entity.id,
                    default_seller_id=seller.id,
                    price_tax_mode="tax_inclusive",
                    default_tax_rate=Decimal("0.13"),
                    is_enabled=True,
                    confirmation_status="confirmed",
                )
            )
        product = Product(
            customer_id=second_customer.id,
            product_code="PARTNER-BOX",
            customer_material_code="PARTNER-BOX",
            product_name="合作客户纸箱",
            box_category="normal",
            cost_unit_price=Decimal("1.00"),
        )
        db.add(product)
        db.flush()
        order = Order(
            order_number="SO-PARTNER",
            customer_id=second_customer.id,
            customer_po="PO-PARTNER",
            order_date=date(2026, 8, 1),
            delivery_date=date(2026, 8, 15),
            status="pending_reconciliation",
            payment_status="unpaid",
            total_amount=Decimal("120.00"),
        )
        db.add(order)
        db.flush()
        order_item = OrderItem(
            order_id=order.id,
            product_id=product.id,
            quantity=12,
            delivered_quantity=12,
            unit_price=Decimal("10.00"),
            subtotal=Decimal("120.00"),
            material_status="received",
            snapshot_product_code="PARTNER-BOX",
            snapshot_product_name="合作客户纸箱",
            price_tax_mode_snapshot="tax_inclusive",
            tax_rate_snapshot=Decimal("0.13"),
        )
        db.add(order_item)
        db.flush()
        delivery = Delivery(
            delivery_number="DN-PARTNER",
            customer_id=second_customer.id,
            delivery_date=date(2026, 8, 15),
            status="dispatched",
            total_quantity=12,
        )
        db.add(delivery)
        db.flush()
        delivery_item = DeliveryItem(
            delivery_id=delivery.id,
            order_item_id=order_item.id,
            delivered_quantity=12,
        )
        db.add(delivery_item)
        db.flush()
        receipt = ReturnReceipt(
            delivery_id=delivery.id,
            actual_received_date=date(2026, 8, 15),
            reconciliation_month="2026-08",
            status="confirmed",
        )
        db.add(receipt)
        db.flush()
        receipt_item = ReturnReceiptItem(
            return_receipt_id=receipt.id,
            delivery_item_id=delivery_item.id,
            actual_received_quantity=12,
        )
        db.add(receipt_item)
        db.commit()
        delivery_id = delivery.id
        second_customer_id = second_customer.id
        entity_id = entity.id

    with TestClient(app) as client:
        _login(client)
        options = client.get(
            "/api/finance/statement-customers",
            params={"statement_month": "2026-08"},
        )
        assert options.status_code == 200, options.text
        grouped = next(
            row for row in options.json()["items"] if row["settlement_entity_id"] == entity_id
        )
        assert grouped["name"] == "合作纸箱厂有限公司"
        assert grouped["customer_ids"] == [1, second_customer_id]
        pending = client.get(
            "/api/finance/pending_statements",
            params={"customer_id": 1, "statement_month": "2026-08"},
        )
        assert pending.status_code == 200, pending.text
        assert pending.json()["settlement_name"] == "合作纸箱厂有限公司"
        assert [row["delivery_id"] for row in pending.json()["deliveries"]] == [delivery_id]
        created = client.post(
            "/api/finance/statements",
            json={
                "customer_id": 1,
                "statement_month": "2026-08",
                "delivery_ids": [delivery_id],
                "return_receipt_item_ids": [],
                "customer_charge_ids": [],
                "idempotency_key": "p1130-settlement-entity",
            },
        )
        assert created.status_code == 201, created.text
        with factory() as db:
            statement = db.get(Statement, created.json()["id"])
            item = db.scalar(
                select(StatementItem).where(StatementItem.statement_id == statement.id)
            )
            assert statement.settlement_entity_id == entity_id
            assert statement.settlement_name_snapshot == "合作纸箱厂有限公司"
            assert statement.statement_cycle_start_day_snapshot == 20
            assert statement.settlement_customer_ids_snapshot_json == f"[1, {second_customer_id}]"
            assert item.source_customer_id == second_customer_id


def test_statement_customer_exports_are_compact_sorted_and_have_matching_totals(
    p1_130_app,
) -> None:
    app, _factory = p1_130_app
    with TestClient(app) as client:
        _login(client)
        excel = client.get(
            "/api/finance/statements/1/customer-export.xlsx",
            params={"sort_by": "order_number"},
        )
        assert excel.status_code == 200, excel.text
        assert "思迈尔" in unquote(excel.headers["content-disposition"])
        workbook = load_workbook(BytesIO(excel.content), data_only=True)
        sheet = workbook["对账单"]
        assert [sheet.cell(4, column).value for column in range(1, 9)] == [
            "送货日期",
            "送货单号",
            "客户单号",
            "存货编码",
            "产品名称",
            "数量",
            "含税单价",
            "含税金额",
        ]
        assert sheet["C5"].value == "PO-A"
        assert sheet["C6"].value == "PO-B"
        assert Decimal(str(sheet["F7"].value)) == Decimal("15")
        assert Decimal(str(sheet["H7"].value)) == Decimal("200")
        pdf = client.get("/api/finance/statements/1/customer-export.pdf")
        assert pdf.status_code == 200, pdf.text
        assert pdf.content.startswith(b"%PDF")


def test_payables_feed_aging_structure_and_six_month_trend(p1_130_app) -> None:
    app, _factory = p1_130_app
    with TestClient(app) as client:
        _login(client)
        created = client.post(
            "/api/finance/payables",
            json={
                "supplier_id": None,
                "counterparty_name": "苏州纸板供应商",
                "category": "material",
                "document_number": "AP-001",
                "document_date": "2026-08-01",
                "due_date": "2026-08-10",
                "amount": "1200.00",
                "note": "八月纸板款",
                "idempotency_key": "p1130-payable-001",
            },
        )
        assert created.status_code == 201, created.text
        confirmed = client.post(
            f"/api/finance/payables/{created.json()['id']}/confirm",
            json={"expected_version": 1},
        )
        assert confirmed.status_code == 200, confirmed.text
        overview = client.get(
            "/api/finance/overview", params={"through_month": "2026-08"}
        )
        assert overview.status_code == 200, overview.text
        august = overview.json()["trend"][-1]
        assert Decimal(str(august["payable_amount"])) == Decimal("1200.00")
        assert overview.json()["expense_structure"][0]["category"] == "material"
        assert Decimal(
            str(overview.json()["payable_aging"]["overdue_1_30"])
        ) == Decimal("1200.00")
        assert overview.json()["collection_note"].startswith("客户收款不在 ERP")


def test_p1_130_frontend_has_four_simple_finance_workbenches_and_no_receipt_action() -> None:
    index = Path("static/index.html").read_text(encoding="utf-8")
    finance_start = index.index('<template v-else-if="activePage === \'finance\'">')
    finance_end = index.index('<template v-else-if="activePage === \'system\'">')
    finance = index[finance_start:finance_end]
    for label in ("经营概览", "客户对账", "开票任务", "应付支出"):
        assert label in finance
    assert "收款核销" not in finance
    assert "customer-export.${suffix}" in index
    assert "openStatementDispute" in index
