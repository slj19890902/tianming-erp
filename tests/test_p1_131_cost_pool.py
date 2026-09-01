from __future__ import annotations

from collections.abc import Generator
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, sessionmaker


@pytest.fixture()
def p1_131_cost_app(tmp_path: Path):
    from app.api.auth import router as auth_router
    from app.api.cost_accounting import router as cost_router
    from app.api.deps import get_db
    from app.core.database import create_sqlite_engine
    from app.core.security import hash_password
    from app.models import Base
    from app.models.finance_cost import FinanceCostCenter
    from app.models.user import User

    engine = create_sqlite_engine(tmp_path / "p1-131-cost.sqlite3")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add_all(
            [
                User(
                    username="p1131-admin",
                    password_hash=hash_password("RolePass123!"),
                    role="admin",
                    real_name="Admin",
                    must_change_password=False,
                ),
                User(
                    username="p1131-restricted",
                    password_hash=hash_password("RolePass123!"),
                    role="finance",
                    real_name="Restricted Finance",
                    customer_access_mode="selected",
                    must_change_password=False,
                ),
                FinanceCostCenter(
                    code="PROD",
                    name="生产成本",
                    center_type="production",
                    is_active=True,
                ),
                FinanceCostCenter(
                    code="ADMIN",
                    name="管理费用",
                    center_type="administration",
                    is_active=True,
                ),
                FinanceCostCenter(
                    code="UNALLOCATED",
                    name="待分类",
                    center_type="unallocated",
                    is_active=True,
                ),
                FinanceCostCenter(
                    code="WH_DELIVERY",
                    name="仓储配送",
                    center_type="warehouse_delivery",
                    is_active=True,
                ),
                FinanceCostCenter(
                    code="FINANCE",
                    name="财务",
                    center_type="finance",
                    is_active=True,
                ),
            ]
        )
        db.commit()

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(cost_router, prefix="/api/finance")

    def override_get_db() -> Generator[Session, None, None]:
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    yield app, factory
    engine.dispose()


def _login(client: TestClient, username: str = "p1131-admin") -> None:
    response = client.post(
        "/api/auth/login",
        json={"username": username, "password": "RolePass123!"},
    )
    assert response.status_code == 200, response.text


def _entry_payload(*, key: str, amount: str = "0.30", description: str = "生产工资") -> dict:
    return {
        "cost_month": "2026-08",
        "document_date": "2026-08-31",
        "cost_center_id": 1,
        "cost_category": "production_wages",
        "allocation_basis": "unallocated",
        "description": description,
        "counterparty_name": "员工工资",
        "document_number": "PAYROLL-202608",
        "source_reference": "2026年8月工资表",
        "amount": amount,
        "tax_amount": "0.00",
        "note": "财务复核后确认",
        "idempotency_key": key,
    }


def test_wage_categories_keep_production_driver_finance_and_admin_boundaries(
    p1_131_cost_app,
) -> None:
    app, _factory = p1_131_cost_app
    with TestClient(app) as client:
        _login(client)
        metadata = client.get("/api/finance/cost-pool/metadata")
        assert metadata.status_code == 200, metadata.text
        categories = {
            row["value"]: row for row in metadata.json()["categories"]
        }
        assert categories["production_wages"] == {
            "value": "production_wages",
            "label": "生产工资/社保",
            "accounting_class": "manufacturing",
            "default_center": "PROD",
        }
        assert categories["driver_wages"]["accounting_class"] == "selling"
        assert categories["driver_wages"]["default_center"] == "WH_DELIVERY"
        assert categories["finance_wages"]["accounting_class"] == "finance"
        assert categories["finance_wages"]["default_center"] == "FINANCE"
        assert categories["administrative_wages"]["accounting_class"] == "administrative"
        assert "warehouse_wages" not in categories
        centers = {row["code"]: row["id"] for row in metadata.json()["centers"]}

        driver_payload = _entry_payload(
            key="p1131-driver-wages-001", description="司机工资"
        )
        driver_payload.update(
            cost_center_id=centers["WH_DELIVERY"],
            cost_category="driver_wages",
        )
        driver = client.post("/api/finance/cost-pool", json=driver_payload)
        assert driver.status_code == 201, driver.text
        assert driver.json()["accounting_class"] == "selling"

        finance_payload = _entry_payload(
            key="p1131-finance-wages-001", description="财务工资"
        )
        finance_payload.update(
            cost_center_id=centers["FINANCE"],
            cost_category="finance_wages",
        )
        finance = client.post("/api/finance/cost-pool", json=finance_payload)
        assert finance.status_code == 201, finance.text
        assert finance.json()["accounting_class"] == "finance"

        wrong_center = dict(driver_payload)
        wrong_center.update(
            idempotency_key="p1131-driver-wages-wrong-center",
            cost_center_id=centers["PROD"],
        )
        rejected = client.post("/api/finance/cost-pool", json=wrong_center)
        assert rejected.status_code == 409


def test_cost_pool_decimal_idempotency_version_and_missing_cost_blocker(
    p1_131_cost_app,
) -> None:
    from app.models.finance_cost import FinanceCostPoolEntry

    app, factory = p1_131_cost_app
    with TestClient(app) as client:
        _login(client)
        created = client.post(
            "/api/finance/cost-pool",
            json=_entry_payload(key="p1131-create-001"),
        )
        assert created.status_code == 201, created.text
        assert Decimal(str(created.json()["amount"])) == Decimal("0.30")
        assert created.json()["accounting_class"] == "manufacturing"
        assert created.json()["status"] == "draft"

        replay = client.post(
            "/api/finance/cost-pool",
            json=_entry_payload(key="p1131-create-001"),
        )
        assert replay.status_code == 201, replay.text
        assert replay.json()["id"] == created.json()["id"]

        conflict = client.post(
            "/api/finance/cost-pool",
            json=_entry_payload(key="p1131-create-001", amount="0.31"),
        )
        assert conflict.status_code == 409, conflict.text

        confirmed = client.post(
            f"/api/finance/cost-pool/{created.json()['id']}/confirm",
            json={
                "expected_version": 1,
                "idempotency_key": "p1131-confirm-001",
            },
        )
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["status"] == "confirmed"
        assert confirmed.json()["version"] == 2

        confirm_replay = client.post(
            f"/api/finance/cost-pool/{created.json()['id']}/confirm",
            json={
                "expected_version": 1,
                "idempotency_key": "p1131-confirm-001",
            },
        )
        assert confirm_replay.status_code == 200, confirm_replay.text
        assert confirm_replay.json()["version"] == 2

        forbidden_edit = client.put(
            f"/api/finance/cost-pool/{created.json()['id']}",
            json={
                **_entry_payload(key="p1131-update-001", amount="1.00"),
                "expected_version": 2,
            },
        )
        assert forbidden_edit.status_code == 409, forbidden_edit.text

        summary = client.get(
            "/api/finance/cost-pool/summary", params={"month": "2026-08"}
        )
        assert summary.status_code == 200, summary.text
        body = summary.json()
        assert Decimal(str(body["confirmed_total"])) == Decimal("0.30")
        assert body["close_ready"] is False
        assert {
            blocker["code"] for blocker in body["blockers"]
        } >= {
            "manufacturing_cost_unallocated",
            "month_close_workflow_pending",
        }
        assert body["material_cost"]["lineage_ready"] is True
        assert body["material_cost"]["total_delivery_lines"] == 0

    with factory() as db:
        assert db.scalar(select(func.count()).select_from(FinanceCostPoolEntry)) == 1


def test_cost_pool_company_scope_is_fail_closed(p1_131_cost_app) -> None:
    app, _factory = p1_131_cost_app
    with TestClient(app) as client:
        _login(client, "p1131-restricted")
        assert client.get(
            "/api/finance/cost-pool", params={"month": "2026-08"}
        ).status_code == 403
        assert client.get(
            "/api/finance/cost-pool/summary", params={"month": "2026-08"}
        ).status_code == 403
        assert client.get(
            "/api/finance/material-cost/coverage", params={"month": "2026-08"}
        ).status_code == 403
        assert client.get(
            "/api/finance/cost-pool/export", params={"month": "2026-08"}
        ).status_code == 403
        assert client.post(
            "/api/finance/cost-pool",
            json=_entry_payload(key="p1131-restricted-create"),
        ).status_code == 403


def test_cost_pool_audit_failure_rolls_back_business_and_idempotency(
    p1_131_cost_app,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.api import cost_accounting as cost_api
    from app.models.finance import FinanceIdempotencyRecord
    from app.models.finance_cost import FinanceCostPoolEntry

    app, factory = p1_131_cost_app

    def fail_audit(*_args, **_kwargs):
        raise RuntimeError("audit unavailable")

    monkeypatch.setattr(cost_api, "append_audit_event", fail_audit)
    with TestClient(app, raise_server_exceptions=False) as client:
        _login(client)
        response = client.post(
            "/api/finance/cost-pool",
            json=_entry_payload(key="p1131-audit-failure"),
        )
        assert response.status_code == 500

    with factory() as db:
        assert db.scalar(select(func.count()).select_from(FinanceCostPoolEntry)) == 0
        assert db.scalar(select(func.count()).select_from(FinanceIdempotencyRecord)) == 0


def test_cost_pool_amount_limit_and_required_trimmed_text(
    p1_131_cost_app,
) -> None:
    from app.models.finance_cost import FinanceCostPoolEntry

    app, factory = p1_131_cost_app
    with TestClient(app) as client:
        _login(client)
        maximum = client.post(
            "/api/finance/cost-pool",
            json=_entry_payload(
                key="p1131-amount-maximum",
                amount="999999999999.99",
            ),
        )
        assert maximum.status_code == 201, maximum.text
        assert Decimal(str(maximum.json()["amount"])) == Decimal(
            "999999999999.99"
        )

        over_limit = client.post(
            "/api/finance/cost-pool",
            json=_entry_payload(
                key="p1131-amount-over-limit",
                amount="1000000000000.00",
            ),
        )
        assert over_limit.status_code == 422, over_limit.text

        for field, key in (
            ("description", "p1131-blank-description"),
            ("source_reference", "p1131-blank-source"),
        ):
            payload = _entry_payload(key=key)
            payload[field] = "   "
            rejected = client.post("/api/finance/cost-pool", json=payload)
            assert rejected.status_code == 422, rejected.text

    with factory() as db:
        assert db.scalar(select(func.count()).select_from(FinanceCostPoolEntry)) == 1


def _import_workbook_bytes(*, formula_as_text: bool = True) -> bytes:
    workbook = load_workbook(BytesIO(_blank_template_bytes()))
    sheet = workbook["成本费用导入"]
    sheet.append(
        [
            "2026-08",
            date(2026, 8, 31),
            "administrative_expense",
            "ADMIN",
            "=SUM(1,1)",
            Decimal("100.20"),
            Decimal("0.00"),
            "办公用品店",
            "EXP-001",
            "unallocated",
            "Excel 导入",
        ]
    )
    if formula_as_text:
        # Keep the formula-looking value as literal text so this workbook tests
        # export injection escaping.  A separate case below leaves the cell as
        # a real Excel formula and verifies that import rejects it.
        sheet.cell(sheet.max_row, 5).data_type = "s"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


_TEMPLATE_CACHE: bytes | None = None


def _blank_template_bytes() -> bytes:
    assert _TEMPLATE_CACHE is not None
    return _TEMPLATE_CACHE


def test_cost_pool_excel_template_preview_apply_export_and_formula_safety(
    p1_131_cost_app,
) -> None:
    from app.models.finance_cost import FinanceCostPoolEntry

    global _TEMPLATE_CACHE
    app, factory = p1_131_cost_app
    with TestClient(app) as client:
        _login(client)
        template = client.get("/api/finance/cost-pool/template")
        assert template.status_code == 200, template.text
        _TEMPLATE_CACHE = template.content

        formula_content = _import_workbook_bytes(formula_as_text=False)
        formula_preview = client.post(
            "/api/finance/cost-pool/import",
            files={
                "file": (
                    "formula-costs.xlsx",
                    formula_content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"apply": "false"},
        )
        assert formula_preview.status_code == 200, formula_preview.text
        assert formula_preview.json()["row_count"] == 0
        assert formula_preview.json()["error_count"] == 1
        assert formula_preview.json()["errors"][0]["row"] == 2
        assert "不支持公式" in formula_preview.json()["errors"][0]["message"]
        assert "说明*" in formula_preview.json()["errors"][0]["message"]

        formula_apply = client.post(
            "/api/finance/cost-pool/import",
            files={
                "file": (
                    "formula-costs.xlsx",
                    formula_content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"apply": "true", "idempotency_key": "p1131-formula-import"},
        )
        assert formula_apply.status_code == 422, formula_apply.text
        assert formula_apply.json()["detail"]["message"] == "导入文件仍有错误，未写入任何数据"

        content = _import_workbook_bytes()

        preview = client.post(
            "/api/finance/cost-pool/import",
            files={
                "file": (
                    "costs.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"apply": "false"},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["row_count"] == 1
        assert preview.json()["error_count"] == 0
        assert Decimal(str(preview.json()["total_amount"])) == Decimal("100.20")

        applied = client.post(
            "/api/finance/cost-pool/import",
            files={
                "file": (
                    "costs.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"apply": "true", "idempotency_key": "p1131-import-001"},
        )
        assert applied.status_code == 200, applied.text
        assert applied.json()["applied"] is True
        entry_id = applied.json()["created_ids"][0]

        entries = client.get(
            "/api/finance/cost-pool", params={"month": "2026-08"}
        )
        assert entries.status_code == 200, entries.text
        imported = next(
            item for item in entries.json()["items"] if item["id"] == entry_id
        )
        original_source_reference = imported["source_reference"]
        original_source_fingerprint = imported["source_fingerprint"]
        assert original_source_reference.startswith("costs.xlsx:")
        assert "#row:2" in original_source_reference
        assert original_source_fingerprint

        updated = client.put(
            f"/api/finance/cost-pool/{entry_id}",
            json={
                "cost_month": imported["cost_month"],
                "document_date": imported["document_date"],
                "cost_center_id": imported["cost_center_id"],
                "cost_category": imported["cost_category"],
                "allocation_basis": imported["allocation_basis"],
                "description": imported["description"],
                "counterparty_name": imported["counterparty_name"],
                "document_number": imported["document_number"],
                "source_reference": "尝试覆盖 Excel 导入来源",
                "amount": str(imported["amount"]),
                "tax_amount": str(imported["tax_amount"]),
                "note": "导入后人工复核",
                "expected_version": imported["version"],
                "idempotency_key": "p1131-import-update",
            },
        )
        assert updated.status_code == 200, updated.text
        assert updated.json()["version"] == 2
        assert updated.json()["source_reference"] == original_source_reference
        assert updated.json()["source_fingerprint"] == original_source_fingerprint
        assert updated.json()["source_fingerprint"]

        confirmed = client.post(
            f"/api/finance/cost-pool/{entry_id}/confirm",
            json={
                "expected_version": 2,
                "idempotency_key": "p1131-import-confirm",
            },
        )
        assert confirmed.status_code == 200, confirmed.text

        centers = client.get(
            "/api/finance/cost-centers", params={"include_inactive": "true"}
        )
        assert centers.status_code == 200, centers.text
        admin_center = next(
            item
            for item in centers.json()["items"]
            if item["id"] == imported["cost_center_id"]
        )
        assert admin_center["code"] == "ADMIN"
        renamed_disabled = client.put(
            f"/api/finance/cost-centers/{admin_center['id']}",
            json={
                "name": "管理费用（历史停用）",
                "is_active": False,
                "expected_version": admin_center["version"],
                "idempotency_key": "p1131-admin-disable",
            },
        )
        assert renamed_disabled.status_code == 200, renamed_disabled.text
        assert renamed_disabled.json()["name"] == "管理费用（历史停用）"
        assert renamed_disabled.json()["is_active"] is False

        replay = client.post(
            "/api/finance/cost-pool/import",
            files={
                "file": (
                    "costs.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"apply": "true", "idempotency_key": "p1131-import-001"},
        )
        assert replay.status_code == 200, replay.text
        assert replay.json()["applied"] is True
        assert replay.json()["created_ids"] == [entry_id]
        assert replay.json()["file_hash"] == applied.json()["file_hash"]
        assert Decimal(str(replay.json()["total_amount"])) == Decimal("100.20")

        duplicate_file = client.post(
            "/api/finance/cost-pool/import",
            files={
                "file": (
                    "costs.xlsx",
                    content,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            data={"apply": "true", "idempotency_key": "p1131-import-002"},
        )
        assert duplicate_file.status_code == 409, duplicate_file.text

        exported = client.get(
            "/api/finance/cost-pool/export",
            params={"month": "2026-08", "status": "confirmed"},
        )
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content), data_only=False)
        details = workbook["成本费用明细"]
        assert details["G2"].value == "'=SUM(1,1)"
        assert Decimal(str(details["J2"].value)) == Decimal("100.2")
        assert details["N1"].value == "来源引用"
        assert details["N2"].value == original_source_reference
        assert details["O1"].value == "导入行防重指纹"
        assert details["O2"].value == original_source_fingerprint
        assert workbook["月度汇总"]["B2"].value == 100.2

    with factory() as db:
        assert db.scalar(select(func.count()).select_from(FinanceCostPoolEntry)) == 1
        stored = db.get(FinanceCostPoolEntry, entry_id)
        assert stored is not None
        assert stored.source_reference == original_source_reference
        assert stored.source_fingerprint == original_source_fingerprint
