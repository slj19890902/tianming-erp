from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.test_p1_131_cost_pool import (
    _entry_payload,
    _login,
    p1_131_cost_app,
)


INDEX = Path("static/index.html").read_text(encoding="utf-8")


def _summary_rows(workbook) -> dict[str, tuple[object, object]]:
    return {
        str(label): (value, note)
        for label, value, note in workbook["老板月报"].iter_rows(
            min_col=1,
            max_col=3,
            values_only=True,
        )
        if label
    }


def _seed_monthly_report_facts(factory) -> None:
    from app.models.customer import Customer
    from app.models.finance import Statement
    from app.models.finance_payable import FinancePayable
    from app.models.user import User

    with factory() as db:
        admin = db.query(User).filter(User.username == "p1131-admin").one()
        customer = Customer(
            name="昆山华诚电子有限公司",
            chinese_short_name="华诚电子",
            customer_code="P1132",
            statement_cycle_start_day=20,
            is_active=True,
        )
        db.add(customer)
        db.flush()
        db.add_all(
            [
                Statement(
                    statement_number="ST-P1-132-CONFIRMED",
                    customer_id=customer.id,
                    settlement_name_snapshot="合作纸箱厂统一购方",
                    statement_month="2026-08",
                    total_receivable=Decimal("100.00"),
                    total_gross_profit=Decimal("0.00"),
                    invoiced_amount=Decimal("40.00"),
                    settled_amount=Decimal("100.00"),
                    status="settled",
                    confirmation_status="confirmed",
                    confirmed_by=admin.id,
                    confirmed_at=datetime(2026, 8, 31, 9, 0, 0),
                ),
                Statement(
                    statement_number="ST-P1-132-DRAFT",
                    customer_id=customer.id,
                    statement_month="2026-08",
                    total_receivable=Decimal("999.00"),
                    total_gross_profit=Decimal("0.00"),
                    invoiced_amount=Decimal("0.00"),
                    status="unsettled",
                    confirmation_status="draft",
                ),
                FinancePayable(
                    counterparty_name="员工工资",
                    category="wages",
                    document_number="PAYROLL-202608",
                    document_date=date(2026, 8, 31),
                    due_date=date(2026, 9, 10),
                    amount=Decimal("40.00"),
                    status="confirmed",
                    note="同一工资事实也进入管理成本池，但月报不得重复相加",
                    idempotency_key="p1132-payable-confirmed",
                    confirmed_by=admin.id,
                    confirmed_at=datetime(2026, 8, 31, 10, 0, 0),
                    created_by=admin.id,
                ),
                FinancePayable(
                    counterparty_name="不应进入月报的作废记录",
                    category="other",
                    document_number="VOID-202608",
                    document_date=date(2026, 8, 31),
                    amount=Decimal("888.00"),
                    status="voided",
                    note="作废记录不得导出",
                    idempotency_key="p1132-payable-voided",
                    voided_by=admin.id,
                    voided_at=datetime(2026, 8, 31, 11, 0, 0),
                    created_by=admin.id,
                ),
            ]
        )
        db.commit()


def test_management_report_has_five_sheets_and_keeps_ledgers_separate(
    p1_131_cost_app,
) -> None:
    app, factory = p1_131_cost_app
    from app.models.audit import OperationLog
    from app.api.finance import current_customer_months
    from app.models.user import User

    _seed_monthly_report_facts(factory)
    with factory() as db:
        admin = db.query(User).filter(User.username == "p1131-admin").one()
        workbench = current_customer_months(
            statement_month="2026-08",
            balance_type="pending_invoice",
            customer_id=None,
            page=1,
            page_size=50,
            db=db,
            user=admin,
        )
    assert Decimal(str(workbench["summary"]["pending_invoice_amount"])) == Decimal(
        "60.00"
    )
    statement_rows = workbench["items"][0]["statements"]
    draft_statement = next(
        row for row in statement_rows if row["statement_number"] == "ST-P1-132-DRAFT"
    )
    assert Decimal(str(draft_statement["pending_invoice_amount"])) == Decimal("0.00")
    assert draft_statement["invoice_status"] == "not_ready"

    with TestClient(app) as client:
        _login(client)
        payload = _entry_payload(
            key="p1132-cost-create",
            amount="40.00",
            description="2026年8月生产工资",
        )
        payload["document_number"] = "PAYROLL-202608"
        payload["source_reference"] = "PAYROLL-202608"
        created = client.post("/api/finance/cost-pool", json=payload)
        assert created.status_code == 201, created.text
        confirmed = client.post(
            f"/api/finance/cost-pool/{created.json()['id']}/confirm",
            json={
                "expected_version": 1,
                "idempotency_key": "p1132-cost-confirm",
            },
        )
        assert confirmed.status_code == 200, confirmed.text
        with factory() as db:
            operation_log_count_before_export = db.query(OperationLog).count()

        cost_workpaper = client.get(
            "/api/finance/cost-pool/export",
            params={"month": "2026-08", "status": "confirmed"},
        )
        assert cost_workpaper.status_code == 200, cost_workpaper.text
        with factory() as db:
            assert db.query(OperationLog).count() == operation_log_count_before_export

        response = client.get(
            "/api/finance/management-report/export",
            params={"month": "2026-08"},
        )
        assert response.status_code == 200, response.text
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        with factory() as db:
            assert db.query(OperationLog).count() == operation_log_count_before_export

    assert workbook.sheetnames == ["老板月报", "客户应收", "应付支出", "成本费用", "补充材料成本依据"]
    summary = _summary_rows(workbook)
    assert Decimal(str(summary["已确认对账收入"][0])) == Decimal("100.00")
    assert Decimal(str(summary["已登记开票"][0])) == Decimal("40.00")
    assert Decimal(str(summary["待开票"][0])) == Decimal("60.00")
    assert summary["待开票"][1] == "不代表客户欠款或银行未到账"

    # The same payroll fact deliberately exists in AP and in the cost pool.
    # The management workbook must keep the two ledgers visible but separate.
    assert Decimal(str(summary["已确认应付 / 支出"][0])) == Decimal("40.00")
    assert summary["已确认应付 / 支出"][1] == "独立应付台账，不与成本池相加"
    assert Decimal(str(summary["已确认成本费用"][0])) == Decimal("40.00")
    assert summary["已确认成本费用"][1] == "独立管理成本池"
    assert "应付与成本合计" not in summary

    receivable_row = tuple(
        workbook["客户应收"].iter_rows(min_row=2, max_row=2, values_only=True)
    )[0]
    assert receivable_row == ("合作纸箱厂统一购方", 1, 100, 40, 60)
    payable_rows = list(workbook["应付支出"].iter_rows(min_row=2, values_only=True))
    assert len(payable_rows) == 1
    assert payable_rows[0][1] == "员工工资"
    assert Decimal(str(payable_rows[0][4])) == Decimal("40.00")
    cost_rows = list(workbook["成本费用"].iter_rows(min_row=2, values_only=True))
    assert len(cost_rows) == 1
    assert cost_rows[0][1] == "已确认"
    assert Decimal(str(cost_rows[0][6])) == Decimal("40.00")


def test_management_report_export_rejects_missing_or_insufficient_permission(
    p1_131_cost_app,
) -> None:
    from app.api.finance import require_company_finance_read
    from app.core.security import hash_password
    from app.models.user import User

    app, factory = p1_131_cost_app
    with factory() as db:
        db.add(
            User(
                username="p1132-warehouse",
                password_hash=hash_password("RolePass123!"),
                role="workshop",
                real_name="Warehouse Operator",
                must_change_password=False,
            )
        )
        db.commit()
    with factory() as db:
        restricted = db.query(User).filter(User.username == "p1131-restricted").one()
        with pytest.raises(HTTPException) as denied_overview:
            require_company_finance_read(user=restricted, db=db)
        assert denied_overview.value.status_code == 403

    with TestClient(app) as client:
        anonymous = client.get(
            "/api/finance/management-report/export",
            params={"month": "2026-08"},
        )
        assert anonymous.status_code == 401

        _login(client, "p1132-warehouse")
        denied = client.get(
            "/api/finance/management-report/export",
            params={"month": "2026-08"},
        )
        assert denied.status_code == 403

        _login(client, "p1131-restricted")
        company_scope_denied = client.get(
            "/api/finance/management-report/export",
            params={"month": "2026-08"},
        )
        assert company_scope_denied.status_code == 403


def test_finance_monthly_workbench_has_four_actions_and_mobile_two_columns() -> None:
    start = INDEX.index('<div class="panel finance-month-workbench">')
    end = INDEX.index('<div class="finance-overview-grid"', start)
    workbench = INDEX[start:end]
    assert workbench.count('class="finance-month-action"') == 4
    for action, label in (
        ("pending_reconciliation", "客户对账"),
        ("pending_invoice", "开票准备"),
        ("cost_drafts", "费用复核"),
        ("close_readiness", "月报检查"),
    ):
        assert f"openFinanceMonthlyAction('{action}')" in workbench
        assert label in workbench
    assert 'v-if="canExportFinanceCosts"' in workbench
    assert 'v-else-if="!financeOverviewState.error"' in INDEX
    assert "该月对账已开票" in INDEX
    assert "该月确认对账对应开票" in INDEX

    action_start = INDEX.index("openFinanceMonthlyAction(action) {")
    action_end = INDEX.index("async exportFinanceManagementReport()", action_start)
    action_body = INDEX[action_start:action_end]
    assert 'this.financeFilters.balance_type = action' in action_body
    assert 'return this.setFinanceView("current")' in action_body
    assert 'action === "cost_drafts" ? "draft" : ""' in action_body
    assert 'return this.setFinanceView("expenses")' in action_body

    export_start = action_end
    export_end = INDEX.index("financeCostStatusAmount(status)", export_start)
    export_body = INDEX[export_start:export_end]
    assert "if (!this.canExportFinanceCosts" in export_body
    assert 'axios.get("/api/finance/management-report/export"' in export_body
    assert 'responseType:"blob"' in export_body

    overview_start = INDEX.index("async loadFinanceOverview(sessionContext = null) {")
    overview_end = INDEX.index("financeTrendHeight(value)", overview_start)
    overview_body = INDEX[overview_start:overview_end]
    assert 'axios.get("/api/finance/cost-pool/summary"' in overview_body
    assert ".catch(()=>({data:null}))" not in overview_body

    breakpoint = INDEX.index("@media (max-width: 980px)")
    mobile_css = INDEX[breakpoint : INDEX.index(".table-wrap", breakpoint)]
    compact_css = "".join(mobile_css.split())
    assert (
        ".finance-month-actions{grid-template-columns:repeat(2,minmax(0,1fr))}"
        in compact_css
    )
