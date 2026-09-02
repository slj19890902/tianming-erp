from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.test_p1_131_cost_pool import _login, p1_131_cost_app


def _enable_router(app) -> None:
    from app.api.finance_simplified import router

    app.include_router(router, prefix="/api/finance")


def _centers(factory) -> dict[str, int]:
    from app.models.finance_cost import FinanceCostCenter

    with factory() as db:
        return {row.code: row.id for row in db.query(FinanceCostCenter).all()}


def _rule_payload(
    *,
    key: str,
    rule_type: str,
    name: str,
    center_id: int,
    **values,
) -> dict:
    return {
        "rule_type": rule_type,
        "name": name,
        "role_name": values.get("role_name"),
        "cost_center_id": center_id,
        "start_month": values.get("start_month", "2026-01"),
        "end_month": values.get("end_month"),
        "base_amount": values.get("base_amount", "0"),
        "allowance_amount": values.get("allowance_amount", "0"),
        "employer_social_amount": values.get("employer_social_amount", "0"),
        "deduction_amount": values.get("deduction_amount", "0"),
        "annual_amount": values.get("annual_amount", "0"),
        "monthly_amount": values.get("monthly_amount", "0"),
        "due_day": values.get("due_day", 20),
        "note": values.get("note"),
        "is_active": values.get("is_active", True),
        "idempotency_key": key,
    }


def test_recurring_wages_rent_and_vehicle_loan_generate_monthly_drafts(
    p1_131_cost_app,
) -> None:
    app, factory = p1_131_cost_app
    _enable_router(app)
    centers = _centers(factory)
    with TestClient(app) as client:
        _login(client)
        employee = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="p1144-rule-employee-001",
                rule_type="employee_wage",
                name="张师傅",
                role_name="车间",
                center_id=centers["PROD"],
                base_amount="5000",
                allowance_amount="300",
                employer_social_amount="700",
                deduction_amount="100",
            ),
        )
        assert employee.status_code == 201, employee.text
        assert employee.json()["cost_category"] == "production_wages"

        future_employee = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="p1144-rule-future-employee",
                rule_type="employee_wage",
                name="未来入职员工",
                role_name="车间",
                center_id=centers["PROD"],
                start_month="2027-01",
                base_amount="5000",
            ),
        )
        assert future_employee.status_code == 201, future_employee.text

        rent = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="p1144-rule-rent-001",
                rule_type="annual_rent",
                name="厂房房租",
                center_id=centers["PROD"],
                annual_amount="380000",
            ),
        )
        assert rent.status_code == 201, rent.text
        loan = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="p1144-rule-loan-001",
                rule_type="fixed_monthly",
                name="车辆月供",
                center_id=centers["FINANCE"],
                monthly_amount="8000",
            ),
        )
        assert loan.status_code == 201, loan.text
        wrong_wage_center = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="p1144-rule-wrong-center",
                rule_type="employee_wage",
                name="错误成本中心",
                center_id=centers["FINANCE"],
                base_amount="1000",
            ),
        )
        assert wrong_wage_center.status_code == 409

        generated = client.post(
            "/api/finance/simple-finance/recurring-generate",
            json={"month": "2026-11", "idempotency_key": "p1144-generate-202611"},
        )
        assert generated.status_code == 200, generated.text
        assert generated.json()["created_count"] == 3
        assert generated.json()["active_employee_count"] == 1

        replay = client.post(
            "/api/finance/simple-finance/recurring-generate",
            json={"month": "2026-11", "idempotency_key": "p1144-generate-202611"},
        )
        assert replay.status_code == 200
        assert replay.json() == generated.json()

        summary = client.get(
            "/api/finance/simple-finance/summary", params={"month": "2026-11"}
        )
        assert summary.status_code == 200, summary.text
        groups = summary.json()["primary_groups"]
        assert Decimal(str(groups["personnel"]["draft_amount"])) == Decimal("5900.00")
        assert Decimal(str(groups["fixed_operating"]["draft_amount"])) == Decimal(
            "39666.67"
        )
        assert summary.json()["active_employee_count"] == 1

        december = client.post(
            "/api/finance/simple-finance/recurring-generate",
            json={"month": "2026-12", "idempotency_key": "p1144-generate-202612"},
        )
        assert december.status_code == 200, december.text

    from app.models.finance_cost import FinanceCostPoolEntry

    with factory() as db:
        entries = db.query(FinanceCostPoolEntry).filter_by(cost_month="2026-12").all()
        rent_entry = next(row for row in entries if row.cost_category == "factory_rent")
        assert rent_entry.amount == Decimal("31666.63")
        assert len({row.source_fingerprint for row in entries}) == 3


def test_finance_text_identifiers_reject_values_that_become_blank_after_trim(
    p1_131_cost_app,
) -> None:
    app, factory = p1_131_cost_app
    _enable_router(app)
    centers = _centers(factory)
    with TestClient(app) as client:
        _login(client)
        blank_key = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="        ",
                rule_type="fixed_monthly",
                name="车辆月供",
                center_id=centers["FINANCE"],
                monthly_amount="1000",
            ),
        )
        assert blank_key.status_code == 422

        blank_name = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="p1144-blank-name",
                rule_type="fixed_monthly",
                name=" ",
                center_id=centers["FINANCE"],
                monthly_amount="1000",
            ),
        )
        assert blank_name.status_code == 422

        blank_bill = client.post(
            "/api/finance/simple-finance/acceptances",
            json={
                "bill_number": " ",
                "customer_id": 1,
                "amount": "100.00",
                "received_date": "2026-09-01",
                "maturity_date": "2099-12-01",
                "idempotency_key": "p1144-blank-bill",
            },
        )
        assert blank_bill.status_code == 422


def test_recurring_generate_cas_cannot_overwrite_concurrently_confirmed_cost(
    p1_131_cost_app,
) -> None:
    from fastapi import HTTPException
    from starlette.requests import Request

    from app.api.finance_simplified import RecurringGenerate, generate_recurring_drafts
    from app.core.time_contract import beijing_now_naive
    from app.models.finance_cost import FinanceCostPoolEntry
    from app.models.user import User

    app, factory = p1_131_cost_app
    _enable_router(app)
    centers = _centers(factory)
    with TestClient(app) as client:
        _login(client)
        created = client.post(
            "/api/finance/simple-finance/recurring-rules",
            json=_rule_payload(
                key="p1144-cas-rule-create",
                rule_type="fixed_monthly",
                name="并发车辆月供",
                center_id=centers["FINANCE"],
                monthly_amount="8000",
            ),
        )
        assert created.status_code == 201, created.text
        generated = client.post(
            "/api/finance/simple-finance/recurring-generate",
            json={"month": "2026-11", "idempotency_key": "p1144-cas-generate-one"},
        )
        assert generated.status_code == 200, generated.text
        entry_id = generated.json()["generated_entry_ids"][0]
        update_payload = _rule_payload(
            key="p1144-cas-rule-update",
            rule_type="fixed_monthly",
            name="并发车辆月供",
            center_id=centers["FINANCE"],
            monthly_amount="9000",
        )
        update_payload["expected_version"] = 1
        updated = client.put(
            f"/api/finance/simple-finance/recurring-rules/{created.json()['id']}",
            json=update_payload,
        )
        assert updated.status_code == 200, updated.text

    stale_session = factory()
    winning_session = factory()
    try:
        stale_entry = stale_session.get(FinanceCostPoolEntry, entry_id)
        stale_user = stale_session.query(User).filter_by(username="p1131-admin").one()
        assert stale_entry.status == "draft" and stale_entry.version == 1
        stale_session.commit()

        winner_entry = winning_session.get(FinanceCostPoolEntry, entry_id)
        winner_user = winning_session.query(User).filter_by(username="p1131-admin").one()
        winner_entry.status = "confirmed"
        winner_entry.confirmed_by = winner_user.id
        winner_entry.confirmed_at = beijing_now_naive()
        winner_entry.version = 2
        winning_session.commit()

        request = Request(
            {
                "type": "http",
                "method": "POST",
                "path": "/api/finance/simple-finance/recurring-generate",
                "headers": [],
                "query_string": b"",
                "client": ("test", 1),
                "server": ("test", 80),
            }
        )
        with pytest.raises(HTTPException) as error:
            generate_recurring_drafts(
                payload=RecurringGenerate(
                    month="2026-11", idempotency_key="p1144-cas-generate-two"
                ),
                request=request,
                db=stale_session,
                user=stale_user,
            )
        assert error.value.status_code == 409
    finally:
        stale_session.close()
        winning_session.close()

    with factory() as db:
        protected = db.get(FinanceCostPoolEntry, entry_id)
        assert protected.status == "confirmed"
        assert protected.amount == Decimal("8000.00")
        assert protected.version == 2


def test_utility_meter_uses_invoice_fact_and_carries_previous_reading(
    p1_131_cost_app,
) -> None:
    app, factory = p1_131_cost_app
    _enable_router(app)
    centers = _centers(factory)
    with TestClient(app) as client:
        _login(client)
        created = client.post(
            "/api/finance/simple-finance/utility-readings",
            json={
                "cost_month": "2026-08",
                "utility_type": "electricity",
                "cost_center_id": centers["PROD"],
                "previous_reading": "100.100",
                "current_reading": "120.225",
                "unit_price": "1.2000",
                "paid_amount": "0",
                "idempotency_key": "p1144-electricity-create",
            },
        )
        assert created.status_code == 200, created.text
        assert Decimal(str(created.json()["usage_quantity"])) == Decimal("20.125")
        assert Decimal(str(created.json()["calculated_amount"])) == Decimal("24.15")
        assert Decimal(str(created.json()["recognized_amount"])) == Decimal("24.15")

        updated = client.post(
            "/api/finance/simple-finance/utility-readings",
            json={
                "cost_month": "2026-08",
                "utility_type": "electricity",
                "cost_center_id": centers["PROD"],
                "previous_reading": "100.100",
                "current_reading": "120.225",
                "unit_price": "1.2000",
                "invoice_number": "ELEC-202608",
                "invoice_date": "2026-08-31",
                "invoice_amount": "25.00",
                "paid_amount": "0",
                "expected_version": 1,
                "idempotency_key": "p1144-electricity-invoice",
            },
        )
        assert updated.status_code == 200, updated.text
        assert Decimal(str(updated.json()["difference_amount"])) == Decimal("0.85")
        assert updated.json()["version"] == 2

        confirmed = client.post(
            f"/api/finance/cost-pool/{updated.json()['cost_pool_entry_id']}/confirm",
            json={
                "expected_version": 2,
                "idempotency_key": "p1144-electricity-confirm",
            },
        )
        assert confirmed.status_code == 200, confirmed.text
        assert confirmed.json()["status"] == "confirmed"

        paid = client.post(
            "/api/finance/simple-finance/utility-readings",
            json={
                "cost_month": "2026-08",
                "utility_type": "electricity",
                "cost_center_id": centers["PROD"],
                "previous_reading": "100.100",
                "current_reading": "120.225",
                "unit_price": "1.2000",
                "invoice_number": "ELEC-202608",
                "invoice_date": "2026-08-31",
                "invoice_amount": "25.00",
                "paid_amount": "25.00",
                "payment_date": "2026-09-05",
                "expected_version": 2,
                "idempotency_key": "p1144-electricity-payment",
            },
        )
        assert paid.status_code == 200, paid.text
        assert paid.json()["version"] == 3
        assert Decimal(str(paid.json()["paid_amount"])) == Decimal("25.00")

        blocked_after_confirm = client.post(
            "/api/finance/simple-finance/utility-readings",
            json={
                "cost_month": "2026-08",
                "utility_type": "electricity",
                "cost_center_id": centers["PROD"],
                "previous_reading": "100.100",
                "current_reading": "121.000",
                "unit_price": "1.2000",
                "invoice_number": "ELEC-202608",
                "invoice_date": "2026-08-31",
                "invoice_amount": "25.00",
                "paid_amount": "25.00",
                "payment_date": "2026-09-05",
                "expected_version": 3,
                "idempotency_key": "p1144-electricity-locked",
            },
        )
        assert blocked_after_confirm.status_code == 409

        next_month = client.get(
            "/api/finance/simple-finance/utility-readings", params={"month": "2026-09"}
        )
        assert next_month.status_code == 200
        assert Decimal(
            str(next_month.json()["recommended_previous_readings"]["electricity"])
        ) == Decimal("120.225")

    from app.models.finance_cost import FinanceCostPoolEntry

    with factory() as db:
        entry = db.query(FinanceCostPoolEntry).filter_by(
            source_fingerprint="utility:electricity:2026-08"
        ).one()
        assert entry.amount == Decimal("25.00")
        assert entry.source_type == "utility_reading"
        assert entry.status == "confirmed"


def test_utility_save_cas_cannot_overwrite_concurrently_confirmed_cost(
    p1_131_cost_app,
) -> None:
    from fastapi import HTTPException
    from starlette.requests import Request

    from app.api.finance_simplified import UtilityReadingPayload, save_utility_reading
    from app.core.time_contract import beijing_now_naive
    from app.models.finance_cost import FinanceCostPoolEntry
    from app.models.finance_simplified import FinanceUtilityReading
    from app.models.user import User

    app, factory = p1_131_cost_app
    _enable_router(app)
    centers = _centers(factory)
    with TestClient(app) as client:
        _login(client)
        invalid_payment_date = client.post(
            "/api/finance/simple-finance/utility-readings",
            json={
                "cost_month": "2026-08",
                "utility_type": "electricity",
                "cost_center_id": centers["PROD"],
                "previous_reading": "100.100",
                "current_reading": "120.225",
                "unit_price": "1.2000",
                "paid_amount": "0",
                "payment_date": "2026-09-05",
                "idempotency_key": "p1144-electricity-invalid-payment-date",
            },
        )
        assert invalid_payment_date.status_code == 422

        created = client.post(
            "/api/finance/simple-finance/utility-readings",
            json={
                "cost_month": "2026-10",
                "utility_type": "water",
                "cost_center_id": centers["PROD"],
                "previous_reading": "12345.678",
                "current_reading": "12346.001",
                "unit_price": "4.0000",
                "paid_amount": "0",
                "idempotency_key": "p1144-water-cas-create",
            },
        )
        assert created.status_code == 200, created.text
        reading_id = created.json()["id"]
        entry_id = created.json()["cost_pool_entry_id"]

    stale_session = factory()
    winning_session = factory()
    try:
        stale_reading = stale_session.get(FinanceUtilityReading, reading_id)
        stale_entry = stale_session.get(FinanceCostPoolEntry, entry_id)
        stale_user = stale_session.query(User).filter_by(username="p1131-admin").one()
        assert stale_reading.version == 1 and stale_entry.status == "draft"
        stale_session.commit()

        winner_entry = winning_session.get(FinanceCostPoolEntry, entry_id)
        winner_user = winning_session.query(User).filter_by(username="p1131-admin").one()
        winner_entry.status = "confirmed"
        winner_entry.confirmed_by = winner_user.id
        winner_entry.confirmed_at = beijing_now_naive()
        winner_entry.version = 2
        winning_session.commit()

        request = Request(
            {
                "type": "http",
                "method": "POST",
                "path": "/api/finance/simple-finance/utility-readings",
                "headers": [],
                "query_string": b"",
                "client": ("test", 1),
                "server": ("test", 80),
            }
        )
        with pytest.raises(HTTPException) as error:
            save_utility_reading(
                payload=UtilityReadingPayload(
                    cost_month="2026-10",
                    utility_type="water",
                    cost_center_id=centers["PROD"],
                    previous_reading=Decimal("12345.678"),
                    current_reading=Decimal("12346.001"),
                    unit_price=Decimal("4.0000"),
                    invoice_number="WATER-202610",
                    invoice_date=date(2026, 10, 31),
                    invoice_amount=Decimal("2.00"),
                    paid_amount=Decimal("0"),
                    expected_version=1,
                    idempotency_key="p1144-water-cas-update",
                ),
                request=request,
                db=stale_session,
                user=stale_user,
            )
        assert error.value.status_code == 409
    finally:
        stale_session.close()
        winning_session.close()

    with factory() as db:
        protected_entry = db.get(FinanceCostPoolEntry, entry_id)
        protected_reading = db.get(FinanceUtilityReading, reading_id)
        assert protected_entry.status == "confirmed"
        assert protected_entry.amount == Decimal("1.29")
        assert protected_entry.version == 2
        assert protected_reading.invoice_number is None
        assert protected_reading.version == 1


def _seed_payable_statement(factory) -> tuple[int, int]:
    from app.models.customer import Customer
    from app.models.finance_payable import FinancePayable
    from app.models.supplier import Supplier
    from app.models.supplier_settlement import (
        SupplierMonthlyAdjustment,
        SupplierMonthlyInvoice,
        SupplierMonthlyStatement,
    )
    from app.models.user import User

    with factory() as db:
        admin = db.query(User).filter_by(username="p1131-admin").one()
        customer = Customer(
            name="昆山华诚电子有限公司",
            chinese_short_name="华诚",
            customer_code="P1144-CUSTOMER",
            is_active=True,
        )
        supplier = Supplier(
            standard_name="上游纸板厂",
            normalized_name="上游纸板厂",
            display_name="纸板厂",
            is_active=True,
        )
        db.add_all([customer, supplier])
        db.flush()
        payable = FinancePayable(
            supplier_id=supplier.id,
            counterparty_name=supplier.standard_name,
            category="material",
            document_number="SUP-ST-202608",
            document_date=date(2026, 8, 20),
            amount=Decimal("1000.00"),
            status="confirmed",
            idempotency_key="p1144-seeded-payable",
            created_by=admin.id,
        )
        db.add(payable)
        db.flush()
        statement = SupplierMonthlyStatement(
            statement_number="SUP-ST-202608",
            supplier_id=supplier.id,
            supplier_name_snapshot=supplier.standard_name,
            settlement_month="2026-08",
            period_start=date(2026, 7, 21),
            period_end=date(2026, 8, 20),
            currency="CNY",
            tax_basis="tax_inclusive",
            status="invoiced_pending_payment",
            erp_amount=Decimal("980.00"),
            adjustment_amount=Decimal("20.00"),
            adjusted_amount=Decimal("1000.00"),
            supplier_statement_number="BILL-202608",
            supplier_statement_date=date(2026, 8, 21),
            supplier_statement_amount=Decimal("1050.00"),
            confirmed_amount=Decimal("1000.00"),
            invoice_allocated_amount=Decimal("1000.00"),
            paid_amount=Decimal("0.00"),
            finance_payable_id=payable.id,
            generated_by=admin.id,
        )
        db.add(statement)
        db.flush()
        db.add(
            SupplierMonthlyInvoice(
                statement_id=statement.id,
                supplier_id=supplier.id,
                invoice_number="INV-202608",
                invoice_date=date(2026, 8, 25),
                received_date=date(2026, 8, 26),
                invoice_total_amount=Decimal("1000.00"),
                allocated_amount=Decimal("1000.00"),
                tax_amount=Decimal("115.04"),
                created_by=admin.id,
            )
        )
        db.add(
            SupplierMonthlyAdjustment(
                statement_id=statement.id,
                difference_type="price",
                amount=Decimal("20.00"),
                note="供应商调价",
                created_by=admin.id,
            )
        )
        db.commit()
        return customer.id, statement.id


def test_material_summary_includes_mold_ink_and_keeps_zero_supplier_bill(
    p1_131_cost_app,
) -> None:
    from app.models.finance_payable import FinancePayable
    from app.models.supplier_settlement import SupplierMonthlyStatement
    from app.models.user import User

    app, factory = p1_131_cost_app
    _enable_router(app)
    _customer_id, statement_id = _seed_payable_statement(factory)
    with factory() as db:
        admin = db.query(User).filter_by(username="p1131-admin").one()
        statement = db.get(SupplierMonthlyStatement, statement_id)
        statement.supplier_statement_amount = Decimal("0.00")
        db.add(
            FinancePayable(
                counterparty_name="油墨供应商",
                category="material",
                document_number="INK-202608",
                document_date=date(2026, 8, 18),
                amount=Decimal("250.00"),
                status="confirmed",
                idempotency_key="p1144-standalone-ink",
                created_by=admin.id,
            )
        )
        db.commit()

    with TestClient(app) as client:
        _login(client)
        summary = client.get(
            "/api/finance/simple-finance/summary", params={"month": "2026-08"}
        )
        assert summary.status_code == 200, summary.text
        material = summary.json()["primary_groups"]["material"]
        assert Decimal(str(material["erp_amount"])) == Decimal("1250.00")
        assert Decimal(str(material["supplier_settlement_erp_amount"])) == Decimal(
            "1000.00"
        )
        assert Decimal(str(material["mold_ink_other_amount"])) == Decimal("250.00")
        assert Decimal(str(material["supplier_statement_amount"])) == Decimal("0.00")
        assert Decimal(str(material["statement_difference_amount"])) == Decimal(
            "-1000.00"
        )

        exported = client.get(
            "/api/finance/simple-finance/material-export", params={"month": "2026-08"}
        )
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content), data_only=True)
        assert workbook["供应商核票汇总"]["E2"].value == 0
        assert workbook["供应商核票汇总"]["F2"].value == -1000


def test_acceptance_supplier_payment_requires_company_cost_permission(
    p1_131_cost_app,
) -> None:
    from app.core.security import hash_password
    from app.models.access_control import UserPermissionOverride
    from app.models.user import User

    app, factory = p1_131_cost_app
    _enable_router(app)
    customer_id, statement_id = _seed_payable_statement(factory)
    with factory() as db:
        admin = db.query(User).filter_by(username="p1131-admin").one()
        restricted = User(
            username="p1144-no-cost",
            password_hash=hash_password("RolePass123!"),
            role="finance",
            real_name="No Cost Finance",
            customer_access_mode="all",
            must_change_password=False,
        )
        db.add(restricted)
        db.flush()
        db.add(
            UserPermissionOverride(
                user_id=restricted.id,
                permission_code="cost.view",
                is_allowed=False,
                granted_by=admin.id,
            )
        )
        db.commit()

    with TestClient(app) as client:
        _login(client)
        created = client.post(
            "/api/finance/simple-finance/acceptances",
            json={
                "bill_number": "ACCEPT-NO-COST",
                "customer_id": customer_id,
                "amount": "100.00",
                "received_date": "2026-09-01",
                "maturity_date": "2027-03-01",
                "idempotency_key": "p1144-no-cost-create",
            },
        )
        assert created.status_code == 201, created.text

        _login(client, "p1144-no-cost")
        denied = client.post(
            f"/api/finance/simple-finance/acceptances/{created.json()['id']}/endorse",
            json={
                "expected_version": 1,
                "supplier_statement_id": statement_id,
                "endorsement_date": "2026-09-02",
                "idempotency_key": "p1144-no-cost-endorse",
            },
        )
        assert denied.status_code == 403


def test_acceptance_idempotency_replays_before_customer_state_revalidation(
    p1_131_cost_app,
) -> None:
    from app.models.customer import Customer

    app, factory = p1_131_cost_app
    _enable_router(app)
    customer_id, _statement_id = _seed_payable_statement(factory)
    payload = {
        "bill_number": "ACCEPT-IDEMPOTENT",
        "customer_id": customer_id,
        "amount": "300.00",
        "received_date": "2026-09-01",
        "maturity_date": "2027-03-01",
        "idempotency_key": "p1144-acceptance-idempotent",
    }
    with TestClient(app) as client:
        _login(client)
        first = client.post("/api/finance/simple-finance/acceptances", json=payload)
        assert first.status_code == 201, first.text
        with factory() as db:
            customer = db.get(Customer, customer_id)
            customer.is_active = False
            db.commit()
        replay = client.post("/api/finance/simple-finance/acceptances", json=payload)
        assert replay.status_code == 201, replay.text
        assert replay.json() == first.json()


def test_acceptance_rejects_impossible_endorsement_and_early_maturity_dates(
    p1_131_cost_app,
) -> None:
    app, factory = p1_131_cost_app
    _enable_router(app)
    customer_id, statement_id = _seed_payable_statement(factory)
    with TestClient(app) as client:
        _login(client)
        created = client.post(
            "/api/finance/simple-finance/acceptances",
            json={
                "bill_number": "ACCEPT-DATE-GUARD",
                "customer_id": customer_id,
                "amount": "100.00",
                "received_date": "2026-09-01",
                "maturity_date": "2099-12-01",
                "idempotency_key": "p1144-date-guard-create",
            },
        )
        assert created.status_code == 201, created.text

        before_received = client.post(
            f"/api/finance/simple-finance/acceptances/{created.json()['id']}/endorse",
            json={
                "expected_version": 1,
                "supplier_statement_id": statement_id,
                "endorsement_date": "2026-08-31",
                "idempotency_key": "p1144-date-guard-endorse",
            },
        )
        assert before_received.status_code == 409

        early_maturity = client.post(
            f"/api/finance/simple-finance/acceptances/{created.json()['id']}/transition",
            json={
                "expected_version": 1,
                "action": "matured",
                "idempotency_key": "p1144-date-guard-mature",
            },
        )
        assert early_maturity.status_code == 409

        current = client.get("/api/finance/simple-finance/acceptances")
        assert current.status_code == 200
        guarded = next(
            row for row in current.json()["items"] if row["id"] == created.json()["id"]
        )
        assert guarded["status"] == "held"
        assert guarded["version"] == 1


def test_acceptance_cannot_exceed_confirmed_or_invoiced_supplier_balance(
    p1_131_cost_app,
) -> None:
    from app.models.finance_simplified import FinanceAcceptanceNote
    from app.models.supplier_settlement import SupplierMonthlyPayment

    app, factory = p1_131_cost_app
    _enable_router(app)
    customer_id, statement_id = _seed_payable_statement(factory)
    with TestClient(app) as client:
        _login(client)
        metadata = client.get("/api/finance/simple-finance/metadata")
        assert metadata.status_code == 200
        statement_option = next(
            row for row in metadata.json()["supplier_statements"] if row["id"] == statement_id
        )
        assert Decimal(str(statement_option["remaining_amount"])) == Decimal("1000.00")
        assert Decimal(str(statement_option["invoice_available_amount"])) == Decimal(
            "1000.00"
        )
        assert Decimal(str(statement_option["available_payment_amount"])) == Decimal(
            "1000.00"
        )

        created = client.post(
            "/api/finance/simple-finance/acceptances",
            json={
                "bill_number": "ACCEPT-OVER-BALANCE",
                "customer_id": customer_id,
                "amount": "1100.00",
                "received_date": "2026-09-01",
                "maturity_date": "2027-03-01",
                "idempotency_key": "p1144-over-balance-create",
            },
        )
        assert created.status_code == 201, created.text
        blocked = client.post(
            f"/api/finance/simple-finance/acceptances/{created.json()['id']}/endorse",
            json={
                "expected_version": 1,
                "supplier_statement_id": statement_id,
                "endorsement_date": "2026-09-02",
                "idempotency_key": "p1144-over-balance-endorse",
            },
        )
        assert blocked.status_code == 422, blocked.text

    with factory() as db:
        note = db.query(FinanceAcceptanceNote).filter_by(
            bill_number="ACCEPT-OVER-BALANCE"
        ).one()
        assert note.status == "held"
        assert note.version == 1
        assert db.query(SupplierMonthlyPayment).count() == 0


def test_supplier_payment_compare_and_swap_rejects_stale_competing_writer(
    p1_131_cost_app,
) -> None:
    from app.models.supplier_settlement import (
        SupplierMonthlyPayment,
        SupplierMonthlyStatement,
    )
    from app.models.user import User
    from app.services.supplier_monthly_settlement import (
        SupplierSettlementError,
        add_payment,
    )

    _app, factory = p1_131_cost_app
    _customer_id, statement_id = _seed_payable_statement(factory)
    stale_session = factory()
    winning_session = factory()
    try:
        stale_row = stale_session.get(SupplierMonthlyStatement, statement_id)
        assert stale_row.version == 1
        stale_session.commit()  # keep the v1 identity while ending its read transaction

        winner = winning_session.query(User).filter_by(username="p1131-admin").one()
        winning_row, _payment = add_payment(
            winning_session,
            statement_id=statement_id,
            expected_version=1,
            payment_date=date(2026, 9, 2),
            amount=Decimal("100.00"),
            reference="第一笔银行付款",
            user=winner,
        )
        assert winning_row.version == 2
        winning_session.commit()

        stale_user = stale_session.query(User).filter_by(username="p1131-admin").one()
        with pytest.raises(SupplierSettlementError) as error:
            add_payment(
                stale_session,
                statement_id=statement_id,
                expected_version=1,
                payment_date=date(2026, 9, 2),
                amount=Decimal("100.00"),
                reference="并发旧版本付款",
                user=stale_user,
            )
        assert error.value.code == "SUPPLIER_SETTLEMENT_STALE"
        stale_session.rollback()
    finally:
        stale_session.close()
        winning_session.close()

    with factory() as db:
        statement = db.get(SupplierMonthlyStatement, statement_id)
        assert statement.version == 2
        assert statement.paid_amount == Decimal("100.00")
        payments = db.query(SupplierMonthlyPayment).all()
        assert len(payments) == 1
        assert payments[0].amount == Decimal("100.00")


def test_supplier_invoice_stale_writer_cannot_overwrite_payment_state(
    p1_131_cost_app,
) -> None:
    from app.models.supplier_settlement import (
        SupplierMonthlyInvoice,
        SupplierMonthlyStatement,
    )
    from app.models.user import User
    from app.services.supplier_monthly_settlement import (
        SupplierSettlementError,
        add_invoice,
        add_payment,
    )

    _app, factory = p1_131_cost_app
    _customer_id, statement_id = _seed_payable_statement(factory)
    stale_session = factory()
    winning_session = factory()
    try:
        stale_row = stale_session.get(SupplierMonthlyStatement, statement_id)
        assert stale_row.version == 1
        stale_session.commit()

        winner = winning_session.query(User).filter_by(username="p1131-admin").one()
        add_payment(
            winning_session,
            statement_id=statement_id,
            expected_version=1,
            payment_date=date(2026, 9, 2),
            amount=Decimal("100.00"),
            reference="先发生的付款",
            user=winner,
        )
        winning_session.commit()

        stale_user = stale_session.query(User).filter_by(username="p1131-admin").one()
        with pytest.raises(SupplierSettlementError) as error:
            add_invoice(
                stale_session,
                statement_id=statement_id,
                expected_version=1,
                invoice_number="INV-STALE-WRITER",
                invoice_date=date(2026, 9, 2),
                received_date=date(2026, 9, 2),
                invoice_total_amount=Decimal("100.00"),
                allocated_amount=Decimal("100.00"),
                tax_amount=Decimal("11.50"),
                note="并发旧版本发票",
                user=stale_user,
            )
        assert error.value.code == "SUPPLIER_SETTLEMENT_STALE"
        stale_session.rollback()
    finally:
        stale_session.close()
        winning_session.close()

    with factory() as db:
        statement = db.get(SupplierMonthlyStatement, statement_id)
        assert statement.status == "partial_payment"
        assert statement.version == 2
        assert statement.invoice_allocated_amount == Decimal("1000.00")
        assert (
            db.query(SupplierMonthlyInvoice)
            .filter_by(invoice_number="INV-STALE-WRITER")
            .count()
            == 0
        )


def test_acceptance_endorsement_is_traceable_non_cash_supplier_payment_and_export(
    p1_131_cost_app,
) -> None:
    app, factory = p1_131_cost_app
    _enable_router(app)
    customer_id, statement_id = _seed_payable_statement(factory)
    with TestClient(app) as client:
        _login(client)
        created = client.post(
            "/api/finance/simple-finance/acceptances",
            json={
                "bill_number": "ACCEPT-202608-001",
                "customer_id": customer_id,
                "amount": "600.00",
                "received_date": "2026-08-28",
                "maturity_date": "2027-02-28",
                "note": "客户承兑",
                "idempotency_key": "p1144-acceptance-create",
            },
        )
        assert created.status_code == 201, created.text
        assert created.json()["status"] == "held"

        metadata = client.get("/api/finance/simple-finance/metadata")
        assert metadata.status_code == 200, metadata.text
        option = next(
            row for row in metadata.json()["supplier_statements"] if row["id"] == statement_id
        )
        assert Decimal(str(option["invoice_available_amount"])) == Decimal("1000.00")

        endorsed = client.post(
            f"/api/finance/simple-finance/acceptances/{created.json()['id']}/endorse",
            json={
                "expected_version": 1,
                "supplier_statement_id": statement_id,
                "endorsement_date": "2026-09-01",
                "idempotency_key": "p1144-acceptance-endorse",
            },
        )
        assert endorsed.status_code == 200, endorsed.text
        assert endorsed.json()["acceptance"]["status"] == "endorsed"
        assert endorsed.json()["supplier_statement"]["status"] == "partial_payment"
        assert Decimal(str(endorsed.json()["supplier_statement"]["paid_amount"])) == Decimal(
            "600.00"
        )
        assert "不是银行现金付款" in endorsed.json()["message"]

        blocked = client.post(
            f"/api/finance/simple-finance/acceptances/{created.json()['id']}/transition",
            json={
                "expected_version": 2,
                "action": "returned",
                "idempotency_key": "p1144-acceptance-return-blocked",
            },
        )
        assert blocked.status_code == 409

        summary = client.get(
            "/api/finance/simple-finance/summary", params={"month": "2026-08"}
        )
        assert summary.status_code == 200, summary.text
        material = summary.json()["primary_groups"]["material"]
        assert Decimal(str(material["erp_amount"])) == Decimal("1000.00")
        assert Decimal(str(material["invoice_amount"])) == Decimal("1000.00")
        assert Decimal(str(material["statement_difference_amount"])) == Decimal("50.00")
        assert Decimal(str(material["invoice_difference_amount"])) == Decimal("0.00")
        assert Decimal(
            str(material["differences"][0]["invoice_difference_amount"])
        ) == Decimal("0.00")
        assert material["differences"][0]["adjustments"][0]["note"] == "供应商调价"

        exported = client.get(
            "/api/finance/simple-finance/material-export", params={"month": "2026-08"}
        )
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content), data_only=True)
        assert workbook.sheetnames == [
            "供应商核票汇总",
            "纸板与外购包材明细",
            "差异明细",
        ]
        assert workbook["供应商核票汇总"]["A2"].value == "上游纸板厂"
        assert workbook["供应商核票汇总"]["J1"].value == "已结算（含承兑）"
        assert workbook["供应商核票汇总"]["J2"].value == 600
        assert workbook["供应商核票汇总"]["K2"].value == 0
        assert workbook["供应商核票汇总"]["L2"].value == 600
        difference_labels = {
            row[2]
            for row in workbook["差异明细"].iter_rows(min_row=2, values_only=True)
            if row[2]
        }
        assert "供应商账单与ERP调整后金额" in difference_labels
        assert "价格差异" in difference_labels

    from app.models.finance_simplified import FinanceAcceptanceNote
    from app.models.supplier_settlement import SupplierMonthlyPayment

    with factory() as db:
        note = db.query(FinanceAcceptanceNote).filter_by(
            bill_number="ACCEPT-202608-001"
        ).one()
        payment = db.get(SupplierMonthlyPayment, note.supplier_payment_id)
        assert payment.payment_method == "acceptance"
        assert payment.acceptance_note_id == note.id
        assert payment.reference == "承兑背书 ACCEPT-202608-001"
