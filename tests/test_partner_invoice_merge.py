from decimal import Decimal
from io import BytesIO
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from test_fin001_invoice_tasks import fin001_app, _login, _complete_invoice_profile


def prepare_partner(client, factory, *, create_tasks=True):
    from app.models.customer import Customer
    from app.models.delivery import Delivery, DeliveryItem
    from app.models.finance import ReturnReceipt, ReturnReceiptItem, Statement, StatementItem
    from app.models.invoice_task import CustomerInvoiceProfile, CustomerInvoiceItemRule, FinanceSettlementEntity

    seller_id = _complete_invoice_profile(client)
    with factory() as db:
        child = Customer(customer_code='PARTNER-CHILD-2', name='合作子客户二')
        db.add(child)
        db.flush()
        entity = FinanceSettlementEntity(entity_code='MERGE-PARTNER', entity_name='天美德风',
            tax_no='913200000000000088', default_seller_id=seller_id,
            is_enabled=True, confirmation_status='confirmed')
        db.add(entity)
        db.flush()
        profile = db.scalar(select(CustomerInvoiceProfile).where(CustomerInvoiceProfile.customer_id == 1))
        profile.settlement_entity_id = entity.id
        db.add(CustomerInvoiceProfile(customer_id=child.id, settlement_entity_id=entity.id,
            default_seller_id=seller_id, invoice_title='天美德风', tax_no=entity.tax_no,
            default_tax_rate=Decimal('.13'), confirmation_status='confirmed', is_enabled=True))
        rule = db.scalar(select(CustomerInvoiceItemRule).where(CustomerInvoiceItemRule.customer_id == 1))
        db.add(CustomerInvoiceItemRule(customer_id=child.id, project_name=rule.project_name,
            tax_classification_code=rule.tax_classification_code, unit=rule.unit, tax_rate=rule.tax_rate,
            spec_source=rule.spec_source, fill_unit_price=rule.fill_unit_price, confirmation_status='confirmed'))
        original = db.get(Statement, 1)
        original.settlement_entity_id = entity.id
        original.settlement_name_snapshot = entity.entity_name
        original.settlement_customer_ids_snapshot_json = '[1]'
        original.confirmation_status = 'confirmed'
        original.version = 2
        db.scalar(select(StatementItem).where(StatementItem.statement_id == 1)).source_customer_id = 1
        delivery = Delivery(delivery_number='PARTNER-D2', customer_id=child.id,
            delivery_date=original.created_at.date(), status='dispatched', total_quantity=20)
        db.add(delivery)
        db.flush()
        item = DeliveryItem(delivery_id=delivery.id, source_type='unordered_finished', product_id=1,
            delivered_quantity=20, product_code_snapshot='PARTNER-BOX-2', product_name_snapshot='合作纸箱二',
            unit_snapshot='个', unit_price_snapshot=Decimal('11.30'), price_source='manual')
        db.add(item)
        db.flush()
        receipt = ReturnReceipt(delivery_id=delivery.id, actual_received_date=delivery.delivery_date, status='confirmed')
        db.add(receipt)
        db.flush()
        receipt_item = ReturnReceiptItem(return_receipt_id=receipt.id, delivery_item_id=item.id, actual_received_quantity=20)
        db.add(receipt_item)
        db.flush()
        statement = Statement(statement_number='ST-PARTNER-2', customer_id=child.id,
            settlement_entity_id=entity.id, settlement_name_snapshot=entity.entity_name,
            settlement_customer_ids_snapshot_json=json.dumps([child.id]), statement_month='2026-08',
            total_receivable=Decimal('226'), total_gross_profit=0, confirmation_status='confirmed', version=2)
        db.add(statement)
        db.flush()
        db.add(StatementItem(statement_id=statement.id, source_customer_id=child.id,
            return_receipt_item_id=receipt_item.id, actual_received_quantity=20,
            unit_price_snapshot=Decimal('11.30'), unit_cost_snapshot=Decimal('1'),
            receivable_amount=Decimal('226'), gross_profit_amount=0,
            price_tax_mode_snapshot='tax_inclusive', tax_rate_snapshot=Decimal('.13')))
        db.commit()
    tasks = []
    if not create_tasks:
        return tasks
    for sid in (1, 2):
        result = client.post(f'/api/finance/statements/{sid}/invoice-tasks',
            json={'expected_version':2,'idempotency_key':f'partner-invoice-{sid}'})
        assert result.status_code == 201, result.text
        tasks.append(result.json())
    return tasks


def test_two_separate_partner_statements_become_one_invoice_file(fin001_app):
    from app.models.finance import Invoice, Statement
    from app.models.invoice_task import FinanceInvoiceTask
    app, factory = fin001_app
    with TestClient(app) as client:
        _login(client)
        tasks = prepare_partner(client, factory)
        preview = client.get('/api/finance/statements/1/partner-invoice-preview')
        assert preview.status_code == 200, preview.text
        data = preview.json()
        assert len(data['statements']) == 2
        assert Decimal(str(data['total_amount'])) == 339
        payload = {'scope_hash':data['scope_hash'], 'idempotency_key':'merge-two-partner-bills'}
        merged = client.post('/api/finance/statements/1/partner-invoice-tasks', json=payload)
        assert merged.status_code == 201, merged.text
        task = merged.json()
        assert len(task['source_statements']) == 2
        assert Decimal(str(task['total_amount'])) == 339
        replay = client.post('/api/finance/statements/1/partner-invoice-tasks', json=payload)
        assert replay.status_code == 201 and replay.json()['id'] == task['id']
        exported = client.get(f"/api/finance/invoice-tasks/{task['id']}/tax-template.xlsx")
        assert exported.status_code == 200, exported.text
        workbook = load_workbook(BytesIO(exported.content))
        sheet = workbook.active
        quantities = [row[4] for row in sheet.iter_rows(min_row=4, values_only=True) if len(row)>4 and isinstance(row[4], (int,float))]
        assert sorted(quantities) == [10,20]
        task = client.get(f"/api/finance/invoice-tasks/{task['id']}").json()
        result_payload = {'status':'issued','invoice_number':'PARTNER-ONE-INVOICE',
            'invoice_date':'2026-09-22','expected_version':task['version'],
            'expected_ledger_version':task['ledger_version'],
            'expected_ledger_versions':{str(s['statement_id']):s['ledger_version'] for s in task['source_statements']}}
        issued = client.post(f"/api/finance/invoice-tasks/{task['id']}/result",json=result_payload)
        assert issued.status_code == 200, issued.text
        assert client.post(f"/api/finance/invoice-tasks/{task['id']}/result",json=result_payload).status_code == 200
        for sid, amount in [(1,113),(2,226)]:
            detail = client.get(f'/api/finance/statements/{sid}').json()
            assert Decimal(str(detail['invoiced_amount'])) == amount
            assert detail['invoices'][0]['invoice_number'] == 'PARTNER-ONE-INVOICE'
            assert Decimal(str(detail['invoices'][0]['invoice_amount'])) == amount
        with factory() as db:
            assert db.query(Invoice).count() == 1
            assert Decimal(str(db.scalar(select(Invoice)).invoice_amount)) == 339
            assert all(db.get(FinanceInvoiceTask,t['id']).status == 'voided' for t in tasks)
            assert all(db.get(Statement,sid).version == 2 for sid in (1,2))


def merge(client, key='merge-partner-test'):
    preview = client.get('/api/finance/statements/1/partner-invoice-preview')
    assert preview.status_code == 200, preview.text
    return client.post('/api/finance/statements/1/partner-invoice-tasks',
        json={'scope_hash':preview.json()['scope_hash'], 'idempotency_key':key})


def issue_payload(task):
    return {'status':'issued','invoice_number':'ONE-PARTNER-INVOICE','invoice_date':'2026-09-22',
        'expected_version':task['version'],'expected_ledger_version':task['ledger_version'],
        'expected_ledger_versions':{str(s['statement_id']):s['ledger_version'] for s in task['source_statements']}}


def test_missing_tasks_are_prepared_once_and_old_download_is_blocked(fin001_app):
    from app.models.invoice_task import FinanceInvoiceTask
    app, factory = fin001_app
    with TestClient(app) as client:
        _login(client)
        prepare_partner(client, factory, create_tasks=False)
        response = merge(client)
        assert response.status_code == 201, response.text
        task = response.json()
        assert len(task['source_statements']) == 2
        retry = merge(client, 'second-click-new-key')
        assert retry.status_code == 201 and retry.json()['id'] == task['id']
        for sid in (1,2):
            reuse = client.post(f'/api/finance/statements/{sid}/invoice-tasks',
                json={'expected_version':2, 'idempotency_key':f'recreate-{sid}'})
            assert reuse.status_code == 201 and reuse.json()['id'] == task['id']
        with factory() as db:
            assert db.query(FinanceInvoiceTask).count() == 3
            old = db.scalars(select(FinanceInvoiceTask).where(FinanceInvoiceTask.status == 'voided')).all()
            for row in old:
                assert client.get(f'/api/finance/invoice-tasks/{row.id}/tax-template.xlsx').status_code == 409


@pytest.mark.parametrize('change', ['draft','ledger','month','buyer','partial'])
def test_preflight_or_concurrent_scope_change_never_silently_omits_member(fin001_app, change):
    from app.models.finance import Statement
    from app.models.invoice_task import FinanceInvoiceTask
    app, factory = fin001_app
    with TestClient(app) as client:
        _login(client)
        tasks = prepare_partner(client,factory)
        preview = client.get('/api/finance/statements/1/partner-invoice-preview').json()
        with factory() as db:
            second = db.get(Statement,2)
            if change == 'draft': second.confirmation_status = 'draft'
            if change == 'ledger': second.ledger_version += 1
            if change == 'month': second.statement_month = '2026-09'
            if change == 'partial': second.invoiced_amount = 1
            if change == 'buyer':
                task = db.get(FinanceInvoiceTask,tasks[1]['id'])
                buyer = json.loads(task.buyer_snapshot_json)
                buyer['tax_no'] = 'OTHER-TAX-NO'
                task.buyer_snapshot_json = json.dumps(buyer)
            db.commit()
        result = client.post('/api/finance/statements/1/partner-invoice-tasks',
            json={'scope_hash':preview['scope_hash'],'idempotency_key':'stale-preview-merge'})
        assert result.status_code == 409, result.text
        with factory() as db:
            assert db.query(FinanceInvoiceTask).count() == 2
            assert all(db.get(FinanceInvoiceTask,t['id']).status=='draft' for t in tasks)


def test_incomplete_single_download_is_rejected_and_failure_can_retry(fin001_app,monkeypatch):
    from app.api import invoice_tasks as api
    app, factory = fin001_app
    with TestClient(app) as client:
        _login(client)
        tasks = prepare_partner(client,factory)
        ready = client.post(f"/api/finance/invoice-tasks/{tasks[0]['id']}/confirm",json={'expected_version':1})
        assert ready.status_code == 200
        assert client.get(f"/api/finance/invoice-tasks/{tasks[0]['id']}/tax-template.xlsx").status_code == 409
        task = merge(client).json()
        real_export = api.generate_invoice_tax_template
        def fail(**kwargs): raise api.InvoiceTaxTemplateError('isolated export failure')
        monkeypatch.setattr(api,'generate_invoice_tax_template',fail)
        assert client.get(f"/api/finance/invoice-tasks/{task['id']}/tax-template.xlsx").status_code == 409
        assert client.get(f"/api/finance/invoice-tasks/{task['id']}").json()['status'] == 'ready'
        monkeypatch.setattr(api,'generate_invoice_tax_template',real_export)
        assert client.get(f"/api/finance/invoice-tasks/{task['id']}/tax-template.xlsx").status_code == 200


@pytest.mark.parametrize('failure', ['ledger','audit','duplicate'])
def test_invoice_registration_is_atomic_across_all_statements(fin001_app,monkeypatch,failure):
    from app.api import invoice_tasks as api
    from app.models.finance import Invoice, Statement
    from app.models.invoice_task import FinanceInvoiceTask
    from datetime import date
    app,factory=fin001_app
    with TestClient(app,raise_server_exceptions=False) as client:
        _login(client)
        prepare_partner(client,factory)
        task=merge(client).json()
        assert client.get(f"/api/finance/invoice-tasks/{task['id']}/tax-template.xlsx").status_code==200
        task=client.get(f"/api/finance/invoice-tasks/{task['id']}").json()
        payload=issue_payload(task)
        original_audit=api._audit
        if failure=='ledger':
            with factory() as db:
                db.get(Statement,2).ledger_version+=1
                db.commit()
        if failure=='audit':
            def fail(*args,**kwargs): raise RuntimeError('audit unavailable')
            monkeypatch.setattr(api,'_audit',fail)
        if failure=='duplicate':
            with factory() as db:
                db.add(Invoice(statement_id=1,invoice_number=payload['invoice_number'],invoice_date=date(2026,9,22),invoice_amount=1))
                db.commit()
        response=client.post(f"/api/finance/invoice-tasks/{task['id']}/result",json=payload)
        assert response.status_code==(500 if failure=='audit' else 409),response.text
        with factory() as db:
            assert all(db.get(Statement,sid).invoiced_amount==0 for sid in (1,2))
            assert db.get(FinanceInvoiceTask,task['id']).status=='exported'
        monkeypatch.setattr(api,'_audit',original_audit)
        current=client.get(f"/api/finance/invoice-tasks/{task['id']}").json()
        retry=issue_payload(current)
        retry['invoice_number']='NEW-UNIQUE-NUMBER'
        assert client.post(f"/api/finance/invoice-tasks/{task['id']}/result",json=retry).status_code==200
        assert client.post(f"/api/finance/invoice-tasks/{task['id']}/void",json={'expected_version':2}).status_code==409
        for sid in (1,2):
            assert client.post(f'/api/finance/statements/{sid}/reopen',json={'expected_version':2,'reason':'不能更改真实已开票'}).status_code==409


def test_cancel_merge_then_modify_only_one_child_and_remerge(fin001_app):
    from app.models.finance import Statement, StatementItem
    app,factory=fin001_app
    with TestClient(app) as client:
        _login(client)
        prepare_partner(client,factory)
        task=merge(client).json()
        adjustment={'expected_version':2,'reason':'子客户一单价更正','update_lines':[{'statement_item_id':1,'unit_price':'12.50'}]}
        assert client.post('/api/finance/statements/1/adjust-dispute',json=adjustment).status_code==409
        result=client.post(f"/api/finance/invoice-tasks/{task['id']}/void",json={'expected_version':task['version']})
        assert result.status_code==200,result.text
        result=client.post('/api/finance/statements/1/adjust-dispute',json=adjustment)
        assert result.status_code==200,result.text
        with factory() as db:
            assert db.get(Statement,2).confirmation_status=='confirmed'
            assert db.get(Statement,2).version==2
            assert db.get(StatementItem,2).unit_price_snapshot==Decimal('11.30')
        assert client.post('/api/finance/statements/1/confirm',json={'expected_version':3}).status_code==200
        result=merge(client,'after-child-price-change')
        assert result.status_code==201,result.text
        assert Decimal(str(result.json()['total_amount']))==351


def test_two_simultaneous_merges_create_only_one_active_task(fin001_app):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from app.models.invoice_task import FinanceInvoiceTask
    app,factory=fin001_app
    with TestClient(app) as setup:
        _login(setup)
        prepare_partner(setup,factory)
        scope=setup.get('/api/finance/statements/1/partner-invoice-preview').json()['scope_hash']
    barrier=Barrier(2)
    def attempt():
        with TestClient(app) as client:
            _login(client)
            barrier.wait()
            return client.post('/api/finance/statements/1/partner-invoice-tasks',json={'scope_hash':scope,'idempotency_key':'concurrent-identical-merge'})
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses=list(pool.map(lambda _:attempt(), range(2)))
    assert [r.status_code for r in responses]==[201,201]
    assert len({r.json()['id'] for r in responses})==1
    with factory() as db:
        assert db.query(FinanceInvoiceTask).filter(FinanceInvoiceTask.status!='voided').count()==1


def test_all_member_permissions_and_frozen_membership_are_required(fin001_app):
    from app.models.access_control import UserCustomerScope
    from app.models.user import User
    from app.models.invoice_task import CustomerInvoiceProfile
    app,factory=fin001_app
    with TestClient(app) as client:
        _login(client)
        prepare_partner(client,factory)
        task=merge(client).json()
        with factory() as db:
            # Changing today's cooperation setting must not reinterpret existing bills.
            db.scalar(select(CustomerInvoiceProfile).where(CustomerInvoiceProfile.customer_id==2)).settlement_entity_id=None
            user=db.scalar(select(User).where(User.username=='fin001-finance'))
            user.customer_access_mode='selected'
            db.add(UserCustomerScope(user_id=user.id,customer_id=1))
            db.commit()
        assert client.get('/api/finance/statements/1/partner-invoice-preview').status_code==403
        assert client.get(f"/api/finance/invoice-tasks/{task['id']}").status_code==403
        assert all(t['id']!=task['id'] for t in client.get('/api/finance/invoice-tasks').json()['items'])
        _login(client,'fin001-admin')
        preview=client.get('/api/finance/statements/1/partner-invoice-preview')
        assert preview.status_code==200
        assert len(preview.json()['statements'])==2


def test_merge_audit_failure_rolls_back_superseded_tasks(fin001_app,monkeypatch):
    from app.api import invoice_tasks as api
    from app.models.invoice_task import FinanceInvoiceTask, FinanceInvoiceTaskStatement
    app,factory=fin001_app
    with TestClient(app,raise_server_exceptions=False) as client:
        _login(client)
        original=prepare_partner(client,factory)
        audit=api._audit
        def fail(*args,**kwargs):
            if kwargs['action']=='MERGE_PARTNER_INVOICE_TASKS': raise RuntimeError('isolated audit failure')
            return audit(*args,**kwargs)
        monkeypatch.setattr(api,'_audit',fail)
        assert merge(client).status_code==500
        with factory() as db:
            assert db.query(FinanceInvoiceTask).count()==2
            assert db.query(FinanceInvoiceTaskStatement).count()==0
            assert all(db.get(FinanceInvoiceTask,t['id']).status=='draft' for t in original)
