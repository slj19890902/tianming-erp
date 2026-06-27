from __future__ import annotations

from collections.abc import Generator
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker


@pytest.fixture()
def phase12_app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from app.api.auth import router as auth_router
    from app.api.customers import router as customers_router
    from app.api.deps import get_db
    from app.api.finance import router as finance_router
    from app.api.incoming import router as incoming_router
    from app.api.materials import router as materials_router
    from app.api.products import router as products_router
    from app.core.database import create_sqlite_engine
    from app.core.security import hash_password
    from app.models import Base
    from app.models.customer import Customer
    from app.models.delivery import Delivery, DeliveryItem
    from app.models.finance import (
        ReturnReceipt,
        ReturnReceiptItem,
        Statement,
        StatementItem,
    )
    from app.models.material import Material
    from app.models.order import Order, OrderItem
    from app.models.product import Product
    from app.models.user import User

    upload_dir = tmp_path / "uploads"
    monkeypatch.setenv("ERP_DRAWING_DIR", str(upload_dir))
    engine = create_sqlite_engine(tmp_path / "phase12.sqlite3")
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    with session_factory() as session:
        users = [
            User(
                username=role,
                password_hash=hash_password("RolePass123!"),
                role=role,
                real_name=role,
                display_name=role,
                must_change_password=False,
            )
            for role in ("admin", "finance", "sales", "workshop")
        ]
        active = Customer(
            customer_number=1,
            customer_code="ACTIVE",
            name="苏州正常客户",
            payment_term_days=30,
            credit_limit=0,
            delivery_method="配送",
        )
        inactive = Customer(
            customer_number=2,
            customer_code="STOP",
            name="昆山停用客户",
            payment_term_days=30,
            credit_limit=0,
            delivery_method="物流",
            status="inactive",
            is_active=False,
        )
        material = Material(
            code="WCX1",
            layer_count=5,
            flute_type="AB",
            basis_weight_description="供应商A 170g/130g/80g/170g/150g 高强",
        )
        session.add_all([*users, active, inactive, material])
        session.flush()
        product = Product(
            customer_id=active.id,
            product_code="21301028",
            customer_material_code="21301028",
            product_name="中性外箱",
            material_id=material.id,
            box_category="normal",
            cost_unit_price=Decimal("2.00"),
        )
        session.add(product)
        session.flush()
        order = Order(
            order_number="PO-20260614-001",
            customer_id=active.id,
            order_date=date(2026, 6, 14),
            delivery_date=date(2026, 6, 21),
            status="pending_production",
            payment_status="unpaid",
            total_amount=Decimal("300"),
        )
        session.add(order)
        session.flush()
        item = OrderItem(
            order_id=order.id,
            product_id=product.id,
            quantity=100,
            delivered_quantity=80,
            unit_price=Decimal("3.00"),
            subtotal=Decimal("300"),
            material_status="pending",
            requisition_status="已报料",
            snapshot_product_code="21301028",
            snapshot_product_name="中性外箱",
            snapshot_material="WCX1",
        )
        session.add(item)
        session.flush()
        delivery = Delivery(
            delivery_number="DH-20260614-001",
            customer_id=active.id,
            delivery_date=date(2026, 6, 14),
            status="dispatched",
            total_quantity=80,
            dispatched_at=datetime(2026, 6, 14, 9, 0),
        )
        session.add(delivery)
        session.flush()
        delivery_item = DeliveryItem(
            delivery_id=delivery.id,
            order_item_id=item.id,
            delivered_quantity=80,
        )
        session.add(delivery_item)
        session.flush()
        receipt = ReturnReceipt(
            delivery_id=delivery.id,
            actual_received_date=date(2026, 6, 14),
            signed_by="王经理",
            status="confirmed",
        )
        session.add(receipt)
        session.flush()
        receipt_item = ReturnReceiptItem(
            return_receipt_id=receipt.id,
            delivery_item_id=delivery_item.id,
            actual_received_quantity=78,
            difference_reason="压坏2个",
        )
        session.add(receipt_item)
        session.flush()
        statement = Statement(
            statement_number="ST-202606-001",
            customer_id=active.id,
            statement_month="2026-06",
            total_receivable=Decimal("234"),
            total_gross_profit=Decimal("78"),
            status="unsettled",
        )
        session.add(statement)
        session.flush()
        session.add(
            StatementItem(
                statement_id=statement.id,
                return_receipt_item_id=receipt_item.id,
                actual_received_quantity=78,
                unit_price_snapshot=Decimal("3.00"),
                unit_cost_snapshot=Decimal("2.00"),
                receivable_amount=Decimal("234"),
                gross_profit_amount=Decimal("78"),
            )
        )
        session.commit()

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(customers_router, prefix="/api/master/customers")
    app.include_router(materials_router, prefix="/api/master/materials")
    app.include_router(products_router, prefix="/api/master/products")
    app.include_router(incoming_router, prefix="/api/incoming")
    app.include_router(finance_router, prefix="/api/finance")

    def override_get_db() -> Generator[Session, None, None]:
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    return app, session_factory, upload_dir


def _login(client: TestClient, role: str = "admin") -> None:
    response = client.post(
        "/api/auth/login",
        json={"username": role, "password": "RolePass123!"},
    )
    assert response.status_code == 200


def test_customer_list_hides_inactive_and_delete_blocks_open_order(phase12_app):
    app, session_factory, _ = phase12_app
    with TestClient(app) as client:
        _login(client)
        visible = client.get("/api/master/customers")
        all_rows = client.get(
            "/api/master/customers",
            params={"include_inactive": True},
        )
        blocked = client.delete("/api/master/customers/1")
        removed = client.delete("/api/master/customers/2")

    assert visible.json()["total"] == 1
    assert all_rows.json()["total"] == 2
    assert blocked.status_code == 400
    assert removed.status_code == 204
    with session_factory() as session:
        customer = session.get(__import__("app.models.customer", fromlist=["Customer"]).Customer, 2)
        assert customer is not None
        assert customer.is_active is False


def test_admin_can_reenable_inactive_customer(phase12_app):
    from app.models.audit import OperationLog
    from app.models.customer import Customer

    app, session_factory, _ = phase12_app
    with TestClient(app) as client:
        _login(client)
        enabled = client.put(
            "/api/master/customers/2/status",
            json={"is_active": True},
        )
        visible = client.get("/api/master/customers")

    assert enabled.status_code == 200, enabled.text
    assert enabled.json()["is_active"] is True
    assert enabled.json()["status"] == "active"
    assert visible.json()["total"] == 2
    with session_factory() as session:
        assert session.get(Customer, 2).is_active is True
        actions = list(
            session.scalars(
                select(OperationLog.action).where(
                    OperationLog.resource == "Customer"
                )
            )
        )
    assert "ENABLE" in actions


def test_material_normalizes_weight_and_rejects_unknown_flute(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "admin")
        created = client.post(
            "/api/master/materials",
            json={
                "code": "TEST-AB",
                "layer_count": 5,
                "flute_type": "AB",
                "basis_weight_description": "嘉林亿170克/130g/80克/170g/150g",
            },
        )
        invalid = client.post(
            "/api/master/materials",
            json={"code": "BAD", "flute_type": "BC"},
        )

    assert created.status_code == 201, created.text
    assert created.json()["basis_weight_description"] == "170g/130g/80g/170g/150g"
    assert invalid.status_code == 422


def test_product_drawing_upload_saves_compressed_files_not_base64(phase12_app):
    app, _, upload_dir = phase12_app
    image = Image.new("RGB", (2400, 1600), "white")
    source = BytesIO()
    image.save(source, format="PNG")
    with TestClient(app) as client:
        _login(client, "admin")
        response = client.post(
            "/api/master/products/1/drawing",
            files={"file": ("drawing.png", source.getvalue(), "image/png")},
        )

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["drawing_path"].startswith("/static/uploads/drawings/")
    assert data["thumbnail_path"].startswith("/static/uploads/drawings/")
    files = list(upload_dir.glob("*"))
    assert len(files) == 2
    assert all(path.stat().st_size < len(source.getvalue()) for path in files)
    assert "base64" not in str(data).lower()


def test_wms_receive_synchronizes_requisition_status(phase12_app):
    from app.models.order import OrderItem

    app, session_factory, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "workshop")
        response = client.put("/api/incoming/receive/1")

    assert response.status_code == 200
    assert response.json()["material_status"] == "received"
    assert response.json()["requisition_status"] == "已入库"
    with session_factory() as session:
        assert session.get(OrderItem, 1).requisition_status == "已入库"


def test_statement_excel_export_is_valid_workbook(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    assert sheet["A1"].value == "月结对账单"
    assert "苏州正常客户" in sheet["A2"].value
    assert sheet.max_row >= 7


# ── v0.20.2 导出完善测试 ────────────────────────────────────────────────────────


def test_customer_abbr_extracts_short_name():
    from app.api.finance import _customer_abbr

    assert _customer_abbr("苏州天明包装有限公司") == "天明"
    assert _customer_abbr("苏州天华超净科技股份有限公司") == "天华"
    assert _customer_abbr("天华超净科技（苏州）有限公司") == "天华"
    assert _customer_abbr("昆山华诚电子有限公司") == "华诚"
    assert _customer_abbr("苏州思迈尔包装有限公司") == "思迈"
    assert _customer_abbr("上海威力科技有限公司") == "威力"
    assert _customer_abbr("正常两字") == "正常"


def test_customer_abbr_falls_back_to_ke_hu_when_empty():
    from app.api.finance import _customer_abbr

    # 苏州 + 有限公司 → 去掉后剩余无中文字符，应回退为"客户"
    assert _customer_abbr("苏州有限公司") == "客户"


def test_safe_filename_removes_illegal_chars():
    from app.api.finance import _safe_filename

    assert _safe_filename("天明2026/06对账单") == "天明202606对账单"
    assert _safe_filename('file*name?"<>|end') == "filenameend"
    assert _safe_filename("no\\backslash:here") == "nobackslashhere"


def test_export_content_disposition_uses_rfc5987_with_ascii_fallback(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    assert response.status_code == 200
    cd = response.headers["content-disposition"]
    assert 'filename="statement.xlsx"' in cd, f"ASCII fallback missing from: {cd}"
    assert "filename*=UTF-8''" in cd, f"RFC 5987 encoding missing from: {cd}"


def test_export_filename_uses_customer_abbr_and_month(phase12_app):
    # Fixture customer "苏州正常客户" → abbr "正常"; month "2026-06"
    # Expected filename: 正常2026-06对账单.xlsx
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    assert response.status_code == 200
    cd = response.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd
    encoded_part = cd.split("filename*=UTF-8''")[1].split(";")[0].strip()
    decoded = unquote(encoded_part)
    assert decoded == "正常2026-06对账单.xlsx", f"Got filename: {decoded}"


def test_export_excel_contains_all_required_columns(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    header_row = [cell.value for cell in sheet[4]]
    required = {
        "客户名称", "客户单号", "存货编码", "送货日期", "送货单号",
        "产品名称", "规格型号", "材质", "实际签收数量",
        "单价", "金额", "备注", "开票状态", "对账状态", "结清状态",
    }
    missing = required - set(header_row)
    assert not missing, f"缺少列：{missing}"


def test_export_uses_actual_received_quantity_not_delivered(phase12_app):
    # Fixture: delivered_quantity=80, actual_received_quantity=78
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[4]]
    qty_col = headers.index("实际签收数量") + 1
    assert sheet.cell(5, qty_col).value == 78, "应使用实际签收数量 78，不是送货数量 80"


def test_export_amount_equals_receivable_amount(phase12_app):
    # Fixture: actual_received_quantity=78, unit_price=3.00, receivable_amount=234
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[4]]
    amount_col = headers.index("金额") + 1
    assert float(sheet.cell(5, amount_col).value) == 234.0


def test_export_remarks_spec_material_correct(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[4]]
    remarks_col = headers.index("备注") + 1
    material_col = headers.index("材质") + 1
    assert sheet.cell(5, remarks_col).value == "压坏2个"
    assert sheet.cell(5, material_col).value == "WCX1"


def test_export_returns_404_for_nonexistent_statement(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/999/export")

    assert response.status_code == 404


def test_export_does_not_write_to_database(phase12_app):
    from app.models.audit import OperationLog

    app, session_factory, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        with session_factory() as session:
            before_count = session.query(OperationLog).count()
        client.get("/api/finance/statements/1/export")

    with session_factory() as session:
        after_count = session.query(OperationLog).count()

    assert after_count == before_count, "导出操作不应写入数据库"


def _get_header_col(sheet, name: str) -> int:
    headers = [cell.value for cell in sheet[4]]
    return headers.index(name) + 1


def test_export_product_code_uses_snapshot_not_current(phase12_app):
    # Fixture: OrderItem.snapshot_product_code="21301028"
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "存货编码")
    assert sheet.cell(5, col).value == "21301028"


def test_export_invoice_status_partial(phase12_app):
    from app.models.finance import Statement

    app, session_factory, _ = phase12_app
    with session_factory() as session:
        stmt = session.get(Statement, 1)
        stmt.invoiced_amount = Decimal("100")
        session.commit()

    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "开票状态")
    assert sheet.cell(5, col).value == "部分开票"


def test_export_invoice_status_full(phase12_app):
    from app.models.finance import Statement

    app, session_factory, _ = phase12_app
    with session_factory() as session:
        stmt = session.get(Statement, 1)
        stmt.invoiced_amount = Decimal("234")
        session.commit()

    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "开票状态")
    assert sheet.cell(5, col).value == "已开票"


def test_export_settlement_status_not_received(phase12_app):
    # Default fixture: settled_amount=0
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "结清状态")
    assert sheet.cell(5, col).value == "未收款"


def test_export_settlement_status_partial(phase12_app):
    from app.models.finance import Statement

    app, session_factory, _ = phase12_app
    with session_factory() as session:
        stmt = session.get(Statement, 1)
        stmt.settled_amount = Decimal("100")
        session.commit()

    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "结清状态")
    assert sheet.cell(5, col).value == "部分收款"


def test_export_settlement_status_full(phase12_app):
    from app.models.finance import Statement

    app, session_factory, _ = phase12_app
    with session_factory() as session:
        stmt = session.get(Statement, 1)
        stmt.settled_amount = Decimal("234")
        session.commit()

    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "结清状态")
    assert sheet.cell(5, col).value == "已结清"


def test_export_reconciliation_status_is_always_reconciled(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "对账状态")
    assert sheet.cell(5, col).value == "已对账"


def test_export_original_fields_not_lost(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    headers = {cell.value for cell in sheet[4]}
    for field in ("送货日期", "送货单号", "客户单号", "产品名称", "规格型号",
                  "实际签收数量", "单价", "金额", "备注"):
        assert field in headers, f"原有字段丢失：{field}"


def test_export_column_order_exact(phase12_app):
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    actual = [cell.value for cell in sheet[4]]
    expected = [
        "客户名称", "客户单号", "存货编码", "送货日期", "送货单号",
        "产品名称", "规格型号", "材质", "实际签收数量", "单价", "金额",
        "备注", "开票状态", "对账状态", "结清状态",
    ]
    assert actual == expected, f"列顺序不符\n预期：{expected}\n实际：{actual}"


def test_export_invoice_status_zero_is_not_invoiced(phase12_app):
    # Default fixture: invoiced_amount=0 (default)
    app, _, _ = phase12_app
    with TestClient(app) as client:
        _login(client, "finance")
        response = client.get("/api/finance/statements/1/export")

    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    col = _get_header_col(sheet, "开票状态")
    assert sheet.cell(5, col).value == "未开票"
