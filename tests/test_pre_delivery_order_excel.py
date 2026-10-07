from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.services.excel_document_import import ExcelImportError, parse_excel_document


def purchase_order_file(customer_code="YG", *, requested=700, ordered=600, po="PO-TEST"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "包装箱格式 (1)"
    sheet["A3"] = "■購買要求書"
    sheet["O2"], sheet["P2"] = "管理NO.:", po
    sheet["L4"] = date(2026, 9, 30)
    for column, value in enumerate(["NO", "契约NO", "品目号", "部品名", "品番/图番", "要求数", "单位", "LOTS", "LOTS2", "發注数", "币种", "单价", "金额", "YKE着" if customer_code == "YG" else "KEW着"], 1):
        sheet.cell(10, column, value)
    for column, value in enumerate([1, None, "80010631", "测试纸箱", "00632094", requested, "只", None, None, ordered, "CNY", 2.865, ordered * 2.865], 1):
        sheet.cell(11, column, value)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("customer_code", ["YG", "GY"])
def test_same_purchase_workbook_supports_order_and_pre_delivery(customer_code):
    content = purchase_order_file(customer_code)
    order = parse_excel_document(content, "同昌订单.xlsx", document_type="order")
    delivery = parse_excel_document(content, "同昌订单.xlsx", document_type="pre_delivery")
    assert order.customer_code == delivery.customer_code == customer_code
    assert order.source_hash == delivery.source_hash
    assert delivery.document_type == "pre_delivery"
    assert delivery.customer_po == delivery.rows[0].customer_po == "PO-TEST"
    assert delivery.rows[0].requested_quantity == 700
    assert delivery.rows[0].order_quantity == order.rows[0].order_quantity == 600
    assert delivery.rows[0].drawing_number == "00632094"
    assert (delivery.rows[0].source_row, delivery.rows[0].source_no) == (11, "1")
    assert delivery.rows[0].issues
    assert delivery.rows[0].cells["requested_quantity"].raw_value == 700
    assert delivery.rows[0].cells["requested_quantity"].coordinate == "F11"
    assert "预送数量" in delivery.rows[0].issues[0]


def test_purchase_order_customer_evidence_cannot_cross_or_guess_customer():
    with pytest.raises(ExcelImportError, match="不一致"):
        parse_excel_document(purchase_order_file("GY"), "研光.xlsx", document_type="pre_delivery", expected_customer_code="YG")
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(purchase_order_file()))
    book.active["A1"] = "光洋"
    buffer = BytesIO(); book.save(buffer)
    with pytest.raises(ExcelImportError, match="客户.*不唯一"):
        parse_excel_document(buffer.getvalue(), "订单.xlsx", document_type="pre_delivery")


def test_invalid_purchase_order_quantity_is_not_reinterpreted_as_delivery_layout():
    with pytest.raises(ExcelImportError, match="数量必须是正整数"):
        parse_excel_document(purchase_order_file(ordered=-1), "订单.xlsx", document_type="pre_delivery")
