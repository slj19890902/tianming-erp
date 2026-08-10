from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture()
def product_workbook_app(tmp_path, monkeypatch):
    from app.api.auth import router as auth_router
    from app.api.deps import get_db
    from app.api.product_import import router as product_import_router
    from app.core.database import create_sqlite_engine
    from app.core.security import hash_password
    from app.models import Base
    from app.models.customer import Customer
    from app.models.material import Material
    from app.models.supplier import Supplier
    from app.models.user import User

    monkeypatch.setenv("ERP_FILE_STORAGE_DIR", str(tmp_path / "private_uploads"))
    monkeypatch.setenv("ERP_UPLOAD_TEMP_DIR", str(tmp_path / "temporary_uploads"))
    engine = create_sqlite_engine(tmp_path / "product-workbook.sqlite3")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add_all(
            [
                User(
                    username="product-import-admin",
                    password_hash=hash_password("WorkbookPass123!"),
                    role="admin",
                    real_name="常用箱管理员",
                    must_change_password=False,
                ),
                User(
                    username="product-import-sales",
                    password_hash=hash_password("WorkbookPass123!"),
                    role="sales",
                    real_name="业务",
                    must_change_password=False,
                ),
                Customer(
                    customer_number=9001,
                    customer_code="P1-22-CUSTOMER",
                    name="P1-22测试客户",
                ),
                Supplier(
                    standard_name="测试供应商",
                    normalized_name="测试供应商",
                    display_name="测试供应商",
                    is_active=True,
                ),
                Material(
                    code="C4C",
                    supplier_name="测试供应商",
                    layer_count=3,
                    basis_weight_description="80g/100g/80g",
                    is_active=True,
                    version=1,
                ),
            ]
        )
        db.commit()

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(product_import_router, prefix="/api/master/products")

    def override_get_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    app.state.session_factory = factory
    with factory() as db:
        app.state.customer_id = db.scalar(
            select(Customer.id).where(Customer.customer_code == "P1-22-CUSTOMER")
        )
    yield app
    engine.dispose()


def _login(client: TestClient, username: str) -> None:
    response = client.post(
        "/api/auth/login",
        json={"username": username, "password": "WorkbookPass123!"},
    )
    assert response.status_code == 200


def _download(client: TestClient, customer_id: int) -> bytes:
    response = client.get(
        "/api/master/products/import-template.xlsx",
        params={"customer_id": customer_id},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(EXCEL_MIME)
    return response.content


def _filled_workbook(content: bytes, *, drawing_filename: str | None = None) -> bytes:
    auto_embed = drawing_filename is None
    workbook = load_workbook(BytesIO(content))
    try:
        product = workbook["样品录入"]
        product["B2"] = "ERP-BOX-001"
        product["C2"] = "390×225×330"
        product["D2"] = "B"
        product["E2"] = "模切"
        product["F2"] = "是"
        product["G2"] = "黑色"
        product["H2"] = "粘合"
        product["I2"] = "否"
        product["J2"] = (
            "ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg"
            if auto_embed
            else drawing_filename
        )
        product["K2"] = "MOLD-P1-22-001"
        product["M2"] = "新增"
        product["P2"] = "CUSTOMER-BOX-001"
        product["Q2"] = "测试模切常用箱"
        product["R2"] = "C4C"
        product["T2"] = 3.25
        mold = workbook["模具档案"]
        mold["A2"] = "新增"
        mold["B2"] = "MOLD-P1-22-001"
        mold["C2"] = "测试模切版"
        mold["D2"] = "模具架-A01"
        output = BytesIO()
        workbook.save(output)
        result = output.getvalue()
        return _with_embedded_drawings(result) if auto_embed else result
    finally:
        workbook.close()


def _jpeg_bytes(color: str) -> bytes:
    from PIL import Image

    image = Image.new("RGB", (120, 80), color)
    output = BytesIO()
    image.save(output, format="JPEG", quality=78)
    return output.getvalue()


def _with_embedded_drawings(
    content: bytes,
    *,
    anchors: tuple[str, ...] = ("J2", "J2"),
    colors: tuple[str, ...] | None = None,
) -> bytes:
    workbook = load_workbook(BytesIO(content))
    streams: list[BytesIO] = []
    try:
        sheet = workbook["样品录入"]
        for index, anchor in enumerate(anchors):
            color = colors[index] if colors is not None else ("white" if index % 2 == 0 else "gray")
            stream = BytesIO(_jpeg_bytes(color))
            streams.append(stream)
            image = ExcelImage(stream)
            image.width = 120
            image.height = 80
            sheet.add_image(image, anchor)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
        for stream in streams:
            stream.close()


def test_multiple_volumes_merge_repeated_rows_and_cross_volume_images(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        template = _filled_workbook(
            _download(client, customer_id),
            drawing_filename="IMG_2001.jpg；IMG_2002.jpg",
        )
        volume_one = _with_embedded_drawings(
            template,
            anchors=("J2",),
            colors=("white",),
        )
        volume_two = _with_embedded_drawings(
            template,
            anchors=("J2",),
            colors=("gray",),
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files=[
                ("files", ("第001卷.xlsx", volume_one, EXCEL_MIME)),
                ("files", ("第002卷.xlsx", volume_two, EXCEL_MIME)),
            ],
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is True, payload
        assert payload["summary"]["workbooks"] == 2
        assert payload["summary"]["create_products"] == 1
        assert payload["summary"]["drawings"] == 2


def test_repeated_auto_import_reuses_product_mold_and_skips_duplicate_images(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.product import Product
    from app.models.product_drawing import ProductDrawing

    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _filled_workbook(_download(client, customer_id))
        workbook = load_workbook(BytesIO(content))
        try:
            workbook["样品录入"]["M2"] = "自动"
            stream = BytesIO()
            workbook.save(stream)
            content = stream.getvalue()
        finally:
            workbook.close()

        first_preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("首次.xlsx", content, EXCEL_MIME)},
        )
        assert first_preview.status_code == 200, first_preview.text
        assert first_preview.json()["valid"] is True, first_preview.json()
        first_apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": first_preview.json()["preview_token"]},
        )
        assert first_apply.status_code == 200, first_apply.text

        second_preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("重复上传.xlsx", content, EXCEL_MIME)},
        )
        assert second_preview.status_code == 200, second_preview.text
        second_payload = second_preview.json()
        assert second_payload["valid"] is True, second_payload
        assert second_payload["summary"]["drawing_only_products"] == 1, repr(
            second_payload["items"][0]["changed_fields"]
        )
        second_apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": second_payload["preview_token"]},
        )
        assert second_apply.status_code == 200, second_apply.text
        assert second_apply.json()["summary"]["drawings"] == 0
        assert second_apply.json()["summary"]["skipped_drawings"] == 2

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        product = db.scalar(select(Product).where(Product.product_code == "ERP-BOX-001"))
        assert product.version == 1
        assert db.scalar(
            select(func.count(ProductDrawing.id)).where(
                ProductDrawing.product_id == product.id
            )
        ) == 2


def test_multiple_volumes_block_conflicting_repeated_product_fields(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        first = _filled_workbook(_download(client, customer_id))
        workbook = load_workbook(BytesIO(first))
        try:
            workbook["样品录入"]["D2"] = "E"
            stream = BytesIO()
            workbook.save(stream)
            second = stream.getvalue()
        finally:
            workbook.close()
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files=[
                ("files", ("第001卷.xlsx", first, EXCEL_MIME)),
                ("files", ("第002卷.xlsx", second, EXCEL_MIME)),
            ],
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any("不同卷" in item["message"] for item in payload["errors"])


def test_duplicate_material_code_requires_supplier_qualification(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.material import Material
    from app.models.supplier import Supplier

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        db.add_all(
            [
                Supplier(
                    standard_name="另一供应商",
                    normalized_name="另一供应商",
                    display_name="另一供应商",
                    is_active=True,
                ),
                Material(
                    code="C4C",
                    supplier_name="另一供应商",
                    layer_count=3,
                    basis_weight_description="90g/110g/90g",
                    is_active=True,
                    version=1,
                ),
            ]
        )
        db.commit()

    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        ambiguous = _filled_workbook(_download(client, customer_id))
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("未指定供应商.xlsx", ambiguous, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any("供应商|材质代码" in item["message"] for item in payload["errors"])

        workbook = load_workbook(BytesIO(ambiguous))
        try:
            workbook["样品录入"]["R2"] = "测试供应商|C4C"
            stream = BytesIO()
            workbook.save(stream)
            qualified = stream.getvalue()
        finally:
            workbook.close()
        qualified_preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("已指定供应商.xlsx", qualified, EXCEL_MIME)},
        )
        assert qualified_preview.status_code == 200, qualified_preview.text
        assert qualified_preview.json()["valid"] is True, qualified_preview.json()


def _without_worksheet_dimensions(content: bytes) -> bytes:
    source = BytesIO(content)
    output = BytesIO()
    with ZipFile(source, "r") as input_zip, ZipFile(
        output,
        "w",
        compression=ZIP_DEFLATED,
    ) as output_zip:
        for item in input_zip.infolist():
            data = input_zip.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(
                ".xml"
            ):
                data = re.sub(br"<dimension ref=\"[^\"]+\"\s*/>", b"", data)
            output_zip.writestr(item, data)
    return output.getvalue()


def _rewrite_single_zip_member(
    content: bytes,
    *,
    member_pattern: str,
    transform,
) -> bytes:
    output = BytesIO()
    matched = 0
    with ZipFile(BytesIO(content), "r") as input_zip, ZipFile(
        output,
        "w",
        compression=ZIP_DEFLATED,
    ) as output_zip:
        for item in input_zip.infolist():
            data = input_zip.read(item.filename)
            if re.fullmatch(member_pattern, item.filename):
                matched += 1
                data = transform(data)
            output_zip.writestr(item, data)
    assert matched == 1
    return output.getvalue()


def test_template_is_customer_locked_printable_and_exportable(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-sales")
        content = _download(client, customer_id)
        workbook = load_workbook(BytesIO(content), read_only=False)
        try:
            assert workbook.sheetnames == [
                "使用说明",
                "样品录入",
                "模具档案",
                "导入信息",
            ]
            assert workbook["导入信息"].sheet_state == "hidden"
            assert workbook["导入信息"]["B2"].value == customer_id
            assert tuple(
                cell.value for cell in workbook["样品录入"][1]
            ) == (
                "样品号",
                    "手写型号*",
                "尺寸(mm)",
                "楞型",
                "成型方式*",
                "印刷*",
                "印刷颜色",
                "结合方式*",
                "二次粘合*",
                "图纸文件名(多个用分号)",
                "模具编号",
                "现场备注",
                "操作*",
                "系统ID",
                "当前版本",
                "客户料号*",
                "产品名称*",
                "材质代码",
                "单位",
                "默认含税单价",
                "报料长(mm)",
                "报料宽(mm)",
                "启用",
            )
            assert workbook["样品录入"].freeze_panes == "M2"
            assert workbook["样品录入"].max_row == 41
            assert workbook["样品录入"]["A2"].value == "YP001"
            assert workbook["样品录入"]["B2"].value is None
            assert workbook["样品录入"]["G2"].value == "黑色"
            assert workbook["样品录入"]["I2"].value == "否"
            assert (
                workbook["样品录入"]["J2"].value
                == "YP001_实物.jpg；YP001_展开.jpg"
            )
            assert workbook["样品录入"].column_dimensions["J"].width == 48
            assert workbook["样品录入"].row_dimensions[2].height == 88
            assert "一次选择多卷" in workbook["使用说明"]["B10"].value
            assert workbook["样品录入"].print_area == "'样品录入'!$A$1:$I$23"
        finally:
            workbook.close()


def test_preview_is_admin_only_and_requires_locked_customer(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-sales")
        content = _filled_workbook(_download(client, customer_id))
        forbidden = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("常用箱.xlsx", content, EXCEL_MIME)},
        )
        assert forbidden.status_code == 403

        _login(client, "product-import-admin")
        mismatch = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id + 1},
            files={"file": ("常用箱.xlsx", content, EXCEL_MIME)},
        )
        assert mismatch.status_code in {403, 404}


def test_preview_then_apply_creates_product_mold_and_version_atomically(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.audit import OperationLog
    from app.models.master_data_object_version import MasterDataObjectVersion
    from app.models.mold_tool import MoldTool
    from app.models.product import Product

    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _without_worksheet_dimensions(
            _filled_workbook(_download(client, customer_id))
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("常用箱.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True
        summary = preview.json()["summary"]
        assert summary["workbooks"] == 1
        assert summary["create_products"] == 1
        assert summary["update_products"] == 0
        assert summary["drawing_only_products"] == 0
        assert summary["create_molds"] == 1
        assert summary["drawings"] == 2
        item = preview.json()["items"][0]
        assert item["action"] == "新增"
        assert item["product_code"] == "ERP-BOX-001"
        assert item["missing_fields"] == ["报料长", "报料宽", "压线类型"]

        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 200
        assert apply.json()["ok"] is True

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        product = db.scalar(
            select(Product).where(Product.product_code == "ERP-BOX-001")
        )
        mold = db.scalar(
            select(MoldTool).where(MoldTool.mold_code == "MOLD-P1-22-001")
        )
        assert product is not None
        assert product.customer_id == customer_id
        assert product.customer_material_code == "CUSTOMER-BOX-001"
        assert product.material.code == "C4C"
        assert product.layer_count == 3
        assert product.flute_type == "B"
        assert product.length_mm == 390
        assert product.width_mm == 225
        assert product.height_mm == 330
        assert product.print_content == "单色印刷"
        assert product.printing_colors == "黑色"
        assert product.production_process == "模切、印刷、粘合"
        assert product.mold_tool_id == mold.id
        assert mold.rack_location == "模具架-A01"
        assert db.scalar(
            select(func.count(MasterDataObjectVersion.id)).where(
                MasterDataObjectVersion.object_type == "product",
                MasterDataObjectVersion.object_id == product.id,
            )
        ) == 1
        assert db.scalar(
            select(func.count(OperationLog.id)).where(
                OperationLog.resource == "ProductImport",
                OperationLog.action == "BATCH_IMPORT",
            )
        ) == 1


def test_apply_detects_post_preview_conflict_and_writes_nothing(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.product import Product

    customer_id = product_workbook_app.state.customer_id
    factory = product_workbook_app.state.session_factory
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _filled_workbook(_download(client, customer_id))
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("常用箱.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True

        with factory() as db:
            db.add(
                Product(
                    customer_id=customer_id,
                    product_code="ERP-BOX-001",
                    customer_material_code="OTHER-CODE",
                    product_name="测试模切常用箱",
                    box_category="normal",
                )
            )
            db.commit()

        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 409

    with factory() as db:
        assert db.scalar(
            select(func.count(Product.id)).where(
                Product.customer_material_code == "CUSTOMER-BOX-001"
            )
        ) == 0


def test_import_allows_same_customer_code_when_product_name_is_distinct(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.product import Product

    customer_id = product_workbook_app.state.customer_id
    factory = product_workbook_app.state.session_factory
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _filled_workbook(_download(client, customer_id))

        with factory() as db:
            db.add(
                Product(
                    customer_id=customer_id,
                    product_code="ERP-BOX-001",
                    customer_material_code="CUSTOMER-BOX-001",
                    product_name="同码外箱",
                    box_category="normal",
                )
            )
            db.commit()

        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("常用箱.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True
        applied = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert applied.status_code == 200, applied.text

    with factory() as db:
        rows = db.scalars(
            select(Product)
            .where(
                Product.customer_id == customer_id,
                Product.product_code == "ERP-BOX-001",
            )
            .order_by(Product.id)
        ).all()
        assert [row.product_name for row in rows] == ["同码外箱", "测试模切常用箱"]


def test_drawing_named_in_workbook_is_bound_to_imported_product(
    product_workbook_app: FastAPI,
) -> None:
    from PIL import Image

    from app.models.product import Product
    from app.models.product_drawing import ProductDrawing
    from app.services.secure_uploads import resolve_stored_reference

    customer_id = product_workbook_app.state.customer_id
    image = Image.new("RGB", (80, 60), "white")
    image_bytes = BytesIO()
    image.save(image_bytes, format="JPEG")
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _filled_workbook(
            _download(client, customer_id),
            drawing_filename="ERP-BOX-001_印刷图.jpg",
        )
        workbook = load_workbook(BytesIO(content))
        workbook["样品录入"]["L2"] = "图片要求=1张（已确认）"
        stream = BytesIO()
        workbook.save(stream)
        content = stream.getvalue()
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files=[
                ("file", ("常用箱.xlsx", content, EXCEL_MIME)),
                (
                    "drawings",
                    (
                        "ERP-BOX-001_印刷图.jpg",
                        image_bytes.getvalue(),
                        "image/jpeg",
                    ),
                ),
            ],
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True
        assert preview.json()["summary"]["drawings"] == 1
        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 200, apply.text

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        product = db.scalar(
            select(Product).where(Product.product_code == "ERP-BOX-001")
        )
        drawing = db.scalar(
            select(ProductDrawing).where(ProductDrawing.product_id == product.id)
        )
        assert drawing is not None
        assert resolve_stored_reference(drawing.image_path).is_file()
        assert resolve_stored_reference(drawing.thumbnail_path).is_file()


def test_two_embedded_drawings_are_bound_without_separate_upload(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.product import Product
    from app.models.product_drawing import ProductDrawing
    from app.services.secure_uploads import resolve_stored_reference

    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _with_embedded_drawings(
            _filled_workbook(
                _download(client, customer_id),
                drawing_filename="ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg",
            )
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("内嵌两图.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True, preview.text
        assert preview.json()["summary"]["drawings"] == 2
        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 200, apply.text

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        product = db.scalar(
            select(Product).where(Product.product_code == "ERP-BOX-001")
        )
        drawings = list(
            db.scalars(
                select(ProductDrawing)
                .where(ProductDrawing.product_id == product.id)
                .order_by(ProductDrawing.id)
            ).all()
        )
        assert len(drawings) == 2
        assert all(resolve_stored_reference(item.image_path).is_file() for item in drawings)
        assert all(
            resolve_stored_reference(item.thumbnail_path).is_file() for item in drawings
        )


def test_second_embedded_drawing_failure_rolls_back_database_and_saved_files(
    product_workbook_app: FastAPI,
    monkeypatch,
    tmp_path,
) -> None:
    from app.api import product_import
    from app.models.mold_tool import MoldTool
    from app.models.product import Product
    from app.models.product_drawing import ProductDrawing

    real_save = product_import.save_product_drawing_files
    calls = 0

    def fail_on_second_save(**kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("simulated second drawing failure")
        return real_save(**kwargs)

    monkeypatch.setattr(
        product_import,
        "save_product_drawing_files",
        fail_on_second_save,
    )
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _with_embedded_drawings(
            _filled_workbook(
                _download(client, customer_id),
                drawing_filename="ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg",
            )
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("两图回滚.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True
        with pytest.raises(RuntimeError, match="second drawing failure"):
            client.post(
                "/api/master/products/import/apply",
                json={"preview_token": preview.json()["preview_token"]},
            )

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        assert db.scalar(select(func.count(Product.id))) == 0
        assert db.scalar(select(func.count(MoldTool.id))) == 0
        assert db.scalar(select(func.count(ProductDrawing.id))) == 0
    private_root = tmp_path / "private_uploads"
    assert not list((private_root / "drawings").glob("*.webp"))
    assert not list((private_root / "drawing_thumbnails").glob("*.webp"))


@pytest.mark.parametrize("anchors", [("J2",), ("J2", "J2", "J2")])
def test_embedded_mode_rejects_new_row_without_exactly_two_images(
    product_workbook_app: FastAPI,
    anchors: tuple[str, ...],
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _with_embedded_drawings(
            _filled_workbook(
                _download(client, customer_id),
                drawing_filename="ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg",
            ),
            anchors=anchors,
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("内嵌一图.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any(
            "新增行必须恰好2张" in item["message"]
            for item in payload["errors"]
        )


def test_embedded_mode_accepts_one_image_with_explicit_confirmed_marker(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _filled_workbook(
            _download(client, customer_id),
            drawing_filename="ERP-BOX-001_图1.jpg",
        )
        workbook = load_workbook(BytesIO(content))
        try:
            workbook["样品录入"]["L2"] = "图片要求=1张（已确认）"
            output = BytesIO()
            workbook.save(output)
            content = output.getvalue()
        finally:
            workbook.close()
        content = _with_embedded_drawings(content, anchors=("J2",))

        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("内嵌单图已确认.xlsx", content, EXCEL_MIME)},
        )

        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is True, payload["errors"]
        assert payload["summary"]["drawings"] == 1


def test_embedded_mode_rejects_image_anchored_outside_column_j(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _with_embedded_drawings(
            _filled_workbook(
                _download(client, customer_id),
                drawing_filename="ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg",
            ),
            anchors=("J2", "I2"),
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("锚点错列.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any(
            "起始锚点必须放在本行 J 列" in item["message"]
            for item in payload["errors"]
        )


def test_same_row_cannot_mix_embedded_and_external_drawings(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _with_embedded_drawings(
            _filled_workbook(
                _download(client, customer_id),
                drawing_filename="ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg",
            )
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files=[
                ("file", ("混用.xlsx", content, EXCEL_MIME)),
                (
                    "drawings",
                    (
                        "ERP-BOX-001_实物.jpg",
                        _jpeg_bytes("white"),
                        "image/jpeg",
                    ),
                ),
            ],
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any(
            "同一行不能混用" in item["message"]
            for item in payload["errors"]
        )


def test_embedded_image_pixel_limit_blocks_before_drawing_storage(
    product_workbook_app: FastAPI,
    monkeypatch,
) -> None:
    from app.services import product_import_workbook

    monkeypatch.setattr(
        product_import_workbook,
        "MAX_EMBEDDED_IMAGE_PIXELS",
        100,
    )
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _with_embedded_drawings(
            _filled_workbook(
                _download(client, customer_id),
                drawing_filename="ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg",
            )
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("像素过大.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any(
            "像素尺寸过大" in item["message"]
            for item in payload["errors"]
        )


@pytest.mark.parametrize(
    ("member_pattern", "transform", "expected_message"),
    [
        (
            r"xl/drawings/_rels/drawing[0-9]+\.xml\.rels",
            lambda data: re.sub(
                br'Target="[^"]*media/[^"]+"',
                b'Target="https://example.invalid/image.jpg" TargetMode="External"',
                data,
                count=1,
            ),
            "外部链接",
        ),
        (
            r"xl/drawings/drawing[0-9]+\.xml",
            lambda data: re.sub(
                br"(</(?:xdr:)?wsDr>)",
                br"<sp/>\1",
                data,
                count=1,
            ),
            "只允许普通图片",
        ),
    ],
)
def test_embedded_drawing_container_rejects_active_relationships_and_shapes(
    product_workbook_app: FastAPI,
    member_pattern: str,
    transform,
    expected_message: str,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        content = _with_embedded_drawings(
            _filled_workbook(
                _download(client, customer_id),
                drawing_filename="ERP-BOX-001_实物.jpg；ERP-BOX-001_展开.jpg",
            )
        )
        malicious = _rewrite_single_zip_member(
            content,
            member_pattern=member_pattern,
            transform=transform,
        )
        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("不安全内嵌图.xlsx", malicious, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any(expected_message in item["message"] for item in payload["errors"])


def test_preview_derives_black_printing_and_two_dimension_liner(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.product import Product

    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        workbook = load_workbook(BytesIO(_download(client, customer_id)))
        try:
            sheet = workbook["样品录入"]
            sheet["B2"] = "LINER-HAND-001"
            sheet["C2"] = "299×209mm"
            sheet["D2"] = "E"
            sheet["E2"] = "无需"
            sheet["F2"] = "是"
            sheet["G2"] = ""
            sheet["H2"] = "无需结合"
            sheet["I2"] = "否"
            sheet["J2"] = "LINER-HAND-001_实物.jpg；LINER-HAND-001_展开.jpg"
            sheet["M2"] = "新增"
            sheet["P2"] = "CUSTOMER-LINER-001"
            sheet["Q2"] = "测试衬板"
            output = BytesIO()
            workbook.save(output)
        finally:
            workbook.close()

            content = _with_embedded_drawings(output.getvalue())
            preview = client.post(
                "/api/master/products/import/preview",
                data={"customer_id": customer_id},
                files={"file": ("衬板.xlsx", content, EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True
        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 200, apply.text

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        product = db.scalar(
            select(Product).where(Product.product_code == "LINER-HAND-001")
        )
        assert product is not None
        assert product.customer_material_code == "CUSTOMER-LINER-001"
        assert (product.length_mm, product.width_mm, product.height_mm) == (
            299,
            209,
            None,
        )
        assert (product.flute_type, product.layer_count) == ("E", 3)
        assert product.printing_colors == "黑色"
        assert product.production_process == "印刷"


def test_preview_rejects_secondary_gluing_without_die_cut_gluing(
    product_workbook_app: FastAPI,
) -> None:
    customer_id = product_workbook_app.state.customer_id
    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        workbook = load_workbook(
            BytesIO(_filled_workbook(_download(client, customer_id)))
        )
        try:
            sheet = workbook["样品录入"]
            sheet["E2"] = "开槽"
            sheet["H2"] = "打钉"
            sheet["I2"] = "是"
            sheet["K2"] = ""
            output = BytesIO()
            workbook.save(output)
        finally:
            workbook.close()

        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("错误二次粘合.xlsx", output.getvalue(), EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        payload = preview.json()
        assert payload["valid"] is False
        assert any(
            "二次粘合=是时" in item["message"]
            for item in payload["errors"]
        )


def test_update_preserves_unrepresented_process_and_crease_fields(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.product import Product

    customer_id = product_workbook_app.state.customer_id
    factory = product_workbook_app.state.session_factory
    with factory() as db:
        db.add(
            Product(
                customer_id=customer_id,
                product_code="HAND-EXISTING-001",
                customer_material_code="CUSTOMER-EXISTING-001",
                product_name="已有开槽纸箱",
                length_mm=400,
                width_mm=300,
                height_mm=200,
                box_category="normal",
                box_style="模切内盒",
                print_content="无印刷",
                printing_colors="黑色",
                production_process="覆膜、开槽、钉箱、粘箱",
                report_length_mm=800,
                report_width_mm=500,
                crease_type="压线",
                crease_left_mm=100,
                crease_middle_mm=300,
                crease_right_mm=100,
                layer_count=3,
                version=1,
            )
        )
        db.commit()

    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        workbook = load_workbook(BytesIO(_download(client, customer_id)))
        try:
            sheet = workbook["样品录入"]
            assert sheet["B2"].value == "HAND-EXISTING-001"
            assert sheet["D2"].value is None
            assert sheet["E2"].value == "开槽"
            assert sheet["F2"].value == "否"
            assert sheet["G2"].value is None
            assert sheet["H2"].value == "打钉"
            assert sheet["M2"].value == "自动"
            sheet["M2"] = "更新"
            sheet["T2"] = 4.5
            output = BytesIO()
            workbook.save(output)
        finally:
            workbook.close()

        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("更新已有.xlsx", output.getvalue(), EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True, preview.json()
        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 200, apply.text

    with factory() as db:
        product = db.scalar(
            select(Product).where(Product.product_code == "HAND-EXISTING-001")
        )
        assert product.sale_unit_price == pytest.approx(4.5)
        assert product.print_content == "无印刷"
        assert product.printing_colors == "黑色"
        assert product.production_process == "覆膜、开槽、钉箱、粘箱"
        assert product.layer_count == 3
        assert product.flute_type is None
        assert product.box_style is None
        assert product.crease_type == "压线"
        assert (
            product.crease_left_mm,
            product.crease_middle_mm,
            product.crease_right_mm,
        ) == (100, 300, 100)


def test_secondary_gluing_requires_die_cut_category_in_all_payload_writes() -> None:
    from pydantic import ValidationError

    from app.api.products import ProductPayload, _secondary_gluing_error

    with pytest.raises(ValidationError, match="箱类别必须为模切"):
        ProductPayload(
            customer_id=1,
            product_code="SECONDARY-NORMAL",
            customer_material_code="CUSTOMER-SECONDARY-NORMAL",
            product_name="错误二次粘合",
            box_category="normal",
            box_style="模切内盒",
            production_process="模切、粘合、二次粘合",
        )

    valid = ProductPayload(
        customer_id=1,
        product_code="SECONDARY-DIE-CUT",
        customer_material_code="CUSTOMER-SECONDARY-DIE-CUT",
        product_name="正确二次粘合",
        box_category="die_cut",
        box_style="模切内盒",
        production_process="模切、粘合、二次粘合",
    )
    assert valid.box_category == "die_cut"
    legacy_glue = ProductPayload(
        customer_id=1,
        product_code="SECONDARY-LEGACY-GLUE",
        customer_material_code="CUSTOMER-SECONDARY-LEGACY-GLUE",
        product_name="旧粘箱别名二次粘合",
        box_category="die_cut",
        box_style="模切内盒",
        production_process="模切、粘箱、二次粘合",
    )
    assert legacy_glue.box_style == "模切内盒"
    assert _secondary_gluing_error(
        production_process="模切、粘合、二次粘合",
        box_style="模切内盒",
        box_category="normal",
    )
    printed = ProductPayload(
        customer_id=1,
        product_code="PRINT-CONSISTENT",
        customer_material_code="CUSTOMER-PRINT-CONSISTENT",
        product_name="印刷一致性",
        box_category="normal",
        production_process="印刷",
        print_content="无印刷",
    )
    assert printed.print_content == "单色印刷"
    assert printed.printing_colors == "黑色"


def test_sync_fields_rejects_secondary_gluing_for_normal_product(
    product_workbook_app: FastAPI,
) -> None:
    from fastapi import HTTPException

    from app.api.products import SyncFieldsPayload, sync_product_fields
    from app.models.product import Product
    from app.models.user import User

    factory = product_workbook_app.state.session_factory
    with factory() as db:
        product = Product(
            customer_id=product_workbook_app.state.customer_id,
            product_code="SYNC-SECONDARY-NORMAL",
            customer_material_code="CUSTOMER-SYNC-SECONDARY-NORMAL",
            product_name="同步校验普通箱",
            box_category="normal",
            version=1,
        )
        db.add(product)
        db.commit()
        db.refresh(product)
        user = db.scalar(
            select(User).where(User.username == "product-import-admin")
        )
        with pytest.raises(HTTPException, match="箱类别必须为模切") as exc_info:
            sync_product_fields(
                product.id,
                SyncFieldsPayload(
                    fields={
                        "box_style": "模切内盒",
                        "production_process": "模切、粘合、二次粘合",
                    },
                    expected_version=product.version,
                    change_reason="测试二次粘合统一校验",
                ),
                db=db,
                user=user,
            )
        assert exc_info.value.status_code == 400
        db.refresh(product)
        assert product.production_process is None


def test_update_one_process_dimension_preserves_other_parallel_processes(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.product import Product

    customer_id = product_workbook_app.state.customer_id
    factory = product_workbook_app.state.session_factory
    with factory() as db:
        db.add(
            Product(
                customer_id=customer_id,
                product_code="PROCESS-DIMENSION-001",
                customer_material_code="CUSTOMER-PROCESS-DIMENSION-001",
                product_name="并行结合工艺纸箱",
                box_category="normal",
                print_content="无印刷",
                production_process="覆膜、开槽、钉箱、粘箱",
                version=1,
            )
        )
        db.commit()

    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        workbook = load_workbook(BytesIO(_download(client, customer_id)))
        try:
            sheet = workbook["样品录入"]
            assert sheet["E2"].value == "开槽"
            assert sheet["F2"].value == "否"
            assert sheet["H2"].value == "打钉"
            sheet["F2"] = "是"
            sheet["G2"] = "红色"
            sheet["M2"] = "更新"
            output = BytesIO()
            workbook.save(output)
        finally:
            workbook.close()

        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("只改印刷.xlsx", output.getvalue(), EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True, preview.json()
        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 200, apply.text

    with factory() as db:
        product = db.scalar(
            select(Product).where(
                Product.product_code == "PROCESS-DIMENSION-001"
            )
        )
        assert product.production_process == "覆膜、开槽、钉箱、粘箱、印刷"
        assert product.print_content == "单色印刷"
        assert product.printing_colors == "红色"


def test_secondary_gluing_with_parallel_nailing_round_trips_as_gluing(
    product_workbook_app: FastAPI,
) -> None:
    from app.models.mold_tool import MoldTool
    from app.models.product import Product

    customer_id = product_workbook_app.state.customer_id
    factory = product_workbook_app.state.session_factory
    with factory() as db:
        mold = MoldTool(
            mold_code="MOLD-ROUNDTRIP-001",
            mold_name="二次粘合回导模具",
            rack_location="模具架-RT01",
            is_active=True,
        )
        db.add(mold)
        db.flush()
        db.add(
            Product(
                customer_id=customer_id,
                product_code="SECONDARY-ROUNDTRIP-001",
                customer_material_code="CUSTOMER-SECONDARY-ROUNDTRIP-001",
                product_name="带打钉记录的二次粘合内盒",
                box_category="die_cut",
                box_style="模切内盒",
                production_process="模切、打钉、粘合、二次粘合",
                mold_tool_id=mold.id,
                version=1,
            )
        )
        db.commit()

    with TestClient(product_workbook_app) as client:
        _login(client, "product-import-admin")
        workbook = load_workbook(BytesIO(_download(client, customer_id)))
        try:
            sheet = workbook["样品录入"]
            assert sheet["E2"].value == "模切"
            assert sheet["H2"].value == "粘合"
            assert sheet["I2"].value == "是"
            sheet["M2"] = "更新"
            sheet["T2"] = 2.75
            output = BytesIO()
            workbook.save(output)
        finally:
            workbook.close()

        preview = client.post(
            "/api/master/products/import/preview",
            data={"customer_id": customer_id},
            files={"file": ("二次粘合回导.xlsx", output.getvalue(), EXCEL_MIME)},
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["valid"] is True, preview.json()
        apply = client.post(
            "/api/master/products/import/apply",
            json={"preview_token": preview.json()["preview_token"]},
        )
        assert apply.status_code == 200, apply.text

    with factory() as db:
        product = db.scalar(
            select(Product).where(
                Product.product_code == "SECONDARY-ROUNDTRIP-001"
            )
        )
        assert product.sale_unit_price == pytest.approx(2.75)
        assert product.production_process == "模切、打钉、粘合、二次粘合"
