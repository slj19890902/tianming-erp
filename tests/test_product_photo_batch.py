from __future__ import annotations

import csv
from io import BytesIO
import json
from pathlib import Path
import random
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image
import pytest

from app.services.product_import_workbook import (
    GUIDE_SHEET,
    INFO_SHEET,
    MOLD_HEADERS,
    MOLD_SHEET,
    PRODUCT_HEADERS,
    PRODUCT_SHEET,
    SINGLE_PHOTO_CONFIRMED_REMARK,
    TEMPLATE_VERSION,
    _preflight_product_xlsx_container,
)
from app.services.product_photo_batch import (
    ProductPhotoBatchError,
    build_product_photo_workbooks,
)


def _image_bytes(
    color: str,
    *,
    image_format: str = "JPEG",
    size: tuple[int, int] = (640, 480),
) -> bytes:
    image = Image.new("RGB", size, color)
    try:
        output = BytesIO()
        image.save(output, format=image_format, quality=92)
        return output.getvalue()
    finally:
        image.close()


def _write_image(
    path: Path,
    color: str | tuple[int, int, int],
    *,
    image_format: str | None = None,
    size: tuple[int, int] = (640, 480),
) -> None:
    image = Image.new("RGB", size, color)
    try:
        image.save(path, format=image_format, quality=92)
    finally:
        image.close()


def _write_noise_image(path: Path, seed: int) -> None:
    rng = random.Random(seed)
    image = Image.frombytes("RGB", (512, 512), rng.randbytes(512 * 512 * 3))
    try:
        image.save(path, format="JPEG", quality=95)
    finally:
        image.close()


def _write_template(
    path: Path,
    rows: list[dict[str, object]],
) -> None:
    workbook = Workbook()
    guide = workbook.active
    guide.title = GUIDE_SHEET
    guide.append(["说明", "测试"])
    product = workbook.create_sheet(PRODUCT_SHEET)
    product.append(list(PRODUCT_HEADERS))
    for row in rows:
        values: list[object] = [None] * len(PRODUCT_HEADERS)
        for column, value in row.items():
            values[PRODUCT_HEADERS.index(column)] = value
        product.append(values)
    mold = workbook.create_sheet(MOLD_SHEET)
    mold.append(list(MOLD_HEADERS))
    info = workbook.create_sheet(INFO_SHEET)
    info.append(["模板版本", "客户ID", "客户名称", "生成时间"])
    info.append([TEMPLATE_VERSION, 0, "测试客户", "2026-07-30"])
    workbook.save(path)
    workbook.close()


def _complete_row(
    sample_id: str,
    product_code: str,
    *,
    drawing_names: str = "",
    action: str = "新增",
    mold_code: str = "",
) -> dict[str, object]:
    return {
        "样品号": sample_id,
        "手写型号/ERP存货编码*": product_code,
        "成型方式*": "无需",
        "印刷*": "否",
        "结合方式*": "无需结合",
        "二次粘合*": "否",
        "图纸文件名(多个用分号)": drawing_names,
        "模具编号": mold_code,
        "操作*": action,
        "客户料号*": f"C-{product_code}",
        "产品名称*": f"产品-{product_code}",
        "单位": "只",
        "启用": "是",
    }


def _drawing_rows(path: Path) -> dict[int, list[int]]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        result: dict[int, list[int]] = {}
        for drawing in workbook[PRODUCT_SHEET]._images:
            marker = drawing.anchor._from
            assert int(marker.col) == 9
            result.setdefault(int(marker.row) + 1, []).append(int(marker.colOff))
        return result
    finally:
        workbook.close()


def test_batch_tool_matches_explicit_j_names_and_preserves_action(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    row = _complete_row(
        "YP001",
        "ERP-001",
        drawing_names="IMG_0001.JPG；IMG_0002.PNG",
        action="",
    )
    row["现场备注"] = "仅1张：待补第2张/确认样品归属，ERP必须阻止导入"
    _write_template(template, [row])
    _write_image(photos / "IMG_0001.JPG", "red")
    _write_image(photos / "IMG_0002.PNG", "blue", image_format="PNG")

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        batch_id="explicit",
    )

    assert result.sample_count == 1
    assert result.photo_count == 2
    assert len(result.volumes) == 1
    volume = result.volumes[0].path
    assert volume.stat().st_size < 16 * 1024 * 1024
    _preflight_product_xlsx_container(volume.read_bytes())
    workbook = load_workbook(volume, read_only=False, data_only=False)
    try:
        sheet = workbook[PRODUCT_SHEET]
        assert sheet["J2"].value == "YP001_图1.jpg;YP001_图2.jpg"
        assert sheet["L2"].value is None
        assert sheet["M2"].value is None
        assert len(sheet._images) == 2
    finally:
        workbook.close()
    offsets = _drawing_rows(volume)[2]
    assert len(offsets) == 2
    assert offsets[0] != offsets[1]


def test_batch_tool_auto_matches_sample_role_names_and_heic(
    tmp_path: Path,
) -> None:
    pytest.importorskip("pillow_heif")
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    _write_template(template, [_complete_row("YP002", "ERP-002", action="")])
    _write_image(
        photos / "YP002_图1.HEIC",
        "green",
        image_format="HEIF",
        size=(900, 1200),
    )
    _write_image(photos / "YP002_图2.jpg", "yellow", size=(1200, 900))

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        batch_id="heic",
    )

    assert result.photo_count == 2
    assert result.compressed_photo_bytes < result.source_photo_bytes
    workbook = load_workbook(result.volumes[0].path, read_only=False)
    try:
        sheet = workbook[PRODUCT_SHEET]
        assert sheet["J2"].value == "YP002_图1.jpg;YP002_图2.jpg"
        assert len(sheet._images) == 2
    finally:
        workbook.close()


def test_batch_tool_accepts_three_column_mapping_csv(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    mapping = tmp_path / "mapping.csv"
    photos.mkdir()
    _write_template(template, [_complete_row("YP003", "ERP-003", action="")])
    _write_image(photos / "IMG_1001.JPG", "purple")
    _write_image(photos / "IMG_1002.JPG", "orange")
    with mapping.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("样品号", "照片1", "照片2"))
        writer.writerow(("YP003", "IMG_1001.JPG", "IMG_1002.JPG"))

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        mapping_path=mapping,
        batch_id="csv",
    )

    assert result.sample_count == 1
    assert result.report_csv.is_file()
    report = result.report_csv.read_text(encoding="utf-8-sig")
    assert "IMG_1001.JPG" in report
    assert "IMG_1002.JPG" in report


def test_batch_tool_accepts_gbk_mapping_with_extra_column(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    mapping = tmp_path / "mapping.csv"
    photos.mkdir()
    _write_template(template, [_complete_row("YP003", "ERP-003", action="")])
    _write_image(photos / "IMG_1001.JPG", "purple")
    _write_image(photos / "IMG_1002.JPG", "orange")
    mapping.write_bytes(
        (
            "样品号,照片1,照片2,\r\n"
            "YP003,IMG_1001.JPG,IMG_1002.JPG,80012178\r\n"
        ).encode("gb18030")
    )

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        mapping_path=mapping,
        batch_id="gbk-csv",
    )

    assert result.sample_count == 1
    assert result.photo_count == 2


def test_cli_result_json_preserves_chinese_for_windows_launcher(
    tmp_path: Path,
) -> None:
    from scripts.prepare_product_photo_workbooks import _write_result_payload

    result_path = tmp_path / "result.json"
    _write_result_payload(
        {
            "ok": False,
            "message": "批量图片处理发现 5 个阻断问题",
            "issue_count": 5,
        },
        result_path,
    )

    raw = result_path.read_bytes()
    assert raw.startswith(b"{")
    payload = json.loads(raw.decode("utf-8"))
    assert payload["message"] == "批量图片处理发现 5 个阻断问题"


def test_csv_photo2_no_need_creates_confirmed_single_photo_row(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    mapping = tmp_path / "mapping.csv"
    photos.mkdir()
    row = _complete_row("YP023", "ERP-023", action="")
    row["现场备注"] = "客户样品仅保留正面"
    _write_template(template, [row])
    _write_image(photos / "IMG_2001.JPG", "purple")
    mapping.write_bytes(
        (
            "样品号,照片1,照片2\r\n"
            "YP023,IMG_2001.JPG,无需\r\n"
        ).encode("gb18030")
    )

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        mapping_path=mapping,
        batch_id="single-photo",
    )

    assert result.sample_count == 1
    assert result.photo_count == 1
    workbook = load_workbook(result.volumes[0].path, read_only=False)
    try:
        sheet = workbook[PRODUCT_SHEET]
        assert sheet["J2"].value == "YP023_图1.jpg"
        assert sheet["L2"].value == (
            f"客户样品仅保留正面；{SINGLE_PHOTO_CONFIRMED_REMARK}"
        )
        assert len(sheet._images) == 1
    finally:
        workbook.close()
    assert len(_drawing_rows(result.volumes[0].path)[2]) == 1


def test_csv_blank_photo2_still_blocks_as_missing(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    mapping = tmp_path / "mapping.csv"
    photos.mkdir()
    _write_template(template, [_complete_row("YP023", "ERP-023", action="")])
    _write_image(photos / "IMG_2001.JPG", "purple")
    mapping.write_text(
        "样品号,照片1,照片2\nYP023,IMG_2001.JPG,\n",
        encoding="utf-8-sig",
    )

    with pytest.raises(ProductPhotoBatchError) as error_info:
        build_product_photo_workbooks(
            template_path=template,
            photo_dir=photos,
            output_dir=output,
            mapping_path=mapping,
            batch_id="blank-photo2",
        )

    assert "PHOTO_MISSING" in {
        issue.code for issue in error_info.value.issues
    }


def test_mapping_is_authoritative_and_ignores_unmapped_placeholder_rows(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    mapping = tmp_path / "mapping.csv"
    photos.mkdir()
    mapped = _complete_row("YP001", "ERP-001", action="")
    filler = _complete_row(
        "YP002",
        "ERP-002",
        drawing_names="YP099_实物.jpg；YP099_展开.jpg",
        action="",
    )
    _write_template(template, [mapped, filler])
    _write_image(photos / "IMG_0001.JPG", "red")
    _write_image(photos / "IMG_0002.JPG", "blue")
    _write_image(photos / "IMG_NOT_RECORDED_YET.JPG", "green")
    mapping.write_text(
        "样品号,照片1,照片2\nYP001,IMG_0001.JPG,IMG_0002.JPG\n",
        encoding="utf-8-sig",
    )

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        mapping_path=mapping,
        batch_id="authoritative",
    )

    assert result.sample_count == 1
    assert result.volumes[0].sample_ids == ("YP001",)
    report = result.report_csv.read_text(encoding="utf-8-sig")
    assert "IMG_NOT_RECORDED_YET.JPG" not in report


def test_csv_mapping_replaces_existing_embedded_preview_images(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    mapping = tmp_path / "mapping.csv"
    photos.mkdir()
    _write_template(template, [_complete_row("YP001", "ERP-001", action="")])
    workbook = load_workbook(template, read_only=False)
    streams: list[BytesIO] = []
    try:
        sheet = workbook[PRODUCT_SHEET]
        for index, color in enumerate(("white", "gray")):
            stream = BytesIO(_image_bytes(color))
            streams.append(stream)
            image = ExcelImage(stream)
            image.anchor = OneCellAnchor(
                _from=AnchorMarker(
                    col=9,
                    colOff=pixels_to_EMU(index * 140),
                    row=1,
                    rowOff=0,
                ),
                ext=XDRPositiveSize2D(
                    cx=pixels_to_EMU(120),
                    cy=pixels_to_EMU(80),
                ),
            )
            sheet.add_image(image)
        workbook.save(template)
    finally:
        workbook.close()
        for stream in streams:
            stream.close()
    _write_image(photos / "IMG_0001.JPG", "red")
    _write_image(photos / "IMG_0002.JPG", "blue")
    mapping.write_text(
        "样品号,照片1,照片2\nYP001,IMG_0001.JPG,IMG_0002.JPG\n",
        encoding="utf-8-sig",
    )

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        mapping_path=mapping,
        batch_id="replace-preview",
    )

    report = result.report_csv.read_text(encoding="utf-8-sig")
    assert "IMG_0001.JPG" in report
    assert "IMG_0002.JPG" in report
    assert "Excel已有内嵌图" not in report


def test_early_template_error_always_writes_issue_report(
    tmp_path: Path,
) -> None:
    template = tmp_path / "invalid.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    template.write_bytes(b"not-an-xlsx")

    with pytest.raises(ProductPhotoBatchError) as error_info:
        build_product_photo_workbooks(
            template_path=template,
            photo_dir=photos,
            output_dir=output,
            batch_id="early-error",
        )

    assert error_info.value.report_path == (
        output / "批量图片处理_early-error_异常报告.csv"
    )
    assert error_info.value.report_path.is_file()
    assert {issue.code for issue in error_info.value.issues} == {
        "INPUT_VALIDATION"
    }


def test_repeated_error_report_does_not_overwrite_previous_run(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    _write_template(template, [_complete_row("YP001", "ERP-001", action="")])
    _write_image(photos / "YP001_图1.jpg", "red")

    reports: list[Path] = []
    for _attempt in range(2):
        with pytest.raises(ProductPhotoBatchError) as error_info:
            build_product_photo_workbooks(
                template_path=template,
                photo_dir=photos,
                output_dir=output,
                batch_id="same-batch",
            )
        assert error_info.value.report_path is not None
        reports.append(error_info.value.report_path)

    assert reports[0].name == "批量图片处理_same-batch_异常报告.csv"
    assert reports[1].name == "批量图片处理_same-batch_异常报告_002.csv"
    assert reports[0].is_file()
    assert reports[1].is_file()


def test_late_xlsx_preflight_error_is_converted_and_reported(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.services import product_photo_batch as batch_module
    from app.services.secure_uploads import UploadValidationError

    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    _write_template(template, [_complete_row("YP001", "ERP-001", action="")])
    _write_image(photos / "YP001_图1.jpg", "red")
    _write_image(photos / "YP001_图2.jpg", "blue")
    original_preflight = batch_module._preflight_product_xlsx_container
    call_count = 0

    def fail_second_preflight(content: bytes) -> None:
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            original_preflight(content)
            return
        raise UploadValidationError("测试用输出阶段错误")

    monkeypatch.setattr(
        batch_module,
        "_preflight_product_xlsx_container",
        fail_second_preflight,
    )

    with pytest.raises(ProductPhotoBatchError) as error_info:
        build_product_photo_workbooks(
            template_path=template,
            photo_dir=photos,
            output_dir=output,
            batch_id="late-preflight",
        )

    assert error_info.value.report_path is not None
    assert error_info.value.report_path.is_file()
    assert "未通过 ERP XLSX 安全检查" in error_info.value.issues[0].message
    assert list(output.glob("*.xlsx")) == []


def test_product_xlsx_preflight_allows_empty_directory_entries(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    repacked = tmp_path / "repacked.xlsx"
    _write_template(template, [_complete_row("YP001", "ERP-001", action="")])
    with ZipFile(repacked, "w", compression=ZIP_DEFLATED) as output_zip:
        output_zip.writestr("xl/drawings/", b"")
        output_zip.writestr("xl/drawings/_rels/", b"")
        output_zip.writestr("xl/media/", b"")
        with ZipFile(template, "r") as input_zip:
            for item in input_zip.infolist():
                output_zip.writestr(item, input_zip.read(item.filename))

    _preflight_product_xlsx_container(repacked.read_bytes())


def test_product_xlsx_preflight_rejects_directory_entry_with_payload(
    tmp_path: Path,
) -> None:
    from app.services.secure_uploads import UploadValidationError

    template = tmp_path / "template.xlsx"
    repacked = tmp_path / "repacked.xlsx"
    _write_template(template, [_complete_row("YP001", "ERP-001", action="")])
    with ZipFile(repacked, "w", compression=ZIP_DEFLATED) as output_zip:
        output_zip.writestr("xl/drawings/_rels/", b"not-empty")
        with ZipFile(template, "r") as input_zip:
            for item in input_zip.infolist():
                output_zip.writestr(item, input_zip.read(item.filename))

    with pytest.raises(UploadValidationError, match="目录成员包含异常数据"):
        _preflight_product_xlsx_container(repacked.read_bytes())


@pytest.mark.parametrize(
    ("second_photo", "expected_code"),
    [
        (None, "PHOTO_MISSING"),
        ("same", "PHOTO_CONTENT_DUPLICATE"),
    ],
)
def test_batch_tool_blocks_missing_or_duplicate_second_photo_without_output(
    tmp_path: Path,
    second_photo: str | None,
    expected_code: str,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    _write_template(template, [_complete_row("YP004", "ERP-004", action="")])
    first = _image_bytes("white")
    (photos / "YP004_图1.jpg").write_bytes(first)
    if second_photo == "same":
        (photos / "YP004_图2.jpg").write_bytes(first)

    with pytest.raises(ProductPhotoBatchError) as error_info:
        build_product_photo_workbooks(
            template_path=template,
            photo_dir=photos,
            output_dir=output,
            batch_id="blocked",
        )

    assert expected_code in {issue.code for issue in error_info.value.issues}
    assert error_info.value.report_path is not None
    assert error_info.value.report_path.is_file()
    assert list(output.glob("*.xlsx")) == []


def test_batch_tool_splits_without_changing_non_volume_rows(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    rows = []
    for index, color_pair in enumerate(
        (("red", "blue"), ("green", "yellow"), ("purple", "orange")),
        start=1,
    ):
        sample_id = f"YP{index:03d}"
        rows.append(_complete_row(sample_id, f"ERP-{index:03d}"))
        _write_image(photos / f"{sample_id}_图1.jpg", color_pair[0])
        _write_image(photos / f"{sample_id}_图2.jpg", color_pair[1])
    _write_template(template, rows)

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        batch_id="split",
        max_samples_per_volume=1,
    )

    assert len(result.volumes) == 3
    assert [volume.sample_ids for volume in result.volumes] == [
        ("YP001",),
        ("YP002",),
        ("YP003",),
    ]
    for volume_index, volume in enumerate(result.volumes, start=1):
        workbook = load_workbook(volume.path, read_only=False, data_only=False)
        try:
            sheet = workbook[PRODUCT_SHEET]
            assert len(sheet._images) == 2
            for row_number in range(2, 5):
                if row_number == volume_index + 1:
                    assert sheet.cell(row_number, 13).value == "新增"
                else:
                    assert sheet.cell(row_number, 13).value is None
        finally:
            workbook.close()
        assert volume.size < 16 * 1024 * 1024


def test_batch_tool_blocks_unused_supported_photo(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    _write_template(template, [_complete_row("YP005", "ERP-005", action="")])
    _write_image(photos / "YP005_图1.jpg", "red")
    _write_image(photos / "YP005_图2.jpg", "blue")
    _write_image(photos / "orphan.jpg", "black")

    with pytest.raises(ProductPhotoBatchError) as error_info:
        build_product_photo_workbooks(
            template_path=template,
            photo_dir=photos,
            output_dir=output,
            batch_id="unused",
        )

    assert "PHOTO_UNUSED" in {issue.code for issue in error_info.value.issues}


def test_batch_tool_handles_200_samples_and_400_photos_in_four_volumes(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    rows = []
    for sample_index in range(200):
        sample_id = f"YP{sample_index + 1:04d}"
        rows.append(_complete_row(sample_id, f"ERP-{sample_index + 1:04d}"))
        for role in (1, 2):
            color_index = sample_index * 2 + role - 1
            color = (
                color_index % 256,
                (color_index // 256) * 127,
                (color_index * 37) % 256,
            )
            _write_image(
                photos / f"{sample_id}_图{role}.jpg",
                color,
                size=(80, 60),
            )
    _write_template(template, rows)

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        batch_id="large",
        max_samples_per_volume=50,
    )

    assert result.sample_count == 200
    assert result.photo_count == 400
    assert len(result.volumes) == 4
    all_sample_ids = [
        sample_id
        for volume in result.volumes
        for sample_id in volume.sample_ids
    ]
    assert len(all_sample_ids) == len(set(all_sample_ids)) == 200
    for volume in result.volumes:
        assert len(volume.sample_ids) == 50
        assert volume.size < 16 * 1024 * 1024
        workbook = load_workbook(volume.path, read_only=False)
        try:
            assert len(workbook[PRODUCT_SHEET]._images) == 100
        finally:
            workbook.close()


def test_batch_tool_splits_again_by_actual_xlsx_size(
    tmp_path: Path,
) -> None:
    template = tmp_path / "template.xlsx"
    photos = tmp_path / "photos"
    output = tmp_path / "output"
    photos.mkdir()
    rows = []
    for index in range(8):
        sample_id = f"SIZE{index + 1:03d}"
        rows.append(_complete_row(sample_id, f"ERP-SIZE-{index + 1:03d}"))
        _write_noise_image(photos / f"{sample_id}_图1.jpg", index * 2)
        _write_noise_image(photos / f"{sample_id}_图2.jpg", index * 2 + 1)
    _write_template(template, rows)

    result = build_product_photo_workbooks(
        template_path=template,
        photo_dir=photos,
        output_dir=output,
        batch_id="actual-size",
        target_xlsx_bytes=1 * 1024 * 1024,
        max_samples_per_volume=200,
    )

    assert len(result.volumes) > 1
    assert sum(len(volume.sample_ids) for volume in result.volumes) == 8
    assert all(volume.size <= 1 * 1024 * 1024 for volume in result.volumes)
