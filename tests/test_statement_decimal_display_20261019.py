from decimal import Decimal
from io import BytesIO

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from tests.test_p1_130_statement_invoice_finance import p1_130_app, _login


@pytest.mark.parametrize('value,expected', [
    ('222.525', '222.525'), ('222.5', '222.500'), ('222.000', '222'),
    ('0', '0'), ('-0.0004', '0'), ('-2.5', '-2.500'),
    ('1.2345', '1.235'), ('-1.2345', '-1.235'), ('222.9996', '223'),
])
def test_display_rule_fractional_three_places_integer_plain(value, expected):
    from app.services.reconciliation_number_display import format_reconciliation_number
    assert format_reconciliation_number(Decimal(value)) == expected


def test_reconciliation_read_projection_does_not_round_saved_third_decimal():
    from app.api.finance import _money_value
    assert _money_value(Decimal('222.525')) == Decimal('222.525')
    assert _money_value(Decimal('0.004')) == Decimal('0.004')


@pytest.mark.parametrize('suffix,sheet_name,amount_cell,total_cell', [
    ('customer-export.xlsx', '对账单', 'H5', 'H7'),
    ('export', '月结对账单', 'N6', 'N9'),
])
def test_statement_excel_retains_numeric_precision_and_freezes_original_facts(
        p1_130_app, suffix, sheet_name, amount_cell, total_cell):
    from app.models.finance import Statement, StatementItem
    app, factory = p1_130_app
    def facts(db):
        bill = db.get(Statement, 1)
        return (bill.version, bill.total_receivable, bill.confirmation_status,
                [(r.id, r.unit_price_snapshot, r.receivable_amount) for r in
                 db.scalars(select(StatementItem).where(StatementItem.statement_id == 1).order_by(StatementItem.id))])
    with factory() as db:
        line = db.scalar(select(StatementItem).where(StatementItem.statement_id == 1).order_by(StatementItem.id))
        line.unit_price_snapshot = Decimal('22.2525')
        line.receivable_amount = Decimal('222.525')
        db.get(Statement, 1).total_receivable = Decimal('322.525')
        db.commit()
        before = facts(db)
    with TestClient(app) as client:
        _login(client)
        response = client.get('/api/finance/statements/1/' + suffix)
        assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content), data_only=False)[sheet_name]
    assert sheet[amount_cell].value == 222.525
    assert sheet[amount_cell].data_type == 'n'
    assert sheet[amount_cell].number_format == '0.000'
    assert sheet[total_cell].value == 322.525
    assert sheet[total_cell].number_format == '0.000'
    # The other line and its integer quantity must not acquire decimal zeros.
    integer_cell = 'H6' if suffix.startswith('customer') else 'N7'
    assert sheet[integer_cell].value == 100
    assert sheet[integer_cell].number_format == '0'
    with factory() as db:
        assert facts(db) == before


def test_statement_pdf_uses_three_decimals_and_plain_integers():
    from app.services.statement_pdf import render_customer_statement_pdf
    content = render_customer_statement_pdf(
        title='对账单', subtitle='虚构精度验证客户', price_label='单价', amount_label='金额',
        rows=[dict(quantity=1, unit_price=Decimal('222.525'), amount=Decimal('222.525')),
              dict(quantity=2, unit_price=Decimal('50'), amount=Decimal('100'))],
        quantity_total=3, amount_total=Decimal('322.525'))
    with fitz.open(stream=content, filetype='pdf') as document:
        words = [w[4] for page in document for w in page.get_text('words')]
    assert words.count('222.525') == 2
    assert '322.525' in words
    assert '100' in words and '50' in words
    assert '100.00' not in words and '50.000' not in words


def test_comparison_report_formats_numeric_values_without_changing_matching():
    from app.services.customer_statement_check import export_report
    report = dict(customer_name='虚构客户', statement_month='2026-10', statement_number='TEST',
        matching_fields=['存货编码'], complete=True, errors=[], items=[dict(status='一致',
        delivery_number='TEST', customer_po='000001', product_code='=CODE', product_name='测试',
        customer_quantity='222.500', erp_quantity='222.500', quantity_difference='0',
        customer_amount='222.525', erp_amount='222.525', amount_difference='0',
        customer_prices='222.500', erp_prices='222.500', customer_rows=[5], statement_item_ids=[1])])
    sheet = load_workbook(export_report(report), data_only=False)['对账差异']
    assert sheet['F5'].value == 222.5 and sheet['F5'].number_format == '0.000'
    assert sheet['I5'].value == 222.525 and sheet['I5'].number_format == '0.000'
    assert sheet['H5'].value == 0 and sheet['H5'].number_format == '0'
    assert sheet['D5'].value == '=CODE' and sheet['D5'].data_type == 's'
