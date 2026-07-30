from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import func, select
from sqlalchemy.orm import sessionmaker


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook_bytes(
    *,
    paper_rows: list[tuple] | None = None,
    material_rows: list[tuple] | None = None,
) -> bytes:
    from app.services.supplier_material_workbook import (
        MATERIAL_HEADERS,
        MATERIAL_SHEET,
        PAPER_HEADERS,
        PAPER_SHEET,
    )

    workbook = Workbook()
    paper_sheet = workbook.active
    paper_sheet.title = PAPER_SHEET
    paper_sheet.append(PAPER_HEADERS)
    for row in paper_rows or []:
        paper_sheet.append(row)
    material_sheet = workbook.create_sheet(MATERIAL_SHEET)
    material_sheet.append(MATERIAL_HEADERS)
    for row in material_rows or []:
        material_sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.fixture()
def supplier_workbook_app(tmp_path):
    from app.api.auth import router as auth_router
    from app.api.deps import get_db
    from app.api.materials import router as materials_router
    from app.core.database import create_sqlite_engine
    from app.core.security import hash_password
    from app.models import Base
    from app.models.material import Material
    from app.models.supplier import Supplier, SupplierAlias
    from app.models.supplier_paper_code import SupplierPaperCode
    from app.models.user import User
    from app.services.supplier_master import normalize_supplier_identity

    engine = create_sqlite_engine(tmp_path / "supplier-material-workbook.sqlite3")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, expire_on_commit=False)
    with factory() as db:
        db.add_all(
            [
                User(
                    username="workbook-admin",
                    password_hash=hash_password("WorkbookPass123!"),
                    role="admin",
                    real_name="材质管理员",
                    must_change_password=False,
                ),
                User(
                    username="workbook-sales",
                    password_hash=hash_password("WorkbookPass123!"),
                    role="sales",
                    real_name="销售",
                    must_change_password=False,
                ),
                Supplier(
                    standard_name="苏州嘉林亿",
                    normalized_name=normalize_supplier_identity("苏州嘉林亿"),
                    display_name="嘉林亿",
                    business_code="JLY",
                    normalized_business_code="JLY",
                    sort_order=10,
                    is_active=True,
                    version=1,
                    aliases=[
                        SupplierAlias(
                            alias_name="嘉林亿",
                            normalized_alias=normalize_supplier_identity("嘉林亿"),
                        ),
                        SupplierAlias(
                            alias_name="苏州嘉林亿纸业有限公司",
                            normalized_alias=normalize_supplier_identity(
                                "苏州嘉林亿纸业有限公司"
                            ),
                        ),
                    ],
                ),
                Supplier(
                    standard_name="昆山鸣朋",
                    normalized_name=normalize_supplier_identity("昆山鸣朋"),
                    display_name="鸣朋",
                    business_code="MP",
                    normalized_business_code="MP",
                    sort_order=20,
                    is_active=True,
                    version=1,
                    aliases=[
                        SupplierAlias(
                            alias_name="鸣朋",
                            normalized_alias=normalize_supplier_identity("鸣朋"),
                        ),
                        SupplierAlias(
                            alias_name="昆山鸣朋纸业有限公司",
                            normalized_alias=normalize_supplier_identity(
                                "昆山鸣朋纸业有限公司"
                            ),
                        ),
                    ],
                ),
                Supplier(
                    standard_name="苏州佳丰",
                    normalized_name=normalize_supplier_identity("苏州佳丰"),
                    display_name="佳丰",
                    business_code="JF",
                    normalized_business_code="JF",
                    sort_order=90,
                    is_active=False,
                    version=1,
                    aliases=[
                        SupplierAlias(
                            alias_name="佳丰",
                            normalized_alias=normalize_supplier_identity("佳丰"),
                        )
                    ],
                ),
                SupplierPaperCode(
                    supplier_name="苏州嘉林亿",
                    code_char="C",
                    paper_name="牛卡C",
                    gram_weight=80,
                    is_active=True,
                ),
                SupplierPaperCode(
                    supplier_name="苏州嘉林亿",
                    code_char="4",
                    paper_name="瓦纸4",
                    gram_weight=100,
                    is_active=True,
                ),
                Material(
                    code="C4C",
                    supplier_name="苏州嘉林亿",
                    layer_count=3,
                    basis_weight_description="80g/100g/80g",
                    paper_composition=(
                        "面纸:C=80g 牛卡C | 瓦楞纸:4=100g 瓦纸4 | "
                        "里纸:C=80g 牛卡C"
                    ),
                    quote_price=Decimal("1.2000"),
                    quote_date=date(2026, 7, 1),
                    price_unit="元/㎡",
                    is_active=True,
                    version=1,
                ),
            ]
        )
        db.commit()

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(materials_router, prefix="/api/master/materials")

    def override_get_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_get_db
    app.state.session_factory = factory
    yield app
    engine.dispose()


def _login(client: TestClient, username: str) -> None:
    response = client.post(
        "/api/auth/login",
        json={"username": username, "password": "WorkbookPass123!"},
    )
    assert response.status_code == 200


def _upload(client: TestClient, content: bytes):
    return client.post(
        "/api/master/materials/import/preview",
        files={"file": ("供应商材质.xlsx", content, EXCEL_MIME)},
    )


def test_supplier_input_cleaning_does_not_hardcode_aliases() -> None:
    from app.services.supplier_material_workbook import normalize_supplier_name

    assert normalize_supplier_name(" 鸣朋 ") == "鸣朋"
    assert normalize_supplier_name("昆山鸣朋纸业有限公司") == (
        "昆山鸣朋纸业有限公司"
    )
    assert normalize_supplier_name("嘉林亿") == "嘉林亿"
    assert normalize_supplier_name("未来纸板供应商") == "未来纸板供应商"


def test_template_exports_blank_two_sheet_entry_form_and_is_admin_only(
    supplier_workbook_app: FastAPI,
) -> None:
    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-sales")
        assert (
            client.get("/api/master/materials/import-template.xlsx").status_code
            == 403
        )

        _login(client, "workbook-admin")
        response = client.get("/api/master/materials/import-template.xlsx")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith(EXCEL_MIME)
        workbook = load_workbook(BytesIO(response.content), read_only=False)
        try:
            assert workbook.sheetnames == ["基础纸种", "组合材质"]
            assert tuple(cell.value for cell in workbook["基础纸种"][1]) == (
                "系统ID",
                "供应商标准名",
                "基础代码",
                "纸种名称",
                "克重(g)",
                "等级/说明",
                "纸张用途",
                "备注",
                "启用",
            )
            assert tuple(cell.value for cell in workbook["组合材质"][1]) == (
                "系统ID",
                "当前版本",
                "供应商标准名",
                "组合代码",
                "层数",
                "平方单价(元/㎡)",
                "报价日期",
                "价格单位",
                "备注",
                "启用",
            )
            assert workbook["基础纸种"].max_row == 1
            assert workbook["组合材质"].max_row == 1
            assert "主档中维护的别名" in workbook["基础纸种"]["B1"].comment.text
            assert "元/平方米" in workbook["组合材质"]["F1"].comment.text
        finally:
            workbook.close()


def test_preview_then_apply_is_atomic_versioned_and_normalizes_aliases(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.audit import OperationLog
    from app.models.master_data_object_version import MasterDataObjectVersion
    from app.models.material import Material
    from app.models.material_price_history import MaterialPriceHistory
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    with factory() as db:
        paper_c = db.scalar(
            select(SupplierPaperCode).where(SupplierPaperCode.code_char == "C")
        )
        material = db.scalar(select(Material).where(Material.code == "C4C"))
        paper_c_id = paper_c.id
        material_id = material.id
        material_version = material.version

    content = _workbook_bytes(
        paper_rows=[
            (
                paper_c_id,
                "嘉林亿",
                "C",
                "牛卡C",
                90,
                "A级",
                "面纸",
                "批量调整克重",
                "是",
            ),
            (
                None,
                "嘉林亿",
                "6",
                "高强瓦纸6",
                130,
                "A级",
                "芯纸",
                None,
                "是",
            ),
        ],
        material_rows=[
            (
                material_id,
                material_version,
                "嘉林亿",
                "C4C",
                3,
                10,
                "2026-07-28",
                "元/㎡",
                "大幅价格变化也由预览确认",
                "是",
            ),
            (
                None,
                None,
                "嘉林亿",
                "C6C",
                3,
                1.58,
                "2026-07-28",
                "元/㎡",
                "新增组合",
                "是",
            ),
        ],
    )

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        preview = _upload(client, content)
        assert preview.status_code == 200
        body = preview.json()
        assert body["valid"] is True
        assert body["summary"] == {
            "paper_create": 1,
            "paper_update": 1,
            "paper_unchanged": 0,
            "material_create": 1,
            "material_update": 1,
            "material_unchanged": 0,
        }
        assert {row["supplier_name"] for row in body["materials"]} == {
            "苏州嘉林亿"
        }

        with factory() as db:
            unchanged = db.get(SupplierPaperCode, paper_c_id)
            assert unchanged.gram_weight == 80
            assert db.scalar(select(Material).where(Material.code == "C6C")) is None

        applied = client.post(
            "/api/master/materials/import/apply",
            json={"preview_token": body["preview_token"]},
        )
        assert applied.status_code == 200
        assert applied.json()["paper_created"] == 1
        assert applied.json()["paper_updated"] == 1
        assert applied.json()["material_created"] == 1
        assert applied.json()["material_updated"] == 1

    with factory() as db:
        changed_paper = db.get(SupplierPaperCode, paper_c_id)
        assert changed_paper.supplier_name == "苏州嘉林亿"
        assert changed_paper.gram_weight == 90
        existing = db.get(Material, material_id)
        assert existing.version == material_version + 1
        assert existing.quote_price == Decimal("10.0000")
        assert existing.basis_weight_description == "90g/100g/90g"
        created = db.scalar(select(Material).where(Material.code == "C6C"))
        assert created is not None
        assert created.supplier_name == "苏州嘉林亿"
        assert created.quote_price == Decimal("1.5800")
        assert created.basis_weight_description == "90g/130g/90g"
        assert (
            db.scalar(
                select(func.count())
                .select_from(MasterDataObjectVersion)
                .where(MasterDataObjectVersion.object_type == "material")
            )
            >= 3
        )
        assert (
            db.scalar(
                select(func.count())
                .select_from(OperationLog)
                .where(OperationLog.resource == "Material")
            )
            >= 2
        )
        price_history = db.scalar(
            select(MaterialPriceHistory).where(
                MaterialPriceHistory.material_id == material_id
            )
        )
        assert price_history.old_price == Decimal("1.2000")
        assert price_history.new_price == Decimal("10.0000")
        assert price_history.adjust_reason == "供应商材质 Excel 批量维护"


def test_invalid_preview_and_global_code_conflict_write_nothing(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.material import Material
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    invalid = _workbook_bytes(
        paper_rows=[
            (
                None,
                "鸣朋",
                "X",
                "测试纸",
                120,
                None,
                None,
                None,
                "是",
            )
        ],
        material_rows=[
            (
                None,
                None,
                "鸣朋",
                "X9X",
                3,
                1.2,
                None,
                "元/㎡",
                None,
                "是",
            )
        ],
    )
    conflict = _workbook_bytes(
        paper_rows=[],
        material_rows=[
            (
                None,
                None,
                "鸣朋",
                "C4C",
                3,
                1.2,
                None,
                "元/㎡",
                None,
                "是",
            )
        ],
    )

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        invalid_response = _upload(client, invalid)
        assert invalid_response.status_code == 200
        invalid_body = invalid_response.json()
        assert invalid_body["valid"] is False
        assert "preview_token" not in invalid_body
        assert "缺少启用的基础代码：9" in invalid_body["errors"][0]["message"]

        conflict_response = _upload(client, conflict)
        assert conflict_response.status_code == 200
        conflict_body = conflict_response.json()
        assert conflict_body["valid"] is False
        assert "全局唯一" in conflict_body["errors"][0]["message"]

    with factory() as db:
        assert db.scalar(
            select(SupplierPaperCode).where(
                SupplierPaperCode.supplier_name == "昆山鸣朋"
            )
        ) is None
        assert db.scalar(select(Material).where(Material.code == "X9X")) is None


def test_formula_workbook_and_non_admin_are_rejected(
    supplier_workbook_app: FastAPI,
) -> None:
    content = _workbook_bytes()
    workbook = load_workbook(BytesIO(content))
    workbook["基础纸种"]["B2"] = "=1+1"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    formula_content = output.getvalue()

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-sales")
        assert _upload(client, content).status_code == 403
        assert (
            client.post(
                "/api/master/materials/import/apply",
                json={"preview_token": "not-authorized"},
            ).status_code
            == 403
        )

        _login(client, "workbook-admin")
        rejected = _upload(client, formula_content)
        assert rejected.status_code == 400
        assert rejected.json()["detail"]["code"] == (
            "SUPPLIER_MATERIAL_UPLOAD_INVALID"
        )


def test_unknown_and_inactive_suppliers_are_rejected_before_preview_token(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.material import Material
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    unknown = _workbook_bytes(
        paper_rows=[
            (
                None,
                "未来纸板供应商",
                "X",
                "测试纸",
                120,
                None,
                None,
                None,
                "是",
            )
        ]
    )
    inactive = _workbook_bytes(
        paper_rows=[
            (
                None,
                "佳丰",
                "Y",
                "停用供应商测试纸",
                120,
                None,
                None,
                None,
                "是",
            )
        ]
    )

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        unknown_response = _upload(client, unknown)
        assert unknown_response.status_code == 200
        unknown_body = unknown_response.json()
        assert unknown_body["valid"] is False
        assert "preview_token" not in unknown_body
        assert any(
            "尚未建档" in error["message"]
            for error in unknown_body["errors"]
        )

        inactive_response = _upload(client, inactive)
        assert inactive_response.status_code == 200
        inactive_body = inactive_response.json()
        assert inactive_body["valid"] is False
        assert "preview_token" not in inactive_body
        assert any(
            "已停用" in error["message"]
            for error in inactive_body["errors"]
        )

    with factory() as db:
        assert db.scalar(
            select(SupplierPaperCode).where(
                SupplierPaperCode.code_char.in_(["X", "Y"])
            )
        ) is None
        assert db.scalar(
            select(Material).where(Material.code.in_(["X", "Y"]))
        ) is None


def test_enabled_dynamic_alias_resolves_to_supplier_standard_name(
    supplier_workbook_app: FastAPI,
) -> None:
    content = _workbook_bytes(
        paper_rows=[
            (
                None,
                "苏州嘉林亿纸业有限公司",
                "6",
                "高强瓦纸6",
                130,
                None,
                None,
                None,
                "是",
            )
        ]
    )

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        response = _upload(client, content)
        assert response.status_code == 200
        body = response.json()
        assert body["valid"] is True
        assert body["paper_codes"][0]["supplier_name"] == "苏州嘉林亿"


def test_existing_historical_material_keeps_inactive_or_unknown_supplier(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.material import Material
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    with factory() as db:
        db.add_all(
            [
                SupplierPaperCode(
                    supplier_name="苏州佳丰",
                    code_char="J",
                    paper_name="佳丰面纸",
                    gram_weight=120,
                    is_active=True,
                ),
                SupplierPaperCode(
                    supplier_name="苏州佳丰",
                    code_char="4",
                    paper_name="佳丰瓦纸",
                    gram_weight=100,
                    is_active=True,
                ),
                SupplierPaperCode(
                    supplier_name="历史未建档供应商",
                    code_char="U",
                    paper_name="历史面纸",
                    gram_weight=130,
                    is_active=True,
                ),
                SupplierPaperCode(
                    supplier_name="历史未建档供应商",
                    code_char="5",
                    paper_name="历史瓦纸",
                    gram_weight=105,
                    is_active=True,
                ),
                Material(
                    code="J4J",
                    supplier_name="苏州佳丰",
                    layer_count=3,
                    quote_price=Decimal("1.1000"),
                    price_unit="元/㎡",
                    is_active=True,
                    version=1,
                ),
                Material(
                    code="U5U",
                    supplier_name="历史未建档供应商",
                    layer_count=3,
                    quote_price=Decimal("1.3000"),
                    price_unit="元/㎡",
                    is_active=True,
                    version=1,
                ),
            ]
        )
        db.commit()
        jiafeng = db.scalar(select(Material).where(Material.code == "J4J"))
        unknown = db.scalar(select(Material).where(Material.code == "U5U"))
        rows = [
            (
                jiafeng.id,
                jiafeng.version,
                "佳丰",
                "J4J",
                3,
                1.1,
                "2026-07-30",
                "元/㎡",
                "仅更新历史备注",
                "是",
            ),
            (
                unknown.id,
                unknown.version,
                "历史未建档供应商",
                "U5U",
                3,
                1.3,
                "2026-07-30",
                "元/㎡",
                "未建档历史备注",
                "是",
            ),
        ]

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        preview = _upload(client, _workbook_bytes(material_rows=rows))
        assert preview.status_code == 200
        body = preview.json()
        assert body["valid"] is True
        assert body["summary"]["material_update"] == 2
        applied = client.post(
            "/api/master/materials/import/apply",
            json={"preview_token": body["preview_token"]},
        )
        assert applied.status_code == 200
        assert applied.json()["material_updated"] == 2

    with factory() as db:
        jiafeng = db.scalar(select(Material).where(Material.code == "J4J"))
        unknown = db.scalar(select(Material).where(Material.code == "U5U"))
        assert jiafeng.supplier_name == "苏州佳丰"
        assert jiafeng.remarks == "仅更新历史备注"
        assert jiafeng.version == 2
        assert unknown.supplier_name == "历史未建档供应商"
        assert unknown.remarks == "未建档历史备注"
        assert unknown.version == 2


def test_existing_historical_paper_code_can_be_maintained_only_by_system_id(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    with factory() as db:
        inactive = SupplierPaperCode(
            supplier_name="苏州佳丰",
            code_char="J",
            paper_name="佳丰历史纸",
            gram_weight=120,
            remark="旧备注",
            is_active=True,
        )
        unknown = SupplierPaperCode(
            supplier_name="历史未建档供应商",
            code_char="U",
            paper_name="未建档历史纸",
            gram_weight=130,
            remark="旧备注",
            is_active=True,
        )
        db.add_all([inactive, unknown])
        db.commit()
        rows = [
            (
                inactive.id,
                "苏州佳丰",
                "J",
                "佳丰历史纸",
                120,
                None,
                None,
                "新备注",
                "是",
            ),
            (
                unknown.id,
                "历史未建档供应商",
                "U",
                "未建档历史纸",
                130,
                None,
                None,
                "未建档新备注",
                "是",
            ),
        ]

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        preview = _upload(client, _workbook_bytes(paper_rows=rows))
        assert preview.status_code == 200
        body = preview.json()
        assert body["valid"] is True
        assert body["summary"]["paper_update"] == 2
        applied = client.post(
            "/api/master/materials/import/apply",
            json={"preview_token": body["preview_token"]},
        )
        assert applied.status_code == 200
        assert applied.json()["paper_updated"] == 2

    with factory() as db:
        inactive = db.scalar(
            select(SupplierPaperCode).where(
                SupplierPaperCode.supplier_name == "苏州佳丰",
                SupplierPaperCode.code_char == "J",
            )
        )
        unknown = db.scalar(
            select(SupplierPaperCode).where(
                SupplierPaperCode.supplier_name == "历史未建档供应商",
                SupplierPaperCode.code_char == "U",
            )
        )
        assert inactive.remark == "新备注"
        assert unknown.remark == "未建档新备注"


def test_historical_paper_code_cannot_be_created_or_change_supplier(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    with factory() as db:
        historical = SupplierPaperCode(
            supplier_name="苏州佳丰",
            code_char="J",
            paper_name="佳丰历史纸",
            gram_weight=120,
            is_active=True,
        )
        db.add(historical)
        db.commit()
        historical_id = historical.id

    cases = [
        (
            historical_id,
            "历史未建档供应商",
            "J",
            "供应商变更必须拒绝",
        ),
        (
            None,
            "苏州佳丰",
            "N",
            "停用供应商新增必须拒绝",
        ),
    ]
    for system_id, supplier_name, code_char, paper_name in cases:
        content = _workbook_bytes(
            paper_rows=[
                (
                    system_id,
                    supplier_name,
                    code_char,
                    paper_name,
                    120,
                    None,
                    None,
                    None,
                    "是",
                )
            ]
        )
        with TestClient(supplier_workbook_app) as client:
            _login(client, "workbook-admin")
            response = _upload(client, content)
            assert response.status_code == 200
            body = response.json()
            assert body["valid"] is False
            assert "preview_token" not in body

    with factory() as db:
        historical = db.get(SupplierPaperCode, historical_id)
        assert historical.supplier_name == "苏州佳丰"
        assert historical.paper_name == "佳丰历史纸"
        assert db.scalar(
            select(SupplierPaperCode).where(
                SupplierPaperCode.supplier_name == "苏州佳丰",
                SupplierPaperCode.code_char == "N",
            )
        ) is None


def test_historical_material_cannot_change_to_unknown_or_inactive_supplier(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.material import Material

    factory = supplier_workbook_app.state.session_factory
    with factory() as db:
        material = db.scalar(select(Material).where(Material.code == "C4C"))
        material_id = material.id
        version = material.version

    for supplier_name in ("佳丰", "未来纸板供应商"):
        content = _workbook_bytes(
            material_rows=[
                (
                    material_id,
                    version,
                    supplier_name,
                    "C4C",
                    3,
                    1.2,
                    "2026-07-30",
                    "元/㎡",
                    "禁止改供应商",
                    "是",
                )
            ]
        )
        with TestClient(supplier_workbook_app) as client:
            _login(client, "workbook-admin")
            response = _upload(client, content)
            assert response.status_code == 200
            body = response.json()
            assert body["valid"] is False
            assert "preview_token" not in body
            assert any(
                marker in error["message"]
                for error in body["errors"]
                for marker in ("已停用", "尚未建档")
            )

    with factory() as db:
        material = db.get(Material, material_id)
        assert material.supplier_name == "苏州嘉林亿"
        assert material.version == version


def test_supplier_change_after_preview_blocks_apply_before_any_write(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.material import Material
    from app.models.supplier import Supplier
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    content = _workbook_bytes(
        paper_rows=[
            (
                None,
                "鸣朋",
                "M",
                "鸣朋测试纸",
                110,
                None,
                None,
                None,
                "是",
            )
        ]
    )

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        preview = _upload(client, content)
        assert preview.status_code == 200
        token = preview.json()["preview_token"]

        with factory() as db:
            supplier = db.scalar(
                select(Supplier).where(Supplier.standard_name == "昆山鸣朋")
            )
            supplier.is_active = False
            supplier.version += 1
            db.commit()

        applied = client.post(
            "/api/master/materials/import/apply",
            json={"preview_token": token},
        )
        assert applied.status_code == 409
        assert applied.json()["detail"]["code"] == (
            "SUPPLIER_MATERIAL_SUPPLIER_STALE"
        )

    with factory() as db:
        assert db.scalar(
            select(SupplierPaperCode).where(
                SupplierPaperCode.supplier_name == "昆山鸣朋",
                SupplierPaperCode.code_char == "M",
            )
        ) is None
        assert db.scalar(select(Material).where(Material.code == "M")) is None


def test_stale_material_version_blocks_entire_apply_before_paper_write(
    supplier_workbook_app: FastAPI,
) -> None:
    from app.models.material import Material
    from app.models.supplier_paper_code import SupplierPaperCode

    factory = supplier_workbook_app.state.session_factory
    with factory() as db:
        material = db.scalar(select(Material).where(Material.code == "C4C"))
        material_id = material.id
        version = material.version

    content = _workbook_bytes(
        paper_rows=[
            (
                None,
                "鸣朋",
                "M",
                "鸣朋测试纸",
                110,
                None,
                None,
                None,
                "是",
            )
        ],
        material_rows=[
            (
                material_id,
                version,
                "嘉林亿",
                "C4C",
                3,
                1.3,
                None,
                "元/㎡",
                None,
                "是",
            )
        ],
    )

    with TestClient(supplier_workbook_app) as client:
        _login(client, "workbook-admin")
        preview = _upload(client, content)
        assert preview.status_code == 200
        token = preview.json()["preview_token"]

        with factory() as db:
            material = db.get(Material, material_id)
            material.version += 1
            db.commit()

        applied = client.post(
            "/api/master/materials/import/apply",
            json={"preview_token": token},
        )
        assert applied.status_code == 409
        assert applied.json()["detail"]["code"] == (
            "SUPPLIER_MATERIAL_PREVIEW_STALE"
        )

    with factory() as db:
        assert db.scalar(
            select(SupplierPaperCode).where(
                SupplierPaperCode.supplier_name == "昆山鸣朋",
                SupplierPaperCode.code_char == "M",
            )
        ) is None
