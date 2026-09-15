from io import BytesIO
import json
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from tests.test_phase11_requisition import requisition_app, _login


def sample():
    return dict(order_number='SRO-TEST-0001', supplier_name='示例供应商',created_at='2026-09-15',
                sender={'company_name':'苏州天明包装有限公司','address':'苏州','phone':'123'},
                lines=[dict(report_length_mm=870,report_width_mm=970,crease_display='净料',
                            material_display='BC14C / AB',requisition_qty=31,remark='=1+1',
                            source_items=[{'customer_po':'PRIVATE-PO','location':'PRIVATE-LOCATION'}])])


def test_supplier_export_whitelist_numbers_and_formula_safety():
    from app.services.supplier_purchase_view import supplier_xlsx, supplier_pdf, purchase_date
    assert purchase_date({'created_at':'2026-09-14T23:00:00Z'})=='2026-09-15'
    purchase=sample(); xlsx=supplier_xlsx(purchase)
    ws=load_workbook(BytesIO(xlsx)).active
    assert ws['E6'].value == 31 and ws['E6'].data_type=='n'
    assert ws['F6'].value=='=1+1' and ws['F6'].data_type=='s'
    assert 'PRIVATE' not in str(list(ws.values))
    assert ws.print_title_rows=='$1:$5'
    pdf=supplier_pdf(purchase); reader=PdfReader(BytesIO(pdf))
    text=''.join(page.extract_text() for page in reader.pages)
    assert 'PRIVATE' not in text and '870' in text and '31' in text and '示例供应商' in text
    assert len(reader.pages)==1


def test_export_multipage_repeats_table_header():
    from app.services.supplier_purchase_view import supplier_pdf
    purchase=sample();purchase['lines']=purchase['lines']*70
    reader=PdfReader(BytesIO(supplier_pdf(purchase)))
    assert len(reader.pages)>1
    for page in reader.pages:
        assert '纸板长宽' in page.extract_text()


def test_authorized_endpoints_and_customer_scope_are_read_only(requisition_app):
    from app.models.supplier_requisition_order import SupplierRequisitionOrder, SupplierRequisitionOrderItem
    from app.models.order import Order
    from app.api import requisition
    from fastapi import HTTPException
    app, factory=requisition_app
    with factory() as db:
        db.get(Order,1).customer_po='CUSTOMER-001'
        order=SupplierRequisitionOrder(order_number='SRO-TEST-1',supplier_name='嘉林亿',requisition_qty=31,total_quantity=31)
        order.items=[SupplierRequisitionOrderItem(order_item_id=1,product_code='21301028',quantity=31,requisition_qty=31,
            report_length_mm=870,report_width_mm=970,material_code_snapshot='BC14C',layer_count_snapshot=5,flute_type_snapshot='AB')]
        db.add(order);db.commit();order_id=order.id
    paths=[f'/api/requisition/supplier-orders/{order_id}/{kind}' for kind in ['pdf','xlsx','internal-trace']]
    with TestClient(app) as client:
        for path in paths: assert client.get(path).status_code==401
        _login(client,'admin')
        with factory() as db:
            before='\n'.join(db.connection().connection.driver_connection.iterdump())
        for path in paths:
            response=client.get(path);assert response.status_code==200,response.text[:300]
        trace=client.get(paths[-1]).json()
        assert trace['sources'][0]['customer_po']=='CUSTOMER-001'
        assert trace['sources'][0]['inventory_records']==[]
        assert trace['sources'][0]['order_purpose_sheet_qty'] is None
        with factory() as db:
            after='\n'.join(db.connection().connection.driver_connection.iterdump())
        assert before==after
        original=requisition._require_supplier_order_customer_access
        def denied(*args): raise HTTPException(status_code=403,detail='outside customer scope')
        requisition._require_supplier_order_customer_access=denied
        try:
            for path in paths: assert client.get(path).status_code==403
        finally: requisition._require_supplier_order_customer_access=original


def test_inventory_trace_separates_raw_semi_finished_and_released(requisition_app):
    from datetime import date, datetime
    from app.models.warehouse_inventory import WarehouseLocation, InventoryLot, InventoryReservation, SemiFinishedInventoryDetail, FinishedGoodsInventoryDetail
    from app.services.supplier_purchase_view import internal_trace
    _, factory=requisition_app
    with factory() as db:
        place=WarehouseLocation(location_code='TEST-A',location_name='F货架L003 二层1格',warehouse_type='shared',warehouse_floor=3)
        db.add(place);db.flush()
        for i,kind in enumerate(['finished','net_sheet','raw_board'],1):
            lot=InventoryLot(lot_number=f'TEST-{i}',inventory_type='finished' if kind=='finished' else 'semi_finished',
                warehouse_location_id=place.id,unit='boxes' if kind=='finished' else 'sheets',source_type='manual',
                stock_date=date(2026,9,15),last_movement_at=datetime(2026,9,15),quantity_reserved=7,quantity_consumed=2)
            if kind=='finished':
                lot.finished_detail=FinishedGoodsInventoryDetail(owner_customer_id=1,product_id=1,inventory_code_snapshot='21301028',product_name_snapshot='箱')
            else:
                lot.semi_finished_detail=SemiFinishedInventoryDetail(material_code_snapshot='CCC',normalized_material_code='CCC',
                    layer_count=3,flute_type='B',board_length_mm=800,board_width_mm=600,sheet_type=kind)
            db.add(lot);db.flush()
            db.add(InventoryReservation(reservation_number=f'R-{i}',inventory_lot_id=lot.id,
                reservation_type='finished_order' if kind=='finished' else 'semi_order',order_id=1,order_item_id=1,
                reserved_stock_quantity=10,released_stock_quantity=1,consumed_stock_quantity=2,status='partial'))
            db.add(InventoryReservation(reservation_number=f'VOID-{i}',inventory_lot_id=lot.id,
                reservation_type='semi_order',order_id=1,order_item_id=1,reserved_stock_quantity=10,released_stock_quantity=10,status='released'))
        db.commit()
        data=internal_trace(db,{'lines':[{'source_items':[{'id':1,'order_item_id':1}]}]})
        records=data['sources'][0]['inventory_records']
        assert [row['kind'] for row in records]==['finished','semi','raw']
        assert all(row['pending_qty']==7 and row['consumed_qty']==2 for row in records)
        assert all('L003' in row['location'] and row['location_id']==place.id for row in records)
        assert all('cost' not in json.dumps(row) for row in records)
