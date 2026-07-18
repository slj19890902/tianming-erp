"""v0.20.3 公司信息维护及打印联动 — UAT tests"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import event
from sqlalchemy.orm import sessionmaker


# ─────────────────────────────────────────────────────────────
# Fixture
# ─────────────────────────────────────────────────────────────

@pytest.fixture()
def company_app(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    from app.api.auth import router as auth_router
    from app.api.deliveries import router as deliveries_router
    from app.api.finance import router as finance_router
    from app.api.system import router as system_router
    from app.api.deps import get_db
    from app.core.database import create_sqlite_engine
    from app.core.security import hash_password
    from app.models import Base
    from app.models.company_config import CompanyConfig
    from app.models.customer import Customer
    from app.models.delivery import Delivery, DeliveryItem
    from app.models.finance import (
        ReturnReceipt, ReturnReceiptItem, Statement, StatementItem,
    )
    from app.models.material import Material
    from app.models.order import Order, OrderItem
    from app.models.product import Product
    from app.models.user import User

    monkeypatch.setenv("ERP_DRAWING_DIR", str(tmp_path / "uploads"))
    engine = create_sqlite_engine(tmp_path / "v0203.sqlite3")
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)

    with session_factory() as session:
        users = [
            User(
                username=role,
                password_hash=hash_password("RolePass123!"),
                role=role,
                real_name=role,
                display_name=role,
                must_change_password=False,
            )
            for role in ("admin", "finance", "sales")
        ]
        customer = Customer(
            customer_number=1,
            customer_code="CUST",
            name="苏州正常客户有限公司",
            payment_term_days=30,
            credit_limit=0,
            delivery_method="配送",
        )
        material = Material(
            code="WCX1",
            layer_count=5,
            flute_type="AB",
            basis_weight_description="高强",
        )
        session.add_all([*users, customer, material])
        session.flush()

        product = Product(
            customer_id=customer.id,
            product_code="P001",
            customer_material_code="P001",
            product_name="外箱",
            material_id=material.id,
            box_category="normal",
            cost_unit_price=Decimal("3.00"),
        )
        session.add(product)
        session.flush()

        order = Order(
            order_number="PO-2026-001",
            customer_id=customer.id,
            order_date=date(2026, 6, 1),
            delivery_date=date(2026, 6, 14),
            status="pending_production",
            payment_status="unpaid",
            total_amount=Decimal("300"),
        )
        session.add(order)
        session.flush()

        item = OrderItem(
            order_id=order.id,
            product_id=product.id,
            quantity=100,
            delivered_quantity=80,
            unit_price=Decimal("3.00"),
            subtotal=Decimal("300"),
            material_status="pending",
            requisition_status="已报料",
            snapshot_product_code="P001",
            snapshot_product_name="外箱",
            snapshot_material="WCX1",
        )
        session.add(item)
        session.flush()

        delivery = Delivery(
            delivery_number="DH-2026-001",
            customer_id=customer.id,
            delivery_date=date(2026, 6, 14),
            status="dispatched",
            total_quantity=80,
            dispatched_at=datetime(2026, 6, 14, 9, 0),
        )
        session.add(delivery)
        session.flush()

        d_item = DeliveryItem(
            delivery_id=delivery.id,
            order_item_id=item.id,
            delivered_quantity=80,
        )
        session.add(d_item)
        session.flush()

        receipt = ReturnReceipt(
            delivery_id=delivery.id,
            actual_received_date=date(2026, 6, 14),
            signed_by="王经理",
            status="confirmed",
        )
        session.add(receipt)
        session.flush()

        r_item = ReturnReceiptItem(
            return_receipt_id=receipt.id,
            delivery_item_id=d_item.id,
            actual_received_quantity=80,
            difference_reason=None,
        )
        session.add(r_item)
        session.flush()

        statement = Statement(
            statement_number="ST-202606-001",
            customer_id=customer.id,
            statement_month="2026-06",
            total_receivable=Decimal("240.00"),
            invoiced_amount=Decimal("0"),
            settled_amount=Decimal("0"),
            status="unsettled",
        )
        session.add(statement)
        session.flush()

        s_item = StatementItem(
            statement_id=statement.id,
            return_receipt_item_id=r_item.id,
            unit_price_snapshot=Decimal("3.00"),
            unit_cost_snapshot=Decimal("2.00"),
            actual_received_quantity=80,
            receivable_amount=Decimal("240.00"),
            gross_profit_amount=Decimal("80.00"),
        )
        session.add(s_item)
        session.commit()

    app = FastAPI()

    def override_get_db():
        with session_factory() as db:
            yield db

    from app.api.deps import get_db as _get_db
    app.dependency_overrides[_get_db] = override_get_db
    app.include_router(auth_router, prefix="/api/auth")
    app.include_router(system_router, prefix="/api/system")
    app.include_router(finance_router, prefix="/api/finance")
    app.include_router(deliveries_router, prefix="/api/deliveries")

    return app, session_factory


def _login(client: TestClient, role: str = "admin") -> None:
    resp = client.post(
        "/api/auth/login",
        json={"username": role, "password": "RolePass123!"},
    )
    assert resp.status_code == 200


# ─────────────────────────────────────────────────────────────
# 公司信息 API
# ─────────────────────────────────────────────────────────────

def test_get_company_returns_empty_on_fresh_db(company_app):
    app, _, = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        resp = client.get("/api/system/company")
    assert resp.status_code == 200
    assert resp.json()["company_name"] == ""


def test_get_company_requires_admin(company_app):
    """公司信息包含税号及银行账号，只允许管理员读取。"""
    app, _ = company_app
    with TestClient(app) as client:
        resp = client.get("/api/system/company")
    assert resp.status_code == 401


def test_get_company_does_not_write_when_config_row_is_missing(company_app):
    """只读 GET 不得为了补空配置行而执行 INSERT。"""
    app, session_factory = company_app
    engine = session_factory.kw["bind"]
    write_statements: list[str] = []

    with TestClient(app) as client:
        _login(client, "admin")

        def record_writes(
            _conn, _cursor, statement, _parameters, _context, _executemany
        ):
            normalized = statement.lstrip().upper()
            if normalized.startswith(("INSERT", "UPDATE", "DELETE")):
                write_statements.append(normalized)

        event.listen(engine, "before_cursor_execute", record_writes)
        try:
            resp = client.get("/api/system/company")
        finally:
            event.remove(engine, "before_cursor_execute", record_writes)

    assert resp.status_code == 200
    assert write_statements == []


def test_put_company_rejects_blank_company_name(company_app):
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        resp = client.put(
            "/api/system/company",
            json={"company_name": "   "},
        )
    assert resp.status_code == 422


def test_put_company_requires_admin(company_app):
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "finance")
        resp = client.put(
            "/api/system/company",
            json={"company_name": "苏州天明包装有限公司"},
        )
    assert resp.status_code == 403


def test_put_company_saves_all_fields(company_app):
    app, session_factory = company_app
    payload = {
        "company_name": "苏州天明包装有限公司",
        "short_name": "天明包装",
        "address": "苏州市相城区渭塘镇",
        "phone": "0512-12345678",
        "fax": "0512-87654321",
        "tax_number": "9132059412345678XY",
        "bank_name": "工商银行苏州支行",
        "bank_account": "6202001234567890",
        "contact_person": "张总",
        "contact_phone": "13800138000",
    }
    with TestClient(app) as client:
        _login(client, "admin")
        resp = client.put("/api/system/company", json=payload)
    assert resp.status_code == 200
    data = resp.json()
    assert data["company_name"] == "苏州天明包装有限公司"
    assert data["short_name"] == "天明包装"
    assert data["tax_number"] == "9132059412345678XY"
    assert data["bank_name"] == "工商银行苏州支行"
    assert data["updated_at"] is not None


def test_put_company_persists_and_get_reflects_update(company_app):
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        client.put(
            "/api/system/company",
            json={"company_name": "苏州天明包装有限公司", "address": "苏州市相城区"},
        )
        resp = client.get("/api/system/company")
    assert resp.json()["company_name"] == "苏州天明包装有限公司"
    assert resp.json()["address"] == "苏州市相城区"


def test_put_company_strips_whitespace(company_app):
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        resp = client.put(
            "/api/system/company",
            json={"company_name": "  苏州天明  ", "phone": "  0512-123  "},
        )
    assert resp.json()["company_name"] == "苏州天明"
    assert resp.json()["phone"] == "0512-123"


def test_put_company_writes_audit_log(company_app):
    from app.models.audit import OperationLog
    app, session_factory = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        client.put(
            "/api/system/company",
            json={"company_name": "天明包装"},
        )
    with session_factory() as session:
        log = (
            session.query(OperationLog)
            .filter_by(action="UPDATE_COMPANY_CONFIG")
            .first()
        )
    assert log is not None
    assert "天明包装" in log.details


# ─────────────────────────────────────────────────────────────
# 对账单 Excel 联动
# ─────────────────────────────────────────────────────────────

def test_statement_excel_row3_empty_when_no_company(company_app):
    """公司信息未设置时第 3 行为空字符串，不崩溃。"""
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "finance")
        resp = client.get("/api/finance/statements/1/export")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.content))
    sheet = wb.active
    assert sheet["A3"].value == "" or sheet["A3"].value is None


def test_statement_excel_row3_contains_company_name(company_app):
    """公司信息设置后第 3 行包含供方名称。"""
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        client.put(
            "/api/system/company",
            json={
                "company_name": "苏州天明包装有限公司",
                "address": "苏州市相城区渭塘镇",
                "tax_number": "9132059412345678XY",
            },
        )
        _login(client, "finance")
        resp = client.get("/api/finance/statements/1/export")
    wb = load_workbook(BytesIO(resp.content))
    sheet = wb.active
    row3 = sheet["A3"].value or ""
    assert "苏州天明包装有限公司" in row3
    assert "苏州市相城区渭塘镇" in row3
    assert "9132059412345678XY" in row3


def test_statement_excel_header_row_is_row5(company_app):
    """设置公司信息后表头仍在第 5 行。"""
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        client.put("/api/system/company", json={"company_name": "天明"})
        _login(client, "finance")
        resp = client.get("/api/finance/statements/1/export")
    wb = load_workbook(BytesIO(resp.content))
    sheet = wb.active
    headers = [c.value for c in sheet[5]]
    assert "客户名称" in headers
    assert "实际签收数量" in headers


# ─────────────────────────────────────────────────────────────
# 送货单打印联动
# ─────────────────────────────────────────────────────────────

def test_print_endpoint_returns_sender_field(company_app):
    """/api/deliveries/{id}/print 返回 sender 字段。"""
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        resp = client.get("/api/deliveries/1/print")
    assert resp.status_code == 200
    data = resp.json()
    assert "sender" in data
    assert "company_name" in data["sender"]


def test_print_endpoint_sender_empty_when_no_company(company_app):
    """公司信息未设置时 sender.company_name 为空字符串。"""
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        resp = client.get("/api/deliveries/1/print")
    assert resp.json()["sender"]["company_name"] == ""


def test_print_endpoint_sender_reflects_company_info(company_app):
    """设置公司信息后打印接口返回最新数据。"""
    app, _ = company_app
    with TestClient(app) as client:
        _login(client, "admin")
        client.put(
            "/api/system/company",
            json={
                "company_name": "苏州天明包装有限公司",
                "tax_number": "9132059412345678XY",
                "bank_name": "工商银行苏州支行",
            },
        )
        resp = client.get("/api/deliveries/1/print")
    sender = resp.json()["sender"]
    assert sender["company_name"] == "苏州天明包装有限公司"
    assert sender["tax_number"] == "9132059412345678XY"
    assert sender["bank_name"] == "工商银行苏州支行"


def test_delivery_print_page_uses_sender_from_api():
    project_root = Path(__file__).resolve().parents[1]
    source = (project_root / "static" / "delivery-print.html").read_text(
        encoding="utf-8"
    )
    assert 'data-field="senderCompanyName"' in source
    assert 'data-field="senderContact"' in source
    assert "const sender = data.sender || {}" in source
    assert "sender.company_name" in source
    assert "sender.address" in source
    assert "sender.phone" in source
    assert '<h1 class="company">苏州天明包装有限公司</h1>' not in source
